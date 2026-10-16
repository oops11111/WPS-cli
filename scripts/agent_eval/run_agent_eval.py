"""Run fixed tasks through the claude CLI against the wps-agent MCP server, then judge the workspace state.

The agent is the real `claude -p` process. Each task runs in its own workspace, which is also the MCP
server's working directory, so the server's state files live there. A task can list setup steps that
prepare the workspace before the agent starts (register, back up, change or hold the file, write an HTML
table). The judge reads the workspace state (documents, operations, backups, file hashes, output files)
and the agent's tool calls from the CLI transcript. It does not read the agent's reply.

Usage:
    python scripts/agent_eval/run_agent_eval.py [--tasks config/agent_eval_tasks.json] [--out artifacts/agent-eval] [--only ID] [--model MODEL] [--max-turns N] [--timeout SECONDS]

Exit code: 0 when every task passes, 1 when any fails, 2 when the claude CLI is not available.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import os
import shutil
import subprocess
import sys
import time
from collections.abc import Callable
from pathlib import Path
from typing import Any

REPO_ROOT = Path(__file__).resolve().parents[2]
SRC_DIR = REPO_ROOT / "src"
MCP_SERVER_NAME = "wps-agent"
ALLOWED_TOOLS = f"mcp__{MCP_SERVER_NAME}__*"

sys.path.insert(0, str(SRC_DIR))
from wps_ai_agent_cli.backups import create_backup, list_backups  # noqa: E402
from wps_ai_agent_cli.open_document_hints import OPEN_DOCUMENT_ALTERNATIVES  # noqa: E402
from wps_ai_agent_cli.operations import get_operation  # noqa: E402
from wps_ai_agent_cli.sessions import (  # noqa: E402
    get_document,
    register_document,
    stable_document_id,
)

TOOL_PREFIX = f"mcp__{MCP_SERVER_NAME}__wps_agent_"
# File-based write tools: wrong on a document open in WPS, where the selection tools apply (P9-002).
FILE_BASED_WRITE_TOOLS = frozenset(command.replace("-", "_") for command in OPEN_DOCUMENT_ALTERNATIVES)
HTML_TABLE = (
    "<html><body><h2>Sales</h2><table>"
    "<tr><th>Region</th><th>Q1</th></tr><tr><td>North</td><td>120</td></tr><tr><td>South</td><td>95</td></tr>"
    "</table></body></html>"
)


def file_sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


WRITER_FIXTURE_PARAGRAPHS = ("Quarterly report", "Status: DRAFT", "Owner: Finance")
PRESENTATION_FIXTURE_TEXTS = ("Quarterly review", "Status: DRAFT")
# table_report.docx: one heading paragraph and one table.
WRITER_TABLE_HEADING = "Regional status"
WRITER_TABLE_ROWS = (("Region", "Q1", "Status"), ("North", "120", "DRAFT"), ("South", "95", "DRAFT"))
# status.xlsx: the same table on one sheet, for the spreadsheet selection tasks (P7-003).
SPREADSHEET_STATUS_SHEET = "Status"
SPREADSHEET_STATUS_ROWS = (("Region", "Q1", "Status"), ("North", 120, "DRAFT"), ("South", 95, "DRAFT"))
# Chinese fixtures for the Chinese-content tasks (P8-004): a Chinese file name, Chinese text and, for the
# workbook, a Chinese sheet name. Same layout as report.docx and status.xlsx.
WRITER_ZH_FIXTURE = "季度报告.docx"
WRITER_ZH_PARAGRAPHS = ("季度报告", "状态：草稿", "负责人：财务部")
# Characters GBK cannot encode, for the text export task (P9-004): an emoji and a CJK character outside the BMP.
WRITER_BEYOND_GBK_FIXTURE = "核对清单.docx"
WRITER_BEYOND_GBK_PARAGRAPHS = ("核对清单", "✅ 已核对：𠮷野家", "负责人：财务部")
WRITER_PARAGRAPHS = {WRITER_ZH_FIXTURE: WRITER_ZH_PARAGRAPHS, WRITER_BEYOND_GBK_FIXTURE: WRITER_BEYOND_GBK_PARAGRAPHS}
SPREADSHEET_ZH_FIXTURE = "状态表.xlsx"
SPREADSHEET_ZH_SHEET = "汇总"
SPREADSHEET_ZH_ROWS = (("地区", "一季度", "状态"), ("华北", 120, "草稿"), ("华南", 95, "草稿"))


def write_fixture(workspace: Path, name: str) -> str:
    """Write the task's fixture, chosen by its name, else by its suffix: .xlsx, .docx or .pptx."""
    path = workspace / name
    suffix = path.suffix.lower()
    tables = {
        "status.xlsx": (SPREADSHEET_STATUS_SHEET, SPREADSHEET_STATUS_ROWS),
        SPREADSHEET_ZH_FIXTURE: (SPREADSHEET_ZH_SHEET, SPREADSHEET_ZH_ROWS),
    }
    if suffix == ".xlsx" and path.name in tables:
        from openpyxl import Workbook

        title, rows = tables[path.name]
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = title
        for row in rows:
            sheet.append(list(row))
        workbook.save(path)
    elif suffix == ".xlsx":
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active["A1"] = "seed"
        workbook.save(path)
    elif suffix == ".docx" and path.name == "table_report.docx":
        from docx import Document

        document = Document()
        document.add_paragraph(WRITER_TABLE_HEADING)
        table = document.add_table(rows=len(WRITER_TABLE_ROWS), cols=len(WRITER_TABLE_ROWS[0]))
        for row, values in zip(table.rows, WRITER_TABLE_ROWS, strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
        document.save(path)
    elif suffix == ".docx":
        from docx import Document

        document = Document()
        for text in WRITER_PARAGRAPHS.get(path.name, WRITER_FIXTURE_PARAGRAPHS):
            document.add_paragraph(text)
        document.save(path)
    elif suffix == ".pptx":
        from pptx import Presentation
        from pptx.util import Inches

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        for index, text in enumerate(PRESENTATION_FIXTURE_TEXTS):
            box = slide.shapes.add_textbox(Inches(1), Inches(1 + index), Inches(8), Inches(1))
            box.text_frame.text = text
        presentation.save(path)
    else:
        raise ValueError(f"unsupported fixture type {suffix}")
    return file_sha256(path)


def text_file_encoding(path: Path) -> str:
    """utf-16 or utf-8-sig by byte order mark, utf-8 when it decodes strictly, else other (a code page)."""
    raw = path.read_bytes()
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return "utf-16"
    if raw[:3] == b"\xef\xbb\xbf":
        return "utf-8-sig"
    try:
        raw.decode("utf-8")
    except UnicodeDecodeError:
        return "other"
    return "utf-8"


def document_text(path: Path) -> str:
    """All text in a .docx (paragraphs and tables), a .pptx (every shape) or a .txt, joined by newlines.

    A .txt from WPS is UTF-16 with a byte order mark (Unicode text) or, as a fallback, the system code page:
    GBK on Chinese Windows, which is what Chinese text that is not UTF-8 is read as (P8-004).
    """
    suffix = path.suffix.lower()
    if suffix == ".txt":
        raw = path.read_bytes()
        if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
            return raw.decode("utf-16")
        try:
            return raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            return raw.decode("gbk", errors="replace")
    if suffix == ".docx":
        from docx import Document

        document = Document(path)
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            parts += [cell.text for row in table.rows for cell in row.cells]
        return "\n".join(parts)
    if suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation(path)
        return "\n".join(
            shape.text_frame.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame
        )
    raise ValueError(f"unsupported document type {suffix}")


def write_mcp_config(workspace: Path) -> Path:
    config = {
        "mcpServers": {
            MCP_SERVER_NAME: {
                "command": sys.executable,
                "args": ["-m", "wps_ai_agent_cli", "mcp-server"],
                "cwd": str(workspace),
                "env": {"PYTHONPATH": str(SRC_DIR)},
            }
        }
    }
    path = workspace.parent / f"{workspace.name}.mcp.json"
    path.write_text(json.dumps(config, indent=2), encoding="utf-8")
    return path


def hold_file(path: Path) -> tuple[Callable[[], None] | None, str | None]:
    """Hold the file so the server cannot read it, as WPS does with an open document.

    Returns (release, None), or (None, reason) when this machine cannot hold a file.
    """
    if os.name == "nt":
        import ctypes
        from ctypes import wintypes

        kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel32.CreateFileW.argtypes = [
            wintypes.LPCWSTR,
            wintypes.DWORD,
            wintypes.DWORD,
            wintypes.LPVOID,
            wintypes.DWORD,
            wintypes.DWORD,
            wintypes.HANDLE,
        ]
        kernel32.CreateFileW.restype = wintypes.HANDLE
        # GENERIC_READ, no sharing, OPEN_EXISTING, FILE_ATTRIBUTE_NORMAL: every other open is refused.
        handle = kernel32.CreateFileW(str(path), 0x80000000, 0, None, 3, 0x80, None)
        if handle == wintypes.HANDLE(-1).value:
            return None, f"CreateFileW failed with error {ctypes.get_last_error()}"
        return (lambda: kernel32.CloseHandle(handle)), None
    if hasattr(os, "geteuid") and os.geteuid() == 0:
        return None, "running as root, which ignores file permissions, so the file cannot be held"
    mode = path.stat().st_mode
    path.chmod(0)
    return (lambda: path.chmod(mode)), None


# How open_in_wps starts WPS. Both run the open command registered for the file's extension, the command a
# double-click runs; "automation" adds /Automation to it. Neither changes the system's open command. The shell
# (os.startfile) is not used: on the P6-008 machine it showed the "Open with" chooser for .xlsx and .pptx
# instead of running that command.
WPS_LAUNCH_MODES = ("default", "automation")
WPS_LAUNCH = "default"
WPS_DOCUMENT_PROCESSES = ("wps", "et", "wpp")

_WPS_DOCUMENT_SCRIPT = """
$ErrorActionPreference = 'Stop'
$params = @'
__PARAMS_JSON__
'@ | ConvertFrom-Json
__ATTACH__
function Out-Result($value) { $value | ConvertTo-Json -Depth 5; exit 0 }
$app = Get-RunningWpsApp @($params.prog_ids)
if ($null -eq $app) { Out-Result ([pscustomobject]@{ ok = $true; status = (Get-WpsMissingStatus @($params.prog_ids)) }) }
$collectionName = @{ writer = 'Documents'; spreadsheets = 'Workbooks'; presentation = 'Presentations' }[[string]$params.component]
if ($params.action -eq 'new_untitled') {
  # A new document that has never been saved, with some text, for the Save As tasks (P7-003).
  $text = [string]$params.text
  $new = $app.($collectionName).Add()
  if ($params.component -eq 'writer') {
    $new.Content.Text = $text
  } elseif ($params.component -eq 'spreadsheets') {
    $new.Worksheets.Item(1).Range('A1').Value2 = $text
  } else {
    $slide = $new.Slides.Add(1, 12)
    $box = $slide.Shapes.AddTextbox(1, 50, 50, 600, 60)
    $box.TextFrame.TextRange.Text = $text
  }
  Out-Result ([pscustomobject]@{
    ok = $true; status = 'ok'; name = [string]$new.Name; full_name = [string]$new.FullName; saved = [bool]$new.Saved
  })
}
if ($params.action -eq 'close_untitled') {
  $collection = $app.($collectionName)
  for ($i = [int]$collection.Count; $i -ge 1; $i--) {
    $item = $collection.Item($i)
    if ([string]$item.Name -ne [string]$params.name) { continue }
    $onDisk = $false
    try { $onDisk = [System.IO.File]::Exists([string]$item.FullName) } catch { }
    if ($onDisk) { continue }
    if ($params.component -eq 'writer') { $item.Close(0) }
    elseif ($params.component -eq 'spreadsheets') { $item.Close($false) }
    else { $item.Saved = -1; $item.Close() }
    Out-Result ([pscustomobject]@{ ok = $true; status = 'closed' })
  }
  Out-Result ([pscustomobject]@{ ok = $true; status = 'not_found' })
}
if ($params.component -eq 'spreadsheets') {
  $doc = Find-OpenWorkbook $app $params.path
  if ($null -eq $doc) { Out-Result ([pscustomobject]@{ ok = $true; status = 'not_open' }) }
  if ($params.action -eq 'close_discard') {
    # Close($false) failed with RPC_E_SERVERFAULT on WPS 12.8 (P7-007): mark it saved and close with no argument.
    $doc.Saved = $true
    try { [void]$doc.Close() } catch { [void]$doc.Close($false) }
    Out-Result ([pscustomobject]@{ ok = $true; status = 'closed' })
  }
  if ($params.action -ne 'select') { Out-Result ([pscustomobject]@{ ok = $true; status = 'unsupported' }) }
  # Select a cell range through COM and read its address back (P7-003).
  try { $doc.Activate() } catch { }
  if ([string]::IsNullOrEmpty([string]$params.sheet)) { $sheet = $doc.Worksheets.Item(1) }
  else { $sheet = $doc.Worksheets.Item([string]$params.sheet) }
  $sheet.Activate()
  # [void]: WPS Spreadsheets returns True from Range.Select(), and a bare call puts it on stdout
  # ahead of the JSON, which _run_script then reads as a failure. Measured on WPS 12.8.2.18913;
  # Activate() above puts nothing there.
  [void]$sheet.Range([string]$params.range).Select()
  $selection = $app.ActiveWindow.Selection
  Out-Result ([pscustomobject]@{
    ok = $true; status = 'ok'; address = [string]$selection.Address($false, $false)
    sheet = [string]$selection.Worksheet.Name; saved = [bool]$doc.Saved
  })
}
if ($params.component -eq 'presentation') {
  $doc = Find-OpenPresentation $app $params.path
  if ($null -eq $doc) { Out-Result ([pscustomobject]@{ ok = $true; status = 'not_open' }) }
  if ($params.action -eq 'close_discard') {
    $doc.Saved = -1
    $doc.Close()
    Out-Result ([pscustomobject]@{ ok = $true; status = 'closed' })
  }
  if ($params.action -ne 'select') { Out-Result ([pscustomobject]@{ ok = $true; status = 'unsupported' }) }
  # TextRange.Find fails with an unexpected HRESULT on WPS 12.8 (P6-008 rerun), so search the text in
  # PowerShell and select the characters by position (Characters is 1-based).
  $wanted = [string]$params.text
  for ($s = 1; $s -le [int]$doc.Slides.Count; $s++) {
    $slide = $doc.Slides.Item($s)
    for ($h = 1; $h -le [int]$slide.Shapes.Count; $h++) {
      $shape = $slide.Shapes.Item($h)
      if (-not [bool]$shape.HasTextFrame) { continue }
      $range = $shape.TextFrame.TextRange
      $at = ([string]$range.Text).IndexOf($wanted, [System.StringComparison]::Ordinal)
      if ($at -lt 0) { continue }
      try { $app.ActiveWindow.View.GotoSlide($s) } catch { }
      # With collapse 'end', a caret right after the text (P7-004 insert task).
      if ($params.collapse -eq 'end') { $range.Characters($at + 1 + $wanted.Length, 0).Select() }
      else { $range.Characters($at + 1, $wanted.Length).Select() }
      $selection = $app.ActiveWindow.Selection
      $selectionType = $null
      $start = $null
      $length = $null
      try { $selectionType = [int]$selection.Type } catch { }
      try { $start = [int]$selection.TextRange.Start; $length = [int]$selection.TextRange.Length } catch { }
      $end = $null
      if ($null -ne $start -and $null -ne $length) { $end = $start + $length }
      Out-Result ([pscustomobject]@{
        ok = $true; status = 'ok'; selection_text = [string]$selection.TextRange.Text
        selection_type = $selectionType; start = $start; end = $end; saved = [bool]$doc.Saved
      })
    }
  }
  Out-Result ([pscustomobject]@{ ok = $true; status = 'text_not_found' })
}
$doc = Find-OpenWriterDocument $app $params.path
if ($null -eq $doc) { Out-Result ([pscustomobject]@{ ok = $true; status = 'not_open' }) }
if ($params.action -eq 'close_discard') {
  $doc.Close(0)
  Out-Result ([pscustomobject]@{ ok = $true; status = 'closed' })
}
if ($params.action -eq 'append_unsaved') {
  $range = $doc.Content
  $range.InsertParagraphAfter()
  $range.InsertAfter([string]$params.text)
  Out-Result ([pscustomobject]@{ ok = $true; status = 'ok'; saved = [bool]$doc.Saved })
}
$range = $doc.Content
$found = $range.Find.Execute([string]$params.text)
if (-not $found) { Out-Result ([pscustomobject]@{ ok = $true; status = 'text_not_found' }) }
if ($params.collapse -eq 'end') { $range.Collapse(0) }
$range.Select()
$selection = $doc.ActiveWindow.Selection
Out-Result ([pscustomobject]@{
  ok = $true; status = 'ok'; selection_text = [string]$selection.Text
  start = [int]$selection.Start; end = [int]$selection.End; saved = [bool]$doc.Saved
})
"""


def wps_process_ids() -> set[int]:
    """Ids of the WPS document processes (wps, et, wpp) running now."""
    from wps_ai_agent_cli.powershell_runner import run_powershell_command

    names = ",".join(WPS_DOCUMENT_PROCESSES)
    completed = run_powershell_command(
        f"Get-Process -Name {names} -ErrorAction SilentlyContinue | ForEach-Object {{ $_.Id }}", timeout_seconds=30
    )
    return {int(token) for token in completed.stdout.split() if token.strip().isdigit()}


COMPONENT_PROCESSES = {"writer": "wps", "spreadsheets": "et", "presentation": "wpp"}


def component_process_ids(component: str) -> set[int]:
    """Ids of the processes of one WPS component (wps for Writer, et for Spreadsheets, wpp for Presentation)."""
    from wps_ai_agent_cli.powershell_runner import run_powershell_command

    completed = run_powershell_command(
        f"Get-Process -Name {COMPONENT_PROCESSES[component]} -ErrorAction SilentlyContinue | ForEach-Object {{ $_.Id }}",
        timeout_seconds=30,
    )
    return {int(token) for token in completed.stdout.split() if token.strip().isdigit()}


_QUIT_SCRIPT = """
$app = Get-RunningWpsApp @($params.prog_ids)
if ($null -eq $app) {
  [pscustomobject]@{ status = 'unreachable' } | ConvertTo-Json -Depth 3
  exit 0
}
try { $app.DisplayAlerts = $false } catch { }
$app.Quit()
[pscustomobject]@{ status = 'quit' } | ConvertTo-Json -Depth 3
"""


def quit_wps_app(component: str) -> str:
    """Call Quit() on the component's running WPS: 'quit', 'unreachable', or why it failed.

    A WPS started by opening a file registers nothing once its last document is closed (P8-005), so it is often
    unreachable here; close_main_windows is the other normal way to end it.
    """
    from wps_ai_agent_cli.capabilities import WPS_COMPONENTS
    from wps_ai_agent_cli.open_document_scripts import _SCRIPT_HEAD
    from wps_ai_agent_cli.open_documents import _run_script
    from wps_ai_agent_cli.wps_script_snippets import ATTACH_RUNNING_INSTANCE

    script = _SCRIPT_HEAD + ATTACH_RUNNING_INSTANCE + _QUIT_SCRIPT
    result = _run_script(script, {"prog_ids": list(WPS_COMPONENTS[component])}, timeout_seconds=60)
    if not result["ok"]:
        return "failed: " + "; ".join(str(error.get("message")) for error in result["errors"])
    return str((result["data"] or {}).get("status") or "no output")


def close_main_windows(ids: set[int]) -> None:
    """Close the processes' main windows, as a user closing WPS would; WPS then exits normally."""
    from wps_ai_agent_cli.powershell_runner import run_powershell_command

    listed = ",".join(map(str, sorted(ids)))
    run_powershell_command(
        f"Get-Process -Id {listed} -ErrorAction SilentlyContinue | ForEach-Object {{ [void]$_.CloseMainWindow() }}",
        timeout_seconds=30,
    )


def end_wps_normally(component: str, ids: set[int], timeout: float = 20.0) -> tuple[list[int], str]:
    """Let the component's WPS exit by itself: Quit(), else close its main windows. Returns (exited ids, how).

    A WPS ended with Stop-Process while it still had a document open treats the next start as a recovery and
    reopens that document; after Phase 7 and Phase 8 runs, every Windows run had to start with closing those by
    hand (P9-005). A WPS that exits normally reopens nothing.
    """
    how = quit_wps_app(component)
    if how != "quit":
        close_main_windows(ids)
        how = "main window closed" if how == "unreachable" else f"main window closed after Quit() {how}"
    else:
        how = "Quit()"
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline and component_process_ids(component) & ids:
        time.sleep(1)
    return sorted(ids - component_process_ids(component)), how


def stop_started_wps(component: str, started: set[int], rounds: int = 3, settle: float = 3.0) -> str:
    """End the component's WPS processes after a run that started one, while WPS lists no document for it.

    It first lets WPS exit normally (end_wps_normally) and stops only the processes still running after that.

    Not only the processes the run started: when one is stopped, WPS can start another in its place. In P7-007
    stopping et.exe made WPS start "et.exe recover1 -cpid=<the stopped id>", which reopened status.xlsx in a
    recovery workspace without registering for automation; every later launch handed the file to it, so the
    next runs could not reach their workbook. So every process of the component is stopped, in up to `rounds`
    rounds, until none comes back. This is safe because the Windows check's prerequisites require that no WPS
    process runs when a check starts: any process of the component now belongs to this run.
    """
    if not started:
        return "none started"
    from wps_ai_agent_cli.open_documents import list_open_documents
    from wps_ai_agent_cli.powershell_runner import run_powershell_command

    _ok, data, _errors = list_open_documents(component)
    if data.get("documents"):
        # Name them: P7-007 left processes running and the record could not say which document kept them.
        names = ", ".join(str(item.get("name") or item.get("full_name") or "?") for item in data["documents"])
        return f"left running: WPS still lists {names}"
    exited: list[int] = []
    how = ""
    running = component_process_ids(component)
    if running:
        exited, how = end_wps_normally(component, running)
    stopped: list[int] = []
    for _round in range(rounds):
        ids = component_process_ids(component)
        if not ids:
            break
        stopped += sorted(ids)
        listed = ",".join(map(str, sorted(ids)))
        run_powershell_command(f"Stop-Process -Id {listed} -Force -ErrorAction SilentlyContinue", timeout_seconds=30)
        deadline = time.monotonic() + 20
        while time.monotonic() < deadline and component_process_ids(component) & ids:
            time.sleep(1)
        # Give WPS the moment it takes to start a replacement, so the next round sees it.
        time.sleep(settle)
    left = sorted(component_process_ids(component))
    parts = []
    if exited:
        parts.append(f"exited {','.join(map(str, exited))} ({how})")
    if stopped:
        tried = f" (still running after {how})" if how else ""
        parts.append(f"stopped {','.join(map(str, stopped))}{tried}")
    summary = "; ".join(parts) or "none running"
    return summary + (f"; still running after {rounds} rounds: {','.join(map(str, left))}" if left else "")


def registered_open_command(extension: str) -> str:
    """The open command registered for an extension, for example: "C:\\...\\wps.exe" "%1"."""
    import ctypes
    from ctypes import wintypes

    assoc_str_command = 1
    size = wintypes.DWORD(2048)
    buffer = ctypes.create_unicode_buffer(size.value)
    result = ctypes.windll.shlwapi.AssocQueryStringW(
        0, assoc_str_command, extension, "open", buffer, ctypes.byref(size)
    )
    if result != 0 or not buffer.value:
        raise RuntimeError(f"no open command is registered for {extension} (HRESULT {result & 0xFFFFFFFF:#010x})")
    return buffer.value


def launch_command_line(template: str, path: Path, automation: bool) -> str:
    """Turn a registered open command into the command line to run for the file.

    %1 (or %L) becomes the file, other %-placeholders are dropped, and with automation /Automation goes right
    after the program, as in: wps.exe /Automation "<file>".
    """
    template = template.strip()
    if template.startswith('"'):
        end = template.index('"', 1)
        program, rest = template[: end + 1], template[end + 1 :]
    else:
        program, _, rest = template.partition(" ")
        rest = " " + rest if rest else ""
    file = str(path)
    if "%1" in rest or "%L" in rest or "%l" in rest:
        for token in ("%1", "%L", "%l"):
            rest = rest.replace(token, file)
    else:
        rest += f' "{file}"'
    rest = " ".join(word for word in rest.split(" ") if not (word.startswith("%") or word.startswith('"%')))
    command = program + (" /Automation" if automation else "") + (" " + rest.strip() if rest.strip() else "")
    return command


def start_registered_command(path: Path, automation: bool) -> str:
    """Run the file's registered open command, as a double-click does; returns the command line."""
    command = launch_command_line(registered_open_command(path.suffix), path.resolve(), automation)
    flags = getattr(subprocess, "DETACHED_PROCESS", 0) | getattr(subprocess, "CREATE_NEW_PROCESS_GROUP", 0)
    subprocess.Popen(command, creationflags=flags, close_fds=True)  # noqa: S603 - the registered WPS command
    return command


def wps_document_action(component: str, path: Path, action: str, **params: Any) -> dict[str, Any]:
    """Select text in, or add unsaved text to, a document open in WPS. Returns the script's JSON."""
    from wps_ai_agent_cli.capabilities import WPS_COMPONENTS
    from wps_ai_agent_cli.open_documents import _run_script
    from wps_ai_agent_cli.wps_script_snippets import ATTACH_RUNNING_INSTANCE

    script = _WPS_DOCUMENT_SCRIPT.replace("__ATTACH__", ATTACH_RUNNING_INSTANCE)
    result = _run_script(
        script,
        {
            "component": component,
            "prog_ids": list(WPS_COMPONENTS[component]),
            "path": str(path.resolve()),
            "action": action,
            **params,
        },
        timeout_seconds=60,
    )
    if not result["ok"]:
        return {"ok": False, "status": "script_failed", "errors": result["errors"]}
    return result["data"] or {"ok": False, "status": "no_output"}


def open_in_wps(
    path: Path, component: str, timeout: float = 45.0, poll: float = 2.0
) -> tuple[Callable[[], None] | None, str | None]:
    """Open the file in WPS (see WPS_LAUNCH) and wait until WPS lists it.

    Returns (release, None) or (None, reason) when this machine cannot do it. Release closes the document
    without saving, and only if WPS reports no unsaved changes; then it ends the WPS processes that were
    started for it, if WPS lists no other document, so that the next run starts from the same state.
    """
    if os.name != "nt":
        return None, "opening a document in WPS needs Windows"
    from wps_ai_agent_cli.open_documents import list_open_documents

    before = wps_process_ids()
    target = str(path.resolve()).casefold()
    start_registered_command(path, automation=WPS_LAUNCH == "automation")
    started: set[int] = set()

    def release() -> str:
        return release_opened_fixture(component, path, started)

    deadline = time.monotonic() + timeout
    last_errors: list[dict[str, Any]] = []
    while time.monotonic() < deadline:
        ok, data, last_errors = list_open_documents(component)
        listed = {str(Path(item.get("full_name") or "").resolve()).casefold() for item in data.get("documents", [])}
        if ok and target in listed:
            started.update(wps_process_ids() - before)
            return release, None
        time.sleep(poll)
    started.update(wps_process_ids() - before)
    release()
    detail = "; ".join(f"{error.get('code')}: {error.get('message')}" for error in last_errors)
    # Which WPS processes were there: in P7-007 the runs after a stopped Spreadsheets instance could not reach
    # the new one, and the record could not say what was running.
    running = ",".join(map(str, sorted(wps_process_ids()))) or "none"
    ours = ",".join(map(str, sorted(started))) or "none"
    earlier = ",".join(map(str, sorted(before))) or "none"
    processes = f" [WPS processes: {running}; started for this run: {ours}; running before it: {earlier}]"
    return (
        None,
        f"WPS did not list {path.name} as open within {timeout:.0f} s (launch mode {WPS_LAUNCH})"
        + (f": {detail}" if detail else f"; is WPS the default app for {path.suffix}?")
        + processes,
    )


def release_opened_fixture(component: str, path: Path, started: set[int]) -> str:
    """Close the fixture open_in_wps opened, then end the WPS processes started for it.

    A run that stopped early can leave the fixture with unsaved changes (P6-008 rerun). It is the runner's own
    file in a task workspace, so it is closed without saving; left open, it would also keep its WPS process
    from being ended.
    """
    from wps_ai_agent_cli.open_documents import close_unchanged_document

    status, error = close_unchanged_document(component, path, do_close=True)
    closed = status or f"{(error or {}).get('code')} ({(error or {}).get('message')})"
    if error and error.get("code") == "DOCUMENT_HAS_UNSAVED_CHANGES":
        closed = "discarded: " + str(wps_document_action(component, path, "close_discard").get("status"))
    # Recorded with the run, so that a WPS process left running afterwards can be explained (P6-008 rerun).
    return f"fixture {closed}; processes {stop_started_wps(component, started)}"


def prepare_open_document(
    step: str, fixture: Path, component: str, options: dict[str, Any], state: dict[str, Any] | None = None
) -> str | None:
    """Run a step that needs the document open in WPS. Returns a reason when it could not be done.

    `state` collects what a step made that the run needs later: the name of an untitled document.
    """
    state = {} if state is None else state
    if step == "select_text" and component == "spreadsheets":
        selection = options.get("selection") or {}
        wanted = str(selection.get("range") or "").replace("$", "").upper()
        payload = wps_document_action(component, fixture, "select", range=wanted, sheet=selection.get("sheet"))
        if payload.get("status") != "ok":
            return f"select_text: {payload.get('status')} {payload.get('errors') or ''}".strip()
        got = str(payload.get("address") or "").replace("$", "").upper()
        if got != wanted:
            return f"select_text: WPS selected {got!r}, expected {wanted!r}"
        sheet = selection.get("sheet")
        if sheet and payload.get("sheet") != sheet:
            return f"select_text: WPS selected cells on sheet {payload.get('sheet')!r}, expected {sheet!r}"
        return None
    if step == "untitled_document":
        text = str(options.get("untitled_text") or "Draft notes")
        payload = wps_document_action(component, fixture, "new_untitled", text=text)
        name = str(payload.get("name") or "")
        if payload.get("status") != "ok" or not name:
            return f"untitled_document: {payload.get('status')} {payload.get('errors') or ''}".strip()
        state["untitled_name"] = name
        return None
    if step == "select_text":
        selection = options.get("selection") or {}
        text = str(selection.get("text") or "")
        collapse = selection.get("collapse")
        payload = wps_document_action(component, fixture, "select", text=text, collapse=collapse)
        if payload.get("status") != "ok":
            return f"select_text: {payload.get('status')} {payload.get('errors') or ''}".strip()
        if collapse == "end":
            # A collapsed selection reports the character after the caret as its text, so compare positions.
            if payload.get("start") is None or payload.get("start") != payload.get("end"):
                return f"select_text: expected a caret, WPS reports {payload.get('start')}..{payload.get('end')}"
            if component == "presentation" and payload.get("selection_type") != 3:
                return f"select_text: WPS reports selection type {payload.get('selection_type')}, not a caret in text"
            return None
        if component == "presentation" and payload.get("selection_type") != 3:
            # The selection tools take only a text selection (type 3); anything else is a setup failure.
            return f"select_text: WPS reports selection type {payload.get('selection_type')}, not a text selection"
        if str(payload.get("selection_text") or "").rstrip("\r") != text:
            return f"select_text: WPS selected {payload.get('selection_text')!r}, expected {text!r}"
        return None
    if step == "append_unsaved":
        payload = wps_document_action(component, fixture, "append_unsaved", text=str(options.get("unsaved_text") or ""))
        if payload.get("status") != "ok" or payload.get("saved") is not False:
            return f"append_unsaved: {payload.get('status')}, saved={payload.get('saved')}"
        return None
    raise ValueError(f"unknown setup step {step}")


OPEN_DOCUMENT_STEPS = ("select_text", "append_unsaved", "untitled_document")


def close_workspace_documents(component: str, workspace: Path, fixture: Path, untitled_name: str | None) -> str:
    """Close what a run left open in WPS besides the fixture: the untitled document, or the file it was saved as.

    Both are the runner's own: the untitled document was made by the untitled_document step, and anything else
    WPS lists inside the task workspace was saved there during the run. They are closed without saving.
    """
    from wps_ai_agent_cli.open_documents import list_open_documents

    notes = []
    if untitled_name:
        payload = wps_document_action(component, fixture, "close_untitled", name=untitled_name)
        notes.append(f"untitled {untitled_name} {payload.get('status')}")
    root = str(workspace.resolve()).casefold()
    _ok, data, _errors = list_open_documents(component)
    for item in data.get("documents", []):
        full = str(item.get("full_name") or "")
        if not full:
            continue
        path = Path(full)
        resolved = str(path.resolve()).casefold()
        if resolved.startswith(root) and resolved != str(fixture.resolve()).casefold():
            notes.append(f"{path.name} {wps_document_action(component, path, 'close_discard').get('status')}")
    return "; ".join(notes) or "no other documents"


def documents_outside(workspace: Path) -> list[dict[str, Any]] | str:
    """What WPS lists as open, in any component, outside the task workspace: documents it restored (P8-007).

    WPS reopens the documents of its previous session when it starts; in P7-007 they showed up in the agent's
    open_documents results and had to be closed by hand. They belong to the machine's user, so the runner
    only records them. Returns the reason instead when the list failed.
    """
    from wps_ai_agent_cli.open_documents import list_open_documents

    ok, data, errors = list_open_documents()
    if not ok:
        return "; ".join(f"{error.get('code')}: {error.get('message')}" for error in errors) or "list failed"
    root = str(workspace.resolve()).casefold()
    outside = []
    for item in data.get("documents", []):
        full = str(item.get("full_name") or "")
        inside = bool(full) and os.path.isabs(full) and str(Path(full).resolve()).casefold().startswith(root)
        if not inside:
            outside.append({"component": item.get("component"), "name": item.get("name"), "full_name": full})
    return outside


def run_setup(
    steps: list[str],
    workspace: Path,
    fixture: Path,
    component: str,
    options: dict[str, Any] | None = None,
    state: dict[str, Any] | None = None,
) -> tuple[str, Callable[[], None] | None, str | None]:
    """Prepare the workspace.

    Returns (hash of the file as the agent receives it, release for a held or opened file, skip reason). The
    hash is taken before a hold or an open, because a held file cannot be read; "hold" must be the last step,
    and only steps that work on the open document (select_text, append_unsaved, untitled_document) may follow
    "open_in_wps". `state` receives what the steps made, such as untitled_name.
    """
    state = {} if state is None else state
    for index, step in enumerate(steps):
        if step == "register":
            ok, record, errors = register_document(component, str(fixture), workspace)
            if not ok:
                raise RuntimeError(f"setup register failed: {errors}")
        elif step == "backup":
            document_id = stable_document_id(component, str(fixture))
            ok, _result, errors, _replayed = create_backup(document_id, "setup-backup", workspace=workspace)
            if not ok:
                raise RuntimeError(f"setup backup failed: {errors}")
        elif step == "modify":
            from openpyxl import load_workbook

            workbook = load_workbook(fixture)
            workbook.active["A1"] = "changed by mistake"
            workbook.save(fixture)
        elif step == "html_table":
            (workspace / "table.html").write_text(HTML_TABLE, encoding="utf-8")
        elif step in ("hold", "open_in_wps"):
            later = steps[index + 1 :]
            allowed = OPEN_DOCUMENT_STEPS if step == "open_in_wps" else ()
            if any(item not in allowed for item in later):
                raise ValueError(f"setup step {step} can only be followed by {list(allowed)}, not {later}")
            start_sha = file_sha256(fixture)
            opened, reason = hold_file(fixture) if step == "hold" else open_in_wps(fixture, component)
            if opened is None:
                return start_sha, None, reason
            release = opened
            if step == "open_in_wps":
                # Before the later steps, so that an untitled document the runner makes is not counted.
                state["documents_outside"] = documents_outside(workspace)
            if "untitled_document" in later:

                def release(opened: Callable[[], Any] = opened) -> str:
                    others = close_workspace_documents(component, workspace, fixture, state.get("untitled_name"))
                    return f"{others}; {opened()}"

            for item in later:
                reason = prepare_open_document(item, fixture, component, options or {}, state)
                if reason:
                    release()
                    return start_sha, None, reason
            return start_sha, release, None
        else:
            raise ValueError(f"unknown setup step {step}")
    return file_sha256(fixture), None, None


def _error_codes(content: Any) -> list[str]:
    """Error codes from a wps-agent tool result, which is the CLI envelope as JSON text."""
    if isinstance(content, list):
        content = "".join(part.get("text", "") for part in content if isinstance(part, dict))
    try:
        envelope = json.loads(content) if isinstance(content, str) else content
    except json.JSONDecodeError:
        return []
    if not isinstance(envelope, dict):
        return []
    return [str(error.get("code")) for error in envelope.get("errors") or [] if isinstance(error, dict)]


def parse_transcript(lines: list[str]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    """Return the final result event and every tool call, in order, from stream-json output."""
    result: dict[str, Any] = {}
    calls: list[dict[str, Any]] = []
    by_id: dict[str, dict[str, Any]] = {}
    for line in lines:
        try:
            event = json.loads(line)
        except json.JSONDecodeError:
            continue
        if event.get("type") == "result":
            result = event
            continue
        message = event.get("message") or {}
        content = message.get("content") if isinstance(message, dict) else None
        if not isinstance(content, list):
            continue
        for block in content:
            if block.get("type") == "tool_use":
                call = {
                    "name": block.get("name"),
                    "input": block.get("input") or {},
                    "is_error": None,
                    "error_codes": [],
                }
                calls.append(call)
                by_id[block.get("id")] = call
            elif block.get("type") == "tool_result" and block.get("tool_use_id") in by_id:
                call = by_id[block["tool_use_id"]]
                call["error_codes"] = _error_codes(block.get("content"))
                call["is_error"] = bool(block.get("is_error")) or bool(call["error_codes"])
    return result, calls


def tool_metrics(
    calls: list[dict[str, Any]], expected_tools: list[str] | None = None, open_document: bool = False
) -> dict[str, Any]:
    """Counts for one run. With expected_tools, also the wps-agent calls to any other tool (off-path calls).

    For a task whose document is open in WPS, wrong_tool_calls counts the calls to file-based write tools, which
    cannot change an open document (P9-002); it is None for other tasks. Off-path and wrong-tool calls are
    reported, not judged: a task passes or fails on its checks alone.
    """
    wps_calls = [call for call in calls if str(call["name"]).startswith(TOOL_PREFIX)]
    used = [str(call["name"]).removeprefix(TOOL_PREFIX) for call in wps_calls]
    metrics = {
        "wps_tool_calls": len(wps_calls),
        "wps_tool_errors": sum(bool(call["is_error"]) for call in wps_calls),
        "other_tool_calls": len(calls) - len(wps_calls),
        "tools_used": sorted(set(used)),
        "expected_tools": sorted(expected_tools) if expected_tools is not None else None,
        "off_path_calls": None,
        "off_path_tools": None,
        "wrong_tool_calls": sum(tool in FILE_BASED_WRITE_TOOLS for tool in used) if open_document else None,
    }
    if expected_tools is not None:
        off_path = [tool for tool in used if tool not in expected_tools]
        metrics["off_path_calls"] = len(off_path)
        metrics["off_path_tools"] = sorted(set(off_path))
    return metrics


def run_claude(
    prompt: str, mcp_config: Path, workspace: Path, model: str | None, max_turns: int, timeout: int
) -> dict[str, Any]:
    command = [
        # On Windows the CLI is claude.CMD, which subprocess only finds by its full name (P6-008).
        shutil.which("claude") or "claude",
        "-p",
        prompt,
        "--output-format",
        "stream-json",
        "--verbose",
        "--mcp-config",
        str(mcp_config),
        "--strict-mcp-config",
        "--allowedTools",
        ALLOWED_TOOLS,
        "--max-turns",
        str(max_turns),
    ]
    if model:
        command += ["--model", model]
    try:
        # The transcript is UTF-8 whatever the console code page is (P6-008).
        completed = subprocess.run(
            command,
            cwd=workspace,
            capture_output=True,
            encoding="utf-8",
            errors="replace",
            timeout=timeout,
        )
    except subprocess.TimeoutExpired:
        return {"ran": False, "error": f"timed out after {timeout}s"}
    payload, calls = parse_transcript(completed.stdout.splitlines())
    return {
        "ran": True,
        "exit_code": completed.returncode,
        "is_error": payload.get("is_error"),
        "num_turns": payload.get("num_turns"),
        "duration_ms": payload.get("duration_ms"),
        "total_cost_usd": payload.get("total_cost_usd"),
        "result_excerpt": str(payload.get("result", ""))[:500],
        "stderr_excerpt": completed.stderr[-500:],
        "tool_calls": calls,
    }


def agent_error(agent: dict[str, Any]) -> str | None:
    """The CLI's own error when the agent could not work at all, for example an API error such as 402
    Insufficient Balance (P6-008 rerun). Such a run still fails, but says nothing about the tools."""
    excerpt = str(agent.get("result_excerpt") or "")
    if agent.get("is_error") and excerpt.startswith("API Error"):
        return excerpt[:200]
    return None


def _calls_to(calls: list[dict[str, Any]], tool: str) -> list[dict[str, Any]]:
    return [call for call in calls if call["name"] == TOOL_PREFIX + tool]


def judge(
    task: dict[str, Any],
    workspace: Path,
    fixture: Path,
    fixture_sha: str,
    component: str,
    seed_sha: str | None = None,
    calls: list[dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    """Check the workspace state and tool calls the task asks for. A check never reads the agent's reply.

    fixture_sha is the file as the agent received it, after setup; seed_sha is the file before setup.
    """
    seed_sha = seed_sha or fixture_sha
    calls = calls or []
    results = []
    document = get_document(stable_document_id(component, str(fixture)), workspace)
    document_id = document["document_id"] if document else None
    for check in task["checks"]:
        kind = check["kind"]
        if kind == "document_registered":
            passed, detail = document is not None, document_id or "not registered"
        elif kind == "operation_recorded":
            record = get_operation(check["request_id"], workspace)
            passed = bool(record) and record.get("command") == check["command"]
            detail = record.get("command") if record else "no record"
        elif kind == "no_operation":
            passed = get_operation(check["request_id"], workspace) is None
            detail = "none recorded" if passed else "a record exists"
        elif kind == "backup_count":
            count = 0
            if document_id:
                ok, data, _errors = list_backups(document_id, workspace)
                count = len(data.get("backups", [])) if ok else -1
            passed, detail = count == check["count"], count
        elif kind == "file_unchanged":
            passed = file_sha256(fixture) == fixture_sha
            detail = "unchanged" if passed else "changed"
        elif kind == "file_matches_seed":
            passed = file_sha256(fixture) == seed_sha
            detail = "matches the seed" if passed else "differs from the seed"
        elif kind == "tool_called":
            matching = [
                call
                for call in _calls_to(calls, check["tool"])
                if all(call["input"].get(key) == value for key, value in check.get("input", {}).items())
            ]
            passed, detail = len(matching) >= check.get("min", 1), f"{len(matching)} matching call(s)"
        elif kind == "tool_not_called":
            count = len(_calls_to(calls, check["tool"]))
            passed, detail = count == 0, f"{count} call(s)"
        elif kind == "tool_error_seen":
            passed = any(check["code"] in call["error_codes"] for call in calls)
            detail = "seen" if passed else "not seen"
        elif kind == "max_tool_calls":
            count = tool_metrics(calls)["wps_tool_calls"]
            passed, detail = count <= check["count"], f"{count} wps-agent call(s)"
        elif kind == "file_text":
            path = workspace / check["path"]
            if not path.is_file():
                passed, detail = False, "file missing"
            else:
                text = document_text(path)
                missing = [item for item in check.get("contains", []) if item not in text]
                present = [item for item in check.get("absent", []) if item in text]
                passed = not missing and not present
                detail = "ok" if passed else {"missing": missing, "present": present}
        elif kind == "text_encoding":
            path = workspace / check["path"]
            encoding = text_file_encoding(path) if path.is_file() else None
            passed = encoding in check["encodings"]
            detail = encoding or "file missing"
        elif kind == "xlsx_sheets":
            from openpyxl import load_workbook

            path = workspace / check["path"]
            if not path.is_file():
                passed, detail = False, "file missing"
            else:
                workbook = load_workbook(path, read_only=True)
                try:
                    names = list(workbook.sheetnames)
                finally:
                    workbook.close()
                passed, detail = names == check["names"], names
        elif kind == "docx_table_cell":
            from docx import Document

            path = workspace / check["path"]
            if not path.is_file():
                passed, detail = False, "file missing"
            else:
                tables = Document(path).tables
                try:
                    value = tables[check["table"] - 1].cell(check["row"] - 1, check["column"] - 1).text
                except IndexError:
                    value = None
                passed, detail = value == check["text"], value
        elif kind == "pdf_file":
            path = workspace / check["path"]
            if not path.is_file():
                passed, detail = False, "file missing"
            else:
                head = path.read_bytes()[:5]
                size = path.stat().st_size
                passed = head == b"%PDF-" and size >= check.get("min_bytes", 1)
                detail = {"bytes": size, "starts_with_pdf_header": head == b"%PDF-"}
        elif kind == "xlsx_cell":
            from openpyxl import load_workbook

            path = workspace / check["path"]
            if not path.is_file():
                passed, detail = False, "output missing"
            else:
                workbook = load_workbook(path, read_only=True)
                try:
                    if "sheet" in check and check["sheet"] not in workbook.sheetnames:
                        value, detail = None, f"sheet {check['sheet']} missing"
                    else:
                        sheet = workbook[check["sheet"]] if "sheet" in check else workbook.worksheets[0]
                        value = detail = sheet[check["cell"]].value
                finally:
                    workbook.close()
                passed = value == check["value"]
        else:
            passed, detail = False, f"unknown check kind {kind}"
        results.append(
            {"kind": kind, "passed": passed, "detail": detail, **{k: v for k, v in check.items() if k != "kind"}}
        )
    return results


def wps_available(component: str) -> bool:
    """True when this machine has a registered WPS ProgID for the component (the project's own probe)."""
    from wps_ai_agent_cli.capabilities import probe_wps_capabilities

    selected = probe_wps_capabilities().get("components", {}).get(component, {}).get("selected_prog_id")
    return bool(selected)


def run_task(
    task: dict[str, Any],
    fixture_name: str,
    component: str,
    root: Path,
    model: str | None,
    max_turns: int,
    timeout: int,
    runner: Callable[..., dict[str, Any]] = run_claude,
    has_wps: Callable[[str], bool] = wps_available,
) -> dict[str, Any]:
    if task.get("requires_wps") and not has_wps(component):
        return {
            "id": task["id"],
            "passed": False,
            "skipped": f"needs WPS {component}, which is not available here",
            "agent": {},
            "checks": [],
        }
    workspace = root / task["id"]
    if workspace.exists():
        shutil.rmtree(workspace)
    workspace.mkdir(parents=True)
    fixture = workspace / fixture_name
    seed_sha = write_fixture(workspace, fixture_name)
    state: dict[str, Any] = {}
    fixture_sha, release, skip_reason = run_setup(
        task.get("setup", []), workspace, fixture, component, task, state=state
    )
    if skip_reason:
        return {"id": task["id"], "passed": False, "skipped": skip_reason, "agent": {}, "checks": []}
    # WPS names an untitled document itself (for example "文字文稿2"), so the prompt can only name it now.
    prompt = task["prompt"].replace("{untitled_name}", state.get("untitled_name", "{untitled_name}"))
    cleanup = None
    try:
        mcp_config = write_mcp_config(workspace)
        agent = runner(prompt, mcp_config, workspace, model, max_turns, timeout)
    finally:
        if release is not None:
            cleanup = release()
    calls = agent.get("tool_calls") or []
    checks = judge(task, workspace, fixture, fixture_sha, component, seed_sha=seed_sha, calls=calls)
    result = {
        "id": task["id"],
        "passed": bool(agent.get("ran")) and all(check["passed"] for check in checks),
        "agent": agent,
        "metrics": tool_metrics(calls, task.get("expected_tools"), opens_document(task)),
        "checks": checks,
    }
    error = agent_error(agent)
    if error:
        result["agent_error"] = error
    if cleanup:
        result["cleanup"] = cleanup
    if state.get("untitled_name"):
        result["untitled_name"] = state["untitled_name"]
    if "documents_outside" in state:
        result["documents_outside"] = state["documents_outside"]
    return result


def opens_document(task: dict[str, Any]) -> bool:
    """Whether the task's document is open in WPS while the agent works on it."""
    return "open_in_wps" in task.get("setup", [])


def run_task_repeated(
    task: dict[str, Any],
    fixture_name: str,
    component: str,
    root: Path,
    model: str | None,
    max_turns: int,
    timeout: int,
    repeat: int = 1,
    min_passes: int = 1,
    runner: Callable[..., dict[str, Any]] = run_claude,
    has_wps: Callable[[str], bool] = wps_available,
) -> dict[str, Any]:
    """Run a task `repeat` times, each in its own workspace. It passes when at least `min_passes` runs pass.

    A skip on the first run (no WPS, or a file that cannot be held) does not depend on the run, so it ends the
    task. A setup that fails only on a later run (for example WPS not listing a document it opened before) is
    that run failing: it counts as a failed run with `setup_failed`, and the earlier runs are kept.
    """
    runs = []
    for index in range(1, repeat + 1):
        result = run_task(
            task,
            fixture_name,
            component,
            root / f"run-{index}",
            model,
            max_turns,
            timeout,
            runner=runner,
            has_wps=has_wps,
        )
        if result.get("skipped"):
            if not runs:
                return {"id": task["id"], "passed": False, "skipped": result["skipped"], "runs": []}
            result = {
                "id": task["id"],
                "passed": False,
                "setup_failed": result["skipped"],
                "agent": {},
                "metrics": tool_metrics([], task.get("expected_tools"), opens_document(task)),
                "checks": [],
            }
        runs.append(result)
    passes = sum(run["passed"] for run in runs)
    return {
        "id": task["id"],
        "passed": passes >= min_passes,
        "repeat": repeat,
        "min_passes": min_passes,
        "passes": passes,
        "pass_rate": passes / repeat,
        "metrics": {
            "wps_tool_calls": sum(run["metrics"]["wps_tool_calls"] for run in runs),
            "wps_tool_errors": sum(run["metrics"]["wps_tool_errors"] for run in runs),
            "off_path_calls": (
                sum(run["metrics"]["off_path_calls"] for run in runs) if "expected_tools" in task else None
            ),
            "off_path_tools": (
                sorted({tool for run in runs for tool in run["metrics"]["off_path_tools"]})
                if "expected_tools" in task
                else None
            ),
            "wrong_tool_calls": (
                sum(run["metrics"].get("wrong_tool_calls") or 0 for run in runs) if opens_document(task) else None
            ),
        },
        "runs": runs,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--tasks", default=str(REPO_ROOT / "config" / "agent_eval_tasks.json"))
    parser.add_argument("--out", default=str(REPO_ROOT / "artifacts" / "agent-eval"))
    parser.add_argument("--only", action="append", default=[], help="Run only the task with this id (repeatable).")
    parser.add_argument("--model", default=None)
    parser.add_argument("--max-turns", type=int, default=15)
    parser.add_argument("--timeout", type=int, default=600)
    parser.add_argument("--repeat", type=int, default=1, help="Run each task this many times (default 1).")
    parser.add_argument(
        "--min-passes", type=int, default=None, help="Runs that must pass for a task to pass (default: all of them)."
    )
    parser.add_argument(
        "--wps-launch",
        choices=WPS_LAUNCH_MODES,
        default="default",
        help="How the open_in_wps step starts WPS: default opens the file like a double-click; automation "
        "starts the associated program with /Automation (the system open command is not changed).",
    )
    args = parser.parse_args(argv)
    global WPS_LAUNCH
    WPS_LAUNCH = args.wps_launch
    if args.repeat < 1:
        parser.error("--repeat must be at least 1")
    min_passes = args.repeat if args.min_passes is None else args.min_passes
    if not 1 <= min_passes <= args.repeat:
        parser.error("--min-passes must be between 1 and --repeat")

    if shutil.which("claude") is None:
        print("claude CLI not found on PATH.", file=sys.stderr)
        return 2
    for stream in (sys.stdout, sys.stderr):
        # Check details can hold text the console code page cannot encode; never fail on printing them.
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(errors="replace")

    spec = json.loads(Path(args.tasks).read_text(encoding="utf-8"))
    tasks = [task for task in spec["tasks"] if not args.only or task["id"] in args.only]
    stamp = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    root = Path(args.out) / f"run-{stamp}"
    root.mkdir(parents=True, exist_ok=True)

    results = [
        run_task_repeated(
            task,
            task.get("fixture", spec["fixture"]),
            task.get("component", spec["component"]),
            root / "work",
            args.model,
            args.max_turns,
            args.timeout,
            repeat=args.repeat,
            min_passes=min_passes,
        )
        for task in tasks
    ]
    summary = summarize(results, stamp, args.model)
    artifact = root / "summary.json"
    artifact.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    for result in results:
        if result.get("skipped"):
            print(f"SKIP {result['id']}: {result['skipped']}")
            continue
        metrics = result.get("metrics", {})
        print(
            f"{'PASS' if result['passed'] else 'FAIL'} {result['id']} "
            f"({result['passes']}/{result['repeat']} runs passed, need {result['min_passes']}; "
            f"wps-agent calls {metrics.get('wps_tool_calls')}, errors {metrics.get('wps_tool_errors')}, "
            f"off-path calls {metrics.get('off_path_calls')}, wrong-tool calls {metrics.get('wrong_tool_calls')})"
        )
        for index, run in enumerate(result["runs"], 1):
            label = f"  run {index}: " if result["repeat"] > 1 else ""
            if run.get("setup_failed"):
                print(f"    {label}BAD setup: {run['setup_failed']}")
            if run.get("agent_error"):
                print(f"    {label}BAD agent did not run: {run['agent_error']}")
            for check in run["checks"]:
                print(f"    {label}{'ok ' if check['passed'] else 'BAD'} {check['kind']}: {check['detail']}")
    print(f"{summary['passed']}/{summary['total']} passed, {summary['skipped']} skipped; artifact: {artifact}")
    return 0 if summary["passed"] == summary["total"] else 1


def summarize(results: list[dict[str, Any]], stamp: str, model: str | None) -> dict[str, Any]:
    ran = [result for result in results if not result.get("skipped")]
    return {
        "generated_at": stamp,
        "claude_cli": "claude",
        "model": model or "CLI default",
        "wps_launch": WPS_LAUNCH,
        "passed": sum(result["passed"] for result in ran),
        "total": len(ran),
        "skipped": len(results) - len(ran),
        "wps_tool_calls": sum(result.get("metrics", {}).get("wps_tool_calls", 0) for result in ran),
        "wps_tool_errors": sum(result.get("metrics", {}).get("wps_tool_errors", 0) for result in ran),
        "off_path_calls": sum(result.get("metrics", {}).get("off_path_calls") or 0 for result in ran),
        "wrong_tool_calls": sum(result.get("metrics", {}).get("wrong_tool_calls") or 0 for result in ran),
        "tasks_without_expected_tools": [
            result["id"] for result in ran if result.get("metrics", {}).get("off_path_calls") is None
        ],
        "results": results,
    }


if __name__ == "__main__":
    raise SystemExit(main())
