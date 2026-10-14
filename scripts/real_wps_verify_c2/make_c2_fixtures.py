"""Write the two scratch workbooks used by c2-post-open.ps1. Never touches user documents.

usage: python make_c2_fixtures.py <work_dir>        (needs openpyxl; pip install -e ".[test]" covers it)
"""

import sys
from pathlib import Path

import openpyxl

out = Path(sys.argv[1])
out.mkdir(parents=True, exist_ok=True)
for name in ("c2_clean.xlsx", "c2_dirty.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = 1
    ws["B1"] = "scratch"
    wb.save(out / name)
print("fixtures written to", out)
