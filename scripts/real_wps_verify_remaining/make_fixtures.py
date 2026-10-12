"""Write scratch fixtures for remaining real-WPS checks. Never touches user documents.

usage: python make_fixtures.py <work_dir>
needs: openpyxl, python-docx, python-pptx  (pip install -e ".[test]")
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
from docx import Document
from pptx import Presentation
from pptx.util import Inches

out = Path(sys.argv[1])
out.mkdir(parents=True, exist_ok=True)

# Spreadsheet: inspect + write targets
for name, a1 in (
    ("r1_inspect_free.xlsx", 11),
    ("r1_inspect_hold.xlsx", 22),
    ("r1_inspect_wps.xlsx", 33),
    ("r3_alerts.xlsx", 44),
    ("r4_race.xlsx", 55),
    ("r5_other_free.xlsx", 65),
    ("r5_other_hold.xlsx", 66),
):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = a1
    ws["B1"] = "scratch"
    wb.save(out / name)

# Writer
for name, text in (
    ("r2_writer_clean.docx", "ALPHA clean scratch"),
    ("r2_writer_hold.docx", "ALPHA hold scratch"),
):
    doc = Document()
    doc.add_paragraph(text)
    doc.save(out / name)

# Presentation
for name, text in (
    ("r2_wpp_clean.pptx", "ALPHA clean slide"),
    ("r2_wpp_hold.pptx", "ALPHA hold slide"),
):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text_frame.text = text
    prs.save(out / name)

print("fixtures written to", out)
