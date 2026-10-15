from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from .document_text import count_text_in_docx
from .errors import INPUT_FILE_NOT_FOUND
from .sessions import get_document
from .writer_ops import BODY_PARTS


def _normalize(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return "" if value is None else str(value)


# Agents often called validate-document with only a document_id, as a general "check the document" step, and
# got INVALID_ARGUMENT without being told what to pass (P6-008, P7-006). The messages below say what each
# component needs, with an example, and where to go to read a document instead.
_READ_INSTEAD = "To read the document instead of checking one value, use snapshot-document (snapshot_document)."
_WRITER_NEEDS = (
    "Writer validation needs contains (--contains): text that must appear in the document body, "
    'for example contains="Status: FINAL".'
)
_SPREADSHEET_NEEDS = (
    'Spreadsheet validation needs both cell (--cell), an A1 address on the first sheet such as "B2", '
    'and equals (--equals), the value that cell must hold as text, for example cell="B2" equals="42".'
)


_A1_CELL = re.compile(r"^\$?[A-Za-z]{1,3}\$?[1-9][0-9]*$")


def _invalid(document: dict[str, Any], message: str) -> tuple[bool, dict[str, Any], list[dict[str, str]]]:
    return False, {"document": document}, [{"code": "INVALID_ARGUMENT", "message": message}]


def validate_document(
    document_id: str,
    contains: str | None = None,
    cell: str | None = None,
    equals: str | None = None,
    workspace: str | Path = ".",
) -> tuple[bool, dict[str, Any], list[dict[str, str]]]:
    document = get_document(document_id, workspace)
    if not document:
        return (
            False,
            {},
            [{"code": "DOCUMENT_NOT_FOUND", "message": f"Document not registered: {document_id}"}],
        )

    path = Path(document["path"])
    if not path.exists():
        return (
            False,
            {"document": document},
            [{"code": INPUT_FILE_NOT_FOUND, "message": f"Input file not found: {path}"}],
        )

    if document["component"] == "writer":
        if not contains:
            given = " cell and equals are for spreadsheets." if cell or equals is not None else ""
            return _invalid(document, f"{_WRITER_NEEDS}{given} {_READ_INSTEAD}")
        count = count_text_in_docx(path, contains, parts=BODY_PARTS)
        result = {
            "document_id": document_id,
            "component": "writer",
            "path": str(path),
            "scope": "body",
            "contains": contains,
            "count": count,
        }
        if count <= 0:
            message = (
                f"The text {contains!r} was not found in the document body. The check is an exact, case-sensitive "
                "match within one paragraph, spaces included."
            )
            return False, result, [{"code": "VALIDATION_FAILED", "message": message}]
        return True, result, []

    if document["component"] == "spreadsheets":
        if not cell or equals is None:
            missing = " and ".join(name for name, absent in (("cell", not cell), ("equals", equals is None)) if absent)
            given = " contains is for Writer documents." if contains else ""
            return _invalid(document, f"{_SPREADSHEET_NEEDS} Missing: {missing}.{given} {_READ_INSTEAD}")
        if not _A1_CELL.match(cell):
            return _invalid(
                document,
                f'cell {cell!r} is not an A1 address. Give one cell on the first sheet, such as "B2", '
                "without a sheet name or a range.",
            )
        try:
            import openpyxl  # noqa: F401

            from .ooxml import load_workbook_guarded as load_workbook
        except ImportError:
            return (
                False,
                {"document": document},
                [{"code": "DEPENDENCY_MISSING", "message": "openpyxl is required for spreadsheet validation."}],
            )

        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            sheet = workbook.worksheets[0]
            value = sheet[cell].value
        finally:
            workbook.close()

        normalized = _normalize(value)
        passed = normalized == equals
        result = {
            "document_id": document_id,
            "component": "spreadsheets",
            "path": str(path),
            "sheet": sheet.title,
            "cell": cell,
            "expected": equals,
            "actual": normalized,
        }
        if not passed:
            message = f"Cell {cell} on sheet {sheet.title!r} holds {normalized!r}, not {equals!r}."
            return False, result, [{"code": "VALIDATION_FAILED", "message": message}]
        return True, result, []

    return (
        False,
        {"document": document},
        [
            {
                "code": "UNSUPPORTED_COMPONENT",
                "message": (
                    f"validate-document checks Writer text and spreadsheet cells only, not {document['component']}. "
                    + _READ_INSTEAD
                ),
            }
        ],
    )
