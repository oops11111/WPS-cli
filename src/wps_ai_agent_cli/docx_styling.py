"""Computed CSS styles applied to python-docx objects (P13-004).

html_editable builds the document's structure; the functions here give each paragraph, run, table, cell, image
and section the look the browser (or the built-in cascade) computed for the HTML element it came from: fonts,
sizes, colours, weight, slant, decoration, letter spacing, alignment, line height, spacing before and after,
indents, shading, borders, page breaks, column widths, cell padding and alignment, image sizes, and the page
size and margins of @page. In appearance mode (decision 9), absolutely positioned elements keep their position
as framed paragraphs.
"""

from __future__ import annotations

import math
import re
from dataclasses import dataclass, field
from typing import Any

from .html_styles import StyledDocument, _walk, color, first_font, is_bold, px

EMU_PER_PX = 9525
_TRANSPARENT_PARENTS = {"html", "body", "main", "document"}
PRINT_MARGIN_MM = 10
_PAGE_SIZES = {"a4": (210, 297), "a3": (297, 420), "a5": (148, 210), "letter": (215.9, 279.4), "legal": (215.9, 355.6)}


@dataclass
class Context:
    styled: StyledDocument | None
    appearance: bool = False
    last_space_after: float = 0.0
    carry_before: float = 0.0  # a table's margin-bottom, in points, for the paragraph after it
    applied: set[str] = field(default_factory=set)


def _qn(tag: str):
    from docx.oxml.ns import qn

    return qn(tag)


def _element(tag: str, **attrs: str):
    from docx.oxml import OxmlElement

    element = OxmlElement(tag)
    for key, value in attrs.items():
        element.set(_qn(key), value)
    return element


def run_props(context: Context, node) -> dict[str, Any]:
    if context.styled is None or node is None:
        return {}
    style = context.styled.style(node)
    size = px(style.get("font-size")) or 16.0
    props: dict[str, Any] = {
        "font": first_font(style.get("font-family")),
        "size": round(size * 0.75, 1),
        "color": color(style.get("color")),
        "bold": is_bold(style),
        "italic": style.get("font-style") in {"italic", "oblique"},
    }
    # Decorations, script position and an inline background are not inherited, but show on the text inside:
    # take them from the inline elements around the text too.
    chain = context.styled.inline_ancestors(node)
    inline = [n for n in chain if context.styled.style(n).get("display", "inline") == "inline"]
    decoration = " ".join(context.styled.style(n).get("text-decoration-line", "") for n in chain)
    if "underline" in decoration:
        props["underline"] = True
    if "line-through" in decoration:
        props["strike"] = True
    for element in inline:
        shade = color(context.styled.style(element).get("background-color"))
        if shade:
            props["shade"] = shade
            break
    spacing = px(style.get("letter-spacing"), size)
    if spacing:
        props["spacing"] = spacing * 0.75
    vertical = next(
        (v for v in (context.styled.style(n).get("vertical-align") for n in inline) if v in {"sub", "super"}), None
    )
    if vertical in {"sub", "super"}:
        props["subscript" if vertical == "sub" else "superscript"] = True
    context.applied.add("character")
    return props


def _style_turns_on(run, name: str) -> bool:
    from .office_html import _inherited

    try:
        return bool(_inherited(run, run._parent, name))
    except AttributeError:
        return False


def apply_run(run, props: dict[str, Any]) -> None:
    from docx.shared import Pt, RGBColor

    if not props:
        return
    font = run.font
    if props.get("font"):
        font.name = props["font"]
        fonts = run._r.get_or_add_rPr().find(_qn("w:rFonts"))
        if fonts is None:
            fonts = _element("w:rFonts")
            run._r.get_or_add_rPr().insert(0, fonts)
        for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(_qn(key), props["font"])
    if props.get("size"):
        font.size = Pt(props["size"])
    if props.get("color"):
        font.color.rgb = RGBColor.from_string(props["color"])
    for name in ("bold", "italic"):
        if props.get(name):
            setattr(font, name, True)
        elif name in props and _style_turns_on(run, name):
            # The paragraph style (Quote, Caption, a heading) would add it; the HTML did not ask for it.
            setattr(font, name, False)
    if props.get("underline"):
        font.underline = True
    if props.get("strike"):
        font.strike = True
    if props.get("subscript"):
        font.subscript = True
    if props.get("superscript"):
        font.superscript = True
    if props.get("shade"):
        run._r.get_or_add_rPr().append(
            _element("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": props["shade"]})
        )
    if props.get("spacing"):
        run._r.get_or_add_rPr().append(_element("w:spacing", **{"w:val": str(int(props["spacing"] * 20))}))


def _container_with(context: Context, node, ancestors, name: str):
    """The nearest enclosing block (not the page body) whose computed style has a visible value for name."""
    for ancestor in reversed(ancestors):
        if ancestor.tag in _TRANSPARENT_PARENTS:
            return None
        style = context.styled.style(ancestor)
        if name == "background-color" and color(style.get(name)):
            return ancestor
        if name == "border" and any(_border(style, side) for side in ("top", "right", "bottom", "left")):
            return ancestor
    return None


def _border(style: dict[str, str], side: str) -> tuple[float, str, str] | None:
    kind = style.get(f"border-{side}-style", "none")
    width = px(style.get(f"border-{side}-width")) or 0
    if kind in {"none", "hidden", ""} or width <= 0:
        return None
    word = {"dashed": "dashed", "dotted": "dotted", "double": "double"}.get(kind, "single")
    return width, word, color(style.get(f"border-{side}-color")) or "000000"


def _borders_xml(tag: str, borders: dict[str, tuple[float, str, str]]):
    container = _element(tag)
    for side in ("top", "left", "bottom", "right"):
        if side in borders:
            width, kind, hue = borders[side]
            container.append(
                _element(
                    f"w:{side}",
                    **{"w:val": kind, "w:sz": str(max(2, int(width * 0.75 * 8))), "w:space": "1", "w:color": hue},
                )
            )
    return container


def _apply_padding(form, props, holder_style: dict[str, str], borders: dict, fill: str | None) -> None:
    """The box's padding: the border's distance from the text, and the right indent by the padding and border so
    that the box's edge stays where the browser draws it. A shaded box with padding but no border gets a border in
    its own fill colour, the only way Word draws shading beyond the text (Windows check on f8ed7bd: the manual's
    code block and note had no padding in WPS)."""
    from docx.shared import Pt

    pad = {side: px(holder_style.get(f"padding-{side}")) or 0.0 for side in ("top", "left", "bottom", "right")}
    if not borders and fill and any(pad.values()):
        borders = {side: (0.5, "single", fill) for side in ("top", "left", "bottom", "right")}
    if not borders:
        return
    container = _borders_xml("w:pBdr", borders)
    for edge in container:
        side = edge.tag.rsplit("}", 1)[1]
        edge.set(_qn("w:space"), str(max(0, min(31, int(round(pad.get(side, 0.0) * 0.75))))))
    props.append(container)
    # The left inset is already in the paragraph's indent (apply_paragraph); the right one is not.
    right = pad["right"] + (borders["right"][0] if "right" in borders else 0.0)
    if right:
        form.right_indent = Pt((form.right_indent.pt if form.right_indent is not None else 0.0) + right * 0.75)


def _left_offset(context: Context, node, ancestors) -> float | None:
    """How far the element's border edge starts right of the page's content edge, in CSS px (its own border and
    padding are added by the caller)."""
    styled = context.styled
    box = styled.box(node)
    if box is not None:
        body = next((a for a in ancestors if a.tag == "body"), None)
        body_box = styled.box(body) if body is not None else None
        base = body_box["x"] + (px(styled.style(body).get("padding-left")) or 0) if body_box else 8.0
        return max(0.0, box["x"] - base)
    total = 0.0
    for element in [*ancestors, node]:
        if element.tag in {"html", "body", "document", "ul", "ol"}:
            continue
        style = styled.style(element)
        total += px(style.get("margin-left")) or 0
        if element is not node:
            total += (px(style.get("padding-left")) or 0) + (_border(style, "left") or (0, "", ""))[0]
    return total


_BLOCK_CHILDREN = {"p", "div", "ul", "ol", "li", "table", "pre", "blockquote", "h1", "h2", "h3", "h4", "h5", "h6"}


def _measured_pitch(node, line: float, context: Context) -> float:
    """The line pitch the browser drew: the element's content height over its number of lines. A line with a
    subscript or superscript is taller than line-height says (the report's H2O line was 3 px taller, Windows check
    on 3095499). Only for an element with text and inline content alone, and only when it is near the CSS line."""
    box = context.styled.box(node) if context.styled is not None else None
    if box is None or any(getattr(c, "tag", None) in _BLOCK_CHILDREN for c in getattr(node, "children", [])):
        return line
    style = context.styled.style(node)
    height = box["height"] - sum(
        (px(style.get(name)) or 0)
        for name in ("padding-top", "padding-bottom", "border-top-width", "border-bottom-width")
    )
    lines = max(1, round(height / line))
    pitch = height / lines
    # A subscript and a superscript on one line made it up to 1.47 times the normal line in the browser; without
    # them, keep close to the CSS line.
    stack, raised = [node], False
    while stack and not raised:
        element = stack.pop()
        raised = element.tag in {"sub", "sup"}
        stack.extend(c for c in element.children if not isinstance(c, str))
    return pitch if line <= pitch < line * (1.5 if raised else 1.25) else line


def _fits_line(node, line: float, context: Context) -> bool:
    """Whether every piece of text in ``node`` is at most ``line`` px high and there is no image or object in it."""
    stack = [node]
    while stack:
        element = stack.pop()
        if element.tag in {"img", "svg", "canvas", "video", "iframe", "object", "embed", "input", "button"}:
            return False
        if (px(context.styled.style(element).get("font-size")) or 16.0) > line:
            return False
        stack.extend(child for child in element.children if not isinstance(child, str))
    return True


def _list_tab(paragraph, left) -> None:
    """The text after the number or bullet starts at the paragraph's indent: the template's list numbering has a
    tab stop of its own (360 twips at the first level), which a narrow bullet stops at, so bullet text started 16
    px left of the browser's in WPS (Windows check on f8ed7bd). Clear those stops and set one at the indent."""
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Twips

    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    style = paragraph.style
    num = style.element.find(f"{w}pPr/{w}numPr/{w}numId") if style is not None else None
    positions: set[int] = set()
    try:
        numbering = paragraph.part.numbering_part.element
    except (AttributeError, KeyError, NotImplementedError):
        numbering = None
    if num is not None and numbering is not None:
        number = num.get(f"{w}val")
        for entry in numbering.findall(f"{w}num"):
            if entry.get(f"{w}numId") != number:
                continue
            abstract = entry.find(f"{w}abstractNumId").get(f"{w}val")
            for definition in numbering.findall(f"{w}abstractNum"):
                if definition.get(f"{w}abstractNumId") == abstract:
                    for tab in definition.findall(f"{w}lvl/{w}pPr/{w}tabs/{w}tab"):
                        positions.add(int(tab.get(f"{w}pos")))
    stops = paragraph.paragraph_format.tab_stops
    for position in sorted(positions):
        if position != left.twips:
            stops.add_tab_stop(Twips(position), WD_TAB_ALIGNMENT.CLEAR)
    stops.add_tab_stop(left, WD_TAB_ALIGNMENT.LEFT)


def apply_paragraph(paragraph, node, ancestors, context: Context, *, list_item: bool = False) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.shared import Pt

    if context.styled is None or node is None:
        return
    style = context.styled.style(node)
    size = px(style.get("font-size")) or 16.0
    form = paragraph.paragraph_format
    align = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "end": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "start": WD_ALIGN_PARAGRAPH.LEFT,
    }.get(style.get("text-align", "start"))
    if align is not None:
        form.alignment = align
    # "normal" is the browser's measured line for the font, when it was laid out (line-height-normal).
    line = px(style.get("line-height"), size, size) or px(style.get("line-height-normal"))
    if line:
        # Exactly the CSS height when nothing in the paragraph is taller than a line: as "at least", WPS and Word
        # give an East Asian font its own taller line (Microsoft YaHei is about 1.7 em), so a 1.5 line came out
        # 1.3 times higher than in the browser. An image or a larger inline font keeps "at least", not to clip it.
        # data-line-rule="exact" (PDF conversion) places lines at a measured pitch either way.
        exact = getattr(node, "attrs", {}).get("data-line-rule") == "exact" or _fits_line(node, line, context)
        if exact and getattr(node, "attrs", {}).get("data-line-rule") != "exact":
            line = _measured_pitch(node, line, context)
        form.line_spacing_rule = WD_LINE_SPACING.EXACTLY if exact else WD_LINE_SPACING.AT_LEAST
        form.line_spacing = Pt(line * 0.75)
    else:
        # CSS "normal" is about 1.2 times the font size.
        form.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        form.line_spacing = 1.0
    top = (px(style.get("margin-top"), size) or 0) * 0.75
    bottom = (px(style.get("margin-bottom"), size) or 0) * 0.75
    # Adjacent vertical margins collapse in CSS: the larger one wins.
    form.space_before = Pt(max(0.0, max(top, context.carry_before) - context.last_space_after))
    form.space_after = Pt(bottom)
    context.last_space_after = bottom
    context.carry_before = 0.0
    if not list_item:
        # The text starts inside the element's own border and padding.
        offset = (_left_offset(context, node, ancestors) or 0.0) + (px(style.get("padding-left")) or 0)
        offset += (_border(style, "left") or (0, "", ""))[0]
        if offset:
            form.left_indent = Pt(offset * 0.75)
    elif getattr(node, "tag", "") == "li":
        # The item's text starts where the browser puts the <li> box (the lists' 40 px padding each), with the
        # number or bullet hanging in the padding before it.
        if context.styled.box(node) is not None:
            offset = _left_offset(context, node, ancestors) or 0.0
        else:
            offset = 40.0 * sum(1 for a in ancestors if a.tag in {"ul", "ol"})
        if offset:
            form.left_indent = Pt(offset * 0.75)
            # The browser puts the marker just before the text ("1. " or "• ", about 1 em), not 24 px before it;
            # a number with more digits needs more room.
            parent = ancestors[-1] if ancestors else None
            items = sum(1 for c in getattr(parent, "children", []) if getattr(c, "tag", None) == "li")
            digits = len(str(items)) if getattr(parent, "tag", "") == "ol" and items else 1
            hanging = size * (1.05 + 0.55 * (digits - 1))
            form.first_line_indent = Pt(-min(offset, hanging) * 0.75)
            _list_tab(paragraph, form.left_indent)
    indent = px(style.get("text-indent"), size)
    if indent:
        form.first_line_indent = Pt(indent * 0.75)
    if style.get("break-before") in {"page", "always", "left", "right"} or style.get("page-break-before") == "always":
        form.page_break_before = True
    props = paragraph._p.get_or_add_pPr()
    _no_grid(props)
    # A table cell's own background and borders belong to the cell (apply_table), not to its paragraphs.
    cell = getattr(node, "tag", "") in {"td", "th"}
    fill = color(style.get("background-color"))
    holder = node if fill else _container_with(context, node, ancestors, "background-color")
    if cell:
        holder = None
    if holder is not None:
        fill = color(context.styled.style(holder).get("background-color"))
        props.append(_element("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill}))
        context.applied.add("shading")
    border_holder = (
        node
        if any(_border(style, s) for s in ("top", "right", "bottom", "left"))
        else (_container_with(context, node, ancestors, "border"))
    )
    box_holder = None if cell else border_holder if border_holder is not None else holder
    borders: dict[str, tuple[float, str, str]] = {}
    if border_holder is not None and not cell:
        holder_style = context.styled.style(border_holder)
        borders = {side: b for side in ("top", "left", "bottom", "right") if (b := _border(holder_style, side))}
        context.applied.add("border")
    if box_holder is not None:
        _apply_padding(form, props, context.styled.style(box_holder), borders, fill)
    if context.appearance and style.get("position") in {"absolute", "fixed"}:
        box = context.styled.box(node)
        if box is not None:
            props.append(
                _element(
                    "w:framePr",
                    **{
                        "w:w": str(int(box["width"] * 15)),
                        "w:x": str(int(box["x"] * 15)),
                        "w:y": str(int(box["y"] * 15)),
                        "w:hAnchor": "page",
                        "w:vAnchor": "page",
                        "w:wrap": "around",
                    },
                )
            )
            context.applied.add("position")
    context.applied.add("paragraph")


def apply_table(table, node, grid, rows, context: Context, page_width_px: float) -> None:
    """Column widths, table width, cell shading, borders, padding and vertical alignment."""
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.shared import Emu, Pt

    # <thead> rows repeat at the top of each page, and say which rows are the header when read back.
    header_rows = sum(
        1
        for part in node.children
        if getattr(part, "tag", "") == "thead"
        for row in part.children
        if getattr(row, "tag", "") == "tr"
    )
    for row in table.rows[: min(header_rows, len(table.rows) - 1)]:
        row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    if context.styled is None:
        table.style = "Table Grid"
        return
    styled = context.styled
    table_style = styled.style(node)
    widths = _column_widths(styled, node, grid, page_width_px)
    if widths:
        table.autofit = False
        total = sum(widths)
        scale = min(1.0, page_width_px / total) if total else 1.0
        # Whole twips, rounded up: Word stores twips, and a width cut short by the conversion could wrap.
        sizes = [Emu(math.ceil(w * scale * 15) * 635) for w in widths]
        for row in table.rows:
            for index, cell in enumerate(row.cells[: len(sizes)]):
                cell.width = sizes[index]
        for index, column in enumerate(table.columns):
            if index < len(sizes):
                column.width = sizes[index]
        context.applied.add("column widths")
    if table_style.get("margin-left") == "auto" or node.attrs.get("align") == "center":
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
    else:
        # The browser starts the table's border at the content edge; Word (compatibility mode 14) starts it
        # one cell margin to the left. Indent by the table's cell margin, set to the cells' own padding.
        first = next((cell for row in grid for cell in row if cell is not None), None)
        pad = px(styled.style(first).get("padding-left")) if first is not None else None
        twips = str(int(round((pad if pad is not None else 1.0) * 15)))
        properties = table._tbl.tblPr
        for old in properties.findall(_qn("w:tblInd")) + properties.findall(_qn("w:tblCellMar")):
            properties.remove(old)
        margins = _element("w:tblCellMar")
        margins.append(_element("w:left", **{"w:w": twips, "w:type": "dxa"}))
        margins.append(_element("w:right", **{"w:w": twips, "w:type": "dxa"}))
        indent = _element("w:tblInd", **{"w:w": twips, "w:type": "dxa"})
        # Schema order in tblPr: ... tblW, jc, tblCellSpacing, tblInd, tblBorders, shd, tblLayout, tblCellMar, tblLook
        anchor = next(
            (
                properties.find(_qn(tag))
                for tag in ("w:tblBorders", "w:shd", "w:tblLayout", "w:tblLook")
                if properties.find(_qn(tag)) is not None
            ),
            None,
        )
        if anchor is not None:
            anchor.addprevious(indent)
        else:
            properties.append(indent)
        look = next(
            (properties.find(_qn(t)) for t in ("w:tblLook", "w:tblCaption") if properties.find(_qn(t)) is not None),
            None,
        )
        if look is not None:
            look.addprevious(margins)
        else:
            properties.append(margins)
    from docx.enum.table import WD_ROW_HEIGHT_RULE

    # With separate borders (the default) the browser puts border-spacing (2 px) between rows, and Word has no
    # equivalent outside cell spacing: count it in each row's height (the report's rows came out 2 px short each in
    # WPS, Windows check on f8ed7bd).
    spacing = 0.0
    collapsed = table_style.get("border-collapse", "separate") == "collapse"
    if not collapsed:
        values = table_style.get("border-spacing", "2px").split()
        spacing = (px(values[-1]) or 0.0) if values else 0.0
    for row, row_node in zip(table.rows, rows, strict=False):
        row_height = px(styled.style(row_node).get("height"))
        if row_height:
            row.height, row.height_rule = Pt((row_height + spacing) * 0.75), WD_ROW_HEIGHT_RULE.AT_LEAST
    has_borders = node.attrs.get("border") not in {None, "", "0"}
    for row_index, row in enumerate(grid):
        for column_index, cell_node in enumerate(row):
            if cell_node is None:
                continue
            cell = table.cell(row_index, column_index)
            style = styled.style(cell_node)
            row_node = rows[row_index] if row_index < len(rows) else None
            fill = color(style.get("background-color")) or (
                color(styled.style(row_node).get("background-color")) if row_node is not None else None
            )
            props = cell._tc.get_or_add_tcPr()
            # Schema order in tcPr: tcW, gridSpan, vMerge, tcBorders, shd, noWrap, tcMar, ..., vAlign.
            borders = {side: b for side in ("top", "left", "bottom", "right") if (b := _border(style, side))}
            if not borders and has_borders:
                borders = {side: (1.0, "single", "808080") for side in ("top", "left", "bottom", "right")}
            if borders:
                props.append(_borders_xml("w:tcBorders", borders))
                context.applied.add("cell borders")
            if fill:
                props.append(_element("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill}))
                context.applied.add("cell shading")
            margins = _element("w:tcMar")
            for side in ("left", "right"):
                value = px(style.get(f"padding-{side}"))
                if value is not None:
                    margins.append(_element(f"w:{side}", **{"w:w": str(int(value * 15)), "w:type": "dxa"}))
            if len(margins):
                props.append(margins)
            # Top and bottom padding become space before the first paragraph and after the last: WPS made every
            # row with a top/bottom cell margin about 7 pt taller than the browser (Windows check on c55d04e),
            # while rows without one (the PDF tables) matched.
            top, bottom = px(style.get("padding-top")), px(style.get("padding-bottom"))
            if collapsed and bottom and "bottom" in borders:
                # Collapsed borders are shared between rows in the browser, while WPS adds each cell's border to
                # its row: every bordered row came out a pixel taller (Windows check on c149caf).
                bottom = max(0.0, bottom - borders["bottom"][0])
            if top:
                cell.paragraphs[0].paragraph_format.space_before = Pt(top * 0.75)
            if bottom:
                cell.paragraphs[-1].paragraph_format.space_after = Pt(bottom * 0.75)
            vertical = {
                "middle": WD_CELL_VERTICAL_ALIGNMENT.CENTER,
                "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM,
                "top": WD_CELL_VERTICAL_ALIGNMENT.TOP,
            }.get(style.get("vertical-align", ""))
            if vertical is not None:
                cell.vertical_alignment = vertical
    # Borders come from the CSS on each cell (or the border attribute): without them the browser draws no
    # grid lines, so the table has no table style either.
    table.style = None


def _column_widths(styled: StyledDocument, node, grid, page_width_px: float) -> list[float]:
    columns = len(grid[0]) if grid else 0
    widths: list[float | None] = [None] * columns
    for row in grid:
        for index, cell in enumerate(row):
            if cell is None or widths[index] is not None:
                continue
            if int(cell.attrs.get("colspan", "1") or "1") != 1:
                continue
            box = styled.box(cell)
            if box is not None:
                # The browser's cell box includes its padding and borders, as a Word cell width does; 1 px more
                # keeps a cell that only just fits from wrapping on a rounding (2 px a column made the contract
                # table 8 px wider in WPS, Windows check on c205851; half a pixel let "Windows 10 / 11" wrap).
                widths[index] = box["width"] + 1
                continue
            declared = cell.attrs.get("width") or styled.style(cell).get("width")
            if declared:
                parent = px(styled.style(node).get("width")) or page_width_px
                value = px(declared if not declared.isdigit() else f"{declared}px", 16, parent)
                if value:
                    widths[index] = value
    # <col style="width:..."> gives the widths of columns no single cell sized.
    cols = [
        col
        for group in node.children
        if getattr(group, "tag", "") == "colgroup"
        for col in group.children
        if getattr(col, "tag", "") == "col"
    ]
    for index, col in enumerate(cols[:columns]):
        if widths[index] is None:
            declared = col.attrs.get("width") or styled.style(col).get("width")
            value = px(declared if declared and not declared.isdigit() else f"{declared}px", 16) if declared else None
            if value:
                widths[index] = value
    if all(width is None for width in widths):
        return []
    known = [w for w in widths if w is not None]
    fallback = (sum(known) / len(known)) if known else 100.0
    return [w if w is not None else fallback for w in widths]


def image_size(context: Context, node) -> tuple[int | None, int | None]:
    """The image's width and height in EMU, from the browser box, the attributes or the style."""
    if context.styled is not None:
        box = context.styled.box(node)
        if box is not None and box["width"] > 0:
            return int(box["width"] * EMU_PER_PX), int(box["height"] * EMU_PER_PX)
        style = context.styled.style(node)
    else:
        style = {}
    width = px(node.attrs.get("width") + "px" if (node.attrs.get("width") or "").isdigit() else style.get("width"))
    height = px(node.attrs.get("height") + "px" if (node.attrs.get("height") or "").isdigit() else style.get("height"))
    return (int(width * EMU_PER_PX) if width else None, int(height * EMU_PER_PX) if height else None)


_AFTER_SNAP_TO_GRID = (
    "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
    "w:rPr", "w:sectPr", "w:pPrChange",
)  # fmt: skip


def _no_grid(props) -> None:
    """snapToGrid off on the paragraph itself: WPS ignores the document default and snaps CJK lines to the
    section's line pitch (18 pt), so a browser's 16.5 pt lines came out 10% taller (Windows check on c117637)."""
    if props.find(_qn("w:snapToGrid")) is not None:
        return
    element = _element("w:snapToGrid", **{"w:val": "0"})
    later = next((child for child in props if child.tag in {_qn(t) for t in _AFTER_SNAP_TO_GRID}), None)
    if later is None:
        props.append(element)
    else:
        later.addprevious(element)


def browser_defaults(document) -> None:
    """Paragraph defaults that lay text out as a browser does, which WPS otherwise does not.

    WPS (like Word with East Asian layout) adds space between Chinese text and Latin letters or digits
    (autoSpaceDE/autoSpaceDN) and snaps lines to the document grid; a browser does neither, so lines grow
    longer and pages overflow. The template's 10 pt after and 1.15 lines become the browser's 0 and single.
    """
    styles = document.styles.element
    defaults = styles.find(_qn("w:docDefaults"))
    if defaults is None:
        return
    paragraph = defaults.find(_qn("w:pPrDefault"))
    if paragraph is None:
        paragraph = _element("w:pPrDefault")
        defaults.append(paragraph)
    props = paragraph.find(_qn("w:pPr"))
    if props is None:
        props = _element("w:pPr")
        paragraph.append(props)
    for tag in ("w:snapToGrid", "w:autoSpaceDE", "w:autoSpaceDN", "w:spacing"):
        for old in props.findall(_qn(tag)):
            props.remove(old)
    # Schema order: ..., autoSpaceDE, autoSpaceDN, ..., snapToGrid, spacing, ... (ind, jc and later stay after).
    later = [child for child in props if child.tag in {_qn(t) for t in ("w:ind", "w:jc", "w:rPr")}]
    for child in later:
        props.remove(child)
    props.append(_element("w:autoSpaceDE", **{"w:val": "0"}))
    props.append(_element("w:autoSpaceDN", **{"w:val": "0"}))
    props.append(_element("w:snapToGrid", **{"w:val": "0"}))
    props.append(_element("w:spacing", **{"w:after": "0", "w:line": "240", "w:lineRule": "auto"}))
    props.extend(later)


def apply_rule(paragraph, node, context: Context) -> None:
    """An <hr>: a one-point line under its margins, not a text line with a border (the manual's rule sat 17 px
    low in WPS, Windows check on c149caf). Its margins collapse with the blocks around it, as in CSS."""
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Pt

    style = context.styled.style(node) if context.styled is not None else {}
    # The browser's default is 0.5em above and below.
    top = (px(style.get("margin-top"), 16.0) if style.get("margin-top") else 8.0) or 0.0
    bottom = (px(style.get("margin-bottom"), 16.0) if style.get("margin-bottom") else 8.0) or 0.0
    form = paragraph.paragraph_format
    form.line_spacing_rule, form.line_spacing = WD_LINE_SPACING.EXACTLY, Pt(1)
    form.space_before = Pt(max(0.0, top * 0.75 - context.last_space_after))
    form.space_after = Pt(bottom * 0.75)
    context.last_space_after = bottom * 0.75
    context.carry_before = 0.0
    _no_grid(paragraph._p.get_or_add_pPr())


def apply_table_spacing(table, node, context: Context) -> None:
    """A table's vertical margins: Word has no space around a table, so its margin-top goes after the paragraph
    before it and its margin-bottom before the paragraph after it, collapsing with theirs as in CSS (Windows check
    on c205851: the contract's signature table lost its 24 pt margin-top)."""
    if context.styled is None:
        return
    style = context.styled.style(node)
    size = px(style.get("font-size")) or 16.0
    top = (px(style.get("margin-top"), size) or 0) * 0.75
    bottom = (px(style.get("margin-bottom"), size) or 0) * 0.75
    previous = table._tbl.getprevious()
    if top and previous is not None and previous.tag == _qn("w:p"):
        spacing = previous.get_or_add_pPr().get_or_add_spacing()
        after = int(spacing.get(_qn("w:after")) or 0) / 20
        if top > after:
            spacing.set(_qn("w:after"), str(int(round(top * 20))))
    context.last_space_after = 0.0
    context.carry_before = bottom


def apply_page(document, styled: StyledDocument | None) -> float:
    """@page size and margins on the first section; returns the text width in CSS px."""
    from docx.shared import Mm

    browser_defaults(document)
    section = document.sections[0]
    # No document grid: its line pitch is what WPS snaps lines to.
    for grid in section._sectPr.findall(_qn("w:docGrid")):
        section._sectPr.remove(grid)
    page = styled.page if styled is not None else {}
    size = page.get("size", "").casefold().split()
    if size:
        named = next((_PAGE_SIZES[token] for token in size if token in _PAGE_SIZES), None)
        lengths = [px(token) for token in size if px(token)]
        if named:
            width, height = named
            if "landscape" in size:
                width, height = height, width
            section.page_width, section.page_height = Mm(width), Mm(height)
        elif len(lengths) == 2:
            section.page_width, section.page_height = Mm(lengths[0] * 25.4 / 96), Mm(lengths[1] * 25.4 / 96)
    else:
        # Without @page size the browser prints on A4 (html-render too), not the template's US Letter.
        section.page_width, section.page_height = Mm(210), Mm(297)
    # Without @page margins, the browser prints with 10 mm margins (and html-render uses the same), not Word's
    # 1 inch; the document follows the browser.
    sides = {k: page.get(f"margin-{k}") for k in ("top", "right", "bottom", "left")}
    # The body's own margin and padding (8 px by default) sit inside the printed page in the browser; the text
    # edge is past them (Windows check on c205851: without them WPS lines were 16 px wider and wrapped later).
    # Top and bottom stay the page's: the first block's space before already counts from the page top.
    body = next((n for n in _walk(styled.root) if n.tag == "body"), None) if styled is not None else None
    inset = {"left": 0.0, "right": 0.0}
    if body is not None:
        body_style = styled.style(body)
        for side in inset:
            inset[side] = (px(body_style.get(f"margin-{side}")) or 0) + (px(body_style.get(f"padding-{side}")) or 0)
    page_margins = {}
    for side in ("top", "right", "bottom", "left"):
        value = px(sides.get(side)) if sides.get(side) else None
        page_margins[side] = value if value is not None else PRINT_MARGIN_MM * 96 / 25.4
        setattr(section, f"{side}_margin", Mm((page_margins[side] + inset.get(side, 0.0)) * 25.4 / 96))
    # A page header (data-page="header") starts where the body's first line does in the browser, inside the body's
    # top margin, and a running footer sits on the bottom margin; Word's default half inch put the header and the
    # text below it 3 px low and the footer 10 px high (Windows check on 3f82375).
    top_inset = (px(styled.style(body).get("margin-top")) or 0) if body is not None else 0.0
    section.header_distance = Mm((page_margins["top"] + top_inset) * 25.4 / 96)
    section.footer_distance = Mm(page_margins["bottom"] * 25.4 / 96)
    width_emu = section.page_width - section.left_margin - section.right_margin
    return width_emu / EMU_PER_PX


def page_role(node) -> str | None:
    """'header' or 'footer' for elements marked as the page header or footer."""
    role = node.attrs.get("data-page", "").casefold()
    if role in {"header", "footer"}:
        return role
    classes = node.attrs.get("class", "").split()
    if "page-header" in classes:
        return "header"
    if "page-footer" in classes:
        return "footer"
    return None


def css_coverage(styled: StyledDocument | None) -> list[str]:
    """CSS the conversion cannot reproduce, for unsupported_css."""
    if styled is None:
        return []
    notes = set()
    text = styled.css_text
    if re.search(r"display\s*:\s*(flex|grid|inline-flex|inline-grid)", text):
        notes.add("flex and grid layout (paragraphs follow document order)")
    if re.search(r"position\s*:\s*(absolute|fixed)", text):
        notes.add("absolute positioning (kept only with --fidelity appearance)")
    if re.search(r"(box-shadow|border-radius|transform|filter|linear-gradient|radial-gradient)", text):
        notes.add("shadows, rounded corners, transforms, filters and gradients")
    if styled.mode == "cascade":
        notes.update(w.split(":", 1)[0] for w in styled.warnings if "Selectors" in w)
    return sorted(notes)
