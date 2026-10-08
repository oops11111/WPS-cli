from __future__ import annotations

import hashlib
from copy import deepcopy
from io import BytesIO
import json
import os
from pathlib import Path
import re
import tempfile
from typing import Any, Iterator
import zipfile


_TOKEN = re.compile(r"\{\{\s*([a-z_]+)\s*\}\}")
_ALLOWED_TOKENS = {"mode", "source_directory", "output_directory", "total", "processed", "passed", "failed", "cancelled", "files_table"}
_FILE_TOKENS = {"file_status", "file_source", "file_output", "file_source_sha256", "file_output_sha256", "file_errors"}
_MAX_TEMPLATE_BYTES = 1024 * 1024
_MAX_MANIFEST_BYTES = 10 * 1024 * 1024
_MAX_REPORT_BYTES = 10 * 1024 * 1024
_MAX_FILES = 100
_MAX_TEMPLATE_ARCHIVE_BYTES = 50 * 1024 * 1024
_MAX_TEMPLATE_ARCHIVE_ENTRIES = 2_000
_MAX_TOTAL_OUTPUT_BYTES = 500 * 1024 * 1024


def _sha256(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _cell(value: Any) -> str:
    text = str(value if value is not None else "")
    return (text.replace("\\", "\\\\").replace("|", "\\|").replace("`", "\\`")
            .replace("[", "\\[").replace("]", "\\]").replace("<", "&lt;").replace(">", "&gt;")
            .replace("\r", " ").replace("\n", " "))


def _docx_paragraphs(document) -> Iterator[Any]:
    def from_table(table):
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
                for nested in cell.tables:
                    yield from from_table(nested)

    yield from document.paragraphs
    for table in document.tables:
        yield from from_table(table)
    for section in document.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            yield from part.paragraphs
            for table in part.tables:
                yield from from_table(table)


def _replace_docx_paragraph(paragraph, values: dict[str, str], allowed: set[str]) -> tuple[bool, str | None]:
    combined = "".join(run.text for run in paragraph.runs)
    tokens = list(_TOKEN.finditer(combined))
    placeholders = re.findall(r"\{\{[^{}]{1,128}\}\}", combined)
    if len(placeholders) > len(tokens):
        return False, "Template contains a malformed placeholder."
    for match in tokens:
        if match.group(1) not in allowed:
            return False, f"Unsupported template token: {match.group(1)}"
    per_run_tokens = sum(len(_TOKEN.findall(run.text)) for run in paragraph.runs)
    if per_run_tokens != len(tokens):
        return False, "A placeholder is split across Word runs; keep each placeholder in one run."
    for run in paragraph.runs:
        run.text = _TOKEN.sub(lambda match: values[match.group(1)], run.text)
    return bool(tokens), None


def _render_docx_report(
    manifest: dict[str, Any], actual: dict[str, int], manifest_file: Path,
    template_file: Path, output_file: Path, manifest_bytes: bytes, template_bytes: bytes,
) -> tuple[bool, dict[str, Any], list[dict[str, str]]]:
    try:
        from docx import Document
        from docx.table import _Row
        document = Document(BytesIO(template_bytes))
    except ImportError:
        return False, {}, [{"code": "CONVERTER_UNAVAILABLE", "message": "Install the html optional dependency: python-docx."}]
    except Exception as exc:  # noqa: BLE001
        return False, {}, [{"code": "INVALID_TEMPLATE", "message": str(exc)[:500]}]

    template_rows = []
    for table in document.tables:
        for row in table.rows:
            row_text = "\n".join(cell.text for cell in row.cells)
            if re.search(r"\{\{\s*file_[a-z_]+\s*\}\}", row_text):
                template_rows.append((table, row._tr))
    if len(template_rows) > 1:
        return False, {}, [{"code": "INVALID_TEMPLATE", "message": "DOCX template must contain at most one repeated file row."}]
    if manifest["files"] and not template_rows:
        return False, {}, [{"code": "INVALID_TEMPLATE", "message": "DOCX template needs one table row with {{file_*}} placeholders."}]

    output_root = Path(manifest["output_directory"]).resolve()
    template_row_element = template_rows[0][1] if template_rows else None
    if template_row_element is not None:
        table = template_rows[0][0]
        for item in manifest["files"]:
            clone = deepcopy(template_row_element)
            template_row_element.addprevious(clone)
            row = _Row(clone, table)
            output_relative = Path(item["output"]).resolve().relative_to(output_root).as_posix()
            row_values = {
                "mode": str(manifest["mode"]), "source_directory": str(manifest["source_directory"]),
                "output_directory": str(manifest["output_directory"]),
                **{key: str(value) for key, value in actual.items()},
                "processed": str(len(manifest["files"])), "cancelled": str(bool(manifest.get("cancelled", False))).lower(),
                "file_status": item["status"], "file_source": str(item.get("relative_source", "")),
                "file_output": output_relative, "file_source_sha256": item["source_sha256"],
                "file_output_sha256": item.get("output_sha256", ""),
                "file_errors": "; ".join(error.get("message", error.get("code", "")) for error in item.get("errors", [])),
            }
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _, error = _replace_docx_paragraph(paragraph, row_values, _ALLOWED_TOKENS | _FILE_TOKENS)
                    if error:
                        return False, {}, [{"code": "INVALID_TEMPLATE", "message": error}]
        template_row_element.getparent().remove(template_row_element)

    global_values = {
        "mode": str(manifest["mode"]), "source_directory": str(manifest["source_directory"]),
        "output_directory": str(manifest["output_directory"]),
        **{key: str(value) for key, value in actual.items()},
        "processed": str(len(manifest["files"])), "cancelled": str(bool(manifest.get("cancelled", False))).lower(),
    }
    replaced_count = 0
    for paragraph in _docx_paragraphs(document):
        replaced, error = _replace_docx_paragraph(paragraph, global_values, _ALLOWED_TOKENS)
        if error:
            return False, {}, [{"code": "INVALID_TEMPLATE", "message": error}]
        replaced_count += int(replaced)
    if not replaced_count and not manifest["files"]:
        return False, {}, [{"code": "INVALID_TEMPLATE", "message": "DOCX template contains no supported summary placeholders."}]

    buffer = BytesIO()
    try:
        document.save(buffer)
    except Exception as exc:  # noqa: BLE001
        return False, {}, [{"code": "REPORT_RENDER_FAILED", "message": str(exc)[:500]}]
    report_bytes = buffer.getvalue()
    if len(report_bytes) > _MAX_REPORT_BYTES:
        return False, {}, [{"code": "REPORT_TOO_LARGE", "message": "Rendered report exceeds 10 MiB."}]
    temporary: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(prefix=f".{output_file.stem}.", suffix=".tmp", dir=output_file.parent, delete=False) as handle:
            temporary = Path(handle.name)
            handle.write(report_bytes)
        os.link(temporary, output_file)
    except OSError as exc:
        return False, {}, [{"code": "REPORT_WRITE_FAILED", "message": str(exc)[:500]}]
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return True, {
        "manifest_path": str(manifest_file), "manifest_sha256": _sha256(manifest_bytes),
        "template_path": str(template_file), "template_sha256": _sha256(template_bytes),
        "output_path": str(output_file), "output_bytes": len(report_bytes),
        "output_sha256": _sha256(report_bytes), "summary": actual,
        "file_count": len(manifest["files"]),
    }, []


def render_batch_template_report(
    manifest_path: str | Path,
    template_path: str | Path,
    output_path: str | Path,
) -> tuple[bool, dict[str, Any], list[dict[str, str]]]:
    manifest_file = Path(manifest_path).expanduser().resolve()
    template_file = Path(template_path).expanduser().resolve()
    output_file = Path(output_path).expanduser().resolve()
    if not manifest_file.is_file() or manifest_file.suffix.casefold() != ".json":
        return False, {}, [{"code": "INVALID_MANIFEST", "message": "Manifest must be an existing JSON file."}]
    template_extension = template_file.suffix.casefold()
    if not template_file.is_file() or template_extension not in {".md", ".markdown", ".txt", ".docx"}:
        return False, {}, [{"code": "INVALID_TEMPLATE", "message": "Template must be an existing Markdown, text, or DOCX file."}]
    allowed_outputs = {".docx"} if template_extension == ".docx" else {".md", ".markdown", ".txt"}
    if output_file.suffix.casefold() not in allowed_outputs:
        return False, {}, [{"code": "INVALID_OUTPUT", "message": "Output extension must match the template type (.docx or Markdown/text)."}]
    if output_file.exists():
        return False, {}, [{"code": "OUTPUT_ALREADY_EXISTS", "message": f"Refusing to overwrite report: {output_file}"}]
    if not output_file.parent.is_dir():
        return False, {}, [{"code": "OUTPUT_DIRECTORY_NOT_FOUND", "message": f"Output directory does not exist: {output_file.parent}"}]
    if manifest_file.stat().st_size > _MAX_MANIFEST_BYTES or template_file.stat().st_size > _MAX_TEMPLATE_BYTES:
        return False, {}, [{"code": "INPUT_TOO_LARGE", "message": "Manifest or template exceeds its size limit."}]

    manifest_bytes = manifest_file.read_bytes()
    template_bytes = template_file.read_bytes()
    if template_extension == ".docx":
        try:
            with zipfile.ZipFile(BytesIO(template_bytes)) as archive:
                infos = archive.infolist()
                expanded_size = sum(info.file_size for info in infos)
                if len(infos) > _MAX_TEMPLATE_ARCHIVE_ENTRIES or expanded_size > _MAX_TEMPLATE_ARCHIVE_BYTES:
                    return False, {}, [{"code": "TEMPLATE_ARCHIVE_TOO_LARGE", "message": "DOCX template exceeds archive entry or 50 MiB expanded-size limits."}]
                if any(info.file_size and (info.compress_size == 0 or info.file_size / info.compress_size > 1_000) for info in infos):
                    return False, {}, [{"code": "TEMPLATE_COMPRESSION_RATIO_EXCEEDED", "message": "DOCX template contains an entry with an unsafe compression ratio."}]
                if "word/document.xml" not in archive.namelist() or "[Content_Types].xml" not in archive.namelist():
                    return False, {}, [{"code": "INVALID_TEMPLATE", "message": "DOCX template is missing required Word package parts."}]
        except (OSError, zipfile.BadZipFile) as exc:
            return False, {}, [{"code": "INVALID_TEMPLATE", "message": str(exc)[:500]}]
    try:
        manifest = json.loads(manifest_bytes.decode("utf-8-sig"))
        template = template_bytes.decode("utf-8-sig") if template_extension != ".docx" else None
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        return False, {}, [{"code": "INVALID_INPUT", "message": str(exc)[:500]}]
    if not isinstance(manifest, dict) or manifest.get("schema_version") != "wps-agent-batch-conversion/v1":
        return False, {}, [{"code": "UNSUPPORTED_MANIFEST", "message": "Expected a wps-agent-batch-conversion/v1 manifest."}]
    if manifest.get("mode") not in {"pdf", "png", "docx"}:
        return False, {}, [{"code": "INVALID_MANIFEST", "message": "Manifest mode is unsupported."}]
    files = manifest.get("files")
    summary = manifest.get("summary")
    if not isinstance(files, list) or len(files) > _MAX_FILES or not isinstance(summary, dict):
        return False, {}, [{"code": "INVALID_MANIFEST", "message": "Manifest file list or summary has an invalid shape."}]
    if any(
        not isinstance(item, dict) or item.get("status") not in {"passed", "failed"}
        or not isinstance(item.get("errors", []), list)
        or any(not isinstance(error, dict) for error in item.get("errors", []))
        for item in files
    ):
        return False, {}, [{"code": "INVALID_MANIFEST", "message": "Every manifest file entry must have passed or failed status."}]
    processed = len(files)
    actual = {
        "total": summary.get("total"),
        "passed": sum(item["status"] == "passed" for item in files),
        "failed": sum(item["status"] == "failed" for item in files),
    }
    cancelled = bool(manifest.get("cancelled", summary.get("cancelled", False)))
    reported_processed = summary.get("processed", processed)
    if (
        any(summary.get(key) != value for key, value in actual.items())
        or reported_processed != processed
        or not isinstance(actual["total"], int)
        or actual["total"] < processed
        or (not cancelled and actual["total"] != processed)
        or (manifest.get("cancelled") is not None and bool(manifest.get("cancelled")) != bool(summary.get("cancelled", manifest.get("cancelled"))))
    ):
        return False, {}, [{"code": "INVALID_MANIFEST", "message": "Manifest summary does not match its per-file results."}]
    source_root = Path(str(manifest.get("source_directory", ""))).resolve()
    output_root = Path(str(manifest.get("output_directory", ""))).resolve()
    total_source_bytes = 0
    total_output_bytes = 0
    verified_paths: list[tuple[dict[str, Any], Path, Path, Path]] = []
    for item in files:
        source_hash = item.get("source_sha256")
        if not isinstance(source_hash, str) or not re.fullmatch(r"[0-9a-f]{64}", source_hash):
            return False, {}, [{"code": "INVALID_MANIFEST", "message": "A file entry has no valid source SHA-256 digest."}]
        relative = Path(str(item.get("relative_source", "")))
        source = Path(str(item.get("source", ""))).resolve()
        expected_source = (source_root / relative).resolve()
        relative_output = relative.with_suffix(f".{manifest['mode']}")
        output = Path(str(item.get("output", ""))).resolve()
        expected_output = (output_root / relative_output).resolve()
        if source != expected_source or output != expected_output or relative.is_absolute() or ".." in relative.parts:
            return False, {}, [{"code": "INVALID_MANIFEST", "message": "A source or output path escapes its declared batch directory."}]
        if not source.is_file() or source.is_symlink():
            return False, {}, [{"code": "SOURCE_UNAVAILABLE", "message": f"Source file is unavailable for provenance verification: {relative.as_posix()}"}]
        source_size = source.stat().st_size
        total_source_bytes += source_size
        if source_size > _MAX_TEMPLATE_BYTES * 10 or total_source_bytes > 500 * 1024 * 1024:
            return False, {}, [{"code": "INPUT_TOO_LARGE", "message": "Current source files exceed batch provenance limits."}]
        if item["status"] == "passed":
            output_hash = item.get("output_sha256")
            if not isinstance(output_hash, str) or not re.fullmatch(r"[0-9a-f]{64}", output_hash):
                return False, {}, [{"code": "INVALID_MANIFEST", "message": "A successful file entry has no valid output SHA-256 digest."}]
            if not output.is_file() or output.is_symlink():
                return False, {}, [{"code": "OUTPUT_UNAVAILABLE", "message": f"Converted output is unavailable: {relative_output.as_posix()}"}]
            total_output_bytes += output.stat().st_size
            if total_output_bytes > _MAX_TOTAL_OUTPUT_BYTES:
                return False, {}, [{"code": "OUTPUT_TOO_LARGE", "message": "Converted outputs exceed the 500 MiB provenance verification limit."}]
        verified_paths.append((item, source, output, relative_output))
    for item, source, output, relative_output in verified_paths:
        if _sha256_file(source) != item["source_sha256"]:
            return False, {}, [{"code": "SOURCE_HASH_MISMATCH", "message": f"Source changed after conversion: {item.get('relative_source', '')}"}]
        if item["status"] == "passed" and _sha256_file(output) != item["output_sha256"]:
            return False, {}, [{"code": "OUTPUT_HASH_MISMATCH", "message": f"Converted output changed after conversion: {relative_output.as_posix()}"}]

    if template_extension == ".docx":
        return _render_docx_report(manifest, actual, manifest_file, template_file, output_file, manifest_bytes, template_bytes)

    rows = ["| Status | Source | Output | Source SHA-256 | Output SHA-256 | Details |", "| --- | --- | --- | --- | --- | --- |"]
    for item in files:
        details = "; ".join(error.get("message", error.get("code", "")) for error in item.get("errors", []) if isinstance(error, dict))
        rows.append("| " + " | ".join(_cell(value) for value in (
            item["status"], item.get("relative_source"), item.get("output"),
            item.get("source_sha256"), item.get("output_sha256", ""), details,
        )) + " |")
    values = {
        "mode": str(manifest.get("mode", "")),
        "source_directory": str(manifest.get("source_directory", "")),
        "output_directory": str(manifest.get("output_directory", "")),
        **{key: str(value) for key, value in actual.items()},
        "processed": str(processed), "cancelled": str(cancelled).lower(),
        "files_table": "\n".join(rows),
    }
    tokens = set(_TOKEN.findall(template))
    unknown = sorted(tokens - _ALLOWED_TOKENS)
    if unknown:
        return False, {}, [{"code": "UNKNOWN_TEMPLATE_TOKEN", "message": "Unsupported template tokens: " + ", ".join(unknown)}]
    report = _TOKEN.sub(lambda match: values[match.group(1)], template)
    if re.search(r"\{\{[^{}]{1,128}\}\}", report):
        return False, {}, [{"code": "INVALID_TEMPLATE_TOKEN", "message": "Template contains an unresolved or malformed placeholder."}]
    report_bytes = report.encode("utf-8")
    if len(report_bytes) > _MAX_REPORT_BYTES:
        return False, {}, [{"code": "REPORT_TOO_LARGE", "message": "Rendered report exceeds 10 MiB."}]

    temporary: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(prefix=f".{output_file.stem}.", suffix=".tmp", dir=output_file.parent, delete=False) as handle:
            temporary = Path(handle.name)
            handle.write(report_bytes)
        os.link(temporary, output_file)
    except OSError as exc:
        return False, {}, [{"code": "REPORT_WRITE_FAILED", "message": str(exc)[:500]}]
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return True, {
        "manifest_path": str(manifest_file), "manifest_sha256": _sha256(manifest_bytes),
        "template_path": str(template_file), "template_sha256": _sha256(template_bytes),
        "output_path": str(output_file), "output_bytes": len(report_bytes),
        "output_sha256": _sha256(report_bytes), "summary": actual,
        "file_count": len(files),
    }, []
