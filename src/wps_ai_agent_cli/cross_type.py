"""Conversions between document kinds (P13-013): DOCX <-> PPTX, and a worksheet range into a DOCX or PPTX table.

Each one builds clean HTML first (office_html) and hands it to html-editable, so the formatting rules are the ones
the HTML conversions already follow:

* DOCX -> PPTX: every Heading 1 or Heading 2 starts a slide with that title (decision 8); the paragraphs, lists,
  tables and images up to the next one are its body, and a body longer than eight lines goes on to a continuation
  slide titled "<title> (续)". Text before the first heading goes on a slide titled with the document title.
* PPTX -> DOCX: every slide becomes a Heading 1 with its text, tables and pictures below in reading order (top to
  bottom, then left to right) and its speaker notes in a paragraph headed "备注".
* XLSX range -> DOCX or PPTX: the range (the used range by default) becomes one native table with the cells'
  displayed values, merged cells and cell formatting.
"""

from __future__ import annotations

import html
import re
import tempfile
from pathlib import Path
from typing import Any

from .conversion_paths import check_paths, missing_dependency

MAX_LINES_PER_SLIDE = 8


def _run_html_conversion(page: str, folder: Path, destination: Path, use_browser: bool = False):
    from .html_office import convert_html_to_office

    handle = tempfile.NamedTemporaryFile(
        "w", encoding="utf-8", dir=folder, prefix=".wps-cross-", suffix=".html", delete=False
    )
    try:
        with handle:
            handle.write(page)
        return convert_html_to_office(
            handle.name, destination, destination.suffix.casefold().lstrip("."), use_browser=use_browser
        )
    finally:
        Path(handle.name).unlink(missing_ok=True)


def _page(title: str, body: str) -> str:
    from .office_html import _page as page

    return page(title, body)


# --------------------------------------------------------------------------------------------- docx -> pptx


def _top_level_blocks(body: str) -> list[str]:
    """The top-level elements of an HTML fragment, as strings."""
    from .html_editable import _Node, _TreeParser

    parser = _TreeParser()
    parser.feed(body)
    parser.close()
    blocks = []
    for child in parser.root.children:
        if isinstance(child, _Node):
            blocks.append(_serialise(child))
        elif child.strip():
            blocks.append(f"<p>{html.escape(child.strip(), quote=False)}</p>")
    return blocks


def _serialise(node) -> str:
    from .html_editable import _VOID, _Node

    attrs = "".join(f' {k}="{html.escape(v, quote=True)}"' for k, v in node.attrs.items())
    if node.tag in _VOID:
        return f"<{node.tag}{attrs}>"
    inner = "".join(_serialise(c) if isinstance(c, _Node) else html.escape(c, quote=False) for c in node.children)
    return f"<{node.tag}{attrs}>{inner}</{node.tag}>"


def _lines(block: str) -> int:
    """How many slide lines a block takes: one per paragraph or list item, a table or image counts as four."""
    if block.startswith(("<table", "<img")) or "<img" in block[:200]:
        return 4
    return max(1, block.count("<li"))


def _split_list(block: str, room: int) -> tuple[str, str] | None:
    """A top-level list cut after ``room`` items, or None."""
    match = re.match(r"<(ul|ol)([^>]*)>(.*)</\1>$", block, flags=re.S)
    if not match:
        return None
    kind, attrs, inner = match.groups()
    inner = re.sub(r">\s+<", "><", inner.strip())
    items = re.findall(r"<li\b.*?</li>", inner, flags=re.S)
    if len(items) <= room or "".join(items) != inner:
        return None
    return (
        f"<{kind}{attrs}>{''.join(items[:room])}</{kind}>",
        f"<{kind}{attrs}>{''.join(items[room:])}</{kind}>",
    )


def docx_to_pptx(input_path: str | Path, output_path: str | Path) -> tuple[bool, dict[str, Any], list[dict[str, str]]]:
    source, destination, errors = check_paths(input_path, output_path, {".docx"}, {".pptx"})
    if errors:
        return False, {}, errors
    try:
        from docx import Document
    except ImportError:
        return missing_dependency("html", "python-docx")
    from .office_html import _Images, docx_to_html

    document = Document(source)
    body, _counts = docx_to_html(document, _Images(destination, "embed"))
    title = document.core_properties.title or source.stem
    slides: list[tuple[str, list[str]]] = []
    for block in _top_level_blocks(body):
        heading = re.match(r"<h([12])\b[^>]*>(.*)</h\1>$", block, flags=re.S)
        if heading:
            slides.append((heading.group(2), []))
            continue
        if not slides:
            slides.append((html.escape(title, quote=False), []))
        if re.match(r"<h[3-6]\b", block):
            block = re.sub(r"^<h[3-6]([^>]*)>(.*)</h[3-6]>$", r"<p\1><strong>\2</strong></p>", block, flags=re.S)
        slides[-1][1].append(block)
    sections = []
    continued = 0
    for slide_title, blocks in slides:
        pages: list[list[str]] = [[]]
        used = 0
        queue = list(blocks)
        while queue:
            block = queue.pop(0)
            size = _lines(block)
            if used + size > MAX_LINES_PER_SLIDE and pages[-1]:
                split = _split_list(block, MAX_LINES_PER_SLIDE - used)
                if split:
                    pages[-1].append(split[0])
                    queue.insert(0, split[1])
                pages.append([])
                used = 0
                if split:
                    continue
            elif size > MAX_LINES_PER_SLIDE:
                split = _split_list(block, MAX_LINES_PER_SLIDE)
                if split:
                    pages[-1].append(split[0])
                    queue.insert(0, split[1])
                    pages.append([])
                    used = 0
                    continue
            pages[-1].append(block)
            used += size
        for number, page in enumerate(pages):
            name = slide_title if number == 0 else f"{slide_title} (续)"
            continued += number > 0
            sections.append(f"<section><h1>{name}</h1>\n" + "\n".join(page) + "\n</section>")
    ok, data, errors = _run_html_conversion(
        _page(title, "\n".join(sections)), source.parent, destination, use_browser=False
    )
    if ok:
        data["input_path"] = str(source)
        data["continuation_slides"] = continued
    return ok, data, errors


# --------------------------------------------------------------------------------------------- pptx -> docx


def pptx_to_docx(input_path: str | Path, output_path: str | Path) -> tuple[bool, dict[str, Any], list[dict[str, str]]]:
    source, destination, errors = check_paths(input_path, output_path, {".pptx"}, {".docx"})
    if errors:
        return False, {}, errors
    try:
        from pptx import Presentation
    except ImportError:
        return missing_dependency("html", "python-pptx")
    from .office_html import _Images, _pptx_runs, _text_frame_html

    presentation = Presentation(source)
    images = _Images(destination, "embed")
    parts: list[str] = []
    for number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None else ""
        parts.append(f"<h1>{html.escape(title or f'幻灯片 {number}', quote=False)}</h1>")
        others = [s for s in slide.shapes if title_shape is None or s.shape_id != title_shape.shape_id]
        for shape in sorted(others, key=lambda s: ((s.top or 0) // 95250, s.left or 0)):
            if shape.has_text_frame and shape.text_frame.text.strip():
                inner, _count = _text_frame_html(shape, {}, {}, False)
                parts.append(inner)
            elif shape.has_table:
                rows = []
                for r, row in enumerate(shape.table.rows):
                    cells = []
                    for cell in row.cells:
                        if cell.is_spanned:
                            continue
                        tag = "th" if r == 0 and shape.table.first_row else "td"
                        span = (f' colspan="{cell.span_width}"' if cell.span_width > 1 else "") + (
                            f' rowspan="{cell.span_height}"' if cell.span_height > 1 else ""
                        )
                        inner = "<br>".join(_pptx_runs(p) for p in cell.text_frame.paragraphs)
                        cells.append(f"<{tag}{span}>{inner}</{tag}>")
                    rows.append("<tr>" + "".join(cells) + "</tr>")
                head = rows[:1] if shape.table.first_row else []
                body_rows = rows[1:] if shape.table.first_row else rows
                parts.append(
                    '<table border="1">'
                    + (f"<thead>{''.join(head)}</thead>" if head else "")
                    + f"<tbody>{''.join(body_rows)}</tbody></table>"
                )
            elif shape.shape_type == 13:
                image = shape.image
                width = round((shape.width or 0) / 9525)
                src = images.source(image.blob, image.content_type, image.filename or "")
                parts.append(f'<p><img src="{src}" alt="{html.escape(shape.name)}" style="width:{width}px"></p>')
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            notes = html.escape(slide.notes_slide.notes_text_frame.text.strip(), quote=False).replace("\n", "<br>")
            parts.append(f'<p style="color:#666666"><em>备注：</em>{notes}</p>')
    title = presentation.core_properties.title or source.stem
    ok, data, errors = _run_html_conversion(_page(title, "\n".join(parts)), source.parent, destination)
    if ok:
        data["input_path"] = str(source)
        data["slides"] = len(presentation.slides)
        data.setdefault("warnings", []).append(
            "Slide positions, animations and transitions have no place in a document and were omitted."
        )
    return ok, data, errors


# ----------------------------------------------------------------------------------------- xlsx range -> table


def xlsx_range_to_table(
    input_path: str | Path,
    output_path: str | Path,
    sheet_name: str | None = None,
    cell_range: str | None = None,
    header_rows: int | None = None,
) -> tuple[bool, dict[str, Any], list[dict[str, str]]]:
    source, destination, errors = check_paths(input_path, output_path, {".xlsx", ".xlsm"}, {".docx", ".pptx"})
    if errors:
        return False, {}, errors
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries
    except ImportError:
        return missing_dependency("html", "openpyxl")
    from .office_html import sheet_table_html

    values = load_workbook(source, data_only=True)
    formulas = load_workbook(source, data_only=False)
    if sheet_name and sheet_name not in values.sheetnames:
        return (
            False,
            {},
            [{"code": "SHEET_NOT_FOUND", "message": f"No sheet {sheet_name!r}; sheets: {values.sheetnames}"}],
        )
    sheet = values[sheet_name or values.sheetnames[0]]
    bounds = None
    if cell_range:
        try:
            min_col, min_row, max_col, max_row = range_boundaries(cell_range.replace("$", ""))
        except (ValueError, TypeError):
            return False, {}, [{"code": "INVALID_RANGE", "message": f"Not a cell range: {cell_range!r} (e.g. A1:D10)."}]
        if None in (min_col, min_row, max_col, max_row):
            return False, {}, [{"code": "INVALID_RANGE", "message": "Give a bounded range such as A1:D10."}]
        if (max_row - min_row + 1) * (max_col - min_col + 1) > 20000:
            return False, {}, [{"code": "RANGE_TOO_LARGE", "message": "The range has more than 20000 cells."}]
        bounds = (min_row, min_col, max_row, max_col)
    counts = {"cells": 0, "formulas": 0, "merged": 0, "uncached_formulas": 0}
    table = sheet_table_html(sheet, formulas[sheet.title], counts, bounds, header_rows)
    title = html.escape(f"{sheet.title}" + (f" {cell_range}" if cell_range else ""), quote=False)
    if destination.suffix.casefold() == ".pptx":
        body = f"<section><h1>{title}</h1>\n{table}\n</section>"
    else:
        body = f"<h2>{title}</h2>\n{table}"
    ok, data, errors = _run_html_conversion(_page(title, body), source.parent, destination)
    warnings = list(data.get("warnings", [])) if ok else []
    if counts["formulas"]:
        warnings.append(f"{counts['formulas']} formula cell(s) were copied as their displayed values.")
    if counts["uncached_formulas"]:
        warnings.append(
            f"{counts['uncached_formulas']} formula cell(s) have no saved result (the workbook was not recalculated "
            "by WPS or Excel) and are empty."
        )
    if ok:
        data.update(
            {
                "input_path": str(source),
                "sheet": sheet.title,
                "range": cell_range or "used range",
                "cells": counts["cells"],
                "merged_ranges": counts["merged"],
                "warnings": warnings,
            }
        )
    return ok, data, errors
