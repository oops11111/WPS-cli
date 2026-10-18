"""DOCX, PPTX and XLSX to clean, self-contained HTML (P13-005, P13-009).

The HTML is meant to come back: html-editable turns it into the same kind of Office file again. Formatting is
written as inline styles on the elements (no class soup); images are data URIs, or files in ``<name>_files/`` with
``images="files"``. Slides become fixed-size ``<section>`` elements with absolutely positioned shapes, speaker notes
in ``<aside class="notes">``, animations as CSS ``@keyframes`` and transitions as ``data-transition``. Workbook
formulas keep their displayed value as the cell text and the formula in ``data-formula`` (decision 7).
"""

from __future__ import annotations

import base64
import html
import mimetypes
import re
from pathlib import Path
from typing import Any

from .conversion_paths import check_paths, missing_dependency

OFFICE_INPUTS = {".docx", ".pptx", ".xlsx", ".xlsm"}
EMU_PER_PX = 9525
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_HIGHLIGHT = {
    "yellow": "FFFF00",
    "green": "00FF00",
    "cyan": "00FFFF",
    "magenta": "FF00FF",
    "blue": "0000FF",
    "red": "FF0000",
    "darkBlue": "000080",
    "darkCyan": "008080",
    "darkGreen": "008000",
    "darkMagenta": "800080",
    "darkRed": "800000",
    "darkYellow": "808000",
    "darkGray": "808080",
    "lightGray": "C0C0C0",
    "black": "000000",
}


def _esc(text: str) -> str:
    return html.escape(text, quote=False)


def _attr(text: str) -> str:
    return html.escape(text, quote=True)


def _style(pairs: dict[str, Any]) -> str:
    body = ";".join(f"{k}:{v}" for k, v in pairs.items() if v not in (None, ""))
    return f' style="{_attr(body)}"' if body else ""


class _Images:
    """Embeds pictures as data URIs, or writes them next to the HTML."""

    def __init__(self, destination: Path, mode: str) -> None:
        self.mode = mode
        self.folder = destination.with_name(destination.stem + "_files")
        self.count = 0

    def source(self, blob: bytes, content_type: str | None, name: str = "") -> str:
        self.count += 1
        content_type = content_type or mimetypes.guess_type(name)[0] or "image/png"
        if self.mode == "files":
            self.folder.mkdir(exist_ok=True)
            extension = mimetypes.guess_extension(content_type) or Path(name).suffix or ".png"
            path = self.folder / f"image{self.count:03d}{extension}"
            path.write_bytes(blob)
            return f"{self.folder.name}/{path.name}"
        return f"data:{content_type};base64,{base64.b64encode(blob).decode('ascii')}"


def _page(title: str, body: str, css: str = "") -> str:
    font = "font-family:'Microsoft YaHei','Noto Sans CJK SC',sans-serif"
    return (
        '<!doctype html>\n<html lang="zh-CN">\n<head>\n<meta charset="utf-8">\n'
        f"<title>{_esc(title)}</title>\n<style>\nbody{{{font};margin:24px}}\n"
        "table{border-collapse:collapse}\n" + css + "</style>\n</head>\n<body>\n" + body + "</body>\n</html>\n"
    )


# ------------------------------------------------------------------------------------------------------ docx


def _style_chain(style):
    while style is not None:
        yield style
        style = style.base_style


def _inherited(run, paragraph, name: str):
    value = getattr(run.font, name)
    if value is not None:
        return value
    for style in _style_chain(
        run.style if run.style is not None and run.style.name != "Default Paragraph Font" else None
    ):
        value = getattr(style.font, name)
        if value is not None:
            return value
    for style in _style_chain(paragraph.style):
        value = getattr(style.font, name)
        if value is not None:
            return value
    return None


def _run_css(run, paragraph, heading: bool) -> dict[str, Any]:
    css: dict[str, Any] = {}
    props = run._r.rPr
    size = _inherited(run, paragraph, "size")
    if size is not None:
        css["font-size"] = f"{size.pt:g}pt"
    family = None
    if props is not None and props.rFonts is not None:
        family = props.rFonts.get(f"{_W}eastAsia") or props.rFonts.get(f"{_W}ascii")
    family = family or _inherited(run, paragraph, "name")
    if family:
        css["font-family"] = f"'{family}'"
    color = None
    try:
        if run.font.color is not None and run.font.color.type is not None and run.font.color.rgb is not None:
            color = str(run.font.color.rgb)
    except (AttributeError, ValueError):
        color = None
    if color is None and not heading:
        for style in _style_chain(paragraph.style):
            try:
                if style.font.color is not None and style.font.color.rgb is not None:
                    color = str(style.font.color.rgb)
                    break
            except (AttributeError, ValueError):
                break
    if color and color != "000000":
        css["color"] = f"#{color}"
    if props is not None:
        highlight = props.find(f"{_W}highlight")
        shading = props.find(f"{_W}shd")
        if highlight is not None and highlight.get(f"{_W}val") in _HIGHLIGHT:
            css["background-color"] = f"#{_HIGHLIGHT[highlight.get(f'{_W}val')]}"
        elif shading is not None and (shading.get(f"{_W}fill") or "auto") not in {"auto", "FFFFFF"}:
            css["background-color"] = f"#{shading.get(f'{_W}fill')}"
        spacing = props.find(f"{_W}spacing")
        if spacing is not None and spacing.get(f"{_W}val"):
            css["letter-spacing"] = f"{int(spacing.get(f'{_W}val')) / 20:g}pt"
    return css


def _wrap_run(text_html: str, run, paragraph, heading: bool) -> str:
    font = run.font
    bold = _inherited(run, paragraph, "bold")
    italic = _inherited(run, paragraph, "italic")
    underline = _inherited(run, paragraph, "underline")
    strike = _inherited(run, paragraph, "strike")
    out = text_html
    css = _run_css(run, paragraph, heading)
    if css:
        # Inside <sub>/<sup>, so that the run's own size wins over the browser's smaller script size.
        out = f"<span{_style(css)}>{out}</span>"
    if font.subscript:
        out = f"<sub>{out}</sub>"
    if font.superscript:
        out = f"<sup>{out}</sup>"
    if strike:
        out = f"<s>{out}</s>"
    if underline:
        out = f"<u>{out}</u>"
    if italic:
        out = f"<em>{out}</em>"
    if bold and not heading:
        out = f"<strong>{out}</strong>"
    elif bold is False and heading:
        out = f'<span style="font-weight:normal">{out}</span>'
    return out


def _run_html(run, paragraph, document, images: _Images, heading: bool) -> str:
    parts: list[str] = []
    for child in run._r.iterchildren():
        tag = child.tag
        if tag == f"{_W}t":
            parts.append(_esc(child.text or ""))
        elif tag == f"{_W}tab":
            parts.append("\t")
        elif tag in {f"{_W}br", f"{_W}cr"}:
            if child.get(f"{_W}type") == "page":
                parts.append('<span style="break-before:page"></span>')
            else:
                parts.append("<br>")
        elif tag == f"{_W}drawing":
            blip = child.find(f".//{_A}blip")
            extent = child.find(".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent")
            if blip is None:
                continue
            rel = blip.get(f"{_R}embed")
            part = document.part.related_parts.get(rel) if rel else None
            if part is None:
                continue
            size = {}
            if extent is not None:
                size["width"] = f"{round(int(extent.get('cx')) / EMU_PER_PX)}px"
                size["height"] = f"{round(int(extent.get('cy')) / EMU_PER_PX)}px"
            described = child.find(".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr")
            alt = described.get("descr") or described.get("name") or "" if described is not None else ""
            parts.append(
                f'<img src="{_attr(images.source(part.blob, part.content_type, str(part.partname)))}" '
                f'alt="{_attr(alt)}"{_style(size)}>'
            )
    text = "".join(parts)
    if not text:
        return ""
    if "<img" in text:
        return text
    return _wrap_run(text, run, paragraph, heading)


def _paragraph_inner(paragraph, document, images: _Images, heading: bool) -> str:
    from docx.text.run import Run

    parts: list[str] = []
    for child in paragraph._p.iterchildren():
        if child.tag == f"{_W}r":
            parts.append(_run_html(Run(child, paragraph), paragraph, document, images, heading))
        elif child.tag == f"{_W}hyperlink":
            rel = child.get(f"{_R}id")
            anchor = child.get(f"{_W}anchor")
            target = paragraph.part.rels[rel].target_ref if rel and rel in paragraph.part.rels else None
            href = target or (f"#{anchor}" if anchor else "")
            inner = "".join(
                _run_html(Run(r, paragraph), paragraph, document, images, heading) for r in child.findall(f"{_W}r")
            )
            parts.append(f'<a href="{_attr(href)}">{inner}</a>' if href else inner)
        elif child.tag in {f"{_W}ins", f"{_W}smartTag", f"{_W}fldSimple"}:
            for r in child.findall(f"{_W}r"):
                parts.append(_run_html(Run(r, paragraph), paragraph, document, images, heading))
    return "".join(parts)


def _paragraph_css(paragraph) -> dict[str, Any]:
    css: dict[str, Any] = {}
    format_ = paragraph.paragraph_format
    alignment = format_.alignment
    if alignment is None:
        for style in _style_chain(paragraph.style):
            if style.paragraph_format.alignment is not None:
                alignment = style.paragraph_format.alignment
                break
    if alignment is not None:
        css["text-align"] = {0: "left", 1: "center", 2: "right", 3: "justify"}.get(int(alignment))
    if format_.left_indent:
        css["margin-left"] = f"{format_.left_indent.pt:g}pt"
    if format_.first_line_indent:
        css["text-indent"] = f"{format_.first_line_indent.pt:g}pt"
    if format_.space_before is not None:
        css["margin-top"] = f"{format_.space_before.pt:g}pt"
    if format_.space_after is not None:
        css["margin-bottom"] = f"{format_.space_after.pt:g}pt"
    if format_.page_break_before:
        css["break-before"] = "page"
    shading = paragraph._p.pPr.find(f"{_W}shd") if paragraph._p.pPr is not None else None
    if shading is not None and (shading.get(f"{_W}fill") or "auto") not in {"auto", "FFFFFF"}:
        css["background-color"] = f"#{shading.get(f'{_W}fill')}"
    return css


def _list_info(paragraph, numbering) -> tuple[str, int] | None:
    """('ul' or 'ol', level) for a list paragraph."""
    props = paragraph._p.pPr
    num = props.numPr if props is not None else None
    style = paragraph.style.name if paragraph.style is not None else ""
    if num is not None and num.numId is not None and num.numId.val != 0:
        level = num.ilvl.val if num.ilvl is not None else 0
        kind = "ul"
        if numbering is not None:
            kind = "ol" if _num_format(numbering, num.numId.val, level) not in {"bullet", "none"} else "ul"
        return kind, level
    match = re.match(r"List (Bullet|Number)(?: (\d))?$", style)
    if match:
        return ("ul" if match.group(1) == "Bullet" else "ol"), int(match.group(2) or 1) - 1
    return None


def _num_format(numbering, num_id: int, level: int) -> str:
    root = numbering.element
    for num in root.findall(f"{_W}num"):
        if num.get(f"{_W}numId") != str(num_id):
            continue
        abstract = num.find(f"{_W}abstractNumId")
        if abstract is None:
            break
        for definition in root.findall(f"{_W}abstractNum"):
            if definition.get(f"{_W}abstractNumId") != abstract.get(f"{_W}val"):
                continue
            for lvl in definition.findall(f"{_W}lvl"):
                if lvl.get(f"{_W}ilvl") == str(level):
                    fmt = lvl.find(f"{_W}numFmt")
                    return fmt.get(f"{_W}val") if fmt is not None else "decimal"
    return "decimal"


def _heading_level(paragraph) -> int | None:
    style = paragraph.style.name if paragraph.style is not None else ""
    if style == "Title":
        return 1
    match = re.match(r"(?:Heading|标题)\s*(\d)$", style)
    if match:
        return min(6, int(match.group(1)))
    props = paragraph._p.pPr
    outline = props.find(f"{_W}outlineLvl") if props is not None else None
    if outline is not None and outline.get(f"{_W}val", "9").isdigit() and int(outline.get(f"{_W}val")) < 6:
        return int(outline.get(f"{_W}val")) + 1
    return None


def _cell_css(cell) -> dict[str, Any]:
    css: dict[str, Any] = {}
    props = cell._tc.tcPr
    if props is None:
        return css
    shading = props.find(f"{_W}shd")
    if shading is not None and (shading.get(f"{_W}fill") or "auto") not in {"auto", "FFFFFF"}:
        css["background-color"] = f"#{shading.get(f'{_W}fill')}"
    align = props.find(f"{_W}vAlign")
    if align is not None:
        css["vertical-align"] = {"center": "middle", "bottom": "bottom", "top": "top"}.get(align.get(f"{_W}val"))
    width = props.find(f"{_W}tcW")
    if width is not None and width.get(f"{_W}type") == "dxa" and width.get(f"{_W}w", "0").isdigit():
        css["width"] = f"{int(width.get(f'{_W}w')) / 20:g}pt"
    borders = props.find(f"{_W}tcBorders")
    if borders is not None:
        for side in ("top", "left", "bottom", "right"):
            edge = borders.find(f"{_W}{side}")
            if edge is not None and edge.get(f"{_W}val") not in {None, "nil", "none"}:
                size = int(edge.get(f"{_W}sz", "4")) / 8
                hue = edge.get(f"{_W}color", "000000")
                css[f"border-{side}"] = f"{size:g}pt solid #{'000000' if hue == 'auto' else hue}"
    return css


def _table_html(table, document, images: _Images, numbering) -> str:
    from docx.table import _Cell

    rows = table._tbl.tr_lst
    grid: list[list[Any]] = []
    for tr in rows:
        cells = []
        for tc in tr.tc_lst:
            props = tc.tcPr
            span = props.gridSpan.val if props is not None and props.gridSpan is not None else 1
            merge = props.vMerge.val if props is not None and props.vMerge is not None else None
            cells.append((tc, span, merge))
        grid.append(cells)
    borders = table._tbl.tblPr.find(f"{_W}tblBorders") if table._tbl.tblPr is not None else None
    bordered = borders is not None and any(
        edge.get(f"{_W}val") not in {None, "nil", "none"} for edge in borders.iterchildren()
    )
    style_name = table.style.name if table.style is not None else ""
    bordered = bordered or "Grid" in style_name
    out = ["<table" + (' border="1"' if bordered else "") + ">"]
    header_rows = 0
    for tr in rows:
        props = tr.trPr
        if props is not None and props.find(f"{_W}tblHeader") is not None:
            header_rows += 1
        else:
            break
    for r, cells in enumerate(grid):
        if r == 0 and header_rows:
            out.append("<thead>")
        if r == header_rows and header_rows:
            out.append("<tbody>")
        out.append("<tr>")
        column = 0
        for tc, span, merge in cells:
            if merge == "continue":
                column += span
                continue
            rowspan = 1
            if merge == "restart":
                for below in grid[r + 1 :]:
                    position = 0
                    found = False
                    for _other, other_span, other_merge in below:
                        if position == column:
                            found = other_merge == "continue"
                            break
                        position += other_span
                    if not found:
                        break
                    rowspan += 1
            cell = _Cell(tc, table)
            tag = "th" if r < header_rows else "td"
            inner = []
            for index, paragraph in enumerate(cell.paragraphs):
                content = _paragraph_inner(paragraph, document, images, False)
                if index:
                    inner.append("<br>")
                inner.append(content)
            for nested in cell.tables:
                inner.append(_table_html(nested, document, images, numbering))
            attrs = ""
            if span > 1:
                attrs += f' colspan="{span}"'
            if rowspan > 1:
                attrs += f' rowspan="{rowspan}"'
            css = _cell_css(cell)
            if cell.paragraphs and cell.paragraphs[0].paragraph_format.alignment is not None:
                css["text-align"] = {0: "left", 1: "center", 2: "right", 3: "justify"}.get(
                    int(cell.paragraphs[0].paragraph_format.alignment)
                )
            out.append(f"<{tag}{attrs}{_style(css)}>{''.join(inner)}</{tag}>")
            column += span
        out.append("</tr>")
        if r == header_rows - 1 and header_rows:
            out.append("</thead>")
    if header_rows and header_rows < len(grid):
        out.append("</tbody>")
    out.append("</table>")
    return "\n".join(out)


def docx_to_html(document, images: _Images) -> tuple[str, dict[str, int]]:
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        numbering = document.part.numbering_part
    except (KeyError, NotImplementedError):
        numbering = None
    out: list[str] = []
    counts = {"headings": 0, "paragraphs": 0, "lists": 0, "tables": 0}
    open_lists: list[str] = []

    def close_lists(to: int = 0) -> None:
        while len(open_lists) > to:
            out.append(f"</li></{open_lists.pop()}>")

    for child in document.element.body.iterchildren():
        if child.tag == f"{_W}p":
            paragraph = Paragraph(child, document)
            info = _list_info(paragraph, numbering)
            level = _heading_level(paragraph)
            inner = _paragraph_inner(paragraph, document, images, level is not None)
            if info is not None and inner:
                kind, depth = info
                if len(open_lists) > depth + 1:
                    close_lists(depth + 1)
                if len(open_lists) == depth + 1 and open_lists[-1] != kind:
                    close_lists(depth)
                if len(open_lists) == depth + 1:
                    out.append("</li>")
                while len(open_lists) < depth + 1:
                    out.append(f"<{kind}>")
                    open_lists.append(kind)
                    if len(open_lists) < depth + 1:
                        out.append('<li style="list-style:none">')
                out.append(f"<li{_style(_paragraph_css(paragraph))}>{inner}")
                counts["lists"] += 1
                continue
            close_lists()
            css = _paragraph_css(paragraph)
            style = paragraph.style.name if paragraph.style is not None else ""
            props = paragraph._p.pPr
            border = props.find(f"{_W}pBdr/{_W}bottom") if props is not None else None
            fonts = {r.font.name for r in paragraph.runs if r.text.strip()}
            if border is not None and not paragraph.text.strip() and "<img" not in inner:
                out.append("<hr>")
            elif style in {"Quote", "Intense Quote"} and inner.strip():
                counts["paragraphs"] += 1
                if out and out[-1].endswith("</blockquote>"):
                    out[-1] = out[-1][: -len("</blockquote>")] + f"<p>{inner}</p></blockquote>"
                else:
                    out.append(f"<blockquote><p>{inner}</p></blockquote>")
            elif fonts and fonts <= {"Consolas", "Courier New", "Menlo", "monospace"} and inner.strip():
                counts["paragraphs"] += 1
                code = html.escape(paragraph.text, quote=False)
                out.append(f"<pre><code>{code}</code></pre>")
            elif level is not None:
                counts["headings"] += 1
                out.append(f"<h{level}{_style(css)}>{inner}</h{level}>")
            elif inner.strip():
                counts["paragraphs"] += 1
                out.append(f"<p{_style(css)}>{inner}</p>")
            elif css.get("break-before"):
                out.append('<div style="break-before:page"></div>')
        elif child.tag == f"{_W}tbl":
            close_lists()
            counts["tables"] += 1
            out.append(_table_html(Table(child, document), document, images, numbering))
    close_lists()
    return "\n".join(out) + "\n", counts


# ------------------------------------------------------------------------------------------------------ pptx

_PRESET_EFFECTS = {
    ("entr", "10"): "fade",
    ("entr", "2"): "fly",
    ("entr", "53"): "zoom",
    ("exit", "10"): "fade",
    ("exit", "2"): "fly",
    ("exit", "53"): "zoom",
    ("emph", "6"): "pulse",
    ("emph", "8"): "spin",
}
_FLY_FROM = {"8": "left", "2": "right", "1": "top", "4": "bottom"}
_TRANSITION_NAMES = {"fade", "push", "wipe", "cover", "split", "dissolve", "zoom", "cut"}


def _keyframes(kind: str, effect: str, direction: str) -> str:
    hidden = {
        "fade": "opacity:0",
        "zoom": "opacity:0;transform:scale(0)",
        "fly": "opacity:0;transform:"
        + {
            "left": "translateX(-100%)",
            "right": "translateX(100%)",
            "top": "translateY(-100%)",
            "bottom": "translateY(100%)",
        }.get(direction, "translateX(-100%)"),
    }.get(effect, "opacity:0")
    if kind == "exit":
        return f"from{{opacity:1}}to{{{hidden}}}"
    if kind == "emphasis":
        return (
            "from{transform:rotate(0)}to{transform:rotate(360deg)}"
            if effect == "spin"
            else ("0%{transform:scale(1)}50%{transform:scale(1.1)}100%{transform:scale(1)}")
        )
    return f"from{{{hidden}}}to{{opacity:1}}"


def _slide_animations(slide) -> dict[tuple[int, int | None], tuple[str, str, str, int, int, str]]:
    """(shape id, paragraph or None) -> (kind, effect, direction, duration ms, delay ms, trigger)."""
    found: dict[tuple[int, int | None], tuple[str, str, str, int, int, str]] = {}
    timing = slide._element.find(f"{_P}timing")
    if timing is None:
        return found
    for node in timing.iter(f"{_P}cTn"):
        preset_class = node.get("presetClass")
        if preset_class not in {"entr", "exit", "emph"}:
            continue
        target = node.find(f".//{_P}spTgt")
        if target is None:
            continue
        paragraph = target.find(f".//{_P}pRg")
        key = (int(target.get("spid")), int(paragraph.get("st")) if paragraph is not None else None)
        kind = {"entr": "entrance", "exit": "exit", "emph": "emphasis"}[preset_class]
        effect = _PRESET_EFFECTS.get((preset_class, node.get("presetID")), "fade")
        direction = _FLY_FROM.get(node.get("presetSubtype", ""), "") if effect == "fly" else ""
        # The longest behaviour is the effect; a <p:set> of visibility lasts 1 ms.
        durations = [int(i.get("dur")) for i in node.iter(f"{_P}cTn") if i is not node and i.get("dur", "").isdigit()]
        duration = max(durations) if durations else 500
        condition = node.find(f"{_P}stCondLst/{_P}cond")
        delay = (
            int(condition.get("delay", "0")) if condition is not None and condition.get("delay", "").isdigit() else 0
        )
        trigger = {"clickEffect": "click", "afterEffect": "after"}.get(node.get("nodeType", ""), "with")
        found[key] = (kind, effect, direction, duration, delay, trigger)
    return found


def _pptx_runs(paragraph, title: bool = False) -> str:
    """A slide paragraph's runs as inline HTML with their bold, italic, underline, size, colour and font."""
    runs = []
    for run in paragraph.runs:
        text = _esc(run.text).replace("\v", "<br>")
        if not text:
            continue
        font = run.font
        if font.bold and not title:
            text = f"<strong>{text}</strong>"
        if font.italic:
            text = f"<em>{text}</em>"
        if font.underline:
            text = f"<u>{text}</u>"
        css: dict[str, Any] = {}
        if font.size is not None:
            css["font-size"] = f"{font.size.pt:g}pt"
        try:
            if font.color is not None and font.color.type is not None and font.color.rgb is not None:
                css["color"] = f"#{font.color.rgb}"
        except (AttributeError, ValueError):
            pass
        east = run._r.rPr.find(f"{_A}ea") if run._r.rPr is not None else None
        family = (east.get("typeface") if east is not None else None) or font.name
        if family and not family.startswith("+"):
            css["font-family"] = f"'{family}'"
        if css:
            text = f"<span{_style(css)}>{text}</span>"
        runs.append(text)
    return "".join(runs)


def _text_frame_html(shape, animations, names: dict[str, str], title: bool) -> tuple[str, int]:
    """The shape's paragraphs as HTML: a heading for a title, list items for bulleted paragraphs."""
    pieces: list[str] = []
    depth = 0
    paragraphs = 0
    for index, paragraph in enumerate(shape.text_frame.paragraphs):
        content = _pptx_runs(paragraph, title)
        if not content:
            continue
        paragraphs += 1
        props = paragraph._p.pPr
        bulleted = props is not None and props.find(f"{_A}buChar") is not None
        if not bulleted and props is not None and props.find(f"{_A}buNone") is None:
            try:
                bulleted = shape.is_placeholder and shape.placeholder_format.idx == 1
            except ValueError:
                bulleted = False
        css = {}
        if paragraph.alignment is not None:
            css["text-align"] = {1: "left", 2: "center", 3: "right", 4: "justify"}.get(int(paragraph.alignment))
        animation = animations.get((shape.shape_id, index))
        attrs = _style(css)
        if animation:
            attrs = _merge_style(attrs + _animation_attrs(animation, names))
        if title:
            pieces.append(f"<h2{attrs}>{content}</h2>")
        elif bulleted:
            level = min(paragraph.level, 4)
            if depth > level + 1:
                while depth > level + 1:
                    pieces.append("</li></ul>")
                    depth -= 1
            if depth == level + 1:
                pieces.append("</li>")
            while depth < level + 1:
                pieces.append("<ul>")
                depth += 1
                if depth < level + 1:
                    pieces.append('<li style="list-style:none">')
            pieces.append(f"<li{attrs}>{content}")
        else:
            while depth:
                pieces.append("</li></ul>")
                depth -= 1
            pieces.append(f"<p{attrs}>{content}</p>")
    while depth:
        pieces.append("</li></ul>")
        depth -= 1
    return "".join(pieces), paragraphs


def _animation_attrs(animation, names: dict[str, str]) -> str:
    kind, effect, direction, duration, delay, trigger = animation
    key = f"{kind}-{effect}-{direction}".rstrip("-")
    name = names.setdefault(key, f"wps-{key}")
    css = f"animation:{name} {duration}ms ease {delay}ms both"
    extra = f' data-trigger="{trigger}"' if trigger != "with" else ""
    return f' data-animation="{key}"{extra}' + f' style="{css}"'


def _merge_style(attrs: str) -> str:
    """Two style attributes from _style and _animation_attrs become one."""
    styles = re.findall(r' style="([^"]*)"', attrs)
    if len(styles) < 2:
        return attrs
    rest = re.sub(r' style="[^"]*"', "", attrs)
    return rest + f' style="{";".join(styles)}"'


def pptx_to_html(presentation, images: _Images) -> tuple[str, str, dict[str, int]]:
    width = presentation.slide_width / EMU_PER_PX
    height = presentation.slide_height / EMU_PER_PX
    names: dict[str, str] = {}
    sections: list[str] = []
    counts = {"slides": 0, "tables": 0, "images": 0, "animations": 0, "notes": 0}
    for slide in presentation.slides:
        counts["slides"] += 1
        animations = _slide_animations(slide)
        counts["animations"] += len(animations)
        transition = slide._element.find(f"{_P}transition")
        attrs = ""
        if transition is not None and len(transition):
            name = transition[0].tag.rsplit("}", 1)[-1]
            if name in _TRANSITION_NAMES:
                attrs += f' data-transition="{name}"'
        css = {"position": "relative", "width": f"{width:g}px", "height": f"{height:g}px", "overflow": "hidden"}
        background = slide._element.find(f"{_P}cSld/{_P}bg//{_A}srgbClr")
        if background is not None:
            css["background-color"] = f"#{background.get('val')}"
        body: list[str] = []
        title_id = slide.shapes.title.shape_id if slide.shapes.title is not None else None
        ordered = sorted(slide.shapes, key=lambda s: (s.shape_id != title_id, s.top or 0, s.left or 0))
        for shape in ordered:
            box = {
                "position": "absolute",
                "left": f"{(shape.left or 0) / EMU_PER_PX:.0f}px",
                "top": f"{(shape.top or 0) / EMU_PER_PX:.0f}px",
                "width": f"{(shape.width or 0) / EMU_PER_PX:.0f}px",
                "height": f"{(shape.height or 0) / EMU_PER_PX:.0f}px",
                "margin": "0",
            }
            animation = animations.get((shape.shape_id, None))
            anim_attrs = _animation_attrs(animation, names) if animation else ""
            if shape.has_text_frame and shape.text_frame.text.strip():
                inner, paragraphs = _text_frame_html(shape, animations, names, shape.shape_id == title_id)
                if anim_attrs and paragraphs == 1:
                    # One paragraph: the effect goes on it, where html-editable looks for a title's or a line's.
                    match = re.match(r"<(h2|p|ul><li)([^>]*)>", inner)
                    if match:
                        opening = f"<{match.group(1)}{_merge_style(match.group(2) + anim_attrs)}>"
                        inner = opening + inner[match.end() :]
                        anim_attrs = ""
                if shape.shape_id == title_id:
                    inner = inner.replace("<h2", "<h1", 1).replace("</h2>", "</h1>", 1)
                fill = None
                try:
                    if shape.fill.type == 1:
                        fill = f"#{shape.fill.fore_color.rgb}"
                except (AttributeError, TypeError, ValueError):
                    fill = None
                if fill:
                    box["background-color"] = fill
                body.append("<div" + _merge_style(_style(box) + anim_attrs) + ">" + inner + "</div>")
            elif shape.has_table:
                counts["tables"] += 1
                rows = []
                for r, row in enumerate(shape.table.rows):
                    cells = []
                    for cell in row.cells:
                        if cell.is_spanned:
                            continue
                        tag = "th" if r == 0 and shape.table.first_row else "td"
                        span = ""
                        if cell.span_width > 1:
                            span += f' colspan="{cell.span_width}"'
                        if cell.span_height > 1:
                            span += f' rowspan="{cell.span_height}"'
                        inner = "<br>".join(_pptx_runs(p) for p in cell.text_frame.paragraphs)
                        cells.append(f"<{tag}{span}>{inner}</{tag}>")
                    rows.append("<tr>" + "".join(cells) + "</tr>")
                table_css = dict(box)
                body.append(
                    '<table border="1"'
                    + _merge_style(_style(table_css) + anim_attrs)
                    + ">"
                    + "".join(rows)
                    + "</table>"
                )
            elif shape.shape_type == 13:
                counts["images"] += 1
                image = shape.image
                source = images.source(image.blob, image.content_type, image.filename or "")
                body.append(
                    f'<img src="{_attr(source)}" alt="{_attr(shape.name)}"'
                    + _merge_style(_style(box) + anim_attrs)
                    + ">"
                )
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            counts["notes"] += 1
            body.append(f'<aside class="notes">{_esc(slide.notes_slide.notes_text_frame.text)}</aside>')
        sections.append(f"<section{attrs}{_style(css)}>\n" + "\n".join(body) + "\n</section>")
    keyframes = "".join(f"@keyframes {name}{{{_keyframes(*_split_key(key))}}}\n" for key, name in sorted(names.items()))
    css = "section{margin:0 auto 24px;box-shadow:0 0 0 1px #ccc}\nh1,h2,p,ul{margin:0}\naside.notes{display:none}\n"
    return "\n".join(sections) + "\n", css + keyframes, counts


def _split_key(key: str) -> tuple[str, str, str]:
    parts = key.split("-")
    return parts[0], parts[1], parts[2] if len(parts) > 2 else ""


# ------------------------------------------------------------------------------------------------------ xlsx


def _xlsx_cell_css(cell) -> dict[str, Any]:
    css: dict[str, Any] = {}
    font = cell.font
    if font is not None:
        if font.b:
            css["font-weight"] = "bold"
        if font.i:
            css["font-style"] = "italic"
        decoration = [name for name, on in (("underline", font.u), ("line-through", font.strike)) if on]
        if decoration:
            css["text-decoration"] = " ".join(decoration)
        if font.color is not None and font.color.type == "rgb" and font.color.rgb not in {None, "FF000000"}:
            css["color"] = f"#{str(font.color.rgb)[-6:]}"
        if font.sz and abs(float(font.sz) - 11) > 0.01:
            css["font-size"] = f"{float(font.sz):g}pt"
    fill = cell.fill
    if fill is not None and fill.fill_type == "solid" and fill.start_color is not None:
        if fill.start_color.type == "rgb" and fill.start_color.rgb not in {None, "00000000"}:
            css["background-color"] = f"#{str(fill.start_color.rgb)[-6:]}"
    border = cell.border
    if border is not None:
        for side in ("top", "left", "bottom", "right"):
            edge = getattr(border, side)
            if edge is not None and edge.style:
                width = {"thick": 3, "medium": 2}.get(edge.style, 1)
                kind = edge.style if edge.style in {"dashed", "dotted", "double"} else "solid"
                hue = "000000"
                if edge.color is not None and edge.color.type == "rgb" and edge.color.rgb:
                    hue = str(edge.color.rgb)[-6:]
                css[f"border-{side}"] = f"{width}px {kind} #{hue}"
    alignment = cell.alignment
    if alignment is not None:
        if alignment.horizontal in {"center", "right", "left", "justify"}:
            css["text-align"] = alignment.horizontal
        if alignment.vertical in {"center", "top", "bottom"}:
            css["vertical-align"] = "middle" if alignment.vertical == "center" else alignment.vertical
    return css


def sheet_table_html(
    sheet, formula_sheet, counts: dict[str, int], bounds: tuple[int, int, int, int] | None = None, header_rows=None
) -> str:
    """One worksheet range as an HTML table: display values, data-formula, merged cells, styles, column widths.

    ``bounds`` is (first row, first column, last row, last column), 1-based; the used range by default.
    """
    from openpyxl.utils import get_column_letter

    from .fidelity import _cell_display

    if bounds is None:
        max_row, max_col = sheet.max_row, sheet.max_column
        while max_row > 0 and all(
            sheet.cell(max_row, c).value is None and formula_sheet.cell(max_row, c).value is None
            for c in range(1, max_col + 1)
        ):
            max_row -= 1
        bounds = (1, 1, max_row, max_col)
    first_row, first_col, last_row, last_col = bounds
    merged = {}
    covered = set()
    for merged_range in sheet.merged_cells.ranges:
        top, left = max(merged_range.min_row, first_row), max(merged_range.min_col, first_col)
        bottom, right = min(merged_range.max_row, last_row), min(merged_range.max_col, last_col)
        if top > bottom or left > right:
            continue
        counts["merged"] += 1
        merged[(top, left)] = (bottom - top + 1, right - left + 1)
        for r in range(top, bottom + 1):
            for c in range(left, right + 1):
                if (r, c) != (top, left):
                    covered.add((r, c))
    if header_rows is None:
        frozen = sheet.freeze_panes
        header_rows = max(0, int(re.sub(r"\D", "", frozen) or 1) - first_row) if frozen else 0
    out = ['<table border="1">', "<colgroup>"]
    for c in range(first_col, last_col + 1):
        width = sheet.column_dimensions[get_column_letter(c)].width or 8.43
        out.append(f'<col style="width:{round(width * 7 + 5)}px">')
    out.append("</colgroup>")
    for r in range(first_row, last_row + 1):
        index = r - first_row + 1
        if index == 1 and header_rows:
            out.append("<thead>")
        out.append("<tr>")
        for c in range(first_col, last_col + 1):
            if (r, c) in covered:
                continue
            cell = sheet.cell(r, c)
            raw = formula_sheet.cell(r, c).value
            attrs = ""
            if isinstance(raw, str) and raw.startswith("=") and formula_sheet.cell(r, c).data_type == "f":
                attrs += f' data-formula="{_attr(raw)}"'
                counts["formulas"] += 1
                if cell.value is None:
                    counts["uncached_formulas"] += 1
            text = _cell_display(cell) if cell.value is not None else ""
            if cell.value is None and isinstance(raw, str) and not raw.startswith("="):
                text = raw
            if (r, c) in merged:
                rows, columns = merged[(r, c)]
                attrs += (f' rowspan="{rows}"' if rows > 1 else "") + (f' colspan="{columns}"' if columns > 1 else "")
            if text:
                counts["cells"] += 1
            tag = "th" if index <= header_rows else "td"
            css = _xlsx_cell_css(cell)
            if tag == "th" and css.get("font-weight") == "bold":
                css.pop("font-weight")
            body = _esc(text).replace("\n", "<br>")
            if cell.hyperlink is not None and cell.hyperlink.target:
                body = f'<a href="{_attr(cell.hyperlink.target)}">{body}</a>'
            out.append(f"<{tag}{attrs}{_style(css)}>{body}</{tag}>")
        out.append("</tr>")
        if index == header_rows:
            out.append("</thead>")
    out.append("</table>")
    return "\n".join(out)


def xlsx_to_html(values, formulas, images: _Images) -> tuple[str, dict[str, int]]:
    out: list[str] = []
    counts = {"sheets": 0, "cells": 0, "formulas": 0, "merged": 0, "uncached_formulas": 0}
    for sheet in values.worksheets:
        counts["sheets"] += 1
        out.append(f"<h2>{_esc(sheet.title)}</h2>")
        out.append(sheet_table_html(sheet, formulas[sheet.title], counts))
        for image in getattr(sheet, "_images", []):
            try:
                blob = image._data()
            except Exception:  # noqa: BLE001 - an image openpyxl cannot read back
                continue
            out.append(f'<p><img src="{_attr(images.source(blob, None, "image." + (image.format or "png")))}"></p>')
    return "\n".join(out) + "\n", counts


# ---------------------------------------------------------------------------------------------------- entry


def office_to_html(
    input_path: str | Path, output_path: str | Path, images: str = "embed"
) -> tuple[bool, dict[str, Any], list[dict[str, str]]]:
    source, destination, errors = check_paths(input_path, output_path, OFFICE_INPUTS, {".html", ".htm"})
    if errors:
        return False, {}, errors
    if images not in {"embed", "files"}:
        return False, {}, [{"code": "INVALID_ARGUMENT", "message": "images must be embed or files."}]
    store = _Images(destination, images)
    suffix = source.suffix.casefold()
    warnings: list[str] = []
    css = ""
    try:
        if suffix == ".docx":
            try:
                from docx import Document
            except ImportError:
                return missing_dependency("html", "python-docx")
            document = Document(source)
            body, counts = docx_to_html(document, store)
            title = document.core_properties.title or source.stem
            section = document.sections[0] if document.sections else None
            if section is not None and section.page_width:
                margin = section.left_margin.mm if section.left_margin else 25
                css = f"@page{{size:{section.page_width.mm:.0f}mm {section.page_height.mm:.0f}mm;margin:{margin:.0f}mm}}\n"
        elif suffix == ".pptx":
            try:
                from pptx import Presentation
            except ImportError:
                return missing_dependency("html", "python-pptx")
            presentation = Presentation(source)
            body, css, counts = pptx_to_html(presentation, store)
            title = presentation.core_properties.title or source.stem
        else:
            try:
                from openpyxl import load_workbook
            except ImportError:
                return missing_dependency("html", "openpyxl")
            values = load_workbook(source, data_only=True)
            formulas = load_workbook(source, data_only=False)
            body, counts = xlsx_to_html(values, formulas, store)
            title = source.stem
            if counts["uncached_formulas"]:
                warnings.append(
                    f"{counts['uncached_formulas']} formula cell(s) have no saved result (the file was not "
                    "recalculated by WPS or Excel); they are empty in the HTML and keep data-formula."
                )
    except Exception as exc:  # noqa: BLE001 - a damaged or unusual Office file
        return False, {}, [{"code": "INPUT_READ_FAILED", "message": f"Could not read {source.name}: {exc}"}]
    destination.write_text(_page(title, body, css), encoding="utf-8")
    return (
        True,
        {
            "input_path": str(source),
            "output_path": str(destination),
            "format": "html",
            "source_format": suffix.lstrip("."),
            "images": store.count,
            "image_mode": images,
            "image_folder": str(store.folder) if images == "files" and store.count else None,
            "mapped_objects": counts,
            "warnings": warnings,
        },
        [],
    )
