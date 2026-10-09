from __future__ import annotations

import weakref
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

MAX_PART_BYTES = 64 * 1024 * 1024
MAX_TOTAL_UNCOMPRESSED_BYTES = 512 * 1024 * 1024
MAX_ENTRIES = 20_000

_checked_archives: "weakref.WeakSet[ZipFile]" = weakref.WeakSet()


class OoxmlTooLargeError(Exception):
    code = "INPUT_TOO_LARGE"

    def __init__(self, message: str, details: dict[str, Any]):
        super().__init__(message)
        self.details = details


def check_archive_limits(archive: ZipFile) -> None:
    """Reject packages whose declared uncompressed size or entry count is abnormal.

    ZipFile enforces the declared per-entry size while reading, so a forged
    header cannot make an entry expand beyond what is checked here.
    """
    if archive in _checked_archives:
        return
    infos = archive.infolist()
    if len(infos) > MAX_ENTRIES:
        raise OoxmlTooLargeError(
            f"Package has {len(infos)} entries; the limit is {MAX_ENTRIES}.",
            {"entries": len(infos), "limit": MAX_ENTRIES},
        )
    total = sum(info.file_size for info in infos)
    if total > MAX_TOTAL_UNCOMPRESSED_BYTES:
        raise OoxmlTooLargeError(
            f"Package expands to {total} bytes; the limit is {MAX_TOTAL_UNCOMPRESSED_BYTES}.",
            {"uncompressed_bytes": total, "limit": MAX_TOTAL_UNCOMPRESSED_BYTES},
        )
    _checked_archives.add(archive)


def read_zip_part(archive: ZipFile, name: str, max_bytes: int = MAX_PART_BYTES) -> bytes:
    check_archive_limits(archive)
    info = archive.getinfo(name)
    if info.file_size > max_bytes:
        raise OoxmlTooLargeError(
            f"Package part {name} expands to {info.file_size} bytes; the limit is {max_bytes}.",
            {"part": name, "uncompressed_bytes": info.file_size, "limit": max_bytes},
        )
    return archive.read(name)


def parse_xml_part(archive: ZipFile, name: str, max_bytes: int = MAX_PART_BYTES) -> ET.Element:
    return ET.fromstring(read_zip_part(archive, name, max_bytes))


def guard_package(path: str | Path) -> None:
    try:
        with ZipFile(path) as archive:
            check_archive_limits(archive)
    except BadZipFile:
        return


def load_workbook_guarded(path: str | Path, *args: Any, **kwargs: Any):
    from openpyxl import load_workbook

    guard_package(path)
    return load_workbook(path, *args, **kwargs)
