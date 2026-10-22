from __future__ import annotations

import base64
import hashlib
import os
import re
import tempfile
from dataclasses import dataclass, field
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit

from .html_text import decode_html_bytes

_VOID = {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "param", "source", "track", "wbr"}
_SKIP = {"script", "style", "noscript", "template", "svg", "iframe", "object", "canvas"}
_BLOCKS = {"p", "div", "section", "article", "header", "footer", "main", "aside", "blockquote", "pre", "address"}
_HEADINGS = {f"h{i}": i for i in range(1, 7)}
_MAX_HTML_BYTES = 10 * 1024 * 1024
_MAX_IMAGE_BYTES = 20 * 1024 * 1024
_MAX_TOTAL_IMAGE_BYTES = 50 * 1024 * 1024
_MAX_IMAGES = 100


def _safe_link_target(value: str) -> bool:
    stripped = value.strip()
    if "\\" in stripped or any(ord(char) < 32 for char in stripped):
        return False
    parsed = urlsplit(stripped)
    return parsed.scheme.casefold() in {"http", "https", "mailto", "tel"} or (
        not parsed.scheme and not value.strip().startswith("//")
    )


def _bookmark_name(object_id: str) -> str:
    return "wps_" + hashlib.sha256(object_id.encode("utf-8")).hexdigest()[:32]


@dataclass
class _Node:
    tag: str
    attrs: dict[str, str] = field(default_factory=dict)
    children: list[_Node | str] = field(default_factory=list)


class _TreeParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.root = _Node("document")
        self.stack = [self.root]
        self.suppressed = 0
        self.unsupported: set[str] = set()
        self.warnings: set[str] = set()
        self.unsupported_elements: set[str] = set()

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag == "style":
            self.unsupported.add("embedded CSS")
        elif tag == "script":
            self.warnings.add("Scripts and executable content were omitted.")
        elif tag in {"noscript", "template", "svg", "iframe", "object", "canvas"}:
            self.unsupported_elements.add(tag)
        if self.suppressed:
            if tag in _SKIP:
                self.suppressed += 1
            return
        if tag in _SKIP:
            self.suppressed = 1
            return
        if tag in {"li", "p", "tr", "td", "th"}:
            while len(self.stack) > 1 and self.stack[-1].tag == tag:
                self.stack.pop()
        node = _Node(tag, {key.lower(): value or "" for key, value in attrs})
        self.stack[-1].children.append(node)
        if tag not in _VOID:
            self.stack.append(node)

    def handle_startendtag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        self.handle_starttag(tag, attrs)
        if tag.lower() not in _VOID:
            self.handle_endtag(tag)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if self.suppressed:
            if tag in _SKIP:
                self.suppressed -= 1
            return
        for index in range(len(self.stack) - 1, 0, -1):
            if self.stack[index].tag == tag:
                del self.stack[index:]
                break

    def handle_data(self, data: str) -> None:
        if self.suppressed and self.stack[-1].tag == "document":
            return
        if not self.suppressed and data:
            self.stack[-1].children.append(data)


def _text(node: _Node) -> str:
    pieces: list[str] = []
    for child in node.children:
        if isinstance(child, str):
            pieces.append(child)
        elif child.tag == "br":
            pieces.append("\n")
        elif child.tag == "img":
            alt = child.attrs.get("alt", "").strip()
            if alt:
                pieces.append(alt)
        elif child.tag == "a":
            label = _text(child)
            href = child.attrs.get("href", "").strip()
            pieces.append(f"{label} ({href})" if href and href not in label else label)
        else:
            pieces.append(_text(child))
    return re.sub(r"[\t\r\f\v ]+", " ", "".join(pieces)).strip()


def _walk(node: _Node):
    yield node
    for child in node.children:
        if isinstance(child, _Node):
            yield from _walk(child)


def table_rows(table: _Node) -> list[_Node]:
    """The table's own <tr> rows, in thead, tbody and tfoot order, without the rows of nested tables."""
    groups: dict[str, list[_Node]] = {"thead": [], "tbody": [], "tfoot": []}
    for child in table.children:
        if not isinstance(child, _Node):
            continue
        if child.tag == "tr":
            groups["tbody"].append(child)
        elif child.tag in groups:
            groups[child.tag] += [row for row in child.children if isinstance(row, _Node) and row.tag == "tr"]
    return groups["thead"] + groups["tbody"] + groups["tfoot"]


def _span(cell: _Node, name: str) -> int:
    try:
        return max(1, min(int(cell.attrs.get(name, "1") or "1"), 1000))
    except ValueError:
        return 1


def table_grid(table: _Node) -> tuple[list[list[_Node | None]], list[tuple[int, int, int, int]]]:
    """Lay the table's cells out on a grid, honouring rowspan and colspan.

    Returns the grid, where a spanned cell's node sits in its top-left slot and the slots it covers hold None,
    and the merged ranges as (first row, first column, last row, last column), zero-based and inclusive.
    """
    grid: list[list[_Node | None]] = []
    covered: set[tuple[int, int]] = set()
    merges: list[tuple[int, int, int, int]] = []
    for row_index, row in enumerate(table_rows(table)):
        while len(grid) <= row_index:
            grid.append([])
        column = 0
        for cell in row.children:
            if not isinstance(cell, _Node) or cell.tag not in {"td", "th"}:
                continue
            while (row_index, column) in covered:
                column += 1
            rows, columns = _span(cell, "rowspan"), _span(cell, "colspan")
            for r in range(row_index, row_index + rows):
                for c in range(column, column + columns):
                    covered.add((r, c))
            while len(grid[row_index]) < column:
                grid[row_index].append(None)
            grid[row_index][column : column + 1] = [cell]
            if rows > 1 or columns > 1:
                merges.append((row_index, column, row_index + rows - 1, column + columns - 1))
            column += columns
    height = max([len(grid)] + [merge[2] + 1 for merge in merges])
    width = max([len(row) for row in grid] + [merge[3] + 1 for merge in merges] + [0])
    grid += [[] for _ in range(height - len(grid))]
    return [row + [None] * (width - len(row)) for row in grid], merges


_DATA_IMAGES: dict[str, Path] = {}


def _data_image(raw: str) -> Path | None:
    """A PNG, JPEG or GIF data: URI written to a temporary file (removed when the process exits)."""
    if raw in _DATA_IMAGES:
        return _DATA_IMAGES[raw]
    import atexit
    import tempfile

    header, _, encoded = raw.partition(",")
    mime = header[5:].split(";", 1)[0].casefold()
    suffix = {"image/png": ".png", "image/jpeg": ".jpg", "image/gif": ".gif"}.get(mime)
    if suffix is None or ";base64" not in header.casefold():
        return None
    try:
        data = base64.b64decode(encoded, validate=True)
    except (ValueError, base64.binascii.Error):
        return None
    if not data or len(data) > _MAX_IMAGE_BYTES:
        return None
    handle = tempfile.NamedTemporaryFile(prefix="wps-data-image-", suffix=suffix, delete=False)
    with handle:
        handle.write(data)
    path = Path(handle.name)
    atexit.register(path.unlink, missing_ok=True)
    _DATA_IMAGES[raw] = path
    return path


def local_image(source: Path, raw: str) -> tuple[Path | None, str | None]:
    """The local image file an <img src> names, inside the HTML file's folder, or the warning why not.

    A PNG, JPEG or GIF data: URI is decoded to a temporary file.
    """
    raw = raw.strip()
    if raw.startswith("data:"):
        path = _data_image(raw)
        return (path, None) if path else (None, "HTML data image is invalid or exceeds its size limit.")
    if not raw or re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", raw):
        return None, f"Image with non-local or empty source omitted: {raw or '(empty)'}"
    try:
        image = (source.parent / raw.split("#", 1)[0].split("?", 1)[0]).resolve(strict=True)
        image.relative_to(source.parent.resolve())
        if not image.is_file() or image.stat().st_size > _MAX_IMAGE_BYTES:
            raise ValueError("image size or type limit exceeded")
    except (OSError, ValueError):
        return None, f"Local image unavailable or outside input directory: {raw}"
    return image, None


def cell_text(cell: _Node) -> str:
    """A table cell's text for a spreadsheet or slide: <br> as a line break, link labels only, no nested tables."""
    pieces: list[str] = []

    def collect(node: _Node) -> None:
        for child in node.children:
            if isinstance(child, str):
                pieces.append(re.sub(r"\s+", " ", child))
            elif child.tag == "br":
                pieces.append("\n")
            elif child.tag == "img":
                pieces.append(child.attrs.get("alt", ""))
            elif child.tag != "table":
                collect(child)

    collect(cell)
    return "\n".join(line.strip() for line in "".join(pieces).split("\n")).strip()


def has_nested_table(cell: _Node) -> bool:
    return any(node.tag == "table" for node in _walk(cell) if node is not cell)


# Inline elements and the run formatting they stand for.
INLINE_FORMATS = {
    "b": "bold",
    "strong": "bold",
    "i": "italic",
    "em": "italic",
    "cite": "italic",
    "dfn": "italic",
    "var": "italic",
    "u": "underline",
    "ins": "underline",
    "s": "strike",
    "del": "strike",
    "strike": "strike",
    "sub": "subscript",
    "sup": "superscript",
    "code": "code",
    "kbd": "code",
    "samp": "code",
    "tt": "code",
    "mark": "highlight",
}
_NAMED_COLORS = {
    "black": "000000",
    "white": "FFFFFF",
    "red": "FF0000",
    "green": "008000",
    "blue": "0000FF",
    "gray": "808080",
    "grey": "808080",
    "orange": "FFA500",
    "purple": "800080",
}


def style_formats(style: str) -> dict[str, Any]:
    """The run formatting an inline style attribute asks for: weight, slant, decoration and colour."""
    formats: dict[str, Any] = {}
    for declaration in style.split(";"):
        name, _, value = declaration.partition(":")
        name, value = name.strip().casefold(), value.strip().casefold()
        if name == "font-weight" and (value in {"bold", "bolder"} or value.isdigit() and int(value) >= 600):
            formats["bold"] = True
        elif name == "font-style" and value in {"italic", "oblique"}:
            formats["italic"] = True
        elif name in {"text-decoration", "text-decoration-line"}:
            if "underline" in value:
                formats["underline"] = True
            if "line-through" in value:
                formats["strike"] = True
        elif name == "font-size":
            size = re.fullmatch(r"(\d+(?:\.\d+)?)(pt|px)", value)
            if size:
                formats["size"] = float(size.group(1)) * (0.75 if size.group(2) == "px" else 1.0)
        elif name == "font-family":
            family = value.split(",")[0].strip().strip("'\"")
            if family and family not in {"sans-serif", "serif", "monospace", "inherit"}:
                formats["font"] = declaration.partition(":")[2].split(",")[0].strip().strip("'\"")
        elif name == "color":
            match = re.fullmatch(r"#([0-9a-f]{6}|[0-9a-f]{3})", value)
            if match:
                digits = match.group(1)
                formats["color"] = (digits if len(digits) == 6 else "".join(ch * 2 for ch in digits)).upper()
            elif value in _NAMED_COLORS:
                formats["color"] = _NAMED_COLORS[value]
    return formats


def validate_office_paths(
    input_path: str | Path, output_path: str | Path, suffix: str = ".docx"
) -> tuple[Path, Path, list[dict[str, str]]]:
    """Check an HTML source and a new output of the given suffix. Shared by the docx, xlsx and pptx converters."""
    source = Path(input_path).expanduser().resolve()
    destination = Path(output_path).expanduser().resolve()
    if source.suffix.casefold() not in {".html", ".htm"}:
        return source, destination, [{"code": "INVALID_INPUT", "message": "Input must be an .html or .htm file."}]
    if not source.is_file():
        return source, destination, [{"code": "INPUT_FILE_NOT_FOUND", "message": f"HTML input not found: {source}"}]
    if source.stat().st_size > _MAX_HTML_BYTES:
        return source, destination, [{"code": "INPUT_TOO_LARGE", "message": "HTML input exceeds the 10 MiB limit."}]
    if destination.suffix.casefold() != suffix or destination == source:
        return source, destination, [{"code": "INVALID_OUTPUT", "message": f"Output must be a separate {suffix} file."}]
    if destination.exists():
        return (
            source,
            destination,
            [{"code": "OUTPUT_ALREADY_EXISTS", "message": f"Refusing to overwrite existing output: {destination}"}],
        )
    if not destination.parent.is_dir():
        return (
            source,
            destination,
            [
                {
                    "code": "OUTPUT_DIRECTORY_NOT_FOUND",
                    "message": f"Output directory does not exist: {destination.parent}",
                }
            ],
        )
    return source, destination, []


def convert_html_editable(
    input_path: str | Path, output_path: str | Path, fidelity: str = "editable", use_browser: bool = True
) -> tuple[bool, dict[str, Any], list[dict[str, str]]]:
    """HTML to an editable docx with the look the browser computes (P13-004).

    fidelity "editable" keeps everything as flowing, editable paragraphs; "appearance" also keeps absolutely
    positioned elements at their place, as framed paragraphs (decision 9 of docs/PHASE13_PLAN.md).
    """
    source, destination, path_errors = validate_office_paths(input_path, output_path, ".docx")
    if path_errors:
        return False, {}, path_errors
    try:
        from docx import Document
        from docx.shared import Inches
    except ImportError:
        return (
            False,
            {},
            [{"code": "CONVERTER_UNAVAILABLE", "message": "Install the html optional dependency: python-docx."}],
        )

    from . import docx_styling
    from .html_styles import load as load_styles

    _html_text, source_encoding, decode_warning = decode_html_bytes(source.read_bytes())
    styled = load_styles(source, use_browser=use_browser)
    parser = styled.parser
    context = docx_styling.Context(styled, appearance=fidelity == "appearance")
    nodes = list(_walk(parser.root))
    parents = {id(child): node for node in nodes for child in node.children if isinstance(child, _Node)}

    def ancestors(node: _Node) -> list[_Node]:
        chain = []
        while id(node) in parents:
            node = parents[id(node)]
            chain.append(node)
        return list(reversed(chain))

    html_root = next((node for node in nodes if node.tag == "html"), None)
    meta = next(
        (node for node in nodes if node.tag == "meta" and node.attrs.get("name", "").casefold() == "wps-agent-schema"),
        None,
    )
    schema_version = (html_root.attrs.get("data-wps-schema") if html_root else None) or (
        meta.attrs.get("content") if meta else None
    )
    owned_schema = schema_version == "wps-agent-html/v1"
    if owned_schema:
        from .html_roundtrip import build_html_roundtrip_mapping

        valid, _, validation_errors = build_html_roundtrip_mapping(source)
        if not valid:
            return False, {}, validation_errors
    unsupported_css = docx_styling.css_coverage(styled)
    supported_tags = (
        _VOID
        | _SKIP
        | _BLOCKS
        | set(_HEADINGS)
        | {
            "html",
            "head",
            "body",
            "title",
            "span",
            "a",
            "b",
            "strong",
            "i",
            "em",
            "u",
            "s",
            "small",
            "sub",
            "sup",
            "code",
            "abbr",
            "cite",
            "del",
            "ins",
            "q",
            "ul",
            "ol",
            "li",
            "table",
            "thead",
            "tbody",
            "tfoot",
            "tr",
            "td",
            "th",
            "figure",
            "figcaption",
            "caption",
            "details",
            "summary",
            "nav",
            "dl",
            "dt",
            "dd",
            "mark",
            "kbd",
            "samp",
            "var",
            "dfn",
            "strike",
            "tt",
            "colgroup",
            "col",
        }
    )
    warnings = sorted(
        parser.warnings
        | {
            f"Unsupported element <{node.tag}> omitted or flattened to text."
            for node in nodes
            if node is not parser.root and node.tag not in supported_tags
        }
        | {f"Unsupported element <{tag}> omitted." for tag in parser.unsupported_elements}
        | {"Scripts and executable content were omitted." for node in nodes if node.tag == "script"}
    )
    if decode_warning:
        warnings.append(decode_warning)
    document = Document()
    page_width_px = docx_styling.apply_page(document, styled)
    bookmark_counter = 0
    bookmark_count = 0

    def add_bookmark(paragraph_element, object_id: str) -> None:
        nonlocal bookmark_counter, bookmark_count
        if not owned_schema or not object_id:
            return
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        bookmark_counter += 1
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(bookmark_counter))
        start.set(qn("w:name"), _bookmark_name(object_id))
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(bookmark_counter))
        paragraph_element.append(start)
        paragraph_element.append(end)
        bookmark_count += 1

    image_nodes: list[tuple[_Node, Path | bytes]] = []
    image_bytes = 0
    root = source.parent.resolve()
    for node in nodes:
        if node.tag != "img":
            continue
        if len(image_nodes) >= _MAX_IMAGES:
            warnings.append(f"Image count truncated to {_MAX_IMAGES}.")
            continue
        raw = node.attrs.get("src", "").strip()
        if raw.startswith("data:"):
            try:
                header, encoded = raw.split(",", 1)
                mime = header[5:].split(";", 1)[0].casefold()
                if mime not in {"image/png", "image/jpeg", "image/gif"} or ";base64" not in header.casefold():
                    raise ValueError("unsupported data image format")
                image_data = base64.b64decode(encoded, validate=True)
                if (
                    not image_data
                    or len(image_data) > _MAX_IMAGE_BYTES
                    or image_bytes + len(image_data) > _MAX_TOTAL_IMAGE_BYTES
                ):
                    raise ValueError("image size limit exceeded")
                image_nodes.append((node, image_data))
                image_bytes += len(image_data)
            except (ValueError, base64.binascii.Error):
                warnings.append("HTML data image is invalid or exceeds its size limit.")
            continue
        if not raw or re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", raw):
            warnings.append(f"Image with non-local or empty source omitted: {raw or '(empty)'}")
            continue
        try:
            image = (source.parent / raw.split("#", 1)[0].split("?", 1)[0]).resolve(strict=True)
            image.relative_to(root)
            size = image.stat().st_size
            if not image.is_file() or size > _MAX_IMAGE_BYTES or image_bytes + size > _MAX_TOTAL_IMAGE_BYTES:
                raise ValueError("image size or type limit exceeded")
            image_nodes.append((node, image))
            image_bytes += size
        except (OSError, ValueError):
            warnings.append(f"Local image unavailable or outside input directory: {raw}")
    image_map = {id(node): image for node, image in image_nodes}
    counts = {"paragraphs": 0, "headings": 0, "lists": 0, "tables": 0, "images": 0}

    def add_image(node: _Node, paragraph=None) -> None:
        from docx.shared import Emu

        image = image_map.get(id(node))
        if image is not None:
            picture_source = BytesIO(image) if isinstance(image, bytes) else str(image)
            width, height = docx_styling.image_size(context, node)
            limit = int(page_width_px * docx_styling.EMU_PER_PX)
            if width is None:
                width, height = min(Inches(5.5), limit), None
            elif width > limit:
                width, height = limit, (int(height * limit / width) if height else None)
            size = {"width": Emu(width)} | ({"height": Emu(height)} if height else {})
            if paragraph is None:
                paragraph = document.add_paragraph()
                holder = parents.get(id(node))
                if holder is not None and holder.tag in {"figure", "p", "div", "section", "body"}:
                    docx_styling.apply_paragraph(paragraph, holder, ancestors(holder), context, list_item=True)
            picture_run = paragraph.add_run()
            picture_run.add_picture(picture_source, **size)
            described = picture_run._r.find(
                ".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"
            )
            if described is not None and node.attrs.get("alt"):
                described.set("descr", node.attrs["alt"])
            add_bookmark(paragraph._p, node.attrs.get("data-wps-object-id", ""))
            counts["images"] += 1

    def apply_format(run, formats: dict[str, Any]) -> None:
        from docx.enum.text import WD_COLOR_INDEX
        from docx.shared import RGBColor

        font = run.font
        if formats.get("bold"):
            font.bold = True
        if formats.get("italic"):
            font.italic = True
        if formats.get("underline"):
            font.underline = True
        if formats.get("strike"):
            font.strike = True
        if formats.get("subscript"):
            font.subscript = True
        if formats.get("superscript"):
            font.superscript = True
        if formats.get("code"):
            font.name = "Consolas"
        if formats.get("highlight"):
            font.highlight_color = WD_COLOR_INDEX.YELLOW
        if formats.get("color"):
            font.color.rgb = RGBColor.from_string(formats["color"])

    def add_text_with_links(
        paragraph, node: _Node, formats: dict[str, Any] | None = None, preserve: bool = False, skip=()
    ) -> None:
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        formats = formats or {}
        node_props = docx_styling.run_props(context, node)
        for child in node.children:
            if isinstance(child, str):
                # HTML reads CRLF and CR as one line break; python-docx would make each of \r and \n a break (a
                # blank line in every code block of a file checked out with CRLF endings, Windows check on f8ed7bd).
                text = re.sub(r"\r\n?", "\n", child) if preserve else re.sub(r"\s+", " ", child)
                if text:
                    run = paragraph.add_run(text)
                    docx_styling.apply_run(run, node_props)
                    apply_format(run, formats)
            elif child.tag in skip:
                continue
            elif child.tag == "br":
                paragraph.add_run().add_break()
            elif child.tag == "img":
                add_image(child, paragraph)
            elif child.tag == "a":
                label = re.sub(r"\s+", " ", _text(child))
                href = child.attrs.get("href", "").strip()
                if label and href and _safe_link_target(href):
                    relation = paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True)
                    hyperlink = OxmlElement("w:hyperlink")
                    hyperlink.set(qn("r:id"), relation)
                    run = OxmlElement("w:r")
                    props = OxmlElement("w:rPr")
                    for name in ("bold", "italic"):
                        if formats.get(name):
                            props.append(OxmlElement("w:b" if name == "bold" else "w:i"))
                    color = OxmlElement("w:color")
                    color.set(qn("w:val"), "0563C1")
                    underline = OxmlElement("w:u")
                    underline.set(qn("w:val"), "single")
                    props.append(color)
                    props.append(underline)
                    run.append(props)
                    text = OxmlElement("w:t")
                    text.text = label
                    text.set(qn("xml:space"), "preserve")
                    run.append(text)
                    hyperlink.append(run)
                    paragraph._p.append(hyperlink)
                    if owned_schema and child.attrs.get("data-wps-object-id"):
                        add_bookmark(paragraph._p, child.attrs["data-wps-object-id"])
                else:
                    run = paragraph.add_run(label)
                    docx_styling.apply_run(run, docx_styling.run_props(context, child))
                    apply_format(run, formats)
                    if href and not _safe_link_target(href):
                        warnings.append("Unsafe hyperlink target omitted.")
            else:
                inner = dict(formats)
                if child.tag in INLINE_FORMATS:
                    inner[INLINE_FORMATS[child.tag]] = True
                if child.attrs.get("style"):
                    inner.update(style_formats(child.attrs["style"]))
                add_text_with_links(paragraph, child, inner, preserve or child.tag == "pre", skip)

    def trim(paragraph) -> None:
        """Drop the whitespace the HTML source left at the paragraph's ends (runs and links, in order)."""
        from docx.oxml.ns import qn
        from docx.text.run import Run

        def ends(reverse: bool):
            items = [child for child in paragraph._p if child.tag in {qn("w:r"), qn("w:hyperlink")}]
            return reversed(items) if reverse else iter(items)

        for reverse in (False, True):
            for element in ends(reverse):
                if element.tag != qn("w:r") or element.find(qn("w:t")) is None:
                    break
                run = Run(element, paragraph)
                text = run.text.rstrip() if reverse else run.text.lstrip()
                if text:
                    run.text = text
                    break
                if element.find(qn("w:br")) is not None or element.find(qn("w:drawing")) is not None:
                    break
                element.getparent().remove(element)

    def add_paragraph(
        node: _Node, style: str | None = None, skip=(), formats: dict[str, Any] | None = None, container=None
    ):
        preserve = node.tag == "pre"
        paragraph = (container or document).add_paragraph(style=style)
        docx_styling.apply_paragraph(
            paragraph, node, ancestors(node), context, list_item=bool(style and style.startswith("List"))
        )
        add_bookmark(paragraph._p, node.attrs.get("data-wps-object-id", ""))
        add_text_with_links(paragraph, node, formats, preserve, skip)
        if preserve:
            runs = [run for run in paragraph.runs if run.text]
            if runs:
                runs[0].text = runs[0].text.lstrip("\r\n")
                runs[-1].text = runs[-1].text.rstrip()
        else:
            trim(paragraph)
        return paragraph

    def list_style(tag: str, level: int) -> str:
        base = "List Number" if tag == "ol" else "List Bullet"
        return base if level == 1 else f"{base} {min(level, 3)}"

    def add_list(node: _Node, level: int = 1) -> None:
        for item in node.children:
            if not (isinstance(item, _Node) and item.tag == "li"):
                continue
            nested = [child for child in item.children if isinstance(child, _Node) and child.tag in {"ul", "ol"}]
            if _text(_Node("li", children=[c for c in item.children if c not in nested])):
                add_paragraph(item, list_style(node.tag, level), skip=("ul", "ol"))
                counts["lists"] += 1
            for child in nested:
                add_list(child, level + 1)

    def add_table(node: _Node) -> None:
        grid, merges = table_grid(node)
        columns = len(grid[0]) if grid else 0
        if not grid or not columns:
            return
        for caption in (c for c in node.children if isinstance(c, _Node) and c.tag == "caption"):
            if _text(caption):
                add_paragraph(caption, "Caption")
        table = document.add_table(rows=len(grid), cols=columns)
        table.style = "Table Grid"
        rows = table_rows(node)
        context.last_space_after = 0.0
        for row_index, row in enumerate(grid):
            first_cell_paragraph = table.cell(row_index, 0).paragraphs[0]._p
            if row_index == 0:
                add_bookmark(first_cell_paragraph, node.attrs.get("data-wps-object-id", ""))
            if row_index < len(rows):
                add_bookmark(first_cell_paragraph, rows[row_index].attrs.get("data-wps-object-id", ""))
            for column_index, cell_node in enumerate(row):
                if cell_node is None:
                    continue
                paragraph = table.cell(row_index, column_index).paragraphs[0]
                add_bookmark(paragraph._p, cell_node.attrs.get("data-wps-object-id", ""))
                # A header cell is bold by default; with computed styles the cascade decides (th{font-weight:400}).
                bold = cell_node.tag == "th" and context.styled is None
                docx_styling.apply_paragraph(paragraph, cell_node, ancestors(cell_node), context, list_item=True)
                paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = 0
                add_text_with_links(paragraph, cell_node, {"bold": True} if bold else None, skip=("table",))
                trim(paragraph)
                if has_nested_table(cell_node):
                    warnings.append("A table nested in a table cell was flattened to the cell's text.")
                    paragraph.add_run(
                        " "
                        + " ".join(_text(t) for t in cell_node.children if isinstance(t, _Node) and t.tag == "table")
                    )
        for first_row, first_column, last_row, last_column in merges:
            keep = table.cell(first_row, first_column)
            text_paragraphs = list(keep.paragraphs)
            merged = keep.merge(table.cell(last_row, last_column))
            # merge() appends the covered cells' empty paragraphs; keep only the original one.
            for extra in merged.paragraphs[len(text_paragraphs) :]:
                if not extra.text:
                    extra._p.getparent().remove(extra._p)
        docx_styling.apply_table(table, node, grid, rows, context, page_width_px)
        docx_styling.apply_table_spacing(table, node, context)
        counts["tables"] += 1

    def add_blocks(container: _Node, list_style: str | None = None) -> None:
        for node in container.children:
            if isinstance(node, str):
                text = re.sub(r"\s+", " ", node).strip()
                if text:
                    document.add_paragraph(text)
                    counts["paragraphs"] += 1
                continue
            tag = node.tag
            role = docx_styling.page_role(node)
            if role:
                part = document.sections[0].header if role == "header" else document.sections[0].footer
                blocks = [c for c in node.children if isinstance(c, _Node) and _text(c)] or [node]
                for index, block in enumerate(blocks):
                    paragraph = part.paragraphs[0] if index == 0 else part.add_paragraph()
                    docx_styling.apply_paragraph(paragraph, block, ancestors(block), context, list_item=True)
                    add_text_with_links(paragraph, block)
                    trim(paragraph)
                context.applied.add(f"page {role}")
                continue
            if tag in _HEADINGS:
                if _text(node):
                    paragraph = document.add_heading(level=_HEADINGS[tag])
                    docx_styling.apply_paragraph(paragraph, node, ancestors(node), context)
                    add_bookmark(paragraph._p, node.attrs.get("data-wps-object-id", ""))
                    add_text_with_links(paragraph, node)
                    trim(paragraph)
                    counts["headings"] += 1
            elif tag == "p" or tag in {"blockquote", "pre", "address", "dt", "dd", "summary"}:
                if _text(node) or any(child.tag == "img" for child in _walk(node) if child is not node):
                    if tag == "blockquote" and any(
                        isinstance(child, _Node) and child.tag in _BLOCKS | {"ul", "ol", "table"}
                        for child in node.children
                    ):
                        add_blocks(node)
                        continue
                    style = {"blockquote": "Quote", "pre": "No Spacing"}.get(tag)
                    holder = parents.get(id(node))
                    if style is None and holder is not None and holder.tag == "blockquote":
                        style = "Quote"
                    add_paragraph(node, style)
                    counts["paragraphs"] += 1
            elif tag in {"ul", "ol"}:
                add_list(node)
            elif tag == "table":
                add_table(node)
            elif tag == "img":
                add_image(node)
            elif tag == "hr":
                # A rule: an empty paragraph with a bottom border.
                from docx.oxml import OxmlElement
                from docx.oxml.ns import qn

                rule = document.add_paragraph()
                borders = OxmlElement("w:pBdr")
                bottom = OxmlElement("w:bottom")
                for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "0"), ("w:color", "999999")):
                    bottom.set(qn(key), value)
                borders.append(bottom)
                rule._p.get_or_add_pPr().append(borders)
                docx_styling.apply_rule(rule, node, context)
                counts["rules"] = counts.get("rules", 0) + 1
            elif tag == "figcaption":
                if _text(node):
                    add_paragraph(node, "Caption")
                    counts["paragraphs"] += 1
            elif tag in _BLOCKS or tag in {
                "body",
                "html",
                "main",
                "section",
                "article",
                "figure",
                "details",
                "nav",
                "dl",
            }:
                add_blocks(node)

    try:
        add_blocks(parser.root)
        if not any(counts.values()):
            return (
                False,
                {},
                [{"code": "NO_SUPPORTED_CONTENT", "message": "No supported semantic HTML content was found."}],
            )
        document.core_properties.title = next(
            (_text(node) for node in nodes if node.tag == "title" and _text(node)), "HTML conversion"
        )
        handle = tempfile.NamedTemporaryFile(
            prefix=f".{destination.stem}.", suffix=".docx", dir=destination.parent, delete=False
        )
        temporary = Path(handle.name)
        handle.close()
        document.save(temporary)
        payload = temporary.read_bytes()
        os.link(temporary, destination)
        temporary.unlink()
    except Exception as exc:  # noqa: BLE001
        if "temporary" in locals():
            temporary.unlink(missing_ok=True)
        return False, {}, [{"code": "CONVERSION_FAILED", "message": str(exc)[:500]}]
    warnings = sorted(set(warnings))
    return (
        True,
        {
            "input_path": str(source),
            "output_path": str(destination),
            "bytes": len(payload),
            "sha256": hashlib.sha256(payload).hexdigest(),
            "backend": "python-docx-semantic-mapping",
            "editable": True,
            "mapped_objects": counts,
            "unsupported_css": unsupported_css,
            "warnings": warnings,
            "fidelity": fidelity,
            "style_source": styled.mode,
            "applied_styles": sorted(context.applied),
            "source_encoding": source_encoding,
            "javascript_executed": False,
            "remote_resources_fetched": False,
            "roundtrip_bookmarks": bookmark_count,
        },
        [],
    )
