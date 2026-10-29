"""Convert local semantic HTML into an editable spreadsheet (xlsx) or presentation (pptx).

These converters run offline. They do not use WPS. Each outermost <table> becomes a worksheet, and the
nearest preceding heading names it; rowspan and colspan become merged cells, numbers written with
thousands separators, percent signs, currency symbols or exponents become numbers with a matching number
format, ISO-like dates become dates, and <br> becomes a line break in the cell. Local images are placed
below the table before them. The presentation converter is in html_slides. Anything that cannot be mapped
(cell formulas, nested tables) is reported as a warning rather than dropped silently.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from .html_editable import HTML_SPACE, collapse_space, validate_office_paths
from .html_text import decode_html_bytes

_HEADINGS = {"h1", "h2", "h3", "h4", "h5", "h6"}
_NUMBER = re.compile(r"-?(0|[1-9]\d*)(\.\d+)?")
MAX_SHEETS = 50
MAX_ROWS_PER_SHEET = 10_000
MAX_SLIDES = 200
_SHEET_TITLE_LIMIT = 31  # Excel's limit


def _normalise(text: str) -> str:
    return collapse_space(text).strip(HTML_SPACE)


def _read_html(source: Path) -> tuple[str, str | None]:
    text, _encoding, decode_warning = decode_html_bytes(source.read_bytes())
    return text, decode_warning


_THOUSANDS = re.compile(r"[-+]?\d{1,3}(,\d{3})+(\.\d+)?")
_PERCENT = re.compile(r"[-+]?\d+(\.\d+)?%")
_SCIENTIFIC = re.compile(r"[-+]?\d+(\.\d+)?[eE][-+]?\d+")
_DATE = re.compile(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})")
_CURRENCIES = {"¥": "¥", "￥": "¥", "$": "$", "€": "€", "£": "£"}


def _decimals(text: str) -> int:
    return len(text.split(".", 1)[1].rstrip("%")) if "." in text else 0


def _cell_value(text: str) -> tuple[Any, bool, str | None]:
    """Return (value, is_formula_like, number_format) for a cell's text.

    Plain, thousands-separated, signed, percent, currency and scientific numbers become numbers, and ISO-like
    dates become dates, each with a number format that shows them as written. Leading zeros stay text (an id,
    not a number), and text starting with '=' stays text (a formula is never evaluated from HTML).
    """
    import datetime as dt

    if text.startswith("="):
        return text, True, None
    stripped = text.strip()
    if sum(char.isdigit() for char in stripped) > 15:
        # More digits than a spreadsheet number keeps (15 significant): an ID card or bank card number would lose its
        # last digits, so it stays text.
        return text, False, None
    plain = stripped.lstrip("+")
    if _NUMBER.fullmatch(plain):
        decimals = _decimals(plain)
        return (float(plain) if "." in plain else int(plain)), False, ("0." + "0" * decimals if decimals else None)
    if _THOUSANDS.fullmatch(stripped):
        number = stripped.lstrip("+").replace(",", "")
        decimals = _decimals(number)
        value = float(number) if decimals else int(number)
        return value, False, "#,##0" + ("." + "0" * decimals if decimals else "")
    if _PERCENT.fullmatch(stripped):
        number = stripped.lstrip("+")[:-1]
        decimals = _decimals(number)
        return round(float(number) / 100, decimals + 2), False, "0" + ("." + "0" * decimals if decimals else "") + "%"
    if stripped[:1] in _CURRENCIES:
        value, _formula, number_format = _cell_value(stripped[1:].strip())
        if isinstance(value, (int, float)) and not isinstance(value, bool) and "%" not in (number_format or ""):
            symbol = _CURRENCIES[stripped[0]]
            decimals = number_format.split(".", 1)[1] if number_format and "." in number_format else ""
            return value, False, f'"{symbol}"#,##0' + (f".{decimals}" if decimals else "")
    if _SCIENTIFIC.fullmatch(stripped):
        return float(stripped), False, "0.00E+00"
    match = _DATE.fullmatch(stripped)
    if match:
        separator = re.search(r"[-/.]", stripped).group(0)
        try:
            return dt.date(*(int(part) for part in match.groups())), False, f"yyyy{separator}mm{separator}dd"
        except ValueError:
            pass
    return text, False, None


def _unique_sheet_name(base: str, used: set[str]) -> str:
    # Excel forbids [ ] : * ? / \ in a sheet name; WPS also replaces their full-width forms (：？／＊［］＼) when it
    # saves the workbook, so they are replaced here too and the name stays the same after a save in WPS.
    cleaned = re.sub(r"[\[\]:*?/\\：？／＊［］＼]", " ", base).strip() or "Table"
    cleaned = cleaned[:_SHEET_TITLE_LIMIT]
    name, counter = cleaned, 2
    while name.casefold() in used:
        suffix = f" {counter}"
        name = cleaned[: _SHEET_TITLE_LIMIT - len(suffix)] + suffix
        counter += 1
    used.add(name.casefold())
    return name


def _parse_tree(source: Path):
    from .html_editable import _TreeParser

    text, decode_warning = _read_html(source)
    parser = _TreeParser()
    parser.feed(text)
    parser.close()
    return parser, decode_warning


def _outer_tables(root) -> list[tuple[str | None, Any]]:
    """The outermost tables in document order, each with the text of the nearest heading before it."""
    from .html_editable import _Node, _text

    found: list[tuple[str | None, Any]] = []
    last_heading: list[str | None] = [None]

    def visit(node) -> None:
        for child in node.children:
            if not isinstance(child, _Node):
                continue
            if child.tag in _HEADINGS:
                last_heading[0] = _normalise(_text(child)) or last_heading[0]
            elif child.tag == "table":
                found.append((last_heading[0], child))
            else:
                visit(child)

    visit(root)
    return found


def _style_cell(cell, style: dict[str, str], header: bool) -> set[str]:
    """Font, fill, borders and alignment from the cell's computed style (P13-010); returns what was applied."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    from .html_styles import color, first_font, is_bold, points

    applied: set[str] = set()
    hue = color(style.get("color"))
    size = points(style.get("font-size"))
    family = first_font(style.get("font-family"))
    bold = header or is_bold(style)
    italic = style.get("font-style") == "italic"
    decoration = style.get("text-decoration-line", "") or ""
    font = Font(
        name=family,
        size=round(size, 1) if size and abs(size - 12) > 0.1 else None,
        bold=bold or None,
        italic=italic or None,
        underline="single" if "underline" in decoration else None,
        strike=("line-through" in decoration) or None,
        color=hue if hue and hue != "000000" else None,
    )
    if font != Font():
        cell.font = font
        applied.add("font")
    fill = color(style.get("background-color"))
    if fill:
        cell.fill = PatternFill("solid", start_color=fill, end_color=fill)
        applied.add("fill")
    sides = {}
    for side in ("left", "right", "top", "bottom"):
        width = style.get(f"border-{side}-width", "0px")
        kind = style.get(f"border-{side}-style", "none")
        if kind in {"none", "hidden", ""} or width in {"0px", "0", ""}:
            continue
        try:
            thickness = float(width.rstrip("px"))
        except ValueError:
            thickness = 1.0
        line = "thick" if thickness >= 3 else "medium" if thickness >= 2 else "thin"
        line = {"dashed": "dashed", "dotted": "dotted", "double": "double"}.get(kind, line)
        sides[side] = Side(style=line, color=color(style.get(f"border-{side}-color")) or "000000")
    if sides:
        cell.border = Border(**sides)
        applied.add("border")
    horizontal = {"center": "center", "right": "right", "end": "right", "justify": "justify", "left": "left"}.get(
        style.get("text-align", "")
    )
    vertical = {"middle": "center", "top": "top", "bottom": "bottom"}.get(style.get("vertical-align", ""))
    if horizontal in {"left"} and not header:
        horizontal = None
    if horizontal or vertical:
        old = cell.alignment
        cell.alignment = Alignment(
            horizontal=horizontal or old.horizontal, vertical=vertical or old.vertical, wrap_text=old.wrap_text
        )
        applied.add("alignment")
    return applied


def _text_width(text: str) -> float:
    return max((sum(2 if ord(ch) > 0x2E7F else 1 for ch in line) for line in text.split("\n")), default=0)


def convert_html_to_xlsx(
    input_path: str | Path, output_path: str | Path, use_browser: bool = True
) -> tuple[bool, dict[str, Any], list[dict[str, str]]]:
    source, destination, path_errors = validate_office_paths(input_path, output_path, ".xlsx")
    if path_errors:
        return False, {}, path_errors
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError:
        return False, {}, [{"code": "CONVERTER_UNAVAILABLE", "message": "openpyxl is required for xlsx output."}]
    from .html_editable import _safe_link_target, _walk, cell_text, has_nested_table, local_image, table_grid
    from .html_styles import load as load_styles
    from .html_styles import px

    _text_unused, decode_warning = _read_html(source)
    styled = load_styles(source, use_browser=use_browser)
    parser = styled.parser
    tables = _outer_tables(parser.root)
    if not tables:
        return (
            False,
            {},
            [{"code": "NO_CONTENT", "message": "The HTML contains no <table> elements to convert to a worksheet."}],
        )
    warnings: list[str] = [decode_warning] if decode_warning else []
    if len(tables) > MAX_SHEETS:
        warnings.append(f"Only the first {MAX_SHEETS} tables were converted.")
        tables = tables[:MAX_SHEETS]
    workbook = Workbook()
    workbook.remove(workbook.active)
    used: set[str] = set()
    sheets = []
    formula_like = 0
    restored = 0
    applied: set[str] = set()
    total_rows = 0
    nested = 0
    table_sheets: dict[int, Any] = {}
    for index, (title, table) in enumerate(tables, start=1):
        grid, merges = table_grid(table)
        if len(grid) > MAX_ROWS_PER_SHEET:
            warnings.append(f"Table {index} was cut to {MAX_ROWS_PER_SHEET} rows.")
            grid = grid[:MAX_ROWS_PER_SHEET]
            merges = [merge for merge in merges if merge[2] < MAX_ROWS_PER_SHEET]
        sheet = workbook.create_sheet(_unique_sheet_name(title or f"Table {index}", used))
        table_sheets[id(table)] = sheet
        widths: dict[int, float] = {}
        header_rows = sum(
            1
            for part in table.children
            if getattr(part, "tag", "") == "thead"
            for row_node in part.children
            if getattr(row_node, "tag", "") == "tr"
        )
        for row_index, row in enumerate(grid, start=1):
            for column_index, node in enumerate(row, start=1):
                if node is None:
                    continue
                text = cell_text(node)
                if has_nested_table(node):
                    nested += 1
                    inner = [
                        " | ".join(
                            cell_text(item) for item in row_node.children if getattr(item, "tag", "") in {"td", "th"}
                        )
                        for item_table in _walk(node)
                        if item_table is not node and item_table.tag == "table"
                        for row_node in item_table.children
                        if getattr(row_node, "tag", "") == "tr"
                    ]
                    text = "\n".join(part for part in [text, *inner] if part)
                if not text and not node.attrs.get("data-formula"):
                    continue
                value, is_formula, number_format = _cell_value(text)
                cell = sheet.cell(row=row_index, column=column_index)
                formula = node.attrs.get("data-formula", "").strip()
                if formula.startswith("=") and len(formula) > 1:
                    # Written by office-to-html from a workbook formula: restore it (decision 7).
                    cell.value = formula
                    restored += 1
                else:
                    cell.value = value
                if is_formula and not formula:
                    cell.data_type = "s"
                    formula_like += 1
                if number_format:
                    cell.number_format = number_format
                if node.tag == "th":
                    cell.font = Font(bold=True)
                if "\n" in text:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                applied |= _style_cell(cell, styled.style(node), node.tag == "th")
                span = int(node.attrs.get("colspan", "1") or 1) if node.attrs.get("colspan", "1").isdigit() else 1
                if span == 1:
                    box = styled.box(node)
                    css_width = px(styled.style(node).get("width")) if styled.mode == "cascade" else None
                    width = (box["width"] if box else css_width) or 0
                    estimate = width / 7.0 if width else _text_width(text) + 2
                    widths[column_index] = max(widths.get(column_index, 0), estimate)
                links = [link for link in _walk(node) if link.tag == "a" and link.attrs.get("href")]
                if len(links) == 1 and _safe_link_target(links[0].attrs["href"]):
                    cell.hyperlink = links[0].attrs["href"].strip()
                    cell.style = "Hyperlink"
        for first_row, first_column, last_row, last_column in merges:
            sheet.merge_cells(
                start_row=first_row + 1, start_column=first_column + 1, end_row=last_row + 1, end_column=last_column + 1
            )
            if grid[first_row][first_column].tag == "th":
                sheet.cell(row=first_row + 1, column=first_column + 1).alignment = Alignment(
                    horizontal="center", vertical="center"
                )
        from openpyxl.utils import get_column_letter

        for column_index, width in widths.items():
            sheet.column_dimensions[get_column_letter(column_index)].width = round(min(80, max(6, width)), 1)
        if header_rows and header_rows < len(grid):
            sheet.freeze_panes = f"A{header_rows + 1}"
        total_rows += len(grid)
        sheets.append({"name": sheet.title, "rows": len(grid), "merged_ranges": len(merges)})
    # Each image goes on the sheet of the nearest table before it, below that table.
    images = 0
    current_sheet = None
    for node in _walk(parser.root):
        if id(node) in table_sheets:
            current_sheet = table_sheets[id(node)]
        if node.tag != "img":
            continue
        path, warning = local_image(source, node.attrs.get("src", ""))
        if warning:
            warnings.append(warning)
            continue
        try:
            from openpyxl.drawing.image import Image as SheetImage

            picture = SheetImage(str(path))
        except Exception:  # noqa: BLE001 - Pillow missing or not an image
            warnings.append(f"Image could not be read and was omitted: {node.attrs.get('src', '')}")
            continue
        sheet = current_sheet or workbook.worksheets[0]
        picture.anchor = f"A{sheet.max_row + 2}"
        sheet.add_image(picture)
        images += 1
    if formula_like:
        warnings.append(
            f"{formula_like} cell(s) starting with '=' were written as text, not formulas, as in the HTML. "
            "Leave them as text unless the user asks for working formulas."
        )
    if nested:
        warnings.append(f"{nested} table(s) nested in a cell were written into that cell as text, one row per line.")
    workbook.save(destination)
    return (
        True,
        {
            "input_path": str(source),
            "output_path": str(destination),
            "format": "xlsx",
            "sheets": sheets,
            "rows": total_rows,
            "images": images,
            "formulas_restored": restored,
            "applied_styles": sorted(applied),
            "style_source": styled.mode,
            "warnings": warnings + styled.warnings,
        },
        [],
    )


def convert_html_to_pptx(
    input_path: str | Path, output_path: str | Path
) -> tuple[bool, dict[str, Any], list[dict[str, str]]]:
    from .html_slides import convert_html_to_pptx as convert

    return convert(input_path, output_path)


OUTPUT_FORMATS = ("docx", "xlsx", "pptx")
BACKENDS = {"docx": "python-docx", "xlsx": "openpyxl", "pptx": "python-pptx"}


def convert_html_to_office(
    input_path: str | Path,
    output_path: str | Path,
    output_format: str = "docx",
    fidelity: str = "editable",
    use_browser: bool = True,
) -> tuple[bool, dict[str, Any], list[dict[str, str]]]:
    """Dispatch to the converter for one output format."""
    if output_format == "docx":
        from .html_editable import convert_html_editable

        return convert_html_editable(input_path, output_path, fidelity=fidelity, use_browser=use_browser)
    if output_format == "xlsx":
        return convert_html_to_xlsx(input_path, output_path, use_browser=use_browser)
    if output_format == "pptx":
        from .html_slides import convert_html_to_pptx as convert

        return convert(input_path, output_path, fidelity=fidelity, use_browser=use_browser)
    return (
        False,
        {},
        [
            {
                "code": "INVALID_ARGUMENT",
                "message": f"Unsupported output format: {output_format}. Use docx, xlsx or pptx.",
            }
        ],
    )
