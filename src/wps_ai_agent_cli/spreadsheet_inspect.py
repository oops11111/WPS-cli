from __future__ import annotations

from datetime import date, datetime, time
import json
from pathlib import Path
import subprocess
import tempfile
from typing import Any

from .ooxml import load_workbook_guarded as load_workbook
from openpyxl.utils.cell import get_column_letter
from openpyxl.utils.datetime import to_excel

from .capabilities import powershell_executable, probe_wps_capabilities
from .errors import COM_BACKEND_UNAVAILABLE, COM_OPERATION_FAILED, INPUT_FILE_NOT_FOUND
from .sessions import get_document
from .spreadsheet_ranges import validate_spreadsheet_read_range
from .wps_script_snippets import QUIT_IF_IDLE


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _cache_matches(cached: Any, recalculated: Any, epoch: datetime | None = None) -> bool | None:
    if cached is None or recalculated is None:
        return cached is recalculated
    if isinstance(cached, (int, float)) and isinstance(recalculated, (int, float)):
        return abs(float(cached) - float(recalculated)) <= max(1e-12, abs(float(recalculated)) * 1e-12)
    if isinstance(cached, (datetime, date, time)) and isinstance(recalculated, (int, float)) and epoch is not None:
        return abs(float(to_excel(cached, epoch)) - float(recalculated)) <= max(1e-12, abs(float(recalculated)) * 1e-12)
    return _json_value(cached) == _json_value(recalculated)


def inspect_spreadsheet_range(
    document_id: str,
    range_address: str,
    sheet_name: str | None = None,
    workspace: str | Path = ".",
) -> tuple[bool, dict[str, Any], list[dict[str, str]]]:
    document = get_document(document_id, workspace)
    if not document:
        return False, {}, [{"code": "DOCUMENT_NOT_FOUND", "message": f"Document not registered: {document_id}"}]
    if document["component"] != "spreadsheets":
        return False, {"document": document}, [{"code": "UNSUPPORTED_COMPONENT", "message": "spreadsheet-inspect requires a spreadsheet document."}]
    path = Path(document["path"])
    if not path.exists():
        return False, {"document": document}, [{"code": INPUT_FILE_NOT_FOUND, "message": f"Input file not found: {path}"}]
    bounds, range_error = validate_spreadsheet_read_range(range_address)
    if range_error:
        return False, {}, [range_error]
    range_address = range_address.strip()
    min_col, min_row, max_col, max_row = bounds

    formulas_book = load_workbook(path, data_only=False, read_only=True)
    cached_book = load_workbook(path, data_only=True, read_only=True)
    workbook_epoch = formulas_book.epoch
    try:
        selected_sheet = sheet_name or formulas_book.sheetnames[0]
        if selected_sheet not in formulas_book.sheetnames:
            return False, {"document": document}, [{"code": "SHEET_NOT_FOUND", "message": f"Sheet not found: {selected_sheet}"}]
        formula_sheet = formulas_book[selected_sheet]
        cached_sheet = cached_book[selected_sheet]
        cells = []
        for row in range(min_row, max_row + 1):
            for column in range(min_col, max_col + 1):
                formula_cell = formula_sheet.cell(row, column)
                cached_cell = cached_sheet.cell(row, column)
                formula = formula_cell.value if formula_cell.data_type == "f" else None
                cells.append({
                    "address": f"{get_column_letter(column)}{row}",
                    "formula": formula,
                    "saved_cached_value": _json_value(cached_cell.value),
                    "number_format": formula_cell.number_format,
                    "saved_value_type": cached_cell.data_type,
                })
    finally:
        formulas_book.close()
        cached_book.close()

    capabilities = probe_wps_capabilities()
    prog_id = capabilities["components"]["spreadsheets"]["selected_prog_id"]
    if not prog_id:
        return False, {"document_id": document_id, "path": str(path)}, [{"code": COM_BACKEND_UNAVAILABLE, "message": "No registered WPS ProgID detected for spreadsheets."}]
    params_json = json.dumps({
        "prog_id": prog_id, "path": str(path.resolve()), "sheet": selected_sheet,
        "range": range_address, "start_row": min_row, "start_column": min_col,
        "row_count": max_row - min_row + 1, "column_count": max_col - min_col + 1,
    }, ensure_ascii=True)
    script = f"""
$ErrorActionPreference = 'Stop'
$params = @'
{params_json}
'@ | ConvertFrom-Json
$app = $null
$workbook = $null
try {{
  $app = New-Object -ComObject $params.prog_id
  try {{ $app.Visible = $false }} catch {{ }}
  try {{ $app.DisplayAlerts = $false }} catch {{ }}
  try {{ $app.AskToUpdateLinks = $false }} catch {{ }}
  try {{ $app.EnableEvents = $false }} catch {{ }}
  $workbook = $app.Workbooks.Open($params.path, 0, $true)
  $sheet = $workbook.Worksheets.Item($params.sheet)
  try {{ $app.CalculateFullRebuild() }} catch {{ $app.CalculateFull() }}
  $cells = @()
  for ($r = 0; $r -lt [int]$params.row_count; $r++) {{
    for ($c = 0; $c -lt [int]$params.column_count; $c++) {{
      $cell = $sheet.Cells.Item([int]$params.start_row + $r, [int]$params.start_column + $c)
      $formula = $null
      if ($cell.HasFormula) {{ $formula = [string]$cell.Formula }}
      $raw = $cell.Value2
      $display = [string]$cell.Text
      if ($null -eq $raw) {{ $kind = 'empty' }}
      elseif ($display.StartsWith('#')) {{ $kind = 'error' }}
      elseif ($raw -is [string]) {{ $kind = 'text' }}
      elseif ($raw -is [bool]) {{ $kind = 'boolean' }}
      else {{ $kind = 'number' }}
      $cells += [pscustomobject]@{{
        address = [string]$cell.Address($false, $false)
        formula = $formula
        recalculated_value = $raw
        recalculated_type = $kind
        displayed_text = $display
        number_format = [string]$cell.NumberFormat
      }}
    }}
  }}
  [pscustomobject]@{{ ok = $true; backend = 'powershell-com'; prog_id = $params.prog_id; sheet = $sheet.Name; cells = $cells }} | ConvertTo-Json -Depth 8 -Compress
}} catch {{
  [pscustomobject]@{{ ok = $false; backend = 'powershell-com'; error_type = $_.Exception.GetType().FullName; error_message = $_.Exception.Message; position = $_.InvocationInfo.PositionMessage }} | ConvertTo-Json -Depth 5 -Compress
  exit 2
}} finally {{
  if ($workbook -ne $null) {{ try {{ $workbook.Close($false) }} catch {{ }}; try {{ [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($workbook) }} catch {{ }} }}
  if ($app -ne $null) {{ try {{ {QUIT_IF_IDLE} }} catch {{ }}; try {{ [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($app) }} catch {{ }} }}
  [GC]::Collect(); [GC]::WaitForPendingFinalizers()
}}
"""
    script_path = None
    try:
        with tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False, encoding="utf-8-sig") as script_file:
            script_file.write(script)
            script_path = script_file.name
        completed = subprocess.run(
            [powershell_executable(), "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", script_path],
            check=False, capture_output=True, encoding="utf-8", errors="replace", text=True, timeout=120,
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        return False, {"document_id": document_id, "path": str(path)}, [{"code": COM_OPERATION_FAILED, "message": str(exc)}]
    finally:
        if script_path:
            Path(script_path).unlink(missing_ok=True)
    try:
        payload = json.loads(completed.stdout)
    except json.JSONDecodeError as exc:
        return False, {"document_id": document_id, "path": str(path)}, [{"code": COM_OPERATION_FAILED, "message": f"WPS returned invalid JSON: {exc}"}]
    if completed.returncode != 0 or not payload.get("ok"):
        return False, {"document_id": document_id, "path": str(path), "diagnostic": payload}, [{"code": COM_OPERATION_FAILED, "message": payload.get("error_message", "WPS spreadsheet inspection failed.")}]

    raw_live_cells = payload.get("cells", [])
    if isinstance(raw_live_cells, dict):
        raw_live_cells = [raw_live_cells]
    live_cells = {item["address"].replace("$", ""): item for item in raw_live_cells}
    result_cells = []
    for saved in cells:
        live = live_cells.get(saved["address"], {})
        recalculated = _json_value(live.get("recalculated_value"))
        cached = saved["saved_cached_value"]
        result_cells.append({
            **saved,
            "recalculated_value": recalculated,
            "recalculated_type": live.get("recalculated_type"),
            "displayed_text": live.get("displayed_text"),
            "wps_formula": live.get("formula"),
            "formula_matches_file": live.get("formula") == saved["formula"],
            "cache_matches_recalculated": _cache_matches(cached, recalculated, workbook_epoch) if saved["formula"] else None,
        })
    result = {
        "document_id": document_id, "component": "spreadsheets", "path": str(path),
        "backend": payload.get("backend"), "sheet": selected_sheet, "range": range_address,
        "read_only": True, "saved": False, "calculation_performed_in_memory": True,
        "cells": result_cells,
        "cache_mismatch_count": sum(item["cache_matches_recalculated"] is False for item in result_cells),
    }
    return True, result, []
