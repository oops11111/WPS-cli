"""What an open-document write left in the saved file, so the result proves itself (P10-002).

After a write, WPS has saved the file and the command has read the write back in WPS. Agents still read the document
again to check it (P9-006: 38 of 50 off-path calls came after the write), because the result showed only the
selection. verified_state compares the saved file with the pre-write backup and says which paragraphs, text boxes or
cells changed, from what to what, and that nothing else did.
"""

from __future__ import annotations

import difflib
from pathlib import Path
from typing import Any

from openpyxl.utils.cell import (
    column_index_from_string,
    coordinate_from_string,
    get_column_letter,
    range_boundaries,
)

from .document_text import docx_body_story_paragraphs
from .ooxml import load_workbook_guarded
from .presentation_text import pptx_text_objects

MAX_CHANGES = 20
MAX_TEXT = 500
MAX_SURROUNDING = 10
MAX_SURROUNDING_TEXT = 100

# The commands whose successful result carries verified_state.
VERIFIED_COMMANDS = (
    "writer-selection-replace",
    "writer-selection-insert",
    "writer-selection-delete",
    "presentation-selection-replace",
    "presentation-selection-insert",
    "presentation-selection-delete",
    "spreadsheet-selection-replace",
    "spreadsheet-selection-delete",
    "save-open-document",
    "save-as-open-document",
    "export-open-document-text",
)

VERIFIED_NOTE = (
    "The result is already verified: data.verified_state is read from the saved file after the write (for edits, "
    "compared with the pre-write backup) and names what changed, from what to what, and whether anything else "
    "changed; its check sentence is repeated in the summary. There is no need to call snapshot_document, "
    "validate_document, spreadsheet_read, export_open_document_text or operation to check the write."
)


def add_verified_check(command: str, ok: bool, summary: str, data: Any) -> str:
    """The summary of a successful write, followed by its verified_state check sentence."""
    if not ok or command not in VERIFIED_COMMANDS or not isinstance(data, dict):
        return summary
    state = data.get("verified_state")
    check = state.get("check") if isinstance(state, dict) else None
    if not check or check in summary:
        return summary
    return f"{summary} {check}".strip()


def _clip(value: Any) -> Any:
    if isinstance(value, str) and len(value) > MAX_TEXT:
        return value[:MAX_TEXT] + "…"
    return value


def _quote(value: Any) -> str:
    if value is None:
        return "(empty)"
    text = str(_clip(value))
    return f"'{text}'"


def _paragraph_changes(before: list[str], after: list[str]) -> list[dict[str, Any]]:
    changes: list[dict[str, Any]] = []
    matcher = difflib.SequenceMatcher(a=before, b=after, autojunk=False)
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            continue
        for offset in range(max(i2 - i1, j2 - j1)):
            old = before[i1 + offset] if i1 + offset < i2 else None
            new = after[j1 + offset] if j1 + offset < j2 else None
            changes.append({"paragraph": (j1 + offset + 1) if new is not None else None, "before": old, "after": new})
    return changes


def _text_object_changes(before: list[dict[str, Any]], after: list[dict[str, Any]]) -> list[dict[str, Any]]:
    def keyed(objects: list[dict[str, Any]]) -> dict[tuple[int, int], str]:
        return {(int(item["slide_index"]), int(item["object_index"])): str(item.get("text") or "") for item in objects}

    old, new = keyed(before), keyed(after)
    changes = []
    for key in sorted(set(old) | set(new)):
        if old.get(key) != new.get(key):
            changes.append({"slide": key[0], "text_box": key[1], "before": old.get(key), "after": new.get(key)})
    return changes


def _workbook_cells(path: Path) -> dict[tuple[str, str], Any]:
    workbook = load_workbook_guarded(path)
    try:
        return {
            (sheet.title, cell.coordinate): cell.value
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        }
    finally:
        workbook.close()


def _cell_changes(before: dict[tuple[str, str], Any], after: dict[tuple[str, str], Any]) -> list[dict[str, Any]]:
    def order(key: tuple[str, str]) -> tuple[str, int, int]:
        column, row = coordinate_from_string(key[1])
        return key[0], row, column_index_from_string(column)

    changes = []
    for key in sorted(set(before) | set(after), key=order):
        if before.get(key) != after.get(key):
            changes.append({"sheet": key[0], "cell": key[1], "before": before.get(key), "after": after.get(key)})
    return changes


def file_changes(path: str | Path, backup_path: str | Path) -> tuple[str, list[dict[str, Any]]] | None:
    """The paragraphs, text boxes or cells that differ between the backup and the saved file; None without a reader."""
    current, backup = Path(path), Path(backup_path)
    suffix = current.suffix.lower()
    if suffix == ".docx":
        return "paragraphs", _paragraph_changes(docx_body_story_paragraphs(backup), docx_body_story_paragraphs(current))
    if suffix == ".pptx":
        return "text_boxes", _text_object_changes(pptx_text_objects(backup), pptx_text_objects(current))
    if suffix == ".xlsx":
        return "cells", _cell_changes(_workbook_cells(backup), _workbook_cells(current))
    return None


def _in_selection(change: dict[str, Any], sheet: str | None, address: str | None) -> bool:
    if not sheet or not address or change.get("sheet") != sheet:
        return False
    column, row = coordinate_from_string(change["cell"])
    column_number = column_index_from_string(column)
    for area in str(address).replace("$", "").split(","):
        min_col, min_row, max_col, max_row = range_boundaries(area.strip())
        if min_col <= column_number <= max_col and min_row <= row <= max_row:
            return True
    return False


def _describe(kind: str, change: dict[str, Any]) -> str:
    if kind == "paragraphs":
        where = f"paragraph {change['paragraph']}" if change.get("paragraph") else "a paragraph (removed)"
    elif kind == "text_boxes":
        where = f"text box {change['text_box']} on slide {change['slide']}"
    else:
        where = f"{change['sheet']}!{change['cell']}"
    return f"{where} changed from {_quote(change['before'])} to {_quote(change['after'])}"


def _state(
    path: Path,
    kind: str,
    changes: list[dict[str, Any]],
    passed: bool,
    expectation: str,
    others: int = 0,
) -> dict[str, Any]:
    """`others` counts the changes the write does not explain; they are listed after the ones it does."""
    shown = [{key: _clip(value) for key, value in change.items()} for change in changes[:MAX_CHANGES]]
    if changes:
        listed = "; ".join(_describe(kind, change) for change in changes[:3])
        more = f"; and {len(changes) - 3} more" if len(changes) > 3 else ""
        found = f"{listed}{more}"
    else:
        found = f"no {kind.replace('_', ' ')} changed"
    unit = kind.replace("_", " ")
    if passed and not others:
        verdict = f"This matches the write ({expectation}); no other content changed."
    elif passed:
        verdict = (
            f"The write is there ({expectation}); the file also has {others} other change(s) in its {unit} that the "
            "write does not explain, which WPS can make when it saves (see changes)."
        )
    else:
        verdict = f"This does not match the write ({expectation})."
    return {
        "source": "saved_file",
        "compared_with": "pre-write backup",
        "path": str(path),
        "unit": kind,
        "changes": shown,
        "change_count": len(changes),
        "other_changes": others,
        "changes_truncated": len(changes) > MAX_CHANGES,
        "passed": passed,
        "check": f"The saved file was compared with the pre-write backup: {found}. {verdict}",
    }


def _unavailable(path: Path, reason: str) -> dict[str, Any]:
    return {
        "source": "wps_read_back",
        "path": str(path),
        "passed": None,
        "check": f"The write was read back in WPS after saving; the saved file was not compared ({reason}).",
    }


def text_write_state(path: str | Path, backup_path: str | None, removed: str, added: str) -> dict[str, Any]:
    """Writer and Presentation: one paragraph or text box changed, by replacing `removed` with `added`."""
    target = Path(path)
    if not backup_path:
        return _unavailable(target, "no pre-write backup")
    try:
        compared = file_changes(target, backup_path)
    except Exception as exc:  # a file the reader cannot parse leaves the state unknown, not wrong
        return _unavailable(target, f"{type(exc).__name__}: {exc}")
    if compared is None:
        return _unavailable(target, f"no reader for {target.suffix or 'this file type'}")
    kind, changes = compared

    def made_by_the_write(change: dict[str, Any]) -> bool:
        before, after = change["before"], change["after"]
        return (
            before is not None
            and after is not None
            and (not removed or removed in before)
            and (not added or added in after)
            and len(after) - len(before) == len(added) - len(removed)
        )

    matched = [change for change in changes if made_by_the_write(change)][:1]
    passed = bool(matched)
    changes = matched + [change for change in changes if change not in matched]
    unit = "paragraph" if kind == "paragraphs" else "text box"
    if removed and added:
        expectation = f"one {unit} with {_quote(removed)} replaced by {_quote(added)}"
    elif added:
        expectation = f"one {unit} with {_quote(added)} inserted"
    else:
        expectation = f"one {unit} with {_quote(removed)} deleted"
    return _state(target, kind, changes, passed, expectation, others=len(changes) - len(matched))


def _bounds(address: str) -> tuple[int, int, int, int]:
    areas = [range_boundaries(area.strip()) for area in str(address).replace("$", "").split(",") if area.strip()]
    return (
        min(area[0] for area in areas),
        min(area[1] for area in areas),
        max(area[2] for area in areas),
        max(area[3] for area in areas),
    )


def _short(value: Any) -> Any:
    if isinstance(value, str) and len(value) > MAX_SURROUNDING_TEXT:
        return value[:MAX_SURROUNDING_TEXT] + "…"
    return value


def surroundings(
    before: dict[tuple[str, str], Any], after: dict[tuple[str, str], Any], sheet: str, address: str
) -> dict[str, Any]:
    """The cells around a written range in the saved file: one row and column on each side, and the header row.

    P10-006: after a spreadsheet write, agents read the range around it (A1:D5) to see that the header and the
    neighbouring cells were untouched; verified_state listed only the cells that changed (P11-003).
    """
    min_col, min_row, max_col, max_row = _bounds(address)
    first_col, first_row = max(1, min_col - 1), max(1, min_row - 1)
    last_col = min(max_col + 1, first_col + MAX_SURROUNDING - 1)
    last_row = min(max_row + 1, first_row + MAX_SURROUNDING - 1)

    def row_values(row: int) -> list[Any]:
        return [
            _short(after.get((sheet, f"{get_column_letter(column)}{row}"))) for column in range(first_col, last_col + 1)
        ]

    changed_outside = []
    for row in range(first_row, last_row + 1):
        for column in range(first_col, last_col + 1):
            coordinate = f"{get_column_letter(column)}{row}"
            inside = _in_selection({"sheet": sheet, "cell": coordinate}, sheet, address)
            if not inside and before.get((sheet, coordinate)) != after.get((sheet, coordinate)):
                changed_outside.append(coordinate)
    area = f"{get_column_letter(first_col)}{first_row}:{get_column_letter(last_col)}{last_row}"
    result: dict[str, Any] = {
        "sheet": sheet,
        "range": area,
        "rows": [row_values(row) for row in range(first_row, last_row + 1)],
        "changed_outside_selection": changed_outside,
        "truncated": last_col < max_col + 1 or last_row < max_row + 1,
    }
    if first_row > 1:
        header = f"{get_column_letter(first_col)}1:{get_column_letter(last_col)}1"
        result["header_range"] = header
        result["header"] = row_values(1)
    return result


def cell_write_state(
    path: str | Path,
    backup_path: str | None,
    sheet: str | None,
    address: str | None,
    *,
    cleared: bool = False,
    object_kind: str | None = None,
    read_back_text: str | None = None,
) -> dict[str, Any]:
    """Spreadsheets: only cells inside the selection changed; for a shape or chart write, no cell changed."""
    target = Path(path)
    if not backup_path:
        return _unavailable(target, "no pre-write backup")
    try:
        compared = file_changes(target, backup_path)
    except Exception as exc:  # see text_write_state
        return _unavailable(target, f"{type(exc).__name__}: {exc}")
    if compared is None:
        return _unavailable(target, f"no reader for {target.suffix or 'this file type'}")
    kind, changes = compared
    if object_kind:
        passed = not changes
        others = len(changes)
        expectation = (
            f"the {object_kind} on sheet {sheet or '?'} was written, so no cell should change; WPS read the "
            f"{object_kind} back as {_quote(read_back_text)} after saving"
        )
    else:
        inside = [change for change in changes if _in_selection(change, sheet, address)]
        others = len(changes) - len(inside)
        passed = not others and (not cleared or all(change["after"] is None for change in changes))
        changes = inside + [change for change in changes if change not in inside]
        action = "cleared" if cleared else "written"
        expectation = f"only cells inside {sheet}!{address} {action}"
    state = _state(target, kind, changes, passed, expectation, others=others)
    if not object_kind and sheet and address and target.suffix.lower() == ".xlsx":
        try:
            around = surroundings(_workbook_cells(Path(backup_path)), _workbook_cells(target), sheet, address)
        except Exception as exc:  # see text_write_state; the comparison above still stands
            state["surroundings_error"] = f"{type(exc).__name__}: {exc}"
        else:
            state["surroundings"] = around
            if not around["changed_outside_selection"]:
                where = around["range"] + (
                    f" and the header row {around['header_range']}" if "header" in around else ""
                )
                state["check"] += f" The cells around it ({where}) are listed in surroundings and did not change."
    return state


def saved_document_state(path: str | Path, backup_path: str | None) -> dict[str, Any]:
    """save-open-document: what the save wrote compared with the file before it."""
    target = Path(path)
    if not backup_path:
        return document_summary_state(target)
    try:
        compared = file_changes(target, backup_path)
    except Exception as exc:  # see text_write_state
        return _unavailable(target, f"{type(exc).__name__}: {exc}")
    if compared is None:
        return _unavailable(target, f"no reader for {target.suffix or 'this file type'}")
    kind, changes = compared
    return _state(target, kind, changes, True, "the save writes what WPS had in the document")


def document_summary_state(path: str | Path) -> dict[str, Any]:
    """save-as-open-document: the new file opens and what it holds."""
    target = Path(path)
    suffix = target.suffix.lower()
    try:
        if suffix == ".docx":
            paragraphs = docx_body_story_paragraphs(target)
            text = "\n".join(paragraphs)
            found = f"a Writer document with {len(paragraphs)} paragraphs"
        elif suffix == ".pptx":
            objects = pptx_text_objects(target)
            text = "\n".join(str(item.get("text") or "") for item in objects)
            slides = len({item["slide_index"] for item in objects})
            found = f"a presentation with text on {slides} slides"
        elif suffix == ".xlsx":
            cells = _workbook_cells(target)
            sheets = sorted({sheet for sheet, _ in cells})
            text = "\t".join(str(value) for value in list(cells.values())[:50])
            found = f"a workbook with {len(cells)} non-empty cells on sheets {', '.join(sheets) or '(none)'}"
        else:
            return _unavailable(target, f"no reader for {suffix or 'this file type'}")
    except Exception as exc:  # see text_write_state
        return _unavailable(target, f"{type(exc).__name__}: {exc}")
    return {
        "source": "saved_file",
        "path": str(target),
        "text_head": _clip(text[:200]),
        "passed": True,
        "check": f"The saved file was read back: {found}, starting {_quote(text[:80])}.",
    }


def text_file_state(path: str | Path, encoding: str | None) -> dict[str, Any]:
    """export-open-document-text: the exported file decodes and what it starts with."""
    target = Path(path)
    if not encoding:
        return {
            "source": "saved_file",
            "path": str(target),
            "passed": False,
            "check": "The exported file was read back, but its encoding could not be recognised.",
        }
    try:
        data = target.read_bytes()
    except OSError as exc:
        return _unavailable(target, f"{type(exc).__name__}: {exc}")
    codec = {"utf-16": "utf-16", "utf-8-sig": "utf-8-sig", "utf-16-le": "utf-16-le"}.get(encoding, encoding)
    text = data.decode(codec, errors="replace").replace("\r\n", "\n")
    return {
        "source": "saved_file",
        "path": str(target),
        "text_encoding": encoding,
        "characters": len(text),
        "text_head": _clip(text[:200]),
        "passed": True,
        "check": f"The exported file was read back as {encoding}: {len(text)} characters, starting {_quote(text[:80])}.",
    }
