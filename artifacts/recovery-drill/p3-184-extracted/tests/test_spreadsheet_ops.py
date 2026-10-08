import hashlib
import os
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Protection
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from wps_ai_agent_cli.sessions import register_document
from wps_ai_agent_cli.spreadsheet_ops import (
    list_spreadsheet_sheets,
    create_spreadsheet_sheet,
    delete_spreadsheet_sheet,
    copy_spreadsheet_sheet,
    set_spreadsheet_sheet_tab_color,
    read_spreadsheet_range,
    rename_spreadsheet_sheet,
    set_spreadsheet_sheet_visibility,
    write_spreadsheet_formulas,
    write_spreadsheet_range,
)


class SpreadsheetOpsTests(unittest.TestCase):
    def test_read_range_preserves_1904_dates_and_locale_tagged_formats(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "dates.xlsx"
            workbook = Workbook()
            workbook.epoch = CALENDAR_MAC_1904
            sheet = workbook.active
            sheet["A1"] = datetime(2024, 2, 29)
            sheet["A1"].number_format = "[$-409]mmm d, yyyy"
            sheet["B1"] = 0.125
            sheet["B1"].number_format = "[$-409]0.0%"
            workbook.save(document)
            workbook.close()
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)

            ok, result, errors = read_spreadsheet_range(record["document_id"], "A1:B1", workspace=workspace)

            self.assertTrue(ok, errors)
            self.assertEqual(result["values"][0], [datetime(2024, 2, 29), 0.125])
            self.assertEqual([item["format_category"] for item in result["cell_metadata"][0]], ["date", "percentage"])
            self.assertIn("[$-409]", result["cell_metadata"][0][0]["number_format"])
            self.assertIn("[$-409]", result["cell_metadata"][0][1]["number_format"])
    def test_inventory_reads_macro_enabled_archive_without_saving_or_wps(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "macro-enabled.xlsm"
            workbook = Workbook()
            workbook.active.title = "MacroData"
            workbook.active["A1"] = "preserve"
            workbook.active["B2"] = "=1+1"
            workbook.save(document)
            workbook.close()
            with ZipFile(document) as archive:
                entries = {name: archive.read(name) for name in archive.namelist()}
            ignored = b"<ignoredErrors>" + b"".join(
                f'<ignoredError sqref="A{row_index}" numberStoredAsText="1"/>'.encode("ascii")
                for row_index in range(1, 1_003)
            ) + b"</ignoredErrors>"
            entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(b"</worksheet>", ignored + b"</worksheet>")
            with ZipFile(document, "w") as archive:
                for member, content in entries.items():
                    archive.writestr(member, content)
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)
            before = hashlib.sha256(document.read_bytes()).hexdigest()
            with patch("wps_ai_agent_cli.spreadsheet_ops.probe_wps_capabilities") as probe:
                ok, result, errors = list_spreadsheet_sheets(record["document_id"], workspace)
            self.assertTrue(ok, errors)
            self.assertEqual(result["sheets"][0]["name"], "MacroData")
            self.assertEqual(result["sheets"][0]["populated_cell_count"], 2)
            self.assertEqual(result["sheets"][0]["formula_count"], 1)
            self.assertEqual(len(result["sheets"][0]["ignored_errors"]), 1_000)
            self.assertEqual(result["sheets"][0]["ignored_error_count"], 1_002)
            self.assertTrue(result["sheets"][0]["ignored_errors_truncated"])
            self.assertFalse(result["saved"])
            probe.assert_not_called()
            self.assertEqual(hashlib.sha256(document.read_bytes()).hexdigest(), before)

    def test_tab_color_dry_run_validates_hex_and_sheet_before_mutation(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "color.xlsx"
            workbook = Workbook()
            workbook.active.title = "Summary"
            workbook.save(document)
            workbook.close()
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)
            with patch("wps_ai_agent_cli.spreadsheet_ops.create_backup") as backup, patch("wps_ai_agent_cli.spreadsheet_ops._run_spreadsheet_tab_color_com") as com:
                ok, result, errors, replayed = set_spreadsheet_sheet_tab_color(record["document_id"], "Summary", "#12abef", "color-dry", True, workspace)
                self.assertTrue(ok, errors)
                self.assertFalse(replayed)
                self.assertEqual(result["color"], "#12ABEF")
                for color in ("red", "#123", "#12345678", "#GG0000", ""):
                    ok, _result, errors, _ = set_spreadsheet_sheet_tab_color(record["document_id"], "Summary", color, f"color-invalid-{color}", workspace=workspace)
                    self.assertFalse(ok)
                    self.assertEqual(errors[0]["code"], "INVALID_TAB_COLOR")
                ok, _result, errors, _ = set_spreadsheet_sheet_tab_color(record["document_id"], "Missing", "#123456", "color-missing", workspace=workspace)
                self.assertFalse(ok)
                self.assertEqual(errors[0]["code"], "SHEET_NOT_FOUND")
                backup.assert_not_called()
                com.assert_not_called()

    def test_tab_color_commit_verifies_readback_and_replay_binding(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "color.xlsx"
            workbook = Workbook()
            workbook.active.title = "Summary"
            workbook.save(document)
            workbook.close()
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)

            def simulated_wps(path, name, color):
                edited = load_workbook(path)
                edited[name].sheet_properties.tabColor = f"FF{color[1:]}" if color != "none" else None
                actual_color = "none" if color == "none" else color.upper()
                edited.save(path)
                edited.close()
                return {"ok": True, "errors": [], "data": {"backend": "mock-wps", "sheet_name": name, "tab_color": actual_color}}

            with patch("wps_ai_agent_cli.spreadsheet_ops._run_spreadsheet_tab_color_com", side_effect=simulated_wps) as com:
                ok, result, errors, replayed = set_spreadsheet_sheet_tab_color(record["document_id"], "Summary", "#abcdef", "color-commit", workspace=workspace)
                self.assertTrue(ok, errors)
                self.assertFalse(replayed)
                self.assertEqual(result["tab_color_read_back"], "#ABCDEF")
                self.assertTrue(result["validation_passed"])
                self.assertTrue(Path(result["backup"]["backup_path"]).exists())
                ok, _result, errors, replayed = set_spreadsheet_sheet_tab_color(record["document_id"], "Summary", "#ABCDEF", "color-commit", workspace=workspace)
                self.assertTrue(ok, errors)
                self.assertTrue(replayed)
                ok, _result, errors, _ = set_spreadsheet_sheet_tab_color(record["document_id"], "Summary", "#001122", "color-commit", workspace=workspace)
                self.assertFalse(ok)
                self.assertEqual(errors[0]["code"], "IDEMPOTENCY_CONFLICT")
                self.assertEqual(com.call_count, 1)
    def test_copy_sheet_dry_run_previews_order_and_rejects_invalid_targets(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "copy.xlsx"
            workbook = Workbook()
            workbook.active.title = "Source"
            workbook.create_sheet("Tail")
            workbook.save(document)
            workbook.close()
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)
            with patch("wps_ai_agent_cli.spreadsheet_ops.create_backup") as backup, patch("wps_ai_agent_cli.spreadsheet_ops._run_spreadsheet_copy_sheet_com") as com:
                ok, result, errors, replayed = copy_spreadsheet_sheet(record["document_id"], "Source", "Copy", "copy-dry", 2, True, workspace)
                self.assertTrue(ok, errors)
                self.assertFalse(replayed)
                self.assertEqual(result["sheet_names_after"], ["Source", "Copy", "Tail"])
                for source, new_name, index, code in (("Missing", "Copy", 1, "SHEET_NOT_FOUND"), ("Source", "Tail", 1, "SHEET_NAME_CONFLICT"), ("Source", "Bad/Name", 1, "INVALID_SHEET_NAME"), ("Source", "Copy", 0, "INVALID_SHEET_INDEX")):
                    ok, _result, errors, _ = copy_spreadsheet_sheet(record["document_id"], source, new_name, f"copy-invalid-{source}-{new_name}", index, workspace=workspace)
                    self.assertFalse(ok)
                    self.assertEqual(errors[0]["code"], code)
                backup.assert_not_called()
                com.assert_not_called()

    def test_copy_sheet_commit_verifies_order_content_and_replay_binding(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "copy.xlsx"
            workbook = Workbook()
            workbook.active.title = "Source"
            workbook.active["A1"] = "copied"
            workbook.active["B1"] = 19
            workbook.create_sheet("Tail")
            workbook.save(document)
            workbook.close()
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)

            def simulated_wps(path, source, new_name, index):
                edited = load_workbook(path)
                copied = edited.copy_worksheet(edited[source])
                copied.title = new_name
                edited._sheets.remove(copied)
                edited._sheets.insert(index - 1, copied)
                names = list(edited.sheetnames)
                edited.save(path)
                edited.close()
                return {"ok": True, "errors": [], "data": {"backend": "mock-wps", "sheet_names": names}}

            with patch("wps_ai_agent_cli.spreadsheet_ops._run_spreadsheet_copy_sheet_com", side_effect=simulated_wps) as com:
                ok, result, errors, replayed = copy_spreadsheet_sheet(record["document_id"], "Source", "Copy", "copy-commit", 2, workspace=workspace)
                self.assertTrue(ok, errors)
                self.assertFalse(replayed)
                self.assertTrue(result["validation_passed"])
                self.assertEqual(result["copy_a1_read_back"], "copied")
                self.assertTrue(Path(result["backup"]["backup_path"]).exists())
                ok, _result, errors, replayed = copy_spreadsheet_sheet(record["document_id"], "Source", "Copy", "copy-commit", 2, workspace=workspace)
                self.assertTrue(ok, errors)
                self.assertTrue(replayed)
                ok, _result, errors, _ = copy_spreadsheet_sheet(record["document_id"], "Source", "Different", "copy-commit", 2, workspace=workspace)
                self.assertFalse(ok)
                self.assertEqual(errors[0]["code"], "IDEMPOTENCY_CONFLICT")
                self.assertEqual(com.call_count, 1)
    def test_delete_sheet_dry_run_and_last_sheet_guard_precede_mutation(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "delete.xlsx"
            workbook = Workbook()
            workbook.active.title = "Keep"
            workbook.create_sheet("Remove")
            workbook.save(document)
            workbook.close()
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)
            with patch("wps_ai_agent_cli.spreadsheet_ops.create_backup") as backup, patch("wps_ai_agent_cli.spreadsheet_ops._run_spreadsheet_delete_sheet_com") as com:
                ok, result, errors, replayed = delete_spreadsheet_sheet(record["document_id"], "Remove", "delete-dry", True, workspace)
                self.assertTrue(ok, errors)
                self.assertFalse(replayed)
                self.assertEqual(result["sheet_names_after"], ["Keep"])
                only_sheet = Path(tmp) / "only.xlsx"
                single = Workbook()
                single.active.title = "Only"
                single.save(only_sheet)
                single.close()
                ok, only_record, errors = register_document("spreadsheets", str(only_sheet), workspace)
                self.assertTrue(ok, errors)
                ok, _result, errors, _ = delete_spreadsheet_sheet(only_record["document_id"], "Only", "delete-last", workspace=workspace)
                self.assertFalse(ok)
                self.assertEqual(errors[0]["code"], "LAST_WORKSHEET")
                backup.assert_not_called()
                com.assert_not_called()

    def test_delete_sheet_commit_verifies_order_and_replay_binding(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "delete.xlsx"
            workbook = Workbook()
            workbook.active.title = "Keep"
            workbook.active["A1"] = "preserved"
            workbook.create_sheet("Remove")
            workbook.create_sheet("Tail")
            workbook.save(document)
            workbook.close()
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)

            def simulated_wps(path, name):
                edited = load_workbook(path)
                del edited[name]
                names = list(edited.sheetnames)
                edited.save(path)
                edited.close()
                return {"ok": True, "errors": [], "data": {"backend": "mock-wps", "sheet_names": names}}

            with patch("wps_ai_agent_cli.spreadsheet_ops._run_spreadsheet_delete_sheet_com", side_effect=simulated_wps) as com:
                ok, result, errors, replayed = delete_spreadsheet_sheet(record["document_id"], "Remove", "delete-commit", workspace=workspace)
                self.assertTrue(ok, errors)
                self.assertFalse(replayed)
                self.assertTrue(result["validation_passed"])
                self.assertTrue(Path(result["backup"]["backup_path"]).exists())
                ok, _result, errors, replayed = delete_spreadsheet_sheet(record["document_id"], "Remove", "delete-commit", workspace=workspace)
                self.assertTrue(ok, errors)
                self.assertTrue(replayed)
                ok, _result, errors, _ = delete_spreadsheet_sheet(record["document_id"], "Tail", "delete-commit", workspace=workspace)
                self.assertFalse(ok)
                self.assertEqual(errors[0]["code"], "IDEMPOTENCY_CONFLICT")
                self.assertEqual(com.call_count, 1)
            readback = load_workbook(document, read_only=True, data_only=True)
            self.assertEqual(readback.sheetnames, ["Keep", "Tail"])
            self.assertEqual(readback["Keep"]["A1"].value, "preserved")
            readback.close()
    def test_sheet_visibility_dry_run_and_last_visible_guard_precede_mutation(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "visibility.xlsx"
            workbook = Workbook()
            workbook.active.title = "Visible"
            hidden = workbook.create_sheet("Hidden")
            hidden.sheet_state = "hidden"
            workbook.save(document)
            workbook.close()
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)
            with patch("wps_ai_agent_cli.spreadsheet_ops.create_backup") as backup, patch("wps_ai_agent_cli.spreadsheet_ops._run_spreadsheet_visibility_com") as com:
                ok, result, errors, replayed = set_spreadsheet_sheet_visibility(record["document_id"], "Hidden", True, "visibility-preview", True, workspace)
                self.assertTrue(ok, errors)
                self.assertFalse(replayed)
                self.assertEqual(result["sheet_states_after"], [{"name": "Visible", "visible": True}, {"name": "Hidden", "visible": True}])
                ok, _result, errors, _ = set_spreadsheet_sheet_visibility(record["document_id"], "Visible", False, "visibility-last", workspace=workspace)
                self.assertFalse(ok)
                self.assertEqual(errors[0]["code"], "LAST_VISIBLE_SHEET")
                backup.assert_not_called()
                com.assert_not_called()

    def test_sheet_visibility_commit_verifies_readback_and_replay_binding(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "visibility.xlsx"
            workbook = Workbook()
            workbook.active.title = "Visible"
            hidden = workbook.create_sheet("Hidden")
            hidden.sheet_state = "hidden"
            workbook.save(document)
            workbook.close()
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)

            def simulated_wps(path, name, visible):
                edited = load_workbook(path)
                edited[name].sheet_state = "visible" if visible else "hidden"
                states = [{"name": sheet.title, "visible": sheet.sheet_state == "visible"} for sheet in edited.worksheets]
                edited.save(path)
                edited.close()
                return {"ok": True, "errors": [], "data": {"backend": "mock-wps", "sheet_states": states}}

            with patch("wps_ai_agent_cli.spreadsheet_ops._run_spreadsheet_visibility_com", side_effect=simulated_wps) as com:
                ok, result, errors, replayed = set_spreadsheet_sheet_visibility(record["document_id"], "Hidden", True, "visibility-commit", workspace=workspace)
                self.assertTrue(ok, errors)
                self.assertFalse(replayed)
                self.assertTrue(result["validation_passed"])
                self.assertTrue(Path(result["backup"]["backup_path"]).exists())
                ok, _result, errors, replayed = set_spreadsheet_sheet_visibility(record["document_id"], "Hidden", True, "visibility-commit", workspace=workspace)
                self.assertTrue(ok, errors)
                self.assertTrue(replayed)
                ok, _result, errors, _ = set_spreadsheet_sheet_visibility(record["document_id"], "Hidden", False, "visibility-commit", workspace=workspace)
                self.assertFalse(ok)
                self.assertEqual(errors[0]["code"], "IDEMPOTENCY_CONFLICT")
                self.assertEqual(com.call_count, 1)
    def test_create_sheet_dry_run_validates_name_position_and_duplicate_without_side_effects(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "create.xlsx"
            workbook = Workbook()
            workbook.active.title = "First"
            workbook.create_sheet("Last")
            workbook.save(document)
            workbook.close()
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)
            with patch("wps_ai_agent_cli.spreadsheet_ops.create_backup") as backup, patch("wps_ai_agent_cli.spreadsheet_ops._run_spreadsheet_create_sheet_com") as com:
                ok, preview, errors, replayed = create_spreadsheet_sheet(record["document_id"], "Middle", "create-dry", 2, True, workspace)
                self.assertTrue(ok, errors)
                self.assertFalse(replayed)
                self.assertEqual(preview["sheet_names_after"], ["First", "Middle", "Last"])
                for name, index, code in (("bad/name", 1, "INVALID_SHEET_NAME"), ("First", 1, "SHEET_NAME_CONFLICT"), ("Valid", 0, "INVALID_SHEET_INDEX"), ("Valid", "2", "INVALID_SHEET_INDEX")):
                    ok, _result, errors, _ = create_spreadsheet_sheet(record["document_id"], name, f"invalid-{name}-{index}", index, workspace=workspace)
                    self.assertFalse(ok)
                    self.assertEqual(errors[0]["code"], code)
                backup.assert_not_called()
                com.assert_not_called()

    def test_create_sheet_commit_verifies_order_and_binds_replay_arguments(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "create.xlsx"
            workbook = Workbook()
            workbook.active.title = "First"
            workbook.active["A1"] = "preserved"
            workbook.create_sheet("Last")
            workbook.save(document)
            workbook.close()
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)

            def simulated_wps(path, name, index):
                edited = load_workbook(path)
                edited.create_sheet(name, index - 1)
                names = list(edited.sheetnames)
                edited.save(path)
                edited.close()
                return {"ok": True, "errors": [], "data": {"backend": "mock-wps", "sheet_names": names}}

            with patch("wps_ai_agent_cli.spreadsheet_ops._run_spreadsheet_create_sheet_com", side_effect=simulated_wps) as com:
                ok, result, errors, replayed = create_spreadsheet_sheet(record["document_id"], "Middle", "create-commit", 2, workspace=workspace)
                self.assertTrue(ok, errors)
                self.assertFalse(replayed)
                self.assertTrue(result["validation_passed"])
                self.assertTrue(Path(result["backup"]["backup_path"]).exists())
                ok, _result, errors, replayed = create_spreadsheet_sheet(record["document_id"], "Middle", "create-commit", 2, workspace=workspace)
                self.assertTrue(ok, errors)
                self.assertTrue(replayed)
                ok, _result, errors, replayed = create_spreadsheet_sheet(record["document_id"], "Other", "create-commit", 2, workspace=workspace)
                self.assertFalse(ok)
                self.assertFalse(replayed)
                self.assertEqual(errors[0]["code"], "IDEMPOTENCY_CONFLICT")
                self.assertEqual(com.call_count, 1)
            readback = load_workbook(document, read_only=True, data_only=True)
            self.assertEqual(readback.sheetnames, ["First", "Middle", "Last"])
            self.assertEqual(readback["First"]["A1"].value, "preserved")
            readback.close()
    def test_read_spreadsheet_range_returns_matrix(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Calculation"
            sheet["A1"] = "Item"
            sheet["B1"] = "Value"
            sheet["A2"] = "Sample"
            sheet["B2"] = 18
            workbook.save(document)
            ok, record, _errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok)

            ok, result, errors = read_spreadsheet_range(
                record["document_id"],
                "A1:B2",
                workspace=workspace,
            )

            self.assertTrue(ok)
            self.assertEqual(errors, [])
            self.assertEqual(result["values"], [["Item", "Value"], ["Sample", 18]])
            self.assertEqual(result["row_count"], 2)
            self.assertEqual(result["column_count"], 2)

    def test_list_spreadsheet_sheets_reports_order_visibility_and_dimensions_read_only(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "sheets.xlsx"
            workbook = Workbook()
            first = workbook.active
            first.title = "Summary"
            first["A1"] = "Title"
            first["C4"] = 7
            first.freeze_panes = "B2"
            first.auto_filter.ref = "A1:C4"
            first.print_area = "A1:C4"
            first.print_title_rows = "1:2"
            first.page_setup.orientation = "landscape"
            first.page_setup.paperSize = "9"
            first.row_breaks.append(Break(id=10))
            first.col_breaks.append(Break(id=3))
            validation = DataValidation(type="whole", operator="between", formula1=1, formula2=10, allow_blank=True)
            validation.add("A1:A10")
            first.add_data_validation(validation)
            first.conditional_formatting.add("C4", CellIsRule(operator="greaterThan", formula=[5]))
            workbook.defined_names.add(DefinedName("Revenue", attr_text="'Summary'!$C$4"))
            first.defined_names.add(DefinedName("LocalTitle", attr_text="'Summary'!$A$1"))
            for name_index in range(1_001):
                workbook.defined_names.add(DefinedName(f"Name{name_index:04d}", attr_text="'Summary'!$A$1"))
            workbook.calculation.calcMode = "manual"
            workbook.calculation.fullCalcOnLoad = False
            workbook.calculation.forceFullCalc = True
            hidden = workbook.create_sheet("Archive")
            hidden.sheet_properties.tabColor = "FF12ABEF"
            hidden["B2"] = "old"
            hidden["C3"] = "unlocked"
            hidden["C3"].protection = Protection(locked=False)
            hidden["D4"] = "=1+1"
            hidden.protection.sheet = True
            hidden.sheet_state = "hidden"
            very_hidden = workbook.create_sheet("System")
            very_hidden.sheet_state = "veryHidden"
            very_hidden.protection.sheet = True
            very_hidden.cell(row=1001, column=101, value="sparse")
            for row_index in range(1, 1003):
                very_hidden.merge_cells(start_row=row_index, start_column=200, end_row=row_index, end_column=201)
                very_hidden.row_breaks.append(Break(id=row_index))
            for row_index in range(1, 1_002):
                rule = DataValidation(type="whole")
                rule.add(f"A{row_index}")
                very_hidden.add_data_validation(rule)
                very_hidden.conditional_formatting.add(f"C{row_index}", CellIsRule(operator="greaterThan", formula=[0]))
            workbook.save(document)
            workbook.close()
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)
            before = hashlib.sha256(document.read_bytes()).hexdigest()

            with patch("wps_ai_agent_cli.spreadsheet_ops.probe_wps_capabilities") as probe:
                ok, result, errors = list_spreadsheet_sheets(record["document_id"], workspace)

            self.assertTrue(ok, errors)
            self.assertTrue(result["read_only"])
            self.assertFalse(result["saved"])
            self.assertEqual(result["sheet_count"], 3)
            self.assertEqual([(sheet["sheet_index"], sheet["name"]) for sheet in result["sheets"]], [(1, "Summary"), (2, "Archive"), (3, "System")])
            self.assertEqual(result["sheets"][0]["dimension"], "A1:C4")
            self.assertEqual(result["sheets"][0]["max_row"], 4)
            self.assertEqual(result["sheets"][0]["max_column"], 3)
            self.assertEqual(result["sheets"][0]["freeze_panes"], "B2")
            self.assertEqual(result["sheets"][0]["autofilter_range"], "A1:C4")
            self.assertEqual(result["sheets"][0]["print_area"], "'Summary'!$A$1:$C$4")
            self.assertEqual(result["sheets"][0]["print_title_rows"], "$1:$2")
            self.assertIsNone(result["sheets"][0]["print_title_columns"])
            self.assertEqual(result["sheets"][0]["page_orientation"], "landscape")
            self.assertEqual(result["sheets"][0]["paper_size"], "9")
            self.assertEqual(result["sheets"][0]["row_breaks"], [10])
            self.assertEqual(result["sheets"][0]["row_break_count"], 1)
            self.assertFalse(result["sheets"][0]["row_breaks_truncated"])
            self.assertEqual(result["sheets"][0]["column_breaks"], [3])
            self.assertEqual(result["sheets"][0]["column_break_count"], 1)
            self.assertEqual(result["sheets"][0]["data_validations"], [{
                "ranges": "A1:A10", "type": "whole", "operator": "between", "allow_blank": True,
            }])
            self.assertEqual(result["sheets"][0]["data_validation_count"], 1)
            self.assertFalse(result["sheets"][0]["data_validations_truncated"])
            self.assertEqual(result["sheets"][0]["conditional_formatting"], [{
                "ranges": "C4", "type": "cellIs", "operator": "greaterThan", "priority": 1,
            }])
            self.assertEqual(result["sheets"][0]["conditional_formatting_rule_count"], 1)
            self.assertFalse(result["sheets"][0]["conditional_formatting_truncated"])
            self.assertEqual(result["defined_name_count"], 1_003)
            self.assertTrue(result["defined_names_truncated"])
            self.assertEqual(len(result["defined_names"]), 1_000)
            self.assertEqual(result["defined_names"][0], {"name": "LocalTitle", "scope": "Summary", "target": "'Summary'!$A$1", "hidden": False})
            self.assertEqual(result["defined_names"][1]["name"], "Name0000")
            self.assertEqual(result["calculation"], {
                "mode": "manual", "full_calc_on_load": False,
                "force_full_calc": True, "calc_on_save": None, "iterate": None,
            })
            self.assertEqual(result["sheets"][0]["merged_ranges"], [])
            self.assertEqual(result["sheets"][0]["merged_range_count"], 0)
            self.assertFalse(result["sheets"][0]["merged_ranges_truncated"])
            self.assertEqual(len(result["sheets"][2]["merged_ranges"]), 1_000)
            self.assertEqual(result["sheets"][2]["merged_range_count"], 1_002)
            self.assertEqual(result["sheets"][2]["merged_ranges"][:2], ["GR1:GS1", "GR2:GS2"])
            self.assertTrue(result["sheets"][2]["merged_ranges_truncated"])
            self.assertEqual(len(result["sheets"][2]["row_breaks"]), 1_000)
            self.assertEqual(result["sheets"][2]["row_break_count"], 1_002)
            self.assertTrue(result["sheets"][2]["row_breaks_truncated"])
            self.assertEqual(len(result["sheets"][2]["data_validations"]), 1_000)
            self.assertEqual(result["sheets"][2]["data_validation_count"], 1_001)
            self.assertTrue(result["sheets"][2]["data_validations_truncated"])
            self.assertEqual(len(result["sheets"][2]["conditional_formatting"]), 1_000)
            self.assertEqual(result["sheets"][2]["conditional_formatting_rule_count"], 1_001)
            self.assertTrue(result["sheets"][2]["conditional_formatting_truncated"])
            self.assertEqual([(sheet["visibility"], sheet["visible"]) for sheet in result["sheets"]], [("visible", True), ("hidden", False), ("veryHidden", False)])
            self.assertEqual([sheet["tab_color"] for sheet in result["sheets"]], ["none", "#12ABEF", "none"])
            self.assertEqual([sheet["protection_enabled"] for sheet in result["sheets"]], [False, True, True])
            self.assertEqual([sheet["protected_cell_count"] for sheet in result["sheets"]], [0, 15, None])
            self.assertFalse(result["sheets"][1]["protection_scan_truncated"])
            self.assertIsNone(result["sheets"][2]["protected_cell_count"])
            self.assertTrue(result["sheets"][2]["protection_scan_truncated"])
            self.assertEqual([sheet["populated_cell_count"] for sheet in result["sheets"]], [2, 3, None])
            self.assertEqual([sheet["formula_count"] for sheet in result["sheets"]], [0, 1, None])
            self.assertEqual([sheet["inventory_scan_truncated"] for sheet in result["sheets"]], [False, False, True])
            probe.assert_not_called()
            self.assertEqual(hashlib.sha256(document.read_bytes()).hexdigest(), before)

    def test_read_spreadsheet_range_preserves_values_and_adds_format_metadata(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "formats.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = datetime(2026, 10, 7)
            sheet["A1"].number_format = "yyyy-mm-dd"
            sheet["B1"] = 0.25
            sheet["B1"].number_format = "0.0%"
            sheet["C1"] = 1234.5
            sheet["C1"].number_format = "$#,##0.00"
            workbook.save(document)
            workbook.close()
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)

            ok, result, errors = read_spreadsheet_range(record["document_id"], "A1:D1", workspace=workspace)

            self.assertTrue(ok, errors)
            self.assertEqual(result["values"][0][:3], [datetime(2026, 10, 7), 0.25, 1234.5])
            self.assertEqual([cell["format_category"] for cell in result["cell_metadata"][0]], ["date", "percentage", "currency", "empty"])
            self.assertTrue(result["cell_metadata"][0][0]["is_date"])
            self.assertEqual(result["cell_metadata"][0][1]["number_format"], "0.0%")
            self.assertEqual(result["cell_metadata"][0][3]["address"], "D1")

    def test_read_rejects_invalid_unbounded_and_oversized_ranges_before_opening_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "sample.xlsx"
            Workbook().save(document)
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)

            for address, expected_code in (
                ("A:A", "INVALID_RANGE"),
                ("1:1", "INVALID_RANGE"),
                ("A0", "INVALID_RANGE"),
                ("A1:XFD1048576", "RANGE_TOO_LARGE"),
            ):
                with self.subTest(address=address), patch("wps_ai_agent_cli.spreadsheet_ops.load_workbook") as open_workbook:
                    ok, _result, errors = read_spreadsheet_range(record["document_id"], address, workspace=workspace)
                    self.assertFalse(ok)
                    self.assertEqual(errors[0]["code"], expected_code)
                    open_workbook.assert_not_called()

    def test_writes_reject_invalid_and_oversized_ranges_before_backup_or_wps(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "sample.xlsx"
            Workbook().save(document)
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)
            cases = (("A:A", "INVALID_RANGE"), ("A1:XFD1048576", "RANGE_TOO_LARGE"))
            for operation_name, operation in (
                ("values", lambda address, request: write_spreadsheet_range(record["document_id"], address, [[1]], request, workspace=workspace)),
                ("formulas", lambda address, request: write_spreadsheet_formulas(record["document_id"], address, [["=1"]], [[1]], request, workspace=workspace)),
            ):
                for address, expected_code in cases:
                    with self.subTest(operation=operation_name, address=address), patch("wps_ai_agent_cli.spreadsheet_ops.create_backup") as backup, patch("wps_ai_agent_cli.spreadsheet_ops._run_spreadsheet_write_com") as value_com, patch("wps_ai_agent_cli.spreadsheet_ops._run_spreadsheet_formula_write_com") as formula_com:
                        ok, _result, errors, replayed = operation(address, f"{operation_name}-{address}")
                        self.assertFalse(ok)
                        self.assertFalse(replayed)
                        self.assertEqual(errors[0]["code"], expected_code)
                        backup.assert_not_called()
                        value_com.assert_not_called()
                        formula_com.assert_not_called()

    def test_rename_sheet_dry_run_and_invalid_names_do_not_back_up_or_launch_wps(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "rename.xlsx"
            workbook = Workbook()
            workbook.active.title = "Current"
            workbook.create_sheet("Archive")
            workbook.save(document)
            workbook.close()
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)

            with patch("wps_ai_agent_cli.spreadsheet_ops.create_backup") as backup, patch("wps_ai_agent_cli.spreadsheet_ops._run_spreadsheet_rename_sheet_com") as com:
                ok, preview, errors, replayed = rename_spreadsheet_sheet(record["document_id"], "Current", "Renamed", "rename-dry", dry_run=True, workspace=workspace)
                self.assertTrue(ok, errors)
                self.assertFalse(replayed)
                self.assertEqual(preview["sheet_names_after"], ["Renamed", "Archive"])
                self.assertFalse((workspace / ".wps-agent" / "backups").exists())

                for invalid in ("", "bad/name", "x" * 32, "History", "Archive"):
                    with self.subTest(invalid=invalid):
                        ok, _result, errors, replayed = rename_spreadsheet_sheet(record["document_id"], "Current", invalid, f"invalid-{invalid}", workspace=workspace)
                        self.assertFalse(ok)
                        self.assertFalse(replayed)
                        self.assertIn(errors[0]["code"], {"INVALID_SHEET_NAME", "SHEET_NAME_CONFLICT"})
                backup.assert_not_called()
                com.assert_not_called()

    def test_rename_sheet_commit_verifies_order_and_binds_replay_arguments(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "rename.xlsx"
            workbook = Workbook()
            workbook.active.title = "Current"
            workbook.active["A1"] = "kept"
            workbook.create_sheet("Archive")
            workbook.save(document)
            workbook.close()
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)

            def simulated_wps(path, old_name, new_name):
                edited = load_workbook(path)
                edited[old_name].title = new_name
                names = list(edited.sheetnames)
                edited.save(path)
                edited.close()
                return {"ok": True, "errors": [], "data": {"backend": "mock-wps", "sheet_names": names}}

            with patch("wps_ai_agent_cli.spreadsheet_ops._run_spreadsheet_rename_sheet_com", side_effect=simulated_wps) as com:
                ok, result, errors, replayed = rename_spreadsheet_sheet(record["document_id"], "Current", "Renamed", "rename-commit", workspace=workspace)
                self.assertTrue(ok, errors)
                self.assertFalse(replayed)
                self.assertTrue(result["validation_passed"])
                self.assertEqual(result["sheet_names_read_back"], ["Renamed", "Archive"])
                self.assertTrue(Path(result["backup"]["backup_path"]).exists())
                ok, replay, errors, replayed = rename_spreadsheet_sheet(record["document_id"], "Current", "Renamed", "rename-commit", workspace=workspace)
                self.assertTrue(ok, errors)
                self.assertTrue(replayed)
                self.assertEqual(replay, result)
                ok, _result, errors, replayed = rename_spreadsheet_sheet(record["document_id"], "Current", "Other", "rename-commit", workspace=workspace)
                self.assertFalse(ok)
                self.assertFalse(replayed)
                self.assertEqual(errors[0]["code"], "IDEMPOTENCY_CONFLICT")
                self.assertEqual(com.call_count, 1)

    def test_write_spreadsheet_range_dry_run_reports_shape(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "sample.xlsx"
            workbook = Workbook()
            workbook.save(document)
            ok, record, _errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok)

            ok, result, errors, replayed = write_spreadsheet_range(
                record["document_id"],
                "A1:B2",
                [["A", "B"], [1, 2]],
                request_id="sheet-write-dry",
                dry_run=True,
                workspace=workspace,
            )

            self.assertTrue(ok)
            self.assertFalse(replayed)
            self.assertEqual(errors, [])
            self.assertEqual(result["row_count"], 2)
            self.assertEqual(result["column_count"], 2)
            self.assertFalse((workspace / ".wps-agent" / "backups").exists())

    def test_write_spreadsheet_range_rejects_shape_mismatch(self):
        ok, result, errors, replayed = write_spreadsheet_range(
            "doc_missing",
            "A1:B2",
            [["A", "B"]],
            request_id="sheet-write-bad-shape",
            dry_run=True,
        )

        self.assertFalse(ok)
        self.assertFalse(replayed)
        self.assertEqual(result["range_shape"], [2, 2])
        self.assertEqual(result["values_shape"], [1, 2])
        self.assertEqual(errors[0]["code"], "RANGE_SHAPE_MISMATCH")

    def test_write_spreadsheet_formulas_dry_run_reports_shape(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "sample.xlsx"
            workbook = Workbook()
            workbook.save(document)
            ok, record, _errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok)

            ok, result, errors, replayed = write_spreadsheet_formulas(
                record["document_id"],
                "C1:C2",
                [["=A1+B1"], ["=A2+B2"]],
                expected_values=[[3], [7]],
                request_id="sheet-formula-dry",
                dry_run=True,
                workspace=workspace,
            )

            self.assertTrue(ok)
            self.assertFalse(replayed)
            self.assertEqual(errors, [])
            self.assertEqual(result["row_count"], 2)
            self.assertEqual(result["column_count"], 1)
            self.assertFalse((workspace / ".wps-agent" / "backups").exists())

    def test_write_spreadsheet_formulas_rejects_non_formula_string(self):
        ok, result, errors, replayed = write_spreadsheet_formulas(
            "doc_missing",
            "A1:A1",
            [["A1+B1"]],
            expected_values=[[3]],
            request_id="sheet-formula-invalid",
            dry_run=True,
        )

        self.assertFalse(ok)
        self.assertFalse(replayed)
        self.assertEqual(result, {})
        self.assertEqual(errors[0]["code"], "INVALID_ARGUMENT")

    def test_write_spreadsheet_formulas_rejects_expected_shape_mismatch(self):
        ok, result, errors, replayed = write_spreadsheet_formulas(
            "doc_missing",
            "A1:A2",
            [["=A1+B1"], ["=A2+B2"]],
            expected_values=[[3]],
            request_id="sheet-formula-bad-expected",
            dry_run=True,
        )

        self.assertFalse(ok)
        self.assertFalse(replayed)
        self.assertEqual(result["formulas_shape"], [2, 1])
        self.assertEqual(result["expected_values_shape"], [1, 1])
        self.assertEqual(errors[0]["code"], "EXPECTED_SHAPE_MISMATCH")


@unittest.skipUnless(os.environ.get("WPS_AGENT_RUN_INTEGRATION") == "1", "Requires explicit local WPS integration run")
class SpreadsheetRenameWpsIntegrationTests(unittest.TestCase):
    def test_wps_renames_sheet_and_preserves_workbook_order(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "rename.xlsx"
            workbook = Workbook()
            workbook.active.title = "Current"
            workbook.active["A1"] = 42
            workbook.create_sheet("Archive")
            workbook.save(document)
            workbook.close()
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)

            ok, result, errors, replayed = rename_spreadsheet_sheet(record["document_id"], "Current", "Renamed", "wps-rename", workspace=workspace)

            self.assertTrue(ok, errors)
            self.assertFalse(replayed)
            self.assertEqual(result["sheet_names_read_back"], ["Renamed", "Archive"])
            readback = load_workbook(document, read_only=True, data_only=True)
            self.assertEqual(readback["Renamed"]["A1"].value, 42)
            self.assertEqual(readback.sheetnames, ["Renamed", "Archive"])
            readback.close()


@unittest.skipUnless(os.environ.get("WPS_AGENT_RUN_INTEGRATION") == "1", "Requires explicit local WPS integration run")
class SpreadsheetCreateWpsIntegrationTests(unittest.TestCase):
    def test_wps_creates_sheet_at_requested_position_and_preserves_existing_data(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "create.xlsx"
            workbook = Workbook()
            workbook.active.title = "First"
            workbook.active["A1"] = 42
            workbook.create_sheet("Last")
            workbook.save(document)
            workbook.close()
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)

            ok, result, errors, replayed = create_spreadsheet_sheet(record["document_id"], "Middle", "wps-create", 2, workspace=workspace)

            self.assertTrue(ok, errors)
            self.assertFalse(replayed)
            self.assertEqual(result["sheet_names_read_back"], ["First", "Middle", "Last"])
            readback = load_workbook(document, read_only=True, data_only=True)
            self.assertEqual(readback["First"]["A1"].value, 42)
            self.assertEqual(readback.sheetnames, ["First", "Middle", "Last"])
            readback.close()


@unittest.skipUnless(os.environ.get("WPS_AGENT_RUN_INTEGRATION") == "1", "Requires explicit local WPS integration run")
class SpreadsheetStructureWpsIntegrationTests(unittest.TestCase):
    def test_wps_updates_visibility_and_deletes_middle_sheet(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "structure.xlsx"
            workbook = Workbook()
            workbook.active.title = "First"
            workbook.active["A1"] = 17
            middle = workbook.create_sheet("Middle")
            middle["A1"] = "remove"
            workbook.create_sheet("Last")
            workbook.save(document)
            workbook.close()
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)

            ok, visibility, errors, replayed = set_spreadsheet_sheet_visibility(record["document_id"], "Middle", False, "wps-hide-middle", workspace=workspace)
            self.assertTrue(ok, errors)
            self.assertFalse(replayed)
            self.assertTrue(visibility["validation_passed"])
            ok, deleted, errors, replayed = delete_spreadsheet_sheet(record["document_id"], "Middle", "wps-delete-middle", workspace=workspace)
            self.assertTrue(ok, errors)
            self.assertFalse(replayed)
            self.assertEqual(deleted["sheet_names_read_back"], ["First", "Last"])
            readback = load_workbook(document, read_only=True, data_only=True)
            self.assertEqual(readback["First"]["A1"].value, 17)
            self.assertEqual(readback.sheetnames, ["First", "Last"])
            readback.close()

    def test_wps_copies_worksheet_content_to_requested_position(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "copy.xlsx"
            workbook = Workbook()
            workbook.active.title = "Source"
            workbook.active["A1"] = "preserve me"
            workbook.active["B1"] = 24
            workbook.create_sheet("Tail")
            workbook.save(document)
            workbook.close()
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)

            ok, result, errors, replayed = copy_spreadsheet_sheet(record["document_id"], "Source", "Copy", "wps-copy", 2, workspace=workspace)

            self.assertTrue(ok, errors)
            self.assertFalse(replayed)
            self.assertTrue(result["validation_passed"])
            self.assertEqual(result["sheet_names_read_back"], ["Source", "Copy", "Tail"])
            readback = load_workbook(document, read_only=True, data_only=True)
            self.assertEqual(readback["Copy"]["A1"].value, "preserve me")
            self.assertEqual(readback["Copy"]["B1"].value, 24)
            readback.close()

    def test_wps_sets_and_clears_worksheet_tab_color(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "color.xlsx"
            workbook = Workbook()
            workbook.active.title = "Summary"
            workbook.save(document)
            workbook.close()
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)

            ok, result, errors, replayed = set_spreadsheet_sheet_tab_color(record["document_id"], "Summary", "#34A1BC", "wps-color", workspace=workspace)

            self.assertTrue(ok, errors)
            self.assertFalse(replayed)
            self.assertEqual(result["tab_color_read_back"], "#34A1BC")
            ok, result, errors, _ = set_spreadsheet_sheet_tab_color(record["document_id"], "Summary", "none", "wps-color-clear", workspace=workspace)
            self.assertTrue(ok, errors)
            self.assertEqual(result["tab_color_read_back"], "none")


if __name__ == "__main__":
    unittest.main()
