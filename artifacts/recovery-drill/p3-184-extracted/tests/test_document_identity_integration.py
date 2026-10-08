import os
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from wps_ai_agent_cli.backups import create_backup
from wps_ai_agent_cli.sessions import file_identity, get_document, register_document
from wps_ai_agent_cli.writer_ops import writer_fill_bookmark


class DocumentIdentityIntegrationTests(unittest.TestCase):
    @unittest.skipUnless(os.environ.get("WPS_AGENT_RUN_INTEGRATION") == "1", "set WPS_AGENT_RUN_INTEGRATION=1 to launch local WPS Writer")
    def test_wps_fill_refreshes_identity_for_next_backup(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "identity-wps.docx"
            document = Document()
            paragraph = document.add_paragraph("Dear ")
            run = paragraph.add_run("old")
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), "1")
            start.set(qn("w:name"), "Client")
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), "1")
            run._r.addprevious(start)
            run._r.addnext(end)
            paragraph.add_run("!")
            document.save(path)
            ok, record, errors = register_document("writer", str(path), tmp)
            self.assertTrue(ok, errors)
            filled, result, fill_errors, _ = writer_fill_bookmark(
                record["document_id"], "Client", "Updated", "identity-wps-fill", workspace=tmp,
            )
            self.assertTrue(filled, (fill_errors, result))
            self.assertTrue(result["readback_passed"])
            current = get_document(record["document_id"], tmp)
            self.assertEqual({key: current[key] for key in file_identity(path)}, file_identity(path))
            backed_up, _, backup_errors, _ = create_backup(
                record["document_id"], "identity-post-fill-backup", workspace=tmp,
            )
            self.assertTrue(backed_up, backup_errors)


if __name__ == "__main__":
    unittest.main()
