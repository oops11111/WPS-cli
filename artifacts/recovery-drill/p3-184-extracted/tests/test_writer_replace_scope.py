import os
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from tests.test_writer_structure import write_document
from wps_ai_agent_cli.document_text import docx_body_paragraph_target, docx_body_paragraphs, docx_body_tables
from wps_ai_agent_cli.sessions import register_document
from wps_ai_agent_cli.writer_ops import writer_replace
from wps_ai_agent_cli.writer_ops import writer_fill_bookmark


class WriterParagraphMappingTests(unittest.TestCase):
    def test_direct_body_target_excludes_table_paragraphs(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "mapping.docx"
            write_document(path, '''<w:p><w:r><w:t>before</w:t></w:r></w:p>
                <w:tbl><w:tr><w:tc><w:p/><w:p/></w:tc></w:tr></w:tbl>
                <w:p><w:r><w:t>target</w:t></w:r></w:p>''')
            self.assertEqual(docx_body_paragraph_target(path, 2), {
                "body_paragraph_count": 2, "expected_text": "target",
            })
            with self.assertRaises(ValueError):
                docx_body_paragraph_target(path, 3)

    def test_bookmark_fill_dry_run_and_preflight_rejections(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bookmark.docx"
            write_document(path, '<w:p><w:r><w:t>Dear </w:t></w:r><w:bookmarkStart w:id="1" w:name="Client"/><w:r><w:t>old</w:t></w:r><w:bookmarkEnd w:id="1"/><w:r><w:t>!</w:t></w:r></w:p>')
            _, record, _ = register_document("writer", str(path), tmp)
            before = path.read_bytes()
            with patch("wps_ai_agent_cli.writer_ops.create_backup") as backup, patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com") as com:
                ok, result, errors, replayed = writer_fill_bookmark(record["document_id"], "Client", "new", "bookmark-preview", dry_run=True, workspace=tmp)
            self.assertTrue(ok)
            self.assertFalse(replayed)
            self.assertEqual(errors, [])
            self.assertEqual(result["current_text_length"], 3)
            self.assertEqual(result["replacement_text_length"], 3)
            backup.assert_not_called()
            com.assert_not_called()
            self.assertEqual(path.read_bytes(), before)
            ok, _, errors, _ = writer_fill_bookmark(record["document_id"], "Missing", "x", "bookmark-missing", dry_run=True, workspace=tmp)
            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "BOOKMARK_NOT_FOUND")

    def test_replaced_file_blocks_bookmark_fill_before_wps(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bookmark.docx"
            replacement = Path(tmp) / "replacement.docx"
            write_document(path, '<w:p><w:bookmarkStart w:id="1" w:name="Client"/>'
                                 '<w:r><w:t>old</w:t></w:r><w:bookmarkEnd w:id="1"/></w:p>')
            _, record, _ = register_document("writer", str(path), tmp)
            replacement.write_bytes(path.read_bytes())
            os.replace(replacement, path)
            with patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com") as com:
                ok, _, errors, _ = writer_fill_bookmark(
                    record["document_id"], "Client", "new", "replaced-file", workspace=tmp,
                )
            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "DOCUMENT_IDENTITY_CHANGED")
            com.assert_not_called()
            _, rebound, _ = register_document("writer", str(path), tmp)
            self.assertEqual(rebound["document_id"], record["document_id"])
            ok, _, errors, _ = writer_fill_bookmark(
                rebound["document_id"], "Client", "new", "rebound-preview", dry_run=True, workspace=tmp,
            )
            self.assertTrue(ok, errors)
            ok, _, errors, _ = writer_fill_bookmark(record["document_id"], "_GoBack", "x", "bookmark-reserved", dry_run=True, workspace=tmp)
            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "INVALID_ARGUMENT")

    def test_bookmark_fill_commit_readback_and_idempotency(self):
        from zipfile import ZipFile

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bookmark.docx"
            write_document(path, '<w:p><w:r><w:t>Dear </w:t></w:r><w:bookmarkStart w:id="1" w:name="Client"/><w:r><w:t>old</w:t></w:r><w:bookmarkEnd w:id="1"/><w:r><w:t>!</w:t></w:r></w:p>')
            _, record, _ = register_document("writer", str(path), tmp)

            def replace_xml(*args):
                with ZipFile(path) as archive:
                    parts = {name: archive.read(name) for name in archive.namelist()}
                parts["word/document.xml"] = parts["word/document.xml"].replace(b">old<", b">Northwind<")
                with ZipFile(path, "w") as archive:
                    for name, content in parts.items():
                        archive.writestr(name, content)
                return {"ok": True, "errors": [], "data": {"backend": "test"}}

            with patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com", side_effect=replace_xml):
                ok, result, errors, replayed = writer_fill_bookmark(record["document_id"], "Client", "Northwind", "bookmark-commit", workspace=tmp)
                self.assertTrue(ok, errors)
                self.assertFalse(replayed)
                self.assertTrue(result["readback_passed"])
                ok, repeated, errors, replayed = writer_fill_bookmark(record["document_id"], "Client", "Northwind", "bookmark-commit", workspace=tmp)
                self.assertTrue(ok)
                self.assertTrue(replayed)
                self.assertEqual(errors, [])
                self.assertTrue(repeated["readback_passed"])
                ok, _result, errors, replayed = writer_fill_bookmark(record["document_id"], "Client", "different", "bookmark-commit", workspace=tmp)
            self.assertFalse(ok)
            self.assertFalse(replayed)
            self.assertEqual(errors[0]["code"], "IDEMPOTENCY_CONFLICT")

    def test_bookmark_fill_allows_only_terminal_table_empty_paragraph_normalization(self):
        from zipfile import ZipFile

        for has_terminal_table, corrupt_table, expected_ok in (
            (True, False, True), (False, False, False), (True, True, False),
        ):
            with self.subTest(has_terminal_table=has_terminal_table, corrupt_table=corrupt_table), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / "bookmark.docx"
                table = '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Cell</w:t></w:r></w:p></w:tc></w:tr></w:tbl>' if has_terminal_table else ''
                write_document(path, '<w:p><w:bookmarkStart w:id="1" w:name="Target"/>'
                                     '<w:r><w:t>old</w:t></w:r><w:bookmarkEnd w:id="1"/></w:p>' + table)
                _, record, _ = register_document("writer", str(path), tmp)

                def replace_xml(*_args):
                    with ZipFile(path) as archive:
                        parts = {name: archive.read(name) for name in archive.namelist()}
                    payload = parts["word/document.xml"].replace(b">old<", b">new<")
                    if corrupt_table:
                        payload = payload.replace(b">Cell<", b">Other<")
                    parts["word/document.xml"] = payload.replace(b"</w:body>", b"<w:p/></w:body>")
                    with ZipFile(path, "w") as archive:
                        for name, content in parts.items():
                            archive.writestr(name, content)
                    return {"ok": True, "errors": [], "data": {"backend": "test"}}

                with patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com", side_effect=replace_xml):
                    ok, result, errors, _ = writer_fill_bookmark(
                        record["document_id"], "Target", "new", "table-tail", workspace=tmp,
                    )
                self.assertEqual(ok, expected_ok)
                self.assertEqual(result["body_tail_normalized"], has_terminal_table)
                self.assertEqual(result["readback_passed"], expected_ok)
                if not expected_ok:
                    self.assertEqual(errors[0]["code"], "VALIDATION_FAILED")

    def test_bookmark_fill_rejects_duplicate_and_cross_paragraph_targets_before_backup(self):
        cases = [
            ('<w:p><w:bookmarkStart w:id="1" w:name="Target"/><w:r><w:t>A</w:t></w:r><w:bookmarkEnd w:id="1"/><w:bookmarkStart w:id="2" w:name="Target"/><w:r><w:t>B</w:t></w:r><w:bookmarkEnd w:id="2"/></w:p>', "BOOKMARK_AMBIGUOUS"),
            ('<w:p><w:bookmarkStart w:id="1" w:name="Target"/><w:r><w:t>A</w:t></w:r></w:p><w:p><w:r><w:t>B</w:t></w:r><w:bookmarkEnd w:id="1"/></w:p>', "BOOKMARK_SCOPE_UNSUPPORTED"),
            ('<w:tbl><w:tr><w:tc><w:p><w:bookmarkStart w:id="1" w:name="Target"/><w:r><w:t>A</w:t></w:r><w:bookmarkEnd w:id="1"/></w:p></w:tc></w:tr></w:tbl>', "BOOKMARK_SCOPE_UNSUPPORTED"),
        ]
        for body, expected_code in cases:
            with self.subTest(code=expected_code), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / "bookmark.docx"
                write_document(path, body)
                _, record, _ = register_document("writer", str(path), tmp)
                with patch("wps_ai_agent_cli.writer_ops.create_backup") as backup, patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com") as com:
                    ok, _, errors, _ = writer_fill_bookmark(record["document_id"], "Target", "X", "bookmark-rejected", workspace=tmp)
                self.assertFalse(ok)
                self.assertEqual(errors[0]["code"], expected_code)
                backup.assert_not_called()
                com.assert_not_called()

    def test_readback_rejects_unexpected_change_outside_selected_paragraph(self):
        from tests.test_writer_ops import write_minimal_docx_paragraphs
        from zipfile import ZipFile

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "scope.docx"
            write_minimal_docx_paragraphs(path, ["needle first", "needle target", "needle last"])
            _, record, _ = register_document("writer", str(path), tmp)

            def corrupt_neighbor(*_args, **_kwargs):
                with ZipFile(path) as archive:
                    parts = {name: archive.read(name) for name in archive.namelist()}
                parts["word/document.xml"] = parts["word/document.xml"].replace(b"needle first", b"CORRUPTED first").replace(b"needle target", b"replacement target")
                with ZipFile(path, "w") as archive:
                    for name, content in parts.items():
                        archive.writestr(name, content)
                return {"ok": True, "errors": [], "data": {"backend": "test", "replace_count": 1}}

            with patch("wps_ai_agent_cli.writer_ops._run_writer_replace_com", side_effect=corrupt_neighbor):
                ok, result, errors, _ = writer_replace(
                    record["document_id"], "needle", "replacement", "unexpected-neighbor",
                    paragraph_index=2, workspace=tmp,
                )
            self.assertFalse(ok)
            self.assertFalse(result["readback_matches_expected"])
            self.assertEqual(result["changed_paragraph_indices"], [1])
            self.assertEqual(errors[0]["code"], "VALIDATION_FAILED")


@unittest.skipUnless(os.environ.get("WPS_AGENT_RUN_INTEGRATION") == "1", "Requires explicit local WPS integration run")
class WriterReplaceScopeIntegrationTests(unittest.TestCase):
    def test_scoped_replacements_preserve_neighboring_paragraphs_and_table(self):
        from docx import Document

        for replacement in ("x", "a much longer replacement", "", "needle plus", "needle"):
            with self.subTest(replacement=replacement), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / "scope.docx"
                document = Document()
                document.add_paragraph("needle before")
                document.add_table(rows=1, cols=1).cell(0, 0).text = "needle table"
                document.add_paragraph("needle needle NEEDLE")
                document.add_paragraph("needle after")
                document.save(path)
                _, record, _ = register_document("writer", str(path), tmp)
                ok, result, errors, _ = writer_replace(
                    record["document_id"], "needle", replacement, "scope-integration",
                    paragraph_index=2, workspace=tmp,
                )
                self.assertTrue(ok, (errors, result))
                self.assertEqual(result["backend_replace_count"], 2)
                self.assertEqual(result["expected_replace_count"], 2)
                self.assertTrue(result["readback_matches_expected"])
                self.assertEqual(docx_body_paragraphs(path), [
                    "needle before", f"{replacement} {replacement} NEEDLE", "needle after",
                ])
                self.assertEqual(docx_body_tables(path), [[["needle table"]]])
                self.assertTrue(result["backup"])

    def test_document_scope_readback_matches_body_and_table_text(self):
        from docx import Document
        from wps_ai_agent_cli.document_text import docx_body_story_paragraphs

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "whole-body.docx"
            document = Document()
            document.add_paragraph("needle first")
            document.add_table(rows=1, cols=1).cell(0, 0).text = "needle cell"
            document.add_paragraph("needle needle last")
            document.save(path)
            _, record, _ = register_document("writer", str(path), tmp)
            ok, result, errors, _ = writer_replace(
                record["document_id"], "needle", "x", "whole-body",
                workspace=tmp,
            )
            self.assertTrue(ok, (errors, result))
            self.assertEqual(result["expected_replace_count"], 4)
            self.assertTrue(result["readback_matches_expected"])
            self.assertEqual(docx_body_story_paragraphs(path), ["x first", "x cell", "x x last"])

    def test_bookmark_fill_preserves_range_and_reads_it_back(self):
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from wps_ai_agent_cli.writer_structure import read_body_bookmark_text

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bookmark.docx"
            document = Document()
            paragraph = document.add_paragraph("Dear ")
            run = paragraph.add_run("old")
            run.bold = True
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), "1")
            start.set(qn("w:name"), "Client")
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), "1")
            run._r.addprevious(start)
            run._r.addnext(end)
            paragraph.add_run("!")
            document.save(path)
            _, record, _ = register_document("writer", str(path), tmp)
            ok, result, errors, _ = writer_fill_bookmark(
                record["document_id"], "Client", "Northwind", "bookmark-wps", workspace=tmp,
            )
            self.assertTrue(ok, (errors, result))
            self.assertTrue(result["readback_passed"])
            bookmark, text = read_body_bookmark_text(path, "Client")
            self.assertTrue(bookmark["body_paragraph_range_supported"])
            self.assertEqual(text, "Northwind")
            reopened = Document(path)
            self.assertEqual(reopened.paragraphs[0].text, "Dear Northwind!")
            self.assertTrue(reopened.paragraphs[0].runs[1].bold)

    def test_paragraph_child_bookmark_fill_preserves_nested_neighbors(self):
        from wps_ai_agent_cli.writer_structure import read_body_bookmark_text, read_supported_bookmark_text

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "nested-copy.docx"
            shutil.copy2("fixtures/phase3/writer_nested_scope_wps_fixture.docx", path)
            tables_before = docx_body_tables(path)
            _, record, _ = register_document("writer", str(path), tmp)
            ok, result, errors, replayed = writer_fill_bookmark(
                record["document_id"], "BodyMark", "Updated", "paragraph-child-wps", workspace=tmp,
            )
            self.assertTrue(ok, (errors, result))
            self.assertFalse(replayed)
            self.assertTrue(result["readback_passed"])
            self.assertTrue(result["body_tail_normalized"])
            self.assertTrue(Path(result["backup"]["backup_path"]).is_file())
            self.assertEqual(read_body_bookmark_text(path, "BodyMark")[1], "Updated")
            self.assertEqual(docx_body_tables(path), tables_before)
            for name, expected in (("CellMark", "CellValue"), ("HeaderMark", "HeaderValue"),
                                   ("FooterMark", "FooterValue")):
                self.assertEqual(read_supported_bookmark_text(path, name)[1], expected)
            ok, repeated, errors, replayed = writer_fill_bookmark(
                record["document_id"], "BodyMark", "Updated", "paragraph-child-wps", workspace=tmp,
            )
            self.assertTrue(ok, errors)
            self.assertTrue(replayed)
            self.assertEqual(repeated["readback_passed"], True)

if __name__ == "__main__":
    unittest.main()
