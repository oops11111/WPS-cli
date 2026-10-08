from __future__ import annotations

import hashlib
import base64
import html
import json
import os
from pathlib import Path
import re
import tempfile
from typing import Any
import zipfile
import xml.etree.ElementTree as ET

from .html_editable import _TreeParser, _bookmark_name, _safe_link_target, _text, _walk, convert_html_editable


SCHEMA_VERSION = "wps-agent-html/v1"
_ID_PATTERN = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._:-]{0,127}$")
_OBJECT_TYPES = {
    "p": "paragraph", "h1": "heading", "h2": "heading", "h3": "heading",
    "h4": "heading", "h5": "heading", "h6": "heading", "li": "list_item",
    "table": "table", "tr": "table_row", "td": "table_cell", "th": "table_header",
    "a": "hyperlink", "img": "image",
}


def build_html_roundtrip_mapping(input_path: str | Path) -> tuple[bool, dict[str, Any], list[dict[str, Any]]]:
    source = Path(input_path).expanduser().resolve()
    if source.suffix.casefold() not in {".html", ".htm"}:
        return False, {}, [{"code": "INVALID_INPUT", "message": "Input must be an .html or .htm file."}]
    if not source.is_file():
        return False, {}, [{"code": "INPUT_FILE_NOT_FOUND", "message": f"HTML input not found: {source}"}]
    raw = source.read_bytes()
    if len(raw) > 10 * 1024 * 1024:
        return False, {}, [{"code": "INPUT_TOO_LARGE", "message": "HTML input exceeds the 10 MiB limit."}]
    parser = _TreeParser()
    parser.feed(raw.decode("utf-8-sig", errors="replace"))
    nodes = list(_walk(parser.root))
    html_root = next((node for node in nodes if node.tag == "html"), None)
    meta = next((node for node in nodes if node.tag == "meta" and node.attrs.get("name", "").casefold() == "wps-agent-schema"), None)
    version = (html_root.attrs.get("data-wps-schema") if html_root else None) or (meta.attrs.get("content") if meta else None)
    if version != SCHEMA_VERSION:
        return False, {}, [{"code": "ROUNDTRIP_SCHEMA_REQUIRED", "message": f"Controlled round-trip requires schema {SCHEMA_VERSION}."}]

    errors: list[dict[str, Any]] = []
    seen: set[str] = set()
    native_counts: dict[str, int] = {}
    mappings: list[dict[str, Any]] = []
    ids_by_node: dict[int, str] = {}
    parents = {id(child): parent for parent in nodes for child in parent.children if not isinstance(child, str)}
    for node in nodes:
        object_id = node.attrs.get("data-wps-object-id", "").strip()
        if object_id:
            if not _ID_PATTERN.fullmatch(object_id):
                errors.append({"code": "INVALID_OBJECT_ID", "message": f"Invalid object ID on <{node.tag}>."})
            elif object_id in seen:
                errors.append({"code": "DUPLICATE_OBJECT_ID", "message": f"Duplicate object ID: {object_id}"})
            else:
                seen.add(object_id)
                ids_by_node[id(node)] = object_id
            if node.tag not in _OBJECT_TYPES:
                errors.append({"code": "UNSUPPORTED_OBJECT_TYPE", "message": f"Object ID is attached to unsupported <{node.tag}>."})
        elif node.tag in _OBJECT_TYPES:
            errors.append({"code": "OBJECT_ID_REQUIRED", "message": f"<{node.tag}> is missing data-wps-object-id."})

    if parser.unsupported:
        errors.append({"code": "UNSUPPORTED_ROUNDTRIP_CSS", "message": "Embedded CSS is not supported by the v1 round-trip contract."})
    if parser.warnings:
        errors.append({"code": "UNSUPPORTED_ROUNDTRIP_FEATURE", "message": "Executable content is not supported by the v1 round-trip contract."})
    if parser.unsupported_elements:
        errors.append({"code": "UNSUPPORTED_ROUNDTRIP_FEATURE", "message": "Embedded active or non-semantic content is not supported by the v1 round-trip contract: " + ", ".join(sorted(parser.unsupported_elements))})

    for node in nodes:
        if node.tag in {"iframe", "object", "canvas", "svg"}:
            errors.append({"code": "UNSUPPORTED_ROUNDTRIP_FEATURE", "message": f"<{node.tag}> is not supported by controlled round-trip."})
        if node.attrs.get("style") or node.attrs.get("class") or (node.tag == "link" and node.attrs.get("rel", "").casefold() == "stylesheet"):
            errors.append({"code": "UNSUPPORTED_ROUNDTRIP_CSS", "message": f"CSS on <{node.tag}> is not supported by the v1 round-trip contract."})
        object_id = ids_by_node.get(id(node))
        if not object_id or node.tag not in _OBJECT_TYPES:
            continue
        object_type = _OBJECT_TYPES[node.tag]
        native_counts[object_type] = native_counts.get(object_type, 0) + 1
        ancestor = parents.get(id(node))
        while ancestor is not None and id(ancestor) not in ids_by_node:
            ancestor = parents.get(id(ancestor))
        parent = ancestor
        mappings.append({
            "object_id": object_id,
            "bookmark_name": _bookmark_name(object_id),
            "html_tag": node.tag,
            "object_type": object_type,
            "native_index": native_counts[object_type],
            "parent_object_id": ids_by_node.get(id(parent)) if parent else None,
            "text_sha256": hashlib.sha256(_text(node).encode("utf-8")).hexdigest(),
            **({"href": node.attrs.get("href", "")} if node.tag == "a" else {}),
        })
    if not mappings:
        errors.append({"code": "NO_MAPPED_OBJECTS", "message": "No ID-bearing supported objects were found."})
    if errors:
        return False, {"schema_version": version, "mapping_count": len(mappings), "errors": errors}, errors
    return True, {
        "input_path": str(source), "schema_version": SCHEMA_VERSION,
        "source_sha256": hashlib.sha256(raw).hexdigest(), "read_only": True,
        "mapping_count": len(mappings), "mappings": mappings,
        "contract": {
            "object_id_attribute": "data-wps-object-id",
            "stable_id_pattern": _ID_PATTERN.pattern,
            "css_policy": "CSS is rejected in schema v1 to avoid untracked layout edits.",
            "ambiguity_policy": "Missing/duplicate IDs and unsupported ID-bearing elements are rejected.",
        },
    }, []


def import_controlled_html(input_path: str | Path, output_path: str | Path) -> tuple[bool, dict[str, Any], list[dict[str, Any]]]:
    source = Path(input_path).expanduser().resolve()
    destination = Path(output_path).expanduser().resolve()
    valid, mapping, errors = build_html_roundtrip_mapping(source)
    if not valid:
        return False, {"roundtrip_mapping": mapping}, errors
    mapping_path = Path(str(destination) + ".wpsmap.json")
    if mapping_path.exists():
        return False, {}, [{"code": "MAPPING_ALREADY_EXISTS", "message": f"Refusing to overwrite mapping: {mapping_path}"}]
    converted, document, conversion_errors = convert_html_editable(source, destination)
    if not converted:
        return False, {}, conversion_errors
    document_identity = destination.stat()
    try:
        with zipfile.ZipFile(destination) as archive:
            xml_root = ET.fromstring(archive.read("word/document.xml"))
        bookmark_names = [element.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}name") for element in xml_root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bookmarkStart")]
        expected = [item["bookmark_name"] for item in mapping["mappings"]]
        if sorted(bookmark_names.count(name) for name in expected) != [1] * len(expected):
            raise ValueError("Generated DOCX bookmark identity does not match owned HTML IDs.")
        mapping["document_path"] = str(destination)
        mapping["document_sha256_at_import"] = document["sha256"]
        mapping["mapping_path"] = str(mapping_path)
        handle = tempfile.NamedTemporaryFile(prefix=f".{destination.name}.", suffix=".json", dir=destination.parent, delete=False)
        temporary = Path(handle.name)
        try:
            handle.write(json.dumps(mapping, ensure_ascii=False, sort_keys=True, indent=2).encode("utf-8"))
        finally:
            handle.close()
        os.link(temporary, mapping_path)
        temporary.unlink()
    except Exception as exc:  # noqa: BLE001
        if "temporary" in locals():
            temporary.unlink(missing_ok=True)
        try:
            current = destination.stat()
            if (current.st_dev, current.st_ino) == (document_identity.st_dev, document_identity.st_ino):
                destination.unlink()
        except OSError:
            pass
        return False, {}, [{"code": "ROUNDTRIP_IMPORT_FAILED", "message": str(exc)[:500]}]
    return True, {
        "document_path": str(destination), "mapping_path": str(mapping_path),
        "object_count": len(mapping["mappings"]), "document_sha256": document["sha256"],
        "editable": True, "bookmark_count": len(expected),
    }, []


def verify_controlled_document(document_path: str | Path, mapping_path: str | Path) -> tuple[bool, dict[str, Any], list[dict[str, Any]]]:
    document = Path(document_path).expanduser().resolve()
    mapping_file = Path(mapping_path).expanduser().resolve()
    if not document.is_file() or not mapping_file.is_file():
        return False, {}, [{"code": "ROUNDTRIP_FILE_NOT_FOUND", "message": "DOCX and mapping sidecar must both exist."}]
    try:
        mapping = json.loads(mapping_file.read_text(encoding="utf-8"))
        if not isinstance(mapping, dict) or mapping.get("schema_version") != SCHEMA_VERSION or not isinstance(mapping.get("mappings"), list):
            raise ValueError("Mapping sidecar has an unsupported schema or shape.")
        if not all(isinstance(item, dict) and isinstance(item.get("bookmark_name"), str) and isinstance(item.get("object_type"), str) for item in mapping["mappings"]):
            raise ValueError("Mapping sidecar contains an invalid object entry.")
        with zipfile.ZipFile(document) as archive:
            root = ET.fromstring(archive.read("word/document.xml"))
    except (OSError, ValueError, KeyError, TypeError, AttributeError, zipfile.BadZipFile, ET.ParseError) as exc:
        return False, {}, [{"code": "ROUNDTRIP_VERIFY_INPUT_INVALID", "message": str(exc)[:500]}]
    w_name = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}name"
    actual = [node.get(w_name) for node in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bookmarkStart")]
    expected = [item.get("bookmark_name") for item in mapping["mappings"]]
    errors: list[dict[str, Any]] = []
    duplicate = sorted({name for name in actual if name and name.startswith("wps_") and actual.count(name) > 1})
    missing = [name for name in expected if actual.count(name) == 0]
    repeated = [name for name in expected if actual.count(name) > 1]
    unexpected = sorted({name for name in actual if name and name.startswith("wps_") and name not in set(expected)})
    if missing:
        errors.append({"code": "ROUNDTRIP_OBJECT_MISSING", "message": "One or more mapped object bookmarks are missing.", "details": missing})
    if repeated or duplicate:
        errors.append({"code": "ROUNDTRIP_OBJECT_DUPLICATED", "message": "One or more mapped object bookmarks are duplicated.", "details": sorted(set(repeated + duplicate))})
    if unexpected:
        errors.append({"code": "ROUNDTRIP_OBJECT_UNMAPPED", "message": "DOCX contains controlled bookmarks that are absent from the mapping sidecar.", "details": unexpected})
    expected_by_type: dict[str, list[str]] = {}
    for item in mapping["mappings"]:
        expected_by_type.setdefault(item["object_type"], []).append(item["bookmark_name"])
    actual_by_type = {
        kind: [name for name in actual if name in names]
        for kind, names in expected_by_type.items()
    }
    reordered = [kind for kind, names in expected_by_type.items() if actual_by_type[kind] != names]
    if reordered:
        errors.append({"code": "ROUNDTRIP_OBJECT_REORDERED", "message": "Mapped objects changed order within a native object type.", "details": reordered})
    current_hash = hashlib.sha256(document.read_bytes()).hexdigest()
    result = {
        "document_path": str(document), "mapping_path": str(mapping_file),
        "schema_version": mapping["schema_version"], "object_count": len(expected),
        "bookmark_count": sum(1 for name in actual if name in set(expected)),
        "document_edited_since_import": current_hash != mapping.get("document_sha256_at_import"),
        "identity_status": "passed" if not errors else "failed",
        "missing_count": len(missing), "duplicate_count": len(set(repeated + duplicate)),
        "unmapped_count": len(unexpected), "reordered_types": reordered,
    }
    return not errors, result, errors


def export_controlled_html(document_path: str | Path, mapping_path: str | Path, output_path: str | Path) -> tuple[bool, dict[str, Any], list[dict[str, Any]]]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Path(document_path).expanduser().resolve()
    mapping_file = Path(mapping_path).expanduser().resolve()
    output = Path(output_path).expanduser().resolve()
    if output.suffix.casefold() not in {".html", ".htm"}:
        return False, {}, [{"code": "INVALID_OUTPUT", "message": "Output must use .html or .htm."}]
    if output.exists():
        return False, {}, [{"code": "OUTPUT_ALREADY_EXISTS", "message": f"Refusing to overwrite existing output: {output}"}]
    if not output.parent.is_dir():
        return False, {}, [{"code": "OUTPUT_DIRECTORY_NOT_FOUND", "message": f"Output directory does not exist: {output.parent}"}]
    verified, identity, errors = verify_controlled_document(document, mapping_file)
    if not verified:
        return False, {"identity": identity}, errors
    try:
        mapping = json.loads(mapping_file.read_text(encoding="utf-8"))
        word = Document(document)
    except Exception as exc:  # noqa: BLE001
        return False, {}, [{"code": "ROUNDTRIP_EXPORT_INPUT_INVALID", "message": str(exc)[:500]}]
    by_bookmark = {entry["bookmark_name"]: entry for entry in mapping["mappings"]}
    by_type: dict[str, list[dict[str, Any]]] = {}
    for entry in mapping["mappings"]:
        by_type.setdefault(entry["object_type"], []).append(entry)
    cursors = {kind: 0 for kind in by_type}
    used: set[str] = set()
    deferred_images: dict[str, list[str]] = {}
    by_object_id = {entry["object_id"]: entry for entry in mapping["mappings"]}
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

    def bookmarks(element) -> list[dict[str, Any]]:
        names = [node.get(f"{{{ns}}}name") for node in element.iter(f"{{{ns}}}bookmarkStart")]
        return [by_bookmark[name] for name in names if name in by_bookmark]

    def take(kind: str, *, element=None) -> dict[str, Any]:
        candidates = bookmarks(element) if element is not None else []
        match = next((entry for entry in candidates if entry["object_type"] == kind), None)
        if match is None:
            raise ValueError(f"A native {kind} object has no matching stable bookmark.")
        if match["object_id"] in used:
            raise ValueError(f"Stable object ID is mapped more than once: {match['object_id']}")
        used.add(match["object_id"])
        return match

    def next_inline(kind: str) -> dict[str, Any]:
        index = cursors.get(kind, 0)
        entries = by_type.get(kind, [])
        if index >= len(entries):
            raise ValueError(f"DOCX contains an unmapped inline {kind} object.")
        cursors[kind] = index + 1
        entry = entries[index]
        if entry["object_id"] in used:
            raise ValueError(f"Stable object ID is mapped more than once: {entry['object_id']}")
        used.add(entry["object_id"])
        return entry

    def text_content(element) -> str:
        parts: list[str] = []
        for child in element.iter():
            if child.tag == f"{{{ns}}}t":
                parts.append(child.text or "")
            elif child.tag == f"{{{ns}}}tab":
                parts.append("\t")
            elif child.tag in {f"{{{ns}}}br", f"{{{ns}}}cr"}:
                parts.append("\n")
        return "".join(parts)

    def export_inlines(paragraph_element) -> str:
        pieces: list[str] = []
        children = list(paragraph_element)
        index = 0
        while index < len(children):
            child = children[index]
            if child.tag in {f"{{{ns}}}pPr", f"{{{ns}}}bookmarkStart", f"{{{ns}}}bookmarkEnd"}:
                index += 1
                continue
            if child.tag == f"{{{ns}}}hyperlink":
                entry = next_inline("hyperlink")
                relationship_id = child.get(f"{{{r_ns}}}id")
                relation = word.part.rels.get(relationship_id) if relationship_id else None
                href = getattr(relation, "target_ref", "")
                if not href:
                    raise ValueError("A mapped hyperlink has no supported target relationship.")
                pieces.append(f'<a data-wps-object-id="{html.escape(entry["object_id"], quote=True)}" href="{html.escape(href, quote=True)}">{html.escape(text_content(child))}</a>')
                index += 1
            elif child.tag == f"{{{ns}}}r":
                field_start = child.find(f"{{{ns}}}fldChar")
                if field_start is not None and field_start.get(f"{{{ns}}}fldCharType") == "begin":
                    instruction = ""
                    visible: list[str] = []
                    separated = False
                    ended = False
                    cursor = index
                    while cursor < len(children):
                        field_child = children[cursor]
                        if field_child.tag == f"{{{ns}}}r":
                            for part in field_child:
                                if part.tag == f"{{{ns}}}instrText":
                                    instruction += part.text or ""
                                elif part.tag == f"{{{ns}}}fldChar":
                                    field_type = part.get(f"{{{ns}}}fldCharType")
                                    separated = separated or field_type == "separate"
                                    if field_type == "end":
                                        ended = True
                                elif separated and part.tag == f"{{{ns}}}t":
                                    visible.append(part.text or "")
                        cursor += 1
                        if ended:
                            break
                    match = re.search(r'^\s*HYPERLINK\s+"([^"\r\n]+)"', instruction, re.IGNORECASE)
                    if not ended or not match or not _safe_link_target(match.group(1)):
                        raise ValueError("Unsupported or unsafe Word field in a mapped paragraph.")
                    entry = next_inline("hyperlink")
                    href = match.group(1)
                    pieces.append(f'<a data-wps-object-id="{html.escape(entry["object_id"], quote=True)}" href="{html.escape(href, quote=True)}">{html.escape("".join(visible))}</a>')
                    index = cursor
                    continue
                drawing = child.find(f"{{{ns}}}drawing")
                if drawing is not None:
                    entry = next_inline("image")
                    blip = drawing.find(f".//{{http://schemas.openxmlformats.org/drawingml/2006/main}}blip")
                    relationship_id = blip.get(f"{{{r_ns}}}embed") if blip is not None else None
                    relation = word.part.related_parts.get(relationship_id) if relationship_id else None
                    if relation is None:
                        raise ValueError("A mapped image has no embedded image part.")
                    mime = relation.content_type
                    if mime not in {"image/png", "image/jpeg", "image/gif"}:
                        raise ValueError(f"Unsupported embedded image type: {mime}")
                    encoded = base64.b64encode(relation.blob).decode("ascii")
                    image_markup = f'<img data-wps-object-id="{html.escape(entry["object_id"], quote=True)}" alt="" src="data:{mime};base64,{encoded}">'
                    parent_id = entry.get("parent_object_id")
                    if parent_id:
                        parent = by_object_id.get(parent_id, {})
                        if parent.get("object_type") not in {"paragraph", "heading", "list_item", "table_cell", "table_header"}:
                            raise ValueError("An image's original parent cannot be represented in owned HTML v1.")
                        deferred_images.setdefault(parent_id, []).append(image_markup)
                    else:
                        pieces.append(image_markup)
                else:
                    pieces.append(html.escape(text_content(child)).replace("\n", "<br>"))
                index += 1
            elif child.tag in {f"{{{ns}}}bookmarkStart", f"{{{ns}}}bookmarkEnd"}:
                index += 1
                continue
            else:
                raise ValueError(f"Unsupported WordprocessingML inline element: {child.tag.rsplit('}', 1)[-1]}")
        return "".join(pieces)

    def paragraph_html(element) -> tuple[str | None, str | None]:
        entry = next((item for item in bookmarks(element) if item["object_type"] in {"paragraph", "heading", "list_item"}), None)
        inline = export_inlines(element)
        image_only = any(item["object_type"] == "image" for item in bookmarks(element))
        if entry is None:
            if image_only:
                return None, inline
            if text_content(element).strip():
                raise ValueError("DOCX contains an untracked body paragraph; refusing lossy export.")
            return None, None
        used.add(entry["object_id"])
        tag = entry["html_tag"]
        if tag == "li":
            style = Paragraph(element, word).style.name
            list_tag = "ol" if "Number" in style else "ul"
            return f'<{list_tag}><li data-wps-object-id="{html.escape(entry["object_id"], quote=True)}">{inline}</li></{list_tag}>', None
        return f'<{tag} data-wps-object-id="{html.escape(entry["object_id"], quote=True)}">{inline}</{tag}>', None

    body = word._element.body
    blocks: list[str] = []
    try:
        for element in body:
            if element.tag == f"{{{ns}}}p":
                markup, image_markup = paragraph_html(element)
                if markup:
                    blocks.append(markup)
                if image_markup:
                    blocks.append(f"<p>{image_markup}</p>")
            elif element.tag == f"{{{ns}}}tbl":
                table = Table(element, word)
                table_entry = take("table", element=element)
                rows_markup = []
                for row in table.rows:
                    row_entry = take("table_row", element=row._tr)
                    cells_markup = []
                    for cell in row.cells:
                        cell_type = "table_header" if any(entry["object_type"] == "table_header" for entry in bookmarks(cell._tc)) else "table_cell"
                        cell_entry = take(cell_type, element=cell._tc)
                        if len(cell.paragraphs) != 1:
                            raise ValueError("Multi-paragraph table cells are not supported by owned schema v1 export.")
                        content = export_inlines(cell.paragraphs[0]._p)
                        tag = "th" if cell_type == "table_header" else "td"
                        cells_markup.append(f'<{tag} data-wps-object-id="{html.escape(cell_entry["object_id"], quote=True)}">{content}</{tag}>')
                    rows_markup.append(f'<tr data-wps-object-id="{html.escape(row_entry["object_id"], quote=True)}">{"".join(cells_markup)}</tr>')
                blocks.append(f'<table data-wps-object-id="{html.escape(table_entry["object_id"], quote=True)}">{"".join(rows_markup)}</table>')
            elif element.tag != f"{{{ns}}}sectPr":
                raise ValueError(f"Unsupported WordprocessingML body element: {element.tag.rsplit('}', 1)[-1]}")
        expected_ids = {entry["object_id"] for entry in mapping["mappings"]}
        if used != expected_ids:
            missing_ids = sorted(expected_ids - used)
            unexpected_ids = sorted(used - expected_ids)
            raise ValueError("Some mapped identities could not be represented in the exported HTML. "
                             f"Missing: {missing_ids}; unexpected: {unexpected_ids}")
        for parent_id, images in deferred_images.items():
            parent = by_object_id[parent_id]
            tag = "th" if parent["object_type"] == "table_header" else "td" if parent["object_type"] == "table_cell" else parent["html_tag"]
            marker = f'</{tag}>'
            target_marker = f'data-wps-object-id="{html.escape(parent_id, quote=True)}"'
            found = False
            for index, block in enumerate(blocks):
                marker_start = block.find(target_marker)
                marker_end = block.find(marker, marker_start)
                if marker_start >= 0 and marker_end >= 0:
                    blocks[index] = block[:marker_end] + "".join(images) + block[marker_end:]
                    found = True
                    break
            if not found:
                raise ValueError(f"Original image parent is absent from exported HTML: {parent_id}")
        title = html.escape(word.core_properties.title or "Controlled document")
        output_html = (
            '<!doctype html><html data-wps-schema="wps-agent-html/v1"><head>'
            f"<meta charset=\"utf-8\"><title>{title}</title></head><body>"
            + "".join(blocks) + "</body></html>"
        )
        output_bytes = output_html.encode("utf-8")
        if len(output_bytes) > 70 * 1024 * 1024:
            raise ValueError("Exported HTML exceeds the 70 MiB output limit.")
        with tempfile.NamedTemporaryFile(prefix=f".{output.stem}.", suffix=output.suffix, dir=output.parent, delete=False) as handle:
            temporary = Path(handle.name)
            handle.write(output_bytes)
        os.link(temporary, output)
        temporary.unlink()
    except Exception as exc:  # noqa: BLE001
        if "temporary" in locals():
            temporary.unlink(missing_ok=True)
        return False, {}, [{"code": "ROUNDTRIP_EXPORT_UNSUPPORTED", "message": str(exc)[:500]}]
    payload = output.read_bytes()
    return True, {
        "document_path": str(document), "mapping_path": str(mapping_file),
        "output_path": str(output), "bytes": len(payload),
        "sha256": hashlib.sha256(payload).hexdigest(),
        "schema_version": SCHEMA_VERSION, "exported_object_count": len(used),
    }, []
