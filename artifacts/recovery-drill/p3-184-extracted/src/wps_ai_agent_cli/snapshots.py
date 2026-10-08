from __future__ import annotations

from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .document_text import docx_body_paragraphs
from .errors import INPUT_FILE_NOT_FOUND
from .presentation_text import PresentationStructureError, pptx_slide_texts, pptx_text_objects
from .sessions import get_document
from .writer_structure import read_writer_structure


SPREADSHEET_ERROR_VALUES = {
    "#DIV/0!",
    "#N/A",
    "#NAME?",
    "#NULL!",
    "#NUM!",
    "#REF!",
    "#SPILL!",
    "#VALUE!",
    "#CALC!",
}


def _preview(text: str, limit: int = 80) -> str:
    normalized = " ".join(text.split())
    if len(normalized) <= limit:
        return normalized
    return f"{normalized[: limit - 3]}..."


def _writer_snapshot(document_id: str, path: Path) -> dict[str, Any]:
    paragraphs = docx_body_paragraphs(path)
    structure = read_writer_structure(path)
    paragraph_items = [
        {
            **structure["paragraphs"][index - 1],
            "paragraph_index": index,
            "text_length": len(text),
            "is_empty": text.strip() == "",
            "preview": _preview(text),
        }
        for index, text in enumerate(paragraphs, start=1)
    ]
    return {
        "document_id": document_id,
        "component": "writer",
        "path": str(path),
        "paragraph_count": len(paragraphs),
        "non_empty_paragraph_count": sum(1 for text in paragraphs if text.strip()),
        "body_text_length": sum(len(text) for text in paragraphs),
        "paragraphs": paragraph_items,
        "heading_count": sum(item["heading_level"] is not None for item in paragraph_items),
        "headings": [item for item in paragraph_items if item["heading_level"] is not None],
        "bookmark_count": len(structure["bookmarks"]),
        "bookmarks": structure["bookmarks"],
        "structure_warnings": structure["warnings"],
        "structure_scope": structure["scope"],
    }


def _spreadsheet_snapshot(document_id: str, path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=False, read_only=True)
    try:
        sheets = []
        total_formula_count = 0
        total_error_count = 0
        for sheet in workbook.worksheets:
            formula_cells = []
            error_cells = []
            non_empty_count = 0
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if value is None:
                        continue
                    non_empty_count += 1
                    if cell.data_type == "f" or (isinstance(value, str) and value.startswith("=")):
                        formula_cells.append({"cell": cell.coordinate, "formula": value})
                    if cell.data_type == "e" or value in SPREADSHEET_ERROR_VALUES:
                        error_cells.append({"cell": cell.coordinate, "value": value})

            total_formula_count += len(formula_cells)
            total_error_count += len(error_cells)
            sheets.append(
                {
                    "name": sheet.title,
                    "max_row": sheet.max_row,
                    "max_column": sheet.max_column,
                    "non_empty_cell_count": non_empty_count,
                    "formula_count": len(formula_cells),
                    "formula_cells": formula_cells,
                    "formula_error_count": len(error_cells),
                    "formula_error_cells": error_cells,
                }
            )
    finally:
        workbook.close()

    return {
        "document_id": document_id,
        "component": "spreadsheets",
        "path": str(path),
        "sheet_count": len(sheets),
        "formula_count": total_formula_count,
        "formula_error_count": total_error_count,
        "sheets": sheets,
    }


def _presentation_snapshot(document_id: str, path: Path) -> dict[str, Any]:
    objects = pptx_text_objects(path)
    by_slide = Counter(int(item["slide_index"]) for item in objects)
    non_empty_by_slide = Counter(
        int(item["slide_index"]) for item in objects if not item["is_empty"]
    )
    slides = []
    for slide in pptx_slide_texts(path):
        slide_index = int(slide["slide_index"])
        slide_objects = [item for item in objects if item["slide_index"] == slide_index]
        slides.append(
            {
                "slide_index": slide_index,
                "text_object_count": by_slide[slide_index],
                "non_empty_text_object_count": non_empty_by_slide[slide_index],
                "text_length": sum(int(item["text_length"]) for item in slide_objects),
                "objects": [
                    {
                        "object_index": item["object_index"],
                        "object_type": item["object_type"],
                        "text_length": item["text_length"],
                        "is_empty": item["is_empty"],
                        "placeholder_type": item["placeholder_type"],
                        "preview": _preview(str(item["text"])),
                    }
                    for item in slide_objects
                ],
            }
        )

    return {
        "document_id": document_id,
        "component": "presentation",
        "path": str(path),
        "slide_count": len(slides),
        "text_object_count": len(objects),
        "non_empty_text_object_count": sum(1 for item in objects if not item["is_empty"]),
        "slides": slides,
    }


def snapshot_document(
    document_id: str,
    workspace: str | Path = ".",
) -> tuple[bool, dict[str, Any], list[dict[str, str]]]:
    document = get_document(document_id, workspace)
    if not document:
        return (
            False,
            {},
            [{"code": "DOCUMENT_NOT_FOUND", "message": f"Document not registered: {document_id}"}],
        )

    path = Path(document["path"])
    if not path.exists():
        return (
            False,
            {"document": document},
            [{"code": INPUT_FILE_NOT_FOUND, "message": f"Input file not found: {path}"}],
        )

    if document["component"] == "writer":
        return True, _writer_snapshot(document_id, path), []
    if document["component"] == "spreadsheets":
        return True, _spreadsheet_snapshot(document_id, path), []
    if document["component"] == "presentation":
        try:
            return True, _presentation_snapshot(document_id, path), []
        except PresentationStructureError as exc:
            return False, {"document": document}, [{"code": "INVALID_PRESENTATION_STRUCTURE", "message": str(exc)}]

    return (
        False,
        {"document": document},
        [{"code": "UNSUPPORTED_COMPONENT", "message": f"Snapshot is not implemented for {document['component']}."}],
    )
