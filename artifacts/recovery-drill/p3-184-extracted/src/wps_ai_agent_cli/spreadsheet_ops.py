from __future__ import annotations

import json
import posixpath
from pathlib import Path
import subprocess
import tempfile
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.styles.numbers import is_date_format

from .backups import create_backup, guarded_com_mutation, verify_post_com_source
from .capabilities import powershell_executable, probe_wps_capabilities
from .errors import COM_BACKEND_UNAVAILABLE, COM_OPERATION_FAILED, INPUT_FILE_NOT_FOUND
from .mutation_lock import coordinated_mutation
from .operations import record_operation, replay_operation
from .sessions import get_document
from .spreadsheet_ranges import validate_spreadsheet_read_range


def _format_category(cell: Any) -> str:
    if cell.value is None:
        return "empty"
    if cell.is_date or is_date_format(cell.number_format):
        return "date"
    number_format = cell.number_format or "General"
    if "%" in number_format:
        return "percentage"
    if any(symbol in number_format for symbol in ("$", "€", "£", "¥", "₹", "₩")):
        return "currency"
    if cell.data_type in {"s", "str", "inlineStr"} or isinstance(cell.value, str):
        return "text"
    return "number"


def read_spreadsheet_range(
    document_id: str,
    range_address: str,
    sheet_name: str | None = None,
    workspace: str | Path = ".",
) -> tuple[bool, dict[str, Any], list[dict[str, str]]]:
    document = get_document(document_id, workspace)
    if not document:
        return (
            False,
            {},
            [{"code": "DOCUMENT_NOT_FOUND", "message": f"Document not registered: {document_id}"}],
        )
    if document["component"] != "spreadsheets":
        return (
            False,
            {"document": document},
            [{"code": "UNSUPPORTED_COMPONENT", "message": "spreadsheet-read requires a spreadsheet document."}],
        )

    path = Path(document["path"])
    if not path.exists():
        return (
            False,
            {"document": document},
            [{"code": INPUT_FILE_NOT_FOUND, "message": f"Input file not found: {path}"}],
        )

    bounds, range_error = validate_spreadsheet_read_range(range_address)
    if range_error:
        return False, {}, [range_error]
    range_address = range_address.strip()

    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                return (
                    False,
                    {"document": document, "sheet_names": workbook.sheetnames},
                    [{"code": "SHEET_NOT_FOUND", "message": f"Sheet not found: {sheet_name}"}],
                )
            sheet = workbook[sheet_name]
        else:
            sheet = workbook.worksheets[0]
        rows = []
        cell_metadata = []
        min_col, min_row, _max_col, _max_row = bounds
        for row_index, row in enumerate(sheet[range_address]):
            rows.append([cell.value for cell in row])
            cell_metadata.append([
                {
                    "address": f"{get_column_letter(min_col + column_index)}{min_row + row_index}",
                    "value_type": cell.data_type,
                    "number_format": cell.number_format,
                    "is_date": bool(cell.is_date or is_date_format(cell.number_format)),
                    "format_category": _format_category(cell),
                }
                for column_index, cell in enumerate(row)
            ])
        return (
            True,
            {
                "document_id": document_id,
                "component": "spreadsheets",
                "path": str(path),
                "sheet": sheet.title,
                "range": range_address,
                "values": rows,
                "cell_metadata": cell_metadata,
                "row_count": len(rows),
                "column_count": len(rows[0]) if rows else 0,
            },
            [],
        )
    finally:
        workbook.close()


def list_spreadsheet_sheets(
    document_id: str,
    workspace: str | Path = ".",
) -> tuple[bool, dict[str, Any], list[dict[str, str]]]:
    document = get_document(document_id, workspace)
    if not document:
        return False, {}, [{"code": "DOCUMENT_NOT_FOUND", "message": f"Document not registered: {document_id}"}]
    if document["component"] != "spreadsheets":
        return False, {"document": document}, [{"code": "UNSUPPORTED_COMPONENT", "message": "spreadsheet-sheets requires a spreadsheet document."}]
    path = Path(document["path"])
    if not path.exists():
        return False, {"document": document}, [{"code": INPUT_FILE_NOT_FOUND, "message": f"Input file not found: {path}"}]

    workbook = load_workbook(path, data_only=False, read_only=False, keep_vba=path.suffix.casefold() == ".xlsm")
    try:
        ignored_errors = _read_ignored_errors(path)
        sheets = []
        for index, sheet in enumerate(workbook.worksheets, start=1):
            dimension = sheet.calculate_dimension()
            scan_area = (sheet.max_row or 1) * (sheet.max_column or 1)
            inventory_scan_truncated = scan_area > 100_000
            populated_cell_count = None if inventory_scan_truncated else 0
            formula_count = None if inventory_scan_truncated else 0
            protected_cell_count = None if inventory_scan_truncated and sheet.protection.sheet else 0
            if not inventory_scan_truncated:
                populated_cell_count = 0
                formula_count = 0
                protected_cell_count = 0
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            populated_cell_count += 1
                            if cell.data_type == "f":
                                formula_count += 1
                        if sheet.protection.sheet and cell.protection.locked:
                            protected_cell_count += 1
            sorted_merged_ranges = sorted(sheet.merged_cells.ranges, key=lambda item: (item.min_row, item.min_col, item.max_row, item.max_col))
            merged_range_count = len(sorted_merged_ranges)
            merged_ranges_truncated = merged_range_count > 1_000
            merged_ranges = [str(item) for item in sorted_merged_ranges[:1_000]]
            row_breaks = sorted({int(item.id) for item in sheet.row_breaks.brk})
            column_breaks = sorted({int(item.id) for item in sheet.col_breaks.brk})
            validations = sorted(
                sheet.data_validations.dataValidation,
                key=lambda item: (str(item.sqref), item.type or "", item.operator or ""),
            )
            validation_count = len(validations)
            validation_records = [{
                "ranges": str(item.sqref),
                "type": item.type,
                "operator": item.operator,
                "allow_blank": bool(item.allow_blank),
            } for item in validations[:1_000]]
            conditional_rules = []
            for conditional_formatting in sheet.conditional_formatting:
                for rule in conditional_formatting.rules:
                    conditional_rules.append({
                        "ranges": str(conditional_formatting.sqref),
                        "type": rule.type,
                        "operator": rule.operator,
                        "priority": rule.priority,
                    })
            conditional_rules.sort(key=lambda item: (item["ranges"], item["priority"] or 0, item["type"] or "", item["operator"] or ""))
            conditional_rule_count = len(conditional_rules)
            ignored_error = ignored_errors.get(sheet.title, {"records": [], "count": 0, "truncated": False})
            sheets.append({
                "sheet_index": index,
                "name": sheet.title,
                "visibility": sheet.sheet_state,
                "visible": sheet.sheet_state == "visible",
                "tab_color": _tab_color_from_cell(sheet.sheet_properties.tabColor),
                "protection_enabled": bool(sheet.protection.sheet),
                "protected_cell_count": protected_cell_count,
                "protection_scan_truncated": inventory_scan_truncated and bool(sheet.protection.sheet),
                "populated_cell_count": populated_cell_count,
                "formula_count": formula_count,
                "inventory_scan_truncated": inventory_scan_truncated,
                "freeze_panes": sheet.freeze_panes.coordinate if hasattr(sheet.freeze_panes, "coordinate") else sheet.freeze_panes,
                "autofilter_range": sheet.auto_filter.ref,
                "print_area": str(sheet.print_area) if sheet.print_area else None,
                "print_title_rows": sheet.print_title_rows,
                "print_title_columns": sheet.print_title_cols,
                "page_orientation": sheet.page_setup.orientation,
                "paper_size": str(sheet.page_setup.paperSize) if sheet.page_setup.paperSize is not None else None,
                "row_breaks": row_breaks[:1_000],
                "row_break_count": len(row_breaks),
                "row_breaks_truncated": len(row_breaks) > 1_000,
                "column_breaks": column_breaks[:1_000],
                "column_break_count": len(column_breaks),
                "column_breaks_truncated": len(column_breaks) > 1_000,
                "data_validations": validation_records,
                "data_validation_count": validation_count,
                "data_validations_truncated": validation_count > 1_000,
                "conditional_formatting": conditional_rules[:1_000],
                "conditional_formatting_rule_count": conditional_rule_count,
                "conditional_formatting_truncated": conditional_rule_count > 1_000,
                "ignored_errors": ignored_error["records"],
                "ignored_error_count": ignored_error["count"],
                "ignored_errors_truncated": ignored_error["truncated"],
                "merged_ranges": merged_ranges,
                "merged_range_count": merged_range_count,
                "merged_ranges_truncated": merged_ranges_truncated,
                "dimension": dimension,
                "max_row": sheet.max_row or 0,
                "max_column": sheet.max_column or 0,
            })
        defined_names = []
        for defined_name in workbook.defined_names.values():
            defined_names.append({
                "name": defined_name.name,
                "scope": "workbook",
                "target": defined_name.attr_text,
                "hidden": bool(defined_name.hidden),
            })
        for sheet in workbook.worksheets:
            for defined_name in sheet.defined_names.values():
                defined_names.append({
                    "name": defined_name.name,
                    "scope": sheet.title,
                    "target": defined_name.attr_text,
                    "hidden": bool(defined_name.hidden),
                })
        defined_names.sort(key=lambda item: (item["scope"].casefold(), item["name"].casefold(), item["target"] or ""))
        defined_name_count = len(defined_names)
        return True, {
            "document_id": document_id,
            "component": "spreadsheets",
            "path": str(path),
            "read_only": True,
            "saved": False,
            "sheet_count": len(sheets),
            "sheets": sheets,
            "defined_names": defined_names[:1_000],
            "defined_name_count": defined_name_count,
            "defined_names_truncated": defined_name_count > 1_000,
            "calculation": {
                "mode": workbook.calculation.calcMode,
                "full_calc_on_load": workbook.calculation.fullCalcOnLoad,
                "force_full_calc": workbook.calculation.forceFullCalc,
                "calc_on_save": workbook.calculation.calcOnSave,
                "iterate": workbook.calculation.iterate,
            },
        }, []
    finally:
        workbook.close()
        if workbook.vba_archive is not None:
            workbook.vba_archive.close()


def _validate_sheet_title(title: str) -> dict[str, str] | None:
    if not title or title.isspace():
        return {"code": "INVALID_SHEET_NAME", "message": "Worksheet name must not be empty."}
    if len(title) > 31:
        return {"code": "INVALID_SHEET_NAME", "message": "Worksheet names are limited to 31 characters."}
    if any(character in title for character in ":\\/?*[]") or any(ord(character) < 32 for character in title) or title.startswith("'") or title.endswith("'"):
        return {"code": "INVALID_SHEET_NAME", "message": "Worksheet name contains a forbidden character or boundary apostrophe."}
    if title.casefold() == "history":
        return {"code": "INVALID_SHEET_NAME", "message": "The reserved worksheet name 'History' is not allowed."}
    return None


def _run_spreadsheet_rename_sheet_com(path: str, old_name: str, new_name: str) -> dict[str, Any]:
    capabilities = probe_wps_capabilities()
    prog_id = capabilities["components"]["spreadsheets"]["selected_prog_id"]
    if not prog_id:
        return {"ok": False, "errors": [{"code": COM_BACKEND_UNAVAILABLE, "message": "No registered WPS ProgID detected for spreadsheets."}], "data": {"capabilities": capabilities}}
    params_json = json.dumps({"prog_id": prog_id, "path": str(Path(path).resolve()), "old_name": old_name, "new_name": new_name})
    script = f"""
$ErrorActionPreference = 'Stop'
$params = @'
{params_json}
'@ | ConvertFrom-Json
$app = $null; $workbook = $null
try {{
  $app = New-Object -ComObject $params.prog_id
  try {{ $app.Visible = $false }} catch {{ }}
  try {{ $app.DisplayAlerts = $false }} catch {{ }}
  $workbook = $app.Workbooks.Open($params.path)
  $sheet = $workbook.Worksheets.Item($params.old_name)
  $sheet.Name = $params.new_name
  $workbook.Save()
  $names = @()
  for ($index = 1; $index -le $workbook.Worksheets.Count; $index++) {{
    $names += [string]$workbook.Worksheets.Item($index).Name
  }}
  [pscustomobject]@{{ ok = $true; backend = 'powershell-com'; prog_id = $params.prog_id; sheet_names = $names }} | ConvertTo-Json -Depth 5 -Compress
}} catch {{
  [pscustomobject]@{{ ok = $false; backend = 'powershell-com'; error_type = $_.Exception.GetType().FullName; error_message = $_.Exception.Message; position = $_.InvocationInfo.PositionMessage }} | ConvertTo-Json -Depth 5 -Compress
  exit 2
}} finally {{
  if ($workbook -ne $null) {{ try {{ $workbook.Close($false) }} catch {{ }}; try {{ [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($workbook) }} catch {{ }} }}
  if ($app -ne $null) {{ try {{ $app.Quit() }} catch {{ }}; try {{ [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($app) }} catch {{ }} }}
  [GC]::Collect(); [GC]::WaitForPendingFinalizers()
}}
"""
    script_path = None
    try:
        with tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False, encoding="utf-8-sig") as stream:
            stream.write(script)
            script_path = stream.name
        completed = subprocess.run(
            [powershell_executable(), "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", script_path],
            check=False, capture_output=True, encoding="utf-8", errors="replace", text=True, timeout=120,
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        return {"ok": False, "errors": [{"code": COM_OPERATION_FAILED, "message": str(exc)}], "data": {}}
    finally:
        if script_path:
            try:
                Path(script_path).unlink(missing_ok=True)
            except OSError:
                pass
    try:
        payload = json.loads(completed.stdout)
    except json.JSONDecodeError as exc:
        return {"ok": False, "errors": [{"code": COM_OPERATION_FAILED, "message": f"WPS returned invalid JSON: {exc}"}], "data": {}}
    if completed.returncode != 0 or not payload.get("ok"):
        return {"ok": False, "errors": [{"code": COM_OPERATION_FAILED, "message": payload.get("error_message", "WPS worksheet rename failed.")}], "data": {"diagnostic": payload}}
    if isinstance(payload.get("sheet_names"), str):
        payload["sheet_names"] = [payload["sheet_names"]]
    return {"ok": True, "errors": [], "data": payload}


@coordinated_mutation
def rename_spreadsheet_sheet(
    document_id: str,
    old_name: str,
    new_name: str,
    request_id: str,
    dry_run: bool = False,
    workspace: str | Path = ".",
) -> tuple[bool, dict[str, Any], list[dict[str, str]], bool]:
    replay = replay_operation(
        request_id, "spreadsheet-rename-sheet",
        {"document_id": document_id, "old_name": old_name, "new_name": new_name, "dry_run": dry_run},
        workspace,
    )
    if replay is not None:
        return replay
    title_error = _validate_sheet_title(new_name)
    if title_error:
        return False, {}, [title_error], False
    document = get_document(document_id, workspace)
    if not document:
        return False, {}, [{"code": "DOCUMENT_NOT_FOUND", "message": f"Document not registered: {document_id}"}], False
    if document["component"] != "spreadsheets":
        return False, {"document": document}, [{"code": "UNSUPPORTED_COMPONENT", "message": "spreadsheet-rename-sheet requires a spreadsheet document."}], False
    path = Path(document["path"])
    if not path.exists():
        return False, {"document": document}, [{"code": INPUT_FILE_NOT_FOUND, "message": f"Input file not found: {path}"}], False

    workbook = load_workbook(path, data_only=False, read_only=True)
    try:
        names = list(workbook.sheetnames)
    finally:
        workbook.close()
    if old_name not in names:
        return False, {"document_id": document_id, "sheet_names": names}, [{"code": "SHEET_NOT_FOUND", "message": f"Worksheet not found: {old_name}"}], False
    if old_name != new_name and any(name.casefold() == new_name.casefold() for name in names):
        return False, {"document_id": document_id, "sheet_names": names}, [{"code": "SHEET_NAME_CONFLICT", "message": f"Worksheet name already exists: {new_name}"}], False
    expected_names = [new_name if name == old_name else name for name in names]
    preview = {
        "document_id": document_id, "path": str(path), "old_name": old_name,
        "new_name": new_name, "sheet_names_before": names,
        "sheet_names_after": expected_names, "would_modify": old_name != new_name,
        "dry_run": dry_run,
    }
    if dry_run:
        return True, preview, [], False
    if old_name == new_name:
        result = {**preview, "dry_run": False, "saved": False, "validation_passed": True}
        record_operation(request_id, "spreadsheet-rename-sheet", result, workspace)
        return True, result, [], False

    backup_ok, backup_result, backup_errors, backup_replayed = create_backup(
        document_id, f"{request_id}:backup", workspace=workspace,
    )
    if not backup_ok:
        return False, {"preview": preview, "backup": backup_result}, backup_errors, False
    rename_result = guarded_com_mutation(
        path, backup_result, lambda: _run_spreadsheet_rename_sheet_com(str(path), old_name, new_name),
    )
    if not rename_result["ok"]:
        return False, {"preview": preview, "backup": backup_result, "rename": rename_result["data"]}, rename_result["errors"], False
    readback = load_workbook(path, data_only=False, read_only=True)
    try:
        actual_names = list(readback.sheetnames)
    finally:
        readback.close()
    validation_passed = actual_names == expected_names and rename_result["data"].get("sheet_names") == expected_names
    result = {
        **preview, "dry_run": False, "saved": True, "backup": backup_result,
        "backup_replayed": backup_replayed, "rename_backend": rename_result["data"].get("backend"),
        "sheet_names_read_back": actual_names, "validation_passed": validation_passed,
    }
    if not validation_passed:
        return False, result, [{"code": "VALIDATION_FAILED", "message": "Worksheet rename order did not match the expected sheet list."}], False
    post_error = verify_post_com_source(path, rename_result)
    if post_error is not None:
        return False, result, [post_error], False
    record_operation(request_id, "spreadsheet-rename-sheet", result, workspace)
    return True, result, [], False


def _run_spreadsheet_create_sheet_com(path: str, name: str, index: int) -> dict[str, Any]:
    capabilities = probe_wps_capabilities()
    prog_id = capabilities["components"]["spreadsheets"]["selected_prog_id"]
    if not prog_id:
        return {"ok": False, "errors": [{"code": COM_BACKEND_UNAVAILABLE, "message": "No registered WPS ProgID detected for spreadsheets."}], "data": {"capabilities": capabilities}}
    params_json = json.dumps({"prog_id": prog_id, "path": str(Path(path).resolve()), "name": name, "index": index})
    script = f"""
$ErrorActionPreference = 'Stop'
$params = @'
{params_json}
'@ | ConvertFrom-Json
$app = $null; $workbook = $null
try {{
  $app = New-Object -ComObject $params.prog_id
  try {{ $app.Visible = $false }} catch {{ }}
  try {{ $app.DisplayAlerts = $false }} catch {{ }}
  $workbook = $app.Workbooks.Open($params.path)
  $count = $workbook.Worksheets.Count
  if ([int]$params.index -eq $count + 1) {{
    $sheet = $workbook.Worksheets.Add($null, $workbook.Worksheets.Item($count))
  }} else {{
    $sheet = $workbook.Worksheets.Add($workbook.Worksheets.Item([int]$params.index))
  }}
  $sheet.Name = $params.name
  $workbook.Save()
  $names = @()
  for ($position = 1; $position -le $workbook.Worksheets.Count; $position++) {{
    $names += [string]$workbook.Worksheets.Item($position).Name
  }}
  [pscustomobject]@{{ ok = $true; backend = 'powershell-com'; prog_id = $params.prog_id; sheet_names = $names; created_index = [int]$params.index }} | ConvertTo-Json -Depth 5 -Compress
}} catch {{
  [pscustomobject]@{{ ok = $false; backend = 'powershell-com'; error_type = $_.Exception.GetType().FullName; error_message = $_.Exception.Message; position = $_.InvocationInfo.PositionMessage }} | ConvertTo-Json -Depth 5 -Compress
  exit 2
}} finally {{
  if ($workbook -ne $null) {{ try {{ $workbook.Close($false) }} catch {{ }}; try {{ [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($workbook) }} catch {{ }} }}
  if ($app -ne $null) {{ try {{ $app.Quit() }} catch {{ }}; try {{ [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($app) }} catch {{ }} }}
  [GC]::Collect(); [GC]::WaitForPendingFinalizers()
}}
"""
    script_path = None
    try:
        with tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False, encoding="utf-8-sig") as stream:
            stream.write(script)
            script_path = stream.name
        completed = subprocess.run(
            [powershell_executable(), "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", script_path],
            check=False, capture_output=True, encoding="utf-8", errors="replace", text=True, timeout=120,
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        return {"ok": False, "errors": [{"code": COM_OPERATION_FAILED, "message": str(exc)}], "data": {}}
    finally:
        if script_path:
            try:
                Path(script_path).unlink(missing_ok=True)
            except OSError:
                pass
    try:
        payload = json.loads(completed.stdout)
    except json.JSONDecodeError as exc:
        return {"ok": False, "errors": [{"code": COM_OPERATION_FAILED, "message": f"WPS returned invalid JSON: {exc}"}], "data": {}}
    if completed.returncode != 0 or not payload.get("ok"):
        return {"ok": False, "errors": [{"code": COM_OPERATION_FAILED, "message": payload.get("error_message", "WPS worksheet creation failed.")}], "data": {"diagnostic": payload}}
    if isinstance(payload.get("sheet_names"), str):
        payload["sheet_names"] = [payload["sheet_names"]]
    return {"ok": True, "errors": [], "data": payload}


@coordinated_mutation
def create_spreadsheet_sheet(
    document_id: str,
    name: str,
    request_id: str,
    index: int | None = None,
    dry_run: bool = False,
    workspace: str | Path = ".",
) -> tuple[bool, dict[str, Any], list[dict[str, str]], bool]:
    replay = replay_operation(
        request_id, "spreadsheet-create-sheet",
        {"document_id": document_id, "name": name, "index": index, "dry_run": dry_run},
        workspace,
    )
    if replay is not None:
        return replay
    title_error = _validate_sheet_title(name)
    if title_error:
        return False, {}, [title_error], False
    document = get_document(document_id, workspace)
    if not document:
        return False, {}, [{"code": "DOCUMENT_NOT_FOUND", "message": f"Document not registered: {document_id}"}], False
    if document["component"] != "spreadsheets":
        return False, {"document": document}, [{"code": "UNSUPPORTED_COMPONENT", "message": "spreadsheet-create-sheet requires a spreadsheet document."}], False
    path = Path(document["path"])
    if not path.exists():
        return False, {"document": document}, [{"code": INPUT_FILE_NOT_FOUND, "message": f"Input file not found: {path}"}], False
    workbook = load_workbook(path, data_only=False, read_only=True)
    try:
        names = list(workbook.sheetnames)
    finally:
        workbook.close()
    if any(existing.casefold() == name.casefold() for existing in names):
        return False, {"document_id": document_id, "sheet_names": names}, [{"code": "SHEET_NAME_CONFLICT", "message": f"Worksheet name already exists: {name}"}], False
    insertion_index = len(names) + 1 if index is None else index
    if isinstance(insertion_index, bool) or not isinstance(insertion_index, int) or insertion_index < 1 or insertion_index > len(names) + 1:
        return False, {"document_id": document_id, "sheet_count": len(names), "index": insertion_index}, [{"code": "INVALID_SHEET_INDEX", "message": f"index must be between 1 and {len(names) + 1}."}], False
    expected_names = list(names)
    expected_names.insert(insertion_index - 1, name)
    preview = {
        "document_id": document_id, "path": str(path), "name": name,
        "index": insertion_index, "requested_index": index,
        "sheet_names_before": names, "sheet_names_after": expected_names,
        "dry_run": dry_run,
    }
    if dry_run:
        return True, preview, [], False
    backup_ok, backup_result, backup_errors, backup_replayed = create_backup(
        document_id, f"{request_id}:backup", workspace=workspace,
    )
    if not backup_ok:
        return False, {"preview": preview, "backup": backup_result}, backup_errors, False
    creation = guarded_com_mutation(
        path, backup_result, lambda: _run_spreadsheet_create_sheet_com(str(path), name, insertion_index),
    )
    if not creation["ok"]:
        return False, {"preview": preview, "backup": backup_result, "creation": creation["data"]}, creation["errors"], False
    readback = load_workbook(path, data_only=False, read_only=True)
    try:
        actual_names = list(readback.sheetnames)
    finally:
        readback.close()
    validation_passed = actual_names == expected_names and creation["data"].get("sheet_names") == expected_names
    result = {
        **preview, "dry_run": False, "backup": backup_result,
        "backup_replayed": backup_replayed, "create_backend": creation["data"].get("backend"),
        "sheet_names_read_back": actual_names, "validation_passed": validation_passed,
    }
    if not validation_passed:
        return False, result, [{"code": "VALIDATION_FAILED", "message": "Created worksheet order did not match the expected sheet list."}], False
    post_error = verify_post_com_source(path, creation)
    if post_error is not None:
        return False, result, [post_error], False
    record_operation(request_id, "spreadsheet-create-sheet", result, workspace)
    return True, result, [], False


def _run_spreadsheet_visibility_com(path: str, name: str, visible: bool) -> dict[str, Any]:
    capabilities = probe_wps_capabilities()
    prog_id = capabilities["components"]["spreadsheets"]["selected_prog_id"]
    if not prog_id:
        return {"ok": False, "errors": [{"code": COM_BACKEND_UNAVAILABLE, "message": "No registered WPS ProgID detected for spreadsheets."}], "data": {"capabilities": capabilities}}
    params_json = json.dumps({"prog_id": prog_id, "path": str(Path(path).resolve()), "name": name, "visible": visible})
    script = f"""
$ErrorActionPreference = 'Stop'
$params = @'
{params_json}
'@ | ConvertFrom-Json
$app = $null; $workbook = $null
try {{
  $app = New-Object -ComObject $params.prog_id
  try {{ $app.Visible = $false }} catch {{ }}
  try {{ $app.DisplayAlerts = $false }} catch {{ }}
  $workbook = $app.Workbooks.Open($params.path)
  $sheet = $workbook.Worksheets.Item($params.name)
  $sheet.Visible = if ([bool]$params.visible) {{ -1 }} else {{ 0 }}
  $workbook.Save()
  $states = @()
  for ($position = 1; $position -le $workbook.Worksheets.Count; $position++) {{
    $item = $workbook.Worksheets.Item($position)
    $states += [pscustomobject]@{{ name = [string]$item.Name; visible = ([int]$item.Visible -eq -1) }}
  }}
  [pscustomobject]@{{ ok = $true; backend = 'powershell-com'; prog_id = $params.prog_id; sheet_states = $states }} | ConvertTo-Json -Depth 5 -Compress
}} catch {{
  [pscustomobject]@{{ ok = $false; error_message = $_.Exception.Message }} | ConvertTo-Json -Depth 5 -Compress
  exit 2
}} finally {{
  if ($workbook -ne $null) {{ try {{ $workbook.Close($false) }} catch {{ }}; try {{ [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($workbook) }} catch {{ }} }}
  if ($app -ne $null) {{ try {{ $app.Quit() }} catch {{ }}; try {{ [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($app) }} catch {{ }} }}
  [GC]::Collect(); [GC]::WaitForPendingFinalizers()
}}
"""
    script_path = None
    try:
        with tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False, encoding="utf-8-sig") as stream:
            stream.write(script)
            script_path = stream.name
        completed = subprocess.run([powershell_executable(), "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", script_path], check=False, capture_output=True, encoding="utf-8", errors="replace", text=True, timeout=120)
    except (OSError, subprocess.TimeoutExpired) as exc:
        return {"ok": False, "errors": [{"code": COM_OPERATION_FAILED, "message": str(exc)}], "data": {}}
    finally:
        if script_path:
            try:
                Path(script_path).unlink(missing_ok=True)
            except OSError:
                pass
    try:
        payload = json.loads(completed.stdout)
    except json.JSONDecodeError as exc:
        return {"ok": False, "errors": [{"code": COM_OPERATION_FAILED, "message": f"WPS returned invalid JSON: {exc}"}], "data": {}}
    if completed.returncode != 0 or not payload.get("ok"):
        return {"ok": False, "errors": [{"code": COM_OPERATION_FAILED, "message": payload.get("error_message", "WPS worksheet visibility update failed.")}], "data": {"diagnostic": payload}}
    if isinstance(payload.get("sheet_states"), dict):
        payload["sheet_states"] = [payload["sheet_states"]]
    return {"ok": True, "errors": [], "data": payload}


@coordinated_mutation
def set_spreadsheet_sheet_visibility(
    document_id: str, sheet_name: str, visible: bool, request_id: str,
    dry_run: bool = False, workspace: str | Path = ".",
) -> tuple[bool, dict[str, Any], list[dict[str, str]], bool]:
    replay = replay_operation(request_id, "spreadsheet-set-sheet-visibility", {
        "document_id": document_id, "sheet_name": sheet_name, "visible": visible, "dry_run": dry_run,
    }, workspace)
    if replay is not None:
        return replay
    if not isinstance(visible, bool):
        return False, {}, [{"code": "INVALID_ARGUMENT", "message": "visible must be a boolean."}], False
    document = get_document(document_id, workspace)
    if not document:
        return False, {}, [{"code": "DOCUMENT_NOT_FOUND", "message": f"Document not registered: {document_id}"}], False
    if document["component"] != "spreadsheets":
        return False, {"document": document}, [{"code": "UNSUPPORTED_COMPONENT", "message": "spreadsheet-set-sheet-visibility requires a spreadsheet document."}], False
    path = Path(document["path"])
    if not path.exists():
        return False, {"document": document}, [{"code": INPUT_FILE_NOT_FOUND, "message": f"Input file not found: {path}"}], False
    workbook = load_workbook(path, data_only=False, read_only=True)
    try:
        states = [{"name": sheet.title, "visible": sheet.sheet_state == "visible"} for sheet in workbook.worksheets]
    finally:
        workbook.close()
    target = next((sheet for sheet in states if sheet["name"] == sheet_name), None)
    if target is None:
        return False, {"document_id": document_id, "sheet_states": states}, [{"code": "SHEET_NOT_FOUND", "message": f"Worksheet not found: {sheet_name}"}], False
    current_visible_count = sum(item["visible"] for item in states)
    if not visible and target["visible"] and current_visible_count <= 1:
        return False, {"document_id": document_id, "sheet_states": states}, [{"code": "LAST_VISIBLE_SHEET", "message": "At least one worksheet must remain visible."}], False
    expected = [dict(item, visible=(visible if item["name"] == sheet_name else item["visible"])) for item in states]
    preview = {"document_id": document_id, "path": str(path), "sheet_name": sheet_name, "visible": visible, "sheet_states_before": states, "sheet_states_after": expected, "dry_run": dry_run}
    if dry_run:
        return True, preview, [], False
    if target["visible"] == visible:
        result = {**preview, "dry_run": False, "saved": False, "validation_passed": True}
        record_operation(request_id, "spreadsheet-set-sheet-visibility", result, workspace)
        return True, result, [], False
    backup_ok, backup_result, backup_errors, backup_replayed = create_backup(document_id, f"{request_id}:backup", workspace=workspace)
    if not backup_ok:
        return False, {"preview": preview, "backup": backup_result}, backup_errors, False
    mutation = guarded_com_mutation(
        path, backup_result, lambda: _run_spreadsheet_visibility_com(str(path), sheet_name, visible),
    )
    if not mutation["ok"]:
        return False, {"preview": preview, "backup": backup_result, "mutation": mutation["data"]}, mutation["errors"], False
    readback = load_workbook(path, data_only=False, read_only=True)
    try:
        actual = [{"name": sheet.title, "visible": sheet.sheet_state == "visible"} for sheet in readback.worksheets]
    finally:
        readback.close()
    validation_passed = actual == expected and mutation["data"].get("sheet_states") == expected
    result = {**preview, "dry_run": False, "saved": True, "backup": backup_result, "backup_replayed": backup_replayed, "visibility_backend": mutation["data"].get("backend"), "sheet_states_read_back": actual, "validation_passed": validation_passed}
    if not validation_passed:
        return False, result, [{"code": "VALIDATION_FAILED", "message": "Worksheet visibility state did not match the expected ordered state list."}], False
    post_error = verify_post_com_source(path, mutation)
    if post_error is not None:
        return False, result, [post_error], False
    record_operation(request_id, "spreadsheet-set-sheet-visibility", result, workspace)
    return True, result, [], False


def _run_spreadsheet_delete_sheet_com(path: str, name: str) -> dict[str, Any]:
    capabilities = probe_wps_capabilities()
    prog_id = capabilities["components"]["spreadsheets"]["selected_prog_id"]
    if not prog_id:
        return {"ok": False, "errors": [{"code": COM_BACKEND_UNAVAILABLE, "message": "No registered WPS ProgID detected for spreadsheets."}], "data": {"capabilities": capabilities}}
    params_json = json.dumps({"prog_id": prog_id, "path": str(Path(path).resolve()), "name": name})
    script = f"""
$ErrorActionPreference = 'Stop'
$params = @'
{params_json}
'@ | ConvertFrom-Json
$app = $null; $workbook = $null
try {{
  $app = New-Object -ComObject $params.prog_id
  try {{ $app.Visible = $false }} catch {{ }}
  $app.DisplayAlerts = $false
  $workbook = $app.Workbooks.Open($params.path)
  $workbook.Worksheets.Item($params.name).Delete()
  $workbook.Save()
  $names = @()
  for ($position = 1; $position -le $workbook.Worksheets.Count; $position++) {{ $names += [string]$workbook.Worksheets.Item($position).Name }}
  [pscustomobject]@{{ ok = $true; backend = 'powershell-com'; prog_id = $params.prog_id; sheet_names = $names }} | ConvertTo-Json -Depth 5 -Compress
}} catch {{
  [pscustomobject]@{{ ok = $false; error_message = $_.Exception.Message }} | ConvertTo-Json -Depth 5 -Compress
  exit 2
}} finally {{
  if ($workbook -ne $null) {{ try {{ $workbook.Close($false) }} catch {{ }}; try {{ [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($workbook) }} catch {{ }} }}
  if ($app -ne $null) {{ try {{ $app.Quit() }} catch {{ }}; try {{ [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($app) }} catch {{ }} }}
  [GC]::Collect(); [GC]::WaitForPendingFinalizers()
}}
"""
    script_path = None
    try:
        with tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False, encoding="utf-8-sig") as stream:
            stream.write(script)
            script_path = stream.name
        completed = subprocess.run([powershell_executable(), "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", script_path], check=False, capture_output=True, encoding="utf-8", errors="replace", text=True, timeout=120)
    except (OSError, subprocess.TimeoutExpired) as exc:
        return {"ok": False, "errors": [{"code": COM_OPERATION_FAILED, "message": str(exc)}], "data": {}}
    finally:
        if script_path:
            try:
                Path(script_path).unlink(missing_ok=True)
            except OSError:
                pass
    try:
        payload = json.loads(completed.stdout)
    except json.JSONDecodeError as exc:
        return {"ok": False, "errors": [{"code": COM_OPERATION_FAILED, "message": f"WPS returned invalid JSON: {exc}"}], "data": {}}
    if completed.returncode != 0 or not payload.get("ok"):
        return {"ok": False, "errors": [{"code": COM_OPERATION_FAILED, "message": payload.get("error_message", "WPS worksheet deletion failed.")}], "data": {"diagnostic": payload}}
    if isinstance(payload.get("sheet_names"), str):
        payload["sheet_names"] = [payload["sheet_names"]]
    return {"ok": True, "errors": [], "data": payload}


@coordinated_mutation
def delete_spreadsheet_sheet(
    document_id: str, sheet_name: str, request_id: str,
    dry_run: bool = False, workspace: str | Path = ".",
) -> tuple[bool, dict[str, Any], list[dict[str, str]], bool]:
    replay = replay_operation(request_id, "spreadsheet-delete-sheet", {
        "document_id": document_id, "sheet_name": sheet_name, "dry_run": dry_run,
    }, workspace)
    if replay is not None:
        return replay
    document = get_document(document_id, workspace)
    if not document:
        return False, {}, [{"code": "DOCUMENT_NOT_FOUND", "message": f"Document not registered: {document_id}"}], False
    if document["component"] != "spreadsheets":
        return False, {"document": document}, [{"code": "UNSUPPORTED_COMPONENT", "message": "spreadsheet-delete-sheet requires a spreadsheet document."}], False
    path = Path(document["path"])
    if not path.exists():
        return False, {"document": document}, [{"code": INPUT_FILE_NOT_FOUND, "message": f"Input file not found: {path}"}], False
    workbook = load_workbook(path, data_only=False, read_only=True)
    try:
        names = list(workbook.sheetnames)
    finally:
        workbook.close()
    if sheet_name not in names:
        return False, {"document_id": document_id, "sheet_names": names}, [{"code": "SHEET_NOT_FOUND", "message": f"Worksheet not found: {sheet_name}"}], False
    if len(names) <= 1:
        return False, {"document_id": document_id, "sheet_names": names}, [{"code": "LAST_WORKSHEET", "message": "A workbook must retain at least one worksheet."}], False
    expected = [name for name in names if name != sheet_name]
    preview = {"document_id": document_id, "path": str(path), "sheet_name": sheet_name, "sheet_names_before": names, "sheet_names_after": expected, "dry_run": dry_run}
    if dry_run:
        return True, preview, [], False
    backup_ok, backup_result, backup_errors, backup_replayed = create_backup(document_id, f"{request_id}:backup", workspace=workspace)
    if not backup_ok:
        return False, {"preview": preview, "backup": backup_result}, backup_errors, False
    mutation = guarded_com_mutation(
        path, backup_result, lambda: _run_spreadsheet_delete_sheet_com(str(path), sheet_name),
    )
    if not mutation["ok"]:
        return False, {"preview": preview, "backup": backup_result, "mutation": mutation["data"]}, mutation["errors"], False
    readback = load_workbook(path, data_only=False, read_only=True)
    try:
        actual = list(readback.sheetnames)
    finally:
        readback.close()
    validation_passed = actual == expected and mutation["data"].get("sheet_names") == expected
    result = {**preview, "dry_run": False, "saved": True, "backup": backup_result, "backup_replayed": backup_replayed, "delete_backend": mutation["data"].get("backend"), "sheet_names_read_back": actual, "validation_passed": validation_passed}
    if not validation_passed:
        return False, result, [{"code": "VALIDATION_FAILED", "message": "Worksheet deletion order did not match the expected sheet list."}], False
    post_error = verify_post_com_source(path, mutation)
    if post_error is not None:
        return False, result, [post_error], False
    record_operation(request_id, "spreadsheet-delete-sheet", result, workspace)
    return True, result, [], False


def _run_spreadsheet_copy_sheet_com(path: str, source_name: str, new_name: str, index: int) -> dict[str, Any]:
    capabilities = probe_wps_capabilities()
    prog_id = capabilities["components"]["spreadsheets"]["selected_prog_id"]
    if not prog_id:
        return {"ok": False, "errors": [{"code": COM_BACKEND_UNAVAILABLE, "message": "No registered WPS ProgID detected for spreadsheets."}], "data": {"capabilities": capabilities}}
    params_json = json.dumps({"prog_id": prog_id, "path": str(Path(path).resolve()), "source_name": source_name, "new_name": new_name, "index": index})
    script = f"""
$ErrorActionPreference = 'Stop'
$params = @'
{params_json}
'@ | ConvertFrom-Json
$app = $null; $workbook = $null
try {{
  $app = New-Object -ComObject $params.prog_id
  try {{ $app.Visible = $false }} catch {{ }}
  try {{ $app.DisplayAlerts = $false }} catch {{ }}
  $workbook = $app.Workbooks.Open($params.path)
  $source = $workbook.Worksheets.Item($params.source_name)
  $count = $workbook.Worksheets.Count
  if ([int]$params.index -eq $count + 1) {{
    $source.Copy($null, $workbook.Worksheets.Item($count))
  }} else {{
    $source.Copy($workbook.Worksheets.Item([int]$params.index))
  }}
  $copy = $workbook.ActiveSheet
  $copy.Name = $params.new_name
  $workbook.Save()
  $names = @()
  for ($position = 1; $position -le $workbook.Worksheets.Count; $position++) {{ $names += [string]$workbook.Worksheets.Item($position).Name }}
  [pscustomobject]@{{ ok = $true; backend = 'powershell-com'; prog_id = $params.prog_id; sheet_names = $names }} | ConvertTo-Json -Depth 5 -Compress
}} catch {{
  [pscustomobject]@{{ ok = $false; error_message = $_.Exception.Message }} | ConvertTo-Json -Depth 5 -Compress
  exit 2
}} finally {{
  if ($workbook -ne $null) {{ try {{ $workbook.Close($false) }} catch {{ }}; try {{ [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($workbook) }} catch {{ }} }}
  if ($app -ne $null) {{ try {{ $app.Quit() }} catch {{ }}; try {{ [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($app) }} catch {{ }} }}
  [GC]::Collect(); [GC]::WaitForPendingFinalizers()
}}
"""
    script_path = None
    try:
        with tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False, encoding="utf-8-sig") as stream:
            stream.write(script)
            script_path = stream.name
        completed = subprocess.run([powershell_executable(), "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", script_path], check=False, capture_output=True, encoding="utf-8", errors="replace", text=True, timeout=120)
    except (OSError, subprocess.TimeoutExpired) as exc:
        return {"ok": False, "errors": [{"code": COM_OPERATION_FAILED, "message": str(exc)}], "data": {}}
    finally:
        if script_path:
            try:
                Path(script_path).unlink(missing_ok=True)
            except OSError:
                pass
    try:
        payload = json.loads(completed.stdout)
    except json.JSONDecodeError as exc:
        return {"ok": False, "errors": [{"code": COM_OPERATION_FAILED, "message": f"WPS returned invalid JSON: {exc}"}], "data": {}}
    if completed.returncode != 0 or not payload.get("ok"):
        return {"ok": False, "errors": [{"code": COM_OPERATION_FAILED, "message": payload.get("error_message", "WPS worksheet copy failed.")}], "data": {"diagnostic": payload}}
    if isinstance(payload.get("sheet_names"), str):
        payload["sheet_names"] = [payload["sheet_names"]]
    return {"ok": True, "errors": [], "data": payload}


@coordinated_mutation
def copy_spreadsheet_sheet(
    document_id: str, source_name: str, new_name: str, request_id: str,
    index: int | None = None, dry_run: bool = False, workspace: str | Path = ".",
) -> tuple[bool, dict[str, Any], list[dict[str, str]], bool]:
    replay = replay_operation(request_id, "spreadsheet-copy-sheet", {
        "document_id": document_id, "source_name": source_name, "new_name": new_name, "index": index, "dry_run": dry_run,
    }, workspace)
    if replay is not None:
        return replay
    title_error = _validate_sheet_title(new_name)
    if title_error:
        return False, {}, [title_error], False
    document = get_document(document_id, workspace)
    if not document:
        return False, {}, [{"code": "DOCUMENT_NOT_FOUND", "message": f"Document not registered: {document_id}"}], False
    if document["component"] != "spreadsheets":
        return False, {"document": document}, [{"code": "UNSUPPORTED_COMPONENT", "message": "spreadsheet-copy-sheet requires a spreadsheet document."}], False
    path = Path(document["path"])
    if not path.exists():
        return False, {"document": document}, [{"code": INPUT_FILE_NOT_FOUND, "message": f"Input file not found: {path}"}], False
    workbook = load_workbook(path, data_only=False, read_only=True)
    try:
        names = list(workbook.sheetnames)
    finally:
        workbook.close()
    if source_name not in names:
        return False, {"document_id": document_id, "sheet_names": names}, [{"code": "SHEET_NOT_FOUND", "message": f"Worksheet not found: {source_name}"}], False
    if any(existing.casefold() == new_name.casefold() for existing in names):
        return False, {"document_id": document_id, "sheet_names": names}, [{"code": "SHEET_NAME_CONFLICT", "message": f"Worksheet name already exists: {new_name}"}], False
    insertion_index = len(names) + 1 if index is None else index
    if isinstance(insertion_index, bool) or not isinstance(insertion_index, int) or insertion_index < 1 or insertion_index > len(names) + 1:
        return False, {"document_id": document_id, "sheet_count": len(names), "index": insertion_index}, [{"code": "INVALID_SHEET_INDEX", "message": f"index must be between 1 and {len(names) + 1}."}], False
    expected = list(names)
    expected.insert(insertion_index - 1, new_name)
    preview = {"document_id": document_id, "path": str(path), "source_name": source_name, "new_name": new_name, "index": insertion_index, "requested_index": index, "sheet_names_before": names, "sheet_names_after": expected, "dry_run": dry_run}
    if dry_run:
        return True, preview, [], False
    backup_ok, backup_result, backup_errors, backup_replayed = create_backup(document_id, f"{request_id}:backup", workspace=workspace)
    if not backup_ok:
        return False, {"preview": preview, "backup": backup_result}, backup_errors, False
    mutation = guarded_com_mutation(
        path, backup_result,
        lambda: _run_spreadsheet_copy_sheet_com(str(path), source_name, new_name, insertion_index),
    )
    if not mutation["ok"]:
        return False, {"preview": preview, "backup": backup_result, "mutation": mutation["data"]}, mutation["errors"], False
    readback = load_workbook(path, data_only=False, read_only=True)
    try:
        actual = list(readback.sheetnames)
        source_value = readback[source_name]["A1"].value
        copied_value = readback[new_name]["A1"].value
    finally:
        readback.close()
    validation_passed = actual == expected and mutation["data"].get("sheet_names") == expected and copied_value == source_value
    result = {**preview, "dry_run": False, "saved": True, "backup": backup_result, "backup_replayed": backup_replayed, "copy_backend": mutation["data"].get("backend"), "sheet_names_read_back": actual, "source_a1_read_back": source_value, "copy_a1_read_back": copied_value, "validation_passed": validation_passed}
    if not validation_passed:
        return False, result, [{"code": "VALIDATION_FAILED", "message": "Copied worksheet order or sampled cell read-back did not match the source."}], False
    post_error = verify_post_com_source(path, mutation)
    if post_error is not None:
        return False, result, [post_error], False
    record_operation(request_id, "spreadsheet-copy-sheet", result, workspace)
    return True, result, [], False


def _normalize_sheet_tab_color(color: str) -> str | None:
    if not isinstance(color, str):
        return None
    if color.casefold() == "none":
        return "none"
    if len(color) != 7 or color[0] != "#" or any(character not in "0123456789abcdefABCDEF" for character in color[1:]):
        return None
    return color.upper()


def _tab_color_from_cell(color: Any) -> str:
    if color is None or color.type is None:
        return "none"
    value = color.rgb if color.type == "rgb" else None
    return f"#{value[-6:].upper()}" if isinstance(value, str) and len(value) >= 6 else "unsupported"


def _read_ignored_errors(path: Path) -> dict[str, dict[str, Any]]:
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    relationship_namespace = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    package_relationship_namespace = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    with ZipFile(path) as archive:
        workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
        relationships_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        relationships = {
            item.attrib["Id"]: item.attrib["Target"]
            for item in relationships_root.findall(f"{package_relationship_namespace}Relationship")
        }
        result: dict[str, dict[str, Any]] = {}
        for sheet in workbook_root.findall(f"{namespace}sheets/{namespace}sheet"):
            relationship_id = sheet.attrib.get(f"{relationship_namespace}id")
            target = relationships.get(relationship_id)
            if not target:
                continue
            member = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
            records = []
            count = 0
            with archive.open(member) as stream:
                for _event, element in ET.iterparse(stream, events=("end",)):
                    if element.tag == f"{namespace}ignoredError":
                        count += 1
                        if len(records) < 1_000:
                            attributes = dict(element.attrib)
                            records.append({
                                "ranges": attributes.pop("sqref", ""),
                                "flags": dict(sorted(attributes.items())),
                            })
                        element.clear()
            records.sort(key=lambda item: (item["ranges"], tuple(item["flags"].items())))
            result[sheet.attrib["name"]] = {
                "records": records,
                "count": count,
                "truncated": count > 1_000,
            }
        return result


def _run_spreadsheet_tab_color_com(path: str, name: str, color: str) -> dict[str, Any]:
    capabilities = probe_wps_capabilities()
    prog_id = capabilities["components"]["spreadsheets"]["selected_prog_id"]
    if not prog_id:
        return {"ok": False, "errors": [{"code": COM_BACKEND_UNAVAILABLE, "message": "No registered WPS ProgID detected for spreadsheets."}], "data": {"capabilities": capabilities}}
    rgb = int(color[1:3], 16) + (int(color[3:5], 16) << 8) + (int(color[5:7], 16) << 16) if color != "none" else None
    params_json = json.dumps({"prog_id": prog_id, "path": str(Path(path).resolve()), "name": name, "color": color, "rgb": rgb})
    script = f"""
$ErrorActionPreference = 'Stop'
$params = @'
{params_json}
'@ | ConvertFrom-Json
$app = $null; $workbook = $null
try {{
  $app = New-Object -ComObject $params.prog_id
  try {{ $app.Visible = $false }} catch {{ }}
  try {{ $app.DisplayAlerts = $false }} catch {{ }}
  $workbook = $app.Workbooks.Open($params.path)
  $sheet = $workbook.Worksheets.Item($params.name)
  if ($params.color -eq 'none') {{ $sheet.Tab.ColorIndex = -4142 }} else {{ $sheet.Tab.Color = [int]$params.rgb }}
  $workbook.Save()
  $actualColor = 'none'
  if ($params.color -ne 'none') {{
    $value = [int]$sheet.Tab.Color
    $red = $value -band 255; $green = ($value -shr 8) -band 255; $blue = ($value -shr 16) -band 255
    $actualColor = ('#{{0:X2}}{{1:X2}}{{2:X2}}' -f $red, $green, $blue)
  }}
  [pscustomobject]@{{ ok = $true; backend = 'powershell-com'; prog_id = $params.prog_id; sheet_name = [string]$sheet.Name; tab_color = $actualColor }} | ConvertTo-Json -Depth 5 -Compress
}} catch {{
  [pscustomobject]@{{ ok = $false; error_message = $_.Exception.Message }} | ConvertTo-Json -Depth 5 -Compress
  exit 2
}} finally {{
  if ($workbook -ne $null) {{ try {{ $workbook.Close($false) }} catch {{ }}; try {{ [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($workbook) }} catch {{ }} }}
  if ($app -ne $null) {{ try {{ $app.Quit() }} catch {{ }}; try {{ [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($app) }} catch {{ }} }}
  [GC]::Collect(); [GC]::WaitForPendingFinalizers()
}}
"""
    script_path = None
    try:
        with tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False, encoding="utf-8-sig") as stream:
            stream.write(script)
            script_path = stream.name
        completed = subprocess.run([powershell_executable(), "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", script_path], check=False, capture_output=True, encoding="utf-8", errors="replace", text=True, timeout=120)
    except (OSError, subprocess.TimeoutExpired) as exc:
        return {"ok": False, "errors": [{"code": COM_OPERATION_FAILED, "message": str(exc)}], "data": {}}
    finally:
        if script_path:
            try:
                Path(script_path).unlink(missing_ok=True)
            except OSError:
                pass
    try:
        payload = json.loads(completed.stdout)
    except json.JSONDecodeError as exc:
        return {"ok": False, "errors": [{"code": COM_OPERATION_FAILED, "message": f"WPS returned invalid JSON: {exc}"}], "data": {}}
    if completed.returncode != 0 or not payload.get("ok"):
        return {"ok": False, "errors": [{"code": COM_OPERATION_FAILED, "message": payload.get("error_message", "WPS tab color update failed.")}], "data": {"diagnostic": payload}}
    return {"ok": True, "errors": [], "data": payload}


@coordinated_mutation
def set_spreadsheet_sheet_tab_color(
    document_id: str, sheet_name: str, color: str, request_id: str,
    dry_run: bool = False, workspace: str | Path = ".",
) -> tuple[bool, dict[str, Any], list[dict[str, str]], bool]:
    replay = replay_operation(request_id, "spreadsheet-set-sheet-tab-color", {
        "document_id": document_id, "sheet_name": sheet_name, "color": color, "dry_run": dry_run,
    }, workspace)
    if replay is not None:
        return replay
    normalized = _normalize_sheet_tab_color(color)
    if normalized is None:
        return False, {}, [{"code": "INVALID_TAB_COLOR", "message": "color must be #RRGGBB or 'none'."}], False
    document = get_document(document_id, workspace)
    if not document:
        return False, {}, [{"code": "DOCUMENT_NOT_FOUND", "message": f"Document not registered: {document_id}"}], False
    if document["component"] != "spreadsheets":
        return False, {"document": document}, [{"code": "UNSUPPORTED_COMPONENT", "message": "spreadsheet-set-sheet-tab-color requires a spreadsheet document."}], False
    path = Path(document["path"])
    if not path.exists():
        return False, {"document": document}, [{"code": INPUT_FILE_NOT_FOUND, "message": f"Input file not found: {path}"}], False
    workbook = load_workbook(path, data_only=False, read_only=False)
    try:
        if sheet_name not in workbook.sheetnames:
            return False, {"document_id": document_id, "sheet_names": workbook.sheetnames}, [{"code": "SHEET_NOT_FOUND", "message": f"Worksheet not found: {sheet_name}"}], False
        current = _tab_color_from_cell(workbook[sheet_name].sheet_properties.tabColor)
    finally:
        workbook.close()
    preview = {"document_id": document_id, "path": str(path), "sheet_name": sheet_name, "color": normalized, "previous_color": current, "dry_run": dry_run}
    if dry_run:
        return True, preview, [], False
    if current == normalized:
        result = {**preview, "dry_run": False, "saved": False, "validation_passed": True}
        record_operation(request_id, "spreadsheet-set-sheet-tab-color", result, workspace)
        return True, result, [], False
    backup_ok, backup_result, backup_errors, backup_replayed = create_backup(document_id, f"{request_id}:backup", workspace=workspace)
    if not backup_ok:
        return False, {"preview": preview, "backup": backup_result}, backup_errors, False
    mutation = guarded_com_mutation(
        path, backup_result, lambda: _run_spreadsheet_tab_color_com(str(path), sheet_name, normalized),
    )
    if not mutation["ok"]:
        return False, {"preview": preview, "backup": backup_result, "mutation": mutation["data"]}, mutation["errors"], False
    readback = load_workbook(path, data_only=False, read_only=False)
    try:
        actual = _tab_color_from_cell(readback[sheet_name].sheet_properties.tabColor)
    finally:
        readback.close()
    wps_color = mutation["data"].get("tab_color", "")
    validation_passed = actual == normalized and mutation["data"].get("sheet_name") == sheet_name and isinstance(wps_color, str) and wps_color.casefold() == normalized.casefold()
    result = {**preview, "dry_run": False, "saved": True, "backup": backup_result, "backup_replayed": backup_replayed, "tab_color_backend": mutation["data"].get("backend"), "tab_color_read_back": actual, "validation_passed": validation_passed}
    if not validation_passed:
        return False, result, [{"code": "VALIDATION_FAILED", "message": "Worksheet tab color read-back did not match the requested color."}], False
    post_error = verify_post_com_source(path, mutation)
    if post_error is not None:
        return False, result, [post_error], False
    record_operation(request_id, "spreadsheet-set-sheet-tab-color", result, workspace)
    return True, result, [], False


def _matrix_shape(values: list[list[Any]]) -> tuple[int, int]:
    if not values:
        return 0, 0
    width = len(values[0])
    if width == 0:
        return len(values), 0
    for row in values:
        if len(row) != width:
            raise ValueError("values_json must be a rectangular 2D array.")
    return len(values), width


def _run_spreadsheet_write_com(
    path: str,
    sheet_name: str | None,
    range_address: str,
    values: list[list[Any]],
) -> dict[str, Any]:
    capabilities = probe_wps_capabilities()
    selected_prog_id = capabilities["components"]["spreadsheets"]["selected_prog_id"]
    if not selected_prog_id:
        return {
            "ok": False,
            "errors": [
                {
                    "code": COM_BACKEND_UNAVAILABLE,
                    "message": "No registered WPS ProgID detected for spreadsheets.",
                }
            ],
            "data": {"capabilities": capabilities},
        }

    min_col, min_row, max_col, max_row = range_boundaries(range_address)
    params = {
        "prog_id": selected_prog_id,
        "path": str(Path(path).resolve()),
        "sheet_name": sheet_name,
        "range": range_address,
        "start_row": min_row,
        "start_col": min_col,
        "row_count": max_row - min_row + 1,
        "column_count": max_col - min_col + 1,
        "values": values,
    }
    params_json = json.dumps(params)
    script = f"""
$ErrorActionPreference = 'Stop'
$params = @'
{params_json}
'@ | ConvertFrom-Json
$app = $null
$workbook = $null
try {{
  $app = New-Object -ComObject $params.prog_id
  try {{ $app.Visible = $false }} catch {{ }}
  $workbook = $app.Workbooks.Open($params.path)
  if ($null -ne $params.sheet_name -and $params.sheet_name -ne '') {{
    $sheet = $workbook.Worksheets.Item($params.sheet_name)
  }} else {{
    $sheet = $workbook.Worksheets.Item(1)
  }}
  $targetRange = $sheet.Range($params.range)
  for ($r = 0; $r -lt [int]$params.row_count; $r++) {{
    $row = @($params.values[$r])
    for ($c = 0; $c -lt [int]$params.column_count; $c++) {{
      $cell = $targetRange.Cells.Item($r + 1, $c + 1)
      $value = $row[$c]
      if ($null -eq $value) {{
        $cell.Value2 = $null
      }} elseif ($value -is [int] -or $value -is [long] -or $value -is [double] -or $value -is [decimal]) {{
        $cell.Value2 = [double]$value
      }} else {{
        $cell.Value2 = [string]$value
      }}
    }}
  }}
  $workbook.Save()
  [pscustomobject]@{{
    ok = $true
    backend = 'powershell-com'
    component = 'spreadsheets'
    prog_id = $params.prog_id
    path = $params.path
    sheet = $sheet.Name
    range = $params.range
    row_count = $params.row_count
    column_count = $params.column_count
  }} | ConvertTo-Json -Depth 5
}} catch {{
  [pscustomobject]@{{
    ok = $false
    backend = 'powershell-com'
    component = 'spreadsheets'
    prog_id = $params.prog_id
    error_type = $_.Exception.GetType().FullName
    error_message = $_.Exception.Message
    position = $_.InvocationInfo.PositionMessage
  }} | ConvertTo-Json -Depth 5
  exit 2
}} finally {{
  if ($workbook -ne $null) {{
    try {{ $workbook.Close($false) }} catch {{ }}
    try {{ [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($workbook) }} catch {{ }}
  }}
  if ($app -ne $null) {{
    try {{ $app.Quit() }} catch {{ }}
    try {{ [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($app) }} catch {{ }}
  }}
  [GC]::Collect()
  [GC]::WaitForPendingFinalizers()
}}
"""
    script_path = None
    try:
        with tempfile.NamedTemporaryFile(
            "w",
            suffix=".ps1",
            delete=False,
            encoding="utf-8-sig",
        ) as script_file:
            script_file.write(script)
            script_path = script_file.name
        completed = subprocess.run(
            [
                powershell_executable(),
                "-NoProfile",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                script_path,
            ],
            check=False,
            capture_output=True,
            encoding="utf-8",
            errors="replace",
            text=True,
            timeout=120,
        )
    finally:
        if script_path:
            try:
                Path(script_path).unlink(missing_ok=True)
            except OSError:
                pass

    try:
        payload = json.loads(completed.stdout)
    except json.JSONDecodeError:
        payload = None

    if completed.returncode != 0 or not isinstance(payload, dict):
        return {
            "ok": False,
            "errors": [
                {
                    "code": COM_OPERATION_FAILED,
                    "message": (
                        payload.get("error_message")
                        if isinstance(payload, dict)
                        else (completed.stderr or completed.stdout).strip()
                    ),
                }
            ],
            "data": {"diagnostic": payload, "backend": "powershell-com"},
        }
    return {"ok": True, "errors": [], "data": payload}


def _run_spreadsheet_formula_write_com(
    path: str,
    sheet_name: str | None,
    range_address: str,
    formulas: list[list[str]],
) -> dict[str, Any]:
    capabilities = probe_wps_capabilities()
    selected_prog_id = capabilities["components"]["spreadsheets"]["selected_prog_id"]
    if not selected_prog_id:
        return {
            "ok": False,
            "errors": [
                {
                    "code": COM_BACKEND_UNAVAILABLE,
                    "message": "No registered WPS ProgID detected for spreadsheets.",
                }
            ],
            "data": {"capabilities": capabilities},
        }

    min_col, min_row, max_col, max_row = range_boundaries(range_address)
    params = {
        "prog_id": selected_prog_id,
        "path": str(Path(path).resolve()),
        "sheet_name": sheet_name,
        "range": range_address,
        "start_row": min_row,
        "start_col": min_col,
        "row_count": max_row - min_row + 1,
        "column_count": max_col - min_col + 1,
        "formulas": formulas,
    }
    params_json = json.dumps(params)
    script = f"""
$ErrorActionPreference = 'Stop'
$params = @'
{params_json}
'@ | ConvertFrom-Json
$app = $null
$workbook = $null
try {{
  $app = New-Object -ComObject $params.prog_id
  try {{ $app.Visible = $false }} catch {{ }}
  $workbook = $app.Workbooks.Open($params.path)
  if ($null -ne $params.sheet_name -and $params.sheet_name -ne '') {{
    $sheet = $workbook.Worksheets.Item($params.sheet_name)
  }} else {{
    $sheet = $workbook.Worksheets.Item(1)
  }}
  $targetRange = $sheet.Range($params.range)
  for ($r = 0; $r -lt [int]$params.row_count; $r++) {{
    $row = @($params.formulas[$r])
    for ($c = 0; $c -lt [int]$params.column_count; $c++) {{
      $formula = [string]$row[$c]
      if (-not $formula.StartsWith('=')) {{
        throw "Formula must start with '=': $formula"
      }}
      $cell = $targetRange.Cells.Item($r + 1, $c + 1)
      $cell.Formula = $formula
    }}
  }}
  try {{ $app.CalculateFull() }} catch {{ try {{ $app.Calculate() }} catch {{ $workbook.RefreshAll() }} }}
  $workbook.Save()
  [pscustomobject]@{{
    ok = $true
    backend = 'powershell-com'
    component = 'spreadsheets'
    prog_id = $params.prog_id
    path = $params.path
    sheet = $sheet.Name
    range = $params.range
    row_count = $params.row_count
    column_count = $params.column_count
    calculation_triggered = $true
  }} | ConvertTo-Json -Depth 5
}} catch {{
  [pscustomobject]@{{
    ok = $false
    backend = 'powershell-com'
    component = 'spreadsheets'
    prog_id = $params.prog_id
    error_type = $_.Exception.GetType().FullName
    error_message = $_.Exception.Message
    position = $_.InvocationInfo.PositionMessage
  }} | ConvertTo-Json -Depth 5
  exit 2
}} finally {{
  if ($workbook -ne $null) {{
    try {{ $workbook.Close($false) }} catch {{ }}
    try {{ [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($workbook) }} catch {{ }}
  }}
  if ($app -ne $null) {{
    try {{ $app.Quit() }} catch {{ }}
    try {{ [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($app) }} catch {{ }}
  }}
  [GC]::Collect()
  [GC]::WaitForPendingFinalizers()
}}
"""
    script_path = None
    try:
        with tempfile.NamedTemporaryFile(
            "w",
            suffix=".ps1",
            delete=False,
            encoding="utf-8-sig",
        ) as script_file:
            script_file.write(script)
            script_path = script_file.name
        completed = subprocess.run(
            [
                powershell_executable(),
                "-NoProfile",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                script_path,
            ],
            check=False,
            capture_output=True,
            encoding="utf-8",
            errors="replace",
            text=True,
            timeout=120,
        )
    finally:
        if script_path:
            try:
                Path(script_path).unlink(missing_ok=True)
            except OSError:
                pass

    try:
        payload = json.loads(completed.stdout)
    except json.JSONDecodeError:
        payload = None

    if completed.returncode != 0 or not isinstance(payload, dict):
        return {
            "ok": False,
            "errors": [
                {
                    "code": COM_OPERATION_FAILED,
                    "message": (
                        payload.get("error_message")
                        if isinstance(payload, dict)
                        else (completed.stderr or completed.stdout).strip()
                    ),
                }
            ],
            "data": {"diagnostic": payload, "backend": "powershell-com"},
        }
    return {"ok": True, "errors": [], "data": payload}


@coordinated_mutation
def write_spreadsheet_range(
    document_id: str,
    range_address: str,
    values: list[list[Any]],
    request_id: str,
    sheet_name: str | None = None,
    dry_run: bool = False,
    workspace: str | Path = ".",
) -> tuple[bool, dict[str, Any], list[dict[str, str]], bool]:
    replay = replay_operation(
        request_id, "spreadsheet-write",
        {"document_id": document_id, "range": range_address, "values": values, "requested_sheet": sheet_name, "dry_run": dry_run},
        workspace,
    )
    if replay is not None:
        return replay

    try:
        row_count, column_count = _matrix_shape(values)
    except ValueError as exc:
        return False, {}, [{"code": "INVALID_ARGUMENT", "message": str(exc)}], False
    if row_count == 0 or column_count == 0:
        return False, {}, [{"code": "INVALID_ARGUMENT", "message": "values_json must not be empty."}], False

    bounds, range_error = validate_spreadsheet_read_range(range_address)
    if range_error:
        return False, {}, [range_error], False
    min_col, min_row, max_col, max_row = bounds
    expected_rows = max_row - min_row + 1
    expected_cols = max_col - min_col + 1
    if (row_count, column_count) != (expected_rows, expected_cols):
        return (
            False,
            {
                "range": range_address,
                "range_shape": [expected_rows, expected_cols],
                "values_shape": [row_count, column_count],
            },
            [{"code": "RANGE_SHAPE_MISMATCH", "message": "values_json shape must match target range."}],
            False,
        )

    document = get_document(document_id, workspace)
    if not document:
        return False, {}, [{"code": "DOCUMENT_NOT_FOUND", "message": f"Document not registered: {document_id}"}], False
    if document["component"] != "spreadsheets":
        return (
            False,
            {"document": document},
            [{"code": "UNSUPPORTED_COMPONENT", "message": "spreadsheet-write requires a spreadsheet document."}],
            False,
        )
    path = Path(document["path"])
    if not path.exists():
        return False, {"document": document}, [{"code": INPUT_FILE_NOT_FOUND, "message": f"Input file not found: {path}"}], False

    preview = {
        "document_id": document_id,
        "component": "spreadsheets",
        "path": str(path),
        "sheet": sheet_name,
        "requested_sheet": sheet_name,
        "range": range_address,
        "values": values,
        "row_count": row_count,
        "column_count": column_count,
        "dry_run": dry_run,
    }
    if dry_run:
        return True, preview, [], False

    backup_ok, backup_result, backup_errors, backup_replayed = create_backup(
        document_id=document_id,
        request_id=f"{request_id}:backup",
        dry_run=False,
        workspace=workspace,
    )
    if not backup_ok:
        return False, {"preview": preview, "backup": backup_result}, backup_errors, False

    write_result = guarded_com_mutation(
        path, backup_result,
        lambda: _run_spreadsheet_write_com(str(path), sheet_name, range_address, values),
    )
    if not write_result["ok"]:
        return False, {"preview": preview, "backup": backup_result, "write": write_result["data"]}, write_result["errors"], False

    read_ok, read_result, read_errors = read_spreadsheet_range(
        document_id,
        range_address,
        sheet_name=sheet_name,
        workspace=workspace,
    )
    if not read_ok:
        return False, {"preview": preview, "backup": backup_result}, read_errors, False
    validation_passed = read_result["values"] == values
    result = {
        **preview,
        "dry_run": False,
        "sheet": read_result["sheet"],
        "backup": backup_result,
        "backup_replayed": backup_replayed,
        "write_backend": write_result["data"].get("backend"),
        "read_back_values": read_result["values"],
        "validation_passed": validation_passed,
    }
    if not validation_passed:
        return (
            False,
            result,
            [{"code": "VALIDATION_FAILED", "message": "Read-back values did not match written values."}],
            False,
        )

    post_error = verify_post_com_source(path, write_result)
    if post_error is not None:
        return False, result, [post_error], False
    record_operation(request_id, "spreadsheet-write", result, workspace)
    return True, result, [], False


@coordinated_mutation
def write_spreadsheet_formulas(
    document_id: str,
    range_address: str,
    formulas: list[list[Any]],
    expected_values: list[list[Any]],
    request_id: str,
    sheet_name: str | None = None,
    dry_run: bool = False,
    workspace: str | Path = ".",
) -> tuple[bool, dict[str, Any], list[dict[str, str]], bool]:
    replay = replay_operation(
        request_id, "spreadsheet-formula-write",
        {"document_id": document_id, "range": range_address, "formulas": formulas, "expected_values": expected_values, "requested_sheet": sheet_name, "dry_run": dry_run},
        workspace,
    )
    if replay is not None:
        return replay

    try:
        row_count, column_count = _matrix_shape(formulas)
    except ValueError as exc:
        return False, {}, [{"code": "INVALID_ARGUMENT", "message": str(exc)}], False
    if row_count == 0 or column_count == 0:
        return False, {}, [{"code": "INVALID_ARGUMENT", "message": "formulas_json must not be empty."}], False
    if any(not isinstance(formula, str) or not formula.startswith("=") for row in formulas for formula in row):
        return (
            False,
            {},
            [{"code": "INVALID_ARGUMENT", "message": "Every formula must be a string starting with '='."}],
            False,
        )

    try:
        expected_rows, expected_cols = _matrix_shape(expected_values)
    except ValueError as exc:
        return False, {}, [{"code": "INVALID_ARGUMENT", "message": str(exc)}], False
    if (expected_rows, expected_cols) != (row_count, column_count):
        return (
            False,
            {
                "formulas_shape": [row_count, column_count],
                "expected_values_shape": [expected_rows, expected_cols],
            },
            [{"code": "EXPECTED_SHAPE_MISMATCH", "message": "expected_values_json shape must match formulas_json."}],
            False,
        )

    bounds, range_error = validate_spreadsheet_read_range(range_address)
    if range_error:
        return False, {}, [range_error], False
    min_col, min_row, max_col, max_row = bounds
    expected_range_rows = max_row - min_row + 1
    expected_range_cols = max_col - min_col + 1
    if (row_count, column_count) != (expected_range_rows, expected_range_cols):
        return (
            False,
            {
                "range": range_address,
                "range_shape": [expected_range_rows, expected_range_cols],
                "formulas_shape": [row_count, column_count],
            },
            [{"code": "RANGE_SHAPE_MISMATCH", "message": "formulas_json shape must match target range."}],
            False,
        )

    document = get_document(document_id, workspace)
    if not document:
        return False, {}, [{"code": "DOCUMENT_NOT_FOUND", "message": f"Document not registered: {document_id}"}], False
    if document["component"] != "spreadsheets":
        return (
            False,
            {"document": document},
            [{"code": "UNSUPPORTED_COMPONENT", "message": "spreadsheet-formula-write requires a spreadsheet document."}],
            False,
        )
    path = Path(document["path"])
    if not path.exists():
        return False, {"document": document}, [{"code": INPUT_FILE_NOT_FOUND, "message": f"Input file not found: {path}"}], False

    preview = {
        "document_id": document_id,
        "component": "spreadsheets",
        "path": str(path),
        "sheet": sheet_name,
        "requested_sheet": sheet_name,
        "range": range_address,
        "formulas": formulas,
        "expected_values": expected_values,
        "row_count": row_count,
        "column_count": column_count,
        "dry_run": dry_run,
    }
    if dry_run:
        return True, preview, [], False

    backup_ok, backup_result, backup_errors, backup_replayed = create_backup(
        document_id=document_id,
        request_id=f"{request_id}:backup",
        dry_run=False,
        workspace=workspace,
    )
    if not backup_ok:
        return False, {"preview": preview, "backup": backup_result}, backup_errors, False

    write_result = guarded_com_mutation(
        path, backup_result,
        lambda: _run_spreadsheet_formula_write_com(str(path), sheet_name, range_address, formulas),
    )
    if not write_result["ok"]:
        return False, {"preview": preview, "backup": backup_result, "write": write_result["data"]}, write_result["errors"], False

    read_ok, read_result, read_errors = read_spreadsheet_range(
        document_id,
        range_address,
        sheet_name=sheet_name,
        workspace=workspace,
    )
    if not read_ok:
        return False, {"preview": preview, "backup": backup_result}, read_errors, False

    validation_passed = read_result["values"] == expected_values
    result = {
        **preview,
        "dry_run": False,
        "sheet": read_result["sheet"],
        "backup": backup_result,
        "backup_replayed": backup_replayed,
        "write_backend": write_result["data"].get("backend"),
        "calculation_triggered": write_result["data"].get("calculation_triggered"),
        "read_back_values": read_result["values"],
        "validation_passed": validation_passed,
    }
    if not validation_passed:
        return (
            False,
            result,
            [{"code": "VALIDATION_FAILED", "message": "Calculated values did not match expected_values_json."}],
            False,
        )

    post_error = verify_post_com_source(path, write_result)
    if post_error is not None:
        return False, result, [post_error], False
    record_operation(request_id, "spreadsheet-formula-write", result, workspace)
    return True, result, [], False
