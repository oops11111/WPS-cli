from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUTPUT = Path("fixtures/phase3/writer_nested_scope_wps_fixture.docx")


def add_bookmark(paragraph, identifier: int, name: str, value: str) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(identifier))
    start.set(qn("w:name"), name)
    paragraph._p.append(start)
    paragraph.add_run(value)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(identifier))
    paragraph._p.append(end)


def main() -> None:
    if OUTPUT.exists():
        raise SystemExit(f"Fixture already exists: {OUTPUT}")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("Nested Scope Audit", level=1)
    body = document.add_paragraph("Body: ")
    add_bookmark(body, 101, "BodyMark", "BodyValue")
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0).paragraphs[0]
    cell.add_run("Cell: ")
    add_bookmark(cell, 102, "CellMark", "CellValue")
    header = document.sections[0].header.paragraphs[0]
    header.add_run("Header: ")
    add_bookmark(header, 103, "HeaderMark", "HeaderValue")
    footer = document.sections[0].footer.paragraphs[0]
    footer.add_run("Footer: ")
    add_bookmark(footer, 104, "FooterMark", "FooterValue")
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
