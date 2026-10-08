from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUTPUT = Path("fixtures/phase3/writer_mixed_run_wps_fixture.docx")


def main() -> None:
    if OUTPUT.exists():
        raise SystemExit(f"Fixture already exists: {OUTPUT}")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    paragraph = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    paragraph.add_run("Before ")

    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), "201")
    start.set(qn("w:name"), "MixedMark")
    paragraph._p.append(start)
    paragraph.add_run("A")

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), "target")
    run = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = "B"
    run.append(value)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    paragraph.add_run("\tC")

    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "201")
    paragraph._p.append(end)
    paragraph.add_run(" After")
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
