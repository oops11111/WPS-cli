from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUTPUT = Path("fixtures/phase3/writer_structure_wps_fixture.docx")


def main() -> None:
    if OUTPUT.exists():
        raise SystemExit(f"Fixture already exists: {OUTPUT}")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("Structure Audit", level=1)
    document.add_heading("Evidence Section", level=2)
    paragraph = document.add_paragraph()
    paragraph.add_run("Before ")
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), "42")
    start.set(qn("w:name"), "AuditMark")
    paragraph._p.append(start)
    paragraph.add_run("Alpha")
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "42")
    paragraph._p.append(end)
    paragraph.add_run(" after.")
    document.add_paragraph("Plain body text.")
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
