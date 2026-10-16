"""The workbook-close diagnostic (P8-005): what each attempt records, the results it writes, and its script."""

from __future__ import annotations

import importlib.util
import io
import json
import os
import shutil
import tempfile
import types
import unittest
from pathlib import Path
from unittest.mock import patch

from tests.test_undo_in_wps_powershell import FAKE_BOOK_PRELUDE
from tests.test_windows_check import FakeMachine, windows_check
from wps_ai_agent_cli import open_documents

REPO_ROOT = Path(__file__).resolve().parents[1]
_spec = importlib.util.spec_from_file_location(
    "workbook_close", REPO_ROOT / "scripts" / "windows_check" / "workbook_close.py"
)
workbook_close = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(workbook_close)
PWSH = os.environ.get("WPS_TEST_PWSH") or shutil.which("pwsh")

RECOVERY = {"name": "et.exe", "id": 9, "title": "", "command_line": r"et.exe recover1 -cpid=8 /w"}


class FakeRunner:
    WPS_LAUNCH = "default"

    def __init__(self, open_reason=None):
        self.open_reason = open_reason
        self.launches = []
        self.released = 0

    def write_fixture(self, workspace, name):
        (workspace / name).write_bytes(b"xlsx")
        return "sha"

    def open_in_wps(self, path, component):
        self.launches.append((self.WPS_LAUNCH, path.name, component))
        if self.open_reason:
            return None, self.open_reason

        def release():
            self.released += 1
            return "fixture closed; processes stopped 8"

        return release, None


def attempted(returned=True, error=None, **extra):
    return {
        "ok": True,
        "status": "attempted",
        "attempt": {
            "attached_via": "running_object_table",
            "in_rot": True,
            "returned": returned,
            "error": error,
            **extra,
        },
    }


class AttemptTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)

    def attempt(self, payload, listed, processes=(), runner=None, launch="default", method="close"):
        runner = runner or FakeRunner()
        with patch.object(open_documents, "_listed_in_wps", return_value=listed):
            result = workbook_close.run_attempt(
                launch,
                method,
                runner,
                self.root,
                lambda: list(processes),
                close=lambda path, name: payload,
                settle=0,
            )
        return result, runner

    def test_a_close_that_returns_and_leaves_nothing_listed_is_closed(self):
        result, runner = self.attempt(attempted(), listed=False)
        self.assertEqual(workbook_close.outcome(result), "closed")
        self.assertEqual(runner.launches, [("default", "close-check.xlsx", "spreadsheets")])
        self.assertEqual(runner.released, 1)
        self.assertEqual(result["cleanup"], "fixture closed; processes stopped 8")

    def test_a_failed_close_records_its_hresult_and_a_recovery_instance(self):
        error = {"hresult": "0x800A01A8", "type": "System.Runtime.InteropServices.COMException", "message": "x"}
        result, _ = self.attempt(attempted(returned=False, error=error), listed=True, processes=[RECOVERY])
        self.assertEqual(workbook_close.outcome(result), "failed, still open")
        self.assertEqual(workbook_close.recovery_processes(result["processes_after"]), ["et.exe 9"])
        row = workbook_close.results_table([result])[2]
        self.assertIn("| failed, still open | 0x800A01A8 |", row)
        self.assertIn("| et.exe 9 |", row)

    def test_a_failure_after_which_wps_lists_nothing_is_failed_but_closed(self):
        result, _ = self.attempt(attempted(returned=False, error={"hresult": "0x80010105"}), listed=False)
        self.assertEqual(workbook_close.outcome(result), "failed, closed anyway")

    def test_a_close_that_leaves_wps_unreachable_is_decided_by_the_scripts_own_look_at_workbooks(self):
        # aed06e6: once its last workbook closed, a double-clicked WPS answered WPS_INSTANCE_NOT_AUTOMATABLE.
        gone, _ = self.attempt(attempted(listed_after=False), listed=None)
        self.assertEqual(workbook_close.outcome(gone), "closed")
        still, _ = self.attempt(attempted(listed_after=True), listed=None)
        self.assertEqual(workbook_close.outcome(still), "returned, still open")
        failed, _ = self.attempt(
            attempted(returned=False, error={"hresult": "0x800A01A8"}, listed_after=None), listed=None
        )
        self.assertEqual(workbook_close.outcome(failed), "failed, still open")

    def test_a_workbook_wps_did_not_open_is_a_setup_failure_and_nothing_is_tried(self):
        closed = []
        with patch.object(open_documents, "_listed_in_wps") as listed:
            result = workbook_close.run_attempt(
                "automation",
                "close",
                FakeRunner(open_reason="WPS did not list close-check.xlsx"),
                self.root,
                list,
                close=lambda path, name: closed.append(name),
                settle=0,
            )
        self.assertEqual(closed, [])
        listed.assert_not_called()
        self.assertEqual(workbook_close.outcome(result), "setup failed")
        self.assertIn("WPS did not list", workbook_close.results_table([result])[2])

    def test_the_launch_mode_is_set_on_the_runner_before_opening(self):
        _result, runner = self.attempt(attempted(), listed=False, launch="automation")
        self.assertEqual(runner.launches[0][0], "automation")

    def test_a_write_then_close_writes_first_and_closes_with_the_base_method(self):
        steps = []
        with patch.object(open_documents, "_listed_in_wps", return_value=False):
            result = workbook_close.run_attempt(
                "default",
                "write_then_close_false",
                FakeRunner(),
                self.root,
                list,
                close=lambda path, name: steps.append(("close", name)) or attempted(),
                write=lambda path: steps.append(("write", path.name)) or {"status": "written", "saved": True},
                settle=0,
            )
        self.assertEqual(steps, [("write", "close-check.xlsx"), ("close", "close_false")])
        self.assertEqual(result["write"]["status"], "written")
        self.assertEqual(workbook_close.outcome(result), "closed")

    def test_a_failed_write_is_the_outcome_whatever_the_close_did(self):
        with patch.object(open_documents, "_listed_in_wps", return_value=True):
            result = workbook_close.run_attempt(
                "default",
                "write_then_close",
                FakeRunner(),
                self.root,
                list,
                close=lambda path, name: attempted(),
                write=lambda path: {"status": "script_failed"},
                settle=0,
            )
        self.assertEqual(workbook_close.outcome(result), "write failed (script_failed)")

    def run_restore(self, record):
        with patch.object(open_documents, "_listed_in_wps", return_value=False):
            return workbook_close.run_attempt(
                "default",
                workbook_close.RESTORE,
                FakeRunner(),
                self.root,
                list,
                close=lambda path, name: self.fail("the restore path never runs a close attempt"),
                restore=lambda path, workspace, write: record,
                settle=0,
            )

    def test_the_restore_attempt_is_restored_only_when_a1_is_the_backups_again(self):
        restored = self.run_restore({"step": "restore", "ok": True, "a1_after_restore": "seed", "error_codes": []})
        self.assertEqual(workbook_close.outcome(restored), "restored")
        refused = self.run_restore(
            {
                "step": "restore",
                "ok": False,
                "a1_after_restore": "written by close-diag",
                "error_codes": ["WPS_CLOSE_REFUSED"],
                "errors": [{"code": "WPS_CLOSE_REFUSED", "message": "WPS refused to close close-check.xlsx"}],
            }
        )
        self.assertEqual(workbook_close.outcome(refused), "restore failed: WPS_CLOSE_REFUSED")
        self.assertIn("WPS_CLOSE_REFUSED: WPS refused", workbook_close.results_table([refused])[2])
        early = self.run_restore({"step": "backup", "errors": []})
        self.assertEqual(workbook_close.outcome(early), "failed before the restore (backup)")

    def test_restore_with_close_backs_up_writes_and_restores_through_the_product(self):
        from openpyxl import Workbook, load_workbook

        book = self.root / "close-check.xlsx"
        workbook = Workbook()
        workbook.active["A1"] = "seed"
        workbook.save(book)

        def write(path):
            changed = load_workbook(path)
            changed.active["A1"] = workbook_close.WRITTEN_TEXT
            changed.save(path)
            return {"status": "written"}

        closed = {"ok": True, "errors": [], "data": {"status": "closed"}}
        with patch.object(open_documents, "_run_component_attached_com", return_value=closed) as close:
            record = workbook_close.restore_with_close(book, self.root, write)
        self.assertEqual(close.call_args.args[:2], ("spreadsheets", "close-unchanged"))
        self.assertEqual(record["a1_after_write"], workbook_close.WRITTEN_TEXT)
        self.assertTrue(record["ok"], record)
        self.assertEqual(record["close_in_wps"], "closed")
        self.assertEqual(record["a1_after_restore"], "seed")

    def test_the_default_plan_tries_every_method_and_a_control_group(self):
        plan = workbook_close.default_plan()
        self.assertEqual(
            {method for launch, method in plan if launch == "default"},
            {*workbook_close.METHODS, *workbook_close.AFTER_WRITE, workbook_close.RESTORE},
        )
        self.assertEqual(set(workbook_close.DESCRIPTIONS), {method for _launch, method in plan})
        self.assertEqual([method for launch, method in plan if launch == "automation"], ["close", "close_false"])
        parser_choices = {"close", "close_false", "close_false_missing", "saved_then_close", "item_by_name"}
        self.assertLessEqual(parser_choices, set(workbook_close.METHODS))


class CloseDiagCommandTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.out = Path(tmp.name) / "results"
        for patcher in (
            patch("sys.stdout", new_callable=io.StringIO),
            patch("sys.stderr", new_callable=io.StringIO),
        ):
            patcher.start()
            self.addCleanup(patcher.stop)

    def diagnostic(self, ran):
        def run_attempt(launch, method, runner, root, processes):
            ran.append((launch, method))
            failed = method != "close" or launch == "automation"
            payload = attempted(
                returned=not failed, error={"hresult": "0x800A01A8", "message": "m"} if failed else None
            )
            return {
                "launch": launch,
                "method": method,
                **payload,
                "listed_by_cli": failed,
                "processes_after": [RECOVERY] if failed else [],
                "cleanup": "ok",
            }

        return types.SimpleNamespace(**{**vars(workbook_close), "run_attempt": run_attempt})

    def test_every_attempt_is_written_to_the_results(self):
        ran = []
        code = windows_check.main(
            ["close-diag", "--out", str(self.out)],
            machine=FakeMachine(),
            runner=object(),
            diagnostic=self.diagnostic(ran),
        )
        self.assertEqual(code, 0)
        self.assertEqual(ran, workbook_close.default_plan())
        text = (self.out / "RESULTS.md").read_text(encoding="utf-8")
        self.assertIn("| default | `close` | running_object_table | True | closed |", text)
        self.assertIn("| automation | `close_false` |", text)
        self.assertIn("0x800A01A8", text)
        self.assertIn("et.exe recover1 -cpid=8 /w", text)
        state = json.loads((self.out / "results.json").read_text(encoding="utf-8"))
        self.assertEqual(len(state["results"]), len(ran))

    def test_method_and_launch_narrow_the_plan(self):
        ran = []
        windows_check.main(
            [
                "close-diag",
                "--out",
                str(self.out),
                "--method",
                "close",
                "--method",
                "item_by_name",
                "--launch",
                "default",
            ],
            machine=FakeMachine(),
            runner=object(),
            diagnostic=self.diagnostic(ran),
        )
        self.assertEqual(ran, [("default", "close"), ("default", "item_by_name")])
        text = (self.out / "RESULTS.md").read_text(encoding="utf-8")
        self.assertIn("close-diag --method close --method item_by_name --launch default", text)

    def test_failed_prerequisites_run_nothing(self):
        ran = []
        code = windows_check.main(
            ["close-diag", "--out", str(self.out)],
            machine=FakeMachine(processes=[{"name": "et.exe", "id": 1, "title": "x", "command_line": "et.exe"}]),
            runner=object(),
            diagnostic=self.diagnostic(ran),
        )
        self.assertEqual(code, 1)
        self.assertEqual(ran, [])


@unittest.skipUnless(PWSH, "PowerShell (pwsh) is required to run the close attempts against a fake WPS")
class CloseAttemptScriptTests(unittest.TestCase):
    """The fake workbook behaves like the one in P7-007: Close() works, Close($false) faults."""

    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.book = Path(tmp.name) / "close-check.xlsx"
        self.book.write_bytes(b"xlsx")
        for patcher in (
            patch("wps_ai_agent_cli.powershell_runner.powershell_executable", return_value=PWSH),
            patch.dict(os.environ, {"WPSFAKE_PATH": str(self.book.resolve()), "WPSFAKE_SAVES": "1"}),
            patch.object(
                open_documents,
                "_render",
                side_effect=lambda script, values: (
                    FAKE_BOOK_PRELUDE + open_documents.json.dumps(values).join(script.split("__PARAMS_JSON__"))
                ),
            ),
        ):
            patcher.start()
            self.addCleanup(patcher.stop)

    def test_a_close_that_returns_is_recorded_with_how_the_script_attached(self):
        payload = workbook_close.run_close_script(self.book, "close")
        self.assertEqual(payload["status"], "attempted", payload)
        attempt = payload["attempt"]
        self.assertTrue(attempt["returned"])
        self.assertIsNone(attempt["error"])
        self.assertEqual(attempt["attached_via"], "active_object")
        self.assertTrue(attempt["saved_before"])
        # The fake keeps listing the workbook, so the check afterwards sees it.
        self.assertTrue(attempt["listed_after"])

    def test_a_faulting_close_records_the_innermost_hresult(self):
        attempt = workbook_close.run_close_script(self.book, "close_false")["attempt"]
        self.assertFalse(attempt["returned"])
        self.assertEqual(attempt["error"]["hresult"], "0x80010105")
        self.assertEqual(attempt["error"]["type"], "System.Runtime.InteropServices.COMException")

    def test_a_workbook_missing_from_the_running_object_table_is_an_error_not_a_crash(self):
        attempt = workbook_close.run_close_script(self.book, "rot_object")["attempt"]
        self.assertFalse(attempt["in_rot"])
        self.assertFalse(attempt["returned"])
        self.assertIn("Running Object Table", attempt["error"]["message"])

    def test_the_write_script_parses(self):
        import subprocess

        from wps_ai_agent_cli.open_document_scripts.spreadsheet import _ATTACHED_SPREADSHEET_WRAPPER

        path = self.book.with_name("write.ps1")
        path.write_text(_ATTACHED_SPREADSHEET_WRAPPER.replace("__ACTION_BODY__", workbook_close._WRITE_BODY))
        check = (
            "$errors = $null; [void][System.Management.Automation.Language.Parser]::ParseFile("
            f"'{path}', [ref]$null, [ref]$errors); $errors.Count"
        )
        completed = subprocess.run(
            [PWSH, "-NoProfile", "-Command", check], capture_output=True, text=True, timeout=120, check=False
        )
        self.assertEqual(completed.stdout.strip(), "0", completed.stdout + completed.stderr)

    def test_every_method_runs_to_a_record(self):
        for method in workbook_close.METHODS:
            with self.subTest(method=method):
                payload = workbook_close.run_close_script(self.book, method)
                self.assertEqual(payload["status"], "attempted", payload)
                self.assertEqual(payload["attempt"]["method"], method)


if __name__ == "__main__":
    unittest.main()
