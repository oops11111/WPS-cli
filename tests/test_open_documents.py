import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from wps_ai_agent_cli import cli, open_documents
from openpyxl import Workbook

from wps_ai_agent_cli.open_documents import (
    delete_presentation_selection,
    delete_spreadsheet_selection,
    delete_writer_selection,
    export_open_document_copy,
    export_open_document_html,
    export_open_document_pdf,
    export_open_document_text,
    insert_presentation_selection,
    insert_writer_selection,
    save_as_open_document,
    save_open_document,
    list_open_documents,
    read_presentation_selection,
    read_spreadsheet_selection,
    read_writer_selection,
    replace_presentation_selection,
    replace_spreadsheet_selection,
    replace_writer_selection,
)
from wps_ai_agent_cli.sessions import register_document, stable_document_id
from tests.test_presentation_ops import write_minimal_pptx
from tests.test_writer_ops import write_minimal_docx_with_table


def com_ok(data):
    return {"ok": True, "errors": [], "data": data}


def selection_payload(text="PENDING_STATUS", start=10, end=None, saved=True, **extra):
    return {
        "ok": True,
        "status": "ok",
        "saved": saved,
        "start": start,
        "end": start + len(text) if end is None else end,
        "selection_text": text,
        "paragraph_index": 3,
        "page_number": 1,
        **extra,
    }


class Fixture:
    def __init__(self, test, component="writer"):
        self.tmp = tempfile.TemporaryDirectory()
        test.addCleanup(self.tmp.cleanup)
        self.workspace = Path(self.tmp.name) / "workspace"
        self.workspace.mkdir()
        self.component = component
        if component == "presentation":
            self.document = Path(self.tmp.name) / "open.pptx"
            write_minimal_pptx(self.document, ["P3-399 export"])
        elif component == "spreadsheets":
            self.document = Path(self.tmp.name) / "open.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "P3-402 mark"
            workbook.save(self.document)
            workbook.close()
        else:
            self.document = Path(self.tmp.name) / "open.docx"
            write_minimal_docx_with_table(self.document)
        ok, record, _errors = register_document(component, str(self.document), self.workspace)
        assert ok
        self.document_id = record["document_id"]


class ListOpenDocumentsTests(unittest.TestCase):
    def test_reports_missing_instance_without_launching_anything(self):
        with patch.object(open_documents, "_run_open_documents_com", return_value=com_ok({"ok": True, "instances": []})) as com:
            ok, data, errors = list_open_documents()
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "NO_RUNNING_WPS_INSTANCE")
        self.assertEqual(data["probed_components"], ["writer", "spreadsheets", "presentation"])
        com.assert_called_once()

    def test_lists_documents_with_registered_id_and_untitled_documents(self):
        fixture = Fixture(self)
        payload = {
            "ok": True,
            "instances": {
                "component": "writer",
                "version": "12.1",
                "documents": [
                    {"name": "open.docx", "full_name": str(fixture.document), "saved": True, "read_only": False, "active": True},
                    {"name": "Document1", "full_name": "Document1", "saved": False, "read_only": False, "active": False},
                ],
            },
        }
        with patch.object(open_documents, "_run_open_documents_com", return_value=com_ok(payload)):
            ok, data, errors = list_open_documents(component="writer", workspace=fixture.workspace)
        self.assertTrue(ok, errors)
        self.assertEqual(data["instances"], [{"component": "writer", "version": "12.1", "document_count": 2}])
        saved, untitled = data["documents"]
        self.assertEqual(saved["document_id"], fixture.document_id)
        self.assertTrue(saved["registered"])
        self.assertTrue(saved["active"])
        self.assertIsNone(untitled["document_id"])
        self.assertFalse(untitled["path_exists"])

    def test_register_flag_registers_saved_documents_only(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        workspace = Path(tmp.name) / "workspace"
        workspace.mkdir()
        document = Path(tmp.name) / "fresh.docx"
        write_minimal_docx_with_table(document)
        payload = {
            "ok": True,
            "instances": [{
                "component": "writer",
                "documents": [{"name": "fresh.docx", "full_name": str(document), "saved": True, "read_only": False, "active": False}],
            }],
        }
        with patch.object(open_documents, "_run_open_documents_com", return_value=com_ok(payload)):
            _ok, without, _errors = list_open_documents(workspace=workspace)
            self.assertIsNone(without["documents"][0]["document_id"])
            ok, data, _errors = list_open_documents(register=True, workspace=workspace)
        self.assertTrue(ok)
        self.assertEqual(data["documents"][0]["document_id"], stable_document_id("writer", str(document)))

    def test_rejects_unknown_component_before_touching_wps(self):
        with patch.object(open_documents, "_run_open_documents_com") as com:
            ok, _data, errors = list_open_documents(component="outlook")
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "UNSUPPORTED_COMPONENT")
        com.assert_not_called()


class SelectionReadTests(unittest.TestCase):
    def test_reads_selection_and_flags_page_as_unverified(self):
        fixture = Fixture(self)
        with patch.object(open_documents, "_run_writer_attached_com", return_value=com_ok(selection_payload())):
            ok, data, errors = read_writer_selection(fixture.document_id, fixture.workspace)
        self.assertTrue(ok, errors)
        self.assertEqual(data["selection_text"], "PENDING_STATUS")
        self.assertFalse(data["collapsed"])
        self.assertEqual(data["paragraph_index"], 3)
        self.assertFalse(data["page_info_verified"])

    def test_truncates_long_selection_text(self):
        fixture = Fixture(self)
        payload = selection_payload(text="x" * 5000)
        with patch.object(open_documents, "_run_writer_attached_com", return_value=com_ok(payload)):
            ok, data, _errors = read_writer_selection(fixture.document_id, fixture.workspace)
        self.assertTrue(ok)
        self.assertEqual(len(data["selection_text"]), 4096)
        self.assertTrue(data["selection_text_truncated"])
        self.assertEqual(data["selection_length"], 5000)

    def test_status_codes_map_to_structured_errors(self):
        fixture = Fixture(self)
        cases = {
            "no_instance": "NO_RUNNING_WPS_INSTANCE",
            "not_open": "DOCUMENT_NOT_OPEN",
            "no_window": "SELECTION_UNAVAILABLE",
        }
        for status, code in cases.items():
            with self.subTest(status=status):
                with patch.object(open_documents, "_run_writer_attached_com", return_value=com_ok({"ok": True, "status": status})):
                    ok, _data, errors = read_writer_selection(fixture.document_id, fixture.workspace)
                self.assertFalse(ok)
                self.assertEqual(errors[0]["code"], code)

    def test_unregistered_and_wrong_component_documents_are_rejected(self):
        fixture = Fixture(self)
        ok, _data, errors = read_writer_selection("doc_missing", fixture.workspace)
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "DOCUMENT_NOT_FOUND")

        sheet = Path(fixture.tmp.name) / "book.xlsx"
        sheet.write_bytes(b"x")
        _ok, record, _errors = register_document("spreadsheets", str(sheet), fixture.workspace)
        ok, _data, errors = read_writer_selection(record["document_id"], fixture.workspace)
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "UNSUPPORTED_COMPONENT")

    def test_com_timeout_is_returned_as_structured_error(self):
        fixture = Fixture(self)
        failure = {"ok": False, "errors": [{"code": "COM_OPERATION_TIMEOUT", "message": "timed out"}], "data": {"timed_out": True}}
        with patch.object(open_documents, "_run_writer_attached_com", return_value=failure):
            ok, data, errors = read_writer_selection(fixture.document_id, fixture.workspace)
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "COM_OPERATION_TIMEOUT")
        self.assertTrue(data["timed_out"])


class SelectionReplaceTests(unittest.TestCase):
    def replace(self, fixture, text="APPROVED", request_id="sel-1", **kwargs):
        return replace_writer_selection(
            document_id=fixture.document_id, text=text, request_id=request_id, workspace=fixture.workspace, **kwargs,
        )

    def fake_attached(self, fixture, selection, write_status="ok", read_back=None, calls=None):
        def fake(action, params):
            if calls is not None:
                calls.append(action)
            if action == "selection-read":
                return com_ok(selection)
            self.assertEqual(action, "selection-replace")
            fixture.document.write_bytes(fixture.document.read_bytes() + b"\0")
            if write_status != "ok":
                return com_ok({"ok": True, "status": write_status})
            return com_ok({
                "ok": True,
                "status": "ok",
                "backend": "fake-wps",
                "read_back_text": params["text"] if read_back is None else read_back,
                "formatting_mode": "uniform_character",
                "formatting_preserved": True,
            })
        return fake

    def test_dry_run_previews_without_backup_or_write(self):
        fixture = Fixture(self)
        calls = []
        with patch.object(open_documents, "_run_writer_attached_com", side_effect=self.fake_attached(fixture, selection_payload(), calls=calls)):
            ok, data, errors, replayed = self.replace(fixture, dry_run=True, expected_selection_text="PENDING_STATUS")
        self.assertTrue(ok, errors)
        self.assertFalse(replayed)
        self.assertEqual(calls, ["selection-read"])
        self.assertTrue(data["would_modify"])
        self.assertEqual(data["selection_text"], "PENDING_STATUS")
        self.assertFalse((fixture.workspace / ".wps-agent" / "backups").exists())

    def test_replaces_with_backup_validation_and_replay(self):
        fixture = Fixture(self)
        calls = []
        original = fixture.document.read_bytes()
        with patch.object(open_documents, "_run_writer_attached_com", side_effect=self.fake_attached(fixture, selection_payload(), calls=calls)):
            ok, data, errors, _replayed = self.replace(fixture, expected_selection_text="PENDING_STATUS")
            self.assertTrue(ok, errors)
            self.assertTrue(data["validation_passed"])
            self.assertTrue(data["backup"]["created"])
            self.assertEqual(Path(data["backup"]["backup_path"]).read_bytes(), original)
            self.assertEqual(data["write_backend"], "fake-wps")
            self.assertFalse(data["selection_preserved"])
            self.assertEqual(data["formatting_mode"], "uniform_character")
            self.assertTrue(data["formatting_preserved"])
            self.assertIn("uniform character formatting", data["note"])

            ok, again, errors, replayed = self.replace(fixture, expected_selection_text="PENDING_STATUS")
            self.assertTrue(ok, errors)
            self.assertTrue(replayed)
            self.assertEqual(again["read_back_text"], "APPROVED")
        self.assertEqual(calls.count("selection-replace"), 1)

    def test_selection_replace_script_snapshots_font(self):
        captured = {}

        def fake_run(script, timeout_seconds=120, **_kwargs):
            captured["script"] = script
            return {"ok": True, "status": "no_instance"}, None

        with patch.object(open_documents, "run_powershell_script", side_effect=fake_run):
            open_documents._run_writer_attached_com(
                "selection-replace",
                {"path": "C:\\docs\\a.docx", "text": "NEW", "expected_text": "OLD", "expected_start": 0, "expected_end": 3},
            )
        script = captured["script"]
        self.assertIn("$range.Font", script)
        self.assertIn("formatting_preserved", script)
        self.assertIn("wdUndefined", script)
        self.assertNotIn(".Quit()", script)

    def test_rejections_happen_before_backup_and_write(self):
        fixture = Fixture(self)
        cases = [
            ("unsaved", {"ok": True, "status": "unsaved"}, {}, "DOCUMENT_HAS_UNSAVED_CHANGES"),
            ("unsaved-flag", selection_payload(saved=False), {}, "DOCUMENT_HAS_UNSAVED_CHANGES"),
            ("read-only", selection_payload(read_only=True), {}, "DOCUMENT_READ_ONLY"),
            ("empty", selection_payload(text="", start=5, end=5), {}, "SELECTION_EMPTY"),
            ("too-large", selection_payload(text="x" * 5000), {}, "SELECTION_TOO_LARGE"),
            ("paragraph-mark", selection_payload(text="line\r", start=4), {}, "SELECTION_SPANS_STRUCTURE"),
            ("object", selection_payload(text="ab", start=4, end=9), {}, "SELECTION_SPANS_STRUCTURE"),
            ("mismatch", selection_payload(), {"expected_selection_text": "something else"}, "SELECTION_MISMATCH"),
        ]
        for name, payload, kwargs, code in cases:
            with self.subTest(name=name):
                calls = []
                with patch.object(open_documents, "_run_writer_attached_com", side_effect=self.fake_attached(fixture, payload, calls=calls)):
                    ok, _data, errors, _replayed = self.replace(fixture, request_id=f"sel-{name}", **kwargs)
                self.assertFalse(ok)
                self.assertEqual(errors[0]["code"], code)
                self.assertNotIn("selection-replace", calls)
        self.assertFalse((fixture.workspace / ".wps-agent" / "backups").exists())

    def test_invalid_replacement_text_is_rejected_before_attaching(self):
        fixture = Fixture(self)
        with patch.object(open_documents, "_run_writer_attached_com") as com:
            for index, value in enumerate(["a\nb", "a\x00b", "x" * 4097]):
                ok, _data, errors, _replayed = self.replace(fixture, text=value, request_id=f"sel-bad-{index}")
                self.assertFalse(ok)
                self.assertEqual(errors[0]["code"], "INVALID_ARGUMENT")
        com.assert_not_called()

    def test_selection_changed_between_preview_and_write_is_reported(self):
        fixture = Fixture(self)
        with patch.object(open_documents, "_run_writer_attached_com", side_effect=self.fake_attached(fixture, selection_payload(), write_status="selection_changed")):
            ok, data, errors, _replayed = self.replace(fixture, request_id="sel-race")
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "SELECTION_MISMATCH")
        self.assertTrue(data["backup"]["created"])

    def test_read_back_mismatch_fails_validation_and_is_not_recorded(self):
        fixture = Fixture(self)
        with patch.object(open_documents, "_run_writer_attached_com", side_effect=self.fake_attached(fixture, selection_payload(), read_back="garbled")):
            ok, data, errors, _replayed = self.replace(fixture, request_id="sel-bad-readback")
            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "VALIDATION_FAILED")
            self.assertFalse(data["validation_passed"])
        from wps_ai_agent_cli.operations import get_operation

        self.assertIsNone(get_operation("sel-bad-readback", fixture.workspace))

    def test_com_timeout_during_write_leaves_structured_error(self):
        import subprocess

        fixture = Fixture(self)

        def fake(action, params):
            if action == "selection-read":
                return com_ok(selection_payload())
            raise subprocess.TimeoutExpired("powershell", 120)

        with patch.object(open_documents, "_run_writer_attached_com", side_effect=fake):
            ok, _data, errors, _replayed = self.replace(fixture, request_id="sel-timeout")
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "COM_OPERATION_TIMEOUT")


class ExportOpenDocumentTests(unittest.TestCase):
    def export(self, fixture, output, request_id="exp-1"):
        return export_open_document_html(
            document_id=fixture.document_id, output=str(output), request_id=request_id, workspace=fixture.workspace,
        )

    def fake_export(self, output, status="ok", write=True, touch_source=None, still_open=True, method="SaveAs2"):
        def fake(action, params):
            self.assertEqual(action, "export-html")
            if write:
                Path(params["output"]).write_text("<html></html>", encoding="utf-8")
                Path(params["output"]).with_name(Path(params["output"]).stem + ".files").mkdir()
            if touch_source is not None:
                touch_source.write_bytes(touch_source.read_bytes() + b"\0")
            return com_ok({"ok": True, "status": status, "export_method": method, "source_still_open": still_open})
        return fake

    def test_exports_html_and_verifies_source_unchanged(self):
        fixture = Fixture(self)
        output = Path(fixture.tmp.name) / "out.html"
        with patch.object(open_documents, "_run_writer_attached_com", side_effect=self.fake_export(output)):
            ok, data, errors, _replayed = self.export(fixture, output)
            self.assertTrue(ok, errors)
            self.assertTrue(data["source_unchanged"])
            self.assertTrue(data["source_still_open"])
            self.assertTrue(data["validation_passed"])
            self.assertEqual(data["format"], "html")
            self.assertEqual(data["component"], "writer")
            self.assertEqual(len(data["companion_directories"]), 1)
            self.assertTrue(data["known_losses"])

            ok, again, _errors, replayed = self.export(fixture, output)
            self.assertTrue(ok)
            self.assertTrue(replayed)
            self.assertEqual(again["output"], data["output"])

    def test_exports_open_presentation_via_save_copy_as(self):
        fixture = Fixture(self, component="presentation")
        output = Path(fixture.tmp.name) / "deck.html"
        before = fixture.document.read_bytes()
        with patch.object(
            open_documents,
            "_run_presentation_attached_com",
            side_effect=self.fake_export(output, method="SaveCopyAsHtml"),
        ) as presentation_com, patch.object(open_documents, "_run_writer_attached_com") as writer_com:
            ok, data, errors, _replayed = self.export(fixture, output, request_id="exp-pptx")
        self.assertTrue(ok, errors)
        self.assertEqual(data["component"], "presentation")
        self.assertEqual(data["export_method"], "SaveCopyAsHtml")
        self.assertTrue(data["source_unchanged"])
        self.assertTrue(data["validation_passed"])
        self.assertEqual(fixture.document.read_bytes(), before)
        presentation_com.assert_called_once()
        writer_com.assert_not_called()

    def test_presentation_export_rejects_zip_bytes_masquerading_as_html(self):
        fixture = Fixture(self, component="presentation")
        output = Path(fixture.tmp.name) / "deck.html"

        def fake(action, params):
            Path(params["output"]).write_bytes(b"PK\x03\x04not-html")
            return com_ok({
                "ok": True,
                "status": "ok",
                "export_method": "SaveCopyAsHtml",
                "source_still_open": True,
            })

        with patch.object(open_documents, "_run_presentation_attached_com", side_effect=fake):
            ok, data, errors, _replayed = self.export(fixture, output, request_id="exp-zip")
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "EXPORT_OUTPUT_NOT_HTML")
        self.assertTrue(data["source_unchanged"])

    def test_exports_open_spreadsheet_via_save_copy_as(self):
        fixture = Fixture(self, component="spreadsheets")
        output = Path(fixture.tmp.name) / "sheet.html"
        before = fixture.document.read_bytes()
        with patch.object(
            open_documents,
            "_run_spreadsheet_attached_com",
            side_effect=self.fake_export(output, method="SaveCopyAsThenSaveAsHtml"),
        ) as spreadsheet_com, patch.object(open_documents, "_run_writer_attached_com") as writer_com, \
                patch.object(open_documents, "_run_presentation_attached_com") as presentation_com:
            ok, data, errors, _replayed = self.export(fixture, output, request_id="exp-xlsx")
        self.assertTrue(ok, errors)
        self.assertEqual(data["component"], "spreadsheets")
        self.assertEqual(data["export_method"], "SaveCopyAsThenSaveAsHtml")
        self.assertTrue(data["source_unchanged"])
        self.assertTrue(data["validation_passed"])
        self.assertEqual(fixture.document.read_bytes(), before)
        spreadsheet_com.assert_called_once()
        writer_com.assert_not_called()
        presentation_com.assert_not_called()

    def test_spreadsheet_export_rejects_zip_bytes_masquerading_as_html(self):
        fixture = Fixture(self, component="spreadsheets")
        output = Path(fixture.tmp.name) / "sheet.html"

        def fake(action, params):
            Path(params["output"]).write_bytes(b"PK\x03\x04not-html")
            return com_ok({
                "ok": True,
                "status": "ok",
                "export_method": "SaveCopyAsHtml",
                "source_still_open": True,
            })

        with patch.object(open_documents, "_run_spreadsheet_attached_com", side_effect=fake):
            ok, data, errors, _replayed = self.export(fixture, output, request_id="exp-xlsx-zip")
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "EXPORT_OUTPUT_NOT_HTML")
        self.assertTrue(data["source_unchanged"])

    def test_spreadsheet_export_script_uses_temp_copy_and_never_quits(self):
        captured = {}

        def fake_run(script, timeout_seconds=120, **_kwargs):
            captured["script"] = script
            return {"ok": True, "status": "no_instance"}, None

        with patch.object(open_documents, "run_powershell_script", side_effect=fake_run):
            result = open_documents._run_spreadsheet_attached_com(
                "export-html",
                {"path": "C:\\docs\\book.xlsx", "output": "C:\\out\\book.html"},
            )
        script = captured["script"]
        self.assertTrue(result["ok"])
        self.assertIn("SaveCopyAs", script)
        self.assertIn("$copy.SaveAs", script)
        self.assertIn("Find-OpenWorkbook", script)
        self.assertNotIn(".Quit()", script)
        self.assertNotIn("$doc.SaveAs", script)

    def test_invalid_outputs_are_rejected_without_calling_wps(self):
        fixture = Fixture(self)
        existing = Path(fixture.tmp.name) / "exists.html"
        existing.write_text("keep", encoding="utf-8")
        cases = [
            (fixture.document, "INVALID_ARGUMENT"),
            (Path(fixture.tmp.name) / "out.pdf", "INVALID_ARGUMENT"),
            (existing, "OUTPUT_EXISTS"),
            (Path(fixture.tmp.name) / "missing-dir" / "out.html", "INVALID_ARGUMENT"),
        ]
        with patch.object(open_documents, "_run_writer_attached_com") as com:
            for index, (output, code) in enumerate(cases):
                with self.subTest(output=str(output)):
                    ok, _data, errors, _replayed = self.export(fixture, output, request_id=f"exp-bad-{index}")
                    self.assertFalse(ok)
                    self.assertEqual(errors[0]["code"], code)
        com.assert_not_called()
        self.assertEqual(existing.read_text(encoding="utf-8"), "keep")

    def test_unsaved_or_closed_documents_are_reported(self):
        fixture = Fixture(self)
        for status, code in (("unsaved", "DOCUMENT_HAS_UNSAVED_CHANGES"), ("not_open", "DOCUMENT_NOT_OPEN"), ("no_instance", "NO_RUNNING_WPS_INSTANCE")):
            with self.subTest(status=status):
                output = Path(fixture.tmp.name) / f"{status}.html"
                with patch.object(open_documents, "_run_writer_attached_com", side_effect=self.fake_export(output, status=status, write=False)):
                    ok, _data, errors, _replayed = self.export(fixture, output, request_id=f"exp-{status}")
                self.assertFalse(ok)
                self.assertEqual(errors[0]["code"], code)
                if status == "unsaved":
                    self.assertIn("save-open-document", errors[0]["message"])
                    self.assertIn("exported", errors[0]["message"])
                    self.assertNotIn("save it in WPS first", errors[0]["message"])

    def test_missing_output_and_changed_source_are_failures(self):
        fixture = Fixture(self)
        missing = Path(fixture.tmp.name) / "missing.html"
        with patch.object(open_documents, "_run_writer_attached_com", side_effect=self.fake_export(missing, write=False)):
            ok, _data, errors, _replayed = self.export(fixture, missing, request_id="exp-missing")
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "EXPORT_OUTPUT_MISSING")

        changed = Path(fixture.tmp.name) / "changed.html"
        with patch.object(open_documents, "_run_writer_attached_com", side_effect=self.fake_export(changed, touch_source=fixture.document)):
            ok, data, errors, _replayed = self.export(fixture, changed, request_id="exp-changed")
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "DOCUMENT_CHANGED_AFTER_COM")
        self.assertFalse(data["source_unchanged"])


class ScriptRenderingTests(unittest.TestCase):
    def test_rendered_scripts_have_no_placeholders_and_never_quit_wps(self):
        for action in ("selection-read", "selection-replace", "export-html"):
            with self.subTest(action=action):
                captured = {}

                def fake_run(script, timeout_seconds=120, **_kwargs):
                    captured["script"] = script
                    return {"ok": True, "status": "no_instance"}, None

                with patch.object(open_documents, "run_powershell_script", side_effect=fake_run):
                    result = open_documents._run_writer_attached_com(action, {"path": "C:\\docs\\a.docx", "text": "it's '@ tricky"})
                script = captured["script"]
                self.assertTrue(result["ok"])
                self.assertNotIn("__PARAMS_JSON__", script)
                self.assertNotIn("__ACTION_BODY__", script)
                self.assertNotIn(".Quit()", script)
                self.assertIn('"prog_ids"', script)
                self.assertIn("\n'@ | ConvertFrom-Json", script)
                if action == "export-html":
                    self.assertNotIn("status = 'read_only'", script)
                    self.assertIn("Documents.Add", script)
                if action == "selection-replace":
                    self.assertIn("status = 'read_only'", script)

    def test_presentation_export_script_uses_save_copy_as_and_never_quits(self):
        captured = {}

        def fake_run(script, timeout_seconds=120, **_kwargs):
            captured["script"] = script
            return {"ok": True, "status": "no_instance"}, None

        with patch.object(open_documents, "run_powershell_script", side_effect=fake_run):
            result = open_documents._run_presentation_attached_com(
                "export-html",
                {"path": "C:\\docs\\deck.pptx", "output": "C:\\docs\\deck.html"},
            )
        script = captured["script"]
        self.assertTrue(result["ok"])
        self.assertIn("Find-OpenPresentation", script)
        self.assertIn("SaveCopyAs([string]$params.output, 12)", script)
        self.assertNotIn("Documents.Add", script)
        self.assertNotIn(".Quit()", script)
        self.assertNotIn("__ACTION_BODY__", script)

    def test_presentation_selection_read_script_never_quits(self):
        captured = {}

        def fake_run(script, timeout_seconds=120, **_kwargs):
            captured["script"] = script
            return {"ok": True, "status": "no_instance"}, None

        with patch.object(open_documents, "run_powershell_script", side_effect=fake_run):
            result = open_documents._run_presentation_attached_com(
                "selection-read",
                {"path": "C:\\docs\\deck.pptx"},
            )
        script = captured["script"]
        self.assertTrue(result["ok"])
        self.assertIn("selection_kind", script)
        self.assertIn("Find-OpenPresentation", script)
        self.assertNotIn(".Quit()", script)

    def test_list_script_probes_every_component_prog_id(self):
        captured = {}

        def fake_run(script, timeout_seconds=120, **_kwargs):
            captured["script"] = script
            return {"ok": True, "instances": []}, None

        with patch.object(open_documents, "run_powershell_script", side_effect=fake_run):
            open_documents._run_open_documents_com(["writer", "presentation"])
        script = captured["script"]
        self.assertIn("kwps.Application", script)
        self.assertIn("kwpp.Application", script)
        self.assertNotIn("ket.Application", script)
        self.assertNotIn(".Quit()", script)


class PresentationSelectionReadTests(unittest.TestCase):
    def test_reads_text_selection_metadata(self):
        fixture = Fixture(self, component="presentation")
        payload = {
            "ok": True,
            "status": "ok",
            "saved": True,
            "read_only": False,
            "selection_type": 3,
            "selection_kind": "text",
            "selection_text": "P3-400 mark",
            "text_range_start": 1,
            "text_range_length": 11,
            "collapsed": False,
            "slide_index": 1,
            "shape_name": "Title 1",
            "shape_count": 0,
        }
        with patch.object(open_documents, "_run_presentation_attached_com", return_value=com_ok(payload)):
            ok, data, errors = read_presentation_selection(fixture.document_id, fixture.workspace)
        self.assertTrue(ok, errors)
        self.assertEqual(data["selection_kind"], "text")
        self.assertEqual(data["selection_text"], "P3-400 mark")
        self.assertEqual(data["text_range_start"], 1)
        self.assertEqual(data["text_range_length"], 11)
        self.assertFalse(data["collapsed"])
        self.assertEqual(data["slide_index"], 1)
        self.assertFalse(data["slide_info_verified"])
        self.assertEqual(data["component"], "presentation")

    def test_rejects_writer_document_ids(self):
        fixture = Fixture(self, component="writer")
        with patch.object(open_documents, "_run_presentation_attached_com") as com:
            ok, _data, errors = read_presentation_selection(fixture.document_id, fixture.workspace)
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "UNSUPPORTED_COMPONENT")
        com.assert_not_called()


class SpreadsheetSelectionReadTests(unittest.TestCase):
    def test_selection_read_script_never_quits(self):
        captured = {}

        def fake_run(script, timeout_seconds=120, **_kwargs):
            captured["script"] = script
            return {"ok": True, "status": "no_instance"}, None

        with patch.object(open_documents, "run_powershell_script", side_effect=fake_run):
            result = open_documents._run_spreadsheet_attached_com(
                "selection-read",
                {"path": "C:\\docs\\book.xlsx"},
            )
        script = captured["script"]
        self.assertTrue(result["ok"])
        self.assertIn("Find-OpenWorkbook", script)
        self.assertIn("selection_kind", script)
        self.assertIn("ShapeRange", script)
        self.assertIn("object_kind", script)
        self.assertIn("has_text_frame", script)
        self.assertIn("TextFrame", script)
        self.assertIn("ChartType", script)
        self.assertIn("HasTitle", script)
        self.assertNotIn(".Quit()", script)

    def test_reads_cell_selection_metadata(self):
        fixture = Fixture(self, component="spreadsheets")
        payload = {
            "ok": True,
            "status": "ok",
            "saved": True,
            "read_only": False,
            "selection_kind": "cell",
            "address": "$A$1",
            "sheet_name": "Sheet1",
            "cell_count": 1,
            "area_count": 1,
            "selection_text": "P3-402 mark",
            "sample_truncated": False,
            "formula": None,
            "object_kind": None,
            "object_name": None,
        }
        with patch.object(open_documents, "_run_spreadsheet_attached_com", return_value=com_ok(payload)):
            ok, data, errors = read_spreadsheet_selection(fixture.document_id, fixture.workspace)
        self.assertTrue(ok, errors)
        self.assertEqual(data["selection_kind"], "cell")
        self.assertEqual(data["address"], "$A$1")
        self.assertEqual(data["selection_text"], "P3-402 mark")
        self.assertEqual(data["component"], "spreadsheets")
        self.assertEqual(data["sample_cell_limit"], 100)
        self.assertIsNone(data["object_kind"])
        self.assertFalse(data["object_info_verified"])

    def test_reads_chart_object_metadata_for_non_range(self):
        fixture = Fixture(self, component="spreadsheets")
        payload = {
            "ok": True,
            "status": "ok",
            "saved": True,
            "read_only": False,
            "selection_kind": "non_range",
            "address": None,
            "sheet_name": "Sheet1",
            "cell_count": 0,
            "area_count": 0,
            "selection_text": "",
            "sample_truncated": False,
            "formula": None,
            "object_kind": "chart",
            "object_name": "Chart 1",
        }
        with patch.object(open_documents, "_run_spreadsheet_attached_com", return_value=com_ok(payload)):
            ok, data, errors = read_spreadsheet_selection(fixture.document_id, fixture.workspace)
        self.assertTrue(ok, errors)
        self.assertEqual(data["selection_kind"], "non_range")
        self.assertEqual(data["object_kind"], "chart")
        self.assertEqual(data["object_name"], "Chart 1")
        self.assertEqual(data["sheet_name"], "Sheet1")
        self.assertFalse(data["has_text_frame"])
        self.assertFalse(data["has_chart_title"])
        self.assertFalse(data["has_axis_title"])
        self.assertFalse(data["has_series_values"])
        self.assertIsNone(data["series_name"])
        self.assertFalse(data["object_info_verified"])
        self.assertIn("ChartTitle", data["object_note"])
        self.assertIn("AxisTitle", data["object_note"])
        self.assertIn("Series Values", data["object_note"])

    def test_reads_chart_title_metadata(self):
        fixture = Fixture(self, component="spreadsheets")
        payload = {
            "ok": True,
            "status": "ok",
            "saved": True,
            "read_only": False,
            "selection_kind": "non_range",
            "address": None,
            "sheet_name": "Sheet1",
            "cell_count": 0,
            "area_count": 0,
            "selection_text": "P3-443 title",
            "sample_truncated": False,
            "formula": None,
            "object_kind": "chart",
            "object_name": "Chart 1",
            "has_chart_title": True,
        }
        with patch.object(open_documents, "_run_spreadsheet_attached_com", return_value=com_ok(payload)):
            ok, data, errors = read_spreadsheet_selection(fixture.document_id, fixture.workspace)
        self.assertTrue(ok, errors)
        self.assertTrue(data["has_chart_title"])
        self.assertFalse(data["has_axis_title"])
        self.assertEqual(data["selection_text"], "P3-443 title")

    def test_reads_axis_title_metadata(self):
        fixture = Fixture(self, component="spreadsheets")
        payload = {
            "ok": True,
            "status": "ok",
            "saved": True,
            "read_only": False,
            "selection_kind": "non_range",
            "address": None,
            "sheet_name": "Sheet1",
            "cell_count": 0,
            "area_count": 0,
            "selection_text": "P3-446 axis",
            "sample_truncated": False,
            "formula": None,
            "object_kind": "chart",
            "object_name": "Chart 1",
            "has_axis_title": True,
        }
        with patch.object(open_documents, "_run_spreadsheet_attached_com", return_value=com_ok(payload)):
            ok, data, errors = read_spreadsheet_selection(fixture.document_id, fixture.workspace)
        self.assertTrue(ok, errors)
        self.assertTrue(data["has_axis_title"])
        self.assertFalse(data["has_chart_title"])
        self.assertFalse(data["has_series_values"])
        self.assertEqual(data["selection_text"], "P3-446 axis")

    def test_reads_series_values_metadata(self):
        fixture = Fixture(self, component="spreadsheets")
        payload = {
            "ok": True,
            "status": "ok",
            "saved": True,
            "read_only": False,
            "selection_kind": "non_range",
            "address": None,
            "sheet_name": "Sheet1",
            "cell_count": 0,
            "area_count": 0,
            "selection_text": "10\t20",
            "sample_truncated": False,
            "formula": None,
            "object_kind": "chart",
            "object_name": "P3-447-Chart",
            "has_series_values": True,
            "series_name": "系列 1",
        }
        with patch.object(open_documents, "_run_spreadsheet_attached_com", return_value=com_ok(payload)):
            ok, data, errors = read_spreadsheet_selection(fixture.document_id, fixture.workspace)
        self.assertTrue(ok, errors)
        self.assertTrue(data["has_series_values"])
        self.assertEqual(data["series_name"], "系列 1")
        self.assertFalse(data["has_chart_title"])
        self.assertFalse(data["has_axis_title"])
        self.assertEqual(data["selection_text"], "10\t20")

    def test_reads_shape_text_frame_metadata(self):
        fixture = Fixture(self, component="spreadsheets")
        payload = {
            "ok": True,
            "status": "ok",
            "saved": True,
            "read_only": False,
            "selection_kind": "non_range",
            "address": None,
            "sheet_name": "Sheet1",
            "cell_count": 0,
            "area_count": 0,
            "selection_text": "P3-442 mark",
            "sample_truncated": False,
            "formula": None,
            "object_kind": "shape",
            "object_name": "P3-442-Rect",
            "has_text_frame": True,
        }
        with patch.object(open_documents, "_run_spreadsheet_attached_com", return_value=com_ok(payload)):
            ok, data, errors = read_spreadsheet_selection(fixture.document_id, fixture.workspace)
        self.assertTrue(ok, errors)
        self.assertEqual(data["object_kind"], "shape")
        self.assertEqual(data["selection_text"], "P3-442 mark")
        self.assertTrue(data["has_text_frame"])

    def test_rejects_writer_document_ids(self):
        fixture = Fixture(self, component="writer")
        with patch.object(open_documents, "_run_spreadsheet_attached_com") as com:
            ok, _data, errors = read_spreadsheet_selection(fixture.document_id, fixture.workspace)
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "UNSUPPORTED_COMPONENT")
        com.assert_not_called()


class SpreadsheetSelectionReplaceTests(unittest.TestCase):
    def _cell_selection(self, text="P3-404 mark", **extra):
        payload = {
            "ok": True,
            "status": "ok",
            "saved": True,
            "read_only": False,
            "selection_kind": "cell",
            "address": "$A$1",
            "sheet_name": "Sheet1",
            "cell_count": 1,
            "area_count": 1,
            "row_count": 1,
            "column_count": 1,
            "selection_text": text,
            "sample_truncated": False,
            "formula": None,
        }
        payload.update(extra)
        return payload

    def test_selection_replace_script_never_quits(self):
        captured = {}

        def fake_run(script, timeout_seconds=120, **_kwargs):
            captured["script"] = script
            return {"ok": True, "status": "no_instance"}, None

        with patch.object(open_documents, "run_powershell_script", side_effect=fake_run):
            result = open_documents._run_spreadsheet_attached_com(
                "selection-replace",
                {
                    "path": "C:\\docs\\book.xlsx",
                    "text": "NEW",
                    "expected_text": "OLD",
                    "expected_address": "$A$1",
                },
            )
        script = captured["script"]
        self.assertTrue(result["ok"])
        self.assertIn("Value2", script)
        self.assertIn("Find-OpenWorkbook", script)
        self.assertIn("$sel.Areas", script)
        self.assertIn("matrix", script)
        self.assertNotIn(".Quit()", script)

    def test_dry_run_does_not_write(self):
        fixture = Fixture(self, component="spreadsheets")
        with patch.object(
            open_documents,
            "_run_spreadsheet_attached_com",
            return_value=com_ok(self._cell_selection()),
        ) as com:
            ok, data, errors, replayed = replace_spreadsheet_selection(
                fixture.document_id,
                "P3-404 done",
                request_id="xlsx-sel-dry",
                expected_selection_text="P3-404 mark",
                dry_run=True,
                workspace=fixture.workspace,
            )
        self.assertTrue(ok, errors)
        self.assertFalse(replayed)
        self.assertTrue(data["dry_run"])
        self.assertTrue(data["would_modify"])
        self.assertEqual(com.call_count, 1)
        self.assertEqual(com.call_args.args[0], "selection-read")

    def test_multi_area_matrix_write_flattens_in_area_order(self):
        fixture = Fixture(self, component="spreadsheets")
        read_payload = self._cell_selection(
            text="a\tb",
            selection_kind="multi_area",
            address="$A$1,$C$1",
            cell_count=2,
            area_count=2,
            row_count=0,
            column_count=0,
        )
        write_payload = {
            "ok": True,
            "status": "ok",
            "backend": "powershell-com",
            "address": "$A$1,$C$1",
            "cell_count": 2,
            "area_count": 2,
            "cells_matched": True,
            "read_back_text": "x",
            "read_back_values": ["x", "y"],
        }
        calls = []

        def fake_com(action, params):
            calls.append(action)
            if action == "selection-read":
                return com_ok(read_payload)
            self.assertEqual(action, "selection-replace")
            self.assertEqual(params["mode"], "matrix")
            self.assertEqual(params["values"], [["x", "y"]])
            self.assertEqual(params["row_count"], 1)
            self.assertEqual(params["column_count"], 2)
            return com_ok(write_payload)

        with patch.object(open_documents, "_run_spreadsheet_attached_com", side_effect=fake_com):
            with patch.object(
                open_documents,
                "create_backup",
                return_value=(True, {"created": True, "backup_path": "b.xlsx"}, [], False),
            ):
                with patch.object(open_documents, "guarded_com_mutation", side_effect=lambda path, backup, fn: fn()):
                    with patch.object(open_documents, "verify_post_com_source", return_value=None):
                        ok, data, errors, _replayed = replace_spreadsheet_selection(
                            fixture.document_id,
                            request_id="xlsx-sel-multi-area-matrix",
                            values=[["x", "y"]],
                            expected_selection_text="a\tb",
                            workspace=fixture.workspace,
                        )
        self.assertTrue(ok, errors)
        self.assertEqual(data["fill_mode"], "matrix")
        self.assertTrue(data["validation_passed"])
        self.assertEqual(data["read_back_values"], ["x", "y"])
        self.assertEqual(calls.count("selection-replace"), 1)

    def test_multi_area_matrix_cell_count_mismatch_before_backup(self):
        fixture = Fixture(self, component="spreadsheets")
        payload = self._cell_selection(
            text="a\tb",
            selection_kind="multi_area",
            address="$A$1,$C$1",
            cell_count=2,
            area_count=2,
            row_count=0,
            column_count=0,
        )
        with patch.object(open_documents, "_run_spreadsheet_attached_com", return_value=com_ok(payload)):
            with patch.object(open_documents, "create_backup") as backup:
                ok, _data, errors, _replayed = replace_spreadsheet_selection(
                    fixture.document_id,
                    request_id="xlsx-sel-multi-area-matrix-shape",
                    values=[["x", "y", "z"]],
                    workspace=fixture.workspace,
                )
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "RANGE_SHAPE_MISMATCH")
        backup.assert_not_called()

    def test_rejects_chart_without_title_api_before_backup(self):
        fixture = Fixture(self, component="spreadsheets")
        payload = self._cell_selection(
            text="",
            selection_kind="non_range",
            address=None,
            cell_count=0,
            area_count=0,
            object_kind="chart",
            object_name="Chart 1",
            sheet_name="Sheet1",
            has_chart_title=False,
        )
        with patch.object(open_documents, "_run_spreadsheet_attached_com", return_value=com_ok(payload)):
            with patch.object(open_documents, "create_backup") as backup:
                ok, data, errors, _replayed = replace_spreadsheet_selection(
                    fixture.document_id,
                    "NEW",
                    request_id="xlsx-sel-chart-reject",
                    workspace=fixture.workspace,
                )
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "OBJECT_SELECTION_UNSUPPORTED")
        self.assertIn("Chart 1", errors[0]["message"])
        self.assertIn("ChartTitle", errors[0]["message"])
        self.assertIn("Series Values", errors[0]["message"])
        self.assertIn("cannot be replaced", errors[0]["message"])
        self.assertEqual(data["selection"]["object_kind"], "chart")
        backup.assert_not_called()

    def test_series_values_replace_dry_run_and_write(self):
        fixture = Fixture(self, component="spreadsheets")
        read_payload = self._cell_selection(
            text="10\t20",
            selection_kind="non_range",
            address=None,
            cell_count=0,
            area_count=0,
            object_kind="chart",
            object_name="P3-447-Chart",
            sheet_name="Sheet1",
            has_series_values=True,
            series_name="系列 1",
        )

        def fake(action, params):
            if action == "selection-read":
                return com_ok(read_payload)
            self.assertEqual(action, "selection-series-values-replace")
            self.assertEqual(params["values"], ["30", "40"])
            self.assertEqual(params["expected_object_name"], "P3-447-Chart")
            return com_ok({
                "ok": True,
                "status": "ok",
                "backend": "powershell-com",
                "object_kind": "chart",
                "object_name": "P3-447-Chart",
                "series_name": "系列 1",
                "cells_matched": True,
                "read_back_text": "30\t40",
            })

        with patch.object(open_documents, "_run_spreadsheet_attached_com", side_effect=fake):
            with patch.object(
                open_documents,
                "create_backup",
                return_value=(True, {"created": True, "backup_path": "b.xlsx"}, [], False),
            ):
                with patch.object(open_documents, "guarded_com_mutation", side_effect=lambda path, backup, fn: fn()):
                    with patch.object(open_documents, "verify_post_com_source", return_value=None):
                        ok, data, errors, _replayed = replace_spreadsheet_selection(
                            fixture.document_id,
                            "30\t40",
                            request_id="xlsx-series-values",
                            expected_selection_text="10\t20",
                            workspace=fixture.workspace,
                        )
        self.assertTrue(ok, errors)
        self.assertEqual(data["fill_mode"], "series_values")
        self.assertEqual(data["read_back_text"], "30\t40")
        self.assertIn("Series.Values", data["note"])

    def test_series_values_rejects_length_mismatch_before_backup(self):
        fixture = Fixture(self, component="spreadsheets")
        payload = self._cell_selection(
            text="10\t20",
            selection_kind="non_range",
            address=None,
            cell_count=0,
            area_count=0,
            object_kind="chart",
            object_name="P3-447-Chart",
            has_series_values=True,
        )
        with patch.object(open_documents, "_run_spreadsheet_attached_com", return_value=com_ok(payload)):
            with patch.object(open_documents, "create_backup") as backup:
                ok, data, errors, _replayed = replace_spreadsheet_selection(
                    fixture.document_id,
                    "30\t40\t50",
                    request_id="xlsx-series-len",
                    workspace=fixture.workspace,
                )
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "RANGE_SHAPE_MISMATCH")
        self.assertEqual(data["expected_count"], 2)
        self.assertEqual(data["provided_count"], 3)
        backup.assert_not_called()

    def test_axis_title_replace_dry_run_and_write(self):
        fixture = Fixture(self, component="spreadsheets")
        read_payload = self._cell_selection(
            text="P3-446 axis",
            selection_kind="non_range",
            address=None,
            cell_count=0,
            area_count=0,
            object_kind="chart",
            object_name="Chart 1",
            sheet_name="Sheet1",
            has_axis_title=True,
        )

        def fake(action, params):
            if action == "selection-read":
                return com_ok(read_payload)
            self.assertEqual(action, "selection-axis-title-replace")
            self.assertEqual(params["text"], "P3-446 done")
            return com_ok({
                "ok": True,
                "status": "ok",
                "backend": "powershell-com",
                "object_kind": "chart",
                "object_name": "Chart 1",
                "cells_matched": True,
                "read_back_text": "P3-446 done",
            })

        with patch.object(open_documents, "_run_spreadsheet_attached_com", side_effect=fake):
            with patch.object(
                open_documents,
                "create_backup",
                return_value=(True, {"created": True, "backup_path": "b.xlsx"}, [], False),
            ):
                with patch.object(open_documents, "guarded_com_mutation", side_effect=lambda path, backup, fn: fn()):
                    with patch.object(open_documents, "verify_post_com_source", return_value=None):
                        ok, data, errors, _replayed = replace_spreadsheet_selection(
                            fixture.document_id,
                            "P3-446 done",
                            request_id="xlsx-axis-title",
                            expected_selection_text="P3-446 axis",
                            workspace=fixture.workspace,
                        )
        self.assertTrue(ok, errors)
        self.assertEqual(data["fill_mode"], "axis_title")
        self.assertEqual(data["read_back_text"], "P3-446 done")
        self.assertIn("AxisTitle", data["note"])

    def test_chart_title_replace_dry_run_and_write(self):
        fixture = Fixture(self, component="spreadsheets")
        read_payload = self._cell_selection(
            text="P3-443 title",
            selection_kind="non_range",
            address=None,
            cell_count=0,
            area_count=0,
            object_kind="chart",
            object_name="Chart 1",
            sheet_name="Sheet1",
            has_chart_title=True,
        )
        calls = []

        def fake(action, params):
            calls.append(action)
            if action == "selection-read":
                return com_ok(read_payload)
            self.assertEqual(action, "selection-chart-title-replace")
            self.assertEqual(params["expected_object_name"], "Chart 1")
            self.assertEqual(params["text"], "P3-443 done")
            fixture.document.write_bytes(fixture.document.read_bytes() + b"\0")
            return com_ok({
                "ok": True,
                "status": "ok",
                "backend": "fake-wps",
                "object_kind": "chart",
                "object_name": "Chart 1",
                "cells_matched": True,
                "read_back_text": "P3-443 done",
            })

        with patch.object(open_documents, "_run_spreadsheet_attached_com", side_effect=fake):
            ok, data, errors, _replayed = replace_spreadsheet_selection(
                fixture.document_id,
                "P3-443 done",
                request_id="xlsx-chart-dry",
                expected_selection_text="P3-443 title",
                dry_run=True,
                workspace=fixture.workspace,
            )
            self.assertTrue(ok, errors)
            self.assertEqual(data["fill_mode"], "chart_title")
            self.assertEqual(calls, ["selection-read"])

            ok, data, errors, _replayed = replace_spreadsheet_selection(
                fixture.document_id,
                "P3-443 done",
                request_id="xlsx-chart-write",
                expected_selection_text="P3-443 title",
                workspace=fixture.workspace,
            )
        self.assertTrue(ok, errors)
        self.assertTrue(data["validation_passed"])
        self.assertEqual(data["read_back_text"], "P3-443 done")
        self.assertIn("ChartTitle", data["note"])
        self.assertEqual(calls.count("selection-chart-title-replace"), 1)

    def test_rejects_shape_without_text_frame_before_backup(self):
        fixture = Fixture(self, component="spreadsheets")
        payload = self._cell_selection(
            text="",
            selection_kind="non_range",
            address=None,
            cell_count=0,
            area_count=0,
            object_kind="shape",
            object_name="Rectangle 1",
            sheet_name="Sheet1",
            has_text_frame=False,
        )
        with patch.object(open_documents, "_run_spreadsheet_attached_com", return_value=com_ok(payload)):
            with patch.object(open_documents, "create_backup") as backup:
                ok, data, errors, _replayed = replace_spreadsheet_selection(
                    fixture.document_id,
                    "NEW",
                    request_id="xlsx-sel-shape-reject",
                    workspace=fixture.workspace,
                )
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "OBJECT_SELECTION_UNSUPPORTED")
        self.assertIn("Rectangle 1", errors[0]["message"])
        self.assertIn("TextFrame", errors[0]["message"])
        self.assertEqual(data["selection"]["object_kind"], "shape")
        backup.assert_not_called()

    def test_shape_text_replace_dry_run_and_write(self):
        fixture = Fixture(self, component="spreadsheets")
        read_payload = self._cell_selection(
            text="P3-442 mark",
            selection_kind="non_range",
            address=None,
            cell_count=0,
            area_count=0,
            object_kind="shape",
            object_name="P3-442-Rect",
            sheet_name="Sheet1",
            has_text_frame=True,
        )
        calls = []

        def fake(action, params):
            calls.append(action)
            if action == "selection-read":
                return com_ok(read_payload)
            self.assertEqual(action, "selection-shape-replace")
            self.assertEqual(params["expected_object_name"], "P3-442-Rect")
            self.assertEqual(params["expected_text"], "P3-442 mark")
            self.assertEqual(params["text"], "P3-442 done")
            fixture.document.write_bytes(fixture.document.read_bytes() + b"\0")
            return com_ok({
                "ok": True,
                "status": "ok",
                "backend": "fake-wps",
                "object_kind": "shape",
                "object_name": "P3-442-Rect",
                "cells_matched": True,
                "read_back_text": "P3-442 done",
            })

        with patch.object(open_documents, "_run_spreadsheet_attached_com", side_effect=fake):
            ok, data, errors, _replayed = replace_spreadsheet_selection(
                fixture.document_id,
                "P3-442 done",
                request_id="xlsx-shape-dry",
                expected_selection_text="P3-442 mark",
                dry_run=True,
                workspace=fixture.workspace,
            )
            self.assertTrue(ok, errors)
            self.assertTrue(data["would_modify"])
            self.assertEqual(data["fill_mode"], "shape_text")
            self.assertEqual(calls, ["selection-read"])

            ok, data, errors, _replayed = replace_spreadsheet_selection(
                fixture.document_id,
                "P3-442 done",
                request_id="xlsx-shape-write",
                expected_selection_text="P3-442 mark",
                workspace=fixture.workspace,
            )
        self.assertTrue(ok, errors)
        self.assertTrue(data["validation_passed"])
        self.assertEqual(data["read_back_text"], "P3-442 done")
        self.assertEqual(data["fill_mode"], "shape_text")
        self.assertIn("TextFrame", data["note"])
        self.assertEqual(calls.count("selection-shape-replace"), 1)

    def test_shape_text_rejects_matrix_before_backup(self):
        fixture = Fixture(self, component="spreadsheets")
        payload = self._cell_selection(
            text="mark",
            selection_kind="non_range",
            address=None,
            cell_count=0,
            area_count=0,
            object_kind="shape",
            object_name="Rect",
            has_text_frame=True,
        )
        with patch.object(open_documents, "_run_spreadsheet_attached_com", return_value=com_ok(payload)):
            with patch.object(open_documents, "create_backup") as backup:
                ok, _data, errors, _replayed = replace_spreadsheet_selection(
                    fixture.document_id,
                    values=[["a"]],
                    request_id="xlsx-shape-matrix",
                    workspace=fixture.workspace,
                )
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "OBJECT_SELECTION_UNSUPPORTED")
        self.assertIn("constant --text", errors[0]["message"])
        backup.assert_not_called()

    def test_rejects_unsaved_before_backup(self):
        fixture = Fixture(self, component="spreadsheets")
        payload = self._cell_selection(saved=False)
        with patch.object(open_documents, "_run_spreadsheet_attached_com", return_value=com_ok(payload)):
            with patch.object(open_documents, "create_backup") as backup:
                ok, data, errors, _replayed = replace_spreadsheet_selection(
                    fixture.document_id,
                    "NEW",
                    request_id="xlsx-sel-unsaved",
                    workspace=fixture.workspace,
                )
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "DOCUMENT_HAS_UNSAVED_CHANGES")
        self.assertFalse(data["selection"]["saved"])
        self.assertIn("save-open-document", errors[0]["message"])
        backup.assert_not_called()

    def test_multi_area_constant_fill_validates_cells_matched(self):
        fixture = Fixture(self, component="spreadsheets")
        read_payload = self._cell_selection(
            text="a\tb",
            selection_kind="multi_area",
            address="$A$1,$C$1",
            cell_count=2,
            area_count=2,
            row_count=0,
            column_count=0,
        )
        write_payload = {
            "ok": True,
            "status": "ok",
            "backend": "powershell-com",
            "address": "$A$1,$C$1",
            "cell_count": 2,
            "area_count": 2,
            "cells_matched": True,
            "read_back_text": "P3-407 fill",
        }

        def fake_com(action, params):
            if action == "selection-read":
                return com_ok(read_payload)
            self.assertEqual(action, "selection-replace")
            self.assertEqual(params["expected_address"], "$A$1,$C$1")
            self.assertEqual(params["expected_text"], "a\tb")
            self.assertEqual(params["max_cells"], 100)
            self.assertEqual(params["mode"], "constant")
            self.assertEqual(params["text"], "P3-407 fill")
            return com_ok(write_payload)

        with patch.object(open_documents, "_run_spreadsheet_attached_com", side_effect=fake_com):
            with patch.object(
                open_documents,
                "create_backup",
                return_value=(True, {"created": True, "backup_path": "b.xlsx"}, [], False),
            ):
                with patch.object(open_documents, "guarded_com_mutation", side_effect=lambda path, backup, fn: fn()):
                    with patch.object(open_documents, "verify_post_com_source", return_value=None):
                        with patch.object(open_documents, "record_operation") as record:
                            ok, data, errors, replayed = replace_spreadsheet_selection(
                                fixture.document_id,
                                "P3-407 fill",
                                request_id="xlsx-sel-multi-area-write",
                                expected_selection_text="a\tb",
                                workspace=fixture.workspace,
                            )
        self.assertTrue(ok, errors)
        self.assertFalse(replayed)
        self.assertTrue(data["validation_passed"])
        self.assertTrue(data["cells_matched"])
        self.assertEqual(data["cell_count"], 2)
        self.assertEqual(data["area_count"], 2)
        self.assertEqual(data["fill_mode"], "constant")
        record.assert_called_once()

    def test_rejects_oversized_range_before_backup(self):
        fixture = Fixture(self, component="spreadsheets")
        payload = self._cell_selection(
            text="x",
            selection_kind="range",
            address="$A$1:$A$200",
            cell_count=200,
            area_count=1,
            sample_truncated=True,
        )
        with patch.object(open_documents, "_run_spreadsheet_attached_com", return_value=com_ok(payload)):
            with patch.object(open_documents, "create_backup") as backup:
                ok, _data, errors, _replayed = replace_spreadsheet_selection(
                    fixture.document_id,
                    "NEW",
                    request_id="xlsx-sel-too-large",
                    workspace=fixture.workspace,
                )
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "SELECTION_TOO_LARGE")
        backup.assert_not_called()

    def test_range_fill_validates_cells_matched(self):
        fixture = Fixture(self, component="spreadsheets")
        read_payload = self._cell_selection(
            text="a\tb",
            selection_kind="range",
            address="$A$1:$B$1",
            cell_count=2,
            area_count=1,
            row_count=1,
            column_count=2,
        )
        write_payload = {
            "ok": True,
            "status": "ok",
            "backend": "powershell-com",
            "address": "$A$1:$B$1",
            "cell_count": 2,
            "cells_matched": True,
            "read_back_text": "P3-405 fill",
        }

        def fake_com(action, params):
            if action == "selection-read":
                return com_ok(read_payload)
            self.assertEqual(action, "selection-replace")
            self.assertEqual(params["expected_address"], "$A$1:$B$1")
            self.assertEqual(params["expected_text"], "a\tb")
            self.assertEqual(params["max_cells"], 100)
            self.assertEqual(params["mode"], "constant")
            return com_ok(write_payload)

        with patch.object(open_documents, "_run_spreadsheet_attached_com", side_effect=fake_com):
            with patch.object(
                open_documents,
                "create_backup",
                return_value=(True, {"created": True, "backup_path": "b.xlsx"}, [], False),
            ):
                with patch.object(open_documents, "guarded_com_mutation", side_effect=lambda path, backup, fn: fn()):
                    with patch.object(open_documents, "verify_post_com_source", return_value=None):
                        with patch.object(open_documents, "record_operation") as record:
                            ok, data, errors, replayed = replace_spreadsheet_selection(
                                fixture.document_id,
                                "P3-405 fill",
                                request_id="xlsx-sel-range-write",
                                expected_selection_text="a\tb",
                                workspace=fixture.workspace,
                            )
        self.assertTrue(ok, errors)
        self.assertFalse(replayed)
        self.assertTrue(data["validation_passed"])
        self.assertTrue(data["cells_matched"])
        self.assertEqual(data["cell_count"], 2)
        self.assertEqual(data["fill_mode"], "constant")
        record.assert_called_once()

    def test_rejects_mismatch_before_backup(self):
        fixture = Fixture(self, component="spreadsheets")
        with patch.object(
            open_documents,
            "_run_spreadsheet_attached_com",
            return_value=com_ok(self._cell_selection()),
        ):
            with patch.object(open_documents, "create_backup") as backup:
                ok, _data, errors, _replayed = replace_spreadsheet_selection(
                    fixture.document_id,
                    "NEW",
                    request_id="xlsx-sel-mismatch",
                    expected_selection_text="other",
                    workspace=fixture.workspace,
                )
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "SELECTION_MISMATCH")
        backup.assert_not_called()

    def test_write_validates_read_back(self):
        fixture = Fixture(self, component="spreadsheets")
        read_payload = self._cell_selection()
        write_payload = {
            "ok": True,
            "status": "ok",
            "backend": "powershell-com",
            "address": "$A$1",
            "read_back_text": "P3-404 done",
        }

        def fake_com(action, params):
            if action == "selection-read":
                return com_ok(read_payload)
            self.assertEqual(action, "selection-replace")
            self.assertEqual(params["expected_text"], "P3-404 mark")
            self.assertEqual(params["expected_address"], "$A$1")
            self.assertEqual(params["text"], "P3-404 done")
            self.assertEqual(params["max_cells"], 100)
            self.assertEqual(params["mode"], "constant")
            return com_ok(write_payload)

        with patch.object(open_documents, "_run_spreadsheet_attached_com", side_effect=fake_com):
            with patch.object(
                open_documents,
                "create_backup",
                return_value=(True, {"created": True, "backup_path": "b.xlsx"}, [], False),
            ):
                with patch.object(open_documents, "guarded_com_mutation", side_effect=lambda path, backup, fn: fn()):
                    with patch.object(open_documents, "verify_post_com_source", return_value=None):
                        with patch.object(open_documents, "record_operation") as record:
                            ok, data, errors, replayed = replace_spreadsheet_selection(
                                fixture.document_id,
                                "P3-404 done",
                                request_id="xlsx-sel-write",
                                expected_selection_text="P3-404 mark",
                                workspace=fixture.workspace,
                            )
        self.assertTrue(ok, errors)
        self.assertFalse(replayed)
        self.assertTrue(data["validation_passed"])
        self.assertEqual(data["read_back_text"], "P3-404 done")
        record.assert_called_once()

    def test_matrix_write_validates_shape_and_read_back(self):
        fixture = Fixture(self, component="spreadsheets")
        read_payload = self._cell_selection(
            text="a\tb",
            selection_kind="range",
            address="$A$1:$B$1",
            cell_count=2,
            area_count=1,
            row_count=1,
            column_count=2,
        )
        write_payload = {
            "ok": True,
            "status": "ok",
            "backend": "powershell-com",
            "address": "$A$1:$B$1",
            "cell_count": 2,
            "cells_matched": True,
            "read_back_text": "c",
            "read_back_values": ["c", "d"],
        }

        def fake_com(action, params):
            if action == "selection-read":
                return com_ok(read_payload)
            self.assertEqual(action, "selection-replace")
            self.assertEqual(params["mode"], "matrix")
            self.assertEqual(params["values"], [["c", "d"]])
            self.assertEqual(params["row_count"], 1)
            self.assertEqual(params["column_count"], 2)
            return com_ok(write_payload)

        with patch.object(open_documents, "_run_spreadsheet_attached_com", side_effect=fake_com):
            with patch.object(
                open_documents,
                "create_backup",
                return_value=(True, {"created": True, "backup_path": "b.xlsx"}, [], False),
            ):
                with patch.object(open_documents, "guarded_com_mutation", side_effect=lambda path, backup, fn: fn()):
                    with patch.object(open_documents, "verify_post_com_source", return_value=None):
                        with patch.object(open_documents, "record_operation") as record:
                            ok, data, errors, replayed = replace_spreadsheet_selection(
                                fixture.document_id,
                                request_id="xlsx-sel-matrix",
                                values=[["c", "d"]],
                                expected_selection_text="a\tb",
                                workspace=fixture.workspace,
                            )
        self.assertTrue(ok, errors)
        self.assertFalse(replayed)
        self.assertEqual(data["fill_mode"], "matrix")
        self.assertTrue(data["validation_passed"])
        self.assertEqual(data["values"], [["c", "d"]])
        record.assert_called_once()

    def test_matrix_shape_mismatch_before_backup(self):
        fixture = Fixture(self, component="spreadsheets")
        payload = self._cell_selection(
            text="a\tb",
            selection_kind="range",
            address="$A$1:$B$1",
            cell_count=2,
            area_count=1,
            row_count=1,
            column_count=2,
        )
        with patch.object(open_documents, "_run_spreadsheet_attached_com", return_value=com_ok(payload)):
            with patch.object(open_documents, "create_backup") as backup:
                ok, _data, errors, _replayed = replace_spreadsheet_selection(
                    fixture.document_id,
                    request_id="xlsx-sel-shape",
                    values=[["c"], ["d"]],
                    workspace=fixture.workspace,
                )
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "RANGE_SHAPE_MISMATCH")
        backup.assert_not_called()

    def test_rejects_text_and_values_together(self):
        fixture = Fixture(self, component="spreadsheets")
        with patch.object(open_documents, "_run_spreadsheet_attached_com") as com:
            ok, _data, errors, _replayed = replace_spreadsheet_selection(
                fixture.document_id,
                "NEW",
                request_id="xlsx-sel-both",
                values=[["x"]],
                workspace=fixture.workspace,
            )
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "INVALID_ARGUMENT")
        com.assert_not_called()


class SpreadsheetSelectionDeleteTests(unittest.TestCase):
    def _cell_selection(self, text="P3-419 mark", **extra):
        payload = {
            "ok": True,
            "status": "ok",
            "saved": True,
            "read_only": False,
            "selection_kind": "cell",
            "address": "$A$1",
            "sheet_name": "Sheet1",
            "cell_count": 1,
            "area_count": 1,
            "row_count": 1,
            "column_count": 1,
            "selection_text": text,
            "sample_truncated": False,
            "formula": None,
        }
        payload.update(extra)
        return payload

    def test_dry_run_and_delete_with_backup_replay(self):
        fixture = Fixture(self, component="spreadsheets")
        calls = []
        original = fixture.document.read_bytes()

        def fake(action, params):
            calls.append(action)
            if action == "selection-read":
                return com_ok(self._cell_selection())
            self.assertEqual(action, "selection-delete")
            fixture.document.write_bytes(fixture.document.read_bytes() + b"\0")
            return com_ok({
                "ok": True,
                "status": "ok",
                "backend": "fake-wps",
                "deleted_text": params["expected_text"],
                "read_back_text": "",
                "read_back_values": [""],
                "cells_matched": True,
            })

        with patch.object(open_documents, "_run_spreadsheet_attached_com", side_effect=fake):
            ok, data, errors, replayed = delete_spreadsheet_selection(
                fixture.document_id, request_id="xlsx-del-1", dry_run=True,
                expected_selection_text="P3-419 mark", workspace=fixture.workspace,
            )
            self.assertTrue(ok, errors)
            self.assertFalse(replayed)
            self.assertEqual(calls, ["selection-read"])
            self.assertTrue(data["would_modify"])

            ok, data, errors, _replayed = delete_spreadsheet_selection(
                fixture.document_id, request_id="xlsx-del-1",
                expected_selection_text="P3-419 mark", workspace=fixture.workspace,
            )
            self.assertTrue(ok, errors)
            self.assertTrue(data["validation_passed"])
            self.assertTrue(data["backup"]["created"])
            self.assertEqual(Path(data["backup"]["backup_path"]).read_bytes(), original)
            self.assertEqual(data["deleted_text"], "P3-419 mark")

            ok, again, errors, replayed = delete_spreadsheet_selection(
                fixture.document_id, request_id="xlsx-del-1",
                expected_selection_text="P3-419 mark", workspace=fixture.workspace,
            )
            self.assertTrue(ok, errors)
            self.assertTrue(replayed)
        self.assertEqual(calls.count("selection-delete"), 1)

    def test_rejects_unsaved_before_backup(self):
        fixture = Fixture(self, component="spreadsheets")
        with patch.object(open_documents, "_run_spreadsheet_attached_com", return_value=com_ok(self._cell_selection(saved=False))):
            with patch.object(open_documents, "create_backup") as backup:
                ok, _data, errors, _replayed = delete_spreadsheet_selection(
                    fixture.document_id, request_id="xlsx-del-unsaved", workspace=fixture.workspace,
                )
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "DOCUMENT_HAS_UNSAVED_CHANGES")
        backup.assert_not_called()

    def test_empty_and_chart_rejected_before_backup(self):
        fixture = Fixture(self, component="spreadsheets")
        cases = [
            ("empty", self._cell_selection(text=""), "SELECTION_EMPTY"),
            (
                "chart-no-title",
                self._cell_selection(
                    text="", selection_kind="non_range", address=None, cell_count=0, area_count=0,
                    object_kind="chart", object_name="Chart 1", has_chart_title=False,
                ),
                "OBJECT_SELECTION_UNSUPPORTED",
            ),
            (
                "shape-no-frame",
                self._cell_selection(
                    text="", selection_kind="non_range", address=None, cell_count=0, area_count=0,
                    object_kind="shape", object_name="Rect", has_text_frame=False,
                ),
                "OBJECT_SELECTION_UNSUPPORTED",
            ),
        ]
        for name, payload, code in cases:
            with self.subTest(name=name):
                with patch.object(open_documents, "_run_spreadsheet_attached_com", return_value=com_ok(payload)):
                    with patch.object(open_documents, "create_backup") as backup:
                        ok, _data, errors, _replayed = delete_spreadsheet_selection(
                            fixture.document_id, request_id=f"xlsx-del-{name}", workspace=fixture.workspace,
                        )
                self.assertFalse(ok)
                self.assertEqual(errors[0]["code"], code)
                if code == "OBJECT_SELECTION_UNSUPPORTED":
                    self.assertIn("cannot be cleared", errors[0]["message"])
                    self.assertNotIn("cannot be replaced", errors[0]["message"])
                backup.assert_not_called()

    def test_shape_text_delete_with_backup(self):
        fixture = Fixture(self, component="spreadsheets")
        calls = []

        def fake(action, params):
            calls.append(action)
            if action == "selection-read":
                return com_ok(self._cell_selection(
                    text="P3-442 mark",
                    selection_kind="non_range",
                    address=None,
                    cell_count=0,
                    area_count=0,
                    object_kind="shape",
                    object_name="P3-442-Rect",
                    has_text_frame=True,
                ))
            self.assertEqual(action, "selection-shape-delete")
            self.assertEqual(params["expected_object_name"], "P3-442-Rect")
            fixture.document.write_bytes(fixture.document.read_bytes() + b"\0")
            return com_ok({
                "ok": True,
                "status": "ok",
                "backend": "fake-wps",
                "deleted_text": "P3-442 mark",
                "read_back_text": "",
                "cells_matched": True,
            })

        with patch.object(open_documents, "_run_spreadsheet_attached_com", side_effect=fake):
            ok, data, errors, _replayed = delete_spreadsheet_selection(
                fixture.document_id,
                request_id="xlsx-shape-del",
                expected_selection_text="P3-442 mark",
                workspace=fixture.workspace,
            )
        self.assertTrue(ok, errors)
        self.assertTrue(data["validation_passed"])
        self.assertEqual(data["fill_mode"], "shape_text")
        self.assertEqual(data["read_back_text"], "")
        self.assertEqual(calls.count("selection-shape-delete"), 1)

    def test_chart_title_delete_with_backup(self):
        fixture = Fixture(self, component="spreadsheets")
        calls = []

        def fake(action, params):
            calls.append(action)
            if action == "selection-read":
                return com_ok(self._cell_selection(
                    text="P3-443 title",
                    selection_kind="non_range",
                    address=None,
                    cell_count=0,
                    area_count=0,
                    object_kind="chart",
                    object_name="Chart 1",
                    has_chart_title=True,
                ))
            self.assertEqual(action, "selection-chart-title-delete")
            fixture.document.write_bytes(fixture.document.read_bytes() + b"\0")
            return com_ok({
                "ok": True,
                "status": "ok",
                "backend": "fake-wps",
                "deleted_text": "P3-443 title",
                "read_back_text": "",
                "cells_matched": True,
            })

        with patch.object(open_documents, "_run_spreadsheet_attached_com", side_effect=fake):
            ok, data, errors, _replayed = delete_spreadsheet_selection(
                fixture.document_id,
                request_id="xlsx-chart-del",
                expected_selection_text="P3-443 title",
                workspace=fixture.workspace,
            )
        self.assertTrue(ok, errors)
        self.assertEqual(data["fill_mode"], "chart_title")
        self.assertTrue(data["validation_passed"])
        self.assertEqual(calls.count("selection-chart-title-delete"), 1)

    def test_series_values_delete_with_backup(self):
        fixture = Fixture(self, component="spreadsheets")
        calls = []

        def fake(action, params):
            calls.append(action)
            if action == "selection-read":
                return com_ok(self._cell_selection(
                    text="10\t20",
                    selection_kind="non_range",
                    address=None,
                    cell_count=0,
                    area_count=0,
                    object_kind="chart",
                    object_name="P3-447-Chart",
                    has_series_values=True,
                    series_name="系列 1",
                ))
            self.assertEqual(action, "selection-series-values-delete")
            self.assertEqual(params["expected_object_name"], "P3-447-Chart")
            fixture.document.write_bytes(fixture.document.read_bytes() + b"\0")
            return com_ok({
                "ok": True,
                "status": "ok",
                "backend": "fake-wps",
                "deleted_text": "10\t20",
                "read_back_text": "\t",
                "cells_matched": True,
            })

        with patch.object(open_documents, "_run_spreadsheet_attached_com", side_effect=fake):
            ok, data, errors, _replayed = delete_spreadsheet_selection(
                fixture.document_id,
                request_id="xlsx-series-del",
                expected_selection_text="10\t20",
                workspace=fixture.workspace,
            )
        self.assertTrue(ok, errors)
        self.assertEqual(data["fill_mode"], "series_values")
        self.assertTrue(data["validation_passed"])
        self.assertEqual(calls.count("selection-series-values-delete"), 1)

    def test_delete_script_never_quits(self):
        captured = {}

        def fake_run(script, timeout_seconds=120, **_kwargs):
            captured["script"] = script
            return {"ok": True, "status": "no_instance"}, None

        with patch.object(open_documents, "run_powershell_script", side_effect=fake_run):
            open_documents._run_spreadsheet_attached_com(
                "selection-delete",
                {"path": "C:\\docs\\book.xlsx", "expected_text": "mark", "expected_address": "$A$1", "max_cells": 100},
            )
        script = captured["script"]
        self.assertIn("[void]$area.ClearContents()", script)
        self.assertIn("Find-OpenWorkbook", script)
        self.assertNotIn(".Quit()", script)


class PresentationSelectionReplaceTests(unittest.TestCase):
    def _text_selection(self, text="P3-401 mark", **extra):
        payload = {
            "ok": True,
            "status": "ok",
            "saved": True,
            "read_only": False,
            "selection_type": 3,
            "selection_kind": "text",
            "selection_text": text,
            "slide_index": 1,
            "shape_name": "Title 1",
            "shape_count": 0,
        }
        payload.update(extra)
        return payload

    def test_selection_replace_script_never_quits(self):
        captured = {}

        def fake_run(script, timeout_seconds=120, **_kwargs):
            captured["script"] = script
            return {"ok": True, "status": "no_instance"}, None

        with patch.object(open_documents, "run_powershell_script", side_effect=fake_run):
            result = open_documents._run_presentation_attached_com(
                "selection-replace",
                {"path": "C:\\docs\\deck.pptx", "text": "NEW", "expected_text": "OLD"},
            )
        script = captured["script"]
        self.assertTrue(result["ok"])
        self.assertIn("TextRange", script)
        self.assertIn("Find-OpenPresentation", script)
        self.assertIn("$tr.Font", script)
        self.assertIn("formatting_preserved", script)
        self.assertNotIn(".Quit()", script)

    def test_dry_run_does_not_write(self):
        fixture = Fixture(self, component="presentation")
        with patch.object(
            open_documents,
            "_run_presentation_attached_com",
            return_value=com_ok(self._text_selection()),
        ) as com:
            ok, data, errors, replayed = replace_presentation_selection(
                fixture.document_id,
                "P3-401 done",
                request_id="pptx-sel-dry",
                expected_selection_text="P3-401 mark",
                dry_run=True,
                workspace=fixture.workspace,
            )
        self.assertTrue(ok, errors)
        self.assertFalse(replayed)
        self.assertTrue(data["dry_run"])
        self.assertTrue(data["would_modify"])
        self.assertEqual(com.call_count, 1)
        self.assertEqual(com.call_args.args[0], "selection-read")

    def test_rejects_non_text_selection_before_backup(self):
        fixture = Fixture(self, component="presentation")
        payload = self._text_selection(text="", selection_kind="shapes", selection_type=2, shape_count=1)
        with patch.object(open_documents, "_run_presentation_attached_com", return_value=com_ok(payload)):
            with patch.object(open_documents, "create_backup") as backup:
                ok, _data, errors, _replayed = replace_presentation_selection(
                    fixture.document_id,
                    "NEW",
                    request_id="pptx-sel-shape",
                    workspace=fixture.workspace,
                )
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "SELECTION_SPANS_STRUCTURE")
        backup.assert_not_called()

    def test_rejects_mismatch_before_backup(self):
        fixture = Fixture(self, component="presentation")
        with patch.object(
            open_documents,
            "_run_presentation_attached_com",
            return_value=com_ok(self._text_selection()),
        ):
            with patch.object(open_documents, "create_backup") as backup:
                ok, _data, errors, _replayed = replace_presentation_selection(
                    fixture.document_id,
                    "NEW",
                    request_id="pptx-sel-mismatch",
                    expected_selection_text="other",
                    workspace=fixture.workspace,
                )
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "SELECTION_MISMATCH")
        backup.assert_not_called()

    def test_rejects_unsaved_before_backup(self):
        fixture = Fixture(self, component="presentation")
        with patch.object(
            open_documents,
            "_run_presentation_attached_com",
            return_value=com_ok(self._text_selection(saved=False)),
        ):
            with patch.object(open_documents, "create_backup") as backup:
                ok, data, errors, _replayed = replace_presentation_selection(
                    fixture.document_id,
                    "NEW",
                    request_id="pptx-sel-unsaved",
                    expected_selection_text="P3-401 mark",
                    workspace=fixture.workspace,
                )
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "DOCUMENT_HAS_UNSAVED_CHANGES")
        self.assertIn("save-open-document", errors[0]["message"])
        self.assertFalse(data["selection"]["saved"])
        backup.assert_not_called()
        self.assertFalse((fixture.workspace / ".wps-agent" / "backups").exists())

    def test_write_validates_read_back(self):
        fixture = Fixture(self, component="presentation")
        read_payload = self._text_selection()
        write_payload = {
            "ok": True,
            "status": "ok",
            "backend": "powershell-com",
            "read_back_text": "P3-401 done",
            "formatting_mode": "uniform_character",
            "formatting_preserved": True,
        }

        def fake_com(action, params):
            if action == "selection-read":
                return com_ok(read_payload)
            self.assertEqual(action, "selection-replace")
            self.assertEqual(params["expected_text"], "P3-401 mark")
            self.assertEqual(params["text"], "P3-401 done")
            return com_ok(write_payload)

        with patch.object(open_documents, "_run_presentation_attached_com", side_effect=fake_com):
            with patch.object(
                open_documents,
                "create_backup",
                return_value=(True, {"created": True, "backup_path": "b.pptx"}, [], False),
            ):
                with patch.object(open_documents, "guarded_com_mutation", side_effect=lambda path, backup, fn: fn()):
                    with patch.object(open_documents, "verify_post_com_source", return_value=None):
                        with patch.object(open_documents, "record_operation") as record:
                            ok, data, errors, replayed = replace_presentation_selection(
                                fixture.document_id,
                                "P3-401 done",
                                request_id="pptx-sel-write",
                                expected_selection_text="P3-401 mark",
                                workspace=fixture.workspace,
                            )
        self.assertTrue(ok, errors)
        self.assertFalse(replayed)
        self.assertTrue(data["validation_passed"])
        self.assertEqual(data["read_back_text"], "P3-401 done")
        self.assertEqual(data["formatting_mode"], "uniform_character")
        self.assertTrue(data["formatting_preserved"])
        self.assertIn("uniform character formatting", data["note"])
        record.assert_called_once()


class PresentationSelectionInsertTests(unittest.TestCase):
    def _caret(self, start=1, **extra):
        payload = {
            "ok": True,
            "status": "ok",
            "saved": True,
            "read_only": False,
            "selection_type": 3,
            "selection_kind": "text",
            "selection_text": "",
            "text_range_start": start,
            "text_range_length": 0,
            "collapsed": True,
            "slide_index": 1,
            "shape_name": "TextBox 1",
            "shape_count": 0,
        }
        payload.update(extra)
        return payload

    def test_dry_run_and_insert_with_backup_replay(self):
        fixture = Fixture(self, component="presentation")
        caret = self._caret(start=1)
        calls = []

        def fake(action, params):
            calls.append(action)
            if action == "selection-read":
                return com_ok(caret)
            self.assertEqual(action, "selection-insert")
            self.assertEqual(params["expected_start"], 1)
            fixture.document.write_bytes(fixture.document.read_bytes() + b"\0")
            return com_ok({
                "ok": True,
                "status": "ok",
                "backend": "fake-wps",
                "start": 1,
                "end": 12,
                "read_back_text": params["text"],
            })

        original = fixture.document.read_bytes()
        with patch.object(open_documents, "_run_presentation_attached_com", side_effect=fake):
            ok, data, errors, replayed = insert_presentation_selection(
                fixture.document_id, "P3-414 mark", request_id="pptx-ins-dry",
                expected_start=1, dry_run=True, workspace=fixture.workspace,
            )
            self.assertTrue(ok, errors)
            self.assertFalse(replayed)
            self.assertEqual(calls, ["selection-read"])
            self.assertTrue(data["would_modify"])

            ok, data, errors, _replayed = insert_presentation_selection(
                fixture.document_id, "P3-414 mark", request_id="pptx-ins-1",
                expected_start=1, workspace=fixture.workspace,
            )
            self.assertTrue(ok, errors)
            self.assertTrue(data["validation_passed"])
            self.assertTrue(data["backup"]["created"])
            self.assertEqual(Path(data["backup"]["backup_path"]).read_bytes(), original)
            self.assertEqual(data["read_back_text"], "P3-414 mark")

            ok, again, errors, replayed = insert_presentation_selection(
                fixture.document_id, "P3-414 mark", request_id="pptx-ins-1",
                expected_start=1, workspace=fixture.workspace,
            )
            self.assertTrue(ok, errors)
            self.assertTrue(replayed)
            self.assertEqual(again["read_back_text"], "P3-414 mark")
        self.assertEqual(calls.count("selection-insert"), 1)

    def test_rejections_before_backup(self):
        fixture = Fixture(self, component="presentation")
        cases = [
            ("not-collapsed", self._caret(start=1, text_range_length=3, collapsed=False, selection_text="abc"), {}, "SELECTION_NOT_COLLAPSED"),
            ("shapes", self._caret(selection_kind="shapes", selection_type=2, collapsed=False, text_range_length=None), {}, "SELECTION_SPANS_STRUCTURE"),
            ("start-mismatch", self._caret(start=2), {"expected_start": 9}, "SELECTION_MISMATCH"),
            ("read-only", self._caret(read_only=True), {}, "DOCUMENT_READ_ONLY"),
            ("unsaved-flag", self._caret(saved=False), {}, "DOCUMENT_HAS_UNSAVED_CHANGES"),
        ]
        for name, payload, kwargs, code in cases:
            with self.subTest(name=name):
                with patch.object(open_documents, "_run_presentation_attached_com", return_value=com_ok(payload)):
                    with patch.object(open_documents, "create_backup") as backup:
                        ok, _data, errors, _replayed = insert_presentation_selection(
                            fixture.document_id, "X", request_id=f"pptx-ins-{name}",
                            workspace=fixture.workspace, **kwargs,
                        )
                self.assertFalse(ok)
                self.assertEqual(errors[0]["code"], code)
                backup.assert_not_called()

    def test_insert_script_never_quits(self):
        captured = {}

        def fake_run(script, timeout_seconds=120, **_kwargs):
            captured["script"] = script
            return {"ok": True, "status": "no_instance"}, None

        with patch.object(open_documents, "run_powershell_script", side_effect=fake_run):
            open_documents._run_presentation_attached_com(
                "selection-insert",
                {"path": "C:\\docs\\deck.pptx", "text": "hi", "expected_start": 1},
            )
        script = captured["script"]
        self.assertIn("not_collapsed", script)
        self.assertIn("$tr.Text = $insert", script)
        self.assertIn("Find-OpenPresentation", script)
        self.assertNotIn(".Quit()", script)


class SelectionDeleteTests(unittest.TestCase):
    def delete(self, fixture, request_id="del-1", **kwargs):
        return delete_writer_selection(
            document_id=fixture.document_id, request_id=request_id, workspace=fixture.workspace, **kwargs,
        )

    def fake_attached(self, fixture, selection, write_status="ok", probe="", calls=None):
        def fake(action, params):
            if calls is not None:
                calls.append(action)
            if action == "selection-read":
                return com_ok(selection)
            self.assertEqual(action, "selection-delete")
            fixture.document.write_bytes(fixture.document.read_bytes() + b"\0")
            if write_status != "ok":
                return com_ok({"ok": True, "status": write_status})
            return com_ok({
                "ok": True,
                "status": "ok",
                "backend": "fake-wps",
                "deleted_text": params["expected_text"],
                "read_back_text": probe,
            })
        return fake

    def test_dry_run_and_delete_with_backup_replay(self):
        fixture = Fixture(self)
        calls = []
        original = fixture.document.read_bytes()
        with patch.object(open_documents, "_run_writer_attached_com", side_effect=self.fake_attached(fixture, selection_payload(), calls=calls)):
            ok, data, errors, replayed = self.delete(fixture, dry_run=True, expected_selection_text="PENDING_STATUS")
            self.assertTrue(ok, errors)
            self.assertFalse(replayed)
            self.assertEqual(calls, ["selection-read"])
            self.assertTrue(data["would_modify"])

            ok, data, errors, _replayed = self.delete(fixture, expected_selection_text="PENDING_STATUS")
            self.assertTrue(ok, errors)
            self.assertTrue(data["validation_passed"])
            self.assertTrue(data["backup"]["created"])
            self.assertEqual(Path(data["backup"]["backup_path"]).read_bytes(), original)
            self.assertEqual(data["deleted_text"], "PENDING_STATUS")

            ok, again, errors, replayed = self.delete(fixture, expected_selection_text="PENDING_STATUS")
            self.assertTrue(ok, errors)
            self.assertTrue(replayed)
            self.assertEqual(again["deleted_text"], "PENDING_STATUS")
        self.assertEqual(calls.count("selection-delete"), 1)

    def test_rejections_before_backup(self):
        fixture = Fixture(self)
        cases = [
            ("empty", selection_payload(text="", start=5, end=5), {}, "SELECTION_EMPTY"),
            ("mismatch", selection_payload(), {"expected_selection_text": "nope"}, "SELECTION_MISMATCH"),
            ("read-only", selection_payload(read_only=True), {}, "DOCUMENT_READ_ONLY"),
            ("unsaved-flag", selection_payload(saved=False), {}, "DOCUMENT_HAS_UNSAVED_CHANGES"),
        ]
        for name, payload, kwargs, code in cases:
            with self.subTest(name=name):
                calls = []
                with patch.object(open_documents, "_run_writer_attached_com", side_effect=self.fake_attached(fixture, payload, calls=calls)):
                    ok, _data, errors, _replayed = self.delete(fixture, request_id=f"del-{name}", **kwargs)
                self.assertFalse(ok)
                self.assertEqual(errors[0]["code"], code)
                self.assertNotIn("selection-delete", calls)
        self.assertFalse((fixture.workspace / ".wps-agent" / "backups").exists())

    def test_delete_script_never_quits(self):
        captured = {}

        def fake_run(script, timeout_seconds=120, **_kwargs):
            captured["script"] = script
            return {"ok": True, "status": "no_instance"}, None

        with patch.object(open_documents, "run_powershell_script", side_effect=fake_run):
            open_documents._run_writer_attached_com(
                "selection-delete",
                {"path": "C:\\docs\\open.docx", "expected_start": 1, "expected_end": 4, "expected_text": "mark"},
            )
        script = captured["script"]
        self.assertIn("$range.Text = ''", script)
        self.assertIn("Find-OpenWriterDocument", script)
        self.assertNotIn(".Quit()", script)


class PresentationSelectionDeleteTests(unittest.TestCase):
    def _text_selection(self, text="P3-416 mark", **extra):
        payload = {
            "ok": True,
            "status": "ok",
            "saved": True,
            "read_only": False,
            "selection_type": 3,
            "selection_kind": "text",
            "selection_text": text,
            "text_range_start": 1,
            "text_range_length": len(text),
            "collapsed": False,
            "slide_index": 1,
            "shape_name": "TextBox 1",
            "shape_count": 0,
        }
        payload.update(extra)
        return payload

    def test_delete_with_backup_and_reject_empty(self):
        fixture = Fixture(self, component="presentation")
        calls = []

        def fake(action, params):
            calls.append(action)
            if action == "selection-read":
                return com_ok(self._text_selection())
            self.assertEqual(action, "selection-delete")
            fixture.document.write_bytes(fixture.document.read_bytes() + b"\0")
            return com_ok({"ok": True, "status": "ok", "backend": "fake-wps", "deleted_text": params["expected_text"], "read_back_text": ""})

        with patch.object(open_documents, "_run_presentation_attached_com", side_effect=fake):
            ok, data, errors, _replayed = delete_presentation_selection(
                fixture.document_id, request_id="pptx-del-1",
                expected_selection_text="P3-416 mark", workspace=fixture.workspace,
            )
        self.assertTrue(ok, errors)
        self.assertTrue(data["validation_passed"])
        self.assertEqual(calls.count("selection-delete"), 1)

        with patch.object(open_documents, "_run_presentation_attached_com", return_value=com_ok(self._text_selection(text="", collapsed=True, text_range_length=0))):
            with patch.object(open_documents, "create_backup") as backup:
                ok, _data, errors, _replayed = delete_presentation_selection(
                    fixture.document_id, request_id="pptx-del-empty", workspace=fixture.workspace,
                )
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "SELECTION_EMPTY")
        backup.assert_not_called()

    def test_rejects_unsaved_before_backup(self):
        fixture = Fixture(self, component="presentation")
        with patch.object(
            open_documents,
            "_run_presentation_attached_com",
            return_value=com_ok(self._text_selection(saved=False)),
        ):
            with patch.object(open_documents, "create_backup") as backup:
                ok, data, errors, _replayed = delete_presentation_selection(
                    fixture.document_id,
                    request_id="pptx-del-unsaved",
                    expected_selection_text="P3-416 mark",
                    workspace=fixture.workspace,
                )
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "DOCUMENT_HAS_UNSAVED_CHANGES")
        self.assertIn("save-open-document", errors[0]["message"])
        self.assertFalse(data["selection"]["saved"])
        backup.assert_not_called()
        self.assertFalse((fixture.workspace / ".wps-agent" / "backups").exists())

    def test_delete_script_never_quits(self):
        captured = {}

        def fake_run(script, timeout_seconds=120, **_kwargs):
            captured["script"] = script
            return {"ok": True, "status": "no_instance"}, None

        with patch.object(open_documents, "run_powershell_script", side_effect=fake_run):
            open_documents._run_presentation_attached_com(
                "selection-delete",
                {"path": "C:\\docs\\deck.pptx", "expected_text": "mark"},
            )
        script = captured["script"]
        self.assertIn("$tr.Text = ''", script)
        self.assertIn("Find-OpenPresentation", script)
        self.assertNotIn(".Quit()", script)


class SelectionInsertTests(unittest.TestCase):
    def insert(self, fixture, text="INSERTED", request_id="ins-1", **kwargs):
        return insert_writer_selection(
            document_id=fixture.document_id, text=text, request_id=request_id, workspace=fixture.workspace, **kwargs,
        )

    def fake_attached(self, fixture, selection, write_status="ok", read_back=None, calls=None):
        def fake(action, params):
            if calls is not None:
                calls.append(action)
            if action == "selection-read":
                return com_ok(selection)
            self.assertEqual(action, "selection-insert")
            self.assertEqual(params["expected_start"], selection["start"])
            fixture.document.write_bytes(fixture.document.read_bytes() + b"\0")
            if write_status != "ok":
                return com_ok({"ok": True, "status": write_status})
            insert = params["text"]
            return com_ok({
                "ok": True,
                "status": "ok",
                "backend": "fake-wps",
                "start": selection["start"],
                "end": selection["start"] + len(insert),
                "read_back_text": insert if read_back is None else read_back,
            })
        return fake

    def test_dry_run_and_insert_with_backup_replay(self):
        fixture = Fixture(self)
        caret = selection_payload(text="", start=12, end=12)
        calls = []
        original = fixture.document.read_bytes()
        with patch.object(open_documents, "_run_writer_attached_com", side_effect=self.fake_attached(fixture, caret, calls=calls)):
            ok, data, errors, replayed = self.insert(fixture, dry_run=True, expected_start=12)
            self.assertTrue(ok, errors)
            self.assertFalse(replayed)
            self.assertEqual(calls, ["selection-read"])
            self.assertTrue(data["would_modify"])
            self.assertFalse((fixture.workspace / ".wps-agent" / "backups").exists())

            ok, data, errors, _replayed = self.insert(fixture, expected_start=12)
            self.assertTrue(ok, errors)
            self.assertTrue(data["validation_passed"])
            self.assertTrue(data["backup"]["created"])
            self.assertEqual(Path(data["backup"]["backup_path"]).read_bytes(), original)
            self.assertEqual(data["read_back_text"], "INSERTED")
            self.assertEqual(data["inserted_end"], 20)

            ok, again, errors, replayed = self.insert(fixture, expected_start=12)
            self.assertTrue(ok, errors)
            self.assertTrue(replayed)
            self.assertEqual(again["read_back_text"], "INSERTED")
        self.assertEqual(calls.count("selection-insert"), 1)

    def test_rejections_happen_before_backup_and_write(self):
        fixture = Fixture(self)
        cases = [
            ("unsaved", {"ok": True, "status": "unsaved"}, {}, "DOCUMENT_HAS_UNSAVED_CHANGES"),
            ("unsaved-flag", selection_payload(text="", start=5, end=5, saved=False), {}, "DOCUMENT_HAS_UNSAVED_CHANGES"),
            ("read-only", selection_payload(text="", start=5, end=5, read_only=True), {}, "DOCUMENT_READ_ONLY"),
            ("not-collapsed", selection_payload(text="hi", start=5), {}, "SELECTION_NOT_COLLAPSED"),
            ("start-mismatch", selection_payload(text="", start=9, end=9), {"expected_start": 3}, "SELECTION_MISMATCH"),
        ]
        for name, payload, kwargs, code in cases:
            with self.subTest(name=name):
                calls = []
                with patch.object(open_documents, "_run_writer_attached_com", side_effect=self.fake_attached(fixture, payload, calls=calls)):
                    ok, _data, errors, _replayed = self.insert(fixture, request_id=f"ins-{name}", **kwargs)
                self.assertFalse(ok)
                self.assertEqual(errors[0]["code"], code)
                self.assertNotIn("selection-insert", calls)
        self.assertFalse((fixture.workspace / ".wps-agent" / "backups").exists())

    def test_invalid_insert_text_is_rejected_before_attaching(self):
        fixture = Fixture(self)
        with patch.object(open_documents, "_run_writer_attached_com") as com:
            for index, value in enumerate(["", "a\nb", "a\x00b", "x" * 4097]):
                ok, _data, errors, _replayed = self.insert(fixture, text=value, request_id=f"ins-bad-{index}")
                self.assertFalse(ok)
                self.assertEqual(errors[0]["code"], "INVALID_ARGUMENT")
        com.assert_not_called()

    def test_insert_script_uses_range_and_never_quits(self):
        captured = {}

        def fake_run(script, timeout_seconds=120, **_kwargs):
            captured["script"] = script
            return {"ok": True, "status": "no_instance"}, None

        with patch.object(open_documents, "run_powershell_script", side_effect=fake_run):
            open_documents._run_writer_attached_com(
                "selection-insert",
                {"path": "C:\\docs\\a.docx", "text": "hi", "expected_start": 4},
            )
        script = captured["script"]
        self.assertIn("$doc.Range($start, $start)", script)
        self.assertIn("$range.Text = $insert", script)
        self.assertIn("expected_start", script)
        self.assertIn("not_collapsed", script)
        self.assertNotIn(".Quit()", script)


class ExportOpenDocumentTextTests(unittest.TestCase):
    def export(self, fixture, output, request_id="txt-1"):
        return export_open_document_text(
            document_id=fixture.document_id, output=str(output), request_id=request_id, workspace=fixture.workspace,
        )

    def fake_export(self, status="ok", write=True, touch_source=None, still_open=True, method="SaveAsUnicodeText", body=b"hello text\n"):
        def fake(action, params):
            self.assertEqual(action, "export-text")
            if write:
                Path(params["output"]).write_bytes(body)
            if touch_source is not None:
                touch_source.write_bytes(touch_source.read_bytes() + b"\0")
            return com_ok({"ok": True, "status": status, "export_method": method, "source_still_open": still_open})
        return fake

    def test_exports_text_and_verifies_source_unchanged(self):
        fixture = Fixture(self)
        output = Path(fixture.tmp.name) / "out.txt"
        before = fixture.document.read_bytes()
        with patch.object(open_documents, "_run_writer_attached_com", side_effect=self.fake_export()):
            ok, data, errors, _replayed = self.export(fixture, output)
            self.assertTrue(ok, errors)
            self.assertEqual(data["format"], "text")
            self.assertEqual(data["component"], "writer")
            self.assertTrue(data["source_unchanged"])
            self.assertTrue(data["validation_passed"])
            self.assertEqual(fixture.document.read_bytes(), before)

            ok, again, _errors, replayed = self.export(fixture, output)
            self.assertTrue(ok)
            self.assertTrue(replayed)
            self.assertEqual(again["output"], data["output"])

    def test_exports_presentation_and_spreadsheet_text(self):
        for component, runner, method in (
            ("presentation", "_run_presentation_attached_com", "ExtractShapeText"),
            ("spreadsheets", "_run_spreadsheet_attached_com", "SaveCopyAsThenUnicodeText"),
        ):
            with self.subTest(component=component):
                fixture = Fixture(self, component=component)
                output = Path(fixture.tmp.name) / f"{component}.txt"
                with patch.object(open_documents, runner, side_effect=self.fake_export(method=method)):
                    ok, data, errors, _replayed = self.export(fixture, output, request_id=f"txt-{component}")
                self.assertTrue(ok, errors)
                self.assertEqual(data["component"], component)
                self.assertEqual(data["export_method"], method)
                self.assertTrue(data["validation_passed"])

    def test_rejects_non_text_bytes_and_bad_suffix(self):
        fixture = Fixture(self)
        with patch.object(open_documents, "_run_writer_attached_com") as com:
            ok, _data, errors, _replayed = self.export(fixture, Path(fixture.tmp.name) / "out.pdf", request_id="txt-ext")
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "INVALID_ARGUMENT")
        com.assert_not_called()

        output = Path(fixture.tmp.name) / "bad.txt"
        with patch.object(open_documents, "_run_writer_attached_com", side_effect=self.fake_export(body=b"PK\x03\x04zip")):
            ok, data, errors, _replayed = self.export(fixture, output, request_id="txt-zip")
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "EXPORT_OUTPUT_NOT_TEXT")
        self.assertTrue(data["source_unchanged"])

    def test_scripts_export_text_without_quitting(self):
        captured = {}

        def fake_run(script, timeout_seconds=120, **_kwargs):
            captured["script"] = script
            return {"ok": True, "status": "no_instance"}, None

        cases = (
            ("writer", "_run_writer_attached_com", "Documents.Add", "Find-OpenWriterDocument"),
            ("presentation", "_run_presentation_attached_com", "ExtractShapeText", "Find-OpenPresentation"),
            ("spreadsheets", "_run_spreadsheet_attached_com", "SaveCopyAs", "Find-OpenWorkbook"),
        )
        for component, runner, marker, finder in cases:
            with self.subTest(component=component):
                with patch.object(open_documents, "run_powershell_script", side_effect=fake_run):
                    getattr(open_documents, runner)(
                        "export-text",
                        {"path": "C:\\docs\\source.bin", "output": "C:\\out\\source.txt"},
                    )
                script = captured["script"]
                self.assertIn(marker, script)
                self.assertIn(finder, script)
                self.assertNotIn(".Quit()", script)
                self.assertNotIn("$doc.Close", script)


class ExportOpenDocumentCopyTests(unittest.TestCase):
    def export(self, fixture, output, request_id="copy-1"):
        return export_open_document_copy(
            document_id=fixture.document_id, output=str(output), request_id=request_id, workspace=fixture.workspace,
        )

    def fake_export(self, status="ok", write=True, touch_source=None, still_open=True, method="SaveCopyAs", body=b"PK\x03\x04native"):
        def fake(action, params):
            self.assertEqual(action, "export-copy")
            if write:
                Path(params["output"]).write_bytes(body)
            if touch_source is not None:
                touch_source.write_bytes(touch_source.read_bytes() + b"\0")
            return com_ok({"ok": True, "status": status, "export_method": method, "source_still_open": still_open})
        return fake

    def test_exports_native_copy_and_verifies_source_unchanged(self):
        fixture = Fixture(self)
        output = Path(fixture.tmp.name) / "out.docx"
        before = fixture.document.read_bytes()
        with patch.object(open_documents, "_run_writer_attached_com", side_effect=self.fake_export()):
            ok, data, errors, _replayed = self.export(fixture, output)
            self.assertTrue(ok, errors)
            self.assertEqual(data["format"], "native-copy")
            self.assertEqual(data["component"], "writer")
            self.assertEqual(data["export_method"], "SaveCopyAs")
            self.assertTrue(data["source_unchanged"])
            self.assertTrue(data["source_still_open"])
            self.assertTrue(data["validation_passed"])
            self.assertEqual(fixture.document.read_bytes(), before)

            ok, again, _errors, replayed = self.export(fixture, output)
            self.assertTrue(ok)
            self.assertTrue(replayed)
            self.assertEqual(again["output"], data["output"])

    def test_exports_presentation_and_spreadsheet_copies(self):
        for component, runner, suffix in (
            ("presentation", "_run_presentation_attached_com", ".pptx"),
            ("spreadsheets", "_run_spreadsheet_attached_com", ".xlsx"),
        ):
            with self.subTest(component=component):
                fixture = Fixture(self, component=component)
                output = Path(fixture.tmp.name) / f"copy{suffix}"
                before = fixture.document.read_bytes()
                with patch.object(open_documents, runner, side_effect=self.fake_export()):
                    ok, data, errors, _replayed = self.export(fixture, output, request_id=f"copy-{component}")
                self.assertTrue(ok, errors)
                self.assertEqual(data["component"], component)
                self.assertTrue(data["validation_passed"])
                self.assertEqual(fixture.document.read_bytes(), before)

    def test_rejects_extension_mismatch_and_non_package_bytes(self):
        fixture = Fixture(self, component="spreadsheets")
        with patch.object(open_documents, "_run_spreadsheet_attached_com") as com:
            ok, _data, errors, _replayed = self.export(fixture, Path(fixture.tmp.name) / "out.pdf", request_id="copy-ext")
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "INVALID_ARGUMENT")
        com.assert_not_called()

        output = Path(fixture.tmp.name) / "bad.xlsx"
        with patch.object(open_documents, "_run_spreadsheet_attached_com", side_effect=self.fake_export(body=b"%PDF-1.4")):
            ok, data, errors, _replayed = self.export(fixture, output, request_id="copy-bad")
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "EXPORT_OUTPUT_NOT_PACKAGE")
        self.assertTrue(data["source_unchanged"])

    def test_scripts_use_save_copy_as_without_quitting(self):
        captured = {}

        def fake_run(script, timeout_seconds=120, **_kwargs):
            captured["script"] = script
            return {"ok": True, "status": "no_instance"}, None

        cases = (
            ("writer", "_run_writer_attached_com", "Find-OpenWriterDocument"),
            ("presentation", "_run_presentation_attached_com", "Find-OpenPresentation"),
            ("spreadsheets", "_run_spreadsheet_attached_com", "Find-OpenWorkbook"),
        )
        for component, runner, finder in cases:
            with self.subTest(component=component):
                with patch.object(open_documents, "run_powershell_script", side_effect=fake_run):
                    getattr(open_documents, runner)(
                        "export-copy",
                        {"path": "C:\\docs\\source.bin", "output": "C:\\out\\source.bin"},
                    )
                script = captured["script"]
                self.assertIn("SaveCopyAs", script)
                self.assertIn(finder, script)
                self.assertNotIn(".Quit()", script)
                self.assertNotIn("$doc.Close", script)


class ExportOpenDocumentPdfTests(unittest.TestCase):
    def export(self, fixture, output, request_id="pdf-1"):
        return export_open_document_pdf(
            document_id=fixture.document_id, output=str(output), request_id=request_id, workspace=fixture.workspace,
        )

    def fake_export(self, status="ok", write=True, touch_source=None, still_open=True, method="ExportAsFixedFormat", body=b"%PDF-1.4\n"):
        def fake(action, params):
            self.assertEqual(action, "export-pdf")
            if write:
                Path(params["output"]).write_bytes(body)
            if touch_source is not None:
                touch_source.write_bytes(touch_source.read_bytes() + b"\0")
            return com_ok({"ok": True, "status": status, "export_method": method, "source_still_open": still_open})
        return fake

    def test_exports_pdf_and_verifies_source_unchanged(self):
        fixture = Fixture(self)
        output = Path(fixture.tmp.name) / "out.pdf"
        before = fixture.document.read_bytes()
        with patch.object(open_documents, "_run_writer_attached_com", side_effect=self.fake_export()):
            ok, data, errors, _replayed = self.export(fixture, output)
            self.assertTrue(ok, errors)
            self.assertEqual(data["format"], "pdf")
            self.assertEqual(data["component"], "writer")
            self.assertEqual(data["export_method"], "ExportAsFixedFormat")
            self.assertTrue(data["source_unchanged"])
            self.assertTrue(data["source_still_open"])
            self.assertTrue(data["validation_passed"])
            self.assertTrue(data["known_losses"])
            self.assertEqual(fixture.document.read_bytes(), before)

            ok, again, _errors, replayed = self.export(fixture, output)
            self.assertTrue(ok)
            self.assertTrue(replayed)
            self.assertEqual(again["output"], data["output"])

    def test_exports_open_presentation_and_spreadsheet(self):
        for component, runner, method in (
            ("presentation", "_run_presentation_attached_com", "SaveCopyAsThenSaveAsPdf"),
            ("spreadsheets", "_run_spreadsheet_attached_com", "ExportAsFixedFormat"),
        ):
            with self.subTest(component=component):
                fixture = Fixture(self, component=component)
                output = Path(fixture.tmp.name) / f"{component}.pdf"
                before = fixture.document.read_bytes()
                with patch.object(open_documents, runner, side_effect=self.fake_export(method=method)):
                    ok, data, errors, _replayed = self.export(fixture, output, request_id=f"pdf-{component}")
                self.assertTrue(ok, errors)
                self.assertEqual(data["component"], component)
                self.assertEqual(data["export_method"], method)
                self.assertTrue(data["validation_passed"])
                self.assertEqual(fixture.document.read_bytes(), before)

    def test_rejects_non_pdf_bytes(self):
        fixture = Fixture(self, component="spreadsheets")
        output = Path(fixture.tmp.name) / "sheet.pdf"
        with patch.object(open_documents, "_run_spreadsheet_attached_com", side_effect=self.fake_export(body=b"PK\x03\x04not-pdf")):
            ok, data, errors, _replayed = self.export(fixture, output, request_id="pdf-zip")
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "EXPORT_OUTPUT_NOT_PDF")
        self.assertTrue(data["source_unchanged"])
        self.assertFalse(data["validation_passed"])

    def test_invalid_outputs_are_rejected_without_calling_wps(self):
        fixture = Fixture(self)
        existing = Path(fixture.tmp.name) / "exists.pdf"
        existing.write_bytes(b"%PDF-1.4 keep")
        cases = [
            (fixture.document, "INVALID_ARGUMENT"),
            (Path(fixture.tmp.name) / "out.html", "INVALID_ARGUMENT"),
            (existing, "OUTPUT_EXISTS"),
            (Path(fixture.tmp.name) / "missing-dir" / "out.pdf", "INVALID_ARGUMENT"),
        ]
        with patch.object(open_documents, "_run_writer_attached_com") as com:
            for index, (output, code) in enumerate(cases):
                with self.subTest(output=str(output)):
                    ok, _data, errors, _replayed = self.export(fixture, output, request_id=f"pdf-bad-{index}")
                    self.assertFalse(ok)
                    self.assertEqual(errors[0]["code"], code)
        com.assert_not_called()
        self.assertTrue(existing.read_bytes().startswith(b"%PDF"))

    def test_changed_source_is_a_failure(self):
        fixture = Fixture(self)
        output = Path(fixture.tmp.name) / "changed.pdf"
        with patch.object(open_documents, "_run_writer_attached_com", side_effect=self.fake_export(touch_source=fixture.document)):
            ok, data, errors, _replayed = self.export(fixture, output, request_id="pdf-changed")
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "DOCUMENT_CHANGED_AFTER_COM")
        self.assertFalse(data["source_unchanged"])

    def test_scripts_export_pdf_without_quitting_or_renaming_the_source(self):
        captured = {}

        def fake_run(script, timeout_seconds=120, **_kwargs):
            captured["script"] = script
            return {"ok": True, "status": "no_instance"}, None

        cases = (
            ("writer", "_run_writer_attached_com", "Documents.Add", "Find-OpenWriterDocument"),
            ("presentation", "_run_presentation_attached_com", "SaveCopyAs", "Find-OpenPresentation"),
            ("spreadsheets", "_run_spreadsheet_attached_com", "ExportAsFixedFormat", "Find-OpenWorkbook"),
        )
        for component, runner, marker, finder in cases:
            with self.subTest(component=component):
                with patch.object(open_documents, "run_powershell_script", side_effect=fake_run):
                    getattr(open_documents, runner)(
                        "export-pdf",
                        {"path": "C:\\docs\\source.bin", "output": "C:\\out\\source.pdf"},
                    )
                script = captured["script"]
                self.assertIn("ExportAsFixedFormat", script)
                self.assertIn(marker, script)
                self.assertIn(finder, script)
                self.assertNotIn(".Quit()", script)
                self.assertNotIn("$doc.SaveAs", script)
                self.assertNotIn("$doc.Close", script)


class SaveOpenDocumentTests(unittest.TestCase):
    def fake_status(self, saved=False, read_only=False, wrote=None, calls=None, touch=None):
        def fake(action, params):
            if calls is not None:
                calls.append((action, bool(params.get("do_save"))))
            self.assertEqual(action, "save")
            if read_only:
                return com_ok({"ok": True, "status": "read_only", "saved": saved, "read_only": True})
            if not params.get("do_save"):
                return com_ok({
                    "ok": True, "status": "ok", "backend": "fake-wps",
                    "saved": saved, "read_only": False, "wrote": False, "source_still_open": True,
                })
            if touch is not None:
                touch.write_bytes(touch.read_bytes() + b"\0")
            saved_after = True
            wrote_after = wrote if wrote is not None else (not saved)
            return com_ok({
                "ok": True, "status": "ok", "backend": "fake-wps",
                "saved": saved_after, "read_only": False, "wrote": wrote_after, "source_still_open": True,
            })
        return fake

    def test_dry_run_and_save_with_backup_replay(self):
        fixture = Fixture(self)
        calls = []
        original = fixture.document.read_bytes()
        with patch.object(
            open_documents,
            "_run_writer_attached_com",
            side_effect=self.fake_status(saved=False, calls=calls, touch=fixture.document),
        ):
            ok, data, errors, replayed = save_open_document(
                fixture.document_id, request_id="save-1", dry_run=True, workspace=fixture.workspace,
            )
            self.assertTrue(ok, errors)
            self.assertFalse(replayed)
            self.assertTrue(data["would_modify"])
            self.assertEqual(calls, [("save", False)])

            ok, data, errors, _replayed = save_open_document(
                fixture.document_id, request_id="save-1", workspace=fixture.workspace,
            )
            self.assertTrue(ok, errors)
            self.assertTrue(data["validation_passed"])
            self.assertTrue(data["wrote"])
            self.assertTrue(data["backup"]["created"])
            self.assertEqual(Path(data["backup"]["backup_path"]).read_bytes(), original)

            ok, again, errors, replayed = save_open_document(
                fixture.document_id, request_id="save-1", workspace=fixture.workspace,
            )
            self.assertTrue(ok, errors)
            self.assertTrue(replayed)
            self.assertTrue(again["wrote"])
        self.assertEqual([item for item in calls if item[1]], [("save", True)])

    def test_already_saved_skips_backup(self):
        fixture = Fixture(self)
        with patch.object(open_documents, "_run_writer_attached_com", side_effect=self.fake_status(saved=True)):
            with patch.object(open_documents, "create_backup") as backup:
                ok, data, errors, _replayed = save_open_document(
                    fixture.document_id, request_id="save-clean", workspace=fixture.workspace,
                )
        self.assertTrue(ok, errors)
        self.assertFalse(data["would_modify"])
        self.assertFalse(data["wrote"])
        backup.assert_not_called()

    def test_read_only_rejected_before_backup(self):
        fixture = Fixture(self)
        with patch.object(open_documents, "_run_writer_attached_com", side_effect=self.fake_status(read_only=True)):
            with patch.object(open_documents, "create_backup") as backup:
                ok, _data, errors, _replayed = save_open_document(
                    fixture.document_id, request_id="save-ro", workspace=fixture.workspace,
                )
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "DOCUMENT_READ_ONLY")
        backup.assert_not_called()

    def test_save_script_never_quits(self):
        captured = {}

        def fake_run(script, timeout_seconds=120, **_kwargs):
            captured["script"] = script
            return {"ok": True, "status": "no_instance"}, None

        with patch.object(open_documents, "run_powershell_script", side_effect=fake_run):
            open_documents._run_writer_attached_com("save", {"path": "C:\\docs\\open.docx", "do_save": True})
        script = captured["script"]
        self.assertIn("$doc.Save()", script)
        self.assertIn("Find-OpenWriterDocument", script)
        self.assertNotIn(".Quit()", script)


class SaveAsOpenDocumentTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.workspace = Path(self.tmp.name) / "workspace"
        self.workspace.mkdir()
        self.output = Path(self.tmp.name) / "untitled-saved.docx"

    def fake_com(self, *, status="ok", do_save_effect=None, **fields):
        def fake(component, name, output, *, do_save):
            self.assertEqual(component, "writer")
            self.assertEqual(name, "Document1")
            base = {
                "ok": True,
                "status": status,
                "backend": "fake-wps",
                "component": "writer",
                "name": name,
                "saved": False,
                "read_only": False,
                "would_modify": True,
                "wrote": False,
                "source_still_open": True,
            }
            base.update(fields)
            if status != "ok":
                return com_ok(base)
            if not do_save:
                return com_ok(base)
            if do_save_effect is not None:
                do_save_effect(output)
            return com_ok({
                **base,
                "wrote": True,
                "saved": True,
                "full_name": str(Path(output).resolve()),
                "path_matched": True,
                "output_exists": True,
                "source_still_open": True,
            })
        return fake

    def test_dry_run_write_register_and_replay(self):
        calls = []

        def effect(output):
            Path(output).write_bytes(b"PK\x03\x04untitled")

        def tracking(component, name, output, *, do_save):
            calls.append(do_save)
            return self.fake_com(do_save_effect=effect)(component, name, output, do_save=do_save)

        with patch.object(open_documents, "_run_save_as_untitled_com", side_effect=tracking):
            ok, data, errors, replayed = save_as_open_document(
                "writer",
                "Document1",
                str(self.output),
                request_id="save-as-1",
                dry_run=True,
                workspace=self.workspace,
            )
            self.assertTrue(ok, errors)
            self.assertFalse(replayed)
            self.assertTrue(data["would_modify"])
            self.assertIsNone(data["backup"])
            self.assertEqual(calls, [False])
            self.assertFalse(self.output.exists())

            ok, data, errors, _replayed = save_as_open_document(
                "writer",
                "Document1",
                str(self.output),
                request_id="save-as-1",
                workspace=self.workspace,
            )
            self.assertTrue(ok, errors)
            self.assertTrue(data["validation_passed"])
            self.assertTrue(data["registered"])
            self.assertTrue(data["document_id"])
            self.assertTrue(self.output.is_file())
            self.assertEqual(calls, [False, False, True])

            ok, again, errors, replayed = save_as_open_document(
                "writer",
                "Document1",
                str(self.output),
                request_id="save-as-1",
                workspace=self.workspace,
            )
            self.assertTrue(ok, errors)
            self.assertTrue(replayed)
            self.assertEqual(again["document_id"], data["document_id"])
        self.assertEqual(calls.count(True), 1)

    def test_rejects_existing_output_and_wrong_extension(self):
        self.output.write_bytes(b"exists")
        with patch.object(open_documents, "_run_save_as_untitled_com") as com:
            ok, _data, errors, _replayed = save_as_open_document(
                "writer", "Document1", str(self.output), request_id="save-as-exists", workspace=self.workspace,
            )
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "OUTPUT_EXISTS")
        com.assert_not_called()

        missing = Path(self.tmp.name) / "bad.txt"
        with patch.object(open_documents, "_run_save_as_untitled_com") as com:
            ok, _data, errors, _replayed = save_as_open_document(
                "writer", "Document1", str(missing), request_id="save-as-ext", workspace=self.workspace,
            )
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "UNSUPPORTED_OUTPUT_EXTENSION")
        com.assert_not_called()

    def test_rejects_already_pathed_and_ambiguous(self):
        with patch.object(
            open_documents,
            "_run_save_as_untitled_com",
            return_value=com_ok({"ok": True, "status": "already_pathed", "match_count": 1, "name": "Document1"}),
        ):
            ok, _data, errors, _replayed = save_as_open_document(
                "writer", "Document1", str(self.output), request_id="save-as-pathed", workspace=self.workspace,
            )
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "DOCUMENT_ALREADY_PATHED")

        with patch.object(
            open_documents,
            "_run_save_as_untitled_com",
            return_value=com_ok({"ok": True, "status": "ambiguous", "match_count": 2, "name": "Document1"}),
        ):
            ok, _data, errors, _replayed = save_as_open_document(
                "writer", "Document1", str(self.output), request_id="save-as-amb", workspace=self.workspace,
            )
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "AMBIGUOUS_OPEN_DOCUMENT")

    def test_save_as_script_never_quits(self):
        captured = {}

        def fake_run(script, timeout_seconds=120, **_kwargs):
            captured["script"] = script
            return {"ok": True, "status": "no_instance"}, None

        with patch.object(open_documents, "run_powershell_script", side_effect=fake_run):
            open_documents._run_save_as_untitled_com("writer", "Document1", "C:\\out\\a.docx", do_save=True)
        script = captured["script"]
        self.assertIn("SaveAs", script)
        self.assertIn("already_pathed", script)
        self.assertIn("ambiguous", script)
        self.assertNotIn(".Quit()", script)


class CliWiringTests(unittest.TestCase):
    def run_cli(self, argv):
        output = io.StringIO()
        code = cli.run(argv, output_stream=output)
        return code, json.loads(output.getvalue())

    def test_open_documents_command_reports_missing_instance(self):
        with patch.object(open_documents, "_run_open_documents_com", return_value=com_ok({"ok": True, "instances": []})):
            code, response = self.run_cli(["open-documents", "--component", "writer", "--request-id", "cli-open"])
        self.assertEqual(code, 0)
        self.assertFalse(response["ok"])
        self.assertEqual(response["errors"][0]["code"], "NO_RUNNING_WPS_INSTANCE")
        self.assertEqual(response["request_id"], "cli-open")

    def test_new_commands_are_exposed_as_mcp_tools(self):
        from wps_ai_agent_cli.mcp_schema import get_mcp_tool_schema

        expected = {
            "open-documents": False,
            "writer-selection-read": False,
            "presentation-selection-read": False,
            "spreadsheet-selection-read": False,
            "spreadsheet-selection-replace": True,
            "spreadsheet-selection-delete": True,
            "presentation-selection-replace": True,
            "presentation-selection-insert": True,
            "writer-selection-replace": True,
            "writer-selection-insert": True,
            "writer-selection-delete": True,
            "presentation-selection-delete": True,
            "save-open-document": True,
            "save-as-open-document": True,
            "export-open-document": False,
            "export-open-document-pdf": False,
            "export-open-document-copy": False,
            "export-open-document-text": False,
        }
        for command, mutates in expected.items():
            schema = get_mcp_tool_schema(command)
            self.assertIsNotNone(schema, command)
            self.assertTrue(schema["requires_wps"])
            self.assertEqual(schema["mutates_document"], mutates)
            self.assertTrue(schema["safety_notes"])


if __name__ == "__main__":
    unittest.main()
