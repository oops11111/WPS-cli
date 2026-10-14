import io
import json
import os
import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from tests.test_wps_quit_guard import CAPABILITIES
from wps_ai_agent_cli import cli, mcp_adapter, powershell_runner, spreadsheet_ops
from wps_ai_agent_cli.powershell_runner import collect_script_warnings, merge_warnings

WARNING = "3 et processes are running but only 1 has a window; the others are assumed to be background processes."


def run_cli(argv):
    stream = io.StringIO()
    code = cli.run(argv, stream)
    return code, json.loads(stream.getvalue())


class MergeWarningsTests(unittest.TestCase):
    def test_merge_adds_deduplicates_and_keeps_existing_entries(self):
        self.assertEqual(merge_warnings({"a": 1}, ["x", "y"]), {"a": 1, "warnings": ["x", "y"]})
        self.assertEqual(merge_warnings({"warnings": ["x"]}, ["x", "y"]), {"warnings": ["x", "y"]})
        self.assertEqual(merge_warnings({"warnings": "x"}, ["y"]), {"warnings": ["x", "y"]})

    def test_merge_leaves_data_alone_without_warnings_or_a_dict(self):
        data = {"a": 1}
        self.assertIs(merge_warnings(data, []), data)
        self.assertEqual(merge_warnings(None, ["x"]), None)

    def test_nested_collectors_hand_their_warnings_to_the_outer_one(self):
        with collect_script_warnings() as outer:
            with collect_script_warnings() as inner:
                powershell_runner._remember_warnings(["one"])
            powershell_runner._remember_warnings(["one", "two"])
        self.assertEqual(inner, ["one"])
        self.assertEqual(outer, ["one", "two"])


class PublicCommandWarningTests(unittest.TestCase):
    """The warnings a WPS script prints must reach the envelope of the public commands, not only the private runners."""

    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.dir = Path(tmp.name)
        previous = os.getcwd()
        os.chdir(self.dir)
        self.addCleanup(os.chdir, previous)
        self.document = self.dir / "book.xlsx"
        workbook = Workbook()
        workbook.active.title = "First"
        workbook.create_sheet("Second")
        workbook.save(self.document)
        workbook.close()
        _code, payload = run_cli(["register-document", "--component", "spreadsheets", "--path", str(self.document)])
        self.document_id = payload["data"]["document"]["document_id"]
        self.stderr = f"WPS_AGENT_WARNING: {WARNING}\n"
        self.returncode = 0
        self.payload = {"ok": True, "backend": "powershell-com"}
        self.edit_file = lambda: None
        for patcher in (
            patch.object(powershell_runner.subprocess, "run", side_effect=self.fake_run),
            patch.object(powershell_runner, "powershell_executable", return_value="pwsh"),
            patch.object(spreadsheet_ops, "probe_wps_capabilities", return_value=CAPABILITIES),
        ):
            patcher.start()
            self.addCleanup(patcher.stop)

    def fake_run(self, args, **kwargs):
        self.edit_file()
        return subprocess.CompletedProcess(args, self.returncode, stdout=json.dumps(self.payload), stderr=self.stderr)

    def rename(self):
        def edit():
            workbook = load_workbook(self.document)
            workbook["First"].title = "Renamed"
            workbook.save(self.document)
            workbook.close()

        self.edit_file = edit
        self.payload = {"ok": True, "backend": "powershell-com", "sheet_names": ["Renamed", "Second"]}
        return [
            "spreadsheet-rename-sheet",
            "--document-id",
            self.document_id,
            "--old-name",
            "First",
            "--new-name",
            "Renamed",
        ]

    def write(self):
        def edit():
            workbook = load_workbook(self.document)
            workbook["First"]["A1"] = 5
            workbook.save(self.document)
            workbook.close()

        self.edit_file = edit
        return ["spreadsheet-write", "--document-id", self.document_id, "--range", "A1:A1", "--values-json", "[[5]]"]

    def test_rename_sheet_envelope_carries_the_script_warnings(self):
        code, payload = run_cli(self.rename())
        self.assertEqual(code, 0, payload)
        self.assertTrue(payload["ok"], payload)
        self.assertTrue(payload["data"]["validation_passed"])
        self.assertEqual(payload["data"]["warnings"], [WARNING])

    def test_spreadsheet_write_envelope_carries_the_script_warnings(self):
        code, payload = run_cli(self.write())
        self.assertEqual(code, 0, payload)
        self.assertTrue(payload["ok"], payload)
        self.assertTrue(payload["data"]["validation_passed"])
        self.assertEqual(payload["data"]["warnings"], [WARNING])

    def test_a_command_that_warned_and_then_failed_still_reports_the_warning(self):
        self.returncode = 2
        self.payload = {"ok": False, "error_message": "boom"}
        code, payload = run_cli(self.rename())
        self.assertFalse(payload["ok"], payload)
        self.assertEqual(payload["data"]["warnings"], [WARNING])

    def test_no_warnings_key_when_the_scripts_reported_none(self):
        self.stderr = ""
        _code, payload = run_cli(self.rename())
        self.assertTrue(payload["ok"], payload)
        self.assertNotIn("warnings", payload["data"])

    def test_mcp_tool_output_carries_the_same_warnings(self):
        argv = self.rename()
        arguments = {"document_id": self.document_id, "old_name": "First", "new_name": "Renamed"}
        ok, result, errors = mcp_adapter.call_mcp_tool(argv[0], arguments)
        self.assertTrue(ok, errors)
        self.assertEqual(result["response"]["data"]["warnings"], [WARNING])


class InspectWarningTests(unittest.TestCase):
    def test_inspect_result_carries_the_script_warnings(self):
        from wps_ai_agent_cli import spreadsheet_inspect

        with tempfile.TemporaryDirectory() as tmp:
            book = Path(tmp) / "book.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = 1
            workbook.save(book)
            workbook.close()
            previous = os.getcwd()
            os.chdir(tmp)
            try:
                _code, registered = run_cli(["register-document", "--component", "spreadsheets", "--path", str(book)])
                document_id = registered["data"]["document"]["document_id"]
                payload = {
                    "ok": True,
                    "backend": "powershell-com",
                    "sheet": "Sheet",
                    "full_recalculation": False,
                    "cells": [
                        {
                            "address": "A1",
                            "formula": None,
                            "recalculated_value": 1,
                            "recalculated_type": "number",
                            "displayed_text": "1",
                            "number_format": "General",
                        }
                    ],
                }
                completed = subprocess.CompletedProcess(
                    [], 0, stdout=json.dumps(payload), stderr=f"WPS_AGENT_WARNING: {WARNING}\n"
                )
                with (
                    patch.object(powershell_runner.subprocess, "run", return_value=completed),
                    patch.object(powershell_runner, "powershell_executable", return_value="pwsh"),
                    patch.object(spreadsheet_inspect, "probe_wps_capabilities", return_value=CAPABILITIES),
                ):
                    ok, result, errors = spreadsheet_inspect.inspect_spreadsheet_range(document_id, "A1:A1")
            finally:
                os.chdir(previous)
        self.assertTrue(ok, errors)
        self.assertEqual(result["warnings"], [WARNING])


if __name__ == "__main__":
    unittest.main()
