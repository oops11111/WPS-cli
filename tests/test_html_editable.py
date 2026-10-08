from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
import os

from docx import Document
from PIL import Image

from wps_ai_agent_cli.html_editable import convert_html_editable
from wps_ai_agent_cli.html_roundtrip import export_controlled_html, build_html_roundtrip_mapping, import_controlled_html, verify_controlled_document
from wps_ai_agent_cli.com_backend import run_com_smoke


class HtmlEditableTests(unittest.TestCase):
    def test_maps_semantics_to_editable_docx_and_reports_css_and_remote_image_loss(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            Image.new("RGB", (40, 20), (210, 30, 40)).save(root / "local.png")
            source = root / "page.html"
            source.write_text(
                '<!doctype html><html><head><title>Project brief</title><style>h1{color:red}</style></head>'
                '<body><h1 class="hero">Plan</h1><p>Hello <a href="https://example.com">Example</a></p>'
                '<ul><li>One</li><li>Two</li></ul><table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table>'
                '<p><img src="local.png" alt="Local"><img src="https://example.invalid/x.png" alt="Remote">'
                '<a href="javascript:alert(1)">Unsafe</a></p>'
                '<script>throw new Error("must not execute")</script></body></html>',
                encoding="utf-8",
            )
            ok, data, errors = convert_html_editable(source, root / "editable.docx")
            self.assertTrue(ok, errors)
            self.assertEqual(data["mapped_objects"], {"paragraphs": 2, "headings": 1, "lists": 2, "tables": 1, "images": 1})
            self.assertIn("embedded CSS", data["unsupported_css"])
            self.assertTrue(any("non-local" in warning for warning in data["warnings"]))
            self.assertIn("Unsafe hyperlink target omitted.", data["warnings"])
            self.assertFalse(data["javascript_executed"])
            document = Document(root / "editable.docx")
            self.assertEqual(document.core_properties.title, "Project brief")
            self.assertEqual(document.paragraphs[0].style.name, "Heading 1")
            self.assertIn("List Bullet", {paragraph.style.name for paragraph in document.paragraphs})
            self.assertEqual(len(document.tables), 1)
            self.assertEqual(document.tables[0].cell(1, 1).text, "2")
            self.assertEqual(len(document.inline_shapes), 1)
            self.assertTrue(any(getattr(rel, "target_ref", "") == "https://example.com" for rel in document.part.rels.values()))
            self.assertFalse(any("javascript:" in getattr(rel, "target_ref", "") for rel in document.part.rels.values()))

    def test_refuses_existing_output_without_modifying_it(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source, output = root / "source.html", root / "existing.docx"
            source.write_text("<p>Text</p>", encoding="utf-8")
            output.write_bytes(b"keep")
            ok, _, errors = convert_html_editable(source, output)
            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "OUTPUT_ALREADY_EXISTS")
            self.assertEqual(output.read_bytes(), b"keep")


@unittest.skipUnless(os.environ.get("WPS_AGENT_RUN_INTEGRATION") == "1", "set WPS_AGENT_RUN_INTEGRATION=1 to launch local WPS Writer")
class HtmlEditableWpsIntegrationTests(unittest.TestCase):
    def test_wps_writer_opens_and_saves_semantically_converted_docx(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source, editable, saved_copy, exported = root / "page.html", root / "page.docx", root / "wps-copy.docx", root / "wps-copy.html"
            source.write_text(
                '<html data-wps-schema="wps-agent-html/v1"><h1 data-wps-object-id="heading">WPS validation</h1>'
                '<p data-wps-object-id="paragraph">Native editable paragraph</p><table data-wps-object-id="table">'
                '<tr data-wps-object-id="row"><td data-wps-object-id="cell">cell</td></tr></table></html>',
                encoding="utf-8",
            )
            ok, data, errors = import_controlled_html(source, editable)
            self.assertTrue(ok, errors)
            result = run_com_smoke("writer", str(editable), str(saved_copy), visible=False)
            self.assertTrue(result["ok"], result)
            self.assertTrue(saved_copy.is_file())
            ok, verified, errors = verify_controlled_document(saved_copy, data["mapping_path"])
            self.assertTrue(ok, errors)
            self.assertEqual(verified["object_count"], 5)
            ok, exported_data, errors = export_controlled_html(saved_copy, data["mapping_path"], exported)
            self.assertTrue(ok, errors)
            self.assertEqual(exported_data["exported_object_count"], 5)
            valid, mapping, errors = build_html_roundtrip_mapping(exported)
            self.assertTrue(valid, errors)
            self.assertEqual({entry["object_id"] for entry in mapping["mappings"]}, {"heading", "paragraph", "table", "row", "cell"})

    def test_wps_rich_media_controlled_roundtrip_preserves_link_and_image_ids(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            Image.new("RGB", (24, 16), (34, 120, 180)).save(root / "mark.png")
            source, editable = root / "media.html", root / "media.docx"
            saved_copy, exported = root / "media-wps.docx", root / "media-roundtrip.html"
            source.write_text(
                '<html data-wps-schema="wps-agent-html/v1"><p data-wps-object-id="p1">'
                'Read <a data-wps-object-id="link1" href="https://example.com/reference">reference</a> '
                '<img data-wps-object-id="image1" src="mark.png" alt="Blue mark"></p></html>',
                encoding="utf-8",
            )
            ok, data, errors = import_controlled_html(source, editable)
            self.assertTrue(ok, errors)
            saved = run_com_smoke("writer", str(editable), str(saved_copy), visible=False)
            self.assertTrue(saved["ok"], saved)
            ok, verified, errors = verify_controlled_document(saved_copy, data["mapping_path"])
            self.assertTrue(ok, errors)
            self.assertEqual(verified["object_count"], 3)
            ok, _, errors = export_controlled_html(saved_copy, data["mapping_path"], exported)
            self.assertTrue(ok, errors)
            valid, mapping, errors = build_html_roundtrip_mapping(exported)
            self.assertTrue(valid, errors)
            self.assertEqual({entry["object_id"] for entry in mapping["mappings"]}, {"p1", "link1", "image1"})
            exported_html = exported.read_text(encoding="utf-8")
            self.assertIn("data:image/png;base64,", exported_html)
            link = next(entry for entry in mapping["mappings"] if entry["object_id"] == "link1")
            self.assertEqual(link.get("href"), "https://example.com/reference")

    def test_wps_lists_and_multirow_table_keep_identity_and_order(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source, editable = root / "structure.html", root / "structure.docx"
            saved_copy, exported = root / "structure-wps.docx", root / "structure-roundtrip.html"
            source.write_text(
                '<html data-wps-schema="wps-agent-html/v1"><h1 data-wps-object-id="title">Structure</h1>'
                '<ul><li data-wps-object-id="bullet1">Alpha</li><li data-wps-object-id="bullet2">Beta</li></ul>'
                '<ol><li data-wps-object-id="number1">One</li><li data-wps-object-id="number2">Two</li></ol>'
                '<table data-wps-object-id="table"><tr data-wps-object-id="row1"><th data-wps-object-id="head1">Name</th></tr>'
                '<tr data-wps-object-id="row2"><td data-wps-object-id="cell1">Value</td></tr></table></html>',
                encoding="utf-8",
            )
            ok, data, errors = import_controlled_html(source, editable)
            self.assertTrue(ok, errors)
            saved = run_com_smoke("writer", str(editable), str(saved_copy), visible=False)
            self.assertTrue(saved["ok"], saved)
            ok, identity, errors = verify_controlled_document(saved_copy, data["mapping_path"])
            self.assertTrue(ok, errors)
            self.assertEqual(identity["object_count"], 10)
            ok, _, errors = export_controlled_html(saved_copy, data["mapping_path"], exported)
            self.assertTrue(ok, errors)
            valid, mapping, errors = build_html_roundtrip_mapping(exported)
            self.assertTrue(valid, errors)
            by_id = {item["object_id"]: item for item in mapping["mappings"]}
            self.assertEqual(set(by_id), {"title", "bullet1", "bullet2", "number1", "number2", "table", "row1", "head1", "row2", "cell1"})
            self.assertEqual([by_id[name]["native_index"] for name in ("bullet1", "bullet2", "number1", "number2")], [1, 2, 3, 4])
            self.assertEqual(by_id["head1"]["object_type"], "table_header")
            self.assertEqual(by_id["cell1"]["object_type"], "table_cell")


if __name__ == "__main__":
    unittest.main()
