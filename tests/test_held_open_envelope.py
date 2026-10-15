"""Gaps A/B: exclusive FileShare.None → held-open / Open() was not called JSON envelope.

Refs: docs/wps-real-verify-remaining/RESULTS-973590b.md
      internal/results-973590b-checklist-verdict.md (store)
"""

from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from tests.test_writer_structure import write_document
from wps_ai_agent_cli.backups import create_backup
from wps_ai_agent_cli.errors import DOCUMENT_HELD_OPEN
from wps_ai_agent_cli.file_hold import (
    held_open_error,
    is_share_conflict,
    probe_file_not_held,
    share_conflict_error,
)
from wps_ai_agent_cli.presentation_ops import presentation_replace
from wps_ai_agent_cli.sessions import register_document
from wps_ai_agent_cli.spreadsheet_inspect import inspect_spreadsheet_range
from wps_ai_agent_cli.spreadsheet_ops import write_spreadsheet_range
from wps_ai_agent_cli.writer_ops import writer_replace


def _assert_held_open_envelope(test: unittest.TestCase, errors: list) -> None:
    test.assertEqual(len(errors), 1, errors)
    test.assertEqual(errors[0]["code"], DOCUMENT_HELD_OPEN)
    message = errors[0]["message"]
    test.assertIn("held open", message)
    test.assertIn("Open() was not called", message)


def _hold_exclusive_win32(path: Path):
    """Open with FileShare.None via CreateFileW; return a closer."""
    import ctypes
    from ctypes import wintypes

    GENERIC_READ = 0x80000000
    FILE_SHARE_NONE = 0
    OPEN_EXISTING = 3
    FILE_ATTRIBUTE_NORMAL = 0x80
    INVALID = 0xFFFFFFFF

    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel32.CreateFileW.argtypes = [
        wintypes.LPCWSTR,
        wintypes.DWORD,
        wintypes.DWORD,
        wintypes.LPVOID,
        wintypes.DWORD,
        wintypes.DWORD,
        wintypes.HANDLE,
    ]
    kernel32.CreateFileW.restype = wintypes.HANDLE
    handle = kernel32.CreateFileW(
        str(path),
        GENERIC_READ,
        FILE_SHARE_NONE,
        None,
        OPEN_EXISTING,
        FILE_ATTRIBUTE_NORMAL,
        None,
    )
    if handle == INVALID or handle == wintypes.HANDLE(-1).value:
        raise OSError(ctypes.get_last_error(), ctypes.FormatError(ctypes.get_last_error()))

    def close():
        kernel32.CloseHandle(handle)

    return close


def _hold_read_share_win32(path: Path):
    """Open with FileShare.Read; return a closer."""
    import ctypes
    from ctypes import wintypes

    GENERIC_READ = 0x80000000
    FILE_SHARE_READ = 0x00000001
    OPEN_EXISTING = 3
    FILE_ATTRIBUTE_NORMAL = 0x80
    INVALID = 0xFFFFFFFF

    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel32.CreateFileW.argtypes = [
        wintypes.LPCWSTR,
        wintypes.DWORD,
        wintypes.DWORD,
        wintypes.LPVOID,
        wintypes.DWORD,
        wintypes.DWORD,
        wintypes.HANDLE,
    ]
    kernel32.CreateFileW.restype = wintypes.HANDLE
    handle = kernel32.CreateFileW(
        str(path),
        GENERIC_READ,
        FILE_SHARE_READ,
        None,
        OPEN_EXISTING,
        FILE_ATTRIBUTE_NORMAL,
        None,
    )
    if handle == INVALID or handle == wintypes.HANDLE(-1).value:
        raise OSError(ctypes.get_last_error(), ctypes.FormatError(ctypes.get_last_error()))

    def close():
        kernel32.CloseHandle(handle)

    return close


class FileHoldHelperTests(unittest.TestCase):
    def test_held_open_message_matches_preopen_guard_contract(self):
        error = held_open_error("document", "sharing violation")
        self.assertEqual(error["code"], DOCUMENT_HELD_OPEN)
        self.assertIn("held open by another handle", error["message"])
        self.assertIn("Open() was not called", error["message"])
        self.assertIn("alerts are turned off", error["message"])

    def test_permission_error_is_treated_as_share_conflict(self):
        self.assertTrue(is_share_conflict(PermissionError("Permission denied")))
        self.assertEqual(share_conflict_error(PermissionError("x"), "document")["code"], DOCUMENT_HELD_OPEN)

    def test_free_file_passes_probe_in_both_modes(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "free.bin"
            path.write_bytes(b"x")
            self.assertIsNone(probe_file_not_held(path, read_only=True))
            self.assertIsNone(probe_file_not_held(path, read_only=False))

    def test_missing_file_is_left_to_later_checks(self):
        with tempfile.TemporaryDirectory() as tmp:
            missing = Path(tmp) / "missing.xlsx"
            self.assertIsNone(probe_file_not_held(missing, read_only=True))
            self.assertIsNone(probe_file_not_held(missing, read_only=False))


class GapAInspectEnvelopeTests(unittest.TestCase):
    def test_inspect_refuses_exclusive_hold_with_held_open_json_before_zip(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "book.xlsx"
            Workbook().save(document)
            _, record, _ = register_document("spreadsheets", str(document), workspace)
            held = held_open_error("document", "exclusive test hold")
            with (
                patch("wps_ai_agent_cli.spreadsheet_inspect.probe_file_not_held", return_value=held) as probe,
                patch("wps_ai_agent_cli.spreadsheet_inspect.load_workbook") as load_workbook,
                patch("wps_ai_agent_cli.spreadsheet_inspect.probe_wps_capabilities") as caps,
            ):
                ok, _result, errors = inspect_spreadsheet_range(
                    record["document_id"],
                    "A1",
                    workspace=workspace,
                )
            self.assertFalse(ok)
            _assert_held_open_envelope(self, errors)
            probe.assert_called_once()
            self.assertEqual(probe.call_args.kwargs["read_only"], True)
            load_workbook.assert_not_called()
            caps.assert_not_called()

    def test_inspect_maps_zip_permission_error_to_held_open_envelope(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "book.xlsx"
            Workbook().save(document)
            _, record, _ = register_document("spreadsheets", str(document), workspace)
            with (
                patch("wps_ai_agent_cli.spreadsheet_inspect.probe_file_not_held", return_value=None),
                patch(
                    "wps_ai_agent_cli.spreadsheet_inspect.load_workbook",
                    side_effect=PermissionError("Permission denied"),
                ),
                patch("wps_ai_agent_cli.spreadsheet_inspect.probe_wps_capabilities") as caps,
            ):
                ok, _result, errors = inspect_spreadsheet_range(
                    record["document_id"],
                    "A1",
                    workspace=workspace,
                )
            self.assertFalse(ok)
            _assert_held_open_envelope(self, errors)
            caps.assert_not_called()

    def test_inspect_read_only_probe_does_not_false_refuse_when_probe_passes(self):
        """FileShare.Read holders must not be treated as exclusive holds at the Zip stage."""
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "book.xlsx"
            Workbook().save(document)
            _, record, _ = register_document("spreadsheets", str(document), workspace)
            fake_book = Workbook()
            with (
                patch("wps_ai_agent_cli.spreadsheet_inspect.probe_file_not_held", return_value=None) as probe,
                patch("wps_ai_agent_cli.spreadsheet_inspect.load_workbook", return_value=fake_book),
                patch(
                    "wps_ai_agent_cli.spreadsheet_inspect.probe_wps_capabilities",
                    return_value={"components": {"spreadsheets": {"selected_prog_id": None}}},
                ),
            ):
                ok, _result, errors = inspect_spreadsheet_range(
                    record["document_id"],
                    "A1",
                    workspace=workspace,
                )
            self.assertEqual(probe.call_args.kwargs["read_only"], True)
            # No ProgID → COM_BACKEND_UNAVAILABLE after Zip, not held-open.
            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "COM_BACKEND_UNAVAILABLE")
            self.assertNotIn("held open", errors[0]["message"])


class GapBWriteEnvelopeTests(unittest.TestCase):
    def test_create_backup_refuses_exclusive_hold_before_identity(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "book.xlsx"
            Workbook().save(document)
            _, record, _ = register_document("spreadsheets", str(document), workspace)
            held = held_open_error("document", "exclusive test hold")
            with (
                patch("wps_ai_agent_cli.backups.probe_file_not_held", return_value=held) as probe,
                patch("wps_ai_agent_cli.backups.file_identity") as identity,
            ):
                ok, _result, errors, _ = create_backup(record["document_id"], "b-held", workspace=workspace)
            self.assertFalse(ok)
            _assert_held_open_envelope(self, errors)
            self.assertEqual(probe.call_args.kwargs["read_only"], False)
            identity.assert_not_called()

    def test_a_backup_for_a_command_working_through_wps_skips_the_exclusive_probe(self):
        # The document of an open-document command is open in WPS by definition; WPS lets others read it (P6-008).
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "book.xlsx"
            Workbook().save(document)
            _, record, _ = register_document("spreadsheets", str(document), workspace)
            held = held_open_error("document", "WPS has it open")
            with patch("wps_ai_agent_cli.backups.probe_file_not_held", return_value=held) as probe:
                ok, result, errors, _ = create_backup(
                    record["document_id"], "b-open", workspace=workspace, source_open_in_wps=True
                )
            self.assertTrue(ok, errors)
            self.assertTrue(result["created"])
            self.assertTrue(Path(result["backup_path"]).is_file())
            probe.assert_not_called()

    def test_a_backup_for_a_command_working_through_wps_still_refuses_an_unreadable_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "book.xlsx"
            Workbook().save(document)
            _, record, _ = register_document("spreadsheets", str(document), workspace)
            with patch("wps_ai_agent_cli.backups.file_identity", side_effect=PermissionError(13, "Permission denied")):
                ok, _result, errors, _ = create_backup(
                    record["document_id"], "b-open-denied", workspace=workspace, source_open_in_wps=True
                )
            self.assertFalse(ok)
            _assert_held_open_envelope(self, errors)

    def test_create_backup_maps_identity_permission_error_to_held_open(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "book.xlsx"
            Workbook().save(document)
            _, record, _ = register_document("spreadsheets", str(document), workspace)
            with (
                patch("wps_ai_agent_cli.backups.probe_file_not_held", return_value=None),
                patch(
                    "wps_ai_agent_cli.backups.file_identity",
                    side_effect=PermissionError("Permission denied"),
                ),
            ):
                ok, _result, errors, _ = create_backup(record["document_id"], "b-perm", workspace=workspace)
            self.assertFalse(ok)
            _assert_held_open_envelope(self, errors)
            self.assertNotEqual(errors[0].get("code"), "DOCUMENT_IDENTITY_UNAVAILABLE")

    def test_spreadsheet_write_surfaces_held_open_from_backup_preflight(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "book.xlsx"
            Workbook().save(document)
            _, record, _ = register_document("spreadsheets", str(document), workspace)
            held = held_open_error("document", "exclusive test hold")
            with (
                patch("wps_ai_agent_cli.backups.probe_file_not_held", return_value=held),
                patch("wps_ai_agent_cli.spreadsheet_ops._run_spreadsheet_write_com") as com,
            ):
                ok, _result, errors, _ = write_spreadsheet_range(
                    record["document_id"],
                    "A1",
                    [["x"]],
                    "w-held",
                    workspace=workspace,
                )
            self.assertFalse(ok)
            _assert_held_open_envelope(self, errors)
            com.assert_not_called()

    def test_writer_replace_refuses_exclusive_hold_before_zip_preview(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "document.docx"
            write_document(path, "<w:p><w:r><w:t>hello</w:t></w:r></w:p>")
            _, record, _ = register_document("writer", str(path), tmp)
            held = held_open_error("document", "exclusive test hold")
            with (
                patch("wps_ai_agent_cli.writer_ops.probe_file_not_held", return_value=held) as probe,
                patch("wps_ai_agent_cli.writer_ops.docx_body_paragraphs") as paragraphs,
                patch("wps_ai_agent_cli.writer_ops.create_backup") as backup,
            ):
                ok, _result, errors, _ = writer_replace(
                    record["document_id"],
                    "hello",
                    "hi",
                    "w-held",
                    workspace=tmp,
                )
            self.assertFalse(ok)
            _assert_held_open_envelope(self, errors)
            self.assertEqual(probe.call_args.kwargs["read_only"], True)
            paragraphs.assert_not_called()
            backup.assert_not_called()

    def test_presentation_replace_refuses_exclusive_hold_before_zip_preview(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "deck.pptx"
            # Minimal zip so register/path.exists works; Zip body is never read when probe refuses.
            path.write_bytes(b"PK\x05\x06" + b"\x00" * 18)
            _, record, _ = register_document("presentation", str(path), tmp)
            held = held_open_error("presentation", "exclusive test hold")
            with (
                patch("wps_ai_agent_cli.presentation_ops.probe_file_not_held", return_value=held) as probe,
                patch("wps_ai_agent_cli.presentation_ops.pptx_slide_texts") as slides,
                patch("wps_ai_agent_cli.presentation_ops.create_backup") as backup,
            ):
                ok, _result, errors, _ = presentation_replace(
                    record["document_id"],
                    "a",
                    "b",
                    "p-held",
                    workspace=tmp,
                )
            self.assertFalse(ok)
            _assert_held_open_envelope(self, errors)
            self.assertEqual(probe.call_args.kwargs["read_only"], True)
            self.assertEqual(probe.call_args.kwargs["what"], "presentation")
            slides.assert_not_called()
            backup.assert_not_called()


@unittest.skipUnless(os.name == "nt", "Win32 FileShare semantics required")
class Win32ShareHoldIntegrationTests(unittest.TestCase):
    def test_exclusive_hold_blocks_read_and_exclusive_probes(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "held.bin"
            path.write_bytes(b"data")
            close = _hold_exclusive_win32(path)
            try:
                refused_read = probe_file_not_held(path, read_only=True, what="document")
                refused_ex = probe_file_not_held(path, read_only=False, what="document")
            finally:
                close()
            self.assertIsNotNone(refused_read)
            self.assertIsNotNone(refused_ex)
            self.assertIn("held open", refused_read["message"])
            self.assertIn("Open() was not called", refused_read["message"])

    def test_read_share_hold_allows_read_probe_but_blocks_exclusive(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "shared.bin"
            path.write_bytes(b"data")
            close = _hold_read_share_win32(path)
            try:
                read_ok = probe_file_not_held(path, read_only=True, what="document")
                exclusive = probe_file_not_held(path, read_only=False, what="document")
            finally:
                close()
            self.assertIsNone(read_ok)
            self.assertIsNotNone(exclusive)
            self.assertIn("held open", exclusive["message"])

    def test_inspect_under_exclusive_hold_returns_json_envelope(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "book.xlsx"
            Workbook().save(document)
            _, record, _ = register_document("spreadsheets", str(document), workspace)
            close = _hold_exclusive_win32(document)
            try:
                with patch("wps_ai_agent_cli.spreadsheet_inspect.probe_wps_capabilities") as caps:
                    ok, _result, errors = inspect_spreadsheet_range(
                        record["document_id"],
                        "A1",
                        workspace=workspace,
                    )
            finally:
                close()
            self.assertFalse(ok)
            _assert_held_open_envelope(self, errors)
            caps.assert_not_called()


if __name__ == "__main__":
    unittest.main()
