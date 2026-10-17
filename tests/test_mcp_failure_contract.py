"""Every MCP tool fails in a way the agent can see (P5-002, Phase 5 exit criterion F1).

Agents read the text of a tool result. Phase 4 found three bugs, only through real runs, where a failure
reached the agent without its error code (#55, #56, #60). These tests make each of the 85 tools fail and
check what the agent gets back:

- a JSON-RPC argument error (-32602), or
- a tool result marked isError whose text lists errors, each with a code and a message, and never
  MCP_TOOL_EXECUTION_FAILED (an uncaught exception) or MCP_RESPONSE_PARSE_FAILED.

PowerShell is pointed at a missing executable, so no test here starts or attaches to WPS, on any machine.
"""

import argparse
import builtins
import json
import os
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from pptx import Presentation

from tests.isolation import use_temporary_cwd
from tests.test_writer_ops import write_minimal_docx
from wps_ai_agent_cli.mcp_schema import list_mcp_tool_schemas
from wps_ai_agent_cli.mcp_server import handle_mcp_request
from wps_ai_agent_cli.sessions import register_document

_restore_cwd = None


def setUpModule():
    global _restore_cwd
    _restore_cwd = use_temporary_cwd()


def tearDownModule():
    _restore_cwd()


UNCAUGHT = {"MCP_TOOL_EXECUTION_FAILED", "MCP_RESPONSE_PARSE_FAILED"}
SHARING_VIOLATION = "The process cannot access the file because it is being used by another process"
PATH_ARGUMENTS = {"input", "path", "docx", "input_dir", "manifest", "template", "request", "mapping"}
# Values that get past argument checks, so that a document tool reaches the document itself.
PLAUSIBLE = {
    "range": "A1",
    "values_json": "[[1]]",
    "formulas_json": '[["=1"]]',
    "expected_values_json": "[[1]]",
    "find": "alpha",
    "replace": "beta",
    "text": "x",
    "sheet_name": "Sheet",
    "name": "Second",
    "new_name": "Second",
    "old_name": "Sheet",
    "source_name": "Sheet",
    "color": "FF0000",
    "table_index": 1,
    "row": 1,
    "column": 1,
    "bookmark_name": "b",
    "visible": False,
    "backup_name": "missing.xlsx",
}


def _no_powershell():
    return patch(
        "wps_ai_agent_cli.powershell_runner.powershell_executable",
        return_value=str(Path("no-such-dir") / "powershell"),
    )


def _call(name, arguments):
    with _no_powershell():
        return handle_mcp_request(
            {"jsonrpc": "2.0", "id": name, "method": "tools/call", "params": {"name": name, "arguments": arguments}}
        )


def _placeholder(name, spec, missing):
    if "enum" in spec:
        return spec["enum"][0]
    kind = spec.get("type")
    if isinstance(kind, list):
        kind = next(item for item in kind if item != "null")
    if name == "document_id":
        return "doc_missing0000000"
    if kind == "string":
        if name in PATH_ARGUMENTS:
            return str(missing / ("in.docx" if name in ("input", "docx") else "in"))
        if name == "output":
            return str(missing / "out" / "out.docx")
        if name == "output_dir":
            return str(missing / "out")
        return "x"
    if kind in ("integer", "number"):
        return spec.get("minimum", 1)
    if kind == "boolean":
        return False
    if kind == "array":
        return []
    if kind == "object":
        return {}
    return "x"


class FailureContractTests(unittest.TestCase):
    def assert_agent_sees_the_failure(self, response):
        """Return "rpc_error", "tool_error" or "succeeded"; fail if the agent gets nothing it can act on."""
        if "error" in response:
            self.assertEqual(response["error"]["code"], -32602, response["error"])
            return "rpc_error"
        result = response["result"]
        text = json.loads(result["content"][0]["text"])
        if not result["isError"]:
            return "succeeded"
        self.assertFalse(text["ok"])
        self.assertTrue(text["errors"], "a failed call must list its errors in the text the agent reads")
        for error in text["errors"]:
            self.assertIsInstance(error.get("code"), str)
            self.assertTrue(error["code"])
            self.assertNotIn(error["code"], UNCAUGHT, error)
            self.assertTrue(error.get("message"), error)
            self.assertNotIn("stderr", error["message"], "the agent cannot read the server's stderr")
        self.assertEqual(result["structuredContent"]["errors"], text["errors"])
        return "tool_error"

    def test_every_tool_fails_visibly_on_missing_inputs(self):
        missing = Path("contract-missing").resolve()
        outcomes = {}
        for tool in list_mcp_tool_schemas():
            schema = tool["input_schema"]
            arguments = {
                name: _placeholder(name, schema["properties"][name], missing) for name in schema.get("required", [])
            }
            with self.subTest(tool=tool["name"]):
                outcome = self.assert_agent_sees_the_failure(_call(tool["name"], arguments))
                if outcome == "succeeded":
                    # Nothing this tool takes can be missing, so check the other way it can fail.
                    response = _call(tool["name"], {**arguments, "unexpected_argument": 1})
                    self.assertEqual(self.assert_agent_sees_the_failure(response), "rpc_error")
                    self.assertEqual(response["error"]["data"][0]["code"], "MCP_ARGUMENTS_UNKNOWN")
                outcomes[tool["name"]] = outcome
        self.assertEqual(len(outcomes), len(list_mcp_tool_schemas()))

    def test_every_document_tool_fails_visibly_on_a_held_document(self):
        documents = self._register_one_document_per_component()
        checked = 0
        for tool in list_mcp_tool_schemas():
            schema = tool["input_schema"]
            if "document_id" not in schema.get("required", []):
                continue
            name = tool["name"]
            component = "writer" if "writer" in name else "presentation" if "presentation" in name else "spreadsheets"
            document_id, path = documents[component]
            arguments = {}
            for argument in schema["required"]:
                spec = schema["properties"][argument]
                if argument == "document_id":
                    arguments[argument] = document_id
                elif argument == "output":
                    arguments[argument] = str(
                        Path("contract-out").resolve() / f"{name}{self._output_suffix(name, path)}"
                    )
                elif "enum" in spec:
                    arguments[argument] = spec["enum"][0]
                else:
                    arguments[argument] = PLAUSIBLE.get(argument, "x")
            with self.subTest(tool=name), _held(path):
                self.assertEqual(self.assert_agent_sees_the_failure(_call(name, arguments)), "tool_error")
                checked += 1
        self.assertGreaterEqual(checked, 30)  # 36 document tools today; guards against the filter going wrong

    def test_writer_structure_reports_a_held_document_as_held(self):
        documents = self._register_one_document_per_component()
        document_id, path = documents["writer"]
        with _held(path):
            response = _call("wps_agent_writer_structure", {"document_id": document_id})

        text = json.loads(response["result"]["content"][0]["text"])
        # It reported WRITER_STRUCTURE_INVALID before, which tells the agent the document is broken.
        self.assertEqual([error["code"] for error in text["errors"]], ["DOCUMENT_HELD_OPEN"])

    def test_rejected_cli_arguments_carry_the_reason(self):
        # No tool's schema lets through arguments the CLI rejects today (restore-backup's exclusive pair was the
        # last, P12-003), so the CLI here gains a required option the schema does not have.
        from wps_ai_agent_cli import cli_parser

        original = cli_parser.build_parser

        def stricter_parser():
            parser = original()
            for action in parser._actions:
                if isinstance(action, argparse._SubParsersAction):
                    action.choices["restore-backup"].add_argument("--confirm-restore", required=True)
            return parser

        with (
            patch("wps_ai_agent_cli.cli_parser.build_parser", stricter_parser),
            patch("wps_ai_agent_cli.cli.build_parser", stricter_parser),
        ):
            response = _call("wps_agent_restore_backup", {"document_id": "doc_missing0000000", "backup_name": "x"})

        text = json.loads(response["result"]["content"][0]["text"])
        self.assertEqual(text["errors"][0]["code"], "MCP_CLI_ARGUMENTS_REJECTED")
        self.assertIn("--confirm-restore", text["errors"][0]["message"])

    def test_process_audit_without_powershell_is_a_tool_error(self):
        text = json.loads(_call("wps_agent_wps_process_audit", {})["result"]["content"][0]["text"])

        self.assertEqual([error["code"] for error in text["errors"]], ["WPS_PROCESS_AUDIT_FAILED"])
        self.assertIn("PowerShell could not be started", text["errors"][0]["message"])

    def _register_one_document_per_component(self):
        folder = Path("contract-docs").resolve()
        folder.mkdir(exist_ok=True)
        book = folder / "book.xlsx"
        workbook = Workbook()
        workbook.active["A1"] = "seed"
        workbook.save(book)
        report = folder / "report.docx"
        write_minimal_docx(report, "alpha")
        deck = folder / "deck.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "alpha"
        presentation.save(deck)
        documents = {}
        for component, path in (("spreadsheets", book), ("writer", report), ("presentation", deck)):
            ok, record, errors = register_document(component, str(path))
            self.assertTrue(ok, errors)
            documents[component] = (record["document_id"], path)
        return documents

    @staticmethod
    def _output_suffix(name, path):
        if "pdf" in name:
            return ".pdf"
        if name == "wps_agent_export_open_document":
            return ".html"
        if "text" in name:
            return ".txt"
        return path.suffix


class _held:
    """Make every open of path fail the way a file held with no sharing does on Windows."""

    def __init__(self, path):
        self.target = Path(path).resolve()
        self.patches = []

    def __enter__(self):
        real_open, real_zip = builtins.open, zipfile.ZipFile

        def refuse(opener):
            def open_or_refuse(file, *args, **kwargs):
                if isinstance(file, (str, os.PathLike)) and Path(file).resolve() == self.target:
                    raise PermissionError(13, SHARING_VIOLATION, str(file))
                return opener(file, *args, **kwargs)

            return open_or_refuse

        self.patches = [
            patch("builtins.open", side_effect=refuse(real_open)),
            patch("io.open", side_effect=refuse(real_open)),
            patch("zipfile.ZipFile", side_effect=refuse(real_zip)),
        ]
        for each in self.patches:
            each.start()
        return self

    def __exit__(self, *exc_info):
        for each in reversed(self.patches):
            each.stop()
        return False


if __name__ == "__main__":
    unittest.main()
