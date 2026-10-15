from __future__ import annotations

import importlib.util
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from wps_ai_agent_cli.backups import create_backup, list_backups, restore_backup
from wps_ai_agent_cli.operations import record_operation
from wps_ai_agent_cli.sessions import register_document, stable_document_id

REPO_ROOT = Path(__file__).resolve().parents[1]
TASKS_PATH = REPO_ROOT / "config" / "agent_eval_tasks.json"

_spec = importlib.util.spec_from_file_location(
    "run_agent_eval", REPO_ROOT / "scripts" / "agent_eval" / "run_agent_eval.py"
)
run_agent_eval = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(run_agent_eval)

KNOWN_CHECK_KINDS = {
    "document_registered",
    "operation_recorded",
    "no_operation",
    "backup_count",
    "file_unchanged",
    "file_matches_seed",
    "tool_called",
    "tool_not_called",
    "tool_error_seen",
    "max_tool_calls",
    "xlsx_cell",
    "file_text",
    "xlsx_sheets",
    "docx_table_cell",
    "pdf_file",
}
WPS_WRITE_TASKS = (
    "wps-spreadsheet-write",
    "wps-writer-replace",
    "wps-presentation-replace",
    "wps-held-spreadsheet-write",
    "wps-spreadsheet-rename-sheet",
    "wps-spreadsheet-create-sheet",
    "wps-writer-table-write",
    "wps-docx-to-pdf",
    "wps-open-document-export-text",
)
KNOWN_SETUP_STEPS = {"register", "backup", "modify", "html_table", "hold", "open_in_wps"}
# Setup steps that leave the fixture held or open for the agent; they must come last.
FINAL_SETUP_STEPS = {"hold", "open_in_wps"}


def _task(task_id: str) -> dict:
    spec = json.loads(TASKS_PATH.read_text(encoding="utf-8"))
    return next(task for task in spec["tasks"] if task["id"] == task_id)


def _fake_agent(actions, tool_calls=None):
    """A stand-in for the claude CLI: it performs the given actions on the workspace, then reports ran."""

    def runner(prompt, mcp_config, workspace, model, max_turns, timeout):
        actions(workspace)
        return {"ran": True, "exit_code": 0, "result_excerpt": "", "tool_calls": tool_calls or []}

    return runner


def _call(tool, error_codes=(), **arguments):
    return {
        "name": f"mcp__wps-agent__wps_agent_{tool}",
        "input": arguments,
        "is_error": bool(error_codes),
        "error_codes": list(error_codes),
    }


def _document_id(workspace: Path) -> str:
    return stable_document_id("spreadsheets", str(workspace / "book.xlsx"))


def _latest_backup(workspace: Path) -> str:
    ok, data, errors = list_backups(_document_id(workspace), workspace)
    assert ok, errors
    return data["backups"][0]["backup_name"]


def _register(workspace: Path):
    ok, record, _ = register_document("spreadsheets", str(workspace / "book.xlsx"), workspace)
    assert ok, record
    return record["document_id"]


class AgentEvalTaskFileTests(unittest.TestCase):
    def test_every_check_kind_in_the_task_file_is_known(self):
        spec = json.loads(TASKS_PATH.read_text(encoding="utf-8"))
        kinds = {check["kind"] for task in spec["tasks"] for check in task["checks"]}
        self.assertLessEqual(kinds, KNOWN_CHECK_KINDS)
        self.assertEqual(len({task["id"] for task in spec["tasks"]}), len(spec["tasks"]))

    def test_every_task_names_real_expected_tools_that_cover_its_tool_checks(self):
        from wps_ai_agent_cli.mcp_schema import list_mcp_tool_schemas

        tools = {schema["name"].removeprefix("wps_agent_") for schema in list_mcp_tool_schemas()}
        spec = json.loads(TASKS_PATH.read_text(encoding="utf-8"))
        for task in spec["tasks"]:
            with self.subTest(task=task["id"]):
                expected = task.get("expected_tools")
                self.assertTrue(expected, "every task lists the tools a correct run needs")
                self.assertLessEqual(set(expected), tools)
                for check in task["checks"]:
                    if check["kind"] == "tool_called":
                        self.assertIn(check["tool"], expected)
                    if check["kind"] == "tool_not_called":
                        self.assertNotIn(check["tool"], expected)

    def test_wps_write_tasks_require_wps_and_name_a_supported_fixture(self):
        spec = json.loads(TASKS_PATH.read_text(encoding="utf-8"))
        tasks = {task["id"]: task for task in spec["tasks"]}
        for task_id in WPS_WRITE_TASKS:
            task = tasks[task_id]
            self.assertTrue(task.get("requires_wps"), task_id)
            fixture = task.get("fixture", spec["fixture"])
            self.assertIn(Path(fixture).suffix, {".xlsx", ".docx", ".pptx"}, task_id)
            self.assertIn(task.get("component", spec["component"]), {"spreadsheets", "writer", "presentation"}, task_id)

    def test_docx_and_pptx_fixtures_hold_the_text_the_tasks_replace(self):
        with tempfile.TemporaryDirectory() as tmp:
            for name in ("report.docx", "deck.pptx"):
                run_agent_eval.write_fixture(Path(tmp), name)
                text = run_agent_eval.document_text(Path(tmp) / name)
                self.assertIn("Status: DRAFT", text, name)

    def test_setup_steps_are_known_and_hold_comes_last(self):
        spec = json.loads(TASKS_PATH.read_text(encoding="utf-8"))
        for task in spec["tasks"]:
            steps = task.get("setup", [])
            self.assertLessEqual(set(steps), KNOWN_SETUP_STEPS, task["id"])
            final = [step for step in steps if step in FINAL_SETUP_STEPS]
            self.assertLessEqual(len(final), 1, task["id"])
            if final:
                self.assertEqual(steps[-1], final[0], task["id"])
            if "open_in_wps" in steps:
                self.assertTrue(task.get("requires_wps"), task["id"])

    def test_the_table_fixture_holds_the_table_the_task_edits(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            run_agent_eval.write_fixture(Path(tmp), "table_report.docx")
            table = Document(Path(tmp) / "table_report.docx").tables[0]
            rows = [[cell.text for cell in row.cells] for row in table.rows]
        self.assertEqual(rows, [list(row) for row in run_agent_eval.WRITER_TABLE_ROWS])
        self.assertEqual(rows[1], ["North", "120", "DRAFT"])


class TranscriptTests(unittest.TestCase):
    def test_tool_calls_and_error_codes_are_read_from_stream_json(self):
        envelope = json.dumps({"ok": False, "errors": [{"code": "DOCUMENT_HELD_OPEN", "message": "held"}]})
        lines = [
            json.dumps({"type": "system", "subtype": "init"}),
            json.dumps(
                {
                    "type": "assistant",
                    "message": {
                        "content": [
                            {"type": "tool_use", "id": "t1", "name": "ToolSearch", "input": {"query": "select:x"}},
                        ]
                    },
                }
            ),
            json.dumps(
                {
                    "type": "assistant",
                    "message": {
                        "content": [
                            {
                                "type": "tool_use",
                                "id": "t2",
                                "name": "mcp__wps-agent__wps_agent_backup_document",
                                "input": {"request_id": "r1"},
                            },
                        ]
                    },
                }
            ),
            json.dumps(
                {
                    "type": "user",
                    "message": {
                        "content": [
                            {
                                "type": "tool_result",
                                "tool_use_id": "t2",
                                "content": [{"type": "text", "text": envelope}],
                            },
                        ]
                    },
                }
            ),
            "not json",
            json.dumps({"type": "result", "num_turns": 3, "total_cost_usd": 0.05, "result": "done"}),
        ]
        result, calls = run_agent_eval.parse_transcript(lines)

        self.assertEqual(result["num_turns"], 3)
        self.assertEqual([call["name"] for call in calls], ["ToolSearch", "mcp__wps-agent__wps_agent_backup_document"])
        self.assertEqual(calls[1]["error_codes"], ["DOCUMENT_HELD_OPEN"])
        self.assertTrue(calls[1]["is_error"])
        metrics = run_agent_eval.tool_metrics(calls)
        self.assertEqual(
            (metrics["wps_tool_calls"], metrics["wps_tool_errors"], metrics["other_tool_calls"]), (1, 1, 1)
        )
        self.assertEqual(metrics["tools_used"], ["backup_document"])


class AgentEvalJudgeTests(unittest.TestCase):
    def test_passes_when_the_expected_state_is_present(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)

            def agent(workspace):
                document_id = _register(workspace)
                ok, _result, errors, _ = create_backup(document_id, "eval-backup-1", workspace=workspace)
                assert ok, errors

            result = run_agent_eval.run_task(
                _task("register-and-backup"), "book.xlsx", "spreadsheets", root, None, 15, 60, runner=_fake_agent(agent)
            )

        self.assertTrue(result["passed"], result)
        self.assertTrue(all(check["passed"] for check in result["checks"]))

    def test_fails_when_the_agent_does_nothing(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = run_agent_eval.run_task(
                _task("register-and-backup"),
                "book.xlsx",
                "spreadsheets",
                Path(tmp),
                None,
                15,
                60,
                runner=_fake_agent(lambda workspace: None),
            )

        self.assertFalse(result["passed"])
        failed = {check["kind"] for check in result["checks"] if not check["passed"]}
        self.assertIn("document_registered", failed)
        self.assertIn("operation_recorded", failed)

    def test_a_refused_restore_leaves_no_operation_and_no_change(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)

            def agent(workspace):
                document_id = _register(workspace)
                ok, _result, _errors, _ = restore_backup(
                    document_id, "eval-restore-1", backup_name="missing.xlsx", workspace=workspace
                )
                assert not ok

            result = run_agent_eval.run_task(
                _task("restore-missing-backup"),
                "book.xlsx",
                "spreadsheets",
                root,
                None,
                15,
                60,
                runner=_fake_agent(agent),
            )

        self.assertTrue(result["passed"], result)

    def test_a_timed_out_agent_is_a_failure_even_if_the_state_looks_right(self):
        def timeout_runner(prompt, mcp_config, workspace, model, max_turns, timeout):
            _register(workspace)
            return {"ran": False, "error": "timed out after 60s"}

        with tempfile.TemporaryDirectory() as tmp:
            result = run_agent_eval.run_task(
                _task("list-backups-readonly"),
                "book.xlsx",
                "spreadsheets",
                Path(tmp),
                None,
                15,
                60,
                runner=timeout_runner,
            )

        self.assertFalse(result["passed"])

    def test_a_retried_backup_with_the_same_request_id_makes_one_backup(self):
        with tempfile.TemporaryDirectory() as tmp:

            def agent(workspace):
                document_id = _register(workspace)
                for _ in range(2):
                    ok, _result, errors, _ = create_backup(document_id, "eval-retry-1", workspace=workspace)
                    assert ok, errors

            calls = [_call("backup_document", request_id="eval-retry-1", document_id="d")] * 2
            result = run_agent_eval.run_task(
                _task("idempotent-backup-retry"),
                "book.xlsx",
                "spreadsheets",
                Path(tmp),
                None,
                15,
                60,
                runner=_fake_agent(agent, calls),
            )

        self.assertTrue(result["passed"], result)

    def test_a_single_backup_call_fails_the_retry_task(self):
        with tempfile.TemporaryDirectory() as tmp:

            def agent(workspace):
                create_backup(_register(workspace), "eval-retry-1", workspace=workspace)

            calls = [_call("backup_document", request_id="eval-retry-1")]
            result = run_agent_eval.run_task(
                _task("idempotent-backup-retry"),
                "book.xlsx",
                "spreadsheets",
                Path(tmp),
                None,
                15,
                60,
                runner=_fake_agent(agent, calls),
            )

        self.assertFalse(result["passed"])
        self.assertEqual([check["kind"] for check in result["checks"] if not check["passed"]], ["tool_called"])

    def test_restoring_the_latest_backup_brings_back_the_seed(self):
        with tempfile.TemporaryDirectory() as tmp:

            def agent(workspace):
                ok, _result, errors, _ = restore_backup(
                    _document_id(workspace),
                    "eval-restore-2",
                    backup_name=_latest_backup(workspace),
                    workspace=workspace,
                )
                assert ok, errors

            calls = [_call("list_backups"), _call("restore_backup", request_id="eval-restore-2")]
            result = run_agent_eval.run_task(
                _task("restore-latest-backup"),
                "book.xlsx",
                "spreadsheets",
                Path(tmp),
                None,
                15,
                60,
                runner=_fake_agent(agent, calls),
            )

        self.assertTrue(result["passed"], result)

    def test_the_restore_task_fails_when_the_mistaken_change_is_left_in_place(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = run_agent_eval.run_task(
                _task("restore-latest-backup"),
                "book.xlsx",
                "spreadsheets",
                Path(tmp),
                None,
                15,
                60,
                runner=_fake_agent(lambda workspace: None),
            )

        failed = {check["kind"] for check in result["checks"] if not check["passed"]}
        self.assertEqual(failed, {"operation_recorded", "file_matches_seed"})

    def test_too_many_tool_calls_fail_the_restore_task(self):
        with tempfile.TemporaryDirectory() as tmp:

            def agent(workspace):
                restore_backup(
                    _document_id(workspace),
                    "eval-restore-2",
                    backup_name=_latest_backup(workspace),
                    workspace=workspace,
                )

            calls = [_call("documents")] * 6 + [_call("restore_backup", request_id="eval-restore-2")]
            result = run_agent_eval.run_task(
                _task("restore-latest-backup"),
                "book.xlsx",
                "spreadsheets",
                Path(tmp),
                None,
                15,
                60,
                runner=_fake_agent(agent, calls),
            )

        self.assertEqual([check["kind"] for check in result["checks"] if not check["passed"]], ["max_tool_calls"])

    def test_dry_run_passes_and_a_real_restore_fails_the_preview_task(self):
        with tempfile.TemporaryDirectory() as tmp:

            def preview(workspace):
                ok, _result, errors, _ = restore_backup(
                    _document_id(workspace),
                    "p1",
                    backup_name=_latest_backup(workspace),
                    dry_run=True,
                    workspace=workspace,
                )
                assert ok, errors

            dry = run_agent_eval.run_task(
                _task("restore-dry-run"),
                "book.xlsx",
                "spreadsheets",
                Path(tmp) / "a",
                None,
                15,
                60,
                runner=_fake_agent(preview, [_call("restore_backup", dry_run=True)]),
            )

            def restore(workspace):
                restore_backup(
                    _document_id(workspace), "p2", backup_name=_latest_backup(workspace), workspace=workspace
                )

            real = run_agent_eval.run_task(
                _task("restore-dry-run"),
                "book.xlsx",
                "spreadsheets",
                Path(tmp) / "b",
                None,
                15,
                60,
                runner=_fake_agent(restore, [_call("restore_backup")]),
            )

        self.assertTrue(dry["passed"], dry)
        # A real restore also takes a safety backup of the current file, so the backup count changes too.
        self.assertEqual(
            {check["kind"] for check in real["checks"] if not check["passed"]},
            {"tool_called", "file_unchanged", "backup_count"},
        )

    def test_held_document_task_is_skipped_when_the_file_cannot_be_held(self):
        with (
            tempfile.TemporaryDirectory() as tmp,
            patch.object(run_agent_eval, "hold_file", return_value=(None, "cannot hold here")),
        ):
            result = run_agent_eval.run_task(
                _task("held-document-backup"),
                "book.xlsx",
                "spreadsheets",
                Path(tmp),
                None,
                15,
                60,
                runner=_fake_agent(lambda workspace: self.fail("agent must not run")),
            )

        self.assertEqual(result["skipped"], "cannot hold here")
        self.assertFalse(result["passed"])

    def test_held_document_task_passes_on_a_reported_refusal_and_releases_the_hold(self):
        released = []
        with (
            tempfile.TemporaryDirectory() as tmp,
            patch.object(run_agent_eval, "hold_file", return_value=(lambda: released.append(True), None)),
        ):
            calls = [_call("backup_document", error_codes=["DOCUMENT_HELD_OPEN"], request_id="eval-held-1")]
            result = run_agent_eval.run_task(
                _task("held-document-backup"),
                "book.xlsx",
                "spreadsheets",
                Path(tmp),
                None,
                15,
                60,
                runner=_fake_agent(lambda workspace: None, calls),
            )

        self.assertTrue(result["passed"], result)
        self.assertEqual(released, [True])

    def test_held_document_task_fails_when_the_refusal_was_never_seen(self):
        with (
            tempfile.TemporaryDirectory() as tmp,
            patch.object(run_agent_eval, "hold_file", return_value=(lambda: None, None)),
        ):
            result = run_agent_eval.run_task(
                _task("held-document-backup"),
                "book.xlsx",
                "spreadsheets",
                Path(tmp),
                None,
                15,
                60,
                runner=_fake_agent(lambda workspace: None, [_call("restore_backup")]),
            )

        self.assertEqual(
            {check["kind"] for check in result["checks"] if not check["passed"]}, {"tool_error_seen", "tool_not_called"}
        )

    @unittest.skipIf(os.name == "nt" or os.geteuid() == 0, "chmod hold needs a non-root POSIX user")
    def test_chmod_hold_makes_the_file_unreadable_until_released(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.xlsx"
            path.write_bytes(b"x")
            release, reason = run_agent_eval.hold_file(path)
            self.assertIsNone(reason)
            with self.assertRaises(PermissionError):
                path.read_bytes()
            release()
            self.assertEqual(path.read_bytes(), b"x")

    def test_html_table_conversion_is_judged_from_the_output_workbook(self):
        from wps_ai_agent_cli.html_office import convert_html_to_office

        with tempfile.TemporaryDirectory() as tmp:

            def agent(workspace):
                convert_html_to_office(workspace / "table.html", workspace / "table.xlsx", "xlsx")

            result = run_agent_eval.run_task(
                _task("html-table-to-xlsx"),
                "book.xlsx",
                "spreadsheets",
                Path(tmp),
                None,
                15,
                60,
                runner=_fake_agent(agent, [_call("html_editable", format="xlsx")]),
            )
            missing = run_agent_eval.run_task(
                _task("html-table-to-xlsx"),
                "book.xlsx",
                "spreadsheets",
                Path(tmp) / "none",
                None,
                15,
                60,
                runner=_fake_agent(lambda workspace: None),
            )

        self.assertTrue(result["passed"], result)
        self.assertEqual(
            [check["detail"] for check in missing["checks"] if check["kind"] == "xlsx_cell"], ["output missing"] * 2
        )

    def test_off_path_calls_are_counted_against_the_expected_tools(self):
        calls = [
            _call("register_document"),
            _call("documents"),
            _call("backup_document"),
            _call("documents"),
            {"name": "ToolSearch", "input": {}, "is_error": False, "error_codes": []},
        ]

        metrics = run_agent_eval.tool_metrics(calls, ["register_document", "backup_document"])

        self.assertEqual(metrics["wps_tool_calls"], 4)
        self.assertEqual(metrics["other_tool_calls"], 1)
        self.assertEqual(metrics["off_path_calls"], 2)
        self.assertEqual(metrics["off_path_tools"], ["documents"])
        self.assertEqual(metrics["expected_tools"], ["backup_document", "register_document"])

    def test_without_expected_tools_off_path_calls_are_not_counted(self):
        metrics = run_agent_eval.tool_metrics([_call("documents")])

        self.assertIsNone(metrics["off_path_calls"])
        self.assertIsNone(metrics["off_path_tools"])

    def test_off_path_calls_are_reported_but_do_not_fail_a_task(self):
        def agent(workspace):
            document_id = _register(workspace)
            create_backup(document_id, "eval-backup-1", workspace=workspace)

        calls = [_call("register_document"), _call("operations"), _call("backup_document", request_id="eval-backup-1")]
        with tempfile.TemporaryDirectory() as tmp:
            result = run_agent_eval.run_task_repeated(
                _task("register-and-backup"),
                "book.xlsx",
                "spreadsheets",
                Path(tmp),
                None,
                15,
                60,
                repeat=2,
                min_passes=2,
                runner=_fake_agent(agent, calls),
            )
        summary = run_agent_eval.summarize([result], "stamp", None)

        self.assertTrue(result["passed"])
        self.assertEqual([run["metrics"]["off_path_calls"] for run in result["runs"]], [1, 1])
        self.assertEqual(result["metrics"]["off_path_calls"], 2)
        self.assertEqual(result["metrics"]["off_path_tools"], ["operations"])
        self.assertEqual(summary["off_path_calls"], 2)
        self.assertEqual(summary["tasks_without_expected_tools"], [])

    def test_summary_counts_skipped_tasks_separately(self):
        results = [
            {"id": "a", "passed": True, "metrics": {"wps_tool_calls": 2, "wps_tool_errors": 0}},
            {"id": "b", "passed": False, "metrics": {"wps_tool_calls": 3, "wps_tool_errors": 1}},
            {"id": "c", "passed": False, "skipped": "cannot hold here"},
        ]
        summary = run_agent_eval.summarize(results, "stamp", None)

        self.assertEqual((summary["passed"], summary["total"], summary["skipped"]), (1, 2, 1))
        self.assertEqual((summary["wps_tool_calls"], summary["wps_tool_errors"]), (5, 1))

    def test_a_wps_task_is_skipped_without_wps_and_the_agent_never_runs(self):
        task = {**_task("register-and-backup"), "requires_wps": True}
        with tempfile.TemporaryDirectory() as tmp:
            result = run_agent_eval.run_task(
                task,
                "book.xlsx",
                "spreadsheets",
                Path(tmp),
                None,
                15,
                60,
                runner=_fake_agent(lambda workspace: self.fail("agent must not run")),
                has_wps=lambda component: False,
            )

        self.assertIn("needs WPS spreadsheets", result["skipped"])
        self.assertFalse(result["passed"])

    def test_a_wps_task_runs_when_wps_is_available(self):
        task = {**_task("register-and-backup"), "requires_wps": True}
        seen = []
        with tempfile.TemporaryDirectory() as tmp:

            def agent(workspace):
                document_id = _register(workspace)
                create_backup(document_id, "eval-backup-1", workspace=workspace)

            result = run_agent_eval.run_task(
                task,
                "book.xlsx",
                "spreadsheets",
                Path(tmp),
                None,
                15,
                60,
                runner=_fake_agent(agent),
                has_wps=lambda component: seen.append(component) or True,
            )

        self.assertEqual(seen, ["spreadsheets"])
        self.assertTrue(result["passed"], result)

    def test_a_task_without_requires_wps_never_probes(self):
        with tempfile.TemporaryDirectory() as tmp:
            run_agent_eval.run_task(
                _task("list-backups-readonly"),
                "book.xlsx",
                "spreadsheets",
                Path(tmp),
                None,
                15,
                60,
                runner=_fake_agent(lambda workspace: None),
                has_wps=lambda component: self.fail("must not probe"),
            )

    def test_repeated_runs_use_separate_workspaces_and_count_passes(self):
        workspaces = []
        outcomes = iter([True, False, True])

        def agent(workspace):
            workspaces.append(workspace)
            if next(outcomes):
                document_id = _register(workspace)
                create_backup(document_id, "eval-backup-1", workspace=workspace)

        with tempfile.TemporaryDirectory() as tmp:
            result = run_agent_eval.run_task_repeated(
                _task("register-and-backup"),
                "book.xlsx",
                "spreadsheets",
                Path(tmp),
                None,
                15,
                60,
                repeat=3,
                min_passes=2,
                runner=_fake_agent(agent),
            )

        self.assertEqual(len(set(workspaces)), 3)
        self.assertEqual((result["passes"], result["repeat"], result["min_passes"]), (2, 3, 2))
        self.assertAlmostEqual(result["pass_rate"], 2 / 3)
        self.assertTrue(result["passed"])
        self.assertEqual([run["passed"] for run in result["runs"]], [True, False, True])

    def test_too_few_passes_fail_the_repeated_task(self):
        outcomes = iter([True, False, False])

        def agent(workspace):
            if next(outcomes):
                create_backup(_register(workspace), "eval-backup-1", workspace=workspace)

        with tempfile.TemporaryDirectory() as tmp:
            result = run_agent_eval.run_task_repeated(
                _task("register-and-backup"),
                "book.xlsx",
                "spreadsheets",
                Path(tmp),
                None,
                15,
                60,
                repeat=3,
                min_passes=2,
                runner=_fake_agent(agent),
            )

        self.assertEqual(result["passes"], 1)
        self.assertFalse(result["passed"])

    def test_a_skip_ends_the_repeated_task_after_one_attempt(self):
        task = {**_task("register-and-backup"), "requires_wps": True}
        probes = []
        with tempfile.TemporaryDirectory() as tmp:
            result = run_agent_eval.run_task_repeated(
                task,
                "book.xlsx",
                "spreadsheets",
                Path(tmp),
                None,
                15,
                60,
                repeat=3,
                min_passes=2,
                runner=_fake_agent(lambda workspace: self.fail("agent must not run")),
                has_wps=lambda component: probes.append(component) or False,
            )

        self.assertEqual(len(probes), 1)
        self.assertIn("skipped", result)
        summary = run_agent_eval.summarize([result], "stamp", None)
        self.assertEqual((summary["total"], summary["skipped"]), (0, 1))

    def test_a_setup_that_fails_on_a_later_run_is_a_failed_run_and_keeps_the_earlier_ones(self):
        # On Windows, a WPS Writer from run 1 can stay registered and hide the document run 2 opens (P5-006).
        reasons = iter([None, "WPS did not list report.docx as open within 45 s", None])
        released = []

        def setup(steps, workspace, fixture, component):
            reason = next(reasons)
            sha = run_agent_eval.file_sha256(fixture)
            if reason:
                return sha, None, reason
            return sha, (lambda: released.append(workspace)), None

        def agent(workspace):
            create_backup(_register(workspace), "eval-backup-1", workspace=workspace)

        with tempfile.TemporaryDirectory() as tmp, patch.object(run_agent_eval, "run_setup", side_effect=setup):
            result = run_agent_eval.run_task_repeated(
                _task("register-and-backup"),
                "book.xlsx",
                "spreadsheets",
                Path(tmp),
                None,
                15,
                60,
                repeat=3,
                min_passes=2,
                runner=_fake_agent(agent),
            )

        self.assertNotIn("skipped", result)
        self.assertEqual([run["passed"] for run in result["runs"]], [True, False, True])
        self.assertIn("did not list", result["runs"][1]["setup_failed"])
        self.assertEqual(result["runs"][1]["metrics"]["wps_tool_calls"], 0)
        self.assertEqual((result["passes"], result["passed"]), (2, True))
        self.assertEqual(len(released), 2)
        summary = run_agent_eval.summarize([result], "stamp", None)
        self.assertEqual((summary["passed"], summary["total"], summary["skipped"]), (1, 1, 0))

    def test_repeat_options_are_validated(self):
        for argv in (["--repeat", "0"], ["--repeat", "2", "--min-passes", "3"], ["--min-passes", "0"]):
            with self.subTest(argv=argv), self.assertRaises(SystemExit) as raised:
                run_agent_eval.main(argv)
            self.assertEqual(raised.exception.code, 2)

    def run_wps_task(self, task_id, root, actions, calls):
        task = _task(task_id)
        fixture = task.get("fixture", "book.xlsx")
        component = task.get("component", "spreadsheets")
        return run_agent_eval.run_task(
            task,
            fixture,
            component,
            root,
            None,
            15,
            60,
            runner=_fake_agent(actions, calls),
            has_wps=lambda component: True,
        )

    @staticmethod
    def simulate_write(workspace, component, fixture, request_id, command, edit):
        """What a successful WPS write leaves behind: a backup, the edited file and an operation record."""
        ok, record, errors = register_document(component, str(workspace / fixture), workspace)
        assert ok, errors
        document_id = record["document_id"]
        ok, _result, errors, _ = create_backup(document_id, f"{request_id}:backup", workspace=workspace)
        assert ok, errors
        edit(workspace / fixture)
        record_operation(request_id, command, {"document_id": document_id}, workspace)

    @staticmethod
    def replace_in_docx(path, old, new):
        from docx import Document

        document = Document(path)
        for paragraph in document.paragraphs:
            if old in paragraph.text:
                paragraph.text = paragraph.text.replace(old, new)
        document.save(path)

    @staticmethod
    def replace_in_pptx(path, old, new):
        from pptx import Presentation

        presentation = Presentation(path)
        for slide in presentation.slides:
            for shape in slide.shapes:
                if shape.has_text_frame and old in shape.text_frame.text:
                    shape.text_frame.text = shape.text_frame.text.replace(old, new)
        presentation.save(path)

    def test_wps_spreadsheet_write_passes_when_b2_is_written_after_a_backup(self):
        from openpyxl import load_workbook

        def edit(path):
            workbook = load_workbook(path)
            workbook.active["B2"] = 42
            workbook.save(path)

        with tempfile.TemporaryDirectory() as tmp:
            result = self.run_wps_task(
                "wps-spreadsheet-write",
                Path(tmp),
                lambda ws: self.simulate_write(
                    ws, "spreadsheets", "book.xlsx", "eval-write-1", "spreadsheet-write", edit
                ),
                [_call("spreadsheet_write", request_id="eval-write-1", range="B2", values_json="[[42]]")],
            )
        self.assertTrue(result["passed"], result)

    def test_wps_spreadsheet_write_fails_when_a1_is_overwritten(self):
        from openpyxl import load_workbook

        def edit(path):
            workbook = load_workbook(path)
            workbook.active["A1"] = 42
            workbook.save(path)

        with tempfile.TemporaryDirectory() as tmp:
            result = self.run_wps_task(
                "wps-spreadsheet-write",
                Path(tmp),
                lambda ws: self.simulate_write(
                    ws, "spreadsheets", "book.xlsx", "eval-write-1", "spreadsheet-write", edit
                ),
                [_call("spreadsheet_write", request_id="eval-write-1")],
            )
        failed = [(check["kind"], check.get("cell")) for check in result["checks"] if not check["passed"]]
        self.assertEqual(failed, [("xlsx_cell", "B2"), ("xlsx_cell", "A1")])

    def test_wps_writer_replace_passes_and_keeps_the_other_paragraphs(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = self.run_wps_task(
                "wps-writer-replace",
                Path(tmp),
                lambda ws: self.simulate_write(
                    ws,
                    "writer",
                    "report.docx",
                    "eval-writer-1",
                    "writer-replace",
                    lambda path: self.replace_in_docx(path, "DRAFT", "FINAL"),
                ),
                [_call("writer_replace", request_id="eval-writer-1", find="DRAFT", replace="FINAL")],
            )
        self.assertTrue(result["passed"], result)

    def test_wps_writer_replace_fails_when_draft_is_left(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = self.run_wps_task(
                "wps-writer-replace",
                Path(tmp),
                lambda ws: self.simulate_write(
                    ws, "writer", "report.docx", "eval-writer-1", "writer-replace", lambda path: None
                ),
                [_call("writer_replace", request_id="eval-writer-1")],
            )
        text_check = next(check for check in result["checks"] if check["kind"] == "file_text")
        self.assertFalse(text_check["passed"])
        self.assertEqual(text_check["detail"], {"missing": ["Status: FINAL"], "present": ["DRAFT"]})

    def test_wps_presentation_replace_passes(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = self.run_wps_task(
                "wps-presentation-replace",
                Path(tmp),
                lambda ws: self.simulate_write(
                    ws,
                    "presentation",
                    "deck.pptx",
                    "eval-presentation-1",
                    "presentation-replace",
                    lambda path: self.replace_in_pptx(path, "DRAFT", "FINAL"),
                ),
                [_call("presentation_replace", request_id="eval-presentation-1", find="DRAFT", replace="FINAL")],
            )
        self.assertTrue(result["passed"], result)

    def test_wps_presentation_replace_fails_when_the_title_is_lost(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = self.run_wps_task(
                "wps-presentation-replace",
                Path(tmp),
                lambda ws: self.simulate_write(
                    ws,
                    "presentation",
                    "deck.pptx",
                    "eval-presentation-1",
                    "presentation-replace",
                    lambda path: self.replace_in_pptx(path, "Quarterly review", "FINAL"),
                ),
                [_call("presentation_replace", request_id="eval-presentation-1")],
            )
        text_check = next(check for check in result["checks"] if check["kind"] == "file_text")
        self.assertEqual(text_check["detail"]["missing"], ["Quarterly review", "Status: FINAL"])

    def test_wps_held_write_passes_on_a_reported_refusal(self):
        with (
            tempfile.TemporaryDirectory() as tmp,
            patch.object(run_agent_eval, "hold_file", return_value=(lambda: None, None)),
        ):
            result = self.run_wps_task(
                "wps-held-spreadsheet-write",
                Path(tmp),
                lambda ws: None,
                [_call("spreadsheet_write", error_codes=["DOCUMENT_HELD_OPEN"], request_id="eval-held-write-1")],
            )
        self.assertTrue(result["passed"], result)

    def test_wps_tasks_are_skipped_without_wps(self):
        with tempfile.TemporaryDirectory() as tmp:
            for task_id in WPS_WRITE_TASKS:
                task = _task(task_id)
                result = run_agent_eval.run_task(
                    task,
                    task.get("fixture", "book.xlsx"),
                    task.get("component", "spreadsheets"),
                    Path(tmp),
                    None,
                    15,
                    60,
                    runner=_fake_agent(lambda ws: self.fail("agent must not run")),
                    has_wps=lambda component: False,
                )
                self.assertIn("needs WPS", result["skipped"], task_id)

    def test_sheets_and_read_passes_when_both_reads_were_made_and_nothing_changed(self):
        calls = [_call("register_document"), _call("spreadsheet_sheets"), _call("spreadsheet_read", range="A1")]
        with tempfile.TemporaryDirectory() as tmp:
            result = self.run_wps_task("spreadsheet-sheets-and-read", Path(tmp), _register, calls)
            missing_read = self.run_wps_task("spreadsheet-sheets-and-read", Path(tmp) / "b", _register, calls[:2])
        self.assertTrue(result["passed"], result)
        self.assertEqual(result["metrics"]["off_path_calls"], 0)
        self.assertEqual(
            [check["tool"] for check in missing_read["checks"] if not check["passed"]], ["spreadsheet_read"]
        )

    def test_html_table_to_docx_is_judged_from_the_output_document(self):
        from docx import Document

        def convert(workspace, rows):
            document = Document()
            table = document.add_table(rows=len(rows), cols=2)
            for row, values in zip(table.rows, rows, strict=True):
                for cell, value in zip(row.cells, values, strict=True):
                    cell.text = value
            document.save(workspace / "table.docx")

        calls = [_call("html_editable", format="docx")]
        rows = [("Region", "Q1"), ("North", "120"), ("South", "95")]
        with tempfile.TemporaryDirectory() as tmp:
            passed = self.run_wps_task("html-table-to-docx", Path(tmp), lambda ws: convert(ws, rows), calls)
            lost_row = self.run_wps_task("html-table-to-docx", Path(tmp) / "b", lambda ws: convert(ws, rows[:2]), calls)
        self.assertTrue(passed["passed"], passed)
        failed = [check for check in lost_row["checks"] if not check["passed"]]
        self.assertEqual([check["kind"] for check in failed], ["file_text"])
        self.assertEqual(failed[0]["detail"]["missing"], ["South", "95"])

    def test_rename_sheet_passes_when_only_the_sheet_name_changed(self):
        from openpyxl import load_workbook

        def rename(path, value="seed"):
            workbook = load_workbook(path)
            workbook.active.title = "Summary"
            workbook.active["A1"] = value
            workbook.save(path)

        calls = [_call("spreadsheet_rename_sheet", request_id="eval-rename-1", old_name="Sheet", new_name="Summary")]
        with tempfile.TemporaryDirectory() as tmp:
            passed = self.run_wps_task(
                "wps-spreadsheet-rename-sheet",
                Path(tmp),
                lambda ws: self.simulate_write(
                    ws, "spreadsheets", "book.xlsx", "eval-rename-1", "spreadsheet-rename-sheet", rename
                ),
                calls,
            )
            changed_cell = self.run_wps_task(
                "wps-spreadsheet-rename-sheet",
                Path(tmp) / "b",
                lambda ws: self.simulate_write(
                    ws,
                    "spreadsheets",
                    "book.xlsx",
                    "eval-rename-1",
                    "spreadsheet-rename-sheet",
                    lambda path: rename(path, "lost"),
                ),
                calls,
            )
        self.assertTrue(passed["passed"], passed)
        self.assertEqual([c["kind"] for c in changed_cell["checks"] if not c["passed"]], ["xlsx_cell"])

    def test_create_sheet_needs_the_new_sheet_after_the_existing_one(self):
        from openpyxl import load_workbook

        def add(path, index=None):
            workbook = load_workbook(path)
            workbook.create_sheet("Q4", index)
            workbook.save(path)

        calls = [_call("spreadsheet_create_sheet", request_id="eval-create-1", name="Q4")]
        with tempfile.TemporaryDirectory() as tmp:
            passed = self.run_wps_task(
                "wps-spreadsheet-create-sheet",
                Path(tmp),
                lambda ws: self.simulate_write(
                    ws, "spreadsheets", "book.xlsx", "eval-create-1", "spreadsheet-create-sheet", add
                ),
                calls,
            )
            first = self.run_wps_task(
                "wps-spreadsheet-create-sheet",
                Path(tmp) / "b",
                lambda ws: self.simulate_write(
                    ws,
                    "spreadsheets",
                    "book.xlsx",
                    "eval-create-1",
                    "spreadsheet-create-sheet",
                    lambda path: add(path, 0),
                ),
                calls,
            )
        self.assertTrue(passed["passed"], passed)
        failed = [check for check in first["checks"] if not check["passed"]]
        self.assertEqual([check["kind"] for check in failed], ["xlsx_sheets"])
        self.assertEqual(failed[0]["detail"], ["Q4", "Sheet"])

    def test_writer_table_write_passes_when_only_the_north_status_changed(self):
        from docx import Document

        def set_cells(path, cells):
            document = Document(path)
            for row, column in cells:
                document.tables[0].cell(row, column).text = "FINAL"
            document.save(path)

        calls = [_call("writer_table_write", request_id="eval-table-1", table_index=1, row=2, column=3)]
        with tempfile.TemporaryDirectory() as tmp:
            passed = self.run_wps_task(
                "wps-writer-table-write",
                Path(tmp),
                lambda ws: self.simulate_write(
                    ws,
                    "writer",
                    "table_report.docx",
                    "eval-table-1",
                    "writer-table-write",
                    lambda path: set_cells(path, [(1, 2)]),
                ),
                calls,
            )
            too_much = self.run_wps_task(
                "wps-writer-table-write",
                Path(tmp) / "b",
                lambda ws: self.simulate_write(
                    ws,
                    "writer",
                    "table_report.docx",
                    "eval-table-1",
                    "writer-table-write",
                    lambda path: set_cells(path, [(1, 2), (2, 2)]),
                ),
                calls,
            )
        self.assertTrue(passed["passed"], passed)
        failed = [(check["kind"], check.get("row")) for check in too_much["checks"] if not check["passed"]]
        self.assertEqual(failed, [("docx_table_cell", 3)])

    def test_docx_to_pdf_needs_a_real_pdf_and_an_unchanged_source(self):
        calls = [_call("convert_smoke", component="writer", format="pdf")]

        def write(workspace, data):
            (workspace / "report.pdf").write_bytes(data)

        with tempfile.TemporaryDirectory() as tmp:
            passed = self.run_wps_task(
                "wps-docx-to-pdf", Path(tmp), lambda ws: write(ws, b"%PDF-1.7\n" + b"0" * 2000), calls
            )
            not_pdf = self.run_wps_task(
                "wps-docx-to-pdf", Path(tmp) / "b", lambda ws: write(ws, b"<html>" * 400), calls
            )
            too_small = self.run_wps_task(
                "wps-docx-to-pdf", Path(tmp) / "c", lambda ws: write(ws, b"%PDF-1.7\n"), calls
            )
        self.assertTrue(passed["passed"], passed)
        for result in (not_pdf, too_small):
            self.assertEqual([check["kind"] for check in result["checks"] if not check["passed"]], ["pdf_file"])

    def test_open_document_text_export_is_read_as_unicode_text(self):
        text = "\r\n".join(run_agent_eval.WRITER_FIXTURE_PARAGRAPHS)
        calls = [_call("open_documents"), _call("export_open_document_text", output="report.txt")]

        def export(workspace, data):
            (workspace / "report.txt").write_bytes(data)

        task = _task("wps-open-document-export-text")
        with (
            tempfile.TemporaryDirectory() as tmp,
            patch.object(run_agent_eval, "open_in_wps", return_value=(lambda: None, None)) as opened,
        ):
            utf16 = self.run_wps_task(
                task["id"],
                Path(tmp),
                lambda ws: export(ws, "\ufeff".encode("utf-16-le") + text.encode("utf-16-le")),
                calls,
            )
            ansi = self.run_wps_task(task["id"], Path(tmp) / "b", lambda ws: export(ws, text.encode("ascii")), calls)
            saved = self.run_wps_task(
                task["id"],
                Path(tmp) / "c",
                lambda ws: export(ws, text.encode("ascii")),
                [*calls, _call("save_open_document")],
            )
        self.assertTrue(utf16["passed"], utf16)
        self.assertTrue(ansi["passed"], ansi)
        self.assertEqual([c["kind"] for c in saved["checks"] if not c["passed"]], ["tool_not_called"])
        self.assertEqual(opened.call_count, 3)

    def test_open_in_wps_is_a_skip_off_windows(self):
        if os.name == "nt":
            self.skipTest("Windows opens the file for real")
        with tempfile.TemporaryDirectory() as tmp:
            result = run_agent_eval.run_task(
                _task("wps-open-document-export-text"),
                "report.docx",
                "writer",
                Path(tmp),
                None,
                15,
                60,
                runner=_fake_agent(lambda ws: self.fail("agent must not run")),
                has_wps=lambda component: True,
            )
        self.assertIn("needs Windows", result["skipped"])

    def test_the_open_document_is_closed_after_the_agent_ran(self):
        released = []
        with (
            tempfile.TemporaryDirectory() as tmp,
            patch.object(run_agent_eval, "open_in_wps", return_value=(lambda: released.append(True), None)),
        ):
            self.run_wps_task("wps-open-document-export-text", Path(tmp), lambda ws: None, [])
        self.assertEqual(released, [True])

    def test_mcp_config_runs_the_server_in_the_task_workspace(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "task"
            workspace.mkdir()
            config = json.loads(run_agent_eval.write_mcp_config(workspace).read_text(encoding="utf-8"))

        server = config["mcpServers"]["wps-agent"]
        self.assertEqual(server["cwd"], str(workspace))
        self.assertEqual(server["args"], ["-m", "wps_ai_agent_cli", "mcp-server"])
        self.assertEqual(server["env"]["PYTHONPATH"], str(REPO_ROOT / "src"))


if __name__ == "__main__":
    unittest.main()
