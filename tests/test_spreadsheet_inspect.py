from __future__ import annotations

from datetime import datetime
import hashlib
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900, to_excel

from wps_ai_agent_cli.sessions import register_document
from wps_ai_agent_cli.spreadsheet_ops import read_spreadsheet_range
from wps_ai_agent_cli.spreadsheet_inspect import _cache_matches, inspect_spreadsheet_range


class SpreadsheetInspectTests(unittest.TestCase):
    def test_cache_comparison_handles_numbers_dates_errors_and_empty_values(self):
        epoch = datetime(1899, 12, 30)
        self.assertTrue(_cache_matches(25, 25.0, epoch))
        self.assertFalse(_cache_matches(25, 26, epoch))
        date_value = datetime(2026, 10, 7)
        self.assertTrue(_cache_matches(date_value, to_excel(date_value, epoch), epoch))
        self.assertTrue(_cache_matches("#DIV/0!", "#DIV/0!", epoch))
        self.assertTrue(_cache_matches(None, None, epoch))
        self.assertFalse(_cache_matches(None, 0, epoch))

    def test_cache_comparison_uses_the_workbook_1904_date_system(self):
        value = datetime(2024, 2, 29, 12, 30)
        mac_serial = to_excel(value, CALENDAR_MAC_1904)
        windows_serial = to_excel(value, CALENDAR_WINDOWS_1900)
        self.assertTrue(_cache_matches(value, mac_serial, CALENDAR_MAC_1904))
        self.assertFalse(_cache_matches(value, mac_serial, CALENDAR_WINDOWS_1900))
        self.assertNotEqual(mac_serial, windows_serial)

    def test_inspect_rejects_invalid_range_without_wps_or_file_changes(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "sample.xlsx"
            Workbook().save(document)
            ok, record, _ = register_document("spreadsheets", str(document), workspace)
            before = hashlib.sha256(document.read_bytes()).hexdigest()

            for address, expected_code in (("A:A", "INVALID_RANGE"), ("A1:XFD1048576", "RANGE_TOO_LARGE")):
                with self.subTest(address=address), patch("wps_ai_agent_cli.spreadsheet_inspect.load_workbook") as open_workbook, patch("wps_ai_agent_cli.spreadsheet_inspect.probe_wps_capabilities") as probe:
                    ok, _result, errors = inspect_spreadsheet_range(
                        record["document_id"], address, workspace=workspace,
                    )
                    self.assertFalse(ok)
                    self.assertEqual(errors[0]["code"], expected_code)
                    open_workbook.assert_not_called()
                    probe.assert_not_called()
            self.assertEqual(hashlib.sha256(document.read_bytes()).hexdigest(), before)


@unittest.skipUnless(os.environ.get("WPS_AGENT_RUN_INTEGRATION") == "1", "Requires explicit local WPS integration run")
class SpreadsheetInspectWpsIntegrationTests(unittest.TestCase):
    def test_1904_date_system_and_locale_tagged_display_are_reported_without_saving(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "locale.xlsx"
            workbook = Workbook()
            workbook.epoch = CALENDAR_MAC_1904
            sheet = workbook.active
            sheet.title = "Locale"
            sheet["A1"] = datetime(2024, 2, 29)
            sheet["A1"].number_format = "[$-409]mmm d, yyyy"
            sheet["B1"] = 0.125
            sheet["B1"].number_format = "[$-407]0,0%"
            sheet["C1"] = "=DATE(2024,2,29)"
            sheet["C1"].number_format = "[$-409]mmm d, yyyy"
            sheet["D1"] = 1234.5
            sheet["D1"].number_format = "[$-407]#.##0,00"
            sheet.column_dimensions["A"].width = 20
            sheet.column_dimensions["C"].width = 20
            sheet.column_dimensions["D"].width = 20
            workbook.save(document)
            workbook.close()
            before = hashlib.sha256(document.read_bytes()).hexdigest()
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)

            ok, result, errors = inspect_spreadsheet_range(record["document_id"], "A1:D1", "Locale", workspace)

            self.assertTrue(ok, errors)
            cells = {cell["address"]: cell for cell in result["cells"]}
            self.assertEqual(cells["A1"]["saved_cached_value"], "2024-02-29T00:00:00")
            self.assertIn("2024", cells["A1"]["displayed_text"])
            self.assertIn("%", cells["B1"]["displayed_text"])
            self.assertTrue(any(character.isdigit() for character in cells["B1"]["displayed_text"]))
            self.assertIn("2024", cells["C1"]["displayed_text"])
            self.assertNotIn("###", cells["C1"]["displayed_text"])
            self.assertNotIn("###", cells["D1"]["displayed_text"])
            self.assertIn("1234", "".join(character for character in cells["D1"]["displayed_text"] if character.isdigit()))
            self.assertIn("[$-407]", cells["B1"]["number_format"])
            self.assertIn("[$-407]", cells["D1"]["number_format"])
            self.assertTrue(result["read_only"])
            self.assertFalse(result["saved"])
            self.assertEqual(hashlib.sha256(document.read_bytes()).hexdigest(), before)

    def test_read_metadata_matches_wps_display_categories(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "formats.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Formats"
            sheet["A1"] = datetime(2026, 10, 7)
            sheet["A1"].number_format = "yyyy-mm-dd"
            sheet["B1"] = 0.25
            sheet["B1"].number_format = "0.0%"
            sheet["C1"] = 1234.5
            sheet["C1"].number_format = "$#,##0.00"
            sheet.column_dimensions["A"].width = 16
            sheet.column_dimensions["B"].width = 12
            sheet.column_dimensions["C"].width = 16
            workbook.save(document)
            workbook.close()
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)

            read_ok, read_result, read_errors = read_spreadsheet_range(record["document_id"], "A1:D1", workspace=workspace)
            inspect_ok, inspect_result, inspect_errors = inspect_spreadsheet_range(record["document_id"], "A1:D1", "Formats", workspace)

            self.assertTrue(read_ok, read_errors)
            self.assertTrue(inspect_ok, inspect_errors)
            metadata = read_result["cell_metadata"][0]
            live = {cell["address"]: cell for cell in inspect_result["cells"]}
            self.assertEqual([cell["format_category"] for cell in metadata], ["date", "percentage", "currency", "empty"])
            self.assertIn("2026", live["A1"]["displayed_text"])
            self.assertIn("%", live["B1"]["displayed_text"])
            self.assertIn("$", live["C1"]["displayed_text"])
            self.assertEqual(live["D1"]["displayed_text"], "")

    def test_reads_formula_cache_recalculation_and_display_without_saving(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            workspace.mkdir()
            document = Path(tmp) / "inspect.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Contract"
            sheet["A1"] = 12.5
            sheet["B1"] = datetime(2026, 10, 7)
            sheet["B1"].number_format = "yyyy-mm-dd"
            sheet["C1"] = "=A1*2"
            sheet["D1"] = "=DATE(2026,10,7)"
            sheet["D1"].number_format = "yyyy-mm-dd"
            sheet["E1"] = "=1/0"
            sheet.column_dimensions["D"].width = 16
            workbook.save(document)
            workbook.close()
            before = hashlib.sha256(document.read_bytes()).hexdigest()
            ok, record, errors = register_document("spreadsheets", str(document), workspace)
            self.assertTrue(ok, errors)

            ok, result, errors = inspect_spreadsheet_range(
                record["document_id"], "A1:F1", "Contract", workspace,
            )

            self.assertTrue(ok, errors)
            cells = {cell["address"]: cell for cell in result["cells"]}
            self.assertEqual(cells["A1"]["recalculated_value"], 12.5)
            self.assertIn("12.5", cells["A1"]["displayed_text"])
            self.assertEqual(cells["C1"]["formula"], "=A1*2")
            self.assertEqual(cells["C1"]["wps_formula"], "=A1*2")
            self.assertEqual(cells["C1"]["recalculated_value"], 25)
            self.assertEqual(cells["C1"]["saved_cached_value"], None)
            self.assertFalse(cells["C1"]["cache_matches_recalculated"])
            self.assertEqual(cells["D1"]["formula"], "=DATE(2026,10,7)")
            self.assertNotIn("###", cells["D1"]["displayed_text"])
            self.assertIn("2026", cells["D1"]["displayed_text"])
            self.assertEqual(cells["E1"]["recalculated_type"], "error")
            self.assertTrue(cells["E1"]["displayed_text"].startswith("#"))
            self.assertEqual(cells["F1"]["recalculated_type"], "empty")
            self.assertEqual(cells["F1"]["displayed_text"], "")
            self.assertTrue(result["read_only"])
            self.assertFalse(result["saved"])
            self.assertEqual(hashlib.sha256(document.read_bytes()).hexdigest(), before)


if __name__ == "__main__":
    unittest.main()
