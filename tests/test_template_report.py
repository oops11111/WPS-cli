import json
import hashlib
from io import BytesIO
import os
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
import zipfile
from unittest.mock import patch

from docx import Document
from wps_ai_agent_cli.template_report import render_batch_template_report
from wps_ai_agent_cli.batch_conversion import convert_html_batch
from wps_ai_agent_cli.com_backend import run_com_smoke
from wps_ai_agent_cli.com_backend import run_conversion_smoke


class TemplateReportTests(unittest.TestCase):
    def _inputs(self, root: Path, template: str):
        source_root, output_root = root / "inputs", root / "artifacts"
        source_root.mkdir()
        output_root.mkdir()
        source_passed, source_failed = source_root / "a.html", source_root / "bad.html"
        output_passed = output_root / "a.docx"
        source_passed.write_bytes(b"source a")
        source_failed.write_bytes(b"source b")
        output_passed.write_bytes(b"docx a")
        digest = lambda path: hashlib.sha256(path.read_bytes()).hexdigest()
        manifest = {
            "schema_version": "wps-agent-batch-conversion/v1",
            "mode": "docx", "source_directory": str(source_root),
            "output_directory": str(output_root),
            "summary": {"total": 2, "passed": 1, "failed": 1},
            "files": [
                {"status": "passed", "relative_source": "a.html", "source": str(source_passed), "output": str(output_passed), "source_sha256": digest(source_passed), "output_sha256": digest(output_passed), "errors": []},
                {"status": "failed", "relative_source": "bad.html", "source": str(source_failed), "output": str(output_root / "bad.docx"), "source_sha256": digest(source_failed), "errors": [{"code": "FAILED", "message": "missing <body> | review"}]},
            ],
        }
        manifest_path, template_path = root / "batch.json", root / "report.md"
        manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
        template_path.write_text(template, encoding="utf-8")
        return manifest_path, template_path

    def test_renders_verified_manifest_summary_and_file_rows(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            manifest, template = self._inputs(root, "# {{mode}}: {{total}} files ({{passed}} passed, {{failed}} failed)\n{{files_table}}")
            output = root / "final.md"
            ok, data, errors = render_batch_template_report(manifest, template, output)
            self.assertTrue(ok, errors)
            report = output.read_text(encoding="utf-8")
            self.assertIn("| passed | a.html |", report)
            self.assertIn("missing &lt;body&gt; \\| review", report)
            self.assertEqual(data["summary"], {"total": 2, "passed": 1, "failed": 1})
            self.assertEqual(len(data["manifest_sha256"]), 64)

    def test_rejects_unknown_placeholders_and_inconsistent_summary(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            manifest, template = self._inputs(root, "{{unknown_token}}")
            ok, _, errors = render_batch_template_report(manifest, template, root / "unknown.md")
            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "UNKNOWN_TEMPLATE_TOKEN")
            payload = json.loads(manifest.read_text(encoding="utf-8"))
            payload["summary"]["failed"] = 0
            manifest.write_text(json.dumps(payload), encoding="utf-8")
            template.write_text("{{total}}", encoding="utf-8")
            ok, _, errors = render_batch_template_report(manifest, template, root / "invalid.md")
            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "INVALID_MANIFEST")

    def test_refuses_to_overwrite_existing_report(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            manifest, template = self._inputs(root, "{{total}}")
            output = root / "existing.md"
            output.write_text("keep", encoding="utf-8")
            ok, _, errors = render_batch_template_report(manifest, template, output)
            self.assertFalse(ok)
            self.assertEqual(output.read_text(encoding="utf-8"), "keep")
            self.assertEqual(errors[0]["code"], "OUTPUT_ALREADY_EXISTS")

    def test_rejects_source_modified_after_manifest_creation(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            manifest, template = self._inputs(root, "{{total}}")
            (root / "inputs" / "a.html").write_bytes(b"changed")
            ok, _, errors = render_batch_template_report(manifest, template, root / "result.md")
            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "SOURCE_HASH_MISMATCH")

    def test_rejects_converted_output_modified_after_manifest_creation(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            manifest, template = self._inputs(root, "{{total}}")
            (root / "artifacts" / "a.docx").write_bytes(b"changed output")
            ok, _, errors = render_batch_template_report(manifest, template, root / "result.md")
            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "OUTPUT_HASH_MISMATCH")

    def test_renders_cancelled_manifest_with_planned_and_processed_counts(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            manifest, template = self._inputs(root, "{{total}} planned; {{processed}} processed; {{cancelled}}")
            payload = json.loads(manifest.read_text(encoding="utf-8"))
            payload["files"] = payload["files"][:1]
            payload["cancelled"] = True
            payload["summary"] = {"total": 2, "processed": 1, "passed": 1, "failed": 0, "cancelled": True}
            manifest.write_text(json.dumps(payload), encoding="utf-8")
            ok, _, errors = render_batch_template_report(manifest, template, root / "cancelled.md")
            self.assertTrue(ok, errors)
            self.assertEqual((root / "cancelled.md").read_text(encoding="utf-8"), "2 planned; 1 processed; true")

    def test_renders_rows_into_docx_template_and_preserves_table_style(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            manifest, _ = self._inputs(root, "unused")
            template_path = root / "report-template.docx"
            template = Document()
            paragraph = template.add_paragraph("Mode: ")
            value = paragraph.add_run("{{mode}}")
            value.bold = True
            table = template.add_table(rows=2, cols=3)
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells, ("Status", "Source", "Output")):
                cell.text = text
            for cell, text in zip(table.rows[1].cells, ("{{file_status}}", "{{file_source}}", "{{file_output}}")):
                cell.text = text
            template.save(template_path)
            output = root / "report.docx"
            ok, data, errors = render_batch_template_report(manifest, template_path, output)
            self.assertTrue(ok, errors)
            rendered = Document(output)
            self.assertEqual(rendered.paragraphs[0].text, "Mode: docx")
            self.assertTrue(rendered.paragraphs[0].runs[1].bold)
            self.assertEqual(len(rendered.tables[0].rows), 3)
            self.assertEqual(rendered.tables[0].style.name, "Table Grid")
            self.assertEqual([row.cells[0].text for row in rendered.tables[0].rows[1:]], ["passed", "failed"])
            self.assertEqual(data["file_count"], 2)

    def test_docx_template_rejects_placeholder_split_across_runs(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            manifest, _ = self._inputs(root, "unused")
            template_path = root / "split.docx"
            template = Document()
            paragraph = template.add_paragraph()
            paragraph.add_run("{{mo")
            paragraph.add_run("de}}")
            row = template.add_table(rows=2, cols=1)
            row.cell(1, 0).text = "{{file_source}}"
            template.save(template_path)
            ok, _, errors = render_batch_template_report(manifest, template_path, root / "split-output.docx")
            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "INVALID_TEMPLATE")

    def test_rejects_docx_zip_bomb_before_document_parsing(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            manifest, _ = self._inputs(root, "unused")
            archive_bytes = BytesIO()
            with zipfile.ZipFile(archive_bytes, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                archive.writestr("[Content_Types].xml", "<Types/>")
                archive.writestr("word/document.xml", "x" * (50 * 1024 * 1024 + 1))
            template = root / "oversized.docx"
            template.write_bytes(archive_bytes.getvalue())
            output = root / "oversized-report.docx"
            ok, _, errors = render_batch_template_report(manifest, template, output)
            self.assertFalse(ok)
            self.assertFalse(output.exists())
            self.assertEqual(errors[0]["code"], "TEMPLATE_ARCHIVE_TOO_LARGE")

    def test_preflights_total_output_hash_budget_before_hashing(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            manifest, template = self._inputs(root, "{{total}}")
            with patch("wps_ai_agent_cli.template_report._MAX_TOTAL_OUTPUT_BYTES", 1):
                ok, _, errors = render_batch_template_report(manifest, template, root / "limited.md")
            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "OUTPUT_TOO_LARGE")
            self.assertFalse((root / "limited.md").exists())


@unittest.skipUnless(os.environ.get("WPS_AGENT_RUN_INTEGRATION") == "1", "set WPS_AGENT_RUN_INTEGRATION=1 to launch local WPS Writer")
class TemplateReportWpsIntegrationTests(unittest.TestCase):
    def test_docx_template_report_opens_and_saves_in_wps(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            manifest, _ = TemplateReportTests()._inputs(root, "unused")
            template_path, report_path = root / "template.docx", root / "report.docx"
            saved_copy = root / "wps-report.docx"
            template = Document()
            template.add_paragraph("Batch total: {{total}}")
            table = template.add_table(rows=2, cols=2)
            table.style = "Table Grid"
            table.rows[0].cells[0].text = "Status"
            table.rows[0].cells[1].text = "Source"
            table.rows[1].cells[0].text = "{{file_status}}"
            table.rows[1].cells[1].text = "{{file_source}}"
            template.save(template_path)
            ok, data, errors = render_batch_template_report(manifest, template_path, report_path)
            self.assertTrue(ok, errors)
            wps_result = run_com_smoke("writer", str(report_path), str(saved_copy), visible=False)
            self.assertTrue(wps_result["ok"], wps_result)
            opened = Document(saved_copy)
            self.assertEqual(opened.paragraphs[0].text, "Batch total: 2")
            self.assertEqual(len(opened.tables[0].rows), 3)
            self.assertTrue(data["output_sha256"])

    def test_multipage_docx_report_retains_all_rows_after_wps_roundtrip(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source_root, output_root = root / "inputs", root / "artifacts"
            source_root.mkdir()
            output_root.mkdir()
            entries = []
            for index in range(50):
                name = f"item-{index:03d}"
                source, artifact = source_root / f"{name}.html", output_root / f"{name}.docx"
                source.write_text(f"source {name}", encoding="utf-8")
                artifact.write_bytes(f"converted {name}".encode("utf-8"))
                entries.append({
                    "status": "passed", "relative_source": source.name,
                    "source": str(source), "output": str(artifact),
                    "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                    "output_sha256": hashlib.sha256(artifact.read_bytes()).hexdigest(), "errors": [],
                })
            manifest = root / "batch.json"
            manifest.write_text(json.dumps({
                "schema_version": "wps-agent-batch-conversion/v1", "mode": "docx",
                "source_directory": str(source_root), "output_directory": str(output_root),
                "summary": {"total": 50, "passed": 50, "failed": 0}, "files": entries,
            }), encoding="utf-8")
            template_path = root / "template.docx"
            template = Document()
            template.add_paragraph("Processed: {{total}}  Passed: {{passed}}")
            table = template.add_table(rows=2, cols=2)
            table.style = "Table Grid"
            table.rows[0].cells[0].text = "Status"
            table.rows[0].cells[1].text = "Source"
            table.rows[1].cells[0].text = "{{file_status}}"
            table.rows[1].cells[1].text = "{{file_source}}"
            template.save(template_path)
            report, saved_copy, pdf = root / "report.docx", root / "wps-report.docx", root / "report.pdf"
            ok, _, errors = render_batch_template_report(manifest, template_path, report)
            self.assertTrue(ok, errors)
            opened = run_com_smoke("writer", str(report), str(saved_copy), visible=False)
            self.assertTrue(opened["ok"], opened)
            document = Document(saved_copy)
            self.assertEqual(document.paragraphs[0].text, "Processed: 50  Passed: 50")
            self.assertEqual(len(document.tables[0].rows), 51)
            self.assertEqual({row.cells[1].text for row in document.tables[0].rows[1:]}, {item["relative_source"] for item in entries})
            converted = run_conversion_smoke("writer", str(saved_copy), str(pdf), "pdf")
            self.assertTrue(converted["ok"], converted)
            from pypdf import PdfReader
            self.assertGreater(len(PdfReader(str(pdf)).pages), 1)

    def test_cancelled_batch_report_keeps_partial_rows_after_wps_roundtrip(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source, output = root / "input", root / "converted"
            source.mkdir()
            for name in ("done", "pending"):
                (source / f"{name}.html").write_text(name, encoding="utf-8")

            def convert(path, destination):
                destination.write_bytes(b"generated docx fixture")
                return True, {}, []

            def progress(completed, total, message):
                return completed < 1

            with patch("wps_ai_agent_cli.batch_conversion.convert_html_editable", side_effect=convert):
                ok, manifest, errors = convert_html_batch(source, output, "docx", progress_callback=progress)
            self.assertFalse(ok)
            self.assertEqual(errors, [])
            template_path = root / "report-template.docx"
            template = Document()
            template.add_paragraph("Planned {{total}}, processed {{processed}}, cancelled {{cancelled}}")
            table = template.add_table(rows=2, cols=2)
            table.style = "Table Grid"
            table.rows[0].cells[0].text = "Status"
            table.rows[0].cells[1].text = "Source"
            table.rows[1].cells[0].text = "{{file_status}}"
            table.rows[1].cells[1].text = "{{file_source}}"
            template.save(template_path)
            report, saved = root / "partial-report.docx", root / "wps-partial-report.docx"
            ok, _, errors = render_batch_template_report(manifest["manifest_path"], template_path, report)
            self.assertTrue(ok, errors)
            opened = run_com_smoke("writer", str(report), str(saved), visible=False)
            self.assertTrue(opened["ok"], opened)
            document = Document(saved)
            self.assertEqual(document.paragraphs[0].text, "Planned 2, processed 1, cancelled true")
            self.assertEqual(len(document.tables[0].rows), 2)
            self.assertEqual(document.tables[0].rows[1].cells[1].text, "done.html")


if __name__ == "__main__":
    unittest.main()
