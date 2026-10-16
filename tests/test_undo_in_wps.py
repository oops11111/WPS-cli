"""A failed write to a document WPS has open is undone inside WPS (P7-001).

WPS holds the file it has open, so copying the pre-write backup over it is always blocked. Each of the eight
open-document write commands now undoes a write that WPS saved but a later check rejected with the inverse edit in
the same WPS instance, then compares the saved file with the backup.
"""

import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from tests.test_open_documents import Fixture, com_ok, selection_payload
from wps_ai_agent_cli import open_documents
from wps_ai_agent_cli.atomic_mutation import atomic_mutation
from wps_ai_agent_cli.errors import ROLLBACK_BLOCKED, ROLLBACK_IN_WPS_FAILED, ROLLED_BACK

RUNNERS = {
    "writer": "_run_writer_attached_com",
    "presentation": "_run_presentation_attached_com",
    "spreadsheets": "_run_spreadsheet_attached_com",
}


def cell_selection(text="P3-402 mark"):
    return {
        "ok": True,
        "status": "ok",
        "saved": True,
        "read_only": False,
        "selection_kind": "cell",
        "address": "$A$1",
        "sheet_name": "Sheet",
        "cell_count": 1,
        "area_count": 1,
        "row_count": 1,
        "column_count": 1,
        "selection_text": text,
        "sample_truncated": False,
        "formula": None,
    }


def text_selection(text="DRAFT", **extra):
    return {
        "ok": True,
        "status": "ok",
        "saved": True,
        "read_only": False,
        "selection_type": 3,
        "selection_kind": "text",
        "selection_text": text,
        "slide_index": 1,
        "shape_name": "Title 1",
        "shape_count": 0,
        **extra,
    }


def shape_capture(before="Status: DRAFT", start=9):
    return {"slide_index": 1, "shape_name": "Title 1", "text_start": start, "shape_text_before": before}


def cell_capture(after, before="P3-402 mark"):
    return {
        "address": "$A$1",
        "sheet_name": "Sheet",
        "before_cells": [{"formula": before, "text_constant": True}],
        "after_formulas": [after],
    }


# Each command with a selection, the write script's payload that fails the command's read-back check, and the
# inverse edit the command must ask WPS for.
CASES = {
    "writer-replace": (
        "writer",
        open_documents.replace_writer_selection,
        {"text": "APPROVED"},
        selection_payload(),
        {"read_back_text": "garbled"},
        {"start": 10, "written_text": "APPROVED", "original_text": "PENDING_STATUS"},
    ),
    "writer-insert": (
        "writer",
        open_documents.insert_writer_selection,
        {"text": "HELLO"},
        selection_payload(text="", start=12, end=12),
        {"read_back_text": "garbled", "start": 12, "end": 17},
        {"start": 12, "written_text": "HELLO", "original_text": ""},
    ),
    "writer-delete": (
        "writer",
        open_documents.delete_writer_selection,
        {},
        selection_payload(),
        {"read_back_text": "PENDING_STATUS"},
        {"start": 10, "written_text": "", "original_text": "PENDING_STATUS"},
    ),
    "presentation-replace": (
        "presentation",
        open_documents.replace_presentation_selection,
        {"text": "FINAL"},
        text_selection(),
        {"read_back_text": "garbled", **shape_capture()},
        {
            "slide_index": 1,
            "shape_name": "Title 1",
            "start": 9,
            "written_text": "FINAL",
            "original_text": "DRAFT",
            "frame_expected": "Status: FINAL",
            "frame_restored": "Status: DRAFT",
        },
    ),
    "presentation-insert": (
        "presentation",
        open_documents.insert_presentation_selection,
        {"text": "DRAFT"},
        text_selection(text="", text_range_start=9, text_range_length=0, collapsed=True),
        {"read_back_text": "garbled", "start": 9, "end": 14, **shape_capture(before="Status: ")},
        {
            "slide_index": 1,
            "shape_name": "Title 1",
            "start": 9,
            "written_text": "DRAFT",
            "original_text": "",
            "frame_expected": "Status: DRAFT",
            "frame_restored": "Status: ",
        },
    ),
    "presentation-delete": (
        "presentation",
        open_documents.delete_presentation_selection,
        {},
        text_selection(),
        {"read_back_text": "DRAFT", **shape_capture()},
        {
            "slide_index": 1,
            "shape_name": "Title 1",
            "start": 9,
            "written_text": "",
            "original_text": "DRAFT",
            "frame_expected": "Status: ",
            "frame_restored": "Status: DRAFT",
        },
    ),
    "spreadsheet-replace": (
        "spreadsheets",
        open_documents.replace_spreadsheet_selection,
        {"text": "NEW"},
        cell_selection(),
        {"cells_matched": False, "read_back_text": "OLD", "cell_count": 1, "area_count": 1, **cell_capture("NEW")},
        cell_capture("NEW"),
    ),
    "spreadsheet-delete": (
        "spreadsheets",
        open_documents.delete_spreadsheet_selection,
        {},
        cell_selection(),
        {"cells_matched": False, "read_back_text": "P3-402 mark", "cell_count": 1, "area_count": 1, **cell_capture("")},
        cell_capture(""),
    ),
}


def object_selection(text, **flags):
    return {
        **cell_selection(text),
        "selection_kind": "non_range",
        "address": None,
        "cell_count": 0,
        "area_count": 0,
        "object_kind": "shape" if flags.get("has_text_frame") else "chart",
        "object_name": "Chart 1",
        **flags,
    }


def object_revert(original, written, read_back, formula=""):
    return {
        "expected_object_name": "Chart 1",
        "original_text": original,
        "written_text": written,
        "read_back_text": read_back,
        "original_formula": formula,
    }


# Shapes and charts (P10-004): the write recorded no cells; the undo finds the object again and puts the text back.
OBJECT_CASES = {
    "shape-text-replace": (
        "spreadsheets",
        open_documents.replace_spreadsheet_selection,
        {"text": "NEW"},
        object_selection("OLD", has_text_frame=True),
        {"cells_matched": False, "read_back_text": "NE"},
        object_revert("OLD", "NEW", "NE"),
        "selection-shape-revert",
    ),
    "chart-title-delete": (
        "spreadsheets",
        open_documents.delete_spreadsheet_selection,
        {},
        object_selection("Q3 sales", has_chart_title=True),
        {"cells_matched": False, "read_back_text": "Q3"},
        object_revert("Q3 sales", "", "Q3"),
        "selection-chart-title-revert",
    ),
    "axis-title-replace": (
        "spreadsheets",
        open_documents.replace_spreadsheet_selection,
        {"text": "Month"},
        object_selection("Quarter", has_axis_title=True),
        {"cells_matched": False, "read_back_text": "Mon"},
        object_revert("Quarter", "Month", "Mon"),
        "selection-axis-title-revert",
    ),
    "series-values-replace": (
        "spreadsheets",
        open_documents.replace_spreadsheet_selection,
        {"text": "30\t40"},
        object_selection("10\t20", has_series_values=True, series_name="Sales"),
        {"cells_matched": False, "read_back_text": "30\t4", "before_formula": "=SERIES(,,Sheet1!$B$2:$B$3,1)"},
        object_revert("10\t20", "30\t40", "30\t4", "=SERIES(,,Sheet1!$B$2:$B$3,1)"),
        "selection-series-values-revert",
    ),
}
CASES.update({name: case[:6] for name, case in OBJECT_CASES.items()})


class UndoInWpsTests(unittest.TestCase):
    def run_case(self, name, *, revert_status="reverted", revert_restores=True, revert_writes=None, write=None):
        component, command, arguments, selection, payload, _expected = CASES[name]
        fixture = Fixture(self, component=component)
        original = fixture.document.read_bytes()
        reverts = []
        self.revert_actions = []

        def fake(action, params):
            if action == "selection-read":
                return com_ok(selection)
            if action.endswith("-revert"):
                self.revert_actions.append(action)
                reverts.append({key: value for key, value in params.items() if key != "path"})
                if revert_writes is not None:
                    revert_writes(fixture.document)
                elif revert_restores:
                    fixture.document.write_bytes(original)
                return com_ok({"ok": True, "status": revert_status})
            fixture.document.write_bytes(original + b"\0")  # WPS saved the write
            return com_ok({"ok": True, "status": "ok", "backend": "fake-wps", **(write or payload)})

        with patch.object(open_documents, RUNNERS[component], side_effect=fake):
            ok, data, errors, _replayed = command(
                document_id=fixture.document_id,
                request_id=f"undo-{name}",
                workspace=fixture.workspace,
                **arguments,
            )
        self.assertFalse(ok)
        return fixture, original, data, errors, reverts

    def test_every_command_undoes_a_rejected_write_in_wps(self):
        for name, case in CASES.items():
            with self.subTest(command=name):
                fixture, original, data, errors, reverts = self.run_case(name)
                self.assertEqual(reverts, [case[5]])
                self.assertEqual(errors[0]["code"], "VALIDATION_FAILED")
                self.assertEqual(errors[-1]["code"], ROLLED_BACK, errors)
                self.assertNotIn(ROLLBACK_BLOCKED, [error["code"] for error in errors])
                rollback = data["rollback"]
                self.assertEqual(rollback["method"], "in_wps")
                self.assertTrue(rollback["rolled_back_in_wps"])
                self.assertTrue(rollback["content_matches_backup"])
                self.assertNotEqual(rollback["file_sha256_before_undo"], rollback["file_sha256_after_undo"])
                self.assertEqual(fixture.document.read_bytes(), original)

    def test_shape_and_chart_writes_are_undone_by_their_own_scripts(self):
        for name, case in OBJECT_CASES.items():
            with self.subTest(command=name):
                _fixture, _original, data, errors, _reverts = self.run_case(name)
                self.assertEqual(self.revert_actions, [case[6]])
                self.assertEqual(errors[-1]["code"], ROLLED_BACK, errors)
                self.assertTrue(data["rollback"]["rolled_back_in_wps"])

    def test_a_chart_that_is_no_longer_selected_is_reported_as_missing(self):
        for status in ("selection_changed", "not_chart_selection", "no_window"):
            with self.subTest(status=status):
                _fixture, _original, data, errors, _reverts = self.run_case(
                    "chart-title-delete", revert_status=status, revert_restores=False
                )
                self.assertEqual(data["rollback"]["status"], "revert_target_missing")
                self.assertEqual(data["rollback"]["detail"], status)
                self.assertEqual(errors[-1]["code"], ROLLBACK_IN_WPS_FAILED)
                self.assertIn("could not be found again", errors[-1]["message"])

    def test_a_post_save_identity_failure_is_undone_the_same_way(self):
        with patch.object(
            open_documents,
            "verify_post_com_source",
            return_value={"code": "DOCUMENT_CHANGED_AFTER_COM", "message": "changed"},
        ):
            _fixture, _original, data, errors, reverts = self.run_case(
                "writer-replace", write={"read_back_text": "APPROVED"}
            )
        self.assertEqual(len(reverts), 1)
        self.assertEqual([error["code"] for error in errors], ["DOCUMENT_CHANGED_AFTER_COM", ROLLED_BACK])
        self.assertTrue(data["rollback"]["rolled_back_in_wps"])

    def test_a_write_that_did_not_stick_counts_as_undone(self):
        _fixture, _original, data, errors, _reverts = self.run_case("presentation-delete", revert_status="unchanged")
        self.assertEqual(errors[-1]["code"], ROLLED_BACK)
        self.assertIn("already in place", errors[-1]["message"])
        self.assertTrue(data["rollback"]["rolled_back_in_wps"])

    def test_text_that_changed_since_the_write_is_left_alone_and_reported(self):
        for status, state in (
            ("revert_mismatch", "no longer hold what it wrote"),
            ("revert_target_missing", "could not be found again"),
            ("revert_read_back_failed", "it was not saved"),
        ):
            with self.subTest(status=status):
                fixture, original, data, errors, _reverts = self.run_case(
                    "writer-replace", revert_status=status, revert_restores=False
                )
                self.assertEqual(errors[-1]["code"], ROLLBACK_IN_WPS_FAILED)
                message = errors[-1]["message"]
                self.assertIn(state, message)
                self.assertIn("Nothing was closed", message)
                self.assertIn(f"restore-backup --document-id {fixture.document_id}", message)
                self.assertIn(data["rollback"]["backup_path"], message)
                self.assertFalse(data["rollback"]["rolled_back_in_wps"])
                self.assertNotIn(ROLLBACK_BLOCKED, [error["code"] for error in errors])
                # atomic_mutation did not try to copy the backup over the file WPS holds.
                self.assertEqual(fixture.document.read_bytes(), original + b"\0")

    def test_a_saved_file_that_still_differs_from_the_backup_is_a_failed_undo(self):
        def save_other_value(path):
            workbook = Workbook()
            workbook.active["A1"] = "something else"
            workbook.save(path)
            workbook.close()

        _fixture, _original, data, errors, _reverts = self.run_case(
            "spreadsheet-replace", revert_writes=save_other_value
        )
        self.assertEqual(errors[-1]["code"], ROLLBACK_IN_WPS_FAILED)
        self.assertEqual(data["rollback"]["status"], "saved_content_differs")
        self.assertFalse(data["rollback"]["content_matches_backup"])
        self.assertIn("does not match the pre-write backup", errors[-1]["message"])

    def test_a_failed_undo_script_leaves_the_state_unknown(self):
        component, command, arguments, selection, payload, _expected = CASES["writer-insert"]
        fixture = Fixture(self, component=component)

        def fake(action, params):
            if action == "selection-read":
                return com_ok(selection)
            if action == "selection-revert":
                return {"ok": False, "errors": [{"code": "COM_OPERATION_FAILED", "message": "RPC busy"}], "data": {}}
            return com_ok({"ok": True, "status": "ok", **payload})

        with patch.object(open_documents, RUNNERS[component], side_effect=fake):
            _ok, data, errors, _replayed = command(
                document_id=fixture.document_id, request_id="undo-crash", workspace=fixture.workspace, **arguments
            )
        self.assertEqual(errors[-1]["code"], ROLLBACK_IN_WPS_FAILED)
        self.assertIn("could not run in WPS (RPC busy)", errors[-1]["message"])
        self.assertEqual(data["rollback"]["status"], "script_failed")

    def test_selections_without_a_recorded_position_are_reported_not_guessed(self):
        # A cell write that recorded no cells, and a presentation write that could not read the shape and records
        # no shape text: nothing tells the undo where to work, so it does not try.
        for name, write in (
            ("spreadsheet-replace", {"cells_matched": False, "read_back_text": "OLD", "cell_count": 1}),
            ("presentation-replace", {"read_back_text": "garbled", "slide_index": 1, "shape_name": "Title 1"}),
        ):
            with self.subTest(command=name):
                _fixture, _original, data, errors, reverts = self.run_case(name, write=write, revert_restores=False)
                self.assertEqual(reverts, [])
                self.assertEqual(data["rollback"]["status"], "unsupported")
                self.assertEqual(errors[-1]["code"], ROLLBACK_IN_WPS_FAILED)
                self.assertIn("not available for this kind of selection", errors[-1]["message"])

    def test_successful_writes_never_run_the_undo(self):
        component, command, arguments, selection, payload, _expected = CASES["writer-replace"]
        fixture = Fixture(self, component=component)
        actions = []

        def fake(action, params):
            actions.append(action)
            if action == "selection-read":
                return com_ok(selection)
            return com_ok({"ok": True, "status": "ok", "read_back_text": params["text"]})

        with patch.object(open_documents, RUNNERS[component], side_effect=fake):
            ok, data, errors, _replayed = command(
                document_id=fixture.document_id, request_id="undo-none", workspace=fixture.workspace, **arguments
            )
        self.assertTrue(ok, errors)
        self.assertNotIn("selection-revert", actions)
        self.assertNotIn("rollback", data)


class RevertPositionTests(unittest.TestCase):
    def test_presentation_positions_count_utf16_units_like_wps(self):
        # An emoji is two UTF-16 units in WPS, so the selection after it starts two places later.
        payload = {"slide_index": 2, "shape_name": "TextBox 3", "text_start": 4, "shape_text_before": "😀 OLD!"}
        revert = open_documents._presentation_revert(payload, "NEW", "OLD")
        self.assertEqual(revert["frame_expected"], "😀 NEW!")
        self.assertEqual(revert["frame_restored"], "😀 OLD!")

    def test_presentation_revert_refuses_a_position_that_does_not_hold_the_original(self):
        payload = {"slide_index": 1, "shape_name": "Title 1", "text_start": 2, "shape_text_before": "Status: DRAFT"}
        self.assertIsNone(open_documents._presentation_revert(payload, "FINAL", "DRAFT"))

    def test_spreadsheet_revert_needs_every_cell_before_and_after(self):
        capture = cell_capture("NEW")
        self.assertIsNotNone(open_documents._spreadsheet_revert(capture))
        self.assertIsNone(open_documents._spreadsheet_revert({**capture, "after_formulas": []}))
        self.assertIsNone(open_documents._spreadsheet_revert({**capture, "sheet_name": None}))


class AtomicMutationLeavesInWpsUndoAloneTests(unittest.TestCase):
    def test_a_result_that_carries_an_in_wps_rollback_is_returned_as_is(self):
        backup = Path("never-read.docx")

        @atomic_mutation
        def command(document_id, request_id, workspace="."):
            result = {"backup": {"backup_path": str(backup)}, "rollback": {"method": "in_wps"}}
            return False, result, [{"code": ROLLBACK_IN_WPS_FAILED, "message": "x"}], False

        with patch("wps_ai_agent_cli.atomic_mutation._undo_write") as undo:
            ok, result, errors, _replayed = command(document_id="doc", request_id="req")
        undo.assert_not_called()
        self.assertFalse(ok)
        self.assertEqual([error["code"] for error in errors], [ROLLBACK_IN_WPS_FAILED])
        self.assertNotIn("open_in_wps", result["rollback"])


if __name__ == "__main__":
    unittest.main()
