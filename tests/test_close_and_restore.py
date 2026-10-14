from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from wps_ai_agent_cli import backups, open_documents
from wps_ai_agent_cli.atomic_mutation import _blocked_guidance
from wps_ai_agent_cli.backups import create_backup, list_backups, restore_backup
from wps_ai_agent_cli.mcp_adapter import build_cli_argv
from wps_ai_agent_cli.operations import get_operation
from wps_ai_agent_cli.sessions import register_document


class CloseAndRestoreTests(unittest.TestCase):
    """restore-backup with close_in_wps, against a stand-in for the WPS close script."""

    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.workspace = Path(tmp.name)
        self.path = self.workspace / "book.xlsx"
        workbook = Workbook()
        workbook.active["A1"] = "seed"
        workbook.save(self.path)
        ok, record, errors = register_document("spreadsheets", str(self.path), self.workspace)
        self.assertTrue(ok, errors)
        self.document_id = record["document_id"]
        ok, result, errors, _ = create_backup(self.document_id, "before-write", workspace=self.workspace)
        self.assertTrue(ok, errors)
        self.backup_name = Path(result["backup_path"]).name
        workbook = load_workbook(self.path)
        workbook.active["A1"] = "rejected"
        workbook.save(self.path)
        self.calls = []

    def fake_close(self, status=None, failure=None):
        def run(component, action, params):
            self.calls.append((component, action, params))
            if failure is not None:
                return {"ok": False, "errors": [failure], "data": {}}
            return {"ok": True, "errors": [], "data": {"status": status}}
        return patch.object(open_documents, "_run_component_attached_com", side_effect=run)

    def restore(self, request_id="restore-1", **kwargs):
        return restore_backup(self.document_id, request_id, backup_name=self.backup_name, workspace=self.workspace, **kwargs)

    def a1(self):
        return load_workbook(self.path, read_only=True).active["A1"].value

    def backup_count(self):
        ok, data, errors = list_backups(self.document_id, self.workspace)
        self.assertTrue(ok, errors)
        return len(data["backups"])

    def test_without_the_flag_wps_is_never_touched(self):
        with patch.object(open_documents, "_run_component_attached_com", side_effect=AssertionError("WPS touched")):
            ok, result, errors, _ = self.restore()
        self.assertTrue(ok, errors)
        self.assertNotIn("close_in_wps", result)
        self.assertEqual(self.a1(), "seed")

    def test_an_open_saved_document_is_closed_then_restored(self):
        with self.fake_close("closed"):
            ok, result, errors, _ = self.restore(close_in_wps=True)
        self.assertTrue(ok, errors)
        self.assertEqual(result["close_in_wps"], "closed")
        self.assertEqual(self.a1(), "seed")
        component, action, params = self.calls[0]
        self.assertEqual((component, action), ("spreadsheets", "close-unchanged"))
        self.assertEqual(params["path"], str(self.path.resolve()))
        self.assertTrue(params["do_close"])

    def test_a_document_that_is_not_open_is_simply_restored(self):
        for status in ("not_open", "no_instance"):
            with self.subTest(status=status), self.fake_close(status):
                ok, result, errors, _ = self.restore(request_id=f"restore-{status}", close_in_wps=True)
            self.assertTrue(ok, errors)
            self.assertEqual(result["close_in_wps"], status)

    def test_unsaved_changes_in_wps_refuse_and_nothing_is_restored(self):
        with self.fake_close("unsaved"):
            ok, _result, errors, _ = self.restore(close_in_wps=True)
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "DOCUMENT_HAS_UNSAVED_CHANGES")
        self.assertEqual(self.a1(), "rejected")
        self.assertIsNone(get_operation("restore-1", self.workspace))
        self.assertEqual(self.backup_count(), 1, "no pre-restore backup may be taken")

    def test_a_failed_close_restores_nothing(self):
        with self.fake_close(failure={"code": "COM_OPERATION_FAILED", "message": "boom"}):
            ok, _result, errors, _ = self.restore(close_in_wps=True)
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "COM_OPERATION_FAILED")
        self.assertEqual(self.a1(), "rejected")

    def test_dry_run_only_asks_whether_it_would_close(self):
        with self.fake_close("open"):
            ok, result, errors, _ = self.restore(close_in_wps=True, dry_run=True)
        self.assertTrue(ok, errors)
        self.assertEqual(result["close_in_wps"], "would_close")
        self.assertFalse(self.calls[0][2]["do_close"])
        self.assertEqual(self.a1(), "rejected")

    def test_the_document_is_closed_before_the_pre_restore_backup(self):
        order = []
        real_create_backup = backups.create_backup

        def recording_backup(*args, **kwargs):
            order.append("backup")
            return real_create_backup(*args, **kwargs)

        def run(component, action, params):
            order.append("close")
            return {"ok": True, "errors": [], "data": {"status": "closed"}}

        with patch.object(open_documents, "_run_component_attached_com", side_effect=run), \
                patch.object(backups, "create_backup", side_effect=recording_backup):
            ok, _result, errors, _ = self.restore(close_in_wps=True)
        self.assertTrue(ok, errors)
        self.assertEqual(order, ["close", "backup"])

    def test_a_repeated_request_is_replayed_without_closing_again(self):
        with self.fake_close("closed"):
            self.restore(close_in_wps=True)
            ok, _result, errors, replayed = self.restore(close_in_wps=True)
        self.assertTrue(ok, errors)
        self.assertTrue(replayed)
        self.assertEqual(len(self.calls), 1)

    def test_a_request_without_the_flag_still_replays(self):
        self.restore()
        ok, _result, errors, replayed = self.restore()
        self.assertTrue(ok, errors)
        self.assertTrue(replayed)


class CloseInWpsEntryPointTests(unittest.TestCase):
    def test_mcp_argument_becomes_the_cli_flag(self):
        ok, argv, _schema, errors = build_cli_argv(
            "wps_agent_restore_backup",
            {"document_id": "doc_1", "backup_name": "b.xlsx", "close_in_wps": True, "request_id": "r1"},
        )
        self.assertTrue(ok, errors)
        self.assertIn("--close-in-wps", argv)

    def test_blocked_rollback_guidance_names_the_flag(self):
        guidance = _blocked_guidance("doc_1", "C:/backups/b.xlsx", True)
        self.assertIn("--close-in-wps", guidance)
        self.assertIn("unsaved", guidance)


if __name__ == "__main__":
    unittest.main()
