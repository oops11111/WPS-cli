import io
import json
import os
import shutil
import tempfile
import unittest
from contextlib import chdir
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from wps_ai_agent_cli.backups import create_backup, guarded_com_mutation
from wps_ai_agent_cli.cli import run
from wps_ai_agent_cli.mcp_adapter import call_mcp_tool
from wps_ai_agent_cli.operations import get_operation, inspect_mutation_request
from wps_ai_agent_cli.sessions import file_identity, register_document
from wps_ai_agent_cli.spreadsheet_ops import _run_spreadsheet_rename_sheet_com, rename_spreadsheet_sheet
from wps_ai_agent_cli.writer_ops import _run_writer_replace_com, writer_replace


@unittest.skipUnless(os.environ.get("WPS_AGENT_RUN_INTEGRATION") == "1", "Requires explicit local WPS integration run")
class MutationRecoveryWpsDrillTests(unittest.TestCase):
    def _retain_evidence(self, component, path, backup, request_id, workspace):
        artifact_root = os.environ.get("WPS_AGENT_DRILL_ARTIFACT_DIR")
        if not artifact_root:
            return
        target = Path(artifact_root) / component
        target.mkdir(parents=True, exist_ok=True)
        current_copy = target / f"current{path.suffix}"
        backup_copy = target / f"backup{path.suffix}"
        shutil.copy2(path, current_copy)
        shutil.copy2(backup["backup_path"], backup_copy)
        report = {
            "component": component,
            "request_id": request_id,
            "status": inspect_mutation_request(request_id, workspace)["status"],
            "current_path": str(current_copy),
            "backup_path": str(backup_copy),
            "current_sha256": file_identity(current_copy)["source_sha256"],
            "backup_sha256": file_identity(backup_copy)["source_sha256"],
            "main_operation_recorded": get_operation(request_id, workspace) is not None,
        }
        (target / "manifest.json").write_text(json.dumps(report, indent=2), encoding="utf-8")

    def _assert_ambiguous(self, workspace, request_id, document_id, backup):
        self.assertIsNone(get_operation(request_id, workspace))
        self.assertEqual(file_identity(Path(backup["backup_path"]))["source_sha256"],
                         backup["source_identity"]["source_sha256"])
        report = inspect_mutation_request(request_id, workspace)
        self.assertEqual(report["status"], "ambiguous")
        self.assertEqual(report["backups"][0]["document_id"], document_id)
        self.assertTrue(any("Do not automatically retry" in item for item in report["recovery_guidance"]))
        with chdir(workspace):
            output = io.StringIO()
            self.assertEqual(run(["mutation-request-inspect", "--request", request_id], output_stream=output), 0)
            cli = json.loads(output.getvalue())
            ok, mcp, errors = call_mcp_tool("wps_agent_mutation_request_inspect", {"request": request_id})
        self.assertTrue(ok, errors)
        self.assertEqual(cli["data"]["mutation_request"], mcp["response"]["data"]["mutation_request"])
        self.assertEqual(cli["data"]["mutation_request"]["status"], "ambiguous")

    def test_writer_saved_without_main_record_is_not_reapplied(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "interrupted.docx"
            document = Document()
            document.add_paragraph("Before")
            document.save(path)
            ok, record, errors = register_document("writer", str(path), tmp)
            self.assertTrue(ok, errors)
            request_id = "drill-writer"
            ok, backup, errors, _ = create_backup(record["document_id"], f"{request_id}:backup", workspace=tmp)
            self.assertTrue(ok, errors)
            saved = guarded_com_mutation(path, backup, lambda: _run_writer_replace_com(str(path), "Before", "After"))
            self.assertTrue(saved["ok"], saved["errors"])
            self.assertIn("After", [paragraph.text for paragraph in Document(path).paragraphs])
            self._assert_ambiguous(tmp, request_id, record["document_id"], backup)
            saved_identity = file_identity(path)
            retry = writer_replace(record["document_id"], "Before", "After", request_id, workspace=tmp)
            self.assertFalse(retry[0])
            self.assertEqual(retry[2][0]["code"], "UNRECORDED_MUTATION_AMBIGUOUS")
            self.assertEqual(file_identity(path), saved_identity)
            self.assertIsNone(get_operation(request_id, tmp))
            self._retain_evidence("writer", path, backup, request_id, tmp)

    def test_spreadsheet_saved_without_main_record_is_not_reapplied(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "interrupted.xlsx"
            workbook = Workbook()
            workbook.active.title = "Before"
            workbook.save(path)
            workbook.close()
            ok, record, errors = register_document("spreadsheets", str(path), tmp)
            self.assertTrue(ok, errors)
            request_id = "drill-spreadsheet"
            ok, backup, errors, _ = create_backup(record["document_id"], f"{request_id}:backup", workspace=tmp)
            self.assertTrue(ok, errors)
            saved = guarded_com_mutation(path, backup, lambda: _run_spreadsheet_rename_sheet_com(str(path), "Before", "After"))
            self.assertTrue(saved["ok"], saved["errors"])
            self._assert_ambiguous(tmp, request_id, record["document_id"], backup)
            saved_identity = file_identity(path)
            retry = rename_spreadsheet_sheet(record["document_id"], "Before", "After", request_id, workspace=tmp)
            self.assertFalse(retry[0])
            self.assertEqual(retry[2][0]["code"], "UNRECORDED_MUTATION_AMBIGUOUS")
            self.assertEqual(file_identity(path), saved_identity)
            self.assertIsNone(get_operation(request_id, tmp))
            self._retain_evidence("spreadsheets", path, backup, request_id, tmp)


if __name__ == "__main__":
    unittest.main()
