import io
import json
import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from wps_ai_agent_cli import ooxml
from wps_ai_agent_cli.cli import run
from wps_ai_agent_cli.document_text import count_text_in_docx, docx_body_paragraphs
from wps_ai_agent_cli.ooxml import OoxmlTooLargeError, load_workbook_guarded, read_zip_part
from wps_ai_agent_cli.writer_structure import read_writer_structure

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def write_docx(path, body_xml, padding=0):
    document = f'<w:document xmlns:w="{W_NS}"><w:body>{body_xml}{" " * padding}</w:body></w:document>'
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
        archive.writestr("word/document.xml", document)


class PartLimitTests(unittest.TestCase):
    def test_oversized_part_is_rejected_before_decompression_into_memory(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bomb.docx"
            write_docx(path, "<w:p/>", padding=ooxml.MAX_PART_BYTES + 1024)
            self.assertLess(path.stat().st_size, 1024 * 1024)
            with self.assertRaises(OoxmlTooLargeError) as caught:
                docx_body_paragraphs(path)
            self.assertEqual(caught.exception.details["part"], "word/document.xml")
            with self.assertRaises(OoxmlTooLargeError):
                read_writer_structure(path)
            with self.assertRaises(OoxmlTooLargeError):
                count_text_in_docx(path, "x")

    def test_normal_document_still_parses(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "ok.docx"
            write_docx(path, "<w:p><w:r><w:t>hello</w:t></w:r></w:p>")
            self.assertEqual(docx_body_paragraphs(path), ["hello"])

    def test_explicit_part_limit_parameter(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("a.xml", "<a>" + "x" * 100 + "</a>")
        with zipfile.ZipFile(buffer) as archive:
            self.assertEqual(len(read_zip_part(archive, "a.xml", max_bytes=1000)), 107)
            with self.assertRaises(OoxmlTooLargeError):
                read_zip_part(archive, "a.xml", max_bytes=50)


class PackageLimitTests(unittest.TestCase):
    def test_total_uncompressed_limit(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("a.xml", "<a/>" * 100)
            archive.writestr("b.xml", "<b/>" * 100)
        with zipfile.ZipFile(buffer) as archive, patch.object(ooxml, "MAX_TOTAL_UNCOMPRESSED_BYTES", 500):
            with self.assertRaises(OoxmlTooLargeError) as caught:
                read_zip_part(archive, "a.xml")
            self.assertIn("uncompressed_bytes", caught.exception.details)

    def test_entry_count_limit(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            for index in range(5):
                archive.writestr(f"p{index}.xml", "<a/>")
        with zipfile.ZipFile(buffer) as archive, patch.object(ooxml, "MAX_ENTRIES", 3):
            with self.assertRaises(OoxmlTooLargeError):
                read_zip_part(archive, "p0.xml")

    def test_workbook_loader_applies_package_guard(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.xlsx"
            workbook = Workbook()
            workbook.save(path)
            workbook.close()
            loaded = load_workbook_guarded(path, read_only=True)
            loaded.close()
            with patch.object(ooxml, "MAX_TOTAL_UNCOMPRESSED_BYTES", 10):
                with self.assertRaises(OoxmlTooLargeError):
                    load_workbook_guarded(path, read_only=True)

    def test_non_zip_files_fall_through_to_the_library_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bad.xlsx"
            path.write_bytes(b"not a zip")
            with self.assertRaises(Exception) as caught:
                load_workbook_guarded(path)
            self.assertNotIsInstance(caught.exception, OoxmlTooLargeError)


class CliStructuredErrorTests(unittest.TestCase):
    def test_cli_returns_input_too_large_instead_of_traceback(self):
        with tempfile.TemporaryDirectory() as tmp:
            previous = os.getcwd()
            os.chdir(tmp)
            try:
                path = Path(tmp) / "bomb.docx"
                write_docx(path, "<w:p/>", padding=ooxml.MAX_PART_BYTES + 1024)
                stream = io.StringIO()
                run(["register-document", "--component", "writer", "--path", str(path)], output_stream=stream)
                document_id = json.loads(stream.getvalue())["data"]["document"]["document_id"]
                stream = io.StringIO()
                code = run(["validate-document", "--document-id", document_id, "--contains", "x", "--request-id", "big-1"], output_stream=stream)
                payload = json.loads(stream.getvalue())
            finally:
                os.chdir(previous)
        self.assertEqual(code, 0)
        self.assertFalse(payload["ok"])
        self.assertEqual(payload["errors"][0]["code"], "INPUT_TOO_LARGE")
        self.assertEqual(payload["request_id"], "big-1")
        self.assertEqual(payload["data"]["part"], "word/document.xml")


if __name__ == "__main__":
    unittest.main()
