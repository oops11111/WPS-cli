from __future__ import annotations

import io
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation

from wps_ai_agent_cli import cli
from wps_ai_agent_cli.html_office import convert_html_to_pptx, convert_html_to_xlsx

TABLES_HTML = """<!doctype html><html><body>
<h2>Sales</h2>
<table><tr><th>Region</th><th>Revenue</th></tr>
<tr><td>North</td><td>1200</td></tr>
<tr><td>South</td><td>007</td></tr>
<tr><td>Total</td><td>=SUM(B2:B3)</td></tr></table>
<table><tr><td>a</td><td>1.5</td></tr></table>
<img src="chart.png">
</body></html>"""

SLIDES_HTML = """<!doctype html><html><body>
<section><h1>Intro</h1><p>Hello team.</p><ul><li>First point</li><li>Second point</li></ul></section>
<section><h2>Plan</h2><p>Ship it.</p></section>
</body></html>"""

HEADINGS_ONLY_HTML = "<html><body><h1>One</h1><p>first</p><h2>Two</h2><p>second</p></body></html>"


class HtmlToXlsxTests(unittest.TestCase):
    def test_tables_become_named_sheets_with_typed_cells(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "page.html"
            source.write_text(TABLES_HTML, encoding="utf-8")
            target = Path(tmp) / "out.xlsx"

            ok, data, errors = convert_html_to_xlsx(source, target)

            self.assertTrue(ok, errors)
            # The second table has no heading of its own, so the nearest preceding heading ("Sales") names it.
            self.assertEqual([sheet["name"] for sheet in data["sheets"]], ["Sales", "Sales 2"])
            workbook = load_workbook(target)
            sales = workbook["Sales"]
            self.assertEqual(sales["A1"].value, "Region")
            self.assertEqual(sales["B2"].value, 1200)
            self.assertEqual(sales["B3"].value, "007")
            self.assertEqual(workbook["Sales 2"]["B1"].value, 1.5)

    def test_formula_like_text_is_written_as_text_and_reported(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "page.html"
            source.write_text(TABLES_HTML, encoding="utf-8")
            target = Path(tmp) / "out.xlsx"

            ok, data, _errors = convert_html_to_xlsx(source, target)

            cell = load_workbook(target)["Sales"]["B4"]
            self.assertEqual(cell.value, "=SUM(B2:B3)")
            self.assertEqual(cell.data_type, "s")
            self.assertTrue(any("written as text" in warning for warning in data["warnings"]))
            self.assertTrue(any("image" in warning for warning in data["warnings"]))

    def test_html_without_tables_is_refused(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "page.html"
            source.write_text("<html><body><p>just text</p></body></html>", encoding="utf-8")

            ok, _data, errors = convert_html_to_xlsx(source, Path(tmp) / "out.xlsx")

            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "NO_CONTENT")
            self.assertFalse((Path(tmp) / "out.xlsx").exists())


class HtmlToPptxTests(unittest.TestCase):
    def test_sections_become_slides_with_title_and_body(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "deck.html"
            source.write_text(SLIDES_HTML, encoding="utf-8")
            target = Path(tmp) / "out.pptx"

            ok, data, errors = convert_html_to_pptx(source, target)

            self.assertTrue(ok, errors)
            self.assertEqual(data["slides"], 2)
            slides = list(Presentation(target).slides)
            self.assertEqual(slides[0].shapes.title.text, "Intro")
            body = [paragraph.text for paragraph in slides[0].placeholders[1].text_frame.paragraphs]
            self.assertEqual(body, ["Hello team.", "First point", "Second point"])
            self.assertEqual(slides[1].shapes.title.text, "Plan")

    def test_headings_start_slides_when_there_are_no_sections(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "deck.html"
            source.write_text(HEADINGS_ONLY_HTML, encoding="utf-8")
            target = Path(tmp) / "out.pptx"

            ok, data, _errors = convert_html_to_pptx(source, target)

            self.assertTrue(ok)
            self.assertEqual(data["slides"], 2)
            titles = [slide.shapes.title.text for slide in Presentation(target).slides]
            self.assertEqual(titles, ["One", "Two"])

    def test_tables_are_reported_as_not_converted_to_slides(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "deck.html"
            source.write_text(TABLES_HTML.replace("<h2>Sales</h2>", "<h1>Sales</h1>"), encoding="utf-8")

            ok, data, _errors = convert_html_to_pptx(source, Path(tmp) / "out.pptx")

            self.assertTrue(ok)
            self.assertTrue(any("table(s) were not converted" in warning for warning in data["warnings"]))


class HtmlOfficeRefusalTests(unittest.TestCase):
    def test_existing_output_is_never_overwritten(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "deck.html"
            source.write_text(SLIDES_HTML, encoding="utf-8")
            target = Path(tmp) / "out.pptx"
            target.write_bytes(b"keep me")

            ok, _data, errors = convert_html_to_pptx(source, target)

            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "OUTPUT_ALREADY_EXISTS")
            self.assertEqual(target.read_bytes(), b"keep me")

    def test_output_extension_must_match_the_converter(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "deck.html"
            source.write_text(SLIDES_HTML, encoding="utf-8")

            ok, _data, errors = convert_html_to_xlsx(source, Path(tmp) / "out.pptx")

            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "INVALID_OUTPUT")


class HtmlEditableFormatCliTests(unittest.TestCase):
    def _run(self, argv):
        stream = io.StringIO()
        code = cli.run(argv, output_stream=stream)
        return code, json.loads(stream.getvalue())

    def test_format_xlsx_through_the_cli_reports_the_backend_and_checks(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "page.html"
            source.write_text(TABLES_HTML, encoding="utf-8")
            target = Path(tmp) / "out.xlsx"

            code, envelope = self._run(
                [
                    "html-editable",
                    "--input",
                    str(source),
                    "--output",
                    str(target),
                    "--format",
                    "xlsx",
                    "--request-id",
                    "req-office-xlsx",
                ]
            )

            self.assertEqual(code, 0)
            self.assertTrue(envelope["ok"], envelope["errors"])
            self.assertEqual(envelope["backend"], "openpyxl")
            self.assertEqual(envelope["validation"]["status"], "passed")
            self.assertTrue(target.exists())

    def test_format_pptx_through_the_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "deck.html"
            source.write_text(SLIDES_HTML, encoding="utf-8")
            target = Path(tmp) / "out.pptx"

            _code, envelope = self._run(
                [
                    "html-editable",
                    "--input",
                    str(source),
                    "--output",
                    str(target),
                    "--format",
                    "pptx",
                    "--request-id",
                    "req-office-pptx",
                ]
            )

            self.assertTrue(envelope["ok"], envelope["errors"])
            self.assertEqual(envelope["backend"], "python-pptx")
            self.assertEqual(envelope["data"]["slides"], 2)


if __name__ == "__main__":
    unittest.main()
