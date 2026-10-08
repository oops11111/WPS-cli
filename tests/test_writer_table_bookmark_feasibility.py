import os
import io
import json
import tempfile
import unittest
from contextlib import chdir
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches
from PIL import Image
from PIL import ImageChops
import pypdfium2

from wps_ai_agent_cli.backups import create_backup
from wps_ai_agent_cli.com_backend import run_conversion_smoke
from wps_ai_agent_cli.document_text import (
    docx_body_drawing_semantics,
    docx_body_link_field_semantics,
    docx_body_table_topology,
)
from wps_ai_agent_cli.sessions import file_identity, register_document
from wps_ai_agent_cli.writer_ops import writer_fill_bookmark
from wps_ai_agent_cli.writer_structure import read_supported_bookmark_text
from wps_ai_agent_cli.cli import run
from wps_ai_agent_cli.mcp_adapter import call_mcp_tool


def write_image_bookmark_document(
    path: Path, bookmark_around_drawing: bool = False, image_count: int = 1,
    anchored: bool = False, cropped: bool = False, image_width_inches: float = 0.25,
    anchor_spacing_inches: float = 0.25,
    anchor_relative_frames: tuple[str, str] = ("column", "paragraph"),
    anchor_offsets_inches: tuple[float, float] = (0.75, 0.0),
    anchor_wrap: str = "square",
    anchor_wrap_text: str | None = "bothSides",
    anchor_wrap_points: tuple[tuple[int, int], ...] | None = None,
    anchor_distances_inches: tuple[float, float, float, float] = (0.0, 0.0, 0.125, 0.125),
    layout_in_cell: str = "1",
    preamble_paragraphs: int = 0,
    postamble_paragraphs: int = 0,
    preamble_images: int = 0,
    flow_text_repetitions: int = 0,
) -> None:
    def image_stream(index: int) -> io.BytesIO:
        stream = io.BytesIO()
        color = ((0, 80, 255), (255, 30, 0), (0, 190, 70))[index % 3]
        Image.new("RGB", (96, 64), color).save(stream, format="PNG")
        stream.seek(0)
        return stream

    document = Document()
    for index in range(preamble_paragraphs):
        label = "PREFACE-START" if index == 0 else f"Preamble paragraph {index + 1}"
        paragraph = document.add_paragraph(f"{label}: pagination fixture text for WPS Writer.")
        if index < preamble_images:
            paragraph.add_run().add_picture(image_stream(index + 1), width=Inches(image_width_inches))
    table = document.add_table(rows=1, cols=1)
    paragraph = table.cell(0, 0).paragraphs[0]
    paragraph.add_run("Before ")
    if bookmark_around_drawing:
        image_run = paragraph.add_run()
        image_run.add_picture(image_stream(0), width=Inches(image_width_inches))
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), "1")
        start.set(qn("w:name"), "ImageMark")
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), "1")
        image_run._r.addprevious(start)
        image_run._r.addnext(end)
    else:
        old = paragraph.add_run("Old")
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), "1")
        start.set(qn("w:name"), "CellMark")
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), "1")
        old._r.addprevious(start)
        old._r.addnext(end)
        paragraph.add_run(" ")
        for index in range(image_count):
            paragraph.add_run().add_picture(image_stream(index), width=Inches(image_width_inches))
        paragraph.add_run(" After")
        if flow_text_repetitions:
            paragraph.add_run(" wrap-flow sample text" * flow_text_repetitions)
    document.save(path)
    if postamble_paragraphs:
        document = Document(path)
        for index in range(postamble_paragraphs):
            label = "AFTER-TABLE-MARKER" if index == 0 else f"Postamble paragraph {index + 1}"
            document.add_paragraph(f"{label}: pagination fixture text for WPS Writer.")
        document.save(path)
    if anchored or cropped:
        document = Document(path)
        inline_drawings = list(document.element.body.iter(qn("wp:inline")))
        if anchored:
            for index, inline in enumerate(inline_drawings):
                anchor = OxmlElement("wp:anchor")
                for name, value in zip(
                    ("distT", "distB", "distL", "distR"),
                    (str(int(Inches(distance).emu)) for distance in anchor_distances_inches),
                ):
                    anchor.set(name, value)
                for name, value in (("simplePos", "0"),
                                    ("relativeHeight", str(100000 + index * 100000)), ("behindDoc", "0"),
                                    ("locked", "0"), ("layoutInCell", layout_in_cell),
                                    ("allowOverlap", "1")):
                    anchor.set(name, value)
                simple = OxmlElement("wp:simplePos")
                simple.set("x", "0")
                simple.set("y", "0")
                anchor.append(simple)
                horizontal = OxmlElement("wp:positionH")
                horizontal.set("relativeFrom", anchor_relative_frames[0])
                horizontal_offset = OxmlElement("wp:posOffset")
                horizontal_offset.text = str(int(Inches(
                    anchor_offsets_inches[0] + index * anchor_spacing_inches
                ).emu))
                horizontal.append(horizontal_offset)
                anchor.append(horizontal)
                vertical = OxmlElement("wp:positionV")
                vertical.set("relativeFrom", anchor_relative_frames[1])
                vertical_offset = OxmlElement("wp:posOffset")
                vertical_offset.text = str(int(Inches(anchor_offsets_inches[1]).emu))
                vertical.append(vertical_offset)
                anchor.append(vertical)
                for child_name in ("wp:extent", "wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
                    child = inline.find(qn(child_name))
                    if child is not None:
                        anchor.append(child)
                wrap = OxmlElement(f"wp:wrap{anchor_wrap.title().replace('_', '')}")
                if anchor_wrap != "none" and anchor_wrap_text is not None:
                    wrap.set("wrapText", anchor_wrap_text)
                if anchor_wrap in {"tight", "through"}:
                    polygon = OxmlElement("wp:wrapPolygon")
                    points = anchor_wrap_points or ((10800, 0), (21600, 10800), (10800, 21600), (0, 10800))
                    start = OxmlElement("wp:start")
                    start.set("x", str(points[0][0]))
                    start.set("y", str(points[0][1]))
                    polygon.append(start)
                    for x, y in points[1:]:
                        point = OxmlElement("wp:lineTo")
                        point.set("x", str(x))
                        point.set("y", str(y))
                        polygon.append(point)
                    wrap.append(polygon)
                anchor.insert(4, wrap)
                inline.getparent().replace(inline, anchor)
        if cropped:
            blip_fill = next(document.element.body.iter(qn("pic:blipFill")))
            source_rect = OxmlElement("a:srcRect")
            source_rect.set("l", "10000")
            source_rect.set("r", "10000")
            blip = blip_fill.find(qn("a:blip"))
            blip_fill.insert(blip_fill.index(blip) + 1, source_rect)
        document.save(path)


def drawing_structure_differences(left, right, path="drawing"):
    if isinstance(left, dict) and isinstance(right, dict):
        differences = []
        for key in sorted(set(left) | set(right)):
            if key not in left or key not in right:
                differences.append((f"{path}.{key}", left.get(key), right.get(key)))
            else:
                differences.extend(drawing_structure_differences(left[key], right[key], f"{path}.{key}"))
        return differences
    if isinstance(left, list) and isinstance(right, list):
        if len(left) != len(right):
            return [(f"{path}.length", len(left), len(right))]
        return [difference for index, (before, after) in enumerate(zip(left, right))
                for difference in drawing_structure_differences(before, after, f"{path}[{index}]")]
    return [] if left == right else [(path, left, right)]


@unittest.skipUnless(os.environ.get("WPS_AGENT_RUN_INTEGRATION") == "1", "Requires explicit local WPS Writer integration run")
class WriterTableBookmarkFeasibilityTests(unittest.TestCase):
    def test_wps_table_bookmark_preserves_concave_tight_and_through_paths(self):
        points = ((0, 0), (21600, 0), (21600, 7200), (7200, 7200), (7200, 21600), (0, 21600))
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for mode in ("tight", "through"):
                with self.subTest(wrap=mode):
                    path = root / f"concave-{mode}.docx"
                    before_pdf = root / f"before-concave-{mode}.pdf"
                    after_pdf = root / f"after-concave-{mode}.pdf"
                    write_image_bookmark_document(
                        path, anchored=True, image_width_inches=1.0,
                        anchor_offsets_inches=(2.0, 0.0), anchor_wrap=mode,
                        anchor_wrap_points=points, flow_text_repetitions=30,
                    )
                    before_semantics = docx_body_drawing_semantics(path)
                    self.assertEqual(before_semantics["unresolved_drawings"], [])
                    self.assertTrue(run_conversion_smoke("writer", str(path), str(before_pdf), "pdf")["ok"])
                    before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
                    before_image = before_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                    before_text = before_pdf_doc[0].get_textpage().get_text_range()
                    before_pdf_doc.close()
                    ok, record, errors = register_document("writer", str(path), tmp)
                    self.assertTrue(ok, errors)
                    ok, result, errors, _ = writer_fill_bookmark(
                        record["document_id"], "CellMark", "Old", f"concave-{mode}", workspace=tmp,
                    )
                    self.assertTrue(ok, (errors, result))
                    self.assertEqual(docx_body_drawing_semantics(path), before_semantics)
                    self.assertTrue(run_conversion_smoke("writer", str(path), str(after_pdf), "pdf")["ok"])
                    after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
                    after_image = after_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                    after_text = after_pdf_doc[0].get_textpage().get_text_range()
                    after_pdf_doc.close()
                    self.assertEqual(before_text, after_text)
                    self.assertIsNone(ImageChops.difference(before_image, after_image).getbbox())

    def test_wps_table_bookmark_accepts_polygon_winding_and_closure_variants(self):
        outline = ((0, 0), (21600, 0), (21600, 7200), (7200, 7200), (7200, 21600), (0, 21600))
        variants = (
            ("ccw-implicit", outline),
            ("cw-implicit", tuple(reversed(outline))),
            ("ccw-explicit", (*outline, outline[0])),
            ("cw-explicit", (outline[0], *reversed(outline[1:]), outline[0])),
        )
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for mode in ("tight", "through"):
                for label, points in variants:
                    with self.subTest(wrap=mode, polygon=label):
                        path = root / f"{mode}-{label}.docx"
                        before_pdf = root / f"before-{mode}-{label}.pdf"
                        after_pdf = root / f"after-{mode}-{label}.pdf"
                        write_image_bookmark_document(
                            path, anchored=True, image_width_inches=1.0,
                            anchor_offsets_inches=(2.0, 0.0), anchor_wrap=mode,
                            anchor_wrap_points=points, flow_text_repetitions=30,
                        )
                        before_semantics = docx_body_drawing_semantics(path)
                        self.assertEqual(before_semantics["unresolved_drawings"], [])
                        self.assertTrue(run_conversion_smoke("writer", str(path), str(before_pdf), "pdf")["ok"])
                        before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
                        before_image = before_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                        before_text = before_pdf_doc[0].get_textpage().get_text_range()
                        before_pdf_doc.close()
                        ok, record, errors = register_document("writer", str(path), tmp)
                        self.assertTrue(ok, errors)
                        ok, result, errors, _ = writer_fill_bookmark(
                            record["document_id"], "CellMark", "Old", f"polygon-{mode}-{label}", workspace=tmp,
                        )
                        self.assertTrue(ok, (errors, result))
                        after_semantics = docx_body_drawing_semantics(path)
                        self.assertEqual(after_semantics["unresolved_drawings"], [])
                        self.assertTrue(run_conversion_smoke("writer", str(path), str(after_pdf), "pdf")["ok"])
                        after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
                        after_image = after_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                        after_text = after_pdf_doc[0].get_textpage().get_text_range()
                        after_pdf_doc.close()
                        self.assertEqual(after_text, before_text)
                        self.assertIsNone(ImageChops.difference(before_image, after_image).getbbox())

    def test_wps_table_bookmark_keeps_cyclic_start_vertex_semantics(self):
        outline = ((0, 0), (21600, 0), (21600, 7200), (7200, 7200), (7200, 21600), (0, 21600))
        starts = (0, 1, 3, 5)
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for mode in ("tight", "through"):
                baseline = None
                for start in starts:
                    points = (*outline[start:], *outline[:start])
                    with self.subTest(wrap=mode, start_vertex=start):
                        path = root / f"{mode}-start-{start}.docx"
                        before_pdf = root / f"before-{mode}-start-{start}.pdf"
                        after_pdf = root / f"after-{mode}-start-{start}.pdf"
                        write_image_bookmark_document(
                            path, anchored=True, image_width_inches=1.0,
                            anchor_offsets_inches=(2.0, 0.0), anchor_wrap=mode,
                            anchor_wrap_points=points, flow_text_repetitions=30,
                        )
                        before_semantics = docx_body_drawing_semantics(path)
                        self.assertEqual(before_semantics["unresolved_drawings"], [])
                        self.assertTrue(run_conversion_smoke("writer", str(path), str(before_pdf), "pdf")["ok"])
                        before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
                        before_image = before_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                        before_text = before_pdf_doc[0].get_textpage().get_text_range()
                        before_pdf_doc.close()
                        if baseline is None:
                            baseline = before_image
                        else:
                            self.assertIsNotNone(ImageChops.difference(baseline, before_image).getbbox())

                        ok, record, errors = register_document("writer", str(path), tmp)
                        self.assertTrue(ok, errors)
                        ok, result, errors, _ = writer_fill_bookmark(
                            record["document_id"], "CellMark", "Old", f"polygon-start-{mode}-{start}", workspace=tmp,
                        )
                        self.assertTrue(ok, (errors, result))
                        self.assertEqual(docx_body_drawing_semantics(path), before_semantics)
                        self.assertTrue(run_conversion_smoke("writer", str(path), str(after_pdf), "pdf")["ok"])
                        after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
                        after_image = after_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                        after_text = after_pdf_doc[0].get_textpage().get_text_range()
                        after_pdf_doc.close()
                        self.assertEqual(after_text, before_text)
                        self.assertIsNone(ImageChops.difference(before_image, after_image).getbbox())

    def test_wps_table_bookmark_audits_collinear_wrap_vertices(self):
        outline = ((0, 0), (21600, 0), (21600, 7200), (7200, 7200), (7200, 21600), (0, 21600))
        with_midpoints = (
            (0, 0), (10800, 0), (21600, 0), (21600, 3600), (21600, 7200),
            (14400, 7200), (7200, 7200), (7200, 14400), (7200, 21600),
            (3600, 21600), (0, 21600), (0, 10800),
        )
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for mode in ("tight", "through"):
                rendered = []
                for label, points in (("original", outline), ("collinear", with_midpoints)):
                    with self.subTest(wrap=mode, polygon=label):
                        path = root / f"{mode}-{label}.docx"
                        before_pdf = root / f"before-{mode}-{label}.pdf"
                        after_pdf = root / f"after-{mode}-{label}.pdf"
                        write_image_bookmark_document(
                            path, anchored=True, image_width_inches=1.0,
                            anchor_offsets_inches=(2.0, 0.0), anchor_wrap=mode,
                            anchor_wrap_points=points, flow_text_repetitions=30,
                        )
                        before_semantics = docx_body_drawing_semantics(path)
                        self.assertEqual(before_semantics["unresolved_drawings"], [])
                        self.assertTrue(run_conversion_smoke("writer", str(path), str(before_pdf), "pdf")["ok"])
                        before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
                        before_image = before_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                        before_text = before_pdf_doc[0].get_textpage().get_text_range()
                        before_pdf_doc.close()
                        rendered.append(before_image)

                        ok, record, errors = register_document("writer", str(path), tmp)
                        self.assertTrue(ok, errors)
                        ok, result, errors, _ = writer_fill_bookmark(
                            record["document_id"], "CellMark", "Old", f"polygon-collinear-{mode}-{label}", workspace=tmp,
                        )
                        self.assertTrue(ok, (errors, result))
                        self.assertEqual(docx_body_drawing_semantics(path), before_semantics)
                        self.assertTrue(run_conversion_smoke("writer", str(path), str(after_pdf), "pdf")["ok"])
                        after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
                        after_image = after_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                        after_text = after_pdf_doc[0].get_textpage().get_text_range()
                        after_pdf_doc.close()
                        self.assertEqual(after_text, before_text)
                        self.assertIsNone(ImageChops.difference(before_image, after_image).getbbox())
                self.assertIsNotNone(ImageChops.difference(rendered[0], rendered[1]).getbbox())

    def test_wps_table_bookmark_round_trips_near_degenerate_and_boundary_polygons(self):
        variants = (
            ("near_duplicate", ((0, 0), (1, 0), (21600, 0), (21600, 7200),
                                (7200, 7200), (7200, 21600), (0, 21600))),
            ("near_bounds", ((1, 1), (21599, 1), (21599, 7199), (7201, 7199),
                             (7201, 21599), (1, 21599))),
        )
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for mode in ("tight", "through"):
                for label, points in variants:
                    with self.subTest(wrap=mode, polygon=label):
                        path = root / f"{mode}-{label}.docx"
                        before_pdf = root / f"before-{mode}-{label}.pdf"
                        after_pdf = root / f"after-{mode}-{label}.pdf"
                        write_image_bookmark_document(
                            path, anchored=True, image_width_inches=1.0,
                            anchor_offsets_inches=(2.0, 0.0), anchor_wrap=mode,
                            anchor_wrap_points=points, flow_text_repetitions=30,
                        )
                        before_semantics = docx_body_drawing_semantics(path)
                        self.assertEqual(before_semantics["unresolved_drawings"], [])
                        self.assertTrue(run_conversion_smoke("writer", str(path), str(before_pdf), "pdf")["ok"])
                        before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
                        before_image = before_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                        before_text = before_pdf_doc[0].get_textpage().get_text_range()
                        before_pdf_doc.close()

                        ok, record, errors = register_document("writer", str(path), tmp)
                        self.assertTrue(ok, errors)
                        ok, result, errors, _ = writer_fill_bookmark(
                            record["document_id"], "CellMark", "Old", f"polygon-{label}-{mode}", workspace=tmp,
                        )
                        self.assertTrue(ok, (errors, result))
                        self.assertEqual(docx_body_drawing_semantics(path), before_semantics)
                        self.assertTrue(run_conversion_smoke("writer", str(path), str(after_pdf), "pdf")["ok"])
                        after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
                        after_image = after_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                        after_text = after_pdf_doc[0].get_textpage().get_text_range()
                        after_pdf_doc.close()
                        self.assertEqual(after_text, before_text)
                        self.assertIsNone(ImageChops.difference(before_image, after_image).getbbox())

    def test_wps_table_bookmark_audits_near_collinear_vertices(self):
        base = ((0, 0), (21600, 0), (21600, 7200), (7200, 7200), (7200, 21600), (0, 21600))
        near_collinear = ((0, 0), (10800, 120), (21600, 0), *base[2:])
        visible_offset = ((0, 0), (10800, 600), (21600, 0), *base[2:])
        variants = (("straight", base), ("offset-120", near_collinear), ("offset-600", visible_offset))
        raster_scale = 4.0
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for mode in ("tight", "through"):
                rendered = []
                for label, points in variants:
                    with self.subTest(wrap=mode, polygon=label):
                        path = root / f"{mode}-{label}.docx"
                        before_pdf = root / f"before-{mode}-{label}.pdf"
                        after_pdf = root / f"after-{mode}-{label}.pdf"
                        write_image_bookmark_document(
                            path, anchored=True, image_width_inches=1.0,
                            anchor_offsets_inches=(2.0, 0.0), anchor_wrap=mode,
                            anchor_wrap_points=points, flow_text_repetitions=30,
                        )
                        before_semantics = docx_body_drawing_semantics(path)
                        self.assertEqual(before_semantics["unresolved_drawings"], [])
                        self.assertTrue(run_conversion_smoke("writer", str(path), str(before_pdf), "pdf")["ok"])
                        before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
                        before_image = before_pdf_doc[0].render(scale=raster_scale).to_pil().convert("RGB")
                        before_text = before_pdf_doc[0].get_textpage().get_text_range()
                        before_pdf_doc.close()
                        rendered.append(before_image)

                        ok, record, errors = register_document("writer", str(path), tmp)
                        self.assertTrue(ok, errors)
                        ok, result, errors, _ = writer_fill_bookmark(
                            record["document_id"], "CellMark", "Old", f"polygon-near-collinear-{mode}-{label}",
                            workspace=tmp,
                        )
                        self.assertTrue(ok, (errors, result))
                        self.assertEqual(docx_body_drawing_semantics(path), before_semantics)
                        self.assertTrue(run_conversion_smoke("writer", str(path), str(after_pdf), "pdf")["ok"])
                        after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
                        after_image = after_pdf_doc[0].render(scale=raster_scale).to_pil().convert("RGB")
                        after_text = after_pdf_doc[0].get_textpage().get_text_range()
                        after_pdf_doc.close()
                        self.assertEqual(after_text, before_text)
                        self.assertIsNone(ImageChops.difference(before_image, after_image).getbbox())
                self.assertIsNone(ImageChops.difference(rendered[0], rendered[1]).getbbox())
                self.assertIsNone(ImageChops.difference(rendered[0], rendered[2]).getbbox())

    def test_wps_table_bookmark_normalizes_integer_wrap_lexical_forms(self):
        import xml.etree.ElementTree as ET
        from zipfile import ZipFile

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            variants = (
                ("+10800", "0"),
                (" 010800 ", "0"),
                ("10800", "-0"),
                ("10800", "+00"),
            )
            for index, (x_lexical, y_lexical) in enumerate(variants):
                with self.subTest(x=x_lexical, y=y_lexical):
                    path = root / f"coordinate-lexical-{index}.docx"
                    before_pdf = root / f"before-lexical-{index}.pdf"
                    after_pdf = root / f"after-lexical-{index}.pdf"
                    write_image_bookmark_document(path, anchored=True, anchor_wrap="tight")
                    with ZipFile(path) as source:
                        parts = {name: source.read(name) for name in source.namelist()}
                    document_xml = ET.fromstring(parts["word/document.xml"])
                    start = next(document_xml.iter(qn("wp:start")))
                    start.set("x", x_lexical)
                    start.set("y", y_lexical)
                    anchor = next(document_xml.iter(qn("wp:anchor")))
                    anchor.set("distL", " +0114300 ")
                    anchor.set("distR", "\t114300\n")
                    parts["word/document.xml"] = ET.tostring(
                        document_xml, encoding="utf-8", xml_declaration=True,
                    )
                    with ZipFile(path, "w") as target:
                        for name, content in parts.items():
                            target.writestr(name, content)

                    before_semantics = docx_body_drawing_semantics(path)
                    self.assertEqual(before_semantics["unresolved_drawings"], [])
                    self.assertTrue(run_conversion_smoke("writer", str(path), str(before_pdf), "pdf")["ok"])
                    before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
                    before_image = before_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                    before_text = before_pdf_doc[0].get_textpage().get_text_range()
                    before_pdf_doc.close()

                    ok, record, errors = register_document("writer", str(path), tmp)
                    self.assertTrue(ok, errors)
                    ok, result, errors, _ = writer_fill_bookmark(
                        record["document_id"], "CellMark", "Old", f"coordinate-lexical-{index}", workspace=tmp,
                    )
                    self.assertTrue(ok, (errors, result))
                    self.assertEqual(docx_body_drawing_semantics(path), before_semantics)
                    self.assertTrue(run_conversion_smoke("writer", str(path), str(after_pdf), "pdf")["ok"])
                    after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
                    after_image = after_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                    after_text = after_pdf_doc[0].get_textpage().get_text_range()
                    after_pdf_doc.close()
                    self.assertEqual(after_text, before_text)
                    self.assertIsNone(ImageChops.difference(before_image, after_image).getbbox())

    def test_wps_table_bookmark_retains_maximum_twip_aligned_wrap_distance(self):
        import xml.etree.ElementTree as ET
        from zipfile import ZipFile

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / "maximum-distance.docx"
            write_image_bookmark_document(path, anchored=True, anchor_wrap="square")
            with ZipFile(path) as source:
                parts = {name: source.read(name) for name in source.namelist()}
            document_xml = ET.fromstring(parts["word/document.xml"])
            next(document_xml.iter(qn("wp:anchor"))).set("distL", "2147483640")
            parts["word/document.xml"] = ET.tostring(document_xml, encoding="utf-8", xml_declaration=True)
            with ZipFile(path, "w") as target:
                for name, content in parts.items():
                    target.writestr(name, content)

            before_semantics = docx_body_drawing_semantics(path)
            self.assertEqual(before_semantics["unresolved_drawings"], [])
            ok, record, errors = register_document("writer", str(path), tmp)
            self.assertTrue(ok, errors)
            ok, result, errors, _ = writer_fill_bookmark(
                record["document_id"], "CellMark", "Old", "maximum-distance", workspace=tmp,
            )
            after_semantics = docx_body_drawing_semantics(path)
            self.assertTrue(ok, (errors, result, drawing_structure_differences(
                before_semantics, after_semantics,
            )))

    def test_wps_table_bookmark_normalizes_omitted_polygon_wrap_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for mode in ("tight", "through"):
                with self.subTest(wrap=mode):
                    omitted = root / f"omitted-{mode}-wrap-text.docx"
                    explicit = root / f"explicit-{mode}-wrap-text.docx"
                    before_pdf = root / f"before-{mode}-wrap-text.pdf"
                    after_pdf = root / f"after-{mode}-wrap-text.pdf"
                    write_image_bookmark_document(
                        omitted, anchored=True, anchor_wrap=mode, anchor_wrap_text=None,
                        flow_text_repetitions=20,
                    )
                    write_image_bookmark_document(
                        explicit, anchored=True, anchor_wrap=mode, anchor_wrap_text="bothSides",
                        flow_text_repetitions=20,
                    )
                    before_semantics = docx_body_drawing_semantics(omitted)
                    self.assertEqual(before_semantics, docx_body_drawing_semantics(explicit))
                    self.assertTrue(run_conversion_smoke(
                        "writer", str(omitted), str(before_pdf), "pdf",
                    )["ok"])
                    before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
                    before_image = before_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                    before_text = before_pdf_doc[0].get_textpage().get_text_range()
                    before_pdf_doc.close()

                    ok, record, errors = register_document("writer", str(omitted), tmp)
                    self.assertTrue(ok, errors)
                    ok, result, errors, _ = writer_fill_bookmark(
                        record["document_id"], "CellMark", "Old", f"polygon-wrap-{mode}",
                        workspace=tmp,
                    )
                    after_semantics = docx_body_drawing_semantics(omitted)
                    self.assertTrue(ok, (errors, result, drawing_structure_differences(
                        before_semantics, after_semantics,
                    )))
                    self.assertEqual(after_semantics, before_semantics)
                    self.assertTrue(run_conversion_smoke(
                        "writer", str(omitted), str(after_pdf), "pdf",
                    )["ok"])
                    after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
                    after_image = after_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                    after_text = after_pdf_doc[0].get_textpage().get_text_range()
                    after_pdf_doc.close()
                    self.assertEqual(after_text, before_text)
                    self.assertIsNone(ImageChops.difference(before_image, after_image).getbbox())

    def test_wps_table_bookmark_reads_wrap_none_extra_wrap_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for wrap_text in ("left", "right", "largest"):
                with self.subTest(wrap_text=wrap_text):
                    path = root / f"wrap-none-{wrap_text}.docx"
                    before_pdf = root / f"before-wrap-none-{wrap_text}.pdf"
                    after_pdf = root / f"after-wrap-none-{wrap_text}.pdf"
                    write_image_bookmark_document(
                        path, anchored=True, anchor_wrap="none", anchor_wrap_text=wrap_text,
                    )
                    before_semantics = docx_body_drawing_semantics(path)
                    self.assertEqual(before_semantics["unresolved_drawings"], [])
                    self.assertTrue(run_conversion_smoke(
                        "writer", str(path), str(before_pdf), "pdf",
                    )["ok"])
                    before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
                    before_image = before_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                    before_text = before_pdf_doc[0].get_textpage().get_text_range()
                    before_pdf_doc.close()

                    ok, record, errors = register_document("writer", str(path), tmp)
                    self.assertTrue(ok, errors)
                    ok, result, errors, _ = writer_fill_bookmark(
                        record["document_id"], "CellMark", "Old", f"wrap-none-{wrap_text}",
                        workspace=tmp,
                    )
                    after_semantics = docx_body_drawing_semantics(path)
                    self.assertTrue(ok, (errors, result, drawing_structure_differences(
                        before_semantics, after_semantics,
                    )))
                    self.assertEqual(after_semantics, before_semantics)
                    self.assertTrue(run_conversion_smoke(
                        "writer", str(path), str(after_pdf), "pdf",
                    )["ok"])
                    after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
                    after_image = after_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                    after_text = after_pdf_doc[0].get_textpage().get_text_range()
                    after_pdf_doc.close()
                    self.assertEqual(after_text, before_text)
                    self.assertIsNone(ImageChops.difference(before_image, after_image).getbbox())

    def test_wps_table_bookmark_normalizes_omitted_top_and_bottom_wrap_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            omitted = root / "omitted-top-bottom-wrap-text.docx"
            explicit = root / "explicit-top-bottom-wrap-text.docx"
            before_pdf = root / "before-top-bottom-wrap-text.pdf"
            after_pdf = root / "after-top-bottom-wrap-text.pdf"
            write_image_bookmark_document(
                omitted, anchored=True, anchor_wrap="top_and_bottom", anchor_wrap_text=None,
                flow_text_repetitions=20,
            )
            write_image_bookmark_document(
                explicit, anchored=True, anchor_wrap="top_and_bottom", anchor_wrap_text="bothSides",
                flow_text_repetitions=20,
            )
            before_semantics = docx_body_drawing_semantics(omitted)
            self.assertEqual(before_semantics, docx_body_drawing_semantics(explicit))
            self.assertTrue(run_conversion_smoke(
                "writer", str(omitted), str(before_pdf), "pdf",
            )["ok"])
            before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
            before_image = before_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
            before_text = before_pdf_doc[0].get_textpage().get_text_range()
            before_pdf_doc.close()

            ok, record, errors = register_document("writer", str(omitted), tmp)
            self.assertTrue(ok, errors)
            ok, result, errors, _ = writer_fill_bookmark(
                record["document_id"], "CellMark", "Old", "top-bottom-wrap-default",
                workspace=tmp,
            )
            after_semantics = docx_body_drawing_semantics(omitted)
            self.assertTrue(ok, (errors, result, drawing_structure_differences(
                before_semantics, after_semantics,
            )))
            self.assertEqual(after_semantics, before_semantics)
            self.assertTrue(run_conversion_smoke(
                "writer", str(omitted), str(after_pdf), "pdf",
            )["ok"])
            after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
            after_image = after_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
            after_text = after_pdf_doc[0].get_textpage().get_text_range()
            after_pdf_doc.close()
            self.assertEqual(after_text, before_text)
            self.assertIsNone(ImageChops.difference(before_image, after_image).getbbox())

    def test_wps_table_bookmark_normalizes_omitted_square_wrap_text(self):
        import xml.etree.ElementTree as ET
        from zipfile import ZipFile

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            omitted = root / "omitted-wrap-text.docx"
            explicit = root / "explicit-wrap-text.docx"
            before_pdf = root / "before-wrap-text.pdf"
            after_pdf = root / "after-wrap-text.pdf"
            write_image_bookmark_document(
                omitted, anchored=True, anchor_wrap="square", anchor_wrap_text=None,
                flow_text_repetitions=20,
            )
            write_image_bookmark_document(
                explicit, anchored=True, anchor_wrap="square", anchor_wrap_text="bothSides",
                flow_text_repetitions=20,
            )
            before_semantics = docx_body_drawing_semantics(omitted)
            self.assertEqual(before_semantics, docx_body_drawing_semantics(explicit))
            self.assertTrue(run_conversion_smoke("writer", str(omitted), str(before_pdf), "pdf")["ok"])
            before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
            before_image = before_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
            before_text = before_pdf_doc[0].get_textpage().get_text_range()
            before_pdf_doc.close()

            ok, record, errors = register_document("writer", str(omitted), tmp)
            self.assertTrue(ok, errors)
            ok, result, errors, _ = writer_fill_bookmark(
                record["document_id"], "CellMark", "Old", "wrap-text-default", workspace=tmp,
            )
            after_semantics = docx_body_drawing_semantics(omitted)
            self.assertTrue(ok, (errors, result, drawing_structure_differences(
                before_semantics, after_semantics,
            )))
            self.assertEqual(after_semantics, before_semantics)
            self.assertTrue(run_conversion_smoke("writer", str(omitted), str(after_pdf), "pdf")["ok"])
            after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
            after_image = after_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
            after_text = after_pdf_doc[0].get_textpage().get_text_range()
            after_pdf_doc.close()
            self.assertEqual(after_text, before_text)
            self.assertIsNone(ImageChops.difference(before_image, after_image).getbbox())

    def test_wps_table_bookmark_keeps_mixed_anchor_wrap_text_independent(self):
        import xml.etree.ElementTree as ET
        from zipfile import ZipFile

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / "mixed-anchor-wrap-text.docx"
            before_pdf = root / "before-mixed-anchor-wrap-text.pdf"
            after_pdf = root / "after-mixed-anchor-wrap-text.pdf"
            write_image_bookmark_document(
                path, anchored=True, image_count=2, anchor_wrap="tight",
                anchor_wrap_text="largest", anchor_spacing_inches=0.6, flow_text_repetitions=20,
            )
            with ZipFile(path) as source:
                parts = {name: source.read(name) for name in source.namelist()}
            document_xml = ET.fromstring(parts["word/document.xml"])
            anchors = list(document_xml.iter(qn("wp:anchor")))
            self.assertEqual(len(anchors), 2)
            first_wrap = next(child for child in anchors[0] if child.tag == qn("wp:wrapTight"))
            first_wrap.tag = qn("wp:wrapSquare")
            first_wrap.clear()
            first_wrap.set("wrapText", "left")
            parts["word/document.xml"] = ET.tostring(
                document_xml, encoding="utf-8", xml_declaration=True,
            )
            with ZipFile(path, "w") as target:
                for name, content in parts.items():
                    target.writestr(name, content)

            before_semantics = docx_body_drawing_semantics(path)
            self.assertEqual(before_semantics["unresolved_drawings"], [])
            self.assertEqual([
                drawing["structure"]["children"][4]["attributes"]["wrapText"]
                for drawing in before_semantics["drawings"]
            ], ["left", "largest"])
            self.assertTrue(run_conversion_smoke(
                "writer", str(path), str(before_pdf), "pdf",
            )["ok"])
            before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
            before_image = before_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
            before_text = before_pdf_doc[0].get_textpage().get_text_range()
            before_pdf_doc.close()

            ok, record, errors = register_document("writer", str(path), tmp)
            self.assertTrue(ok, errors)
            ok, result, errors, _ = writer_fill_bookmark(
                record["document_id"], "CellMark", "Old", "mixed-anchor-wrap-text",
                workspace=tmp,
            )
            after_semantics = docx_body_drawing_semantics(path)
            self.assertTrue(ok, (errors, result, drawing_structure_differences(
                before_semantics, after_semantics,
            )))
            self.assertEqual(after_semantics, before_semantics)
            self.assertTrue(run_conversion_smoke(
                "writer", str(path), str(after_pdf), "pdf",
            )["ok"])
            after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
            after_image = after_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
            after_text = after_pdf_doc[0].get_textpage().get_text_range()
            after_pdf_doc.close()
            self.assertEqual(after_text, before_text)
            self.assertIsNone(ImageChops.difference(before_image, after_image).getbbox())

    def test_wps_table_bookmark_keeps_mixed_anchor_distances_independent(self):
        import xml.etree.ElementTree as ET
        from zipfile import ZipFile

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / "mixed-anchor-distances.docx"
            before_pdf = root / "before-mixed-anchor-distances.pdf"
            after_pdf = root / "after-mixed-anchor-distances.pdf"
            write_image_bookmark_document(
                path, anchored=True, image_count=2, anchor_wrap="tight",
                anchor_wrap_text=None, anchor_spacing_inches=0.6, flow_text_repetitions=20,
            )
            with ZipFile(path) as source:
                parts = {name: source.read(name) for name in source.namelist()}
            document_xml = ET.fromstring(parts["word/document.xml"])
            anchors = list(document_xml.iter(qn("wp:anchor")))
            self.assertEqual(len(anchors), 2)
            first_wrap = next(child for child in anchors[0] if child.tag == qn("wp:wrapTight"))
            first_wrap.tag = qn("wp:wrapSquare")
            first_wrap.clear()
            for anchor in anchors:
                for name in ("distT", "distB", "distL", "distR"):
                    anchor.attrib.pop(name, None)
            anchors[0].set("distT", "635")
            anchors[1].set("distL", "1270")
            parts["word/document.xml"] = ET.tostring(
                document_xml, encoding="utf-8", xml_declaration=True,
            )
            with ZipFile(path, "w") as target:
                for name, content in parts.items():
                    target.writestr(name, content)

            before_semantics = docx_body_drawing_semantics(path)
            self.assertEqual(before_semantics["unresolved_drawings"], [])
            self.assertEqual([
                drawing["wrap_distances"] for drawing in before_semantics["drawings"]
            ], [
                {"distT": "635", "distB": "0", "distL": "114300", "distR": "114300"},
                {"distT": "0", "distB": "0", "distL": "1270", "distR": "114300"},
            ])
            self.assertTrue(run_conversion_smoke(
                "writer", str(path), str(before_pdf), "pdf",
            )["ok"])
            before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
            before_image = before_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
            before_text = before_pdf_doc[0].get_textpage().get_text_range()
            before_pdf_doc.close()

            ok, record, errors = register_document("writer", str(path), tmp)
            self.assertTrue(ok, errors)
            ok, result, errors, _ = writer_fill_bookmark(
                record["document_id"], "CellMark", "Old", "mixed-anchor-distances",
                workspace=tmp,
            )
            after_semantics = docx_body_drawing_semantics(path)
            self.assertTrue(ok, (errors, result, drawing_structure_differences(
                before_semantics, after_semantics,
            )))
            self.assertEqual(after_semantics, before_semantics)
            self.assertTrue(run_conversion_smoke(
                "writer", str(path), str(after_pdf), "pdf",
            )["ok"])
            after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
            after_image = after_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
            after_text = after_pdf_doc[0].get_textpage().get_text_range()
            after_pdf_doc.close()
            self.assertEqual(after_text, before_text)
            self.assertIsNone(ImageChops.difference(before_image, after_image).getbbox())

    def test_wps_table_bookmark_normalizes_defaults_per_mixed_anchor(self):
        import xml.etree.ElementTree as ET
        from zipfile import ZipFile

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / "mixed-anchor-defaults.docx"
            before_pdf = root / "before-mixed-anchor-defaults.pdf"
            after_pdf = root / "after-mixed-anchor-defaults.pdf"
            write_image_bookmark_document(
                path, anchored=True, image_count=2, anchor_wrap="tight",
                anchor_wrap_text=None, anchor_spacing_inches=0.6, flow_text_repetitions=20,
            )
            with ZipFile(path) as source:
                parts = {name: source.read(name) for name in source.namelist()}
            document_xml = ET.fromstring(parts["word/document.xml"])
            anchors = list(document_xml.iter(qn("wp:anchor")))
            self.assertEqual(len(anchors), 2)
            first_wrap = next(child for child in anchors[0] if child.tag == qn("wp:wrapTight"))
            first_wrap.tag = qn("wp:wrapSquare")
            first_wrap.clear()
            for anchor in anchors:
                for name in ("distT", "distB", "distL", "distR"):
                    anchor.attrib.pop(name, None)
            parts["word/document.xml"] = ET.tostring(
                document_xml, encoding="utf-8", xml_declaration=True,
            )
            with ZipFile(path, "w") as target:
                for name, content in parts.items():
                    target.writestr(name, content)

            before_semantics = docx_body_drawing_semantics(path)
            self.assertEqual(before_semantics["unresolved_drawings"], [])
            self.assertEqual(len(before_semantics["drawings"]), 2)
            self.assertEqual([
                drawing["wrap_distances"] for drawing in before_semantics["drawings"]
            ], [
                {"distT": "0", "distB": "0", "distL": "114300", "distR": "114300"}
            ] * 2)
            self.assertTrue(run_conversion_smoke(
                "writer", str(path), str(before_pdf), "pdf",
            )["ok"])
            before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
            before_image = before_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
            before_text = before_pdf_doc[0].get_textpage().get_text_range()
            before_pdf_doc.close()

            ok, record, errors = register_document("writer", str(path), tmp)
            self.assertTrue(ok, errors)
            ok, result, errors, _ = writer_fill_bookmark(
                record["document_id"], "CellMark", "Old", "mixed-anchor-defaults",
                workspace=tmp,
            )
            after_semantics = docx_body_drawing_semantics(path)
            self.assertTrue(ok, (errors, result, drawing_structure_differences(
                before_semantics, after_semantics,
            )))
            self.assertEqual(after_semantics, before_semantics)
            self.assertTrue(run_conversion_smoke(
                "writer", str(path), str(after_pdf), "pdf",
            )["ok"])
            after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
            after_image = after_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
            after_text = after_pdf_doc[0].get_textpage().get_text_range()
            after_pdf_doc.close()
            self.assertEqual(after_text, before_text)
            self.assertIsNone(ImageChops.difference(before_image, after_image).getbbox())

    def test_wps_table_bookmark_keeps_other_explicit_distances_with_default_wrap_text(self):
        import xml.etree.ElementTree as ET
        from zipfile import ZipFile

        defaults = {"distT": "0", "distB": "0", "distL": "114300", "distR": "114300"}
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for explicit_name in ("distB", "distL", "distR"):
                with self.subTest(distance=explicit_name):
                    path = root / f"explicit-{explicit_name}-default-wrap.docx"
                    before_pdf = root / f"before-{explicit_name}.pdf"
                    after_pdf = root / f"after-{explicit_name}.pdf"
                    write_image_bookmark_document(
                        path, anchored=True, anchor_wrap="square", anchor_wrap_text=None,
                        flow_text_repetitions=20,
                    )
                    with ZipFile(path) as source:
                        parts = {name: source.read(name) for name in source.namelist()}
                    document_xml = ET.fromstring(parts["word/document.xml"])
                    anchor = next(document_xml.iter(qn("wp:anchor")))
                    for name in defaults:
                        anchor.attrib.pop(name, None)
                    anchor.set(explicit_name, "635")
                    parts["word/document.xml"] = ET.tostring(
                        document_xml, encoding="utf-8", xml_declaration=True,
                    )
                    with ZipFile(path, "w") as target:
                        for name, content in parts.items():
                            target.writestr(name, content)

                    expected = {**defaults, explicit_name: "635"}
                    before_semantics = docx_body_drawing_semantics(path)
                    self.assertEqual(before_semantics["unresolved_drawings"], [])
                    self.assertEqual(before_semantics["drawings"][0]["wrap_distances"], expected)
                    self.assertTrue(run_conversion_smoke(
                        "writer", str(path), str(before_pdf), "pdf",
                    )["ok"])
                    before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
                    before_image = before_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                    before_text = before_pdf_doc[0].get_textpage().get_text_range()
                    before_pdf_doc.close()

                    ok, record, errors = register_document("writer", str(path), tmp)
                    self.assertTrue(ok, errors)
                    ok, result, errors, _ = writer_fill_bookmark(
                        record["document_id"], "CellMark", "Old",
                        f"explicit-{explicit_name}", workspace=tmp,
                    )
                    after_semantics = docx_body_drawing_semantics(path)
                    self.assertTrue(ok, (errors, result, drawing_structure_differences(
                        before_semantics, after_semantics,
                    )))
                    self.assertEqual(after_semantics, before_semantics)
                    self.assertTrue(run_conversion_smoke(
                        "writer", str(path), str(after_pdf), "pdf",
                    )["ok"])
                    after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
                    after_image = after_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                    after_text = after_pdf_doc[0].get_textpage().get_text_range()
                    after_pdf_doc.close()
                    self.assertEqual(after_text, before_text)
                    self.assertIsNone(ImageChops.difference(before_image, after_image).getbbox())

    def test_wps_table_bookmark_keeps_explicit_distance_with_default_wrap_text(self):
        import xml.etree.ElementTree as ET
        from zipfile import ZipFile

        expected_distances = {
            "distT": "635", "distB": "0", "distL": "114300", "distR": "114300",
        }
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for wrap_mode in ("none", "square", "tight", "through", "top_and_bottom"):
                with self.subTest(wrap=wrap_mode):
                    path = root / f"explicit-distance-default-wrap-{wrap_mode}.docx"
                    before_pdf = root / f"before-explicit-distance-{wrap_mode}.pdf"
                    after_pdf = root / f"after-explicit-distance-{wrap_mode}.pdf"
                    write_image_bookmark_document(
                        path, anchored=True, anchor_wrap=wrap_mode, anchor_wrap_text=None,
                        flow_text_repetitions=20,
                    )
                    with ZipFile(path) as source:
                        parts = {name: source.read(name) for name in source.namelist()}
                    document_xml = ET.fromstring(parts["word/document.xml"])
                    anchor = next(document_xml.iter(qn("wp:anchor")))
                    for name in ("distT", "distB", "distL", "distR"):
                        anchor.attrib.pop(name, None)
                    anchor.set("distT", "635")
                    parts["word/document.xml"] = ET.tostring(
                        document_xml, encoding="utf-8", xml_declaration=True,
                    )
                    with ZipFile(path, "w") as target:
                        for name, content in parts.items():
                            target.writestr(name, content)

                    before_semantics = docx_body_drawing_semantics(path)
                    self.assertEqual(before_semantics["unresolved_drawings"], [])
                    self.assertEqual(before_semantics["drawings"][0]["wrap_distances"], expected_distances)
                    self.assertTrue(run_conversion_smoke(
                        "writer", str(path), str(before_pdf), "pdf",
                    )["ok"])
                    before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
                    before_image = before_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                    before_text = before_pdf_doc[0].get_textpage().get_text_range()
                    before_pdf_doc.close()

                    ok, record, errors = register_document("writer", str(path), tmp)
                    self.assertTrue(ok, errors)
                    ok, result, errors, _ = writer_fill_bookmark(
                        record["document_id"], "CellMark", "Old",
                        f"explicit-distance-{wrap_mode}", workspace=tmp,
                    )
                    after_semantics = docx_body_drawing_semantics(path)
                    self.assertTrue(ok, (errors, result, drawing_structure_differences(
                        before_semantics, after_semantics,
                    )))
                    self.assertEqual(after_semantics, before_semantics)
                    self.assertTrue(run_conversion_smoke(
                        "writer", str(path), str(after_pdf), "pdf",
                    )["ok"])
                    after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
                    after_image = after_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                    after_text = after_pdf_doc[0].get_textpage().get_text_range()
                    after_pdf_doc.close()
                    self.assertEqual(after_text, before_text)
                    self.assertIsNone(ImageChops.difference(before_image, after_image).getbbox())

    def test_wps_table_bookmark_normalizes_combined_omitted_wrap_defaults(self):
        import xml.etree.ElementTree as ET
        from zipfile import ZipFile

        defaults = {"distT": "0", "distB": "0", "distL": "114300", "distR": "114300"}
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for wrap_mode in ("none", "square", "tight", "through", "top_and_bottom"):
                with self.subTest(wrap=wrap_mode):
                    path = root / f"combined-defaults-{wrap_mode}.docx"
                    before_pdf = root / f"before-combined-{wrap_mode}.pdf"
                    after_pdf = root / f"after-combined-{wrap_mode}.pdf"
                    write_image_bookmark_document(
                        path, anchored=True, anchor_wrap=wrap_mode, anchor_wrap_text=None,
                        flow_text_repetitions=20,
                    )
                    with ZipFile(path) as source:
                        parts = {name: source.read(name) for name in source.namelist()}
                    document_xml = ET.fromstring(parts["word/document.xml"])
                    anchor = next(document_xml.iter(qn("wp:anchor")))
                    for name in defaults:
                        anchor.attrib.pop(name, None)
                    parts["word/document.xml"] = ET.tostring(
                        document_xml, encoding="utf-8", xml_declaration=True,
                    )
                    with ZipFile(path, "w") as target:
                        for name, content in parts.items():
                            target.writestr(name, content)

                    before_semantics = docx_body_drawing_semantics(path)
                    self.assertEqual(before_semantics["unresolved_drawings"], [])
                    self.assertEqual(before_semantics["drawings"][0]["wrap_distances"], defaults)
                    self.assertTrue(run_conversion_smoke(
                        "writer", str(path), str(before_pdf), "pdf",
                    )["ok"])
                    before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
                    before_image = before_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                    before_text = before_pdf_doc[0].get_textpage().get_text_range()
                    before_pdf_doc.close()

                    ok, record, errors = register_document("writer", str(path), tmp)
                    self.assertTrue(ok, errors)
                    ok, result, errors, _ = writer_fill_bookmark(
                        record["document_id"], "CellMark", "Old",
                        f"combined-defaults-{wrap_mode}", workspace=tmp,
                    )
                    after_semantics = docx_body_drawing_semantics(path)
                    self.assertTrue(ok, (errors, result, drawing_structure_differences(
                        before_semantics, after_semantics,
                    )))
                    self.assertEqual(after_semantics, before_semantics)
                    self.assertTrue(run_conversion_smoke(
                        "writer", str(path), str(after_pdf), "pdf",
                    )["ok"])
                    after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
                    after_image = after_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                    after_text = after_pdf_doc[0].get_textpage().get_text_range()
                    after_pdf_doc.close()
                    self.assertEqual(after_text, before_text)
                    self.assertIsNone(ImageChops.difference(before_image, after_image).getbbox())

    def test_wps_table_bookmark_normalizes_omitted_zero_wrap_distances(self):
        import xml.etree.ElementTree as ET
        from zipfile import ZipFile

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for wrap_mode in ("none", "square", "tight", "through", "top_and_bottom"):
                reports = []
                for index, omitted in enumerate((True, False)):
                    with self.subTest(wrap=wrap_mode, omitted=omitted):
                        path = root / f"distance-default-{wrap_mode}-{index}.docx"
                        before_pdf = root / f"before-distance-default-{wrap_mode}-{index}.pdf"
                        after_pdf = root / f"after-distance-default-{wrap_mode}-{index}.pdf"
                        write_image_bookmark_document(path, anchored=True, anchor_wrap=wrap_mode)
                        with ZipFile(path) as source:
                            parts = {name: source.read(name) for name in source.namelist()}
                        document_xml = ET.fromstring(parts["word/document.xml"])
                        anchor = next(document_xml.iter(qn("wp:anchor")))
                        for name in ("distT", "distB", "distL", "distR"):
                            if omitted:
                                anchor.attrib.pop(name, None)
                            else:
                                anchor.set(name, "0")
                        parts["word/document.xml"] = ET.tostring(
                            document_xml, encoding="utf-8", xml_declaration=True,
                        )
                        with ZipFile(path, "w") as target:
                            for name, content in parts.items():
                                target.writestr(name, content)

                        before_semantics = docx_body_drawing_semantics(path)
                        self.assertEqual(before_semantics["unresolved_drawings"], [])
                        reports.append(before_semantics)
                        self.assertTrue(run_conversion_smoke("writer", str(path), str(before_pdf), "pdf")["ok"])
                        before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
                        before_image = before_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                        before_text = before_pdf_doc[0].get_textpage().get_text_range()
                        before_pdf_doc.close()

                        ok, record, errors = register_document("writer", str(path), tmp)
                        self.assertTrue(ok, errors)
                        ok, result, errors, _ = writer_fill_bookmark(
                            record["document_id"], "CellMark", "Old",
                            f"distance-default-{wrap_mode}-{index}", workspace=tmp,
                        )
                        after_semantics = docx_body_drawing_semantics(path)
                        self.assertTrue(ok, (errors, result, drawing_structure_differences(
                            before_semantics, after_semantics,
                        )))
                        self.assertEqual(after_semantics, before_semantics)
                        self.assertTrue(run_conversion_smoke("writer", str(path), str(after_pdf), "pdf")["ok"])
                        after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
                        after_image = after_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                        after_text = after_pdf_doc[0].get_textpage().get_text_range()
                        after_pdf_doc.close()
                        self.assertEqual(after_text, before_text)
                        self.assertIsNone(ImageChops.difference(before_image, after_image).getbbox())
                self.assertEqual(reports[0]["drawings"][0]["wrap_distances"], {
                    "distT": "0", "distB": "0", "distL": "114300", "distR": "114300",
                })
                self.assertEqual(reports[1]["drawings"][0]["wrap_distances"], {
                    name: "0" for name in ("distT", "distB", "distL", "distR")
                })
                # The omitted defaults are semantically distinct from explicit
                # zero, but this fixture need not expose that distinction in
                # rendered pixels for every wrap mode.

    def test_wps_table_bookmark_preserves_one_sided_wrap_with_omitted_distances(self):
        import xml.etree.ElementTree as ET
        from zipfile import ZipFile

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for wrap_text in ("left", "right", "largest"):
                with self.subTest(wrap_text=wrap_text):
                    path = root / f"omitted-distance-{wrap_text}.docx"
                    before_pdf = root / f"before-{wrap_text}.pdf"
                    after_pdf = root / f"after-{wrap_text}.pdf"
                    write_image_bookmark_document(
                        path, anchored=True, anchor_wrap="square",
                        anchor_wrap_text=wrap_text, flow_text_repetitions=35,
                    )
                    with ZipFile(path) as source:
                        parts = {name: source.read(name) for name in source.namelist()}
                    document_xml = ET.fromstring(parts["word/document.xml"])
                    anchor = next(document_xml.iter(qn("wp:anchor")))
                    for name in ("distT", "distB", "distL", "distR"):
                        anchor.attrib.pop(name, None)
                    parts["word/document.xml"] = ET.tostring(
                        document_xml, encoding="utf-8", xml_declaration=True,
                    )
                    with ZipFile(path, "w") as target:
                        for name, content in parts.items():
                            target.writestr(name, content)

                    before_semantics = docx_body_drawing_semantics(path)
                    self.assertEqual(before_semantics["unresolved_drawings"], [])
                    wrap = before_semantics["drawings"][0]["structure"]["children"][4]
                    self.assertEqual(wrap["attributes"]["wrapText"], wrap_text)
                    self.assertEqual(before_semantics["drawings"][0]["wrap_distances"], {
                        "distT": "0", "distB": "0", "distL": "114300", "distR": "114300",
                    })
                    self.assertTrue(run_conversion_smoke("writer", str(path), str(before_pdf), "pdf")["ok"])
                    before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
                    before_image = before_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                    before_text = before_pdf_doc[0].get_textpage().get_text_range()
                    before_pdf_doc.close()

                    ok, record, errors = register_document("writer", str(path), tmp)
                    self.assertTrue(ok, errors)
                    ok, result, errors, _ = writer_fill_bookmark(
                        record["document_id"], "CellMark", "Old",
                        f"one-sided-{wrap_text}", workspace=tmp,
                    )
                    after_semantics = docx_body_drawing_semantics(path)
                    self.assertTrue(ok, (errors, result, drawing_structure_differences(
                        before_semantics, after_semantics,
                    )))
                    self.assertEqual(after_semantics, before_semantics)
                    self.assertTrue(run_conversion_smoke("writer", str(path), str(after_pdf), "pdf")["ok"])
                    after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
                    after_image = after_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                    after_text = after_pdf_doc[0].get_textpage().get_text_range()
                    after_pdf_doc.close()
                    self.assertEqual(after_text, before_text)
                    self.assertIsNone(ImageChops.difference(before_image, after_image).getbbox())

    def test_wps_table_bookmark_applies_defaults_per_omitted_distance(self):
        import xml.etree.ElementTree as ET
        from zipfile import ZipFile

        cases = (
            ("omit_top_right", {"distB": "635", "distL": "1270"}),
            ("omit_bottom_left", {"distT": "635", "distR": "1270"}),
            ("omit_top_bottom", {"distL": "635", "distR": "1270"}),
        )
        defaults = {"distT": "0", "distB": "0", "distL": "114300", "distR": "114300"}
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for case_name, explicit in cases:
                with self.subTest(case=case_name):
                    path = root / f"mixed-distance-{case_name}.docx"
                    before_pdf = root / f"before-{case_name}.pdf"
                    after_pdf = root / f"after-{case_name}.pdf"
                    write_image_bookmark_document(path, anchored=True, anchor_wrap="square")
                    with ZipFile(path) as source:
                        parts = {name: source.read(name) for name in source.namelist()}
                    document_xml = ET.fromstring(parts["word/document.xml"])
                    anchor = next(document_xml.iter(qn("wp:anchor")))
                    for name in ("distT", "distB", "distL", "distR"):
                        anchor.attrib.pop(name, None)
                    anchor.attrib.update(explicit)
                    parts["word/document.xml"] = ET.tostring(
                        document_xml, encoding="utf-8", xml_declaration=True,
                    )
                    with ZipFile(path, "w") as target:
                        for name, content in parts.items():
                            target.writestr(name, content)

                    expected = {**defaults, **explicit}
                    before_semantics = docx_body_drawing_semantics(path)
                    self.assertEqual(before_semantics["unresolved_drawings"], [])
                    self.assertEqual(before_semantics["drawings"][0]["wrap_distances"], expected)
                    self.assertTrue(run_conversion_smoke("writer", str(path), str(before_pdf), "pdf")["ok"])
                    before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
                    before_image = before_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                    before_text = before_pdf_doc[0].get_textpage().get_text_range()
                    before_pdf_doc.close()

                    ok, record, errors = register_document("writer", str(path), tmp)
                    self.assertTrue(ok, errors)
                    ok, result, errors, _ = writer_fill_bookmark(
                        record["document_id"], "CellMark", "Old", case_name, workspace=tmp,
                    )
                    after_semantics = docx_body_drawing_semantics(path)
                    self.assertTrue(ok, (errors, result, drawing_structure_differences(
                        before_semantics, after_semantics,
                    )))
                    self.assertEqual(after_semantics, before_semantics)
                    self.assertTrue(run_conversion_smoke("writer", str(path), str(after_pdf), "pdf")["ok"])
                    after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
                    after_image = after_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                    after_text = after_pdf_doc[0].get_textpage().get_text_range()
                    after_pdf_doc.close()
                    self.assertEqual(after_text, before_text)
                    self.assertIsNone(ImageChops.difference(before_image, after_image).getbbox())

    def test_wps_wrap_none_omitted_and_both_sides_render_identically(self):
        import xml.etree.ElementTree as ET
        from docx.oxml.ns import qn
        from zipfile import ZipFile

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            reports = []
            for value in (None, "bothSides"):
                suffix = "omitted" if value is None else "both-sides"
                path = root / f"wrap-none-{suffix}.docx"
                pdf_path = root / f"wrap-none-{suffix}.pdf"
                write_image_bookmark_document(
                    path, anchored=True, anchor_wrap="none", anchor_wrap_text=None,
                    flow_text_repetitions=35,
                )
                if value is not None:
                    with ZipFile(path) as source:
                        parts = {name: source.read(name) for name in source.namelist()}
                    document_xml = ET.fromstring(parts["word/document.xml"])
                    next(document_xml.iter(qn("wp:wrapNone"))).set("wrapText", value)
                    parts["word/document.xml"] = ET.tostring(
                        document_xml, encoding="utf-8", xml_declaration=True,
                    )
                    with ZipFile(path, "w") as target:
                        for name, content in parts.items():
                            target.writestr(name, content)
                semantics = docx_body_drawing_semantics(path)
                self.assertEqual(semantics["unresolved_drawings"], [])
                reports.append(semantics)
                self.assertTrue(run_conversion_smoke(
                    "writer", str(path), str(pdf_path), "pdf",
                )["ok"])
                pdf = pypdfium2.PdfDocument(str(pdf_path))
                reports.append((
                    [pdf[index].get_textpage().get_text_range() for index in range(len(pdf))],
                    [pdf[index].render(scale=1.5).to_pil().convert("RGB") for index in range(len(pdf))],
                ))
                pdf.close()

            self.assertNotEqual(reports[0], reports[2])
            self.assertEqual(reports[1][0], reports[3][0])
            self.assertEqual(len(reports[1][1]), len(reports[3][1]))
            for expected, actual in zip(reports[1][1], reports[3][1]):
                self.assertIsNone(ImageChops.difference(expected, actual).getbbox())

    def test_wps_wrap_none_one_sided_values_render_identically(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            reference_text = None
            reference_image = None
            for wrap_text in ("left", "right", "largest"):
                with self.subTest(wrap_text=wrap_text):
                    path = root / f"wrap-none-render-{wrap_text}.docx"
                    pdf_path = root / f"wrap-none-render-{wrap_text}.pdf"
                    write_image_bookmark_document(
                        path, anchored=True, anchor_wrap="none", anchor_wrap_text=wrap_text,
                        flow_text_repetitions=35,
                    )
                    self.assertTrue(run_conversion_smoke(
                        "writer", str(path), str(pdf_path), "pdf",
                    )["ok"])
                    pdf = pypdfium2.PdfDocument(str(pdf_path))
                    self.assertGreater(len(pdf), 0)
                    page_text = [pdf[index].get_textpage().get_text_range() for index in range(len(pdf))]
                    page_images = [
                        pdf[index].render(scale=1.5).to_pil().convert("RGB")
                        for index in range(len(pdf))
                    ]
                    pdf.close()
                    if reference_text is None:
                        reference_text = page_text
                        reference_image = page_images
                    else:
                        self.assertEqual(page_text, reference_text)
                        self.assertEqual(len(page_images), len(reference_image))
                        for expected, actual in zip(reference_image, page_images):
                            self.assertIsNone(ImageChops.difference(expected, actual).getbbox())

    def test_wps_table_bookmark_preserves_polygon_one_sided_wrap_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for mode in ("tight", "through"):
                for wrap_text in ("left", "right", "largest"):
                    with self.subTest(wrap=mode, wrap_text=wrap_text):
                        path = root / f"{mode}-{wrap_text}-wrap.docx"
                        before_pdf = root / f"before-{mode}-{wrap_text}.pdf"
                        after_pdf = root / f"after-{mode}-{wrap_text}.pdf"
                        write_image_bookmark_document(
                            path, anchored=True, anchor_wrap=mode, anchor_wrap_text=wrap_text,
                            flow_text_repetitions=25,
                        )
                        before_semantics = docx_body_drawing_semantics(path)
                        self.assertEqual(before_semantics["unresolved_drawings"], [])
                        self.assertTrue(run_conversion_smoke(
                            "writer", str(path), str(before_pdf), "pdf",
                        )["ok"])
                        before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
                        before_image = before_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                        before_text = before_pdf_doc[0].get_textpage().get_text_range()
                        before_pdf_doc.close()

                        ok, record, errors = register_document("writer", str(path), tmp)
                        self.assertTrue(ok, errors)
                        ok, result, errors, _ = writer_fill_bookmark(
                            record["document_id"], "CellMark", "Old",
                            f"{mode}-{wrap_text}", workspace=tmp,
                        )
                        after_semantics = docx_body_drawing_semantics(path)
                        self.assertTrue(ok, (errors, result, drawing_structure_differences(
                            before_semantics, after_semantics,
                        )))
                        self.assertEqual(after_semantics, before_semantics)
                        self.assertTrue(run_conversion_smoke(
                            "writer", str(path), str(after_pdf), "pdf",
                        )["ok"])
                        after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
                        after_image = after_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                        after_text = after_pdf_doc[0].get_textpage().get_text_range()
                        after_pdf_doc.close()
                        self.assertEqual(after_text, before_text)
                        self.assertIsNone(ImageChops.difference(before_image, after_image).getbbox())

    def test_wps_table_bookmark_preserves_wrap_side_and_distances(self):
        variants = (
            ("left", (0.0, 0.0, 0.05, 0.3)),
            ("right", (0.15, 0.05, 0.3, 0.05)),
            ("largest", (0.2, 0.2, 0.2, 0.2)),
        )
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for wrap_text, distances in variants:
                with self.subTest(wrap_text=wrap_text):
                    path = root / f"wrap-side-{wrap_text}.docx"
                    before_pdf = root / f"before-{wrap_text}.pdf"
                    after_pdf = root / f"after-{wrap_text}.pdf"
                    write_image_bookmark_document(
                        path, anchored=True, image_width_inches=1.0,
                        anchor_offsets_inches=(2.0, 0.0), anchor_wrap="square",
                        anchor_wrap_text=wrap_text, anchor_distances_inches=distances,
                        flow_text_repetitions=35,
                    )
                    before_semantics = docx_body_drawing_semantics(path)
                    expected_distances = {
                        name: str(int(Inches(value).emu))
                        for name, value in zip(("distT", "distB", "distL", "distR"), distances)
                    }
                    self.assertEqual(before_semantics["drawings"][0]["wrap_distances"], expected_distances)
                    self.assertTrue(run_conversion_smoke("writer", str(path), str(before_pdf), "pdf")["ok"])
                    before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
                    before_image = before_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                    before_text = before_pdf_doc[0].get_textpage().get_text_range()
                    before_pdf_doc.close()

                    ok, record, errors = register_document("writer", str(path), tmp)
                    self.assertTrue(ok, errors)
                    ok, result, errors, _ = writer_fill_bookmark(
                        record["document_id"], "CellMark", "Old", f"wrap-side-{wrap_text}", workspace=tmp,
                    )
                    self.assertTrue(ok, (errors, result))
                    self.assertEqual(docx_body_drawing_semantics(path), before_semantics)
                    self.assertTrue(run_conversion_smoke("writer", str(path), str(after_pdf), "pdf")["ok"])
                    after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
                    after_image = after_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                    after_text = after_pdf_doc[0].get_textpage().get_text_range()
                    after_pdf_doc.close()
                    self.assertEqual(after_text, before_text)
                    self.assertIsNone(ImageChops.difference(before_image, after_image).getbbox())

    def test_wps_table_bookmark_preserves_tight_and_through_wrap_polygons(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for mode in ("tight", "through"):
                with self.subTest(wrap=mode):
                    path = root / f"{mode}-wrap.docx"
                    before_pdf = root / f"before-{mode}.pdf"
                    after_pdf = root / f"after-{mode}.pdf"
                    write_image_bookmark_document(
                        path, anchored=True, image_width_inches=1.0,
                        anchor_offsets_inches=(2.5, 0.0), anchor_wrap=mode,
                    )
                    before_semantics = docx_body_drawing_semantics(path)
                    wrap = before_semantics["drawings"][0]["structure"]["children"][4]
                    self.assertEqual(wrap["tag"].rsplit("}", 1)[-1], f"wrap{mode.title()}")
                    self.assertEqual(wrap["children"][0]["tag"].rsplit("}", 1)[-1], "wrapPolygon")
                    self.assertEqual(len(wrap["children"][0]["children"]), 4)
                    self.assertTrue(run_conversion_smoke("writer", str(path), str(before_pdf), "pdf")["ok"])
                    before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
                    before_image = before_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                    before_text = before_pdf_doc[0].get_textpage().get_text_range()
                    before_pdf_doc.close()

                    ok, record, errors = register_document("writer", str(path), tmp)
                    self.assertTrue(ok, errors)
                    ok, result, errors, _ = writer_fill_bookmark(
                        record["document_id"], "CellMark", "Old", f"{mode}-wrap", workspace=tmp,
                    )
                    self.assertTrue(ok, (errors, result))
                    self.assertEqual(docx_body_drawing_semantics(path), before_semantics)
                    self.assertTrue(run_conversion_smoke("writer", str(path), str(after_pdf), "pdf")["ok"])
                    after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
                    after_image = after_pdf_doc[0].render(scale=1.5).to_pil().convert("RGB")
                    after_text = after_pdf_doc[0].get_textpage().get_text_range()
                    after_pdf_doc.close()
                    self.assertEqual(before_text, after_text)
                    self.assertIsNone(ImageChops.difference(before_image, after_image).getbbox())

    def test_wps_table_bookmark_preserves_multi_anchor_pagination_and_order(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / "multi-anchor-pages.docx"
            before_pdf = root / "before-multi-anchor.pdf"
            after_pdf = root / "after-multi-anchor.pdf"
            write_image_bookmark_document(
                path, anchored=True, image_width_inches=0.9,
                anchor_relative_frames=("column", "paragraph"),
                anchor_offsets_inches=(0.75, 0.0), anchor_spacing_inches=0.25,
                anchor_wrap="square", preamble_paragraphs=36,
                preamble_images=1, postamble_paragraphs=36,
            )
            before_semantics = docx_body_drawing_semantics(path)
            self.assertEqual(len(before_semantics["drawings"]), 2)
            self.assertEqual([item["relative_height_order"] for item in before_semantics["drawings"]], [0, 1])
            self.assertTrue(run_conversion_smoke("writer", str(path), str(before_pdf), "pdf")["ok"])

            def render_report(pdf_path):
                pdf = pypdfium2.PdfDocument(str(pdf_path))
                reports = []
                markers = {}
                for page_index in range(len(pdf)):
                    page = pdf[page_index]
                    image = page.render(scale=1.5).to_pil().convert("RGB")
                    text_page = page.get_textpage()
                    text = text_page.get_text_range()
                    page_markers = {}
                    for marker in ("PREFACE-START", "AFTER-TABLE-MARKER"):
                        start = text.find(marker)
                        if start >= 0:
                            boxes = [text_page.get_charbox(index) for index in range(start, start + len(marker))]
                            page_markers[marker] = tuple(round(value, 2) for value in (
                                min(box[0] for box in boxes), min(box[1] for box in boxes),
                                max(box[2] for box in boxes), max(box[3] for box in boxes),
                            ))
                            markers[marker] = (page_index, page_markers[marker])
                    red, green, blue = image.split()
                    red_mask = ImageChops.multiply(
                        ImageChops.subtract(red, blue).point(lambda value: 255 if value > 100 else 0),
                        red.point(lambda value: 255 if value > 150 else 0),
                    )
                    blue_mask = ImageChops.multiply(
                        ImageChops.subtract(blue, red).point(lambda value: 255 if value > 100 else 0),
                        blue.point(lambda value: 255 if value > 150 else 0),
                    )
                    reports.append({"image": image, "red": red_mask, "blue": blue_mask})
                pdf.close()
                return reports, markers

            before_pages, before_markers = render_report(before_pdf)
            self.assertGreaterEqual(len(before_pages), 2)
            self.assertEqual(set(before_markers), {"PREFACE-START", "AFTER-TABLE-MARKER"})
            red_pages = [index for index, page in enumerate(before_pages) if page["red"].getbbox()]
            blue_pages = [index for index, page in enumerate(before_pages) if page["blue"].getbbox()]
            self.assertEqual(len(red_pages), 1, red_pages)
            self.assertEqual(len(blue_pages), 1, blue_pages)
            self.assertEqual(red_pages[0], 0, red_pages)
            self.assertGreater(blue_pages[0], red_pages[0], blue_pages)

            ok, record, errors = register_document("writer", str(path), tmp)
            self.assertTrue(ok, errors)
            ok, result, errors, _ = writer_fill_bookmark(
                record["document_id"], "CellMark", "Old", "multi-anchor-pages", workspace=tmp,
            )
            self.assertTrue(ok, (errors, result))
            after_semantics = docx_body_drawing_semantics(path)
            self.assertEqual(after_semantics, before_semantics)
            self.assertTrue(run_conversion_smoke("writer", str(path), str(after_pdf), "pdf")["ok"])
            after_pages, after_markers = render_report(after_pdf)
            self.assertEqual(len(after_pages), len(before_pages))
            self.assertEqual(after_markers, before_markers)
            for index, (before, after) in enumerate(zip(before_pages, after_pages)):
                self.assertIsNone(ImageChops.difference(before["image"], after["image"]).getbbox(), index)
                self.assertIsNone(ImageChops.difference(before["red"], after["red"]).getbbox(), index)
                self.assertIsNone(ImageChops.difference(before["blue"], after["blue"]).getbbox(), index)

    def test_wps_table_bookmark_preserves_multi_page_anchor_flow(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / "multi-page-flow.docx"
            before_pdf = root / "before-flow.pdf"
            after_pdf = root / "after-flow.pdf"
            write_image_bookmark_document(
                path, anchored=True, image_width_inches=0.9,
                anchor_relative_frames=("column", "paragraph"),
                anchor_offsets_inches=(2.0, 0.0), anchor_wrap="square",
                preamble_paragraphs=36, postamble_paragraphs=36,
            )
            before_semantics = docx_body_drawing_semantics(path)
            self.assertTrue(run_conversion_smoke("writer", str(path), str(before_pdf), "pdf")["ok"])

            def render_report(pdf_path):
                pdf = pypdfium2.PdfDocument(str(pdf_path))
                page_images = []
                markers = {}
                for page_index in range(len(pdf)):
                    page = pdf[page_index]
                    page_images.append(page.render(scale=1.5).to_pil().convert("RGB"))
                    text_page = page.get_textpage()
                    text = text_page.get_text_range()
                    for marker in ("PREFACE-START", "AFTER-TABLE-MARKER"):
                        start = text.find(marker)
                        if start >= 0:
                            boxes = [text_page.get_charbox(index) for index in range(start, start + len(marker))]
                            markers[marker] = (
                                page_index,
                                tuple(round(value, 2) for value in (
                                    min(box[0] for box in boxes), min(box[1] for box in boxes),
                                    max(box[2] for box in boxes), max(box[3] for box in boxes),
                                )),
                            )
                pdf.close()
                return page_images, markers

            before_pages, before_markers = render_report(before_pdf)
            self.assertGreaterEqual(len(before_pages), 2)
            self.assertEqual(set(before_markers), {"PREFACE-START", "AFTER-TABLE-MARKER"})
            ok, record, errors = register_document("writer", str(path), tmp)
            self.assertTrue(ok, errors)
            ok, result, errors, _ = writer_fill_bookmark(
                record["document_id"], "CellMark", "Old", "multi-page-flow", workspace=tmp,
            )
            self.assertTrue(ok, (errors, result))
            self.assertEqual(docx_body_drawing_semantics(path), before_semantics)
            self.assertTrue(run_conversion_smoke("writer", str(path), str(after_pdf), "pdf")["ok"])
            after_pages, after_markers = render_report(after_pdf)
            self.assertEqual(len(after_pages), len(before_pages))
            self.assertEqual(after_markers, before_markers)
            for index, (before, after) in enumerate(zip(before_pages, after_pages)):
                self.assertEqual(before.size, after.size, index)
                self.assertIsNone(ImageChops.difference(before, after).getbbox(), index)

    def test_wps_table_bookmark_preserves_wrapping_and_edge_clipping(self):
        variants = (
            ("cell_left_none", False, ("column", "paragraph"), (0.0, 0.0), "none", 0.5),
            ("cell_right_square_crop", True, ("column", "paragraph"), (8.5, 0.0), "square", 1.5),
            ("cell_left_negative_crop", True, ("column", "paragraph"), (-0.5, 0.0), "square", 1.5),
            ("cell_middle_top_bottom", False, ("column", "paragraph"), (2.5, 0.0), "top_and_bottom", 1.0),
            ("page_top_left_negative", True, ("page", "page"), (-0.5, -0.25), "square", 1.5),
        )

        def after_word_bounds(pdf_path):
            pdf_doc = pypdfium2.PdfDocument(str(pdf_path))
            text_page = pdf_doc[0].get_textpage()
            text = text_page.get_text_range()
            start = text.find("After")
            self.assertGreaterEqual(start, 0, text)
            boxes = [text_page.get_charbox(index) for index in range(start, start + 5)]
            pdf_doc.close()
            return tuple(round(value, 2) for value in (
                min(box[0] for box in boxes), min(box[1] for box in boxes),
                max(box[2] for box in boxes), max(box[3] for box in boxes),
            ))

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for name, cropped, frames, offsets, wrap, width in variants:
                with self.subTest(variant=name):
                    path = root / f"wrap-{name}.docx"
                    before_pdf = root / f"before-{name}.pdf"
                    after_pdf = root / f"after-{name}.pdf"
                    write_image_bookmark_document(
                        path, anchored=True, cropped=cropped, image_width_inches=width,
                        anchor_relative_frames=frames, anchor_offsets_inches=offsets,
                        anchor_wrap=wrap,
                        layout_in_cell="0" if name == "page_top_left_negative" else "1",
                    )
                    before_semantics = docx_body_drawing_semantics(path)
                    self.assertEqual(before_semantics["unresolved_drawings"], [])
                    self.assertTrue(run_conversion_smoke("writer", str(path), str(before_pdf), "pdf")["ok"])
                    before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
                    before = before_pdf_doc[0].render(scale=2).to_pil().convert("RGB")
                    before_pdf_doc.close()
                    before_text_bounds = after_word_bounds(before_pdf)

                    ok, record, errors = register_document("writer", str(path), tmp)
                    self.assertTrue(ok, errors)
                    ok, result, errors, _ = writer_fill_bookmark(
                        record["document_id"], "CellMark", "Old", f"wrap-edge-{name}", workspace=tmp,
                    )
                    self.assertTrue(ok, (errors, result))
                    self.assertEqual(docx_body_drawing_semantics(path), before_semantics)
                    self.assertTrue(run_conversion_smoke("writer", str(path), str(after_pdf), "pdf")["ok"])
                    after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
                    after = after_pdf_doc[0].render(scale=2).to_pil().convert("RGB")
                    after_pdf_doc.close()
                    self.assertEqual(after_word_bounds(after_pdf), before_text_bounds)

                    def blue_mask(image):
                        red, green, blue = image.split()
                        return ImageChops.multiply(
                            ImageChops.subtract(blue, red).point(lambda value: 255 if value > 100 else 0),
                            blue.point(lambda value: 255 if value > 150 else 0),
                        )

                    before_mask, after_mask = blue_mask(before), blue_mask(after)
                    visible = before_mask.getbbox()
                    self.assertIsNotNone(visible)
                    self.assertIsNone(ImageChops.difference(before_mask, after_mask).getbbox())
                    if name == "cell_right_square_crop":
                        self.assertGreater(visible[2], round(before.width * 0.8), visible)
                    if name == "cell_left_negative_crop":
                        self.assertLess(visible[0], round(before.width * 0.2), visible)
                    if name == "page_top_left_negative":
                        self.assertEqual(visible[0], 0, visible)
                        self.assertEqual(visible[1], 0, visible)
                        self.assertLess(visible[2] - visible[0], round(width * 144) - 20, visible)
                        self.assertLess(visible[3] - visible[1], round(width * 96) - 20, visible)
                    if name in {"cell_right_square_crop", "cell_left_negative_crop"}:
                        self.assertLessEqual(
                            abs((visible[2] - visible[0]) - round(width * 144)), 3, visible,
                        )
                    if name == "page_top_left_negative":
                        self.assertLessEqual(abs((visible[2] - visible[0]) - 144), 3, visible)
                        self.assertLessEqual(abs((visible[3] - visible[1]) - 108), 3, visible)

    def test_wps_table_bookmark_anchor_reference_frames_keep_rendered_edges(self):
        variants = (
            ("column_paragraph", ("column", "paragraph"), (0.0, 0.0)),
            ("page_page", ("page", "page"), (0.0, 0.0)),
        )
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for name, frames, offsets in variants:
                with self.subTest(variant=name):
                    path = root / f"anchor-{name}.docx"
                    before_pdf = root / f"before-{name}.pdf"
                    after_pdf = root / f"after-{name}.pdf"
                    write_image_bookmark_document(
                        path, anchored=True, image_width_inches=0.5,
                        anchor_relative_frames=frames, anchor_offsets_inches=offsets,
                    )
                    before_semantics = docx_body_drawing_semantics(path)
                    self.assertEqual(before_semantics["unresolved_drawings"], [])
                    self.assertEqual(
                        before_semantics["drawings"][0]["structure"]["children"][1]["attributes"]["relativeFrom"],
                        frames[0],
                    )
                    self.assertTrue(run_conversion_smoke("writer", str(path), str(before_pdf), "pdf")["ok"])
                    before_pdf_doc = pypdfium2.PdfDocument(str(before_pdf))
                    self.assertEqual(len(before_pdf_doc), 1)
                    before = before_pdf_doc[0].render(scale=2).to_pil().convert("RGB")
                    before_pdf_doc.close()

                    ok, record, errors = register_document("writer", str(path), tmp)
                    self.assertTrue(ok, errors)
                    ok, result, errors, _ = writer_fill_bookmark(
                        record["document_id"], "CellMark", "New", f"anchor-frame-{name}", workspace=tmp,
                    )
                    self.assertTrue(ok, (errors, result))
                    after_semantics = docx_body_drawing_semantics(path)
                    self.assertEqual(after_semantics, before_semantics)
                    self.assertTrue(run_conversion_smoke("writer", str(path), str(after_pdf), "pdf")["ok"])
                    after_pdf_doc = pypdfium2.PdfDocument(str(after_pdf))
                    self.assertEqual(len(after_pdf_doc), 1)
                    after = after_pdf_doc[0].render(scale=2).to_pil().convert("RGB")
                    after_pdf_doc.close()
                    self.assertEqual(before.size, after.size)

                    def blue_mask(image):
                        red, green, blue = image.split()
                        dominant = ImageChops.subtract(blue, red)
                        return ImageChops.multiply(
                            dominant.point(lambda value: 255 if value > 100 else 0),
                            blue.point(lambda value: 255 if value > 150 else 0),
                        )

                    before_mask, after_mask = blue_mask(before), blue_mask(after)
                    self.assertIsNotNone(before_mask.getbbox())
                    self.assertIsNone(ImageChops.difference(before_mask, after_mask).getbbox())
                    if name == "page_page":
                        page_edge = before_mask.getbbox()
                        self.assertGreater(page_edge[0], 0, page_edge)
                        self.assertGreater(page_edge[1], 0, page_edge)

    def test_wps_table_bookmark_float_images_keep_rendered_layout(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / "overlap.docx"
            before_pdf = root / "before.pdf"
            after_pdf = root / "after.pdf"
            write_image_bookmark_document(
                path, image_count=2, anchored=True, image_width_inches=1.5,
                anchor_spacing_inches=0.25,
            )
            before_conversion = run_conversion_smoke("writer", str(path), str(before_pdf), "pdf")
            self.assertTrue(before_conversion["ok"], before_conversion)
            before_identity = file_identity(path)
            before_document = pypdfium2.PdfDocument(str(before_pdf))
            self.assertEqual(len(before_document), 1)
            before_export = before_document[0].render(scale=2).to_pil().convert("RGB")
            before_document.close()
            self.assertEqual(file_identity(path), before_identity)
            ok, record, errors = register_document("writer", str(path), tmp)
            self.assertTrue(ok, errors)
            ok, result, errors, _ = writer_fill_bookmark(
                record["document_id"], "CellMark", "New", "rendered-anchors", workspace=tmp,
            )
            self.assertTrue(ok, (errors, result))
            self.assertNotEqual(file_identity(path), before_identity)
            drawings_after = docx_body_drawing_semantics(path)
            converted = run_conversion_smoke("writer", str(path), str(after_pdf), "pdf")
            self.assertTrue(converted["ok"], converted)
            after_document = pypdfium2.PdfDocument(str(after_pdf))
            self.assertEqual(len(after_document), 1)
            after_export = after_document[0].render(scale=2).to_pil().convert("RGB")
            after_document.close()
            self.assertEqual(before_export.size, after_export.size)

            def color_mask(image, color):
                red, green, blue = image.split()
                if color == "blue":
                    dominant = ImageChops.subtract(blue, red)
                    strength = blue
                else:
                    dominant = ImageChops.subtract(red, blue)
                    strength = red
                return ImageChops.multiply(
                    dominant.point(lambda value: 255 if value > 100 else 0),
                    strength.point(lambda value: 255 if value > 150 else 0),
                )

            blue_before, blue_after = color_mask(before_export, "blue"), color_mask(after_export, "blue")
            red_before, red_after = color_mask(before_export, "red"), color_mask(after_export, "red")
            self.assertIsNone(ImageChops.difference(blue_before, blue_after).getbbox())
            self.assertIsNone(ImageChops.difference(red_before, red_after).getbbox())
            blue_box, red_box = blue_after.getbbox(), red_after.getbbox()
            self.assertIsNotNone(blue_box)
            self.assertIsNotNone(red_box)
            blue_width_emu = int(drawings_after["drawings"][0]["extent"]["cx"])
            blue_logical_right = blue_box[0] + round(blue_width_emu * 144 / 914400)
            overlap_left, overlap_right = max(blue_box[0], red_box[0]), min(blue_logical_right, red_box[2])
            overlap_top, overlap_bottom = max(blue_box[1], red_box[1]), min(blue_box[3], red_box[3])
            self.assertLess(overlap_left, overlap_right, (blue_box, red_box))
            self.assertLess(overlap_top, overlap_bottom, (blue_box, red_box))
            sample = ((overlap_left + overlap_right) // 2, (overlap_top + overlap_bottom) // 2)
            self.assertGreater(red_after.getpixel(sample), 0)
            self.assertEqual(blue_after.getpixel(sample), 0)

    def test_wps_table_bookmark_preserves_adjacent_inline_drawing(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "drawing-neighbor.docx"
            write_image_bookmark_document(path)
            drawings_before = docx_body_drawing_semantics(path)
            self.assertEqual(len(drawings_before["drawings"]), 1)
            self.assertEqual(drawings_before["unresolved_drawings"], [])
            ok, record, errors = register_document("writer", str(path), tmp)
            self.assertTrue(ok, errors)
            ok, result, errors, _ = writer_fill_bookmark(
                record["document_id"], "CellMark", "New", "drawing-neighbor", workspace=tmp,
            )
            self.assertTrue(ok, (errors, result, drawings_before,
                                 drawing_structure_differences(
                                     drawings_before["drawings"][0]["structure"],
                                     docx_body_drawing_semantics(path)["drawings"][0]["structure"],
                                 )))
            self.assertEqual(docx_body_drawing_semantics(path), drawings_before)
            self.assertEqual(read_supported_bookmark_text(path, "CellMark")[1], "New")
            self.assertEqual(len(Document(path).inline_shapes), 1)

    def test_wps_table_bookmark_preserves_anchored_and_multiple_cropped_images(self):
        for variant, options in (
            ("anchored_crop", {"anchored": True, "cropped": True}),
            ("multiple_crop", {"image_count": 2, "cropped": True}),
            ("multiple_anchor_crop", {"image_count": 2, "anchored": True, "cropped": True}),
        ):
            with self.subTest(variant=variant), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / f"{variant}.docx"
                write_image_bookmark_document(path, **options)
                drawings_before = docx_body_drawing_semantics(path)
                self.assertEqual(drawings_before["unresolved_drawings"], [])
                self.assertEqual(len(drawings_before["drawings"]), options.get("image_count", 1))
                ok, record, errors = register_document("writer", str(path), tmp)
                self.assertTrue(ok, errors)
                ok, result, errors, _ = writer_fill_bookmark(
                    record["document_id"], "CellMark", "New", f"drawing-{variant}", workspace=tmp,
                )
                after = docx_body_drawing_semantics(path)
                self.assertTrue(ok, (errors, result, drawings_before, after))
                self.assertEqual(after, drawings_before)
                self.assertEqual(read_supported_bookmark_text(path, "CellMark")[1], "New")
                expected_type = "anchor" if options.get("anchored") else "inline"
                self.assertEqual([drawing["type"] for drawing in after["drawings"]],
                                 [expected_type] * options.get("image_count", 1))

    def test_wps_table_bookmark_preserves_hyperlink_and_field_neighbors(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "semantic-neighbors.docx"
            document = Document()
            table = document.add_table(rows=1, cols=1)
            paragraph = table.cell(0, 0).paragraphs[0]
            paragraph.add_run("Before ")
            rel_id = paragraph.part.relate_to(
                "https://example.com/audit?source=writer", RT.HYPERLINK, is_external=True,
            )
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), rel_id)
            link_run = OxmlElement("w:r")
            link_text = OxmlElement("w:t")
            link_text.text = "LINK"
            link_run.append(link_text)
            hyperlink.append(link_run)
            paragraph._p.append(hyperlink)
            paragraph.add_run(" ")
            old = paragraph.add_run("Old")
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), "1")
            start.set(qn("w:name"), "CellMark")
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), "1")
            old._r.addprevious(start)
            old._r.addnext(end)
            paragraph.add_run(" ")
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), "MERGEFIELD InvoiceCode")
            field_run = OxmlElement("w:r")
            field_text = OxmlElement("w:t")
            field_text.text = "CODE"
            field_run.append(field_text)
            field.append(field_run)
            paragraph._p.append(field)
            paragraph.add_run(" After")
            document.save(path)
            semantics_before = docx_body_link_field_semantics(path)
            self.assertEqual(semantics_before["hyperlinks"][0]["target"],
                             "https://example.com/audit?source=writer")
            self.assertEqual(semantics_before["fields"][0]["instruction_tokens"],
                             ["MERGEFIELD", "InvoiceCode"])
            ok, record, errors = register_document("writer", str(path), tmp)
            self.assertTrue(ok, errors)
            ok, result, errors, _ = writer_fill_bookmark(
                record["document_id"], "CellMark", "New", "semantic-neighbors", workspace=tmp,
            )
            self.assertTrue(ok, (errors, result, semantics_before,
                                 docx_body_link_field_semantics(path)))
            self.assertEqual(docx_body_link_field_semantics(path), semantics_before)
            self.assertEqual(read_supported_bookmark_text(path, "CellMark")[1], "New")
            ok, result, errors, _ = writer_fill_bookmark(
                record["document_id"], "CellMark", "Again", "semantic-neighbors-second", workspace=tmp,
            )
            self.assertTrue(ok, (errors, result))
            self.assertEqual(docx_body_link_field_semantics(path), semantics_before)
            self.assertEqual(read_supported_bookmark_text(path, "CellMark")[1], "Again")

    def test_wps_complex_table_bookmark_topology(self):
        for variant in ("horizontal_merge", "vertical_merge", "multiple_paragraphs", "styled"):
            with self.subTest(variant=variant), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / f"{variant}.docx"
                document = Document()
                document.add_paragraph("Outside")
                table = document.add_table(rows=2, cols=2)
                if variant == "horizontal_merge":
                    table.cell(0, 0).merge(table.cell(0, 1))
                if variant == "vertical_merge":
                    table.cell(0, 0).merge(table.cell(1, 0))
                if variant == "styled":
                    table.style = "Light Shading Accent 1"
                cell = table.cell(0, 0)
                paragraph = cell.paragraphs[0]
                prefix = paragraph.add_run("Before ")
                old = paragraph.add_run("Old")
                suffix = paragraph.add_run(" After")
                if variant == "styled":
                    prefix.italic = True
                    old.bold = True
                    suffix.italic = True
                start = OxmlElement("w:bookmarkStart")
                start.set(qn("w:id"), "1")
                start.set(qn("w:name"), "CellMark")
                end = OxmlElement("w:bookmarkEnd")
                end.set(qn("w:id"), "1")
                old._r.addprevious(start)
                old._r.addnext(end)
                if variant == "multiple_paragraphs":
                    cell.add_paragraph("Second paragraph")
                if variant != "vertical_merge":
                    table.cell(1, 0).text = "Lower left"
                table.cell(1, 1).text = "Lower right"
                document.save(path)
                topology_before = docx_body_table_topology(path)
                ok, record, errors = register_document("writer", str(path), tmp)
                self.assertTrue(ok, errors)
                ok, result, errors, _ = writer_fill_bookmark(
                    record["document_id"], "CellMark", "New", f"complex-{variant}", workspace=tmp,
                )
                self.assertTrue(ok, (variant, errors, result))
                self.assertEqual(docx_body_table_topology(path), topology_before)
                readback = Document(path)
                self.assertEqual(readback.paragraphs[0].text, "Outside")
                self.assertEqual(readback.tables[0].cell(0, 0).paragraphs[0].text,
                                 "Before New After")
                if variant != "vertical_merge":
                    self.assertEqual(readback.tables[0].cell(1, 0).text, "Lower left")
                self.assertEqual(readback.tables[0].cell(1, 1).text, "Lower right")
                if variant == "multiple_paragraphs":
                    self.assertEqual(readback.tables[0].cell(0, 0).paragraphs[1].text,
                                     "Second paragraph")
                if variant == "styled":
                    self.assertEqual(readback.tables[0].style.name, "Light Shading Accent 1")
                    runs = readback.tables[0].cell(0, 0).paragraphs[0].runs
                    self.assertTrue(any(run.italic and "Before" in run.text for run in runs))
                    self.assertTrue(any(run.italic and "After" in run.text for run in runs))
                    self.assertTrue(any(run.bold and "New" in run.text for run in runs))

    def test_wps_table_cell_bookmark_fill_preserves_neighbors(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "table-bookmark.docx"
            document = Document()
            document.add_paragraph("Outside before")
            table = document.add_table(rows=1, cols=2)
            paragraph = table.cell(0, 0).paragraphs[0]
            paragraph.add_run("Before ")
            old = paragraph.add_run("Old")
            paragraph.add_run(" After")
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), "1")
            start.set(qn("w:name"), "CellMark")
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), "1")
            old._r.addprevious(start)
            old._r.addnext(end)
            table.cell(0, 1).text = "Neighbor"
            document.add_paragraph("Outside after")
            document.save(path)

            bookmark, original = read_supported_bookmark_text(path, "CellMark")
            self.assertEqual(original, "Old")
            self.assertTrue(bookmark["text_range_supported"])
            self.assertFalse(bookmark["body_paragraph_range_supported"])
            ok, record, errors = register_document("writer", str(path), tmp)
            self.assertTrue(ok, errors)
            ok, result, errors, replayed = writer_fill_bookmark(
                record["document_id"], "CellMark", "New", "table-public", workspace=tmp,
            )
            self.assertTrue(ok, (errors, result))
            self.assertFalse(replayed)
            self.assertTrue(result["readback_passed"])
            self.assertEqual(result["scope"], "table_cell")
            backup = result["backup"]
            after_bookmark, after_text = read_supported_bookmark_text(path, "CellMark")
            self.assertTrue(after_bookmark["text_range_supported"])
            self.assertEqual(after_text, "New")
            readback = Document(path)
            self.assertEqual([paragraph.text for paragraph in readback.paragraphs],
                             ["Outside before", "Outside after"])
            self.assertEqual(readback.tables[0].cell(0, 0).text, "Before New After")
            self.assertEqual(readback.tables[0].cell(0, 1).text, "Neighbor")
            self.assertEqual(file_identity(Path(backup["backup_path"]))["source_sha256"],
                             backup["source_identity"]["source_sha256"])
            saved_identity = file_identity(path)
            replay_ok, replay_result, replay_errors, replayed = writer_fill_bookmark(
                record["document_id"], "CellMark", "New", "table-public", workspace=tmp,
            )
            self.assertTrue(replay_ok, replay_errors)
            self.assertTrue(replayed)
            self.assertEqual(replay_result, result)
            self.assertEqual(file_identity(path), saved_identity)
            with chdir(tmp):
                output = io.StringIO()
                self.assertEqual(run([
                    "writer-fill-bookmark", "--document-id", record["document_id"],
                    "--bookmark-name", "CellMark", "--text", "New",
                    "--request-id", "table-public",
                ], output_stream=output), 0)
                cli = json.loads(output.getvalue())
                mcp_ok, mcp, mcp_errors = call_mcp_tool("wps_agent_writer_fill_bookmark", {
                    "document_id": record["document_id"], "bookmark_name": "CellMark",
                    "text": "New", "request_id": "table-public",
                })
            self.assertTrue(mcp_ok, mcp_errors)
            self.assertEqual(cli["data"], mcp["response"]["data"])
            self.assertTrue(cli["data"]["readback_passed"])
            self.assertEqual(file_identity(path), saved_identity)
            backed_up, _, backup_errors, _ = create_backup(
                record["document_id"], "table-post-fill-backup", workspace=tmp,
            )
            self.assertTrue(backed_up, backup_errors)


if __name__ == "__main__":
    unittest.main()
