from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
import json
import zipfile
import xml.etree.ElementTree as ET

from docx import Document
from PIL import Image

from wps_ai_agent_cli.html_roundtrip import build_html_roundtrip_mapping, export_controlled_html, import_controlled_html, verify_controlled_document


class HtmlRoundtripTests(unittest.TestCase):
    def test_owned_schema_produces_stable_identity_map_and_parent_links(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "owned.html"
            source.write_text(
                '<html data-wps-schema="wps-agent-html/v1"><body>'
                '<p data-wps-object-id="p:1">Hello <a data-wps-object-id="link.1" href="https://example.com">world</a></p>'
                '<table data-wps-object-id="table-1"><tr data-wps-object-id="row-1"><td data-wps-object-id="cell-1">value</td></tr></table>'
                '</body></html>',
                encoding="utf-8",
            )
            ok, first, errors = build_html_roundtrip_mapping(source)
            self.assertTrue(ok, errors)
            ok, second, errors = build_html_roundtrip_mapping(source)
            self.assertTrue(ok, errors)
            self.assertEqual(first["mappings"], second["mappings"])
            by_id = {item["object_id"]: item for item in first["mappings"]}
            self.assertEqual(by_id["link.1"]["parent_object_id"], "p:1")
            self.assertEqual(by_id["cell-1"]["parent_object_id"], "row-1")
            self.assertEqual(first["mapping_count"], 5)
            self.assertTrue(first["read_only"])

    def test_rejects_unowned_html_duplicate_ids_and_css(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "owned.html"
            source.write_text('<p data-wps-object-id="p1">Not owned</p>', encoding="utf-8")
            ok, _, errors = build_html_roundtrip_mapping(source)
            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "ROUNDTRIP_SCHEMA_REQUIRED")

            source.write_text(
                '<html data-wps-schema="wps-agent-html/v1"><p data-wps-object-id="p1">A</p>'
                '<p data-wps-object-id="p1">B</p><h1 data-wps-object-id="head" style="color:red">C</h1></html>',
                encoding="utf-8",
            )
            ok, _, errors = build_html_roundtrip_mapping(source)
            self.assertFalse(ok)
            self.assertIn("DUPLICATE_OBJECT_ID", {error["code"] for error in errors})
            self.assertIn("UNSUPPORTED_ROUNDTRIP_CSS", {error["code"] for error in errors})

    def test_requires_ids_for_supported_semantic_nodes_and_rejects_script(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "owned.html"
            source.write_text(
                '<html data-wps-schema="wps-agent-html/v1"><p data-wps-object-id="p1">Text</p>'
                '<table><tr><td>No ids</td></tr></table><script>run()</script><iframe src="x"></iframe></html>',
                encoding="utf-8",
            )
            ok, _, errors = build_html_roundtrip_mapping(source)
            self.assertFalse(ok)
            codes = {error["code"] for error in errors}
            self.assertIn("OBJECT_ID_REQUIRED", codes)
            self.assertIn("UNSUPPORTED_ROUNDTRIP_FEATURE", codes)

    def test_controlled_import_persists_bookmarks_and_verifies_content_edits(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source, document = root / "owned.html", root / "owned.docx"
            source.write_text(
                '<html data-wps-schema="wps-agent-html/v1"><p data-wps-object-id="p1">First</p>'
                '<p data-wps-object-id="p2">Second</p></html>', encoding="utf-8",
            )
            ok, data, errors = import_controlled_html(source, document)
            self.assertTrue(ok, errors)
            mapping_path = Path(data["mapping_path"])
            ok, result, errors = verify_controlled_document(document, mapping_path)
            self.assertTrue(ok, errors)
            self.assertFalse(result["document_edited_since_import"])
            loaded = Document(document)
            loaded.paragraphs[0].runs[0].text = "Edited content"
            loaded.save(document)
            ok, result, errors = verify_controlled_document(document, mapping_path)
            self.assertTrue(ok, errors)
            self.assertTrue(result["document_edited_since_import"])

    def test_identity_verifier_detects_removed_and_reordered_bookmarks(self):
        ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source, document = root / "owned.html", root / "owned.docx"
            source.write_text(
                '<html data-wps-schema="wps-agent-html/v1"><p data-wps-object-id="p1">First</p>'
                '<p data-wps-object-id="p2">Second</p></html>', encoding="utf-8",
            )
            ok, data, errors = import_controlled_html(source, document)
            self.assertTrue(ok, errors)
            mapping = json.loads(Path(data["mapping_path"]).read_text(encoding="utf-8"))
            entries: dict[str, bytes] = {}
            with zipfile.ZipFile(document) as archive:
                entries = {item.filename: archive.read(item.filename) for item in archive.infolist()}
            xml = ET.fromstring(entries["word/document.xml"])
            body = xml.find(f"{{{ns}}}body")
            paragraphs = [child for child in body if child.tag == f"{{{ns}}}p"]
            body.remove(paragraphs[0])
            body.remove(paragraphs[1])
            body.insert(0, paragraphs[1])
            body.insert(1, paragraphs[0])
            entries["word/document.xml"] = ET.tostring(xml, encoding="utf-8", xml_declaration=True)
            with zipfile.ZipFile(document, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                for name, payload in entries.items():
                    archive.writestr(name, payload)
            ok, _, errors = verify_controlled_document(document, data["mapping_path"])
            self.assertFalse(ok)
            self.assertIn("ROUNDTRIP_OBJECT_REORDERED", {error["code"] for error in errors})

            first_name = mapping["mappings"][0]["bookmark_name"]
            with zipfile.ZipFile(document) as archive:
                entries = {item.filename: archive.read(item.filename) for item in archive.infolist()}
            xml = ET.fromstring(entries["word/document.xml"])
            starts = [node for node in xml.iter(f"{{{ns}}}bookmarkStart") if node.get(f"{{{ns}}}name") == first_name]
            self.assertEqual(len(starts), 1)
            bookmark_id = starts[0].get(f"{{{ns}}}id")
            parents = {child: parent for parent in xml.iter() for child in parent}
            ends = [node for node in xml.iter(f"{{{ns}}}bookmarkEnd") if node.get(f"{{{ns}}}id") == bookmark_id]
            for node in [*starts, *ends]:
                parents[node].remove(node)
            entries["word/document.xml"] = ET.tostring(xml, encoding="utf-8", xml_declaration=True)
            with zipfile.ZipFile(document, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                for name, payload in entries.items():
                    archive.writestr(name, payload)
            ok, _, errors = verify_controlled_document(document, data["mapping_path"])
            self.assertFalse(ok)
            self.assertIn("ROUNDTRIP_OBJECT_MISSING", {error["code"] for error in errors})

    def test_export_restores_owned_schema_ids_after_text_edit(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source, document, output = root / "owned.html", root / "owned.docx", root / "roundtrip.html"
            source.write_text(
                '<html data-wps-schema="wps-agent-html/v1"><h1 data-wps-object-id="heading">Title</h1>'
                '<p data-wps-object-id="paragraph">Hello <a data-wps-object-id="link" href="https://example.com">site</a></p>'
                '<ol><li data-wps-object-id="item">Step</li></ol><table data-wps-object-id="table">'
                '<tr data-wps-object-id="row"><th data-wps-object-id="head">H</th><td data-wps-object-id="cell">C</td></tr></table></html>',
                encoding="utf-8",
            )
            ok, imported, errors = import_controlled_html(source, document)
            self.assertTrue(ok, errors)
            word = Document(document)
            word.paragraphs[1].runs[0].text = "Changed "
            word.save(document)
            ok, exported, errors = export_controlled_html(document, imported["mapping_path"], output)
            self.assertTrue(ok, errors)
            valid, after, errors = build_html_roundtrip_mapping(output)
            self.assertTrue(valid, errors)
            before_ids = [item["object_id"] for item in json.loads(Path(imported["mapping_path"]).read_text(encoding="utf-8"))["mappings"]]
            after_ids = [item["object_id"] for item in after["mappings"]]
            self.assertEqual(after_ids, before_ids)
            self.assertIn("Changed", output.read_text(encoding="utf-8"))
            self.assertEqual(exported["exported_object_count"], len(before_ids))

    def test_export_roundtrips_embedded_local_image_as_owned_data_uri(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            Image.new("RGB", (24, 16), (30, 80, 180)).save(root / "image.png")
            source, document, output = root / "owned.html", root / "owned.docx", root / "roundtrip.html"
            source.write_text(
                '<html data-wps-schema="wps-agent-html/v1"><p data-wps-object-id="p1">Before'
                '<img data-wps-object-id="image1" src="image.png" alt="Mark">After</p></html>',
                encoding="utf-8",
            )
            ok, imported, errors = import_controlled_html(source, document)
            self.assertTrue(ok, errors)
            ok, _, errors = export_controlled_html(document, imported["mapping_path"], output)
            self.assertTrue(ok, errors)
            valid, after, errors = build_html_roundtrip_mapping(output)
            self.assertTrue(valid, errors)
            by_id = {entry["object_id"]: entry for entry in after["mappings"]}
            self.assertEqual(by_id["image1"]["parent_object_id"], "p1")
            self.assertIn("data:image/png;base64,", output.read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
