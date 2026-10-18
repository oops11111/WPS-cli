"""The complex HTML fixtures convert without the losses P12-011 found (fixtures/html_complex/)."""

from __future__ import annotations

import datetime as dt
import os
import shutil
import tempfile
import unittest
import xml.etree.ElementTree as etree
import zipfile
from pathlib import Path
from unittest import mock

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from wps_ai_agent_cli.html_editable import convert_html_editable
from wps_ai_agent_cli.html_office import convert_html_to_xlsx
from wps_ai_agent_cli.html_slides import convert_html_to_pptx, parse_css

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures" / "html_complex"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

_NO_BROWSER = mock.patch.dict(os.environ, {"WPS_AGENT_NO_BROWSER": "1"})


def setUpModule():
    # These tests pin the layout-free path; tests/test_html_positioned.py covers the browser layout.
    _NO_BROWSER.start()


def tearDownModule():
    _NO_BROWSER.stop()


def _local(element) -> str:
    return element.tag.rsplit("}", 1)[-1]


class _Converted(unittest.TestCase):
    name = ""
    suffix = ""
    convert = None

    @classmethod
    def setUpClass(cls):
        cls._tmp = tempfile.TemporaryDirectory()
        work = Path(cls._tmp.name)
        for item in FIXTURES.iterdir():
            shutil.copyfile(item, work / item.name)
        cls.output = work / f"{cls.name}.{cls.suffix}"
        ok, cls.data, errors = cls.convert(work / f"{cls.name}.html", cls.output)
        assert ok, errors

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()


class ReportToDocxTests(_Converted):
    name, suffix, convert = "report", "docx", staticmethod(convert_html_editable)

    def setUp(self):
        self.document = Document(self.output)
        self.paragraphs = self.document.paragraphs

    def paragraph(self, start):
        return next(p for p in self.paragraphs if p.text.startswith(start))

    def test_inline_formatting_is_kept(self):
        summary = self.paragraph("本季度营业收入")
        runs = {run.text: run for run in summary.runs}
        self.assertTrue(runs["4,820 万元"].bold)
        self.assertTrue(runs["38.2%"].underline)
        self.assertTrue(runs["1.1"].font.strike)
        self.assertEqual(str(runs["12.5%"].font.color.rgb), "008000")
        self.assertEqual([run.font.subscript for run in summary.runs if run.text == "2"][:1], [True])
        self.assertIn(True, [run.font.superscript for run in summary.runs if run.text == "2"])
        meta = {run.text: run for run in self.paragraph("编制").runs}
        self.assertTrue(meta["v3"].bold)
        self.assertTrue(meta["内部"].italic)
        code = [run for run in self.paragraph("汇率风险").runs if "7.12" in run.text]
        self.assertEqual(code[0].font.name, "Consolas")

    def test_source_whitespace_is_collapsed_but_pre_keeps_its_lines(self):
        self.assertNotIn("\n", self.paragraph("本季度营业收入").text)
        self.assertEqual(self.paragraph("华东区新开门店").text, "华东区新开门店 14 家")
        self.assertEqual(self.paragraph("联系人").text, "联系人：finance@example.com · 电话 021-0000-0000")
        code = self.paragraph("SELECT")
        self.assertEqual(code.text.splitlines()[1], "  FROM sales")

    def test_nested_list_items_are_their_own_second_level_paragraphs(self):
        styles = [(p.style.name, p.text) for p in self.paragraphs if p.style.name.startswith("List")]
        self.assertEqual(
            styles[:5],
            [
                ("List Number", "华东区新开门店 14 家"),
                ("List Bullet 2", "上海 6 家"),
                ("List Bullet 2", "杭州 5 家"),
                ("List Bullet 2", "南京 3 家"),
                ("List Number", "完成 ERP 系统升级"),
            ],
        )

    def test_figure_image_caption_table_caption_and_details_are_kept(self):
        self.assertEqual(len(self.document.inline_shapes), 1)
        captions = [p.text for p in self.paragraphs if p.style.name == "Caption"]
        self.assertEqual(captions, ["表 1 分区域收入（万元）", "图 1 季度收入趋势"])
        self.assertTrue(self.paragraph("计算口径").runs)
        self.assertTrue(any(p.text == "收入按权责发生制确认。" for p in self.paragraphs))
        self.assertFalse(any("caption" in warning or "details" in warning for warning in self.data["warnings"]))

    def test_rowspan_and_colspan_become_merged_cells(self):
        table = self.document.tables[0]
        self.assertEqual(len(table.rows), 7)
        self.assertEqual([cell.text for cell in table.rows[0].cells], ["区域", "收入", "收入", "同比"])
        self.assertEqual([cell.text for cell in table.rows[1].cells], ["区域", "Q2", "Q3", "同比"])
        self.assertIs(table.cell(0, 0)._tc, table.cell(1, 0)._tc)  # one cell spanning two rows
        self.assertIs(table.cell(0, 1)._tc, table.cell(0, 2)._tc)  # one cell spanning two columns
        self.assertEqual(table.cell(2, 2).text, "2,050")
        self.assertTrue(table.cell(0, 0).paragraphs[0].runs[0].bold)

    def test_scripts_and_unsafe_links_stay_out(self):
        text = "\n".join(p.text for p in self.paragraphs)
        self.assertNotIn("changed", text)
        self.assertIn("Unsafe hyperlink target omitted.", self.data["warnings"])
        self.assertFalse(
            any("javascript:" in getattr(rel, "target_ref", "") for rel in self.document.part.rels.values())
        )


class BudgetToXlsxTests(_Converted):
    name, suffix, convert = "budget", "xlsx", staticmethod(convert_html_to_xlsx)

    def setUp(self):
        self.workbook = load_workbook(self.output)

    def test_merged_headers_and_the_numbers_under_them(self):
        sheet = self.workbook["部门费用预算"]
        self.assertEqual(sorted(str(r) for r in sheet.merged_cells.ranges), ["A1:A2", "B1:E1", "F1:F2"])
        self.assertEqual(
            [sheet[c].value for c in ("A1", "B1", "F1", "B2", "E2")], ["部门", "季度预算（元）", "全年", "Q1", "Q4"]
        )
        self.assertEqual([sheet[c].value for c in ("B3", "C3", "D4")], [120000, 135000.5, -5000])
        self.assertIsNone(sheet["C5"].value)
        self.assertTrue(sheet["A1"].font.bold)
        # Formulas are never evaluated from HTML.
        self.assertEqual(sheet["F3"].value, "=SUM(B3:E3)")
        self.assertEqual(sheet["F3"].data_type, "s")

    def test_written_numbers_dates_and_ids_get_their_types(self):
        sheet = self.workbook["人员编制"]
        self.assertEqual(sheet["A2"].value, "000127")
        self.assertEqual(sheet["C2"].value, dt.datetime(2021, 3, 15))
        self.assertEqual(sheet["C2"].number_format, "yyyy-mm-dd")
        self.assertEqual(sheet["C4"].value, dt.datetime(2025, 11, 20))
        self.assertEqual((sheet["D2"].value, sheet["D2"].number_format), (18500, "#,##0"))
        self.assertEqual((sheet["D3"].value, sheet["D3"].number_format), (15000, '"¥"#,##0'))
        self.assertEqual((sheet["E3"].value, sheet["E3"].number_format), (0.95, "0%"))
        self.assertEqual(sheet["D5"].value, 12000.0)
        self.assertEqual(self.workbook["一个名字很长很长很长很长的汇总表 包含 非法字符 以及 方括号"]["B3"].value, "007")

    def test_sheet_names_survive_a_save_in_wps(self):
        # P12-012: WPS replaced the full-width colon with "_" when it saved the workbook.
        for name in self.workbook.sheetnames:
            self.assertFalse(set(name) & set("[]:*?/\\：？／＊［］＼"), name)

    def test_line_breaks_links_nested_tables_and_images(self):
        staff = self.workbook["人员编制"]
        self.assertEqual(staff["B5"].value, "王\n小明")
        self.assertTrue(staff["B5"].alignment.wrap_text)
        self.assertEqual(staff["F5"].hyperlink.target, "https://example.com/wang")
        self.assertEqual(self.workbook["部门费用预算 2"]["A2"].value, "差旅\n嵌套")
        self.assertTrue(any("nested in a cell" in warning for warning in self.data["warnings"]))
        self.assertEqual(self.data["images"], 1)
        self.assertEqual(len(self.workbook.worksheets[-1]._images), 1)


class LaunchToPptxTests(_Converted):
    name, suffix, convert = "launch", "pptx", staticmethod(convert_html_to_pptx)

    def setUp(self):
        self.slides = list(Presentation(self.output).slides)

    def slide_xml(self, number):
        with zipfile.ZipFile(self.output) as archive:
            return etree.fromstring(archive.read(f"ppt/slides/slide{number}.xml"))

    def test_six_slides_with_their_titles(self):
        # A body moved to make room keeps its full width (a partly set placeholder position once zeroed it).
        for slide in self.slides:
            for shape in slide.placeholders:
                self.assertGreater(shape.width, 0)
        titles = [slide.shapes.title.text for slide in self.slides]
        self.assertEqual(titles, ["星河 X1 新品发布会", "为什么是 X1", "规格对比", "销量预测", "价格与上市", "谢谢"])

    def test_tables_images_and_notes_are_converted(self):
        table = next(shape for shape in self.slides[2].shapes if shape.has_table).table
        self.assertEqual([table.cell(1, c).text for c in range(3)], ["续航", "14 小时", "10 小时"])
        self.assertTrue(any(shape.shape_type == 13 for shape in self.slides[3].shapes))  # a picture
        self.assertEqual(self.slides[1].notes_slide.notes_text_frame.text, "演讲者备注：这里停顿，等掌声。")
        # The notes are not body text, and a slide with only a table has no empty body placeholder.
        body = [p.text for p in self.slides[1].placeholders[1].text_frame.paragraphs]
        self.assertEqual(body, ["续航提升 40%", "重量仅 1.1 kg", "支持 Wi-Fi 7"])
        self.assertEqual([shape.has_text_frame for shape in self.slides[2].shapes], [True, False])
        self.assertTrue(self.slides[1].placeholders[1].text_frame.paragraphs[0].runs[1].font.bold)

    def test_css_animations_become_entrance_effects_with_their_delays(self):
        effects = {}
        for number in (1, 2, 5):
            xml = self.slide_xml(number)
            self.assertEqual(len(xml.findall(f"{P}timing")), 1)
            effects[number] = [
                (node.get("presetID"), node.get("nodeType"), node.find(f"{P}stCondLst/{P}cond").get("delay"))
                for node in xml.iter(f"{P}cTn")
                if node.get("presetClass") == "entr"
            ]
        # Slide 1: the title flies in from the left, the date fades in after 0.5 s.
        self.assertEqual(effects[1], [("2", "withEffect", "0"), ("10", "withEffect", "500")])
        # Slide 2: the three list items fly in one after another.
        self.assertEqual(
            effects[2], [("2", "withEffect", "200"), ("2", "withEffect", "400"), ("2", "withEffect", "600")]
        )
        paragraphs = [rng.get("st") for rng in self.slide_xml(2).iter(f"{P}pRg")]
        self.assertEqual(sorted(set(paragraphs)), ["0", "1", "2"])
        builds = self.slide_xml(2).find(f"{P}timing/{P}bldLst/{P}bldP")
        self.assertEqual(builds.get("build"), "p")
        # Slide 5: the animated span fades its paragraph in.
        self.assertEqual(effects[5], [("10", "withEffect", "500")])
        self.assertTrue(any("part of a paragraph" in warning for warning in self.data["warnings"]))

    def test_transitions_follow_data_transition_and_css(self):
        kinds = []
        for number in range(1, 7):
            transition = self.slide_xml(number).find(f"{P}transition")
            kinds.append(_local(transition[0]) if transition is not None else None)
        self.assertEqual(kinds, ["fade", "push", "fade", "fade", "fade", "fade"])
        order = [_local(child) for child in self.slide_xml(1)]
        self.assertLess(order.index("transition"), order.index("timing"))

    def test_what_cannot_be_converted_is_reported(self):
        warnings = "\n".join(self.data["warnings"])
        for expected in ("Scripts", "<svg>", "Hover effects", "teaser.mp4"):
            self.assertIn(expected, warnings)
        self.assertEqual(
            self.data["mapped_objects"],
            {"tables": 1, "images": 1, "media": 0, "notes": 1, "animations": 6, "transitions": 6},
        )


class SlideFeatureTests(unittest.TestCase):
    def convert(self, html, files=()):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        work = Path(tmp.name)
        for name, payload in files:
            (work / name).write_bytes(payload)
        (work / "deck.html").write_text(html, encoding="utf-8")
        ok, data, errors = convert_html_to_pptx(work / "deck.html", work / "deck.pptx")
        self.assertTrue(ok, errors)
        return work / "deck.pptx", data

    def test_keyframes_are_classified_by_what_they_do(self):
        css = parse_css(
            "<style>@keyframes a{from{transform:translateX(50px)}to{transform:none}}"
            "@keyframes b{0%{transform:translateY(-1em)}100%{}}@keyframes c{from{transform:scale(.2)}}"
            "@keyframes d{from{opacity:0}} .x:hover{color:red}"
            "@keyframes e{from{opacity:1}to{opacity:0}}@keyframes f{from{transform:none}to{transform:translateX(-200px)}}"
            "@keyframes g{0%{transform:scale(1)}50%{transform:scale(1.2)}100%{transform:scale(1)}}"
            "@keyframes h{from{transform:rotate(0)}to{transform:rotate(360deg)}}</style>"
        )
        self.assertEqual(
            css.keyframes,
            {
                "a": ("entrance", "fly", "right"),
                "b": ("entrance", "fly", "top"),
                "c": ("entrance", "zoom", ""),
                "d": ("entrance", "fade", ""),
                "e": ("exit", "fade", ""),
                "f": ("exit", "fly", "left"),
                "g": ("emphasis", "pulse", ""),
                "h": ("emphasis", "spin", ""),
            },
        )
        self.assertTrue(css.hover)

    def test_video_shares_the_slide_timing_with_the_animations(self):
        html = (
            "<style>@keyframes z{from{transform:scale(0)}} .z{animation:z 1s 250ms}</style>"
            '<section data-transition="wipe"><h2 class="z">Clip</h2><video src="clip.mp4"></video></section>'
            '<section data-transition="spiral"><h2>Next</h2><ul><li>a<ul><li>b</li></ul></li></ul></section>'
        )
        output, data = self.convert(html, [("clip.mp4", b"\0" * 2000)])
        self.assertEqual(data["mapped_objects"]["media"], 1)
        self.assertIn("Unknown slide transition 'spiral' was omitted.", data["warnings"])
        with zipfile.ZipFile(output) as archive:
            xml = etree.fromstring(archive.read("ppt/slides/slide1.xml"))
        timings = xml.findall(f"{P}timing")
        self.assertEqual(len(timings), 1)
        ids = [node.get("id") for node in timings[0].iter(f"{P}cTn")]
        self.assertEqual(len(ids), len(set(ids)))
        self.assertTrue(any(node.get("presetID") == "53" for node in timings[0].iter(f"{P}cTn")))
        self.assertEqual(_local(xml.find(f"{P}transition")[0]), "wipe")
        levels = [(p.level, p.text) for p in Presentation(output).slides[1].placeholders[1].text_frame.paragraphs]
        self.assertEqual(levels, [(0, "a"), (1, "b")])


if __name__ == "__main__":
    unittest.main()
