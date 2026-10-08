import os
import io
import json
import shutil
import tempfile
import unittest
from contextlib import chdir
from pathlib import Path
from unittest.mock import patch

from tests.test_writer_structure import write_document
from wps_ai_agent_cli.document_text import docx_body_paragraph_target, docx_body_paragraphs, docx_body_tables
from wps_ai_agent_cli.cli import run
from wps_ai_agent_cli.mcp_adapter import call_mcp_tool
from wps_ai_agent_cli.sessions import register_document
from wps_ai_agent_cli.writer_ops import writer_replace
from wps_ai_agent_cli.writer_ops import writer_fill_bookmark


class WriterParagraphMappingTests(unittest.TestCase):
    def test_direct_body_target_excludes_table_paragraphs(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "mapping.docx"
            write_document(path, '''<w:p><w:r><w:t>before</w:t></w:r></w:p>
                <w:tbl><w:tr><w:tc><w:p/><w:p/></w:tc></w:tr></w:tbl>
                <w:p><w:r><w:t>target</w:t></w:r></w:p>''')
            self.assertEqual(docx_body_paragraph_target(path, 2), {
                "body_paragraph_count": 2, "expected_text": "target",
            })
            with self.assertRaises(ValueError):
                docx_body_paragraph_target(path, 3)

    def test_bookmark_fill_dry_run_and_preflight_rejections(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bookmark.docx"
            write_document(path, '<w:p><w:r><w:t>Dear </w:t></w:r><w:bookmarkStart w:id="1" w:name="Client"/><w:r><w:t>old</w:t></w:r><w:bookmarkEnd w:id="1"/><w:r><w:t>!</w:t></w:r></w:p>')
            _, record, _ = register_document("writer", str(path), tmp)
            before = path.read_bytes()
            with patch("wps_ai_agent_cli.writer_ops.create_backup") as backup, patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com") as com:
                ok, result, errors, replayed = writer_fill_bookmark(record["document_id"], "Client", "new", "bookmark-preview", dry_run=True, workspace=tmp)
            self.assertTrue(ok)
            self.assertFalse(replayed)
            self.assertEqual(errors, [])
            self.assertEqual(result["current_text_length"], 3)
            self.assertEqual(result["replacement_text_length"], 3)
            backup.assert_not_called()
            com.assert_not_called()
            self.assertEqual(path.read_bytes(), before)
            ok, _, errors, _ = writer_fill_bookmark(record["document_id"], "Missing", "x", "bookmark-missing", dry_run=True, workspace=tmp)
            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "BOOKMARK_NOT_FOUND")

    def test_replaced_file_blocks_bookmark_fill_before_wps(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bookmark.docx"
            replacement = Path(tmp) / "replacement.docx"
            write_document(path, '<w:p><w:bookmarkStart w:id="1" w:name="Client"/>'
                                 '<w:r><w:t>old</w:t></w:r><w:bookmarkEnd w:id="1"/></w:p>')
            _, record, _ = register_document("writer", str(path), tmp)
            replacement.write_bytes(path.read_bytes())
            os.replace(replacement, path)
            with patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com") as com:
                ok, _, errors, _ = writer_fill_bookmark(
                    record["document_id"], "Client", "new", "replaced-file", workspace=tmp,
                )
            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "DOCUMENT_IDENTITY_CHANGED")
            com.assert_not_called()
            _, rebound, _ = register_document("writer", str(path), tmp)
            self.assertEqual(rebound["document_id"], record["document_id"])
            ok, _, errors, _ = writer_fill_bookmark(
                rebound["document_id"], "Client", "new", "rebound-preview", dry_run=True, workspace=tmp,
            )
            self.assertTrue(ok, errors)
            ok, _, errors, _ = writer_fill_bookmark(record["document_id"], "_GoBack", "x", "bookmark-reserved", dry_run=True, workspace=tmp)
            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "INVALID_ARGUMENT")

    def test_bookmark_fill_commit_readback_and_idempotency(self):
        from zipfile import ZipFile

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bookmark.docx"
            write_document(path, '<w:p><w:r><w:t>Dear </w:t></w:r><w:bookmarkStart w:id="1" w:name="Client"/><w:r><w:t>old</w:t></w:r><w:bookmarkEnd w:id="1"/><w:r><w:t>!</w:t></w:r></w:p>')
            _, record, _ = register_document("writer", str(path), tmp)

            def replace_xml(*args):
                with ZipFile(path) as archive:
                    parts = {name: archive.read(name) for name in archive.namelist()}
                parts["word/document.xml"] = parts["word/document.xml"].replace(b">old<", b">Northwind<")
                with ZipFile(path, "w") as archive:
                    for name, content in parts.items():
                        archive.writestr(name, content)
                return {"ok": True, "errors": [], "data": {"backend": "test"}}

            with patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com", side_effect=replace_xml):
                ok, result, errors, replayed = writer_fill_bookmark(record["document_id"], "Client", "Northwind", "bookmark-commit", workspace=tmp)
                self.assertTrue(ok, errors)
                self.assertFalse(replayed)
                self.assertTrue(result["readback_passed"])
                ok, repeated, errors, replayed = writer_fill_bookmark(record["document_id"], "Client", "Northwind", "bookmark-commit", workspace=tmp)
                self.assertTrue(ok)
                self.assertTrue(replayed)
                self.assertEqual(errors, [])
                self.assertTrue(repeated["readback_passed"])
                ok, _result, errors, replayed = writer_fill_bookmark(record["document_id"], "Client", "different", "bookmark-commit", workspace=tmp)
            self.assertFalse(ok)
            self.assertFalse(replayed)
            self.assertEqual(errors[0]["code"], "IDEMPOTENCY_CONFLICT")

    def test_table_bookmark_fill_rejects_collateral_cell_change(self):
        from zipfile import ZipFile
        from wps_ai_agent_cli.operations import get_operation

        for corrupt_neighbor in (False, True):
            with self.subTest(corrupt_neighbor=corrupt_neighbor), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / "table-bookmark.docx"
                write_document(path, '''
                    <w:p><w:r><w:t>Outside</w:t></w:r></w:p>
                    <w:tbl><w:tr>
                      <w:tc><w:p><w:r><w:t>Before </w:t></w:r>
                        <w:bookmarkStart w:id="1" w:name="CellMark"/>
                        <w:r><w:t>Old</w:t></w:r><w:bookmarkEnd w:id="1"/>
                        <w:r><w:t> After</w:t></w:r></w:p></w:tc>
                      <w:tc><w:p><w:r><w:t>Neighbor</w:t></w:r></w:p></w:tc>
                    </w:tr></w:tbl>''')
                _, record, _ = register_document("writer", str(path), tmp)

                def replace_xml(*_args):
                    with ZipFile(path) as archive:
                        parts = {name: archive.read(name) for name in archive.namelist()}
                    payload = parts["word/document.xml"].replace(b">Old<", b">New<")
                    if corrupt_neighbor:
                        payload = payload.replace(b">Neighbor<", b">Corrupt<")
                    parts["word/document.xml"] = payload
                    with ZipFile(path, "w") as archive:
                        for name, content in parts.items():
                            archive.writestr(name, content)
                    return {"ok": True, "errors": [], "data": {"backend": "test"}}

                with patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com", side_effect=replace_xml):
                    ok, result, errors, _ = writer_fill_bookmark(
                        record["document_id"], "CellMark", "New", "table-commit", workspace=tmp,
                    )
                self.assertEqual(ok, not corrupt_neighbor)
                self.assertEqual(result["readback_passed"], not corrupt_neighbor)
                if corrupt_neighbor:
                    self.assertEqual(errors[0]["code"], "VALIDATION_FAILED")
                    self.assertIsNone(get_operation("table-commit", tmp))
                else:
                    self.assertEqual(errors, [])
                    self.assertIsNotNone(get_operation("table-commit", tmp))

    def test_table_bookmark_fill_rejects_changed_merge_topology(self):
        from zipfile import ZipFile

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "table-topology.docx"
            write_document(path, '<w:tbl><w:tr><w:tc><w:p>'
                                 '<w:bookmarkStart w:id="1" w:name="CellMark"/>'
                                 '<w:r><w:t>Old</w:t></w:r><w:bookmarkEnd w:id="1"/>'
                                 '</w:p></w:tc><w:tc><w:p><w:r><w:t>Neighbor</w:t></w:r>'
                                 '</w:p></w:tc></w:tr></w:tbl>')
            _, record, _ = register_document("writer", str(path), tmp)

            def alter_topology(*_args):
                with ZipFile(path) as archive:
                    parts = {name: archive.read(name) for name in archive.namelist()}
                payload = parts["word/document.xml"].replace(b">Old<", b">New<")
                parts["word/document.xml"] = payload.replace(
                    b"<w:tc>", b'<w:tc><w:tcPr><w:gridSpan w:val="2"/></w:tcPr>', 1,
                )
                with ZipFile(path, "w") as archive:
                    for name, content in parts.items():
                        archive.writestr(name, content)
                return {"ok": True, "errors": [], "data": {"backend": "test"}}

            with patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com", side_effect=alter_topology):
                ok, result, errors, _ = writer_fill_bookmark(
                    record["document_id"], "CellMark", "New", "topology-changed", workspace=tmp,
                )
            self.assertFalse(ok)
            self.assertFalse(result["readback_passed"])
            self.assertEqual(errors[0]["code"], "VALIDATION_FAILED")

    def test_table_bookmark_rejects_link_or_field_target_before_backup(self):
        semantic_ranges = (
            '<w:hyperlink w:anchor="Target"><w:r><w:t>Old</w:t></w:r></w:hyperlink>',
            '<w:fldSimple w:instr="MERGEFIELD Code"><w:r><w:t>Old</w:t></w:r></w:fldSimple>',
            '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
            '<w:r><w:instrText>MERGEFIELD Code</w:instrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
            '<w:r><w:t>Old</w:t></w:r>'
            '<w:r><w:fldChar w:fldCharType="end"/></w:r>',
        )
        for semantic_range in semantic_ranges:
            with self.subTest(semantic_range=semantic_range), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / "semantic-target.docx"
                write_document(path, '<w:tbl><w:tr><w:tc><w:p>'
                                     '<w:bookmarkStart w:id="1" w:name="CellMark"/>'
                                     f'{semantic_range}'
                                     '<w:bookmarkEnd w:id="1"/>'
                                     '</w:p></w:tc></w:tr></w:tbl>')
                _, record, _ = register_document("writer", str(path), tmp)
                with patch("wps_ai_agent_cli.writer_ops.create_backup") as backup, \
                        patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com") as com:
                    ok, _, errors, _ = writer_fill_bookmark(
                        record["document_id"], "CellMark", "New", "semantic-target", workspace=tmp,
                    )
                self.assertFalse(ok)
                self.assertEqual(errors[0]["code"], "BOOKMARK_SCOPE_UNSUPPORTED")
                backup.assert_not_called()
                com.assert_not_called()

    def test_table_bookmark_rejects_unresolved_or_malformed_semantics_before_backup(self):
        malformed_nodes = (
            '<w:hyperlink xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
            ' r:id="rMissing"><w:r><w:t>Link</w:t></w:r></w:hyperlink>',
            '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
            '<w:r><w:instrText>MERGEFIELD Code</w:instrText></w:r>',
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r>',
            '<w:fldSimple><w:r><w:t>Value</w:t></w:r></w:fldSimple>',
            '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
            '<w:r><w:instrText>QUOTE</w:instrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
            '<w:r><w:t>outer</w:t></w:r>'
            '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
            '<w:r><w:instrText>MERGEFIELD Code</w:instrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
            '<w:r><w:t>inner</w:t></w:r>'
            '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
            '<w:r><w:fldChar w:fldCharType="end"/></w:r>',
        )
        for node in malformed_nodes:
            with self.subTest(node=node), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / "malformed-semantic.docx"
                write_document(path, '<w:tbl><w:tr><w:tc><w:p>'
                                     '<w:bookmarkStart w:id="1" w:name="CellMark"/>'
                                     '<w:r><w:t>Old</w:t></w:r><w:bookmarkEnd w:id="1"/>'
                                     f'{node}</w:p></w:tc></w:tr></w:tbl>')
                _, record, _ = register_document("writer", str(path), tmp)
                with patch("wps_ai_agent_cli.writer_ops.create_backup") as backup, \
                        patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com") as com:
                    ok, _, errors, _ = writer_fill_bookmark(
                        record["document_id"], "CellMark", "New", "malformed-semantic", workspace=tmp,
                    )
                self.assertFalse(ok)
                self.assertEqual(errors[0]["code"], "BOOKMARK_SCOPE_UNSUPPORTED")
                backup.assert_not_called()
                com.assert_not_called()

    def test_table_bookmark_rejects_changed_semantic_neighbor(self):
        from zipfile import ZipFile
        from wps_ai_agent_cli.operations import get_operation

        for changed in ("hyperlink", "field"):
            with self.subTest(changed=changed), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / "semantic-neighbor.docx"
                write_document(path, '<w:tbl><w:tr><w:tc><w:p>'
                                     '<w:hyperlink w:anchor="Safe"><w:r><w:t>Link</w:t></w:r></w:hyperlink>'
                                     '<w:bookmarkStart w:id="1" w:name="CellMark"/>'
                                     '<w:r><w:t>Old</w:t></w:r><w:bookmarkEnd w:id="1"/>'
                                     '<w:fldSimple w:instr="MERGEFIELD Code"><w:r><w:t>Code</w:t></w:r></w:fldSimple>'
                                     '</w:p></w:tc></w:tr></w:tbl>')
                _, record, _ = register_document("writer", str(path), tmp)

                def change_xml(*_args):
                    with ZipFile(path) as archive:
                        parts = {name: archive.read(name) for name in archive.namelist()}
                    payload = parts["word/document.xml"].replace(b">Old<", b">New<")
                    if changed == "hyperlink":
                        payload = payload.replace(b'w:anchor="Safe"', b'w:anchor="Other"')
                    else:
                        payload = payload.replace(b'MERGEFIELD Code', b'MERGEFIELD Other')
                    parts["word/document.xml"] = payload
                    with ZipFile(path, "w") as archive:
                        for name, content in parts.items():
                            archive.writestr(name, content)
                    return {"ok": True, "errors": [], "data": {"backend": "test"}}

                with patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com", side_effect=change_xml):
                    ok, result, errors, _ = writer_fill_bookmark(
                        record["document_id"], "CellMark", "New", "semantic-neighbor", workspace=tmp,
                    )
                self.assertFalse(ok)
                self.assertFalse(result["readback_passed"])
                self.assertEqual(errors[0]["code"], "VALIDATION_FAILED")
                self.assertIsNone(get_operation("semantic-neighbor", tmp))

    def test_table_bookmark_rejects_changed_hyperlink_relationship_target(self):
        from zipfile import ZipFile

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "relationship-neighbor.docx"
            write_document(path, '<w:tbl><w:tr><w:tc><w:p>'
                                 '<w:hyperlink xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
                                 ' r:id="rIdLink"><w:r><w:t>Link</w:t></w:r></w:hyperlink>'
                                 '<w:bookmarkStart w:id="1" w:name="CellMark"/>'
                                 '<w:r><w:t>Old</w:t></w:r><w:bookmarkEnd w:id="1"/>'
                                 '</w:p></w:tc></w:tr></w:tbl>')
            with ZipFile(path, "a") as archive:
                archive.writestr("word/_rels/document.xml.rels", '''
                    <Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
                      <Relationship Id="rIdLink" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
                        Target="https://example.com/original" TargetMode="External"/>
                    </Relationships>''')
            _, record, _ = register_document("writer", str(path), tmp)

            def change_target(*_args):
                with ZipFile(path) as archive:
                    parts = {name: archive.read(name) for name in archive.namelist()}
                parts["word/document.xml"] = parts["word/document.xml"].replace(b">Old<", b">New<")
                parts["word/_rels/document.xml.rels"] = parts["word/_rels/document.xml.rels"].replace(
                    b"https://example.com/original", b"https://example.com/changed",
                )
                with ZipFile(path, "w") as archive:
                    for name, content in parts.items():
                        archive.writestr(name, content)
                return {"ok": True, "errors": [], "data": {"backend": "test"}}

            with patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com", side_effect=change_target):
                ok, result, errors, _ = writer_fill_bookmark(
                    record["document_id"], "CellMark", "New", "relationship-changed", workspace=tmp,
                )
            self.assertFalse(ok)
            self.assertFalse(result["readback_passed"])
            self.assertEqual(errors[0]["code"], "VALIDATION_FAILED")

    def test_table_bookmark_rejects_drawing_overlap_before_backup(self):
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "drawing-target.docx"
            write_image_bookmark_document(path, bookmark_around_drawing=True)
            _, record, _ = register_document("writer", str(path), tmp)
            with patch("wps_ai_agent_cli.writer_ops.create_backup") as backup, \
                    patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com") as com:
                ok, _, errors, _ = writer_fill_bookmark(
                    record["document_id"], "ImageMark", "", "drawing-target", workspace=tmp,
                )
            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "BOOKMARK_SCOPE_UNSUPPORTED")
            backup.assert_not_called()
            com.assert_not_called()

    def test_drawing_wrap_text_acceptance_matrix(self):
        import xml.etree.ElementTree as ET
        from docx.oxml.ns import qn
        from zipfile import ZipFile
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document
        from wps_ai_agent_cli.document_text import docx_body_drawing_semantics

        wrap_tags = {
            "none": "wrapNone", "square": "wrapSquare", "tight": "wrapTight",
            "through": "wrapThrough", "top_and_bottom": "wrapTopAndBottom",
        }
        values = (None, "bothSides", "left", "right", "largest", "invalid-enum")
        for mode, tag_name in wrap_tags.items():
            for value in values:
                with self.subTest(wrap=mode, value=value), tempfile.TemporaryDirectory() as tmp:
                    path = Path(tmp) / "wrap-matrix.docx"
                    write_image_bookmark_document(
                        path, anchored=True, anchor_wrap=mode, anchor_wrap_text=None,
                    )
                    with ZipFile(path) as source:
                        parts = {name: source.read(name) for name in source.namelist()}
                    root = ET.fromstring(parts["word/document.xml"])
                    wrap = next(root.iter(qn(f"wp:{tag_name}")))
                    if value is not None:
                        wrap.set("wrapText", value)
                    parts["word/document.xml"] = ET.tostring(
                        root, encoding="utf-8", xml_declaration=True,
                    )
                    with ZipFile(path, "w") as target:
                        for name, content in parts.items():
                            target.writestr(name, content)

                    snapshot = docx_body_drawing_semantics(path)
                    reasons = {issue["reason"] for issue in snapshot["unresolved_drawings"]}
                    if value == "invalid-enum":
                        self.assertIn("invalid_wrap_text", reasons)
                    elif mode == "top_and_bottom" and value in {"left", "right", "largest"}:
                        self.assertIn("unsupported_top_and_bottom_wrap_text", reasons)
                    else:
                        self.assertEqual(snapshot["unresolved_drawings"], [])

                    wrap_snapshot = snapshot["drawings"][0]["structure"]["children"][4]
                    actual = wrap_snapshot["attributes"].get("wrapText")
                    expected = (
                        "bothSides" if value is None and mode != "none" else value
                    )
                    self.assertEqual(actual, expected)

    def test_drawing_snapshot_normalizes_default_top_and_bottom_wrap(self):
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document
        from wps_ai_agent_cli.document_text import docx_body_drawing_semantics

        with tempfile.TemporaryDirectory() as tmp:
            explicit = Path(tmp) / "explicit-wrap.docx"
            omitted = Path(tmp) / "omitted-wrap.docx"
            write_image_bookmark_document(explicit, anchored=True, anchor_wrap="top_and_bottom")
            write_image_bookmark_document(
                omitted, anchored=True, anchor_wrap="top_and_bottom", anchor_wrap_text=None,
            )
            self.assertEqual(docx_body_drawing_semantics(explicit), docx_body_drawing_semantics(omitted))

    def test_drawing_snapshot_normalizes_default_square_wrap_text(self):
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document
        from wps_ai_agent_cli.document_text import docx_body_drawing_semantics

        with tempfile.TemporaryDirectory() as tmp:
            explicit = Path(tmp) / "explicit-square-wrap.docx"
            omitted = Path(tmp) / "omitted-square-wrap.docx"
            write_image_bookmark_document(explicit, anchored=True, anchor_wrap="square")
            write_image_bookmark_document(
                omitted, anchored=True, anchor_wrap="square", anchor_wrap_text=None,
            )
            self.assertEqual(docx_body_drawing_semantics(explicit), docx_body_drawing_semantics(omitted))

    def test_drawing_snapshot_retains_wrap_none_wrap_text_values(self):
        import xml.etree.ElementTree as ET
        from docx.oxml.ns import qn
        from zipfile import ZipFile
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document
        from wps_ai_agent_cli.document_text import docx_body_drawing_semantics

        with tempfile.TemporaryDirectory() as tmp:
            snapshots = {}
            for wrap_text in ("left", "right", "largest"):
                path = Path(tmp) / f"wrap-none-{wrap_text}.docx"
                write_image_bookmark_document(
                    path, anchored=True, anchor_wrap="none", anchor_wrap_text=wrap_text,
                )
                with ZipFile(path) as source:
                    parts = {name: source.read(name) for name in source.namelist()}
                root = ET.fromstring(parts["word/document.xml"])
                next(root.iter(qn("wp:wrapNone"))).set("wrapText", wrap_text)
                parts["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                with ZipFile(path, "w") as target:
                    for name, content in parts.items():
                        target.writestr(name, content)
                snapshot = docx_body_drawing_semantics(path)
                self.assertEqual(snapshot["unresolved_drawings"], [])
                wrap = snapshot["drawings"][0]["structure"]["children"][4]
                self.assertEqual(wrap["attributes"]["wrapText"], wrap_text)
                snapshots[wrap_text] = snapshot

            self.assertNotEqual(snapshots["left"], snapshots["right"])
            self.assertNotEqual(snapshots["right"], snapshots["largest"])

    def test_drawing_snapshot_normalizes_integer_wrap_coordinate_lexical_forms(self):
        import xml.etree.ElementTree as ET
        from zipfile import ZipFile
        from docx.oxml.ns import qn
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document
        from wps_ai_agent_cli.document_text import docx_body_drawing_semantics

        with tempfile.TemporaryDirectory() as tmp:
            canonical = Path(tmp) / "canonical.docx"
            equivalent = Path(tmp) / "equivalent.docx"
            write_image_bookmark_document(canonical, anchored=True, anchor_wrap="tight")
            with ZipFile(canonical) as source:
                parts = {name: source.read(name) for name in source.namelist()}
            root = ET.fromstring(parts["word/document.xml"])
            start = next(root.iter(qn("wp:start")))
            start.set("x", "\t+10800\n")
            start.set("y", "\r -00\t")
            anchor = next(root.iter(qn("wp:anchor")))
            anchor.set("distL", " +0114300 ")
            parts["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            with ZipFile(equivalent, "w") as target:
                for name, content in parts.items():
                    target.writestr(name, content)
            self.assertEqual(
                docx_body_drawing_semantics(canonical),
                docx_body_drawing_semantics(equivalent),
            )
            max_distance = Path(tmp) / "max-distance.docx"
            write_image_bookmark_document(max_distance, anchored=True, anchor_wrap="square")
            with ZipFile(max_distance) as source:
                max_parts = {name: source.read(name) for name in source.namelist()}
            max_root = ET.fromstring(max_parts["word/document.xml"])
            next(max_root.iter(qn("wp:anchor"))).set("distL", "2147483640")
            max_parts["word/document.xml"] = ET.tostring(max_root, encoding="utf-8", xml_declaration=True)
            with ZipFile(max_distance, "w") as target:
                for name, content in max_parts.items():
                    target.writestr(name, content)
            max_semantics = docx_body_drawing_semantics(max_distance)
            self.assertEqual(max_semantics["unresolved_drawings"], [])
            self.assertEqual(max_semantics["drawings"][0]["wrap_distances"]["distL"], "2147483640")

            omitted = Path(tmp) / "omitted-distances.docx"
            explicit_zero = Path(tmp) / "explicit-zero-distances.docx"
            for candidate in (omitted, explicit_zero):
                write_image_bookmark_document(candidate, anchored=True, anchor_wrap="square")
                with ZipFile(candidate) as source:
                    candidate_parts = {name: source.read(name) for name in source.namelist()}
                candidate_root = ET.fromstring(candidate_parts["word/document.xml"])
                candidate_anchor = next(candidate_root.iter(qn("wp:anchor")))
                for name in ("distT", "distB", "distL", "distR"):
                    if candidate == omitted:
                        candidate_anchor.attrib.pop(name, None)
                    else:
                        candidate_anchor.set(name, "0")
                candidate_parts["word/document.xml"] = ET.tostring(
                    candidate_root, encoding="utf-8", xml_declaration=True,
                )
                with ZipFile(candidate, "w") as target:
                    for name, content in candidate_parts.items():
                        target.writestr(name, content)
            omitted_semantics = docx_body_drawing_semantics(omitted)
            zero_semantics = docx_body_drawing_semantics(explicit_zero)
            self.assertEqual(omitted_semantics["drawings"][0]["wrap_distances"], {
                "distT": "0", "distB": "0", "distL": "114300", "distR": "114300",
            })
            self.assertEqual(zero_semantics["drawings"][0]["wrap_distances"], {
                name: "0" for name in ("distT", "distB", "distL", "distR")
            })
            self.assertNotEqual(omitted_semantics, zero_semantics)

    def test_table_bookmark_rejects_invalid_wrap_geometry_before_backup(self):
        import xml.etree.ElementTree as ET
        from docx.oxml.ns import qn
        from zipfile import ZipFile
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document
        from wps_ai_agent_cli.document_text import docx_body_drawing_semantics

        invalid_variants = (
            "missing_polygon", "too_few_points", "coordinate_out_of_range",
            "coordinate_decimal", "coordinate_empty", "coordinate_huge_integer",
            "coordinate_unicode_digits", "coordinate_nbsp", "coordinate_repeated_sign",
            "coordinate_sign_only", "coordinate_embedded_whitespace", "distance_unicode_digits",
            "distance_decimal", "distance_empty", "distance_signed32_overflow",
            "distance_uint32_overflow", "distance_non_twip",
            "distance_huge_integer", "coordinate_negative", "negative_distance",
            "top_bottom_left", "top_bottom_right", "top_bottom_largest",
            "wrap_none_invalid_text",
            "square_invalid_text", "tight_invalid_text", "through_invalid_text",
            "top_bottom_invalid_text",
            "wrap_none_wrong_case", "wrap_none_padded", "wrap_none_empty",
        )
        for variant in invalid_variants:
            with self.subTest(variant=variant), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / f"invalid-{variant}.docx"
                is_top_bottom_side = variant.startswith("top_bottom_")
                invalid_wrap_modes = {
                    "wrap_none_invalid_text": "none",
                    "square_invalid_text": "square",
                    "tight_invalid_text": "tight",
                    "through_invalid_text": "through",
                    "top_bottom_invalid_text": "top_and_bottom",
                    "wrap_none_wrong_case": "none",
                    "wrap_none_padded": "none",
                    "wrap_none_empty": "none",
                }
                invalid_wrap_mode = invalid_wrap_modes.get(variant)
                is_invalid_wrap_enum = invalid_wrap_mode is not None
                write_image_bookmark_document(
                    path, anchored=True,
                    anchor_wrap=invalid_wrap_mode or ("top_and_bottom" if is_top_bottom_side else "tight"),
                    anchor_wrap_text=(
                        variant.removeprefix("top_bottom_") if is_top_bottom_side
                        else {
                            "wrap_none_wrong_case": "Left",
                            "wrap_none_padded": " left ",
                            "wrap_none_empty": "",
                        }.get(variant, "definitely-not-a-wrap-value") if is_invalid_wrap_enum
                        else "bothSides"
                    ),
                )
                with ZipFile(path) as source:
                    parts = {name: source.read(name) for name in source.namelist()}
                root = ET.fromstring(parts["word/document.xml"])
                anchor = next(root.iter(qn("wp:anchor")))
                if is_invalid_wrap_enum:
                    wrap_tag = {
                        "none": "wrapNone", "square": "wrapSquare", "tight": "wrapTight",
                        "through": "wrapThrough", "top_and_bottom": "wrapTopAndBottom",
                    }[invalid_wrap_mode]
                    invalid_value = {
                        "wrap_none_wrong_case": "Left",
                        "wrap_none_padded": " left ",
                        "wrap_none_empty": "",
                    }.get(variant, "definitely-not-a-wrap-value")
                    anchor.find(qn(f"wp:{wrap_tag}")).set("wrapText", invalid_value)
                polygon = anchor.find(f"{qn('wp:wrapTight')}/{qn('wp:wrapPolygon')}")
                if variant == "missing_polygon":
                    anchor.find(qn("wp:wrapTight")).remove(polygon)
                elif variant == "too_few_points":
                    points = polygon.findall(qn("wp:lineTo"))
                    for point in points[1:]:
                        polygon.remove(point)
                elif variant == "coordinate_out_of_range":
                    polygon.find(qn("wp:lineTo")).set("x", "21601")
                elif variant == "coordinate_decimal":
                    polygon.find(qn("wp:lineTo")).set("x", "21600.5")
                elif variant == "coordinate_empty":
                    polygon.find(qn("wp:lineTo")).set("x", "")
                elif variant == "coordinate_huge_integer":
                    polygon.find(qn("wp:lineTo")).set("x", "9" * 200)
                elif variant == "coordinate_unicode_digits":
                    polygon.find(qn("wp:lineTo")).set("x", "٢١٦٠٠")
                elif variant == "coordinate_nbsp":
                    polygon.find(qn("wp:lineTo")).set("x", "\u00a021600\u00a0")
                elif variant == "coordinate_repeated_sign":
                    polygon.find(qn("wp:lineTo")).set("x", "++21600")
                elif variant == "coordinate_sign_only":
                    polygon.find(qn("wp:lineTo")).set("x", "+")
                elif variant == "coordinate_embedded_whitespace":
                    polygon.find(qn("wp:lineTo")).set("x", "21 600")
                elif variant == "coordinate_negative":
                    polygon.find(qn("wp:lineTo")).set("x", "-1")
                elif variant == "distance_unicode_digits":
                    anchor.set("distL", "٢")
                elif variant == "distance_decimal":
                    anchor.set("distL", "1.5")
                elif variant == "distance_empty":
                    anchor.set("distL", "")
                elif variant == "distance_uint32_overflow":
                    anchor.set("distL", "4294967296")
                elif variant == "distance_signed32_overflow":
                    anchor.set("distL", "2147483648")
                elif variant == "distance_huge_integer":
                    anchor.set("distL", "9" * 200)
                elif variant == "distance_non_twip":
                    anchor.set("distL", "1000000")
                elif is_top_bottom_side:
                    pass
                elif is_invalid_wrap_enum:
                    pass
                else:
                    anchor.set("distL", "-1")
                parts["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                with ZipFile(path, "w") as target:
                    for name, content in parts.items():
                        target.writestr(name, content)
                before_bytes = path.read_bytes()
                semantics = docx_body_drawing_semantics(path)
                self.assertTrue(semantics["unresolved_drawings"])
                if is_top_bottom_side:
                    self.assertIn("unsupported_top_and_bottom_wrap_text", {
                        issue["reason"] for issue in semantics["unresolved_drawings"]
                    })
                ok, record, errors = register_document("writer", str(path), tmp)
                self.assertTrue(ok, errors)
                with patch("wps_ai_agent_cli.writer_ops.create_backup") as backup, \
                        patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com") as com:
                    ok, _, errors, _ = writer_fill_bookmark(
                        record["document_id"], "CellMark", "New", f"invalid-wrap-{variant}", workspace=tmp,
                    )
                self.assertFalse(ok)
                self.assertEqual(errors[0]["code"], "BOOKMARK_SCOPE_UNSUPPORTED")
                backup.assert_not_called()
                com.assert_not_called()
                self.assertEqual(path.read_bytes(), before_bytes)

    def test_table_bookmark_rejects_unsupported_wrap_on_any_anchor_before_backup(self):
        import xml.etree.ElementTree as ET
        from docx.oxml.ns import qn
        from zipfile import ZipFile
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document
        from wps_ai_agent_cli.document_text import docx_body_drawing_semantics

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "mixed-supported-unsupported-anchors.docx"
            write_image_bookmark_document(
                path, anchored=True, image_count=2, anchor_wrap="tight",
            )
            with ZipFile(path) as source:
                parts = {name: source.read(name) for name in source.namelist()}
            root = ET.fromstring(parts["word/document.xml"])
            anchors = list(root.iter(qn("wp:anchor")))
            self.assertEqual(len(anchors), 2)
            square = next(child for child in anchors[0] if child.tag == qn("wp:wrapTight"))
            square.tag = qn("wp:wrapSquare")
            square.clear()
            unsupported = next(child for child in anchors[1] if child.tag == qn("wp:wrapTight"))
            unsupported.tag = qn("wp:wrapTopAndBottom")
            unsupported.clear()
            unsupported.set("wrapText", "left")
            parts["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            with ZipFile(path, "w") as target:
                for name, content in parts.items():
                    target.writestr(name, content)

            before_bytes = path.read_bytes()
            semantics = docx_body_drawing_semantics(path)
            self.assertEqual(len(semantics["drawings"]), 2)
            self.assertIn("unsupported_top_and_bottom_wrap_text", {
                issue["reason"] for issue in semantics["unresolved_drawings"]
            })
            ok, record, errors = register_document("writer", str(path), tmp)
            self.assertTrue(ok, errors)
            with patch("wps_ai_agent_cli.writer_ops.create_backup") as backup, \
                    patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com") as com:
                ok, _, errors, _ = writer_fill_bookmark(
                    record["document_id"], "CellMark", "New", "mixed-anchor-unsupported",
                    workspace=tmp,
                )
            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "BOOKMARK_SCOPE_UNSUPPORTED")
            self.assertIn({
                "drawing_type": "anchor",
                "drawing_index": 1,
                "reason": "unsupported_top_and_bottom_wrap_text",
            }, errors[0]["details"])
            backup.assert_not_called()
            com.assert_not_called()
            cli_output = io.StringIO()
            with chdir(tmp):
                run([
                    "writer-fill-bookmark", "--document-id", record["document_id"],
                    "--bookmark-name", "CellMark", "--text", "New", "--dry-run",
                    "--request-id", "multi-anchor-cli-diagnostic",
                ], output_stream=cli_output)
                mcp_ok, mcp_result, mcp_errors = call_mcp_tool(
                    "wps_agent_writer_fill_bookmark",
                    {
                        "document_id": record["document_id"],
                        "bookmark_name": "CellMark",
                        "text": "New",
                        "dry_run": True,
                        "request_id": "multi-anchor-mcp-diagnostic",
                    },
                )
            cli_response = json.loads(cli_output.getvalue())
            expected_detail = {
                "drawing_type": "anchor",
                "drawing_index": 1,
                "reason": "unsupported_top_and_bottom_wrap_text",
            }
            self.assertFalse(cli_response["ok"])
            self.assertEqual(cli_response["errors"][0]["code"], "BOOKMARK_SCOPE_UNSUPPORTED")
            self.assertIn(expected_detail, cli_response["errors"][0]["details"])
            self.assertFalse(mcp_ok)
            self.assertEqual(mcp_errors, [])
            self.assertEqual(
                mcp_result["response"]["errors"][0]["code"], "BOOKMARK_SCOPE_UNSUPPORTED",
            )
            self.assertEqual(mcp_result["response"]["validation"]["status"], "failed")
            self.assertIn(expected_detail, mcp_result["response"]["errors"][0]["details"])
            self.assertEqual(path.read_bytes(), before_bytes)

    def test_preflight_reports_every_unsupported_anchor(self):
        import xml.etree.ElementTree as ET
        from docx.oxml.ns import qn
        from zipfile import ZipFile
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document
        from wps_ai_agent_cli.document_text import docx_body_drawing_semantics

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "two-unsupported-anchors.docx"
            write_image_bookmark_document(
                path, anchored=True, image_count=2, anchor_wrap="tight",
                preamble_paragraphs=1, preamble_images=1,
            )
            with ZipFile(path) as source:
                parts = {name: source.read(name) for name in source.namelist()}
            root = ET.fromstring(parts["word/document.xml"])
            anchors = list(root.iter(qn("wp:anchor")))
            self.assertEqual(len(anchors), 3)
            expected = []
            for index, value in enumerate(("left", "largest"), start=1):
                wrap = next(child for child in anchors[index] if child.tag == qn("wp:wrapTight"))
                wrap.tag = qn("wp:wrapTopAndBottom")
                wrap.clear()
                wrap.set("wrapText", value)
                expected.append({
                    "drawing_type": "anchor",
                    "drawing_index": index,
                    "reason": "unsupported_top_and_bottom_wrap_text",
                })
            parts["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            with ZipFile(path, "w") as target:
                for name, content in parts.items():
                    target.writestr(name, content)

            before_bytes = path.read_bytes()
            semantics = docx_body_drawing_semantics(path)
            self.assertEqual(semantics["unresolved_drawings"], expected)
            ok, record, errors = register_document("writer", str(path), tmp)
            self.assertTrue(ok, errors)
            ok, _, errors, _ = writer_fill_bookmark(
                record["document_id"], "CellMark", "New", "two-unsupported-anchors",
                workspace=tmp,
            )
            self.assertFalse(ok)
            self.assertEqual(errors[0]["details"], expected)
            cli_output = io.StringIO()
            with chdir(tmp):
                run([
                    "writer-fill-bookmark", "--document-id", record["document_id"],
                    "--bookmark-name", "CellMark", "--text", "New", "--dry-run",
                    "--request-id", "two-invalid-cli",
                ], output_stream=cli_output)
                mcp_ok, mcp_result, mcp_errors = call_mcp_tool(
                    "wps_agent_writer_fill_bookmark",
                    {
                        "document_id": record["document_id"],
                        "bookmark_name": "CellMark",
                        "text": "New",
                        "dry_run": True,
                        "request_id": "two-invalid-mcp",
                    },
                )
            cli_response = json.loads(cli_output.getvalue())
            self.assertFalse(cli_response["ok"])
            self.assertEqual(cli_response["errors"][0]["code"], "BOOKMARK_SCOPE_UNSUPPORTED")
            self.assertEqual(cli_response["errors"][0]["details"], expected)
            self.assertFalse(mcp_ok)
            self.assertEqual(mcp_errors, [])
            self.assertEqual(
                mcp_result["response"]["errors"][0]["code"], "BOOKMARK_SCOPE_UNSUPPORTED",
            )
            self.assertEqual(mcp_result["response"]["validation"]["status"], "failed")
            self.assertEqual(mcp_result["response"]["errors"][0]["details"], expected)
            self.assertEqual(path.read_bytes(), before_bytes)

    def test_drawing_diagnostic_index_matches_full_snapshot_order(self):
        import xml.etree.ElementTree as ET
        from docx.oxml.ns import qn
        from zipfile import ZipFile
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document
        from wps_ai_agent_cli.document_text import docx_body_drawing_semantics

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "inline-before-anchors.docx"
            write_image_bookmark_document(
                path, anchored=True, image_count=2, anchor_wrap="tight",
                preamble_paragraphs=1, preamble_images=1,
            )
            with ZipFile(path) as source:
                parts = {name: source.read(name) for name in source.namelist()}
            root = ET.fromstring(parts["word/document.xml"])
            anchors = list(root.iter(qn("wp:anchor")))
            self.assertEqual(len(anchors), 3)
            unsupported = next(child for child in anchors[2] if child.tag == qn("wp:wrapTight"))
            unsupported.tag = qn("wp:wrapTopAndBottom")
            unsupported.clear()
            unsupported.set("wrapText", "left")
            parts["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            with ZipFile(path, "w") as target:
                for name, content in parts.items():
                    target.writestr(name, content)

            before_bytes = path.read_bytes()
            semantics = docx_body_drawing_semantics(path)
            self.assertEqual(len(semantics["drawings"]), 3)
            issue = next(
                item for item in semantics["unresolved_drawings"]
                if item["reason"] == "unsupported_top_and_bottom_wrap_text"
            )
            self.assertEqual(issue["drawing_index"], 2)
            ok, record, errors = register_document("writer", str(path), tmp)
            self.assertTrue(ok, errors)
            ok, _, errors, _ = writer_fill_bookmark(
                record["document_id"], "CellMark", "New", "drawing-index-order", workspace=tmp,
            )
            self.assertFalse(ok)
            self.assertEqual(errors[0]["details"][0]["drawing_index"], 2)
            self.assertEqual(path.read_bytes(), before_bytes)

    def test_invalid_wrap_diagnostic_index_survives_cli_and_mcp(self):
        import xml.etree.ElementTree as ET
        from docx.oxml.ns import qn
        from zipfile import ZipFile
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document
        from wps_ai_agent_cli.document_text import docx_body_drawing_semantics

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "invalid-wrap-index.docx"
            write_image_bookmark_document(
                path, anchored=True, image_count=1, anchor_wrap="tight",
                preamble_paragraphs=1, preamble_images=1,
            )
            with ZipFile(path) as source:
                parts = {name: source.read(name) for name in source.namelist()}
            root = ET.fromstring(parts["word/document.xml"])
            anchor = next(root.iter(qn("wp:anchor")))
            wrap = next(child for child in anchor if child.tag == qn("wp:wrapTight"))
            wrap.set("wrapText", "invalid-enum")
            parts["word/document.xml"] = ET.tostring(
                root, encoding="utf-8", xml_declaration=True,
            )
            with ZipFile(path, "w") as target:
                for name, content in parts.items():
                    target.writestr(name, content)

            before_bytes = path.read_bytes()
            semantics = docx_body_drawing_semantics(path)
            anchor_index = next(
                index for index, drawing in enumerate(semantics["drawings"])
                if drawing["type"] == "anchor"
            )
            expected = [{
                "drawing_type": "anchor",
                "drawing_index": anchor_index,
                "reason": "invalid_wrap_text",
            }]
            self.assertEqual(
                semantics["unresolved_drawings"], expected,
            )
            ok, record, errors = register_document("writer", str(path), tmp)
            self.assertTrue(ok, errors)
            ok, _, errors, _ = writer_fill_bookmark(
                record["document_id"], "CellMark", "New", "invalid-wrap-index",
                workspace=tmp,
            )
            self.assertFalse(ok)
            self.assertEqual(errors[0]["details"], expected)

            cli_output = io.StringIO()
            with chdir(tmp):
                run([
                    "writer-fill-bookmark", "--document-id", record["document_id"],
                    "--bookmark-name", "CellMark", "--text", "New", "--dry-run",
                    "--request-id", "invalid-wrap-index-cli",
                ], output_stream=cli_output)
                mcp_ok, mcp_result, mcp_errors = call_mcp_tool(
                    "wps_agent_writer_fill_bookmark",
                    {
                        "document_id": record["document_id"],
                        "bookmark_name": "CellMark",
                        "text": "New",
                        "dry_run": True,
                        "request_id": "invalid-wrap-index-mcp",
                    },
                )
            cli_response = json.loads(cli_output.getvalue())
            self.assertEqual(cli_response["errors"][0]["details"], expected)
            self.assertFalse(mcp_ok)
            self.assertEqual(mcp_errors, [])
            self.assertEqual(
                mcp_result["response"]["errors"][0]["details"], expected,
            )
            self.assertEqual(path.read_bytes(), before_bytes)

    def test_drawing_diagnostic_reason_schema_matrix(self):
        import xml.etree.ElementTree as ET
        from docx.oxml.ns import qn
        from zipfile import ZipFile
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document
        from wps_ai_agent_cli.document_text import docx_body_drawing_semantics

        cases = (
            ("invalid_relative_height", "relative_height", "invalid_relative_height", {}),
            ("missing_wrap", "missing_wrap", "missing_or_ambiguous_wrap", {}),
            ("invalid_wrap_text", "invalid_wrap_text", "invalid_wrap_text", {}),
            ("unsupported_wrap", "unsupported_wrap", "unsupported_top_and_bottom_wrap_text", {}),
            ("invalid_polygon_shape", "invalid_polygon_shape", "invalid_wrap_polygon_shape", {}),
            ("invalid_polygon_coordinate", "invalid_polygon_coordinate", "invalid_wrap_polygon_coordinate", {}),
            ("polygon_coordinate_range", "polygon_coordinate_range", "wrap_polygon_coordinate_out_of_range", {}),
            ("invalid_distance", "invalid_distance", "invalid_wrap_distance", {"attribute": "distL"}),
            ("missing_image_reference", "missing_image_reference", "missing_image_reference", {}),
            ("missing_relationship", "missing_relationship", "unresolved_drawing_structure", {"relationship_id": True}),
            ("degenerate_polygon", "degenerate_polygon", "degenerate_wrap_polygon", {}),
            ("self_intersecting_polygon", "self_intersecting_polygon", "self_intersecting_wrap_polygon", {}),
        )
        crossing_points = ((0, 0), (21600, 21600), (0, 21600), (21600, 0), (21600, 10800))
        flat_points = ((0, 0), (5400, 0), (10800, 0), (16200, 0), (21600, 0))

        for case_name, mutation, reason, required_fields in cases:
            with self.subTest(case=case_name), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / f"{case_name}.docx"
                write_image_bookmark_document(
                    path, anchored=True, anchor_wrap="tight",
                    anchor_wrap_points=(
                        crossing_points if mutation == "self_intersecting_polygon"
                        else flat_points if mutation == "degenerate_polygon" else None
                    ),
                )
                with ZipFile(path) as source:
                    parts = {name: source.read(name) for name in source.namelist()}
                root = ET.fromstring(parts["word/document.xml"])
                anchor = next(root.iter(qn("wp:anchor")))
                wrap = next(child for child in anchor if child.tag == qn("wp:wrapTight"))

                if mutation == "relative_height":
                    anchor.set("relativeHeight", "invalid")
                elif mutation == "missing_wrap":
                    anchor.remove(wrap)
                elif mutation == "invalid_wrap_text":
                    wrap.set("wrapText", "invalid-enum")
                elif mutation == "unsupported_wrap":
                    wrap.tag = qn("wp:wrapTopAndBottom")
                    wrap.set("wrapText", "left")
                elif mutation == "invalid_polygon_shape":
                    polygon = wrap.find(qn("wp:wrapPolygon"))
                    lines = polygon.findall(qn("wp:lineTo"))
                    for line in lines[1:]:
                        polygon.remove(line)
                elif mutation in {"invalid_polygon_coordinate", "polygon_coordinate_range"}:
                    point = wrap.find(qn("wp:wrapPolygon")).find(qn("wp:start"))
                    point.set("x", "not-an-integer" if mutation == "invalid_polygon_coordinate" else "21601")
                elif mutation == "invalid_distance":
                    anchor.set("distL", "not-an-integer")
                elif mutation == "missing_relationship":
                    rels = ET.fromstring(parts["word/_rels/document.xml.rels"])
                    for relationship in list(rels):
                        rels.remove(relationship)
                    parts["word/_rels/document.xml.rels"] = ET.tostring(
                        rels, encoding="utf-8", xml_declaration=True,
                    )
                elif mutation == "missing_image_reference":
                    for parent in anchor.iter():
                        for blip in list(parent.findall(qn("a:blip"))):
                            parent.remove(blip)

                parts["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                with ZipFile(path, "w") as target:
                    for name, content in parts.items():
                        target.writestr(name, content)

                semantics = docx_body_drawing_semantics(path)
                for item in semantics["unresolved_drawings"]:
                    self.assertIn("drawing_index", item, item)
                issues = [item for item in semantics["unresolved_drawings"] if item.get("reason") == reason]
                self.assertTrue(issues, semantics["unresolved_drawings"])
                for issue in issues:
                    self.assertEqual(issue["drawing_index"], 0)
                    if "drawing_type" in issue:
                        self.assertTrue(issue["drawing_type"])
                    for key, value in required_fields.items():
                        if value is True:
                            self.assertTrue(issue.get(key))
                        else:
                            self.assertEqual(issue[key], value)

    def test_missing_image_relationship_context_survives_cli_and_mcp(self):
        import xml.etree.ElementTree as ET
        from docx.oxml.ns import qn
        from zipfile import ZipFile
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document
        from wps_ai_agent_cli.document_text import docx_body_drawing_semantics

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "missing-image-target.docx"
            write_image_bookmark_document(path)
            with ZipFile(path) as source:
                parts = {name: source.read(name) for name in source.namelist()}
            rels_path = "word/_rels/document.xml.rels"
            rels = ET.fromstring(parts[rels_path])
            image_relationship = next(
                item for item in rels
                if item.get("Type", "").endswith("/image")
            )
            relationship_id = image_relationship.get("Id")
            image_relationship.set("Target", "media/private-missing.png")
            parts[rels_path] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
            with ZipFile(path, "w") as target:
                for name, content in parts.items():
                    target.writestr(name, content)

            before_bytes = path.read_bytes()
            expected = [{
                "relationship_id": relationship_id,
                "target": "media/private-missing.png",
                "drawing_index": 0,
            }]
            semantics = docx_body_drawing_semantics(path)
            self.assertEqual(semantics["unresolved_drawings"], expected)
            ok, record, errors = register_document("writer", str(path), tmp)
            self.assertTrue(ok, errors)
            ok, _, errors, _ = writer_fill_bookmark(
                record["document_id"], "CellMark", "New", "missing-image-context",
                workspace=tmp,
            )
            self.assertFalse(ok)
            self.assertEqual(errors[0]["details"], expected)

            cli_output = io.StringIO()
            with chdir(tmp):
                run([
                    "writer-fill-bookmark", "--document-id", record["document_id"],
                    "--bookmark-name", "CellMark", "--text", "New", "--dry-run",
                    "--request-id", "missing-image-context-cli",
                ], output_stream=cli_output)
                mcp_ok, mcp_result, mcp_errors = call_mcp_tool(
                    "wps_agent_writer_fill_bookmark",
                    {
                        "document_id": record["document_id"],
                        "bookmark_name": "CellMark",
                        "text": "New",
                        "dry_run": True,
                        "request_id": "missing-image-context-mcp",
                    },
                )
            cli_response = json.loads(cli_output.getvalue())
            cli_details = cli_response["errors"][0]["details"]
            mcp_details = mcp_result["response"]["errors"][0]["details"]
            self.assertEqual(cli_details, expected)
            self.assertFalse(mcp_ok)
            self.assertEqual(mcp_errors, [])
            self.assertEqual(mcp_details, expected)
            serialized = json.dumps({"cli": cli_details, "mcp": mcp_details})
            self.assertNotIn(tmp, serialized)
            self.assertNotIn("word/document.xml", serialized)
            self.assertNotIn("private-missing.png bytes", serialized)
            self.assertEqual(path.read_bytes(), before_bytes)

    def test_external_image_relationship_is_metadata_only(self):
        import socket
        import urllib.request
        import xml.etree.ElementTree as ET
        from docx.oxml.ns import qn
        from zipfile import ZipFile
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document
        from wps_ai_agent_cli.document_text import docx_body_drawing_semantics

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "external-image-link.docx"
            write_image_bookmark_document(path)
            with ZipFile(path) as source:
                parts = {name: source.read(name) for name in source.namelist()}
            rels_path = "word/_rels/document.xml.rels"
            rels = ET.fromstring(parts[rels_path])
            image_relationship = next(
                item for item in rels
                if item.get("Type", "").endswith("/image")
            )
            image_relationship.set("Target", "https://example.invalid/image.png")
            image_relationship.set("TargetMode", "External")
            parts[rels_path] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
            root = ET.fromstring(parts["word/document.xml"])
            blip = next(root.iter(qn("a:blip")))
            relationship_id = blip.attrib.pop(qn("r:embed"))
            blip.set(qn("r:link"), relationship_id)
            parts["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            with ZipFile(path, "w") as target:
                for name, content in parts.items():
                    target.writestr(name, content)

            before_bytes = path.read_bytes()
            with patch.object(urllib.request, "urlopen", side_effect=AssertionError("network access")) as urlopen, \
                    patch.object(socket, "create_connection", side_effect=AssertionError("network access")) as connect:
                semantics = docx_body_drawing_semantics(path)
            urlopen.assert_not_called()
            connect.assert_not_called()
            self.assertEqual(semantics["unresolved_drawings"], [])
            self.assertEqual(semantics["drawings"][0]["images"], [{
                "target": "https://example.invalid/image.png",
                "mode": "External",
                "sha256": None,
            }])
            self.assertEqual(path.read_bytes(), before_bytes)

    def test_external_image_target_edge_cases_are_bounded_and_local(self):
        import socket
        import urllib.request
        import xml.etree.ElementTree as ET
        from docx.oxml.ns import qn
        from zipfile import ZipFile
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document
        from wps_ai_agent_cli.document_text import docx_body_drawing_semantics

        cases = (
            ("missing-external", None, "External", True),
            ("malformed-external", "://bad uri", "External", False),
            ("oversized-external", "https://example.invalid/" + ("x" * 65536), "External", False),
            ("internal-traversal", "../../outside.png", None, True),
        )
        for name, target_value, target_mode, expect_unresolved in cases:
            with self.subTest(case=name), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / f"{name}.docx"
                write_image_bookmark_document(path)
                with ZipFile(path) as source:
                    parts = {entry: source.read(entry) for entry in source.namelist()}
                rels_path = "word/_rels/document.xml.rels"
                rels = ET.fromstring(parts[rels_path])
                image_relationship = next(
                    item for item in rels
                    if item.get("Type", "").endswith("/image")
                )
                if target_value is None:
                    image_relationship.attrib.pop("Target", None)
                else:
                    image_relationship.set("Target", target_value)
                if target_mode is None:
                    image_relationship.attrib.pop("TargetMode", None)
                else:
                    image_relationship.set("TargetMode", target_mode)
                parts[rels_path] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
                with ZipFile(path, "w") as target:
                    for entry, content in parts.items():
                        target.writestr(entry, content)

                before_bytes = path.read_bytes()
                read_entries = []
                original_read = ZipFile.read

                def tracked_read(archive, entry, *args, **kwargs):
                    read_entries.append(entry)
                    return original_read(archive, entry, *args, **kwargs)

                with patch.object(urllib.request, "urlopen", side_effect=AssertionError("network access")), \
                        patch.object(socket, "create_connection", side_effect=AssertionError("network access")), \
                        patch.object(ZipFile, "read", new=tracked_read):
                    semantics = docx_body_drawing_semantics(path)

                self.assertEqual(bool(semantics["unresolved_drawings"]), expect_unresolved)
                self.assertTrue(all(item["drawing_index"] == 0 for item in semantics["unresolved_drawings"]))
                if target_mode == "External":
                    self.assertFalse(any(entry.startswith("word/media/") for entry in read_entries))
                if name == "internal-traversal":
                    self.assertNotIn("../outside.png", read_entries)
                self.assertEqual(path.read_bytes(), before_bytes)

    def test_image_relationship_target_mode_controls_package_reads(self):
        import xml.etree.ElementTree as ET
        from zipfile import ZipFile
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document
        from wps_ai_agent_cli.document_text import docx_body_drawing_semantics

        for mode in (None, "Internal", "External", "Unknown"):
            with self.subTest(mode=mode), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / "target-mode.docx"
                write_image_bookmark_document(path)
                with ZipFile(path) as source:
                    parts = {name: source.read(name) for name in source.namelist()}
                rels_path = "word/_rels/document.xml.rels"
                rels = ET.fromstring(parts[rels_path])
                relationship = next(
                    item for item in rels
                    if item.get("Type", "").endswith("/image")
                )
                if mode is None:
                    relationship.attrib.pop("TargetMode", None)
                else:
                    relationship.set("TargetMode", mode)
                parts[rels_path] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
                with ZipFile(path, "w") as target:
                    for name, content in parts.items():
                        target.writestr(name, content)

                read_entries = []
                original_read = ZipFile.read

                def tracked_read(archive, entry, *args, **kwargs):
                    read_entries.append(entry)
                    return original_read(archive, entry, *args, **kwargs)

                with patch.object(ZipFile, "read", new=tracked_read):
                    semantics = docx_body_drawing_semantics(path)
                media_reads = [entry for entry in read_entries if entry.startswith("word/media/")]
                if mode in {None, "Internal"}:
                    self.assertEqual(semantics["unresolved_drawings"], [])
                    self.assertTrue(media_reads)
                    self.assertIsNotNone(semantics["drawings"][0]["images"][0]["sha256"])
                elif mode == "External":
                    self.assertEqual(semantics["unresolved_drawings"], [])
                    self.assertEqual(media_reads, [])
                    self.assertIsNone(semantics["drawings"][0]["images"][0]["sha256"])
                else:
                    self.assertEqual(media_reads, [])
                    self.assertEqual(semantics["unresolved_drawings"], [{
                        "relationship_id": relationship.get("Id"),
                        "target": relationship.get("Target"),
                        "mode": "Unknown",
                        "drawing_index": 0,
                        "reason": "invalid_relationship_target_mode",
                    }])

    def test_internal_image_target_resolution_stays_within_package(self):
        import posixpath
        import xml.etree.ElementTree as ET
        from zipfile import ZipFile
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document
        from wps_ai_agent_cli.document_text import docx_body_drawing_semantics

        cases = ("relative", "rooted", "missing", "traversal")
        for case in cases:
            with self.subTest(case=case), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / "internal-target.docx"
                write_image_bookmark_document(path)
                with ZipFile(path) as source:
                    parts = {name: source.read(name) for name in source.namelist()}
                rels_path = "word/_rels/document.xml.rels"
                rels = ET.fromstring(parts[rels_path])
                relationship = next(
                    item for item in rels
                    if item.get("Type", "").endswith("/image")
                )
                relationship_id = relationship.get("Id")
                original_target = relationship.get("Target")
                package_target = posixpath.normpath(posixpath.join("word", original_target))
                if case == "rooted":
                    relationship.set("Target", f"/{package_target}")
                elif case == "missing":
                    relationship.set("Target", "media/not-present.png")
                elif case == "traversal":
                    relationship.set("Target", "../../outside.png")
                    parts["../outside.png"] = b"must not be read"
                parts[rels_path] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
                with ZipFile(path, "w") as target:
                    for name, content in parts.items():
                        target.writestr(name, content)

                read_entries = []
                original_read = ZipFile.read

                def tracked_read(archive, entry, *args, **kwargs):
                    read_entries.append(entry)
                    return original_read(archive, entry, *args, **kwargs)

                with patch.object(ZipFile, "read", new=tracked_read):
                    semantics = docx_body_drawing_semantics(path)
                media = semantics["drawings"][0]["images"][0]
                if case in {"relative", "rooted"}:
                    self.assertEqual(semantics["unresolved_drawings"], [])
                    self.assertTrue(media["sha256"])
                    self.assertIn(package_target, read_entries)
                elif case == "missing":
                    self.assertEqual(semantics["unresolved_drawings"], [{
                        "relationship_id": relationship_id,
                        "target": "media/not-present.png",
                        "drawing_index": 0,
                    }])
                else:
                    self.assertEqual(semantics["unresolved_drawings"], [{
                        "relationship_id": relationship_id,
                        "target": "../../outside.png",
                        "drawing_index": 0,
                        "reason": "target_outside_package",
                    }])
                    self.assertNotIn("../outside.png", read_entries)

    def test_relationship_boundary_diagnostics_survive_preflight_cli_and_mcp(self):
        import xml.etree.ElementTree as ET
        from zipfile import ZipFile
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document

        cases = (
            (
                "unknown-mode", "TargetMode", "Unknown",
                "media/image1.png", {
                    "relationship_id": None,
                    "target": "media/image1.png",
                    "mode": "Unknown",
                    "drawing_index": 0,
                    "reason": "invalid_relationship_target_mode",
                },
            ),
            (
                "outside-package", "Target", "../../outside.png",
                "TargetMode", None,
            ),
        )
        for case_name, first_attr, first_value, second_attr, second_value in cases:
            with self.subTest(case=case_name), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / f"{case_name}.docx"
                write_image_bookmark_document(path)
                with ZipFile(path) as source:
                    parts = {name: source.read(name) for name in source.namelist()}
                rels_path = "word/_rels/document.xml.rels"
                rels = ET.fromstring(parts[rels_path])
                relationship = next(
                    item for item in rels
                    if item.get("Type", "").endswith("/image")
                )
                relationship_id = relationship.get("Id")
                if first_attr == "TargetMode":
                    relationship.set(first_attr, first_value)
                    expected = dict(second_value)
                    expected["relationship_id"] = relationship_id
                else:
                    relationship.set(first_attr, first_value)
                    expected = {
                        "relationship_id": relationship_id,
                        "target": first_value,
                        "drawing_index": 0,
                        "reason": "target_outside_package",
                    }
                    parts["../outside.png"] = b"must remain unread"
                if second_value is not None and second_attr == "TargetMode":
                    relationship.set(second_attr, second_value)
                parts[rels_path] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
                with ZipFile(path, "w") as target:
                    for name, content in parts.items():
                        target.writestr(name, content)

                before_bytes = path.read_bytes()
                ok, record, errors = register_document("writer", str(path), tmp)
                self.assertTrue(ok, errors)
                cli_output = io.StringIO()
                with patch("wps_ai_agent_cli.writer_ops.create_backup") as backup, \
                        patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com") as com, \
                        chdir(tmp):
                    ok, _, direct_errors, _ = writer_fill_bookmark(
                        record["document_id"], "CellMark", "New", f"{case_name}-direct",
                        workspace=tmp,
                    )
                    run([
                        "writer-fill-bookmark", "--document-id", record["document_id"],
                        "--bookmark-name", "CellMark", "--text", "New", "--dry-run",
                        "--request-id", f"{case_name}-cli",
                    ], output_stream=cli_output)
                    mcp_ok, mcp_result, mcp_errors = call_mcp_tool(
                        "wps_agent_writer_fill_bookmark",
                        {
                            "document_id": record["document_id"],
                            "bookmark_name": "CellMark",
                            "text": "New",
                            "dry_run": True,
                            "request_id": f"{case_name}-mcp",
                        },
                    )
                self.assertFalse(ok)
                self.assertEqual(direct_errors[0]["code"], "BOOKMARK_SCOPE_UNSUPPORTED")
                self.assertEqual(direct_errors[0]["details"], [expected])
                backup.assert_not_called()
                com.assert_not_called()

                cli_response = json.loads(cli_output.getvalue())
                self.assertEqual(cli_response["errors"][0]["code"], "BOOKMARK_SCOPE_UNSUPPORTED")
                self.assertEqual(cli_response["errors"][0]["details"], [expected])
                self.assertFalse(mcp_ok)
                self.assertEqual(mcp_errors, [])
                self.assertEqual(
                    mcp_result["response"]["errors"][0]["code"], "BOOKMARK_SCOPE_UNSUPPORTED",
                )
                self.assertEqual(mcp_result["response"]["errors"][0]["details"], [expected])
                self.assertEqual(path.read_bytes(), before_bytes)

    def test_multiple_relationship_errors_preserve_order_through_cli_and_mcp(self):
        import xml.etree.ElementTree as ET
        from docx.oxml.ns import qn
        from zipfile import ZipFile
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "multiple-relationship-errors.docx"
            write_image_bookmark_document(path, anchored=True, image_count=2)
            with ZipFile(path) as source:
                parts = {name: source.read(name) for name in source.namelist()}
            root = ET.fromstring(parts["word/document.xml"])
            rels_path = "word/_rels/document.xml.rels"
            rels = ET.fromstring(parts[rels_path])
            relationships = {
                item.get("Id"): item for item in rels
                if item.get("Type", "").endswith("/image")
            }
            blip_ids = [blip.get(qn("r:embed")) for blip in root.iter(qn("a:blip"))]
            self.assertEqual(len(blip_ids), 2)
            first = relationships[blip_ids[0]]
            first.set("TargetMode", "Unknown")
            second = relationships[blip_ids[1]]
            second.set("Target", "../../outside.png")
            parts["../outside.png"] = b"must remain unread"
            parts[rels_path] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
            with ZipFile(path, "w") as target:
                for name, content in parts.items():
                    target.writestr(name, content)

            expected = [
                {
                    "relationship_id": blip_ids[0],
                    "target": first.get("Target"),
                    "mode": "Unknown",
                    "drawing_index": 0,
                    "reason": "invalid_relationship_target_mode",
                },
                {
                    "relationship_id": blip_ids[1],
                    "target": "../../outside.png",
                    "drawing_index": 1,
                    "reason": "target_outside_package",
                },
            ]
            before_bytes = path.read_bytes()
            ok, record, errors = register_document("writer", str(path), tmp)
            self.assertTrue(ok, errors)
            cli_output = io.StringIO()
            with patch("wps_ai_agent_cli.writer_ops.create_backup") as backup, \
                    patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com") as com, \
                    chdir(tmp):
                ok, _, direct_errors, _ = writer_fill_bookmark(
                    record["document_id"], "CellMark", "New", "multiple-relationship-direct",
                    workspace=tmp,
                )
                run([
                    "writer-fill-bookmark", "--document-id", record["document_id"],
                    "--bookmark-name", "CellMark", "--text", "New", "--dry-run",
                    "--request-id", "multiple-relationship-cli",
                ], output_stream=cli_output)
                mcp_ok, mcp_result, mcp_errors = call_mcp_tool(
                    "wps_agent_writer_fill_bookmark",
                    {
                        "document_id": record["document_id"],
                        "bookmark_name": "CellMark",
                        "text": "New",
                        "dry_run": True,
                        "request_id": "multiple-relationship-mcp",
                    },
                )
            self.assertFalse(ok)
            self.assertEqual(direct_errors[0]["details"], expected)
            backup.assert_not_called()
            com.assert_not_called()
            cli_response = json.loads(cli_output.getvalue())
            self.assertEqual(cli_response["errors"][0]["details"], expected)
            self.assertFalse(mcp_ok)
            self.assertEqual(mcp_errors, [])
            self.assertEqual(mcp_result["response"]["errors"][0]["details"], expected)
            self.assertEqual(path.read_bytes(), before_bytes)

    def test_relationship_diagnostic_values_are_bounded_deterministically(self):
        import xml.etree.ElementTree as ET
        from docx.oxml.ns import qn
        from zipfile import ZipFile
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document
        from wps_ai_agent_cli.document_text import docx_body_drawing_semantics

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "long-relationship-fields.docx"
            write_image_bookmark_document(path)
            with ZipFile(path) as source:
                parts = {name: source.read(name) for name in source.namelist()}
            rels_path = "word/_rels/document.xml.rels"
            rels = ET.fromstring(parts[rels_path])
            relationship = next(
                item for item in rels
                if item.get("Type", "").endswith("/image")
            )
            old_id = relationship.get("Id")
            long_id = "rId" + ("I" * 1024)
            long_target = "media/" + ("target" * 256) + ".png"
            long_mode = "Unknown" + ("M" * 512)
            relationship.set("Id", long_id)
            relationship.set("Target", long_target)
            relationship.set("TargetMode", long_mode)
            parts[rels_path] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
            root = ET.fromstring(parts["word/document.xml"])
            blip = next(root.iter(qn("a:blip")))
            blip.set(qn("r:embed"), long_id)
            self.assertNotEqual(old_id, long_id)
            parts["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            with ZipFile(path, "w") as target:
                for name, content in parts.items():
                    target.writestr(name, content)

            first = docx_body_drawing_semantics(path)["unresolved_drawings"]
            second = docx_body_drawing_semantics(path)["unresolved_drawings"]
            self.assertEqual(first, second)
            self.assertEqual(len(first), 1)
            issue = first[0]
            for field, original in (
                ("relationship_id", long_id), ("target", long_target), ("mode", long_mode),
            ):
                self.assertEqual(len(issue[field]), 256)
                self.assertEqual(issue[field], original[:253] + "...")
                self.assertTrue(issue[f"{field}_truncated"])
            self.assertEqual(issue["drawing_index"], 0)
            self.assertEqual(issue["reason"], "invalid_relationship_target_mode")
            self.assertLess(len(json.dumps(issue)), 1100)

    def test_bounded_relationship_diagnostic_survives_cli_and_mcp(self):
        import xml.etree.ElementTree as ET
        from docx.oxml.ns import qn
        from zipfile import ZipFile
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bounded-relationship-error.docx"
            write_image_bookmark_document(path)
            with ZipFile(path) as source:
                parts = {name: source.read(name) for name in source.namelist()}
            rels_path = "word/_rels/document.xml.rels"
            rels = ET.fromstring(parts[rels_path])
            relationship = next(
                item for item in rels
                if item.get("Type", "").endswith("/image")
            )
            long_id = "rId" + ("I" * 1024)
            long_target = "media/" + ("target" * 256) + ".png"
            long_mode = "Unknown" + ("M" * 512)
            relationship.set("Id", long_id)
            relationship.set("Target", long_target)
            relationship.set("TargetMode", long_mode)
            parts[rels_path] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
            root = ET.fromstring(parts["word/document.xml"])
            next(root.iter(qn("a:blip"))).set(qn("r:embed"), long_id)
            parts["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            with ZipFile(path, "w") as target:
                for name, content in parts.items():
                    target.writestr(name, content)

            expected = {
                "relationship_id": long_id[:253] + "...",
                "relationship_id_truncated": True,
                "target": long_target[:253] + "...",
                "target_truncated": True,
                "mode": long_mode[:253] + "...",
                "mode_truncated": True,
                "drawing_index": 0,
                "reason": "invalid_relationship_target_mode",
            }
            before_bytes = path.read_bytes()
            ok, record, errors = register_document("writer", str(path), tmp)
            self.assertTrue(ok, errors)
            cli_output = io.StringIO()
            with patch("wps_ai_agent_cli.writer_ops.create_backup") as backup, \
                    patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com") as com, \
                    chdir(tmp):
                ok, _, direct_errors, _ = writer_fill_bookmark(
                    record["document_id"], "CellMark", "New", "bounded-diagnostic-direct",
                    workspace=tmp,
                )
                run([
                    "writer-fill-bookmark", "--document-id", record["document_id"],
                    "--bookmark-name", "CellMark", "--text", "New", "--dry-run",
                    "--request-id", "bounded-diagnostic-cli",
                ], output_stream=cli_output)
                mcp_ok, mcp_result, mcp_errors = call_mcp_tool(
                    "wps_agent_writer_fill_bookmark",
                    {
                        "document_id": record["document_id"],
                        "bookmark_name": "CellMark",
                        "text": "New",
                        "dry_run": True,
                        "request_id": "bounded-diagnostic-mcp",
                    },
                )
            self.assertFalse(ok)
            self.assertEqual(direct_errors[0]["details"], [expected])
            backup.assert_not_called()
            com.assert_not_called()
            cli_response = json.loads(cli_output.getvalue())
            self.assertEqual(cli_response["errors"][0]["details"], [expected])
            self.assertFalse(mcp_ok)
            self.assertEqual(mcp_errors, [])
            self.assertEqual(mcp_result["response"]["errors"][0]["details"], [expected])
            self.assertEqual(path.read_bytes(), before_bytes)

    def test_bounded_relationship_errors_are_isolated_per_drawing(self):
        import xml.etree.ElementTree as ET
        from docx.oxml.ns import qn
        from zipfile import ZipFile
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bounded-multiple-relationship-errors.docx"
            write_image_bookmark_document(path, anchored=True, image_count=2)
            with ZipFile(path) as source:
                parts = {name: source.read(name) for name in source.namelist()}
            root = ET.fromstring(parts["word/document.xml"])
            rels_path = "word/_rels/document.xml.rels"
            rels = ET.fromstring(parts[rels_path])
            relationships = {
                item.get("Id"): item for item in rels
                if item.get("Type", "").endswith("/image")
            }
            blip_ids = [blip.get(qn("r:embed")) for blip in root.iter(qn("a:blip"))]
            self.assertEqual(len(blip_ids), 2)
            long_id = "rId" + ("L" * 600)
            long_target = "media/" + ("long" * 100) + ".png"
            long_mode = "Unknown" + ("X" * 400)
            first = relationships[blip_ids[0]]
            first.set("Id", long_id)
            first.set("Target", long_target)
            first.set("TargetMode", long_mode)
            next(root.iter(qn("a:blip"))).set(qn("r:embed"), long_id)
            second = relationships[blip_ids[1]]
            second.set("Target", "../../outside.png")
            parts["../outside.png"] = b"must remain unread"
            parts[rels_path] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
            parts["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            with ZipFile(path, "w") as target:
                for name, content in parts.items():
                    target.writestr(name, content)

            expected = [
                {
                    "relationship_id": long_id[:253] + "...",
                    "relationship_id_truncated": True,
                    "target": long_target[:253] + "...",
                    "target_truncated": True,
                    "mode": long_mode[:253] + "...",
                    "mode_truncated": True,
                    "drawing_index": 0,
                    "reason": "invalid_relationship_target_mode",
                },
                {
                    "relationship_id": blip_ids[1],
                    "target": "../../outside.png",
                    "drawing_index": 1,
                    "reason": "target_outside_package",
                },
            ]
            before_bytes = path.read_bytes()
            ok, record, errors = register_document("writer", str(path), tmp)
            self.assertTrue(ok, errors)
            cli_output = io.StringIO()
            with patch("wps_ai_agent_cli.writer_ops.create_backup") as backup, \
                    patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com") as com, \
                    chdir(tmp):
                ok, _, direct_errors, _ = writer_fill_bookmark(
                    record["document_id"], "CellMark", "New", "bounded-multiple-direct",
                    workspace=tmp,
                )
                run([
                    "writer-fill-bookmark", "--document-id", record["document_id"],
                    "--bookmark-name", "CellMark", "--text", "New", "--dry-run",
                    "--request-id", "bounded-multiple-cli",
                ], output_stream=cli_output)
                mcp_ok, mcp_result, mcp_errors = call_mcp_tool(
                    "wps_agent_writer_fill_bookmark",
                    {
                        "document_id": record["document_id"],
                        "bookmark_name": "CellMark",
                        "text": "New",
                        "dry_run": True,
                        "request_id": "bounded-multiple-mcp",
                    },
                )
            self.assertFalse(ok)
            self.assertEqual(direct_errors[0]["details"], expected)
            backup.assert_not_called()
            com.assert_not_called()
            cli_response = json.loads(cli_output.getvalue())
            self.assertEqual(cli_response["errors"][0]["details"], expected)
            self.assertFalse(mcp_ok)
            self.assertEqual(mcp_errors, [])
            self.assertEqual(mcp_result["response"]["errors"][0]["details"], expected)
            self.assertEqual(path.read_bytes(), before_bytes)

    def test_total_drawing_diagnostic_payload_has_a_stable_cap(self):
        import xml.etree.ElementTree as ET
        from zipfile import ZipFile
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document
        from wps_ai_agent_cli.document_text import docx_body_drawing_semantics

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "many-invalid-drawings.docx"
            write_image_bookmark_document(path, anchored=True, image_count=150)
            with ZipFile(path) as source:
                parts = {name: source.read(name) for name in source.namelist()}
            rels_path = "word/_rels/document.xml.rels"
            rels = ET.fromstring(parts[rels_path])
            for relationship in rels:
                if relationship.get("Type", "").endswith("/image"):
                    relationship.set("Target", "media/" + ("target" * 80) + ".png")
                    relationship.set("TargetMode", "Unknown" + ("M" * 400))
            parts[rels_path] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
            with ZipFile(path, "w") as target:
                for name, content in parts.items():
                    target.writestr(name, content)

            first = docx_body_drawing_semantics(path)["unresolved_drawings"]
            second = docx_body_drawing_semantics(path)["unresolved_drawings"]
            self.assertEqual(first, second)
            self.assertEqual(len(first), 128)
            self.assertEqual([item["drawing_index"] for item in first[:127]], list(range(127)))
            self.assertTrue(all(item["reason"] == "invalid_relationship_target_mode" for item in first[:127]))
            self.assertEqual(first[-1], {
                "drawing_index": 127,
                "reason": "additional_drawing_diagnostics_truncated",
                "omitted_count": 23,
            })
            self.assertLess(len(json.dumps(first)), 120000)

    def test_diagnostic_truncation_summary_survives_cli_and_mcp(self):
        import xml.etree.ElementTree as ET
        from zipfile import ZipFile
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document
        from wps_ai_agent_cli.document_text import docx_body_drawing_semantics

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "many-invalid-relationships.docx"
            write_image_bookmark_document(path, anchored=True, image_count=150)
            with ZipFile(path) as source:
                parts = {name: source.read(name) for name in source.namelist()}
            rels_path = "word/_rels/document.xml.rels"
            rels = ET.fromstring(parts[rels_path])
            for relationship in rels:
                if relationship.get("Type", "").endswith("/image"):
                    relationship.set("TargetMode", "Unknown")
            parts[rels_path] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
            with ZipFile(path, "w") as target:
                for name, content in parts.items():
                    target.writestr(name, content)

            expected = docx_body_drawing_semantics(path)["unresolved_drawings"]
            self.assertEqual(len(expected), 128)
            self.assertEqual(expected[-1]["reason"], "additional_drawing_diagnostics_truncated")
            before_bytes = path.read_bytes()
            ok, record, errors = register_document("writer", str(path), tmp)
            self.assertTrue(ok, errors)
            cli_output = io.StringIO()
            with patch("wps_ai_agent_cli.writer_ops.create_backup") as backup, \
                    patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com") as com, \
                    chdir(tmp):
                ok, _, direct_errors, _ = writer_fill_bookmark(
                    record["document_id"], "CellMark", "New", "diagnostic-summary-direct",
                    workspace=tmp,
                )
                run([
                    "writer-fill-bookmark", "--document-id", record["document_id"],
                    "--bookmark-name", "CellMark", "--text", "New", "--dry-run",
                    "--request-id", "diagnostic-summary-cli",
                ], output_stream=cli_output)
                mcp_ok, mcp_result, mcp_errors = call_mcp_tool(
                    "wps_agent_writer_fill_bookmark",
                    {
                        "document_id": record["document_id"],
                        "bookmark_name": "CellMark",
                        "text": "New",
                        "dry_run": True,
                        "request_id": "diagnostic-summary-mcp",
                    },
                )
            self.assertFalse(ok)
            self.assertEqual(direct_errors[0]["details"], expected)
            backup.assert_not_called()
            com.assert_not_called()
            cli_response = json.loads(cli_output.getvalue())
            self.assertFalse(cli_response["ok"])
            self.assertEqual(cli_response["validation"]["status"], "failed")
            self.assertEqual(cli_response["errors"][0]["details"], expected)
            self.assertFalse(mcp_ok)
            self.assertEqual(mcp_errors, [])
            self.assertEqual(mcp_result["response"]["validation"]["status"], "failed")
            self.assertEqual(mcp_result["response"]["errors"][0]["details"], expected)
            self.assertEqual(path.read_bytes(), before_bytes)

    def test_drawing_diagnostic_count_cap_boundary(self):
        import xml.etree.ElementTree as ET
        from zipfile import ZipFile
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document
        from wps_ai_agent_cli.document_text import docx_body_drawing_semantics

        for drawing_count in (127, 128, 129):
            with self.subTest(drawing_count=drawing_count), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / f"diagnostic-boundary-{drawing_count}.docx"
                write_image_bookmark_document(
                    path, anchored=True, image_count=drawing_count,
                )
                with ZipFile(path) as source:
                    parts = {name: source.read(name) for name in source.namelist()}
                rels_path = "word/_rels/document.xml.rels"
                rels = ET.fromstring(parts[rels_path])
                for relationship in rels:
                    if relationship.get("Type", "").endswith("/image"):
                        relationship.set("TargetMode", "Unknown")
                parts[rels_path] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
                with ZipFile(path, "w") as target:
                    for name, content in parts.items():
                        target.writestr(name, content)

                issues = docx_body_drawing_semantics(path)["unresolved_drawings"]
                if drawing_count <= 128:
                    self.assertEqual(len(issues), drawing_count)
                    self.assertEqual(
                        [issue["drawing_index"] for issue in issues],
                        list(range(drawing_count)),
                    )
                    self.assertTrue(all(
                        issue["reason"] == "invalid_relationship_target_mode"
                        for issue in issues
                    ))
                else:
                    self.assertEqual(len(issues), 128)
                    self.assertEqual(
                        [issue["drawing_index"] for issue in issues[:127]],
                        list(range(127)),
                    )
                    self.assertEqual(issues[-1], {
                        "drawing_index": 127,
                        "reason": "additional_drawing_diagnostics_truncated",
                        "omitted_count": 2,
                    })

    def test_wrap_polygon_validation_allows_concave_and_rejects_crossing_or_flat_paths(self):
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document
        from wps_ai_agent_cli.document_text import docx_body_drawing_semantics

        concave = ((0, 0), (21600, 0), (21600, 7200), (7200, 7200), (7200, 21600), (0, 21600))
        winding_and_closure_variants = (
            concave,
            tuple(reversed(concave)),
            (*concave, concave[0]),
            (concave[0], *reversed(concave[1:]), concave[0]),
        )
        invalid = {
            "self_intersecting": ((0, 0), (21600, 21600), (0, 21600), (21600, 0), (21600, 10800)),
            "zero_area": ((0, 0), (7200, 0), (14400, 0), (21600, 0)),
        }
        with tempfile.TemporaryDirectory() as tmp:
            valid_path = Path(tmp) / "concave.docx"
            write_image_bookmark_document(
                valid_path, anchored=True, anchor_wrap="tight", anchor_wrap_points=concave,
            )
            self.assertEqual(docx_body_drawing_semantics(valid_path)["unresolved_drawings"], [])
            for index, points in enumerate(winding_and_closure_variants):
                variant_path = Path(tmp) / f"valid-variant-{index}.docx"
                write_image_bookmark_document(
                    variant_path, anchored=True, anchor_wrap="through", anchor_wrap_points=points,
                )
                self.assertEqual(docx_body_drawing_semantics(variant_path)["unresolved_drawings"], [])
            for index in (0, 1, 3, 5):
                points = (*concave[index:], *concave[:index])
                variant_path = Path(tmp) / f"valid-cyclic-start-{index}.docx"
                write_image_bookmark_document(
                    variant_path, anchored=True, anchor_wrap="tight", anchor_wrap_points=points,
                )
                self.assertEqual(docx_body_drawing_semantics(variant_path)["unresolved_drawings"], [])
            collinear = (
                (0, 0), (10800, 0), (21600, 0), (21600, 3600), (21600, 7200),
                (14400, 7200), (7200, 7200), (7200, 14400), (7200, 21600),
                (3600, 21600), (0, 21600), (0, 10800),
            )
            collinear_path = Path(tmp) / "valid-collinear.docx"
            write_image_bookmark_document(
                collinear_path, anchored=True, anchor_wrap="through", anchor_wrap_points=collinear,
            )
            self.assertEqual(docx_body_drawing_semantics(collinear_path)["unresolved_drawings"], [])
            near_degenerate = (
                (0, 0), (1, 0), (21600, 0), (21600, 7200),
                (7200, 7200), (7200, 21600), (0, 21600),
            )
            near_degenerate_path = Path(tmp) / "valid-near-degenerate.docx"
            write_image_bookmark_document(
                near_degenerate_path, anchored=True, anchor_wrap="tight", anchor_wrap_points=near_degenerate,
            )
            self.assertEqual(docx_body_drawing_semantics(near_degenerate_path)["unresolved_drawings"], [])
            near_collinear = (
                (0, 0), (10800, 120), (21600, 0), (21600, 7200),
                (7200, 7200), (7200, 21600), (0, 21600),
            )
            near_collinear_path = Path(tmp) / "valid-near-collinear.docx"
            write_image_bookmark_document(
                near_collinear_path, anchored=True, anchor_wrap="through", anchor_wrap_points=near_collinear,
            )
            self.assertEqual(docx_body_drawing_semantics(near_collinear_path)["unresolved_drawings"], [])

        for variant, points in invalid.items():
            with self.subTest(variant=variant), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / f"{variant}.docx"
                write_image_bookmark_document(
                    path, anchored=True, anchor_wrap="through", anchor_wrap_points=points,
                )
                self.assertTrue(docx_body_drawing_semantics(path)["unresolved_drawings"])
                before_bytes = path.read_bytes()
                ok, record, errors = register_document("writer", str(path), tmp)
                self.assertTrue(ok, errors)
                with patch("wps_ai_agent_cli.writer_ops.create_backup") as backup, \
                        patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com") as com:
                    ok, _, errors, _ = writer_fill_bookmark(
                        record["document_id"], "CellMark", "New", f"wrap-polygon-{variant}", workspace=tmp,
                    )
                self.assertFalse(ok)
                self.assertEqual(errors[0]["code"], "BOOKMARK_SCOPE_UNSUPPORTED")
                backup.assert_not_called()
                com.assert_not_called()
                self.assertEqual(path.read_bytes(), before_bytes)

    def test_table_bookmark_rejects_character_anchor_drift_before_backup(self):
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "character-anchor.docx"
            write_image_bookmark_document(
                path, anchored=True, anchor_relative_frames=("character", "line"),
            )
            _, record, _ = register_document("writer", str(path), tmp)
            with patch("wps_ai_agent_cli.writer_ops.create_backup") as backup, \
                    patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com") as com:
                ok, _, errors, _ = writer_fill_bookmark(
                    record["document_id"], "CellMark", "New", "character-anchor", workspace=tmp,
                )
            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "BOOKMARK_SCOPE_UNSUPPORTED")
            backup.assert_not_called()
            com.assert_not_called()

    def test_table_bookmark_rejects_changed_drawing_relationship(self):
        from zipfile import ZipFile
        from tests.test_writer_table_bookmark_feasibility import write_image_bookmark_document
        from wps_ai_agent_cli.document_text import docx_body_drawing_semantics

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "drawing-neighbor.docx"
            write_image_bookmark_document(path)
            before = docx_body_drawing_semantics(path)
            target = before["drawings"][0]["images"][0]["target"].encode("utf-8")
            _, record, _ = register_document("writer", str(path), tmp)

            def change_relationship(*_args):
                with ZipFile(path) as archive:
                    parts = {name: archive.read(name) for name in archive.namelist()}
                parts["word/document.xml"] = parts["word/document.xml"].replace(b">Old<", b">New<")
                parts["word/_rels/document.xml.rels"] = parts["word/_rels/document.xml.rels"].replace(
                    target, b"media/missing.png",
                )
                with ZipFile(path, "w") as archive:
                    for name, content in parts.items():
                        archive.writestr(name, content)
                return {"ok": True, "errors": [], "data": {"backend": "test"}}

            with patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com", side_effect=change_relationship):
                ok, result, errors, _ = writer_fill_bookmark(
                    record["document_id"], "CellMark", "New", "drawing-relationship", workspace=tmp,
                )
            self.assertFalse(ok)
            self.assertFalse(result["readback_passed"])
            self.assertEqual(errors[0]["code"], "VALIDATION_FAILED")

    def test_bookmark_fill_allows_only_terminal_table_empty_paragraph_normalization(self):
        from zipfile import ZipFile

        for has_terminal_table, corrupt_table, expected_ok in (
            (True, False, True), (False, False, False), (True, True, False),
        ):
            with self.subTest(has_terminal_table=has_terminal_table, corrupt_table=corrupt_table), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / "bookmark.docx"
                table = '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Cell</w:t></w:r></w:p></w:tc></w:tr></w:tbl>' if has_terminal_table else ''
                write_document(path, '<w:p><w:bookmarkStart w:id="1" w:name="Target"/>'
                                     '<w:r><w:t>old</w:t></w:r><w:bookmarkEnd w:id="1"/></w:p>' + table)
                _, record, _ = register_document("writer", str(path), tmp)

                def replace_xml(*_args):
                    with ZipFile(path) as archive:
                        parts = {name: archive.read(name) for name in archive.namelist()}
                    payload = parts["word/document.xml"].replace(b">old<", b">new<")
                    if corrupt_table:
                        payload = payload.replace(b">Cell<", b">Other<")
                    parts["word/document.xml"] = payload.replace(b"</w:body>", b"<w:p/></w:body>")
                    with ZipFile(path, "w") as archive:
                        for name, content in parts.items():
                            archive.writestr(name, content)
                    return {"ok": True, "errors": [], "data": {"backend": "test"}}

                with patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com", side_effect=replace_xml):
                    ok, result, errors, _ = writer_fill_bookmark(
                        record["document_id"], "Target", "new", "table-tail", workspace=tmp,
                    )
                self.assertEqual(ok, expected_ok)
                self.assertEqual(result["body_tail_normalized"], has_terminal_table)
                self.assertEqual(result["readback_passed"], expected_ok)
                if not expected_ok:
                    self.assertEqual(errors[0]["code"], "VALIDATION_FAILED")

    def test_bookmark_fill_rejects_duplicate_and_cross_paragraph_targets_before_backup(self):
        cases = [
            ('<w:p><w:bookmarkStart w:id="1" w:name="Target"/><w:r><w:t>A</w:t></w:r><w:bookmarkEnd w:id="1"/><w:bookmarkStart w:id="2" w:name="Target"/><w:r><w:t>B</w:t></w:r><w:bookmarkEnd w:id="2"/></w:p>', "BOOKMARK_AMBIGUOUS"),
            ('<w:p><w:bookmarkStart w:id="1" w:name="Target"/><w:r><w:t>A</w:t></w:r></w:p><w:p><w:r><w:t>B</w:t></w:r><w:bookmarkEnd w:id="1"/></w:p>', "BOOKMARK_SCOPE_UNSUPPORTED"),
            ('<w:tbl><w:tr><w:tc><w:p/><w:tbl><w:tr><w:tc><w:p><w:bookmarkStart w:id="1" w:name="Target"/><w:r><w:t>A</w:t></w:r><w:bookmarkEnd w:id="1"/></w:p></w:tc></w:tr></w:tbl></w:tc></w:tr></w:tbl>', "BOOKMARK_SCOPE_UNSUPPORTED"),
        ]
        for body, expected_code in cases:
            with self.subTest(code=expected_code), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / "bookmark.docx"
                write_document(path, body)
                _, record, _ = register_document("writer", str(path), tmp)
                with patch("wps_ai_agent_cli.writer_ops.create_backup") as backup, patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com") as com:
                    ok, _, errors, _ = writer_fill_bookmark(record["document_id"], "Target", "X", "bookmark-rejected", workspace=tmp)
                self.assertFalse(ok)
                self.assertEqual(errors[0]["code"], expected_code)
                backup.assert_not_called()
                com.assert_not_called()

    def test_readback_rejects_unexpected_change_outside_selected_paragraph(self):
        from tests.test_writer_ops import write_minimal_docx_paragraphs
        from zipfile import ZipFile

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "scope.docx"
            write_minimal_docx_paragraphs(path, ["needle first", "needle target", "needle last"])
            _, record, _ = register_document("writer", str(path), tmp)

            def corrupt_neighbor(*_args, **_kwargs):
                with ZipFile(path) as archive:
                    parts = {name: archive.read(name) for name in archive.namelist()}
                parts["word/document.xml"] = parts["word/document.xml"].replace(b"needle first", b"CORRUPTED first").replace(b"needle target", b"replacement target")
                with ZipFile(path, "w") as archive:
                    for name, content in parts.items():
                        archive.writestr(name, content)
                return {"ok": True, "errors": [], "data": {"backend": "test", "replace_count": 1}}

            with patch("wps_ai_agent_cli.writer_ops._run_writer_replace_com", side_effect=corrupt_neighbor):
                ok, result, errors, _ = writer_replace(
                    record["document_id"], "needle", "replacement", "unexpected-neighbor",
                    paragraph_index=2, workspace=tmp,
                )
            self.assertFalse(ok)
            self.assertFalse(result["readback_matches_expected"])
            self.assertEqual(result["changed_paragraph_indices"], [1])
            self.assertEqual(errors[0]["code"], "VALIDATION_FAILED")


@unittest.skipUnless(os.environ.get("WPS_AGENT_RUN_INTEGRATION") == "1", "Requires explicit local WPS integration run")
class WriterReplaceScopeIntegrationTests(unittest.TestCase):
    def test_scoped_replacements_preserve_neighboring_paragraphs_and_table(self):
        from docx import Document

        for replacement in ("x", "a much longer replacement", "", "needle plus", "needle"):
            with self.subTest(replacement=replacement), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / "scope.docx"
                document = Document()
                document.add_paragraph("needle before")
                document.add_table(rows=1, cols=1).cell(0, 0).text = "needle table"
                document.add_paragraph("needle needle NEEDLE")
                document.add_paragraph("needle after")
                document.save(path)
                _, record, _ = register_document("writer", str(path), tmp)
                ok, result, errors, _ = writer_replace(
                    record["document_id"], "needle", replacement, "scope-integration",
                    paragraph_index=2, workspace=tmp,
                )
                self.assertTrue(ok, (errors, result))
                self.assertEqual(result["backend_replace_count"], 2)
                self.assertEqual(result["expected_replace_count"], 2)
                self.assertTrue(result["readback_matches_expected"])
                self.assertEqual(docx_body_paragraphs(path), [
                    "needle before", f"{replacement} {replacement} NEEDLE", "needle after",
                ])
                self.assertEqual(docx_body_tables(path), [[["needle table"]]])
                self.assertTrue(result["backup"])

    def test_document_scope_readback_matches_body_and_table_text(self):
        from docx import Document
        from wps_ai_agent_cli.document_text import docx_body_story_paragraphs

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "whole-body.docx"
            document = Document()
            document.add_paragraph("needle first")
            document.add_table(rows=1, cols=1).cell(0, 0).text = "needle cell"
            document.add_paragraph("needle needle last")
            document.save(path)
            _, record, _ = register_document("writer", str(path), tmp)
            ok, result, errors, _ = writer_replace(
                record["document_id"], "needle", "x", "whole-body",
                workspace=tmp,
            )
            self.assertTrue(ok, (errors, result))
            self.assertEqual(result["expected_replace_count"], 4)
            self.assertTrue(result["readback_matches_expected"])
            self.assertEqual(docx_body_story_paragraphs(path), ["x first", "x cell", "x x last"])

    def test_bookmark_fill_preserves_range_and_reads_it_back(self):
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from wps_ai_agent_cli.writer_structure import read_body_bookmark_text

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bookmark.docx"
            document = Document()
            paragraph = document.add_paragraph("Dear ")
            run = paragraph.add_run("old")
            run.bold = True
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), "1")
            start.set(qn("w:name"), "Client")
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), "1")
            run._r.addprevious(start)
            run._r.addnext(end)
            paragraph.add_run("!")
            document.save(path)
            _, record, _ = register_document("writer", str(path), tmp)
            ok, result, errors, _ = writer_fill_bookmark(
                record["document_id"], "Client", "Northwind", "bookmark-wps", workspace=tmp,
            )
            self.assertTrue(ok, (errors, result))
            self.assertTrue(result["readback_passed"])
            bookmark, text = read_body_bookmark_text(path, "Client")
            self.assertTrue(bookmark["body_paragraph_range_supported"])
            self.assertEqual(text, "Northwind")
            reopened = Document(path)
            self.assertEqual(reopened.paragraphs[0].text, "Dear Northwind!")
            self.assertTrue(reopened.paragraphs[0].runs[1].bold)

    def test_paragraph_child_bookmark_fill_preserves_nested_neighbors(self):
        from wps_ai_agent_cli.writer_structure import read_body_bookmark_text, read_supported_bookmark_text

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "nested-copy.docx"
            shutil.copy2("fixtures/phase3/writer_nested_scope_wps_fixture.docx", path)
            tables_before = docx_body_tables(path)
            _, record, _ = register_document("writer", str(path), tmp)
            ok, result, errors, replayed = writer_fill_bookmark(
                record["document_id"], "BodyMark", "Updated", "paragraph-child-wps", workspace=tmp,
            )
            self.assertTrue(ok, (errors, result))
            self.assertFalse(replayed)
            self.assertTrue(result["readback_passed"])
            self.assertTrue(result["body_tail_normalized"])
            self.assertTrue(Path(result["backup"]["backup_path"]).is_file())
            self.assertEqual(read_body_bookmark_text(path, "BodyMark")[1], "Updated")
            self.assertEqual(docx_body_tables(path), tables_before)
            for name, expected in (("CellMark", "CellValue"), ("HeaderMark", "HeaderValue"),
                                   ("FooterMark", "FooterValue")):
                self.assertEqual(read_supported_bookmark_text(path, name)[1], expected)
            ok, repeated, errors, replayed = writer_fill_bookmark(
                record["document_id"], "BodyMark", "Updated", "paragraph-child-wps", workspace=tmp,
            )
            self.assertTrue(ok, errors)
            self.assertTrue(replayed)
            self.assertEqual(repeated["readback_passed"], True)

if __name__ == "__main__":
    unittest.main()
