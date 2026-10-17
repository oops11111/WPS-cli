"""The injected read-back failure (P7-002) and the `undo` subcommand of the Windows check.

On Windows the scenarios run against real WPS. Here a fake WPS stands in: it saves the write to the fixture,
reports it done, and on the undo puts the original file back, the way WPS saves the reverted document.
"""

from __future__ import annotations

import importlib.util
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from tests.test_windows_check import FakeMachine, windows_check
from wps_ai_agent_cli import open_documents

REPO_ROOT = Path(__file__).resolve().parents[1]


def _load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


injection = _load("undo_injection", REPO_ROOT / "scripts" / "windows_check" / "undo_injection.py")
runner_module = _load("run_agent_eval", REPO_ROOT / "scripts" / "agent_eval" / "run_agent_eval.py")


def com_ok(data):
    return {"ok": True, "errors": [], "data": data}


class OfflineRunner:
    """The eval runner's fixture and text helpers, with an open in WPS that only registers the document."""

    write_fixture = staticmethod(runner_module.write_fixture)
    document_text = staticmethod(runner_module.document_text)
    file_sha256 = staticmethod(runner_module.file_sha256)
    WPS_LAUNCH = "default"

    def __init__(self, setup_failure=None):
        self.setup_failure = setup_failure
        self.released = []

    def run_setup(self, steps, workspace, fixture, component, options):
        self.steps = (steps, options)
        if self.setup_failure:
            return runner_module.file_sha256(fixture), None, self.setup_failure
        from wps_ai_agent_cli.sessions import register_document

        register_document(component, str(fixture), workspace)
        return runner_module.file_sha256(fixture), lambda: self.released.append(fixture.name) or "fixture closed", None


def replace_on_disk(path, text):
    """Save the fixture with DRAFT replaced, as WPS saves the write."""
    if path.suffix == ".docx":
        from docx import Document

        document = Document(path)
        for paragraph in document.paragraphs:
            if "DRAFT" in paragraph.text:
                paragraph.text = paragraph.text.replace("DRAFT", text)
        document.save(path)
    else:
        from pptx import Presentation

        presentation = Presentation(path)
        for shape in presentation.slides[0].shapes:
            if "DRAFT" in shape.text_frame.text:
                shape.text_frame.text = shape.text_frame.text.replace("DRAFT", text)
        presentation.save(path)


class FakeWps:
    """Writer or Presentation as the scenario sees it: the selection is "DRAFT" and writes are saved."""

    def __init__(self, component, *, undo_status="reverted", undo_restores=True):
        self.component = component
        self.undo_status = undo_status
        self.undo_restores = undo_restores
        self.actions = []

    def __call__(self, action, params):
        self.actions.append(action)
        path = Path(params["path"])
        if action == "selection-read":
            if self.component == "writer":
                return com_ok(
                    {
                        "ok": True,
                        "status": "ok",
                        "saved": True,
                        "start": 22,
                        "end": 27,
                        "selection_text": "DRAFT",
                        "paragraph_index": 2,
                    }
                )
            return com_ok(
                {
                    "ok": True,
                    "status": "ok",
                    "saved": True,
                    "read_only": False,
                    "selection_type": 3,
                    "selection_kind": "text",
                    "selection_text": "DRAFT",
                    "slide_index": 1,
                    "shape_name": "TextBox 2",
                    "shape_count": 0,
                }
            )
        if action == "selection-replace":
            self.original = path.read_bytes()
            replace_on_disk(path, params["text"])  # WPS saved the write
            data = {"ok": True, "status": "ok", "read_back_text": params["text"]}
            if self.component == "presentation":
                data.update(
                    {"slide_index": 1, "shape_name": "TextBox 2", "text_start": 9, "shape_text_before": "Status: DRAFT"}
                )
            return com_ok(data)
        if action == "selection-revert":
            if self.undo_restores:
                path.write_bytes(self.original)
            return com_ok({"ok": True, "status": self.undo_status, "read_back_text": params.get("original_text")})
        raise AssertionError(f"unexpected action {action}")


class FakeObjectWps:
    """Spreadsheets with a text box or a chart title selected; the write and the undo change the saved file."""

    def __init__(self, scenario, selection_text=None):
        self.selection = injection.SCENARIOS[scenario]["selection"]
        expected = injection.SCENARIOS[scenario]["arguments"]["expected_selection_text"]
        self.selection_text = expected if selection_text is None else selection_text
        self.actions = []

    def __call__(self, action, params):
        self.actions.append(action)
        path = Path(params["path"])
        if action == "selection-read":
            part = self.selection.get("part")
            return com_ok(
                {
                    "ok": True,
                    "status": "ok",
                    "saved": True,
                    "read_only": False,
                    "selection_kind": "non_range",
                    "sheet_name": "Status",
                    "object_kind": "chart" if part else "shape",
                    "object_name": self.selection["object"],
                    "has_text_frame": not part,
                    "has_chart_title": part == "chart_title",
                    "has_axis_title": part == "axis_title",
                    "has_series_values": part == "series",
                    "series_name": "Sales" if part == "series" else None,
                    "selection_text": self.selection_text,
                }
            )
        if action.endswith("-replace"):
            self.original = path.read_bytes()
            part = "xl/charts/chart1.xml" if "chart" in action else "xl/drawings/drawing1.xml"
            written = "\t".join(params["values"]) if "values" in params else params["text"]
            if "values" in params:
                # Literal values replace the series' reference to B2:B3, as WPS saves them.
                rewrite_part(path, part, b"<numRef><f>'Status'!$B$2:$B$3</f>", b"<numLit><f>")
            else:
                rewrite_part(path, part, b"<a:t>DRAFT</a:t>", f"<a:t>{written}</a:t>".encode(), count=1)
            return com_ok({"ok": True, "status": "ok", "read_back_text": written})
        if action.endswith("-revert"):
            self.revert_params = params
            path.write_bytes(self.original)
            return com_ok({"ok": True, "status": "reverted", "read_back_text": params["original_text"]})
        raise AssertionError(f"unexpected action {action}")


def rewrite_part(path, part, old, new, count=-1):
    import zipfile

    with zipfile.ZipFile(path) as archive:
        items = [(item, archive.read(item.filename)) for item in archive.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for item, data in items:
            archive.writestr(item, data.replace(old, new, count) if item.filename == part else data)


class ScenarioTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)

    def run_scenario(self, name, wps, runner=None, saved=True):
        runner = runner or OfflineRunner()
        component = injection.SCENARIOS[name]["component"]

        def listed(component_name):
            documents = [{"full_name": str(path), "saved": saved} for path in self.root.rglob("*.*x")]
            return True, {"documents": documents}, []

        with (
            patch.object(open_documents, injection.RUNNER_FUNCTIONS[component], side_effect=wps),
            patch.object(open_documents, "list_open_documents", side_effect=listed),
        ):
            return injection.run_scenario(name, runner, self.root), runner

    def test_both_scenarios_undo_the_injected_failure_and_pass(self):
        for name in ("writer-replace", "presentation-replace"):
            with self.subTest(scenario=name):
                component = injection.SCENARIOS[name]["component"]
                wps = FakeWps(component)
                result, runner = self.run_scenario(name, wps)
                passed, conditions = injection.verdict(result)
                self.assertTrue(passed, conditions)
                self.assertEqual(wps.actions, ["selection-read", "selection-replace", "selection-revert"])
                self.assertEqual(result["error_codes"], ["VALIDATION_FAILED", "MUTATION_ROLLED_BACK"])
                self.assertEqual(result["real_read_back_text"], "FINAL")
                hashes = result["hashes"]
                self.assertNotEqual(hashes["before_write"], hashes["before_undo"])
                self.assertEqual(hashes["before_write"], hashes["now"])
                self.assertEqual(result["text_after"], result["text_before"])
                self.assertIn("Status: DRAFT", result["text_before"])
                self.assertEqual(result["cleanup"], "fixture closed")
                steps, options = runner.steps
                self.assertEqual(steps, ["register", "open_in_wps", "select_text"])
                self.assertEqual(options, {"selection": {"text": "DRAFT"}})

    def test_a_shape_and_a_chart_title_are_undone_in_wps(self):
        for name, write, revert in (
            ("spreadsheet-shape-text-replace", "selection-shape-replace", "selection-shape-revert"),
            ("spreadsheet-chart-title-replace", "selection-chart-title-replace", "selection-chart-title-revert"),
            ("spreadsheet-axis-title-replace", "selection-axis-title-replace", "selection-axis-title-revert"),
            (
                "spreadsheet-series-values-replace",
                "selection-series-values-replace",
                "selection-series-values-revert",
            ),
        ):
            with self.subTest(scenario=name):
                wps = FakeObjectWps(name)
                result, runner = self.run_scenario(name, wps)
                passed, conditions = injection.verdict(result)
                self.assertTrue(passed, conditions)
                self.assertEqual(wps.actions, ["selection-read", write, revert])
                written = injection.SCENARIOS[name]["arguments"]["text"]
                self.assertEqual(wps.revert_params["written_text"], written)
                self.assertEqual(wps.revert_params["read_back_text"], injection.INJECTED_READ_BACK)
                self.assertEqual(result["text_after"], result["text_before"])
                self.assertIn("DRAFT", result["text_before"])
                _steps, options = runner.steps
                self.assertEqual(options, {"selection": injection.SCENARIOS[name]["selection"]})

    def test_a_refusal_before_the_write_records_what_the_selection_was_read_as(self):
        wps = FakeObjectWps("spreadsheet-chart-title-replace", selection_text="")
        result, _runner = self.run_scenario("spreadsheet-chart-title-replace", wps)
        self.assertEqual(wps.actions, ["selection-read"])
        self.assertEqual(result["error_codes"], ["SELECTION_MISMATCH"])
        self.assertEqual(result["selection"]["object_kind"], "chart")
        self.assertTrue(result["selection"]["has_chart_title"])
        self.assertEqual(result["selection"]["selection_text"], "")

    def test_an_undo_that_did_not_happen_fails_the_scenario(self):
        wps = FakeWps("writer", undo_status="revert_mismatch", undo_restores=False)
        result, _runner = self.run_scenario("writer-replace", wps)
        passed, conditions = injection.verdict(result)
        self.assertFalse(passed)
        self.assertIn("- FAILED: the command reported VALIDATION_FAILED, then MUTATION_ROLLED_BACK", conditions)
        self.assertIn("- FAILED: the text on disk is the text before the write", conditions)
        self.assertEqual(result["error_codes"][-1], "MUTATION_ROLLBACK_IN_WPS_FAILED")
        self.assertEqual(result["cleanup"], "fixture closed")

    def test_a_document_left_unsaved_in_wps_fails_the_scenario(self):
        result, _runner = self.run_scenario("writer-replace", FakeWps("writer"), saved=False)
        passed, conditions = injection.verdict(result)
        self.assertFalse(passed)
        self.assertEqual(
            [line for line in conditions if line.startswith("- FAILED")],
            ["- FAILED: WPS reports the document saved after the undo"],
        )

    def test_a_setup_failure_is_recorded_and_nothing_is_written(self):
        wps = FakeWps("presentation")
        result, _runner = self.run_scenario(
            "presentation-replace", wps, runner=OfflineRunner(setup_failure="select_text: WPS selected ''")
        )
        self.assertEqual(wps.actions, [])
        passed, conditions = injection.verdict(result)
        self.assertFalse(passed)
        self.assertEqual(conditions, ["- FAILED: setup: select_text: WPS selected ''"])

    def test_the_injection_changes_only_a_write_that_wps_reported_done(self):
        def real(action, params):
            return com_ok({"ok": True, "status": params.get("status", "ok"), "read_back_text": "FINAL"})

        with patch.object(open_documents, "_run_writer_attached_com", side_effect=real):
            with injection.read_back_failure("writer") as record:
                read = open_documents._run_writer_attached_com("selection-read", {})
                refused = open_documents._run_writer_attached_com("selection-replace", {"status": "selection_changed"})
                undo = open_documents._run_writer_attached_com("selection-revert", {})
                self.assertFalse(record["injected"])
                write = open_documents._run_writer_attached_com("selection-replace", {})
            after = open_documents._run_writer_attached_com("selection-replace", {})
        self.assertEqual(read["data"]["read_back_text"], "FINAL")
        self.assertEqual(refused["data"]["read_back_text"], "FINAL")
        self.assertEqual(undo["data"]["read_back_text"], "FINAL")
        self.assertEqual(write["data"]["read_back_text"], injection.INJECTED_READ_BACK)
        self.assertTrue(record["injected"])
        self.assertEqual(after["data"]["read_back_text"], "FINAL")


class FakeInjection:
    SCENARIOS = injection.SCENARIOS
    verdict = staticmethod(injection.verdict)

    def __init__(self, blocked=False):
        self.blocked = blocked
        self.ran = []

    def run_scenario(self, name, runner, root):
        self.ran.append(name)
        codes = ["VALIDATION_FAILED", "MUTATION_ROLLBACK_BLOCKED" if self.blocked else "MUTATION_ROLLED_BACK"]
        return {
            "scenario": name,
            "injected": True,
            "error_codes": codes,
            "errors": [{"code": code, "message": f"{code} message"} for code in codes],
            "rollback": {"status": "reverted", "rolled_back_in_wps": not self.blocked, "content_matches_backup": True},
            "real_read_back_text": "FINAL",
            "hashes": {"before_write": "a" * 64, "before_undo": "b" * 64, "after_undo": "c" * 64, "now": "c" * 64},
            "text_before": "Status: DRAFT",
            "text_after": "Status: DRAFT",
            "wps_saved_after": True,
            "cleanup": "fixture closed; processes stopped 12",
        }


class UndoCommandTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.out = Path(tmp.name) / "out"

    def run_undo(self, *extra, machine=None, fake=None):
        fake = fake or FakeInjection()
        code = windows_check.main(
            ["undo", "--out", str(self.out), *extra],
            machine=machine or FakeMachine(),
            runner=OfflineRunner(),
            injection=fake,
        )
        return code, fake

    def test_every_scenario_runs_and_the_results_are_written(self):
        code, fake = self.run_undo()
        self.assertEqual(code, 0)
        self.assertEqual(
            fake.ran,
            [
                "writer-replace",
                "presentation-replace",
                "spreadsheet-shape-text-replace",
                "spreadsheet-chart-title-replace",
                "spreadsheet-axis-title-replace",
                "spreadsheet-series-values-replace",
            ],
        )
        text = (self.out / "RESULTS.md").read_text(encoding="utf-8")
        self.assertIn("# Windows check of the undo in WPS", text)
        self.assertIn("**6/6 scenarios passed.**", text)
        self.assertIn(
            "| `writer-replace` | pass | `VALIDATION_FAILED`, `MUTATION_ROLLED_BACK` | `aaaaaaaaaaaa` |", text
        )
        self.assertIn("fixture closed; processes stopped 12", text)
        self.assertIn("python scripts/windows_check/run_windows_check.py undo\n", text)

    def test_a_blocked_rollback_fails_the_run(self):
        code, _fake = self.run_undo("--scenario", "writer-replace", fake=FakeInjection(blocked=True))
        self.assertEqual(code, 1)
        text = (self.out / "RESULTS.md").read_text(encoding="utf-8")
        self.assertIn("**0/1 scenarios passed.**", text)
        self.assertIn("- FAILED: no MUTATION_ROLLBACK_BLOCKED", text)
        self.assertIn("undo --scenario writer-replace", text)

    def test_failed_prerequisites_run_nothing_and_claude_is_not_needed(self):
        code, fake = self.run_undo(machine=FakeMachine(processes=[{"name": "wps.exe", "id": 4}]))
        self.assertEqual(code, 1)
        self.assertEqual(fake.ran, [])
        code, fake = self.run_undo(machine=FakeMachine(claude=(False, "API Error")))
        self.assertEqual(code, 0)


if __name__ == "__main__":
    unittest.main()
