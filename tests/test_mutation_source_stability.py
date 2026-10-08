import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from tests.test_writer_structure import write_document
from wps_ai_agent_cli.backups import create_backup, guarded_com_mutation, verify_post_com_source
from wps_ai_agent_cli.sessions import register_document
from wps_ai_agent_cli.operations import get_operation
from wps_ai_agent_cli.presentation_ops import presentation_replace
from wps_ai_agent_cli.presentation_text import count_text_in_pptx
from wps_ai_agent_cli.spreadsheet_ops import rename_spreadsheet_sheet
from wps_ai_agent_cli import writer_ops
from wps_ai_agent_cli.writer_ops import writer_fill_bookmark


def replace_with_same_bytes(path: Path) -> None:
    replacement = path.with_name(f"replacement-{path.name}")
    replacement.write_bytes(path.read_bytes())
    os.replace(replacement, path)


class MutationSourceStabilityTests(unittest.TestCase):
    def test_writer_replacement_after_backup_blocks_com(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "writer.docx"
            write_document(path, '<w:p><w:bookmarkStart w:id="1" w:name="Target"/>'
                                 '<w:r><w:t>old</w:t></w:r><w:bookmarkEnd w:id="1"/></w:p>')
            _, record, _ = register_document("writer", str(path), tmp)

            def replace_after_backup(*args, **kwargs):
                result = create_backup(*args, **kwargs)
                self.assertTrue(result[0], result[2])
                replace_with_same_bytes(path)
                return result

            with patch("wps_ai_agent_cli.writer_ops.create_backup", side_effect=replace_after_backup), \
                    patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com") as com:
                ok, _, errors, _ = writer_fill_bookmark(
                    record["document_id"], "Target", "new", "writer-after-backup", workspace=tmp,
                )
            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "DOCUMENT_CHANGED_AFTER_BACKUP")
            com.assert_not_called()

    def test_spreadsheet_replacement_after_backup_blocks_com(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.xlsx"
            workbook = Workbook()
            workbook.active.title = "Original"
            workbook.save(path)
            workbook.close()
            _, record, _ = register_document("spreadsheets", str(path), tmp)

            def replace_after_backup(*args, **kwargs):
                result = create_backup(*args, **kwargs)
                self.assertTrue(result[0], result[2])
                replace_with_same_bytes(path)
                return result

            with patch("wps_ai_agent_cli.spreadsheet_ops.create_backup", side_effect=replace_after_backup), \
                    patch("wps_ai_agent_cli.spreadsheet_ops._run_spreadsheet_rename_sheet_com") as com:
                ok, _, errors, _ = rename_spreadsheet_sheet(
                    record["document_id"], "Original", "Renamed", "sheet-after-backup", workspace=tmp,
                )
            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "DOCUMENT_CHANGED_AFTER_BACKUP")
            com.assert_not_called()

    def test_noop_does_not_rebind_externally_changed_spreadsheet(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.xlsx"
            workbook = Workbook()
            workbook.active.title = "Original"
            workbook.save(path)
            workbook.close()
            _, record, _ = register_document("spreadsheets", str(path), tmp)
            replace_with_same_bytes(path)
            ok, result, errors, _ = rename_spreadsheet_sheet(
                record["document_id"], "Original", "Original", "noop-after-replacement", workspace=tmp,
            )
            self.assertTrue(ok, errors)
            self.assertFalse(result["saved"])
            backed_up, _, backup_errors, _ = create_backup(
                record["document_id"], "after-noop-backup", workspace=tmp,
            )
            self.assertFalse(backed_up)
            self.assertEqual(backup_errors[0]["code"], "DOCUMENT_IDENTITY_CHANGED")

    def test_post_com_change_and_backup_tamper_are_detected(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "source.docx"
            path.write_bytes(b"original")
            _, record, _ = register_document("writer", str(path), tmp)
            ok, backup, errors, _ = create_backup(record["document_id"], "guard-backup", workspace=tmp)
            self.assertTrue(ok, errors)

            def saved_operation():
                path.write_bytes(b"saved by WPS")
                return {"ok": True, "errors": [], "data": {"backend": "controlled"}}

            mutation = guarded_com_mutation(path, backup, saved_operation)
            self.assertTrue(mutation["ok"], mutation["errors"])
            self.assertIsNone(verify_post_com_source(path, mutation))
            path.write_bytes(b"external change")
            self.assertEqual(verify_post_com_source(path, mutation)["code"], "DOCUMENT_CHANGED_AFTER_COM")

            register_document("writer", str(path), tmp)
            ok, backup, errors, _ = create_backup(record["document_id"], "tamper-backup", workspace=tmp)
            self.assertTrue(ok, errors)
            Path(backup["backup_path"]).write_bytes(b"tampered")
            mutation = guarded_com_mutation(path, backup, lambda: self.fail("COM must not run"))
            self.assertFalse(mutation["ok"])
            self.assertEqual(mutation["errors"][0]["code"], "BACKUP_CHANGED_AFTER_CREATION")

    def test_writer_post_com_replacement_fails_even_when_text_matches(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "writer.docx"
            write_document(path, '<w:p><w:bookmarkStart w:id="1" w:name="Target"/>'
                                 '<w:r><w:t>old</w:t></w:r><w:bookmarkEnd w:id="1"/></w:p>')
            _, record, _ = register_document("writer", str(path), tmp)

            def simulated_wps(*_args):
                with ZipFile(path) as archive:
                    parts = {name: archive.read(name) for name in archive.namelist()}
                parts["word/document.xml"] = parts["word/document.xml"].replace(b">old<", b">new<")
                with ZipFile(path, "w") as archive:
                    for name, payload in parts.items():
                        archive.writestr(name, payload)
                return {"ok": True, "errors": [], "data": {"backend": "controlled"}}

            table_reads = 0
            original_tables = writer_ops.docx_body_tables

            def replace_during_readback(source):
                nonlocal table_reads
                table_reads += 1
                value = original_tables(source)
                if table_reads == 2:
                    replace_with_same_bytes(path)
                return value

            with patch("wps_ai_agent_cli.writer_ops._run_writer_bookmark_fill_com", side_effect=simulated_wps), \
                    patch("wps_ai_agent_cli.writer_ops.docx_body_tables", side_effect=replace_during_readback):
                ok, result, errors, _ = writer_fill_bookmark(
                    record["document_id"], "Target", "new", "writer-after-com", workspace=tmp,
                )
            self.assertFalse(ok)
            self.assertTrue(result["readback_passed"])
            self.assertEqual(errors[0]["code"], "DOCUMENT_CHANGED_AFTER_COM")
            self.assertIsNone(get_operation("writer-after-com", tmp))

    @unittest.skipUnless(os.environ.get("WPS_AGENT_RUN_INTEGRATION") == "1", "set WPS_AGENT_RUN_INTEGRATION=1 to launch local WPS")
    def test_presentation_wps_mutation_retains_verified_backup(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "slides.pptx"
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "alpha beta"
            deck.save(path)
            _, record, _ = register_document("presentation", str(path), tmp)
            ok, result, errors, _ = presentation_replace(
                record["document_id"], "alpha", "gamma", "presentation-stability", workspace=tmp,
            )
            self.assertTrue(ok, (errors, result))
            self.assertEqual(count_text_in_pptx(path, "gamma"), 1)
            self.assertTrue(Path(result["backup"]["backup_path"]).is_file())
            backed_up, _, backup_errors, _ = create_backup(
                record["document_id"], "presentation-after-com", workspace=tmp,
            )
            self.assertTrue(backed_up, backup_errors)


if __name__ == "__main__":
    unittest.main()
