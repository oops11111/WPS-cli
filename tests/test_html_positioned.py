"""HTML laid out by the browser (P13-004, P13-006): computed styles for DOCX, positioned slides for PPTX.

These need Node.js, Playwright and Edge or Chromium (WPS_AGENT_EDGE); they are skipped without them.
"""

from __future__ import annotations

import os
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from pptx import Presentation

from wps_ai_agent_cli.fidelity import fidelity_check
from wps_ai_agent_cli.html_office import convert_html_to_office
from wps_ai_agent_cli.html_styles import has_browser

ROOT = Path(__file__).resolve().parents[1]
_BROWSER = mock.patch.dict(os.environ, {"WPS_AGENT_NO_BROWSER": "0"})


def setUpModule():
    _BROWSER.start()


def tearDownModule():
    _BROWSER.stop()


@unittest.skipUnless(has_browser(), "the browser layout needs Node.js, Playwright and Edge or Chromium")
class PositionedTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.work = Path(self._tmp.name)
        for folder in ("fidelity", "html_complex"):
            for item in (ROOT / "fixtures" / folder).iterdir():
                if item.is_file() and not (self.work / item.name).exists():
                    shutil.copyfile(item, self.work / item.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_slides_are_placed_where_the_browser_lays_the_sections_out(self):
        ok, data, errors = convert_html_to_office(self.work / "launch.html", self.work / "launch.pptx", "pptx")
        self.assertTrue(ok, errors)
        presentation = Presentation(self.work / "launch.pptx")
        self.assertEqual((presentation.slide_width, presentation.slide_height), (12192000, 6858000))
        cover = presentation.slides[0]
        title = cover.shapes.title
        subtitle = next(shape for shape in cover.shapes if shape.has_text_frame and "上海" in shape.text_frame.text)
        self.assertFalse(subtitle.is_placeholder)  # a text box at the paragraph's place, below the heading
        self.assertGreaterEqual(subtitle.top, title.top + title.height)
        report = fidelity_check(self.work / "launch.html", self.work / "launch.pptx", layout=False)[1]
        self.assertTrue(report["passed"], report["aspects"])

    def test_appearance_mode_draws_the_decorative_boxes(self):
        ok, data, errors = convert_html_to_office(
            self.work / "training.html", self.work / "training.pptx", "pptx", fidelity="appearance"
        )
        self.assertTrue(ok, errors)
        self.assertGreaterEqual(data["mapped_objects"].get("shapes", 0), 1)
        cover = Presentation(self.work / "training.pptx").slides[0]
        self.assertEqual(str(cover.background.fill.fore_color.rgb), "1F4E79")
        report = fidelity_check(self.work / "training.html", self.work / "training.pptx", layout=False)[1]
        self.assertTrue(report["passed"], report["aspects"])

    def test_docx_takes_the_computed_styles(self):
        ok, data, errors = convert_html_to_office(self.work / "newsletter.html", self.work / "news.docx", "docx")
        self.assertTrue(ok, errors)
        self.assertEqual(data["style_source"], "browser")
        report = fidelity_check(self.work / "newsletter.html", self.work / "news.docx", layout=False)[1]
        self.assertTrue(report["passed"], report["aspects"])


class WithoutBrowserTests(unittest.TestCase):
    def test_appearance_mode_needs_the_browser(self):
        with tempfile.TemporaryDirectory() as tmp, mock.patch.dict(os.environ, {"WPS_AGENT_NO_BROWSER": "1"}):
            source = Path(tmp) / "deck.html"
            shutil.copyfile(ROOT / "fixtures" / "fidelity" / "training.html", source)
            ok, _data, errors = convert_html_to_office(source, Path(tmp) / "deck.pptx", "pptx", fidelity="appearance")
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "BROWSER_BACKEND_UNAVAILABLE")


if __name__ == "__main__":
    unittest.main()
