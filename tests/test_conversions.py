"""Phase 13 conversions: CSV/TSV, Markdown, Office <-> HTML, cross-type, PDF, WPS formats and the convert command."""

from __future__ import annotations

import io
import json
import os
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation

from wps_ai_agent_cli import cli, convert, cross_type, markdown_convert, office_html, pdf_convert, tabular, wps_convert
from wps_ai_agent_cli.fidelity import fidelity_check, read_model, score_data
from wps_ai_agent_cli.html_office import convert_html_to_office, convert_html_to_xlsx
from wps_ai_agent_cli.mcp_adapter import build_cli_argv

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "fixtures" / "fidelity"
COMPLEX = ROOT / "fixtures" / "html_complex"
_NO_BROWSER = mock.patch.dict(os.environ, {"WPS_AGENT_NO_BROWSER": "1"})


def setUpModule():
    _NO_BROWSER.start()


def tearDownModule():
    _NO_BROWSER.stop()


class _Workspace(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.work = Path(self._tmp.name)
        for folder in (FIXTURES, COMPLEX):
            for item in folder.iterdir():
                if item.is_file() and not (self.work / item.name).exists():
                    shutil.copyfile(item, self.work / item.name)

    def tearDown(self):
        self._tmp.cleanup()

    def assertPasses(self, source, target, **options):
        ok, report, errors = fidelity_check(self.work / source, self.work / target, layout=False, **options)
        self.assertTrue(ok, errors)
        self.assertTrue(report["passed"], report["aspects"])
        return report


class TabularTests(_Workspace):
    def test_gbk_csv_becomes_a_typed_workbook_that_keeps_ids_and_formula_text(self):
        ok, data, errors = tabular.csv_to_xlsx(self.work / "sales_gbk.csv", self.work / "sales.xlsx")
        self.assertTrue(ok, errors)
        self.assertEqual((data["encoding"], data["delimiter"], data["header_row"]), ("gb18030", ",", True))
        sheet = load_workbook(self.work / "sales.xlsx").active
        self.assertEqual(sheet["A2"].value, "华东")
        self.assertEqual(sheet["B2"].value, "007")  # an id stays text
        self.assertEqual(sheet["C2"].value, 1200)
        self.assertEqual((sheet["C3"].value, sheet["C3"].number_format), (3.5, "0.00"))
        self.assertEqual((sheet["D2"].value, sheet["D2"].number_format), (0.125, "0.0%"))
        self.assertEqual((sheet["E2"].value, sheet["E2"].number_format), (1234.5, "#,##0.00"))
        self.assertEqual(sheet["F4"].data_type, "s")  # =SUM(...) in a CSV is text, not a formula
        self.assertTrue(sheet["A1"].font.bold)
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertPasses("sales_gbk.csv", "sales.xlsx")

    def test_tsv_and_back_to_csv_keep_every_value(self):
        ok, _data, errors = tabular.csv_to_xlsx(self.work / "staff_utf8.tsv", self.work / "staff.xlsx")
        self.assertTrue(ok, errors)
        ok, data, errors = tabular.xlsx_to_csv(self.work / "staff.xlsx", self.work / "staff.csv")
        self.assertTrue(ok, errors)
        text = (self.work / "staff.csv").read_text(encoding="utf-8-sig")
        self.assertIn("000127,张伟,研发,2021-03-15,18500", text)
        self.assertIn("16800.50", text)
        self.assertIn("Émile O'Brien", text)
        self.assertPasses("staff_utf8.tsv", "staff.xlsx")
        self.assertPasses("staff.xlsx", "staff.csv")

    def test_formulas_without_a_saved_result_are_written_and_reported(self):
        workbook = Workbook()
        workbook.active["A1"] = 2
        workbook.active["A2"] = "=A1*3"
        workbook.save(self.work / "calc.xlsx")
        ok, data, errors = tabular.xlsx_to_csv(self.work / "calc.xlsx", self.work / "calc.csv")
        self.assertTrue(ok, errors)
        self.assertIn("=A1*3", (self.work / "calc.csv").read_text(encoding="utf-8-sig"))
        self.assertTrue(any("no saved result" in w for w in data["warnings"]))

    def test_typed_value_rules(self):
        self.assertEqual(tabular.typed_value("0012"), ("0012", None))
        self.assertEqual(tabular.typed_value("12,345"), (12345, "#,##0"))
        self.assertEqual(tabular.typed_value(" 5"), (" 5", None))
        self.assertEqual(tabular.typed_value("1e5"), ("1e5", None))


class MarkdownTests(_Workspace):
    def test_builtin_reader_covers_the_common_markdown(self):
        body = markdown_convert._builtin((FIXTURES / "guide.md").read_text(encoding="utf-8"))
        for fragment in (
            "<h1>项目周报</h1>",
            "<strong>核心模块</strong>",
            "<s>20%</s>",
            '<a href="https://example.com/design">设计文档</a>',
            "<ul><li>数据层<ul><li>缓存命中率 92%</li><li>索引重建</li></ul></li><li>接口层</li></ul>",
            "<ol><li>第一步：评审",
            '<th style="text-align:center">负责人</th>',
            '<td style="text-align:right">100%</td>',
            "<blockquote><p>注意：下周冻结代码。</p></blockquote>",
            '<pre><code class="language-python">print(&quot;hello&quot;)</code></pre>'.replace("&quot;", '"'),
            '<img src="chart.png" alt="趋势图">',
            "<hr>",
        ):
            self.assertIn(fragment, body)
        setext = markdown_convert._builtin((FIXTURES / "notes.md").read_text(encoding="utf-8"))
        self.assertIn("<h1>会议纪要</h1>", setext)
        self.assertIn("<h2>参会人</h2>", setext)

    def test_markdown_to_docx_and_back(self):
        ok, data, errors = markdown_convert.md_to_office(self.work / "guide.md", self.work / "guide.docx")
        self.assertTrue(ok, errors)
        self.assertPasses("guide.md", "guide.docx")
        document = Document(self.work / "guide.docx")
        self.assertIn("Quote", [p.style.name for p in document.paragraphs])
        ok, data, errors = markdown_convert.to_markdown(self.work / "guide.docx", self.work / "back.md")
        self.assertTrue(ok, errors)
        text = (self.work / "back.md").read_text(encoding="utf-8")
        for fragment in (
            "# 项目周报",
            "**核心模块**",
            "~~20%~~",
            "  - 缓存命中率 92%",
            "| 存储 | 张三 | 100% |",
            "> 注意：下周冻结代码。",
            '```\nprint("hello")\n```',
            "![趋势图](back_files/image001.png)",
            "---",
        ):
            self.assertIn(fragment, text)
        self.assertTrue((self.work / "back_files" / "image001.png").is_file())
        self.assertPasses("guide.docx", "back.md")

    def test_html_tables_with_merged_cells_are_flattened_with_a_warning(self):
        ok, data, errors = markdown_convert.to_markdown(self.work / "contract.html", self.work / "contract.md")
        self.assertTrue(ok, errors)
        self.assertTrue(any("merged cells" in w for w in data["warnings"]))
        self.assertIn(
            "| 签约 | 30% | 90,000 | 合同生效后 5 个工作日内 |", (self.work / "contract.md").read_text("utf-8")
        )


class OfficeHtmlTests(_Workspace):
    def test_docx_round_trip_through_clean_html(self):
        convert_html_to_office(self.work / "report.html", self.work / "report.docx", "docx")
        ok, data, errors = office_html.office_to_html(self.work / "report.docx", self.work / "back.html")
        self.assertTrue(ok, errors)
        page = (self.work / "back.html").read_text(encoding="utf-8")
        self.assertIn("<thead>", page)
        self.assertIn('rowspan="2"', page)
        self.assertIn("data:image/png;base64,", page)
        self.assertNotIn("class=", page.split("<body>")[1])
        self.assertPasses("report.docx", "back.html")
        convert_html_to_office(self.work / "back.html", self.work / "again.docx", "docx")
        self.assertPasses("report.docx", "again.docx")

    def test_pptx_becomes_positioned_sections_with_animations_transitions_and_notes(self):
        convert_html_to_office(self.work / "launch.html", self.work / "launch.pptx", "pptx")
        ok, data, errors = office_html.office_to_html(self.work / "launch.pptx", self.work / "deck.html", "files")
        self.assertTrue(ok, errors)
        page = (self.work / "deck.html").read_text(encoding="utf-8")
        self.assertIn("@keyframes wps-entrance-fly-left", page)
        self.assertIn('data-transition="', page)
        self.assertIn('<aside class="notes">', page)
        self.assertIn("position:absolute", page)
        self.assertTrue(data["image_folder"] and Path(data["image_folder"]).is_dir())
        self.assertPasses("launch.pptx", "deck.html")

    def test_xlsx_formulas_round_trip_through_data_formula(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "预算"
        sheet.append(["项目", "金额"])
        sheet.append(["人力", 100])
        sheet.append(["合计", "=SUM(B2:B2)"])
        sheet.merge_cells("A5:B5")
        sheet["A5"] = "备注"
        sheet.freeze_panes = "A2"
        workbook.save(self.work / "calc.xlsx")
        ok, data, errors = office_html.office_to_html(self.work / "calc.xlsx", self.work / "calc.html")
        self.assertTrue(ok, errors)
        page = (self.work / "calc.html").read_text(encoding="utf-8")
        self.assertIn('data-formula="=SUM(B2:B2)"', page)
        self.assertIn('colspan="2"', page)
        self.assertIn("<thead>", page)
        ok, data, errors = convert_html_to_xlsx(self.work / "calc.html", self.work / "again.xlsx")
        self.assertTrue(ok, errors)
        self.assertEqual(data["formulas_restored"], 1)
        again = load_workbook(self.work / "again.xlsx").active
        self.assertEqual(again["B3"].value, "=SUM(B2:B2)")
        self.assertEqual(again.freeze_panes, "A2")

    def test_html_to_xlsx_applies_cell_styles(self):
        ok, data, errors = convert_html_to_xlsx(self.work / "contract.html", self.work / "contract.xlsx")
        self.assertTrue(ok, errors)
        self.assertTrue({"fill", "border", "font"} <= set(data["applied_styles"]))
        sheet = load_workbook(self.work / "contract.xlsx").worksheets[1]
        self.assertEqual(sheet["A1"].fill.start_color.rgb[-6:], "DDE6F0")
        self.assertEqual(sheet["B3"].border.left.style, "thin")


class CrossTypeTests(_Workspace):
    def test_docx_to_pptx_starts_a_slide_per_heading_and_continues_long_bodies(self):
        document = Document()
        document.add_heading("长章节", 1)
        for index in range(11):
            document.add_paragraph(f"要点 {index + 1}", style="List Bullet")
        document.add_heading("短章节", 2)
        document.add_paragraph("一句话。")
        document.save(self.work / "long.docx")
        ok, data, errors = cross_type.docx_to_pptx(self.work / "long.docx", self.work / "long.pptx")
        self.assertTrue(ok, errors)
        titles = [slide.shapes.title.text for slide in Presentation(self.work / "long.pptx").slides]
        self.assertEqual(titles, ["长章节", "长章节 (续)", "短章节"])
        self.assertEqual(data["continuation_slides"], 1)
        self.assertPasses("long.docx", "long.pptx")

    def test_pptx_to_docx_keeps_slides_in_order_with_tables_and_notes(self):
        ok, data, errors = cross_type.pptx_to_docx(self.work / "deck.pptx", self.work / "deck.docx")
        self.assertTrue(ok, errors)
        document = Document(self.work / "deck.docx")
        headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
        self.assertEqual(headings, ["2026 Q3 经营回顾", "核心指标", "下季度重点"])
        self.assertEqual(document.tables[0].cell(1, 2).text, "4,820")
        self.assertTrue(any(p.text.startswith("备注：重点解释毛利率") for p in document.paragraphs))
        self.assertPasses("deck.pptx", "deck.docx")

    def test_a_worksheet_range_becomes_one_table(self):
        ok, data, errors = cross_type.xlsx_range_to_table(
            self.work / "book.xlsx", self.work / "part.docx", cell_range="A1:C3"
        )
        self.assertTrue(ok, errors)
        table = Document(self.work / "part.docx").tables[0]
        self.assertEqual((len(table.rows), len(table.columns)), (3, 3))
        self.assertEqual(table.cell(1, 1).text, "007")
        bad, _data, errors = cross_type.xlsx_range_to_table(
            self.work / "book.xlsx", self.work / "x.docx", cell_range="A:"
        )
        self.assertFalse(bad)
        self.assertEqual(errors[0]["code"], "INVALID_RANGE")


class PdfTests(_Workspace):
    def test_the_text_layer_gives_headings_lists_tables_and_drops_running_headers(self):
        analysis = pdf_convert.analyse_pdf(self.work / "supplier.pdf")
        self.assertEqual(analysis.headers, ["供应商评估报告 · 内部资料"])
        self.assertEqual(analysis.footers, ["采购中心"])
        items = [item for page in analysis.pages for item in page.items]
        headings = [(i.level, i.text) for i in items if i.kind == "heading"]
        self.assertEqual(headings[0], (1, "供应商评估报告"))
        self.assertIn((2, "二、评分明细"), headings)
        self.assertEqual(sum(i.kind == "list" for i in items), 3)
        tables = [i for i in items if i.kind == "table"]
        self.assertEqual(tables[0].rows[1], ["华星电子", "92", "95", "88", "91.7"])
        self.assertTrue(tables[0].ruled)
        self.assertEqual(tables[1].rows[0], ["事项", "负责人", "期限"])
        self.assertEqual(tables[1].rows[1][:2], ["提交改进计划", "王磊"])
        self.assertFalse(tables[1].ruled)

    def test_pdf_to_docx_xlsx_and_pptx(self):
        for suffix in ("docx", "xlsx", "pptx"):
            with self.subTest(suffix=suffix):
                ok, data, errors = pdf_convert.pdf_to_office(self.work / "supplier.pdf", self.work / f"s.{suffix}")
                self.assertTrue(ok, errors)
                self.assertPasses("supplier.pdf", f"s.{suffix}")
        document = Document(self.work / "s.docx")
        self.assertEqual(document.sections[0].header.paragraphs[0].text, "供应商评估报告 · 内部资料")
        self.assertEqual(load_workbook(self.work / "s.xlsx").sheetnames, ["第1页 表1", "第2页 表2"])
        titles = [slide.shapes.title.text for slide in Presentation(self.work / "s.pptx").slides]
        self.assertEqual(titles, ["供应商评估报告", "三、整改要求"])  # one slide per page

    def test_a_scanned_pdf_needs_ocr(self):
        ok, data, errors = pdf_convert.pdf_to_office(self.work / "supplier_scanned.pdf", self.work / "scan.docx")
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "PDF_NEEDS_OCR")
        self.assertFalse((self.work / "scan.docx").exists())

    def test_ocr_text_boxes_go_through_the_same_line_and_table_reading(self):
        def engine(_image):
            # Boxes in 300 dpi pixels, below the band a running header would sit in: a heading, a paragraph and a
            # three-row, two-column table.
            def box(x0, y0, x1, y1, text):
                return [[[x0, y0], [x1, y0], [x1, y1], [x0, y1]], text, 0.99]

            return [
                box(300, 600, 900, 680, "季度报告"),
                box(300, 750, 1500, 800, "本季度收入稳步增长。"),
                box(300, 900, 500, 950, "地区"),
                box(900, 900, 1100, 950, "收入"),
                box(300, 980, 500, 1030, "华东"),
                box(900, 980, 1100, 1030, "120"),
                box(300, 1060, 500, 1110, "华南"),
                box(900, 1060, 1100, 1110, "95"),
            ], 0.1

        with mock.patch.object(pdf_convert, "_ocr_engine", return_value=engine):
            ok, data, errors = pdf_convert.pdf_to_office(
                self.work / "supplier_scanned.pdf", self.work / "scan.docx", ocr=True
            )
        self.assertTrue(ok, errors)
        self.assertEqual(data["ocr_pages"], 2)
        document = Document(self.work / "scan.docx")
        self.assertEqual(document.paragraphs[0].style.name, "Heading 1")
        self.assertEqual(document.tables[0].cell(2, 1).text, "95")

    def test_ocr_without_the_engine_names_the_extra(self):
        with mock.patch.object(pdf_convert, "_ocr_engine", return_value=None):
            ok, _data, errors = pdf_convert.pdf_to_office(
                self.work / "supplier_scanned.pdf", self.work / "scan.docx", ocr=True
            )
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "CONVERTER_UNAVAILABLE")
        self.assertIn("ocr", errors[0]["message"])


class WpsConvertTests(_Workspace):
    CAPABILITIES = {
        "components": {
            "writer": {"selected_prog_id": "KWPS.Application"},
            "spreadsheets": {"selected_prog_id": "KET.Application"},
            "presentation": {"selected_prog_id": "KWPP.Application"},
        }
    }

    def run_fake(self, source, target, writes=None, **kwargs):
        scripts = []

        def fake_run(script, timeout_seconds, **_kwargs):
            scripts.append(script)
            params = json.loads(script.split("@'\n", 1)[1].split("\n'@", 1)[0])
            written = params["pdf_path"] or params["output_path"]
            if writes is not None:
                shutil.copyfile(writes, written)
            elif Path(written).suffix in {".odt", ".ods", ".odp"}:
                with zipfile.ZipFile(written, "w") as package:  # a real ODF package
                    package.writestr("mimetype", "application/vnd.oasis.opendocument.text")
            else:
                Path(written).write_bytes(b"converted")
            return {"ok": True, "conversion_method": "SaveAs", "written_path": written}, None

        with (
            mock.patch.object(wps_convert, "probe_wps_capabilities", return_value=self.CAPABILITIES),
            mock.patch.object(wps_convert, "run_powershell_script", side_effect=fake_run),
        ):
            result = wps_convert.wps_convert(self.work / source, self.work / target, **kwargs)
        return result, scripts

    def test_save_as_uses_the_format_code_for_the_extension(self):
        cases = {
            ("brief.doc", "out.docx"): ("writer", 16),
            ("brief.docx", "out.doc"): ("writer", 0),
            ("brief.docx", "out.odt"): ("writer", 23),
            ("book.xlsx", "out.xls"): ("spreadsheets", 56),
            ("book.xls", "out.ods"): ("spreadsheets", 60),
            ("deck.pptx", "out.ppt"): ("presentation", 1),
            ("deck.ppt", "out.odp"): ("presentation", 35),
            ("brief.docx", "out.wps"): ("writer", None),
            ("book.xlsx", "out.et"): ("spreadsheets", None),
            ("deck.pptx", "out.dps"): ("presentation", None),
        }
        for (source, target), (component, code) in cases.items():
            with self.subTest(source=source, target=target):
                (ok, data, errors), scripts = self.run_fake(source, target)
                self.assertTrue(ok, errors)
                self.assertEqual((data["component"], data["format_code"]), (component, code))
                self.assertIn("SaveAs", scripts[0])
                self.assertIn(wps_convert.FAMILIES[component]["collection"], scripts[0])
                (self.work / target).unlink()

    def test_pdf_and_png_pages(self):
        (ok, data, errors), scripts = self.run_fake("brief.docx", "brief.pdf")
        self.assertTrue(ok, errors)
        self.assertIn("ExportAsFixedFormat($pdf, 17)", scripts[0])
        (ok, data, errors), _scripts = self.run_fake(
            "deck.pptx", "slide.png", writes=self.work / "supplier.pdf", dpi=50
        )
        self.assertTrue(ok, errors)
        self.assertEqual([Path(p).name for p in data["pages"]], ["slide-001.png", "slide-002.png"])
        self.assertTrue(all(Path(p).is_file() for p in data["pages"]))

    def test_a_busy_wps_is_retried(self):
        replies = [
            (
                None,
                {
                    "ok": False,
                    "errors": [{"code": "COM_OPERATION_FAILED", "message": "0x80010001 (RPC_E_CALL_REJECTED)"}],
                    "data": {},
                },
            ),
            ({"ok": True, "conversion_method": "SaveAs"}, None),
        ]

        def fake_run(script, timeout_seconds, **_kwargs):
            params = json.loads(script.split("@'\n", 1)[1].split("\n'@", 1)[0])
            reply = replies.pop(0)
            if reply[0]:
                Path(params["output_path"]).write_bytes(b"converted")
            return reply

        with (
            mock.patch.object(wps_convert, "probe_wps_capabilities", return_value=self.CAPABILITIES),
            mock.patch.object(wps_convert, "run_powershell_script", side_effect=fake_run) as run,
            mock.patch.object(wps_convert.time, "sleep") as sleep,
        ):
            ok, _data, errors = wps_convert.wps_convert(self.work / "brief.docx", self.work / "out.wps")
        self.assertTrue(ok, errors)
        self.assertEqual(run.call_count, 2)
        sleep.assert_called_once()

    def test_the_window_capture_type_compiles_without_system_drawing(self):
        # Windows check on f8ed7bd: a Graphics parameter in the C# type failed to compile on PowerShell 7 (CS1069).
        script = wps_convert._WINDOW_TITLES_SCRIPT
        csharp = script.split("Add-Type @'", 1)[1].split("'@", 1)[0]
        self.assertNotIn("System.Drawing", csharp)
        self.assertIn("$graphics.GetHdc()", script)

    def test_a_file_wps_still_holds_is_retried_and_a_timeout_says_where_wps_was(self):
        held = {"ok": False, "errors": [{"code": "DOCUMENT_HELD_OPEN", "message": "held"}], "data": {}}
        timeout = {"ok": False, "errors": [{"code": "COM_OPERATION_TIMEOUT", "message": "timed out."}], "data": {}}
        replies = [
            (None, held),
            (None, timeout),
            (
                {
                    "windows": [
                        {"title": "", "class": "Qt5QWindowIcon", "width": 420, "height": 180, "image": "C:/t/w-1.png"}
                    ]
                },
                None,
            ),
        ]

        def fake_run(script, timeout_seconds, **_kwargs):
            if "stage_path" in script:
                params = json.loads(script.split("@'\n", 1)[1].split("\n'@", 1)[0])
                Path(params["stage_path"]).write_text("opening the input", encoding="utf-8")
            return replies.pop(0)

        with (
            mock.patch.object(wps_convert, "probe_wps_capabilities", return_value=self.CAPABILITIES),
            mock.patch.object(wps_convert, "run_powershell_script", side_effect=fake_run) as run,
            mock.patch.object(wps_convert.time, "sleep"),
        ):
            ok, data, errors = wps_convert.wps_convert(self.work / "brief.doc", self.work / "out.docx")
        self.assertFalse(ok)
        self.assertEqual(run.call_count, 3)  # held, timed out, then the windows
        self.assertEqual(errors[0]["code"], "COM_OPERATION_TIMEOUT")
        self.assertIn("WPS was opening the input", errors[0]["message"])
        self.assertIn("Open WPS windows: untitled Qt5QWindowIcon 420x180 (C:/t/w-1.png)", errors[0]["message"])
        self.assertEqual(data["stage"], "opening the input")
        self.assertEqual(data["wps_windows"][0]["image"], "C:/t/w-1.png")

    def test_an_odf_source_wps_will_not_open_is_unsupported_and_its_wps_is_ended(self):
        # Windows check on 3095499: WPS 12.8 asked about every real ODF file as a damaged one.
        timeouts = []

        def fake_run(script, timeout_seconds, **_kwargs):
            timeouts.append(timeout_seconds)
            if "stage_path" in script:
                params = json.loads(script.split("@'\n", 1)[1].split("\n'@", 1)[0])
                Path(params["stage_path"]).write_text("opening the input", encoding="utf-8")
                Path(params["pid_path"]).write_text("4242", encoding="utf-8")
                return None, {"ok": False, "errors": [{"code": "COM_OPERATION_TIMEOUT", "message": "t"}], "data": {}}
            return {"windows": []}, None

        with (
            mock.patch.object(wps_convert, "probe_wps_capabilities", return_value=self.CAPABILITIES),
            mock.patch.object(wps_convert, "run_powershell_script", side_effect=fake_run),
            mock.patch.object(wps_convert.subprocess, "run") as kill,
            mock.patch.object(wps_convert.sys, "platform", "win32"),
        ):
            ok, data, errors = wps_convert.wps_convert(self.work / "brief.odt", self.work / "out.docx")
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "FORMAT_NOT_SUPPORTED_BY_WPS")
        self.assertIn("文件损坏提示", errors[0]["message"])
        self.assertEqual(timeouts[0], wps_convert.ODF_OPEN_TIMEOUT)  # a minute, not three
        self.assertEqual(data["ended_wps"], [4242])
        self.assertEqual(kill.call_args.args[0], ["taskkill", "/F", "/T", "/PID", "4242"])

    def test_an_odf_output_that_is_not_odf_is_removed(self):
        # Windows check on 3095499: WPS "saved as .odt" an OLE compound document.
        (self.work / "ole.bin").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64)
        (ok, _data, errors), _scripts = self.run_fake("brief.docx", "out.odt", writes=self.work / "ole.bin")
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "FORMAT_NOT_SUPPORTED_BY_WPS")
        self.assertFalse((self.work / "out.odt").exists())

    def test_writer_opens_without_the_conversion_prompt(self):
        (ok, _data, errors), scripts = self.run_fake("brief.odt", "out.docx")
        self.assertTrue(ok, errors)
        self.assertIn("Documents.Open($source, $false, $true, $false)", scripts[0])  # read-only

    def test_pairs_wps_cannot_write_are_refused_before_wps_starts(self):
        with mock.patch.object(wps_convert, "run_powershell_script") as run:
            ok, _data, errors = wps_convert.wps_convert(self.work / "brief.docx", self.work / "out.xlsx")
            self.assertFalse(ok)
            self.assertEqual(errors[0]["code"], "FORMAT_NOT_SUPPORTED_BY_WPS")
            ok, _data, errors = wps_convert.wps_convert(self.work / "book.xlsx", self.work / "out.ppt")
            self.assertEqual(errors[0]["code"], "FORMAT_NOT_SUPPORTED_BY_WPS")
            run.assert_not_called()

    def test_without_wps_the_backend_is_unavailable(self):
        missing = {"components": {name: {"selected_prog_id": None} for name in wps_convert.FAMILIES}}
        with mock.patch.object(wps_convert, "probe_wps_capabilities", return_value=missing):
            ok, _data, errors = wps_convert.wps_convert(self.work / "brief.docx", self.work / "out.doc")
        self.assertFalse(ok)
        self.assertEqual(errors[0]["code"], "COM_BACKEND_UNAVAILABLE")


class WpsLayoutTests(_Workspace):
    """What the Windows check on 4b263d9 found: WPS paginated the docx files longer than the browser."""

    def test_docx_pages_and_paragraph_defaults_follow_the_browser(self):
        convert_html_to_office(self.work / "contract.html", self.work / "contract.docx", "docx")
        document = Document(self.work / "contract.docx")
        section = document.sections[0]
        self.assertEqual((round(section.page_width.mm), round(section.page_height.mm)), (210, 297))  # A4
        defaults = document.styles.element.find(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}docDefaults"
        )
        w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        props = defaults.find(f"{w}pPrDefault/{w}pPr")
        for name in ("autoSpaceDE", "autoSpaceDN", "snapToGrid"):
            self.assertEqual(props.find(f"{w}{name}").get(f"{w}val"), "0")
        self.assertEqual(props.find(f"{w}spacing").get(f"{w}after"), "0")
        cell = document.tables[1].cell(2, 0).paragraphs[0]._p
        self.assertIsNone(cell.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr"))

    def test_paragraphs_and_tables_ignore_the_wps_grid_and_cell_margin(self):
        # Windows check on c117637: WPS snapped lines to the 18 pt grid despite the document default, and
        # started tables one cell margin left of the text.
        w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        convert_html_to_office(self.work / "contract.html", self.work / "contract.docx", "docx")
        document = Document(self.work / "contract.docx")
        self.assertIsNone(document.sections[0]._sectPr.find(f"{w}docGrid"))
        for paragraph in document.paragraphs:
            self.assertEqual(paragraph._p.pPr.find(f"{w}snapToGrid").get(f"{w}val"), "0")
        properties = document.tables[1]._tbl.tblPr
        indent = properties.find(f"{w}tblInd").get(f"{w}w")
        self.assertEqual(indent, properties.find(f"{w}tblCellMar/{w}left").get(f"{w}w"))
        self.assertEqual(indent, "120")  # the cells' 8 px padding

    def test_lines_are_exactly_the_css_height_unless_something_is_taller(self):
        # Windows check on 8899219: "at least" 16.5 pt became Microsoft YaHei's own ~1.7 em line in WPS.
        from docx.enum.text import WD_LINE_SPACING

        (self.work / "lines.html").write_text(
            "<html><head><style>body{font-size:16px;line-height:1.5} th{font-weight:400}</style></head><body>"
            "<p>plain</p><p>big <span style='font-size:40px'>word</span></p>"
            "<table><tr><th>head</th></tr></table></body></html>",
            encoding="utf-8",
        )
        convert_html_to_office(self.work / "lines.html", self.work / "lines.docx", "docx", use_browser=False)
        document = Document(self.work / "lines.docx")
        plain, big = document.paragraphs[0].paragraph_format, document.paragraphs[1].paragraph_format
        self.assertEqual(plain.line_spacing_rule, WD_LINE_SPACING.EXACTLY)
        self.assertAlmostEqual(plain.line_spacing.pt, 18, delta=0.1)  # 24 px
        self.assertEqual(big.line_spacing_rule, WD_LINE_SPACING.AT_LEAST)
        self.assertFalse(document.tables[0].cell(0, 0).paragraphs[0].runs[0].bold)  # th{font-weight:400}

    def test_pdf_bottom_margin_does_not_follow_where_the_text_ends(self):
        # Windows check on 8899219: a short last page set a 427 pt bottom margin and pushed text to a new page.
        top, _right, bottom, _left = pdf_convert.analyse_pdf(self.work / "supplier.pdf").margins
        self.assertLessEqual(bottom, top)

    def test_vertical_cell_padding_is_paragraph_spacing_not_a_cell_margin(self):
        # Windows check on c55d04e: rows with a top/bottom cell margin came out about 7 pt taller in WPS.
        w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        convert_html_to_office(self.work / "contract.html", self.work / "contract.docx", "docx", use_browser=False)
        cell = Document(self.work / "contract.docx").tables[1].cell(2, 0)
        margins = cell._tc.tcPr.find(f"{w}tcMar")
        self.assertIsNone(margins.find(f"{w}top"))
        self.assertIsNotNone(margins.find(f"{w}left"))
        form = cell.paragraphs[0].paragraph_format
        self.assertAlmostEqual(form.space_before.pt, 3, delta=0.1)  # 4 px
        self.assertAlmostEqual(form.space_after.pt, 3, delta=0.1)
        names = [child.tag.split("}")[1] for child in cell._tc.tcPr]
        if "shd" in names:
            self.assertLess(names.index("tcBorders"), names.index("shd"))

    def test_the_body_margin_is_part_of_the_page_margin(self):
        # Windows check on c205851: the browser's 8 px body margin was missing, so WPS lines were 16 px wider.
        convert_html_to_office(self.work / "contract.html", self.work / "contract.docx", "docx", use_browser=False)
        section = Document(self.work / "contract.docx").sections[0]
        self.assertAlmostEqual(section.left_margin.mm, 10 + 8 * 25.4 / 96, delta=0.1)
        self.assertAlmostEqual(section.right_margin.mm, 10 + 8 * 25.4 / 96, delta=0.1)
        self.assertAlmostEqual(section.top_margin.mm, 10, delta=0.1)
        # PDF to docx measures its margins from the PDF: no body margin is added to them.
        pdf_convert.pdf_to_office(self.work / "supplier.pdf", self.work / "s.docx")
        left = pdf_convert.analyse_pdf(self.work / "supplier.pdf").margins[3]
        self.assertAlmostEqual(Document(self.work / "s.docx").sections[0].left_margin.pt, round(left), delta=0.6)
        # docx -> HTML -> docx keeps the margin: the HTML's own 24 px body margin is left out of its @page margin.
        office_html.office_to_html(self.work / "contract.docx", self.work / "c.html")
        convert_html_to_office(self.work / "c.html", self.work / "back.docx", "docx", use_browser=False)
        back = Document(self.work / "back.docx").sections[0]
        self.assertAlmostEqual(back.left_margin.mm, section.left_margin.mm, delta=0.2)

    def test_a_table_margin_becomes_the_space_around_it(self):
        # Windows check on c205851: the signature table's margin-top: 24pt was lost.
        (self.work / "gap.html").write_text(
            "<html><head><style>table{margin:24pt 0 30pt}</style></head><body>"
            "<p>before</p><table><tr><td>cell</td></tr></table><p>after</p></body></html>",
            encoding="utf-8",
        )
        convert_html_to_office(self.work / "gap.html", self.work / "gap.docx", "docx", use_browser=False)
        before, after = Document(self.work / "gap.docx").paragraphs[:2]
        self.assertAlmostEqual(before.paragraph_format.space_after.pt, 24, delta=0.1)
        self.assertAlmostEqual(after.paragraph_format.space_before.pt, 30, delta=0.1)

    def test_code_boxes_and_bullets_match_the_browser(self):
        # Windows check on f8ed7bd: CRLF code lines got a blank line between them, padded boxes had no padding,
        # bullet text stopped at the template's 360-twip tab, and rows lost the 2 px border-spacing.
        w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        (self.work / "box.html").write_bytes(
            b"<html><head><style>pre{background:#f2f2f2;padding:6px} .warn{border:1px solid #e06666;padding:8px}"
            b"</style></head><body><pre>one\r\ntwo</pre><p class='warn'>note</p><ul><li>item</li></ul>"
            b"<table><tr style='height:20px'><td>cell</td></tr></table></body></html>"
        )
        convert_html_to_office(self.work / "box.html", self.work / "box.docx", "docx", use_browser=False)
        document = Document(self.work / "box.docx")
        code, note, item = document.paragraphs[:3]
        self.assertEqual(len(code._p.findall(f".//{w}br")), 1)
        self.assertEqual(code.text, "one\ntwo")
        code_border = code._p.pPr.find(f"{w}pBdr/{w}left")
        self.assertEqual(code_border.get(f"{w}color"), "F2F2F2")  # an invisible border carries the padding
        self.assertEqual(code_border.get(f"{w}space"), "4")  # 6 px = 4.5 pt, rounded
        self.assertAlmostEqual(code.paragraph_format.left_indent.pt, 4.5, delta=0.01)  # the 6 px padding
        self.assertEqual(note._p.pPr.find(f"{w}pBdr/{w}top").get(f"{w}space"), "6")
        self.assertAlmostEqual(note.paragraph_format.left_indent.pt, 6.75, delta=0.01)  # 8 px padding + 1 px border
        tabs = [(t.get(f"{w}val"), t.get(f"{w}pos")) for t in item._p.pPr.findall(f"{w}tabs/{w}tab")]
        self.assertIn(("clear", "360"), tabs)
        self.assertIn(("left", str(item.paragraph_format.left_indent.twips)), tabs)
        self.assertAlmostEqual(document.tables[0].rows[0].height.pt, 16.5, delta=0.1)  # 20 px + 2 px spacing

    def test_the_list_marker_hangs_as_far_as_the_browser_puts_it(self):
        # Windows check on 3095499: numbers sat 24 px before the text, the browser's about 0.95 em.
        items = "".join(f"<li>item {n}</li>" for n in range(12))
        (self.work / "marks.html").write_text(
            f"<html><body><ul><li>dot</li></ul><ol>{items}</ol></body></html>", encoding="utf-8"
        )
        convert_html_to_office(self.work / "marks.html", self.work / "marks.docx", "docx", use_browser=False)
        bullet, number = Document(self.work / "marks.docx").paragraphs[:2]
        self.assertAlmostEqual(bullet.paragraph_format.first_line_indent.pt, -16 * 0.95 * 0.75, delta=0.05)
        self.assertAlmostEqual(number.paragraph_format.first_line_indent.pt, -16 * 1.5 * 0.75, delta=0.05)  # "12."

    def test_list_text_starts_where_the_browser_puts_it(self):
        convert_html_to_office(self.work / "contract.html", self.work / "contract.docx", "docx")
        items = [p for p in Document(self.work / "contract.docx").paragraphs if p.style.name.startswith("List")]
        first, nested = items[0].paragraph_format, items[1].paragraph_format
        self.assertAlmostEqual(first.left_indent.pt, 30, delta=0.5)  # 40 px
        self.assertAlmostEqual(nested.left_indent.pt, 60, delta=0.5)
        self.assertLess(first.first_line_indent.pt, 0)  # the number hangs

    def test_pdf_lines_keep_their_measured_pitch(self):
        from docx.enum.text import WD_LINE_SPACING

        pdf_convert.pdf_to_office(self.work / "supplier.pdf", self.work / "s.docx")
        body = [p for p in Document(self.work / "s.docx").paragraphs if p.text.startswith("东方精密需在")][0]
        form = body.paragraph_format
        self.assertEqual(form.line_spacing_rule, WD_LINE_SPACING.EXACTLY)
        self.assertGreater(form.line_spacing.pt, 12)

    def test_pdf_tables_keep_their_row_heights_and_padding(self):
        from docx.shared import Pt

        analysis = pdf_convert.analyse_pdf(self.work / "supplier.pdf")
        table = [i for p in analysis.pages for i in p.items if i.kind == "table"][0]
        self.assertEqual(len(table.row_heights), 6)
        self.assertGreater(table.padding, 3)
        self.assertFalse(table.header_bold)
        pdf_convert.pdf_to_office(self.work / "supplier.pdf", self.work / "s.docx")
        rows = Document(self.work / "s.docx").tables[0].rows
        self.assertAlmostEqual(rows[1].height.pt, Pt(table.row_heights[1]).pt, delta=0.6)
        heading = [p for p in Document(self.work / "s.docx").paragraphs if p.text == "二、评分明细"][0]
        self.assertFalse(heading.runs[0].bold)  # regular in the PDF

    def test_layout_keeps_page_pairs_and_can_render_another_source(self):
        pdf_convert.pdf_to_office(self.work / "supplier.pdf", self.work / "s.docx")
        with mock.patch("wps_ai_agent_cli.fidelity._images_for") as images:
            from PIL import Image

            page = Image.new("L", (40, 60), 255)
            images.side_effect = lambda path, *a: ([page, page], None)
            ok, report, _errors = fidelity_check(
                self.work / "supplier.pdf",
                self.work / "s.docx",
                renderer="libreoffice",
                layout_source=self.work / "supplier_scanned.pdf",
                images_dir=self.work / "pages",
            )
        self.assertTrue(ok)
        self.assertEqual(images.call_args_list[0].args[0], self.work / "supplier_scanned.pdf")
        self.assertEqual(len(report["aspects"]["layout"]["images"]), 2)
        self.assertTrue((self.work / "pages" / "page-001-1.000.png").is_file())


class ConvertDispatchTests(_Workspace):
    def test_routes(self):
        self.assertEqual(convert.routes(".html", ".docx"), [("html-editable", ".docx")])
        self.assertEqual(convert.routes(".pdf", ".html"), [("pdf-to-office", ".docx"), ("office-to-html", ".html")])
        self.assertEqual(convert.routes(".doc", ".md"), [("wps", ".docx"), ("to-markdown", ".md")])
        self.assertEqual(convert.routes(".csv", ".pptx"), [("csv-to-xlsx", ".xlsx"), ("xlsx-range-to-table", ".pptx")])
        self.assertEqual(convert.routes(".docx", ".xlsx"), [("office-to-html", ".html"), ("html-editable", ".xlsx")])
        self.assertEqual(convert.routes(".docx", ".ods"), [])

    def test_two_step_route_through_wps(self):
        def fake_wps(source, target, options):
            shutil.copyfile(self.work / "brief.docx", target)
            return True, {"output_path": str(target), "warnings": ["from wps"]}, []

        with mock.patch.object(convert, "_wps", side_effect=fake_wps):
            ok, data, errors = convert.convert(self.work / "brief.doc", self.work / "brief.md")
        self.assertTrue(ok, errors)
        self.assertEqual(data["route"], ["wps", "to-markdown"])
        self.assertIn("from wps", data["warnings"])
        self.assertIn("# 技术服务合同", (self.work / "brief.md").read_text(encoding="utf-8"))

    def test_the_result_summarises_the_output(self):
        ok, data, errors = convert.convert(self.work / "brief.docx", self.work / "brief.pptx")
        self.assertTrue(ok, errors)
        self.assertEqual(data["output_summary"]["slide_titles"][0], "技术服务合同")
        ok, data, errors = convert.convert(self.work / "sales_gbk.csv", self.work / "s.xlsx")
        self.assertEqual(data["output_summary"]["sheets"][0]["dimensions"], "A1:F4")

    def test_unsupported_pairs_name_the_reason(self):
        ok, _data, errors = convert.convert(self.work / "brief.docx", self.work / "brief.ods")
        self.assertEqual(errors[0]["code"], "FORMAT_NOT_SUPPORTED_BY_WPS")
        ok, _data, errors = convert.convert(self.work / "chart.png", self.work / "chart.docx")
        self.assertEqual(errors[0]["code"], "UNSUPPORTED_CONVERSION")

    def test_cli_convert_with_check_and_fidelity_check(self):
        stream = io.StringIO()
        code = cli.run(
            ["convert", "--input", str(self.work / "api.md"), "--output", str(self.work / "api.html"), "--check"],
            output_stream=stream,
        )
        envelope = json.loads(stream.getvalue())
        self.assertEqual(code, 0, envelope)
        self.assertEqual(envelope["data"]["route"], ["markdown-to-html"])
        self.assertTrue(envelope["data"]["fidelity"]["passed"])
        stream = io.StringIO()
        cli.run(
            [
                "fidelity-check",
                "--source",
                str(self.work / "api.md"),
                "--target",
                str(self.work / "api.html"),
                "--no-layout",
            ],
            output_stream=stream,
        )
        envelope = json.loads(stream.getvalue())
        self.assertEqual(envelope["validation"]["status"], "passed")
        self.assertEqual(envelope["data"]["group"], "C")

    def test_mcp_tools_map_to_the_cli(self):
        ok, argv, _schema, errors = build_cli_argv(
            "wps_agent_convert", {"input": "a.xlsx", "output": "a.docx", "sheet": "预算", "range": "A1:C9"}
        )
        self.assertTrue(ok, errors)
        self.assertEqual(
            argv, ["convert", "--input", "a.xlsx", "--output", "a.docx", "--sheet", "预算", "--range", "A1:C9"]
        )
        ok, argv, schema, errors = build_cli_argv("wps_agent_fidelity_check", {"source": "a", "target": "b"})
        self.assertTrue(ok, errors)
        self.assertFalse(schema["mutates_document"])
        ok, argv, _schema, errors = build_cli_argv(
            "wps_agent_html_editable",
            {"input": "a.html", "output": "a.pptx", "format": "pptx", "fidelity": "appearance"},
        )
        self.assertIn("--fidelity", argv)


class FidelityScoringTests(_Workspace):
    def test_data_compares_values_not_spellings(self):
        from wps_ai_agent_cli.fidelity import Block, Model

        source = Model("html", [Block("table", rows=[["1.2e4", "3.50", "2025/11/20", "x"]])])
        target = Model("xlsx", [Block("table", rows=[["12000", "3.5", "2025-11-20", "x"]])])
        self.assertEqual(score_data(source, target), (1.0, []))

    def test_a_lost_table_fails_the_check(self):
        convert_html_to_office(self.work / "report.html", self.work / "report.docx", "docx")
        document = Document(self.work / "report.docx")
        table = document.tables[0]._tbl
        table.getparent().remove(table)
        document.save(self.work / "report.docx")
        ok, report, _errors = fidelity_check(self.work / "report.html", self.work / "report.docx", layout=False)
        self.assertTrue(ok)
        self.assertFalse(report["passed"])
        self.assertFalse(report["verdict"]["data"])

    def test_pdf_model_leaves_out_running_headers(self):
        model = read_model(self.work / "supplier.pdf")
        texts = [block.text for block in model.blocks]
        self.assertNotIn("采购中心", texts)
        self.assertIn("供应商评估报告", texts)

    @unittest.skipUnless(shutil.which("soffice") or shutil.which("libreoffice"), "LibreOffice is not installed")
    def test_layout_is_measured_with_libreoffice_and_checked_against_a_baseline(self):
        pdf_convert.pdf_to_office(self.work / "supplier.pdf", self.work / "s.docx")
        ok, report, errors = fidelity_check(
            self.work / "supplier.pdf", self.work / "s.docx", renderer="libreoffice", baseline=0.99
        )
        self.assertTrue(ok, errors)
        layout = report["aspects"]["layout"]
        self.assertEqual((layout["status"], layout["renderer"]), ("measured", "libreoffice"))
        self.assertGreater(layout["score"], 0.8)
        self.assertFalse(report["verdict"]["layout"])  # far below the made-up 0.99 baseline


if __name__ == "__main__":
    unittest.main()
