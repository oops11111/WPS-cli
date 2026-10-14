import ast
import io
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from wps_ai_agent_cli import cli, cli_handlers, spreadsheet_ops


def run_cli(argv):
    stream = io.StringIO()
    code = cli.run(argv, stream)
    return code, json.loads(stream.getvalue())


class TaskStatusWiringTests(unittest.TestCase):
    def test_every_task_tracked_handler_passes_the_request_and_command_to_the_tracker(self):
        tree = ast.parse(Path(cli_handlers.__file__).read_text(encoding="utf-8"))
        calls = [
            node
            for node in ast.walk(tree)
            if isinstance(node, ast.Call) and getattr(node.func, "attr", "") == "_with_optional_task_status"
        ]
        self.assertGreater(len(calls), 20)
        for call in calls:
            keywords = {keyword.arg for keyword in call.keywords}
            with self.subTest(line=call.lineno):
                self.assertLessEqual({"task_id", "tracked_command", "operation_request_id"}, keywords)


class SheetCommandCliTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.dir = Path(tmp.name)
        previous = os.getcwd()
        os.chdir(self.dir)
        self.addCleanup(os.chdir, previous)
        self.document = self.dir / "book.xlsx"
        workbook = Workbook()
        workbook.active.title = "First"
        workbook.create_sheet("Second")
        workbook.save(self.document)
        workbook.close()
        code, payload = run_cli(["register-document", "--component", "spreadsheets", "--path", str(self.document)])
        self.assertEqual(code, 0, payload)
        self.document_id = payload["data"]["document"]["document_id"]

    def commands(self):
        doc = ["--document-id", self.document_id]
        return {
            "spreadsheet-rename-sheet": [*doc, "--old-name", "First", "--new-name", "Renamed"],
            "spreadsheet-create-sheet": [*doc, "--name", "Created", "--index", "2"],
            "spreadsheet-set-sheet-visibility": [*doc, "--sheet-name", "Second", "--visible", "false"],
            "spreadsheet-delete-sheet": [*doc, "--sheet-name", "Second"],
            "spreadsheet-copy-sheet": [*doc, "--source-name", "First", "--new-name", "Copy"],
            "spreadsheet-set-sheet-tab-color": [*doc, "--sheet-name", "First", "--color", "#FF0000"],
        }

    def test_every_sheet_command_runs_through_the_cli_without_a_task_id(self):
        for command, args in self.commands().items():
            with self.subTest(command=command):
                code, payload = run_cli([command, *args, "--dry-run"])
                self.assertEqual(code, 0, payload)
                self.assertTrue(payload["ok"], payload)
                self.assertEqual(payload["command"], command)

    def test_every_sheet_command_runs_with_an_explicit_request_id_in_either_position(self):
        for command, args in self.commands().items():
            with self.subTest(command=command, position="after"):
                code, payload = run_cli([command, *args, "--dry-run", "--request-id", f"after-{command}"])
                self.assertTrue(payload["ok"], payload)
                self.assertEqual(payload["request_id"], f"after-{command}")
            with self.subTest(command=command, position="before"):
                code, payload = run_cli(["--request-id", f"before-{command}", command, *args, "--dry-run"])
                self.assertTrue(payload["ok"], payload)
                self.assertEqual(payload["request_id"], f"before-{command}")

    def test_every_sheet_command_runs_with_a_task_id_and_records_the_request_on_the_task(self):
        for index, (command, args) in enumerate(self.commands().items()):
            task_id = f"sheet_task_{index}"
            request_id = f"sheet-req-{index}"
            with self.subTest(command=command):
                code, payload = run_cli([command, *args, "--dry-run", "--task-id", task_id, "--request-id", request_id])
                self.assertEqual(code, 0, payload)
                self.assertTrue(payload["ok"], payload)
                self.assertEqual(payload["command"], command)
                code, status = run_cli(["task-status", "--task-id", task_id])
                self.assertTrue(status["ok"], status)
                task = status["data"]["task_status"]
                self.assertEqual(task["command"], command)
                self.assertEqual(task["request_id"], request_id)
                self.assertEqual(task["document_id"], self.document_id)
                self.assertEqual(task["state"], "succeeded")

    def test_a_failed_sheet_command_marks_the_task_failed_instead_of_crashing(self):
        code, payload = run_cli(
            [
                "spreadsheet-rename-sheet",
                "--document-id",
                self.document_id,
                "--old-name",
                "Missing",
                "--new-name",
                "Whatever",
                "--task-id",
                "sheet_fail",
                "--request-id",
                "sheet-fail-req",
            ]
        )
        self.assertFalse(payload["ok"], payload)
        self.assertEqual(payload["errors"][0]["code"], "SHEET_NOT_FOUND")
        _code, status = run_cli(["task-status", "--task-id", "sheet_fail"])
        self.assertEqual(status["data"]["task_status"]["state"], "failed")
        self.assertEqual(status["data"]["task_status"]["request_id"], "sheet-fail-req")

    def test_rename_runs_end_to_end_through_the_cli_with_a_task_id(self):
        def fake_rename(path, old_name, new_name):
            workbook = load_workbook(path)
            workbook[old_name].title = new_name
            workbook.save(path)
            workbook.close()
            return {"ok": True, "errors": [], "data": {"backend": "fake-com", "sheet_names": [new_name, "Second"]}}

        with patch.object(spreadsheet_ops, "_run_spreadsheet_rename_sheet_com", side_effect=fake_rename):
            code, payload = run_cli(
                [
                    "spreadsheet-rename-sheet",
                    "--document-id",
                    self.document_id,
                    "--old-name",
                    "First",
                    "--new-name",
                    "Renamed",
                    "--task-id",
                    "sheet_real",
                    "--request-id",
                    "sheet-real-req",
                ]
            )
        self.assertEqual(code, 0, payload)
        self.assertTrue(payload["ok"], payload)
        self.assertTrue(payload["data"]["validation_passed"])
        _code, status = run_cli(["task-status", "--task-id", "sheet_real"])
        self.assertEqual(status["data"]["task_status"]["state"], "succeeded")
        self.assertEqual(load_workbook(self.document).sheetnames, ["Renamed", "Second"])


if __name__ == "__main__":
    unittest.main()
